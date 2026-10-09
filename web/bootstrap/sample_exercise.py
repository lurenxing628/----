"""Exercise the isolated sample using the ordinary HTTP business contracts.

No SQL, fabricated schedule, altered worker, or alternate adoption path.
Failures remain in the sample report at the stage where they happened.
"""

import csv
import io
import json
import time
import uuid
from datetime import datetime, timedelta
from pathlib import Path
from urllib.error import HTTPError
from urllib.parse import urlencode
from urllib.request import urlopen

from .sample_constraint_checks import check_candidate, require
from .sample_formal_read import read_formal
from .sample_http import BASE, SampleAPIError, SampleClient, SampleProgress

_TERMINAL = ("complete", "partial", "failed", "cancelled", "interrupted")


def _key():
    return "complex-sample-" + uuid.uuid4().hex


def _query(path, values):
    return path + "?" + urlencode(values)


def _run(client, seed, progress, label, batch_codes=None):
    codes = list(seed["refs"]["batch"]) if batch_codes is None else batch_codes
    value = dict(seed["schedule_window"], batch_refs=sorted(seed["refs"]["batch"][code] for code in codes), hold_window=None)
    preflight = client.post(BASE + "/scheduling/preflight", value)["data"]
    preview = client.post(BASE + "/scheduling/runs/preview", {"input_ref": preflight["input_ref"]})["data"]
    progress.report["current_preflight"] = {"label": label, "counts": preflight["counts"],
        "blockers": preflight["blockers"], "warnings": preflight["warnings"], "preview_context": preview["write_context"]}
    progress.save()
    require(preview["write_context"]["capabilities"]["scheduling.run"] is True,
            "真实预检不允许开始这次排产：{}".format(preview["write_context"]["blocked_reasons"]))
    accepted = client.post(BASE + "/scheduling/runs", {"input_ref": preflight["input_ref"],
                           "write_token": preview["write_context"]["write_token"], "request_key": _key()}, status=202)
    run_ref, started = accepted["run_ref"], time.monotonic()
    deadline = started + 1800
    previous = None
    while True:
        run = client.get(BASE + "/scheduling/runs/" + run_ref)
        marker = (run["state"], run.get("progress"))
        if marker != previous:
            progress.report["current_run"] = {"label": label, "run_ref": run_ref, "state": run["state"],
                                              "progress": run.get("progress"), "elapsed_seconds": round(time.monotonic() - started, 3)}
            progress.save()
            previous = marker
        if run["state"] in _TERMINAL:
            break
        require(time.monotonic() < deadline, "真实排产30分钟后仍未结束，保留该run及进度，不重复受理。")
        time.sleep(0.5)
    require(run["state"] in ("complete", "partial"), f"真实排产没有可读取的完成候选：{json.dumps(run, ensure_ascii=False)}")
    catalog = client.get(BASE + "/scheduling/runs/" + run_ref + "/candidates?size=50")
    complete = [row for row in catalog["candidates"] if row["status"] == "completed" and row["completeness"] == "complete"]
    require(complete, f"真实排产没有完整可行候选：{json.dumps(catalog, ensure_ascii=False)}")
    candidate_ref = complete[0]["candidate_ref"]
    workspace = client.document(BASE + "/scheduling/candidates/" + candidate_ref + "/workspace")
    return {"run": run, "preflight": preflight, "catalog": catalog, "candidate_ref": candidate_ref,
            "workspace": workspace, "selected_batch_codes": codes, "elapsed_seconds": round(time.monotonic() - started, 3)}


def _summary_run(artifacts):
    candidates = artifacts["catalog"]["candidates"]
    skipped = [{"candidate_ref": row["candidate_ref"], "label": row["label"], "status": row["status"],
                "task_count": row["task_count"], "published_reasons": row["blocked_reasons"]}
               for row in candidates if row["status"] == "skipped"]
    return {"run_ref": artifacts["run"]["run_ref"], "state": artifacts["run"]["state"],
            "elapsed_seconds": artifacts["elapsed_seconds"], "preflight_task_count": len(artifacts["preflight"]["tasks"]),
            "candidate_count": artifacts["catalog"]["candidate_count"], "candidates": artifacts["catalog"]["candidates"],
            "candidate_ref": artifacts["candidate_ref"], "task_count": artifacts["workspace"]["data"]["task_count"],
            "selected_batch_count": len(artifacts["selected_batch_codes"]), "skipped_candidates": skipped,
            "optimizer_coverage": "存在未执行候选；只验收实际完整候选，未宣称优化方案全部通过。" if skipped else "本轮候选状态见真实目录记录。"}


def _export_cells(document):
    if "candidate" in document["data"]:
        from core.services.workbench.run.candidate_export import HEADERS, export_rows
    else:
        from core.services.workbench.plan.export import HEADERS, export_rows
    return HEADERS, list(export_rows(document["data"], document["meta"]))


def _decoded_cell(value, fmt):
    if not isinstance(value, str):
        return value
    if fmt == "csv" and value.startswith("'"):
        value = value[1:]
    if value == r"\N":
        return None
    return value[1:] if value.startswith("\\\\") else value


def _same_cell(actual, expected, fmt):
    if type(expected) in (int, float):
        # XLSX numeric cells use the writer's 16 significant digit format.
        # This is its actual file precision, not an arbitrary comparison slack.
        number = float(format(expected, ".16g")) if fmt == "xlsx" else float(expected)
        return actual is not None and float(actual) == number
    return actual == expected


def _read_export_rows(content, fmt):
    if fmt == "csv":
        return list(csv.reader(io.StringIO(content.decode("utf-8-sig"))))
    from openpyxl import load_workbook

    require(content.startswith(b"PK"), "实际XLSX不是有效文件。")
    book = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    try:
        sheet = book.active
        if sheet is None:
            raise RuntimeError("实际XLSX缺少工作表。")
        return [list(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        book.close()


def _download(client, path, fmt, document, folder, name):
    client.request_count += 1
    try:
        response = urlopen(client.base_url + path, timeout=client.timeout)
    except HTTPError as exc:
        body = json.loads(exc.read().decode("utf-8"))
        raise SampleAPIError(path, exc.code, body) from exc
    with response:
        content, headers = response.read(), dict(response.headers)
        require(response.getcode() == 200, "实际文件下载没有成功。")
    if folder is not None:
        folder.mkdir(parents=True, exist_ok=True)
        (folder / (name + "." + fmt)).write_bytes(content)
    rows = _read_export_rows(content, fmt)
    expected = document["data"]["tasks"]
    expected_header, expected_rows = _export_cells(document)
    require(tuple(rows[0]) == expected_header and len(rows) - 1 == len(expected_rows), "导出列或行数与真实候选/正式计划范围不一致。")
    decoded = [[_decoded_cell(value, fmt) for value in row] for row in rows[1:]]
    require(all(len(row) == len(wanted) and all(_same_cell(value, target, fmt) for value, target in zip(row, wanted))
                for row, wanted in zip(decoded, expected_rows)), "导出列内容与本段真实快照的资料或派生结果不一致。")
    header = rows[0]
    columns = {key: header.index(key) for key in ("工序编号", "安排开始", "安排结束")}
    exported = {row[columns["工序编号"]]: (row[columns["安排开始"]], row[columns["安排结束"]]) for row in decoded}
    original = {row["operation_ref"]: (row["start"], row["end"]) for row in expected}
    require(exported == original, "导出工序身份/开完工与当前真实完整范围不一致。")
    return {"bytes": len(content), "rows": len(rows) - 1, "format": fmt,
            "content_type": headers.get("Content-Type"), "content_disposition": headers.get("Content-Disposition"),
            "operation_times_identical": True, "all_columns_match_real_scope": True,
            "file": str(folder / (name + "." + fmt)) if folder else None}


def _task_core(row):
    fields = ("task_ref", "operation_ref", "plan_ref", "batch_id", "sequence", "process_label", "machine_ref", "operator_ref",
              "supplier_ref", "quantity", "batch_quantity", "piece_id", "quantity_basis", "quantity_reason", "start", "end")
    return {key: row[key] for key in fields if key in row}


def _export_document(client, document, base, folder, name):
    source_scope = document["data"].get("scope", {})
    scope = {key: source_scope[key] for key in ("range_start", "range_end") if source_scope.get(key) is not None}
    metadata = {"scope": scope, "snapshot_ref": document["meta"]["snapshot_ref"], "as_of": document["meta"]["as_of"]}
    for fmt in ("csv", "xlsx"):
        args = dict(scope, snapshot_ref=document["meta"]["snapshot_ref"], format=fmt)
        metadata[fmt] = _download(client, _query(base + "/export", args), fmt, document, folder, name)
    return metadata


def _collect_export_tasks(document, wanted, seen):
    for row in document["data"]["tasks"]:
        core = _task_core(row)
        require(core == wanted.get(row["task_ref"]), "本段导出的真实任务与完整正式身份、数量、资源或原始时间不一致。")
        previous = seen.setdefault(row["task_ref"], core)
        require(previous == core, "跨段重叠的同一任务业务资料发生变化，不能去重。")
    return len(document["data"]["tasks"])


def _segmented_export_result(segments, seen, physical_rows):
    result = {"mode": "segmented_http", "segments": segments, "segment_count": len(segments),
              "unique_task_count": len(seen), "physical_rows": physical_rows, "duplicate_overlap_rows": physical_rows - len(seen),
              "scope_rule": "各段派生结果与原文件按各自scope保留；跨段仅去重相同任务业务核心，不合成全局风险或全量快照。"}
    for fmt in ("csv", "xlsx"):
        result[fmt] = {"mode": "segmented_http", "rows": len(seen), "physical_rows": sum(row[fmt]["rows"] for row in segments),
                       "bytes": sum(row[fmt]["bytes"] for row in segments), "segment_count": len(segments),
                       "operation_times_identical": True, "all_columns_match_real_scope": True,
                       "files": [row[fmt]["file"] for row in segments], "single_full_scope_file": False}
    return result


def _exports(client, document, base, folder, name):
    assembly = document.get("assembly")
    if not assembly:
        segment = _export_document(client, document, base, folder, name)
        return {fmt: segment[fmt] for fmt in ("csv", "xlsx")}
    wanted = {row["task_ref"]: _task_core(row) for row in document["data"]["tasks"]}
    segments, seen, physical_rows = [], {}, 0
    for index, segment in enumerate(assembly["segments"], 1):
        filename = name + f"-segment-{index:03d}"
        segments.append(_export_document(client, segment, base, folder, filename))
        physical_rows += _collect_export_tasks(segment, wanted, seen)
    require(seen == wanted, "真实分段导出没有覆盖完整正式计划。")
    return _segmented_export_result(segments, seen, physical_rows)


def _formal(client, plan_ref, expected):
    document = read_formal(client, plan_ref, expected)
    data = document["data"]
    require(data["plan"]["plan_ref"] == plan_ref and data["plan"]["is_current_official"] is True
            and data["tasks_complete"] is True, "采用后的正式计划身份或完整性不成立。")
    fields = ("sequence", "process_label", "piece_id", "quantity", "batch_quantity")
    actual = {row["operation_ref"]: (row["start"], row["end"], row["machine_ref"], row["operator_ref"],
                                    row["supplier_ref"], row["batch_id"], *(row[key] for key in fields)) for row in data["tasks"]}
    wanted = {row["operation_ref"]: (row["start"], row["end"], row["machine"]["ref"] if row["machine"] else None,
                                    row["operator"]["ref"] if row["operator"] else None,
                                    row["supplier"]["ref"] if row["supplier"] else None, row["batch_label"],
                                    *(row[key] for key in fields)) for row in expected}
    require(len(actual) == data["task_count"] == len(data["tasks"]) == len(wanted) == len(expected)
            and actual == wanted, "正式采用/回读改动、重复或遗漏了候选的真实任务安排。")
    return document


def _conflict_trial(client, plan_ref):
    value = {"base": {"plan_ref": plan_ref}}
    preview = client.post(BASE + "/trial/drafts/preview", value)["data"]
    draft = client.command(BASE + "/trial/drafts", preview["write_context"], value)["data"]
    internal = [row for row in draft["tasks"] if row["source"] == "internal" and row["edit_context"]["can_change"]]
    first = internal[0]
    later = next(row for row in internal if row["batch_id"] == first["batch_id"] and row["sequence"] > first["sequence"])
    path = BASE + "/trial/drafts/" + draft["draft_ref"]
    changed = client.command(path + "/change", draft["write_context"], {"task_ref": later["task_ref"],
                             "machine_ref": later["machine_ref"], "operator_ref": later["operator_ref"], "start": first["start"]})["data"]
    require(changed["validation"]["constraints_status"] == "blocked", "故意把后道提前的试调没有显示约束冲突。")
    require(any(row["code"] == "precedence_violation" for row in changed["validation"]["issues"]), "试调没有保留前后道冲突证据。")
    saved = client.command(path + "/save", changed["write_context"], {"name": "复杂样例 · 故意前后道冲突（未采用）"})["data"]
    reopened = client.get(BASE + "/trial/scenarios/" + saved["scenario_ref"])
    adopted = client.post(BASE + "/trial/scenarios/" + saved["scenario_ref"] + "/adopt-preview", {})["data"]
    require(adopted["validation"]["can_adopt"] is False, "故意冲突的试调仍允许正式采用。")
    return {"draft_ref": draft["draft_ref"], "scenario_ref": saved["scenario_ref"], "state": "saved_not_adopted",
            "changed_operation_ref": later["operation_ref"], "validation": reopened["validation"],
            "adoption_validation": adopted["validation"]}


def _partial_report(client, task, quantity):
    path = BASE + "/execution/tasks/" + task["task_ref"]
    document = client.document(path)
    current = document["data"]["task"]
    start = datetime.fromisoformat(current["planned_start"])
    if start.microsecond:
        start = start.replace(microsecond=0) + timedelta(seconds=1)
    planned_end = datetime.fromisoformat(current["planned_end"])
    elapsed = (planned_end - start).total_seconds()
    require(elapsed >= 2 and 0 < quantity < current["quantity"], "所选工序无法形成有真实时间和剩余数量的分次报工。")
    duration = max(1, int(elapsed / 2))
    end = (start + timedelta(seconds=duration)).replace(microsecond=0)
    server_now = datetime.fromisoformat(document["meta"]["as_of"])
    require(planned_end <= server_now, "实际报工只选择已采用计划里真正过去的任务，不把未来计划写成过去执行。")
    value = {"actual_start": start.isoformat(timespec="seconds"), "actual_end": end.isoformat(timespec="seconds"),
             "completed_quantity": quantity, "effective_processing_hours": (end - start).total_seconds() / 3600,
             "actual_machine_ref": current["planned_machine_ref"], "actual_operator_ref": current["planned_operator_ref"],
             "remark": "[复杂样例] 按已采用计划的过去时段保存分次报工，保留剩余数量并核对重排保护"}
    receipt = client.command(path + "/reports", current["execution"]["write_context"], value)
    fresh = client.get(path)["task"]["execution"]
    require(fresh["known_completed_quantity"] == quantity and fresh["remaining_quantity"] == current["quantity"] - quantity,
            "分次报工保存后的数量/剩余数量不一致。")
    require(fresh["execution_state"] == "partial", "分次报工被误读为完全完成或未开工。")
    return {"task_ref": current["task_ref"], "operation_ref": current["operation_ref"], "receipt_ref": receipt["receipt_ref"],
            "known_completed_quantity": quantity, "remaining_quantity": fresh["remaining_quantity"], "input": value,
            "execution_state": fresh["execution_state"], "report_count": len(fresh["reports"]),
            "original_report_ref": fresh["reports"][0]["report_ref"], "batch_id": current["batch_id"],
            "time_basis": "sample_planned_interval_in_past",
            "server_as_of": document["meta"]["as_of"], "planned_start": current["planned_start"], "planned_end": current["planned_end"]}


def _directory(client):
    runs = client.get(BASE + "/scheduling/runs?size=50")
    candidates = []
    for run in runs["runs"]:
        catalog = client.get(BASE + "/scheduling/runs/" + run["run_ref"] + "/candidates?size=50")
        candidates.append({"run_ref": run["run_ref"], "count": catalog["candidate_count"],
                           "candidate_refs": [row["candidate_ref"] for row in catalog["candidates"]]})
    plans = client.get(BASE + "/plans?size=50")
    return {"run_count": runs["run_count"], "runs": [(row["run_ref"], row["state"], row["stage"], row["candidate_count"]) for row in runs["runs"]],
            "candidates": candidates, "plans": [(row["plan_ref"], row["version"], row["kind"], row["is_current_official"]) for row in plans["plans"]]}


def _partial_replan_blocked(client, seed, reports, batch_codes, expected_code):
    before = _directory(client)
    value = dict(seed["schedule_window"], batch_refs=sorted(seed["refs"]["batch"][code] for code in batch_codes), hold_window=None)
    preflight = client.post(BASE + "/scheduling/preflight", value)["data"]
    preview = client.post(BASE + "/scheduling/runs/preview", {"input_ref": preflight["input_ref"]})["data"]
    context = preview["write_context"]
    require(context["capabilities"]["scheduling.run"] is False
            and any(row["code"] == expected_code for row in context["blocked_reasons"]),
            "已有未完成实际工序的重排，没有保留产品明确的复核保护。")
    key = _key()
    try:
        client.post(BASE + "/scheduling/runs", {"input_ref": preflight["input_ref"], "write_token": "blocked-sample-probe", "request_key": key}, status=202)
    except SampleAPIError as rejected:
        require(rejected.status == 409 and rejected.body["error"]["code"] == "constraint_conflict",
                "已知不可排的范围没有在受理入口准确拒绝。")
        rejection = rejected.body
    else:
        raise RuntimeError("已知不可排的范围仍创建了新排产任务。")
    lookup = client.get(BASE + "/scheduling/requests/" + key)
    require(lookup["found"] is False and lookup["run"] is None, "被拒绝的排产创建了新job或受理receipt。")
    require(before == _directory(client), "被拒绝的排产改变了已有run/candidate/plan目录。")
    for report in reports:
        fresh = client.get(BASE + "/execution/tasks/" + report["task_ref"])["task"]["execution"]
        original = next(row for row in fresh["reports"] if row["report_ref"] == report["original_report_ref"])
        require(fresh["known_completed_quantity"] == report["known_completed_quantity"]
                and fresh["remaining_quantity"] == report["remaining_quantity"] and fresh["execution_state"] == "partial"
                and len(fresh["reports"]) == report["report_count"] and all(original[key] == value for key, value in report["input"].items()),
                "被拒的重排改动了原分次报工。")
    return {"expected_rejection": True, "reason": "未完成实际工序缺少可信剩余安排，当前产品不能在此状态下再次排产。",
            "selected_batch_count": len(batch_codes), "selected_operation_count": sum(len(seed["operation_refs"][code]) for code in batch_codes),
            "blockers": context["blocked_reasons"], "partial_actuals_unchanged": True,
            "admission_error": rejection["error"], "new_job_or_receipt": False, "existing_directories_unchanged": True}


def exercise(client, seed_state, report=None):
    """Use an injected client; report may be the injector's SampleProgress."""
    progress = report if report is not None else SampleProgress(None, client)
    progress.report = seed_state
    progress.report.update(state="exercising", exercise_state="running")
    progress.save()
    artifacts = {}
    folder = progress.output.parent / "sample-downloads" if progress.output is not None else None

    def baseline_run():
        artifacts["baseline"] = _run(client, seed_state, progress, "baseline")
        return _summary_run(artifacts["baseline"])

    def verify():
        baseline = artifacts["baseline"]
        checks = check_candidate(baseline["workspace"]["data"], seed_state, baseline["preflight"])
        preview = client.post(BASE + "/scheduling/candidates/" + baseline["candidate_ref"] + "/adopt-preview", {})["data"]
        require(preview["validation"]["can_adopt"] is True and preview["scope_complete"] is True,
                "完整候选没有通过真实采用前的日历/停机/依赖/报工约束复核：{}".format(preview["validation"]))
        require(preview["task_count"] == seed_state["counts"]["operations"], "采用复核没有覆盖全部样例工序。")
        artifacts["adoption_preview"] = preview
        return {"independent_checks": checks, "full_business_validation": preview["validation"], "task_count": preview["task_count"]}

    def adopt():
        baseline = artifacts["baseline"]
        receipt = client.command(BASE + "/scheduling/candidates/" + baseline["candidate_ref"] + "/adopt",
                                 artifacts["adoption_preview"]["write_context"],
                                 {"confirm": True, "reason": "独立复杂样例真实排产验收", "declared_operator": "复杂样例验收"})
        artifacts["plan_ref"] = receipt["data"]["official_plan"]["plan_ref"]
        artifacts["formal"] = _formal(client, artifacts["plan_ref"], baseline["workspace"]["data"]["tasks"])
        return {"result": receipt["result"], "receipt_ref": receipt["receipt_ref"],
                "official_plan": receipt["data"]["official_plan"], "task_count": artifacts["formal"]["data"]["task_count"]}

    def projections():
        dashboard = client.get(BASE + "/dashboard")
        analysis_document = client.document(_query(BASE + "/dashboard/analysis", {"plan_ref": artifacts["plan_ref"]}))
        analysis = analysis_document["data"]
        actual = client.get(_query(BASE + "/actual-gantt", {"plan_ref": artifacts["plan_ref"]}))
        require(dashboard["plan"]["plan_ref"] == artifacts["plan_ref"], "值班台没有读取刚采用的正式计划。")
        require(dashboard["analysis_error"] is None and dashboard["analysis"]["plan"]["plan_ref"] == artifacts["plan_ref"],
                "复杂样例的值班台分析没有正常读取：{}".format(dashboard["analysis_error"]))
        require(analysis["plan"]["plan_ref"] == artifacts["plan_ref"]
                and len(analysis["tasks"]) == seed_state["counts"]["operations"], "全局交付与产能分析没有覆盖当前真实完整正式计划。")
        require(actual["plan"]["plan_ref"] == artifacts["plan_ref"] and actual["items_complete"] is True
                and actual["task_count"] == seed_state["counts"]["operations"], "实际甘特没有读取当前正式计划的完整工序范围。")
        return {"plan_ref": artifacts["plan_ref"], "dashboard_items": len(dashboard["items"]), "dashboard_categories": dashboard["categories"],
                "analysis_task_count": len(analysis["tasks"]), "actual_gantt_summary": {key: actual[key]
                    for key in ("task_count", "items_complete", "report_count", "availability", "scope")},
                "delivery": analysis["deliveries"], "pressure": analysis["pressure"], "resource_pressure": analysis["resources"],
                "global_metrics_source": "/api/workbench/v1/dashboard/analysis", "global_analysis_snapshot": analysis_document["meta"]["snapshot_ref"],
                "resource_count": len(artifacts["formal"]["data"]["resources"])}

    def exports():
        baseline = artifacts["baseline"]
        return {"candidate": _exports(client, baseline["workspace"], BASE + "/scheduling/candidates/" + baseline["candidate_ref"], folder, "baseline-candidate"),
                "official": _exports(client, artifacts["formal"], BASE + "/plans/" + artifacts["plan_ref"], folder, "baseline-official")}

    def second_run():
        again = _run(client, seed_state, progress, "existing-formal-plan-before-partial-reports")
        checks = check_candidate(again["workspace"]["data"], seed_state, again["preflight"])
        preview = client.post(BASE + "/scheduling/candidates/" + again["candidate_ref"] + "/adopt-preview", {})["data"]
        require(preview["validation"]["can_adopt"] is True and preview["scope_complete"] is True
                and preview["task_count"] == seed_state["counts"]["operations"],
                "已有正式计划下的完整再次排产，没有通过真实采用复核。")
        reopened = client.get(BASE + "/scheduling/candidates/" + again["candidate_ref"] + "/workspace")
        require(reopened["tasks"] == again["workspace"]["data"]["tasks"], "保存的第二轮真实候选回读不一致。")
        _formal(client, artifacts["plan_ref"], artifacts["baseline"]["workspace"]["data"]["tasks"])
        artifacts["second"] = again
        return dict(_summary_run(again), independent_checks=checks, adoption_validation=preview["validation"],
                    candidate_reopened=True, formally_adopted=False, original_formal_plan_unchanged=True)

    try:
        progress.step("exercise:real-preflight-worker-candidates", baseline_run)
        progress.step("exercise:complete-candidate-constraints", verify)
        progress.step("exercise:adopt-and-formal-readback", adopt)
        progress.step("exercise:dashboard-gantt-delivery", projections)
        progress.step("exercise:exact-csv-xlsx-exports", exports)
        progress.step("exercise:independent-conflict-trial", lambda: _conflict_trial(client, artifacts["plan_ref"]))
        progress.step("exercise:existing-formal-real-rerun", second_run)
        formal_by_operation = {row["operation_ref"]: row for row in artifacts["formal"]["data"]["tasks"]}
        baseline_by_operation = {row["operation_ref"]: row for row in artifacts["baseline"]["workspace"]["data"]["tasks"]}
        now_document = client.document(_query(BASE + "/execution/tasks", {"plan_ref": artifacts["plan_ref"], "size": 1}))
        server_now = datetime.fromisoformat(now_document["meta"]["as_of"])
        available = [row for row in formal_by_operation.values() if row["sequence"] == 1 and row["quantity"] > 1
                     and baseline_by_operation[row["operation_ref"]]["source"] == "internal"
                     and datetime.fromisoformat(row["end"]) <= server_now]
        available.sort(key=lambda row: (row["start"], row["batch_id"]))
        wanted = []
        skipped = []
        for request in seed_state["blueprint"]["actual_report_requests"]:
            refs = seed_state["operation_refs"][request["batch_code"]]
            ref = refs.get(request["operation_seq"], refs.get(str(request["operation_seq"])))
            if formal_by_operation[ref] in available:
                wanted.append(formal_by_operation[ref])
            else:
                skipped.append({"batch_code": request["batch_code"], "operation_ref": ref,
                                "reason": "此选择的正式任务还在未来，改选已采用计划中真正过去的首道自制任务。",
                                "planned_end": formal_by_operation[ref]["end"]})
        selected = wanted + [row for row in available if row not in wanted]
        report_limit = min(len(seed_state["blueprint"]["actual_report_requests"]), len(seed_state["refs"]["batch"]) - 1)
        selected = selected[:report_limit]
        require(selected, "当前正式计划没有已过去的首道自制任务，尚不能完成真实分次报工验收。")
        progress.report["actual_report_selection"] = {"server_as_of": now_document["meta"]["as_of"],
            "selected_operation_refs": [row["operation_ref"] for row in selected], "skipped_original_selectors": skipped}
        progress.save()
        reports = []
        for task in selected:
            reports.append(progress.step("exercise:partial-report:" + task["batch_id"],
                           lambda task=task: _partial_report(client, task, max(1, task["quantity"] // 3))))
        all_codes = list(seed_state["refs"]["batch"])
        affected = {row["batch_id"] for row in reports}
        remaining = [code for code in all_codes if code not in affected]
        progress.step("exercise:partial-actual-replan-blocked", lambda: _partial_replan_blocked(client, seed_state, reports, all_codes, "execution_review_required"))
        progress.step("exercise:outside-partial-replan-blocked", lambda: _partial_replan_blocked(client, seed_state, reports, remaining, "execution_ledger_requires_reconciliation"))
        _formal(client, artifacts["plan_ref"], artifacts["baseline"]["workspace"]["data"]["tasks"])
        limitation = "未完成实际工序缺少可信剩余安排，当前产品不支持报工后再次排产（即使只选未报工批次）；已验证准确拒绝并保留原计划/报工。"
        progress.report.setdefault("limitations", []).append(limitation)
        progress.report.setdefault("unsupported", []).append(limitation)
        progress.report.update(state="complete", exercise_state="complete", official_plan_ref=artifacts["plan_ref"],
                               baseline_run_ref=artifacts["baseline"]["run"]["run_ref"],
                               second_run_ref=artifacts["second"]["run"]["run_ref"], actual_report_count=len(reports),
                               completion_scope="样例初始化、既有可用流程和真实拒绝护栏验收完成；不表示已支持报工后再次排产。")
        progress.save()
        return progress.report
    except Exception:
        progress.report["exercise_state"] = "failed"
        progress.save()
        raise


def exercise_sample(base_url, seed_state, output=None):
    client = SampleClient(base_url)
    client.request_count = seed_state.get("request_count", 0)
    progress = SampleProgress(None, client)
    progress.started -= seed_state.get("elapsed_seconds", 0)
    progress.output = Path(output) if output is not None else None
    return exercise(client, seed_state, progress)


def main(argv=None):
    """Host acceptance can reopen an existing seed without injecting it again."""
    import argparse

    parser = argparse.ArgumentParser(description="通过真实HTTP验收已注入的独立复杂样例")
    parser.add_argument("--base-url", required=True)
    parser.add_argument("--seed", required=True)
    parser.add_argument("--output")
    arguments = parser.parse_args(argv)
    source = Path(arguments.seed)
    seed = json.loads(source.read_text(encoding="utf-8"))
    report = exercise_sample(arguments.base_url, seed, arguments.output or str(source))
    print(json.dumps({"state": report["state"], "official_plan_ref": report["official_plan_ref"],
                      "actual_report_count": report["actual_report_count"]}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
