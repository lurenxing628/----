"""Verify scoped sample CSV/XLSX downloads against their real HTTP snapshots."""

import csv
import io
import json
from urllib.error import HTTPError
from urllib.parse import urlencode
from urllib.request import urlopen

from .sample_constraint_checks import require
from .sample_http import SampleAPIError


def _export_cells(document):
    if "candidate" in document["data"]:
        from core.services.workbench.run.candidate_export import HEADERS, export_rows
    else:
        from core.services.workbench.plan.export import HEADERS, export_rows
    return HEADERS, list(export_rows(document["data"], document["meta"]))


def _decoded_cell(value, fmt):
    if not isinstance(value, str):
        return value
    if fmt == "csv" and value.startswith("'"):
        value = value[1:]
    if value == r"\N":
        return None
    return value[1:] if value.startswith("\\\\") else value


def _same_cell(actual, expected, fmt):
    if type(expected) in (int, float):
        # XLSX numeric cells use the writer's 16 significant digit format.
        # This is its actual file precision, not an arbitrary comparison slack.
        number = float(format(expected, ".16g")) if fmt == "xlsx" else float(expected)
        return actual is not None and float(actual) == number
    return actual == expected


def _read_export_rows(content, fmt):
    if fmt == "csv":
        return list(csv.reader(io.StringIO(content.decode("utf-8-sig"))))
    from openpyxl import load_workbook

    require(content.startswith(b"PK"), "实际XLSX不是有效文件。")
    book = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    try:
        sheet = book.active
        if sheet is None:
            raise RuntimeError("实际XLSX缺少工作表。")
        return [list(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        book.close()


def _download(client, path, fmt, document, folder, name):
    client.request_count += 1
    try:
        response = urlopen(client.base_url + path, timeout=client.timeout)
    except HTTPError as exc:
        body = json.loads(exc.read().decode("utf-8"))
        raise SampleAPIError(path, exc.code, body) from exc
    with response:
        content, headers = response.read(), dict(response.headers)
        require(response.getcode() == 200, "实际文件下载没有成功。")
    if folder is not None:
        folder.mkdir(parents=True, exist_ok=True)
        (folder / (name + "." + fmt)).write_bytes(content)
    rows = _read_export_rows(content, fmt)
    _check_export_rows(rows, document, fmt)
    return {"bytes": len(content), "rows": len(rows) - 1, "format": fmt,
            "content_type": headers.get("Content-Type"), "content_disposition": headers.get("Content-Disposition"),
            "operation_times_identical": True, "all_columns_match_real_scope": True,
            "file": str(folder / (name + "." + fmt)) if folder else None}


def _check_export_rows(rows, document, fmt):
    expected_header, expected_rows = _export_cells(document)
    require(tuple(rows[0]) == expected_header and len(rows) - 1 == len(expected_rows), "导出列或行数与真实候选/正式计划范围不一致。")
    decoded = [[_decoded_cell(value, fmt) for value in row] for row in rows[1:]]
    require(all(len(row) == len(wanted) and all(_same_cell(value, target, fmt) for value, target in zip(row, wanted))
                for row, wanted in zip(decoded, expected_rows)), "导出列内容与本段真实快照的资料或派生结果不一致。")


def _export_document(client, document, base, folder, name):
    source_scope = document["data"].get("scope", {})
    scope = {key: source_scope[key] for key in ("range_start", "range_end") if source_scope.get(key) is not None}
    metadata = {"scope": scope, "snapshot_ref": document["meta"]["snapshot_ref"], "as_of": document["meta"]["as_of"]}
    for fmt in ("csv", "xlsx"):
        args = dict(scope, snapshot_ref=document["meta"]["snapshot_ref"], format=fmt)
        metadata[fmt] = _download(client, base + "/export?" + urlencode(args), fmt, document, folder, name)
    return metadata


def _segmented_export_result(segments, assembly):
    unique_count = assembly["unique_task_count"]
    overlap_count = assembly["duplicate_overlap_count"]
    physical_rows = unique_count + overlap_count
    result = {"mode": "segmented_http", "segments": segments, "segment_count": len(segments),
              "unique_task_count": unique_count, "physical_rows": physical_rows,
              "duplicate_overlap_rows": overlap_count,
              "scope_rule": "各段派生结果与原文件按各自scope保留；跨段仅去重相同任务业务核心，不合成全局风险或全量快照。"}
    for fmt in ("csv", "xlsx"):
        result[fmt] = {"mode": "segmented_http", "rows": unique_count, "physical_rows": sum(row[fmt]["rows"] for row in segments),
                       "bytes": sum(row[fmt]["bytes"] for row in segments), "segment_count": len(segments),
                       "operation_times_identical": True, "all_columns_match_real_scope": True,
                       "files": [row[fmt]["file"] for row in segments], "single_full_scope_file": False}
    return result


def verify_exports(client, document, base, folder, name):
    assembly = document.get("assembly")
    if not assembly:
        segment = _export_document(client, document, base, folder, name)
        return {fmt: segment[fmt] for fmt in ("csv", "xlsx")}
    segments = [
        _export_document(client, segment, base, folder, f"{name}-segment-{index:03d}")
        for index, segment in enumerate(assembly["segments"], 1)
    ]
    return _segmented_export_result(segments, assembly)
