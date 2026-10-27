"""Run the decision workbench on an isolated database through real HTTP contracts.

Usage: .venv/bin/python scripts/verify_decision_workbench.py serve
Then:  .venv/bin/python scripts/verify_decision_workbench.py verify --instance <instance.json>
The server remains available for browser inspection until its process is stopped.
"""

from __future__ import annotations

import argparse
import atexit
import csv
import io
import json
import os
import re
import sys
import threading
import time
import uuid
from datetime import datetime
from pathlib import Path
from urllib.parse import urlencode
from urllib.request import urlopen

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

API = "/api/workbench/v1/scheduling"
DEFAULT_ROOT = Path("/tmp/aps-decision-workbench-20261013")

# This oracle reads public API fields independently of the product's row generators.
DETAIL_EXPORT_COMMON = ("排产编号", "查看基准永久编号", "所选方案编号")
BATCH_EXPORT_FIELDS = {
    "批次永久编号": "batch_ref", "批次编号": "batch_id", "零件编号": "part_no", "零件名称": "part_label",
    "数量": "quantity", "优先级": "priority", "交期": "due_date", "交期状态": "due_state", "交期缺项": "due_issue",
    "基准完成状态": "before.completion_state", "所选完成状态": "after.completion_state",
    "基准交付状态": "before.risk", "所选交付状态": "after.risk",
    "基准完整结束": "before.planned_finish", "所选完整结束": "after.planned_finish",
    "基准已排部分结束": "before.partial_planned_finish", "所选已排部分结束": "after.partial_planned_finish",
    "结束变化（小时）": "finish_delta_hours", "基准拖期（小时）": "before.delay_hours",
    "所选拖期（小时）": "after.delay_hours", "拖期变化（小时）": "delay_delta_hours", "变化类型": "change_kinds",
    "工序总数": "after.operation_count", "基准已排工序数": "before.scheduled_operation_count",
    "所选已排工序数": "after.scheduled_operation_count", "基准未排工序数": "before.unscheduled_operation_count",
    "所选未排工序数": "after.unscheduled_operation_count", "基准待料暂缓工序数": "before.deferred_operation_count",
    "所选待料暂缓工序数": "after.deferred_operation_count",
    "基准普通漏排工序数": "before.ordinary_unscheduled_operation_count",
    "所选普通漏排工序数": "after.ordinary_unscheduled_operation_count",
    "所选待料暂缓工序编号": "deferred_operation_refs", "所选普通漏排工序编号": "unplanned_operation_refs",
    "基准缺项": "before.issues", "所选缺项": "after.issues",
}
OPERATION_EXPORT_FIELDS = {
    "批次永久编号": "batch_ref", "批次编号": "batch_label", "工序永久编号": "operation_ref", "零件名称": "part_label",
    "工序号": "sequence", "工序名称": "process_label", "单件编号": "piece_id", "数量": "quantity",
    "基准开始": "before.start", "所选开始": "after.start", "开始变化（小时）": "start_delta_hours",
    "基准结束": "before.end", "所选结束": "after.end", "结束变化（小时）": "end_delta_hours",
    "基准设备编号": "before.machine.ref", "基准设备": "before.machine.label",
    "所选设备编号": "after.machine.ref", "所选设备": "after.machine.label", "设备变化": "machine_changed",
    "基准人员编号": "before.operator.ref", "基准人员": "before.operator.label",
    "所选人员编号": "after.operator.ref", "所选人员": "after.operator.label", "人员变化": "operator_changed",
    "基准来源": "before.source", "所选来源": "after.source", "来源变化": "source_changed", "变化类型": "change_kinds",
    "标准准备工时": "effort.setup_standard_hours", "标准单件工时": "effort.unit_standard_hours",
    "执行目标量": "effort.target_quantity", "剩余数量": "effort.remaining_quantity", "标准总工时": "effort.standard_total_hours",
    "工时口径": None, "工时状态": "effort.state", "受理时保护": "constraints.protected_at_admission",
    "基准锁定": "constraints.reference_locked", "所选锁定": "constraints.selected_locked",
    "受理时执行状态": "constraints.execution_state", "受理时已完数量": "constraints.known_completed_quantity",
    "受理时首次实开": "constraints.first_actual_start", "受理时确认完工": "constraints.confirmed_finish",
    "基准多段安排": "before.segments", "所选多段安排": "after.segments",
    "基准缺项": "before.data_gaps", "所选缺项": "after.data_gaps",
}
EXPORT_DETAIL_LABELS = {
    "specified": "有交期", "unspecified": "未设置交期", "invalid": "交期无效", "complete": "完整结束",
    "deferred": "待料暂缓", "incomplete": "普通漏排", "unknown": "缺少资料", "overdue": "拖期", "on_time": "按期",
    "not_applicable": "未设置交期", "earlier": "提前结束", "later": "延后结束", "new_overdue": "新增拖期",
    "relieved_overdue": "解除拖期", "completion_improved": "完整性改善", "completion_worsened": "完整性降低",
    "unchanged": "相同", "added": "新增安排", "removed": "移除安排", "moved": "时间调整", "machine_changed": "设备调整",
    "operator_changed": "人员调整", "source_changed": "来源调整",
}
RESOURCE_EXPORT_COMMON = {
    "排产编号": "$run_ref", "查看基准永久编号": "$baseline", "所选方案编号": "$candidate.candidate_ref",
    "资源永久编号": "resource_ref", "资源类型": "resource_kind", "资源名称": "label",
    "资源视角": "$scope.resource_view", "范围起点（含）": "$scope.range_start", "范围终点（不含）": "$scope.range_end",
    "资源粒度": "$filters.grain", "高负荷阈值（%）": "$filters.pressure_threshold",
    "筛选时段起点（含）": "$filter_window.start", "筛选时段终点（不含）": "$filter_window.end",
    "资源指标版本": "$metric_version",
}
RESOURCE_EXPORT_MEASURES = {
    "available_hours": "可排总时长（小时）", "occupied_hours": "班内已占（小时）",
    "summed_load_hours": "班内工序区间累加（小时）", "outside_calendar_hours": "安排跨度跨非工作时段（小时）",
    "overlap_hours": "班内重复计时（小时）", "utilization_ratio": "占用比例（%）", "demand_ratio": "工序累加时长与可排时长之比（%）",
    "peak_utilization": "时段占用比例（%）", "span_occupied_hours": "保留安排跨度合计（小时）",
    "span_summed_hours": "保留工序跨度累加（小时）", "span_overlap_hours": "保留安排重复计时（小时）",
    "overlap_span_hours": "保留安排重叠时段（小时）",
}
RESOURCE_EXPORT_NUMBERS = {prefix + label: side + "." + field
    for field, label in RESOURCE_EXPORT_MEASURES.items()
    for prefix, side in (("基准", "before"), ("所选", "after"), ("变化", "deltas"))}
RESOURCE_EXPORT_STATES = {
    "基准覆盖状态": "before.state", "所选覆盖状态": "after.state",
    "基准高负荷": "before.high_pressure", "所选高负荷": "after.high_pressure",
    "基准零容量工作证据": "before.zero_capacity_task", "所选零容量工作证据": "after.zero_capacity_task",
    "基准存在冲突": "before.conflict", "所选存在冲突": "after.conflict",
    "基准班表未知天数": "before.unknown_days", "所选班表未知天数": "after.unknown_days",
    "基准零容量天数": "before.zero_capacity_days", "所选零容量天数": "after.zero_capacity_days",
    "基准安排段数": "before.task_count", "所选安排段数": "after.task_count",
    "基准安排工序数": "before.operation_count", "所选安排工序数": "after.operation_count",
    "基准已知可排时长（小时）": "before.known_available_hours", "所选已知可排时长（小时）": "after.known_available_hours",
    "基准已知班内占用（小时）": "before.known_occupied_hours", "所选已知班内占用（小时）": "after.known_occupied_hours",
    "基准已知班内工序累加时长（小时）": "before.known_summed_load_hours", "所选已知班内工序累加时长（小时）": "after.known_summed_load_hours",
    "基准缺项": "before.issues", "所选缺项": "after.issues",
    "基准高负荷时段数": "before.high_pressure_period_count", "所选高负荷时段数": "after.high_pressure_period_count",
    "基准需求超载时段数": "before.demand_overload_period_count", "所选需求超载时段数": "after.demand_overload_period_count",
    "基准班表完整状态": "before.calendar_state", "所选班表完整状态": "after.calendar_state",
    "基准实际工作片段状态": "before.actual_work_state", "所选实际工作片段状态": "after.actual_work_state",
    "基准零容量证据原因": "before.zero_capacity_reason", "所选零容量证据原因": "after.zero_capacity_reason",
    "班表外跨度计算方式": "=保留安排跨度与非工作窗口的交集",
    "冲突计算方式": "=归档可排窗口内不同永久工序的安排重叠",
}
RESOURCE_EXPORT_FIELDS = {**RESOURCE_EXPORT_COMMON,
    **{prefix + ("日峰占用比例（%）" if field == "peak_utilization" else label): side + "." + field
       for field, label in RESOURCE_EXPORT_MEASURES.items()
       for prefix, side in (("基准", "before"), ("所选", "after"), ("变化", "deltas"))},
    **RESOURCE_EXPORT_STATES,
    "正式归档工序数": "$coverage.after.formal_operation_count",
    "所选覆盖正式工序数": "$coverage.after.covered_formal_operation_count",
    "所选完整覆盖正式范围": "$coverage.after.formal_coverage_complete",
    "所选缺失正式工序数": "$coverage.after.missing_formal_operation_refs",
    "所选其他未安排工序数": "$coverage.after.unplanned_operation_refs",
    "所选缺料未安排工序数": "$coverage.after.deferred_operation_refs",
    "筛选时段峰占用比例（%）": "filter_summary.period_peak_utilization",
    "筛选时段日峰占用比例（%）": "filter_summary.daily_peak_utilization",
    "筛选时段高负荷数": "filter_summary.high_pressure_period_count",
    "筛选时段需求超载数": "filter_summary.demand_overload_period_count",
    "筛选时段班表内冲突数": "filter_summary.conflict_period_count",
    "筛选时段高负荷区间": "filter_summary.high_pressure_periods"}
RESOURCE_DAY_EXPORT_FIELDS = {**RESOURCE_EXPORT_COMMON,
    "日期": "date", "时段起点（含）": "start", "时段终点（不含）": "end",
    **RESOURCE_EXPORT_NUMBERS, **RESOURCE_EXPORT_STATES,
    "基准日峰占用比例（%）": "before.daily_peak_utilization", "所选日峰占用比例（%）": "after.daily_peak_utilization",
    "基准冲突段数": "before.conflicts", "所选冲突段数": "after.conflicts",
    "基准冲突工序数": "before.conflicts", "所选冲突工序数": "after.conflicts"}
RESOURCE_COVERAGE_HEADERS = ("排产编号", "查看基准永久编号", "所选方案编号", "比较侧", "分类", "工序永久编号",
    "批次永久编号", "批次编号", "工序号", "工序名称")
RESOURCE_CONFLICT_HEADERS = (*RESOURCE_EXPORT_COMMON, "日期", "比较侧", "冲突起点（含）", "冲突终点（不含）",
    "工序永久编号", "安排永久编号", "批次永久编号", "批次编号", "工序名称")
RESOURCE_METRIC_HEADERS = ("指标", "分组", "单位", "精度", "改善方向", "范围", "来源", "常用指标", "指标版本")
RESOURCE_METRIC_LABELS = {
    "group": {"completeness": "工序安排", "delivery": "交付", "time": "时间", "preparation": "准备与切换",
              "resources": "资源", "stability": "调整"},
    "improvement": {"lower": "越低越好", "higher": "越高越好", "neutral": "按实际需求判断"},
    "scope": {"full_admission": "整次排产范围", "full_candidate": "完整候选安排", "completed_batches_only": "工序已排齐的批次",
              "candidate_search": "本候选计算", "saved_schedule_resources": "算法保存的资源范围",
              "saved_algorithm_batches": "算法评价批次集合", "fixed_resource_window": "相同设备人员与班表范围",
              "fixed_changeover_window": "后一道工序在所选范围内开始的设备工种切换"},
    "source": {"admission": "排产时归档", "candidate_tasks": "保存的候选安排", "archived_delivery": "保存的批次工序安排和交期",
               "saved_algorithm": "排产时算法统计", "candidate_comparison": "按同一道工序比较安排",
               "archived_resource_calendar": "排产时归档班表与保存安排",
               "runtime_operation_effort": "计算时保存的准备和加工所需标准工时",
               "archived_changeover_events": "同一设备上相邻工序的工种变化"},
}


PREPARATION_EXPORT_EFFORT_FIELDS = {
    "来源": "source", "需求数量依据": "quantity_basis", "需求目标量": "target_quantity",
    "执行数量依据": "execution_quantity_basis", "执行目标量": "execution_target_quantity", "执行剩余量": "remaining_quantity",
    "准备标准工时": "setup_standard_hours", "加工标准工时": "processing_standard_hours",
    "所需标准工时": "required_standard_hours", "准备状态": "setup_state", "安排类型": "requirement_role",
    "工序安排状态": "allocation_state", "外协周期天": "external_cycle_days", "外协组编号": "external_group_id",
    "外协合并模式": "external_merge_mode", "工时计算方式": "effort_basis", "工时资料状态": "state",
    "需求来源类型": "evidence_source.kind", "需求来源排产编号": "evidence_source.run_ref",
    "需求来源候选编号": "evidence_source.candidate_ref", "需求来源试调编号": "evidence_source.scenario_ref",
    "需求来源计划编号": "evidence_source.plan_ref", "需求来源正式版本": "evidence_source.version",
    "归档标准准备工时": "standard.setup_standard_hours", "归档标准单件工时": "standard.unit_standard_hours",
    "归档标准目标量": "standard.target_quantity", "归档标准总工时": "standard.standard_total_hours",
    "归档标准状态": "standard.state", "归档标准缺项": "standard.reason", "需求分解缺项": "data_gaps",
}
PREPARATION_EXPORT_FIELDS = {
    "排产编号": "$run_ref", "查看基准永久编号": "$baseline", "所选方案编号": "$candidate.candidate_ref",
    "批次永久编号": "batch_ref", "批次编号": "batch_label", "工序永久编号": "operation_ref", "零件名称": "part_label",
    "工序号": "sequence", "工序名称": "process_label", "单件编号": "piece_id",
    **{prefix + header: side + "." + path for prefix, side in (("基准", "before"), ("所选", "after"))
       for header, path in PREPARATION_EXPORT_EFFORT_FIELDS.items()},
    "基准设备编号": "before.machines.ref", "基准设备": "before.machines.label",
    "所选设备编号": "after.machines.ref", "所选设备": "after.machines.label",
    "基准安排记录": "before.arrangements", "所选安排记录": "after.arrangements",
    **{label + suffix: "deltas." + field + "." + path
       for label, field in (("准备工时变化", "setup_standard_hours"), ("加工工时变化", "processing_standard_hours"),
                            ("所需标准工时变化", "required_standard_hours"))
       for suffix, path in (("", "value"), ("状态", "state"), ("方向", "direction"))},
    "缺少资料": "data_gaps",
}
CHANGEOVER_EXPORT_FIELDS = {
    "排产编号": "$run_ref", "查看基准永久编号": "$baseline", "所选方案编号": "$candidate.candidate_ref",
    "事件编号": "event_ref", "事件侧": "event_side", "设备永久编号": "machine_ref", "设备名称": "machine_label",
    "设备序列位置": "sequence", "前序工序永久编号": "previous_operation_ref", "前序安排永久编号": "previous_row_ref",
    "前序批次永久编号": "previous.batch_ref", "前序批次编号": "previous.batch_label", "前序工序名称": "previous.process_label",
    "前序工种永久编号": "previous.type_ref", "前序工种名称": "previous.type_label",
    "前序开始": "previous_start", "前序结束": "previous_end", "后序工序永久编号": "next_operation_ref",
    "后序安排永久编号": "next_row_ref", "后序批次永久编号": "next.batch_ref", "后序批次编号": "next.batch_label",
    "后序工序名称": "next.process_label", "后序工种永久编号": "next.type_ref", "后序工种名称": "next.type_label",
    "后序开始": "next_start", "后序结束": "next_end", "前序位于窗口之前": "predecessor_outside_window",
    "前后安排重叠": "overlap", "设备安排冲突": "conflict", "事件状态": "state", "事件缺项": "data_gaps",
    "计算方式": None, "事件指标版本": None,
}


RECOMMENDATION_EXPORT_COMMON = ("排产编号", "归档编号", "读取编号", "读取时间", "保存推荐候选编号", "当前查看候选编号", "比较模式",
    "查看基准编号", "查看基准类型", "范围起点（含）", "范围终点（不含）", "资源视角", "推荐版本", "推荐记录格式版本")
RECOMMENDATION_EXPORT_EXPLANATION_FIELDS = ("record_kind", "指标标识", "名称", "范围", "来源", "状态", "变化方向", "单位",
    "基准值", "比较对象值", "差值", "主项类别", "主项顺序", "相关明细类型", "相关明细永久编号",
    "相关明细候选编号", "相关明细页签", "相关明细比较起点", "相关明细比较终点", "相关明细筛选字段", "相关明细筛选值", "说明", "指标版本")
RECOMMENDATION_EXPORT_COMPLETION_FIELDS = {
    "record_kind": None, "安排组号": None, "安排来源类型": None, "候选永久编号": None, "候选名称": None,
    "来源可用": None, "批次永久编号": "batch_ref", "批次编号": "batch_id", "安排分类": None, "双方工序都已排齐": None,
    "批次工序安排状态": "after.completion_state", "交付状态": "after.risk", "交期状态": "due_state", "交期": "due_date",
    "预计完工": "after.planned_finish", "已排部分预计结束": "after.partial_planned_finish", "拖期小时": "after.delay_hours",
    "工序总数": "after.operation_count", "已排工序数": "after.scheduled_operation_count",
    "其他未安排工序数": "after.ordinary_unscheduled_operation_count", "缺料未安排工序数": "after.deferred_operation_count",
    "安排缺少资料": "after.issues",
}
RECOMMENDATION_EXPORT_REMAINING_FIELDS = {
    "record_kind": None, "安排组号": None, "安排来源类型": None, "候选永久编号": None, "候选名称": None, "来源可用": None,
    "批次永久编号": "batch_ref", "批次编号": "batch_label", "工序永久编号": "operation_ref", "零件名称": "part_label",
    "工序号": "sequence", "工序名称": "process_label", "单件编号": "piece_id", "排产时安排状态": None, "排产时说明": None,
    **{header: "after." + path for header, path in PREPARATION_EXPORT_EFFORT_FIELDS.items()},
}
RECOMMENDATION_EXPORT_HEADERS = {
    "推荐说明": RECOMMENDATION_EXPORT_COMMON + RECOMMENDATION_EXPORT_EXPLANATION_FIELDS,
    "工序安排": RECOMMENDATION_EXPORT_COMMON + tuple(RECOMMENDATION_EXPORT_COMPLETION_FIELDS),
    "剩余需求": RECOMMENDATION_EXPORT_COMMON + tuple(RECOMMENDATION_EXPORT_REMAINING_FIELDS),
}
RECOMMENDATION_EXPORT_CSV_HEADERS = tuple(dict.fromkeys(header for headers in RECOMMENDATION_EXPORT_HEADERS.values() for header in headers))
RECOMMENDATION_EXPORT_FIXED_FIELDS = ("generation_policy_version", "selection_policy_version", "primary_objective",
    "target_objectives", "selected_candidate_ref", "objective_winners", "objective_winner_reasons",
    "objective_field_order", "selected_score", "reason_code")


def _save(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")


def _new_root(requested):
    root = Path(requested).resolve()
    for suffix in range(1000):
        candidate = root if suffix == 0 else root.with_name(root.name + "-" + str(suffix))
        try:
            candidate.mkdir(parents=True, exist_ok=False)
            return candidate
        except FileExistsError:
            continue
    raise RuntimeError("隔离实例目录数量已满，请指定新的 --root。")


def _environment(root):
    directories = {"APS_SHARED_DATA_ROOT": root, "APS_LOG_DIR": root / "logs",
                   "APS_BACKUP_DIR": root / "backups", "APS_EXCEL_TEMPLATE_DIR": root / "templates_excel",
                   "APS_SYSTEM_JOURNAL_DIR": root / "system-journal"}
    for key, directory in directories.items():
        directory.mkdir(parents=True, exist_ok=True)
        os.environ[key] = str(directory)
    database = root / "db" / "decision.sqlite"
    database.parent.mkdir(exist_ok=True)
    os.environ["APS_DB_PATH"] = str(database)
    os.environ["APS_ENV"] = "production"
    return database


def _seed(database):
    from core.infrastructure.database import get_connection
    from tests.workbench.identity_metadata_support import insert_row
    from tests.workbench.run_compute_support import RunCase

    conn = get_connection(str(database))
    try:
        conn.execute("INSERT INTO WorkbenchCalendarDefaults(singleton,periods_json) VALUES (1,?)",
                     ('[{"start":"08:00","end":"16:00","day_offset":0}]',))
        for key, label in (("T1", "车削"), ("T2", "铣削")):
            insert_row(conn, "OpTypes", dict(op_type_id=key, name=label))
        for index, label in ((1, "数控车床"), (2, "加工中心")):
            insert_row(conn, "Machines", dict(machine_id="M" + str(index), name=label, op_type_id="T" + str(index)))
            insert_row(conn, "Operators", dict(operator_id="O" + str(index), name="验收计划员 " + str(index)))
            insert_row(conn, "OperatorMachine", dict(operator_id="O" + str(index), machine_id="M" + str(index)))
        insert_row(conn, "Parts", dict(part_no="P1", part_name="决策验收零件"))
        case = RunCase(conn)
        for index, quantity in enumerate((18, 5, 12, 7, 20, 8), 1):
            code = "DC-" + str(index).zfill(2)
            case.batch(code, quantity=quantity, part_name="决策验收零件", due_date="2026-10-" + str(13 + index),
                       priority="urgent" if index in (2, 4) else "normal")
            case.operation(code, seq=1, op_type_name="车削", setup_hours=0.5, unit_hours=0.15 + index * 0.02)
            case.operation(code, seq=2, op_type_id="T2", op_type_name="铣削", machine_id="M2", operator_id="O2",
                           setup_hours=0.75, unit_hours=0.2 + (7 - index) * 0.025)
            if index in (1, 3, 5):
                case.operation(code, seq=3, op_type_name="精车", setup_hours=0.25, unit_hours=0.1)
        case.config(algo_mode="greedy", graph_candidate_weight_count=3, time_budget_seconds=60,
                    ortools_enabled="no", freeze_window_enabled="no")
        return case.settings(*(str(row[0]) for row in conn.execute("SELECT batch_id FROM Batches ORDER BY batch_id")),
                             start_date="2026-10-13", end_date="2026-10-30")
    finally:
        conn.close()


def _start_run(client, settings, evidence, *, allow_partial=False, run_ref=None):
    if run_ref is None:
        preflight = client.post(API + "/preflight", settings)["data"]
        _save(evidence / "preflight.json", {key: preflight[key] for key in
            ("tasks", "scope", "normalized_input", "run_blocked_reasons")})
        preview = client.post(API + "/runs/preview", {"input_ref": preflight["input_ref"]})["data"]
        _save(evidence / "run-preview.json", {"normalized_input": preview["normalized_input"],
            "calendar_check": preview["calendar_check"], "warnings": preview["warnings"],
            "capabilities": preview["write_context"]["capabilities"],
            "blocked_reasons": preview["write_context"]["blocked_reasons"]})
        if preview["write_context"]["capabilities"]["scheduling.run"] is not True:
            raise RuntimeError("隔离排产预检未通过：" + str(preview["write_context"]["blocked_reasons"]))
        accepted = client.post(API + "/runs", {"input_ref": preflight["input_ref"],
            "write_token": preview["write_context"]["write_token"], "request_key": "decision-run-" + uuid.uuid4().hex}, status=202)
        run_ref = accepted["run_ref"]
        _save(evidence / "accepted-run.json", {"run_ref": run_ref})
    started = time.monotonic()
    while True:
        run = client.get(API + "/runs/" + run_ref)
        if run["state"] in ("complete", "partial", "failed", "cancelled", "interrupted"):
            break
        if time.monotonic() - started > 180:
            raise RuntimeError("隔离排产 180 秒后仍在运行，run_ref=" + run_ref)
        time.sleep(0.2)
    _save(evidence / "run.json", run)
    catalog = client.document(API + "/runs/" + run_ref + "/candidates?size=50")
    _save(evidence / "candidate-catalog.json", catalog)
    allowed = ("complete", "partial") if allow_partial else ("complete",)
    if run["state"] not in allowed or catalog["data"]["candidate_count"] != 4:
        raise RuntimeError("隔离实例未产出预期 4 候选，请查看 run.json 与 candidate-catalog.json。")
    return run_ref, [row["candidate_ref"] for row in catalog["data"]["candidates"]]


def _retained_rows(conn):
    tables = ("Batches", "BatchOperations", "ScheduleHistory", "Schedule", "WorkbenchRunJobs", "WorkbenchRunReceipts",
              "WorkbenchRunCandidates", "WorkbenchRunCandidateTasks")
    return {table: {tuple(row) for row in conn.execute("SELECT * FROM " + table)} for table in tables}


def _preparation_fixture_plan():
    def operation(sequence, kind, setup, unit, **extra):
        return dict(seq=sequence, op_type_id="DP-T-" + kind, setup_hours=setup, unit_hours=unit, **extra)

    pieces = [operation(10, "A", 1, 0.5)]
    for piece in ("item-A", "item-B"):
        for sequence, setup in ((20, 0.5), (30, 0.25)):
            pieces.append(operation(sequence, "A", setup, 0.25, piece_id=piece,
                                    op_code="DP-PIECE-" + piece + "-" + str(sequence)))
    pieces.append(operation(40, "A", 0.25, 0.25))
    return {"version": "preparation-real-v1", "types": [
        {"op_type_id": "DP-T-A", "name": "准备验收车削", "category": "internal"},
        {"op_type_id": "DP-T-B", "name": "准备验收铣削", "category": "internal"},
        {"op_type_id": "DP-T-C", "name": "准备验收零工时检验", "category": "internal"},
        {"op_type_id": "DP-T-X", "name": "准备验收外协热处理", "category": "external"}],
        "machines": [{"machine_id": "DP-M1", "name": "准备验收多工种设备", "op_type_id": "DP-T-A",
                      "additional_types": ["DP-T-B", "DP-T-C"]},
                     {"machine_id": "DP-M2", "name": "准备验收分件设备", "op_type_id": "DP-T-A",
                      "additional_types": []},
                     {"machine_id": "DP-M3", "name": "准备验收外协前后设备", "op_type_id": "DP-T-A",
                      "additional_types": ["DP-T-B"]}],
        "parts": [{"part_no": "DP-P", "part_name": "准备需求验收零件"},
                  {"part_no": "DP-XP", "part_name": "外协周期验收零件"}],
        "supplier": {"supplier_id": "DP-S", "name": "准备验收供应商", "op_type_ids": ["DP-T-X"], "default_days": 0.25},
        "batches": [
            {"batch_id": "DP-NORMAL", "part_no": "DP-P", "quantity": 10, "priority": "urgent", "machine_id": "DP-M1",
             "operations": [operation(1, "A", 2, 0.5), operation(2, "A", 0.25, 0.075),
                            operation(3, "B", 0.5, 0.05), operation(4, "C", 0, 0), operation(5, "A", 0.25, 0.075)]},
            {"batch_id": "DP-PIECE", "part_no": "DP-P", "quantity": 2, "priority": "normal", "machine_id": "DP-M2",
             "operations": pieces},
            {"batch_id": "DP-EXTERNAL", "part_no": "DP-XP", "quantity": 2, "priority": "normal", "machine_id": "DP-M3",
             "operations": [operation(1, "A", 0.25, 0.125),
                            operation(2, "X", None, None, source="external", supplier_id="DP-S", ext_days=0.25),
                            operation(3, "B", 0.25, 0.125)]},
            {"batch_id": "DP-UNPLANNED", "part_no": "DP-P", "quantity": 1, "priority": "normal", "machine_id": "DP-M1",
             "operations": [operation(1, "A", 40, 0)]},
            {"batch_id": "DP-STAGE", "part_no": "DP-P", "quantity": 1, "priority": "normal", "machine_id": "DP-M2",
             "operations": [operation(1, "A", 0.5, 0), operation(2, "A", 0.5, 0)]}],
        "material": {"material_id": "DP-MAT", "name": "准备验收后道物料", "unit": "件"},
        "held_batch_id": "DC-05", "held_operation_code": "DC-05-1",
        "settings": {"start_date": "2026-10-13", "end_date": "2026-10-16", "material_strategy": "stage",
                     "hold_window": {"start": "2026-10-13T09:00", "end": "2026-10-13T09:01"}},
        "runtime_config": {"graph_candidate_weight_count": 3, "dispatch_mode": "sgs", "sort_strategy": "priority_first"}}


def _seed_preparation_case(conn, plan):
    from core.infrastructure.transaction import TransactionManager
    from core.services.workbench.batch.materials import WorkbenchBatchMaterialService
    from tests.workbench.identity_metadata_support import insert_row
    from tests.workbench.run_compute_support import RunCase

    case = RunCase(conn)
    type_labels = {row["op_type_id"]: row["name"] for row in plan["types"]}
    for row in plan["types"]:
        insert_row(conn, "OpTypes", row)
    for row in plan["parts"]:
        insert_row(conn, "Parts", row)
    insert_row(conn, "Suppliers", {key: value for key, value in plan["supplier"].items() if key != "op_type_ids"})
    for kind in plan["supplier"]["op_type_ids"]:
        insert_row(conn, "WorkbenchSupplierOpTypes", {"supplier_id": plan["supplier"]["supplier_id"], "op_type_id": kind})
    for number, row in enumerate(plan["machines"], 1):
        insert_row(conn, "Machines", {key: value for key, value in row.items() if key != "additional_types"})
        for kind in row["additional_types"]:
            insert_row(conn, "MachineOpTypes", {"machine_id": row["machine_id"], "op_type_id": kind})
        operator = "DP-O" + str(number)
        insert_row(conn, "Operators", {"operator_id": operator, "name": "准备验收人员 " + str(number)})
        insert_row(conn, "OperatorMachine", {"operator_id": operator, "machine_id": row["machine_id"]})
    part_labels = {row["part_no"]: row["part_name"] for row in plan["parts"]}
    operations, expected = [], []
    for batch in plan["batches"]:
        code = batch["batch_id"]
        case.batch(code, part_no=batch["part_no"], part_name=part_labels[batch["part_no"]], quantity=batch["quantity"],
                   priority=batch["priority"], due_date="2026-10-16")
        for fields in batch["operations"]:
            values = dict(fields, op_type_name=type_labels[fields["op_type_id"]])
            if values.get("source") == "external":
                insert_row(conn, "PartOperations", dict(part_no=batch["part_no"], status="active", **values))
                values.update(machine_id=None, operator_id=None)
            else:
                values.update(machine_id=batch["machine_id"], operator_id="DP-O" + batch["machine_id"][-1])
            op_id = case.operation(code, **values)
            operations.append(op_id)
            operation_ref = conn.execute("SELECT ref FROM WorkbenchPlanSourceRefs WHERE kind='operation' "
                "AND source_key=? AND active=1", (str(op_id),)).fetchone()[0]
            target = 1 if fields.get("piece_id") is not None else batch["quantity"]
            expected.append({"operation_ref": operation_ref, "op_id": op_id, "batch_id": code,
                "sequence": fields["seq"], "piece_id": fields.get("piece_id"), "op_type_id": fields["op_type_id"],
                "source": values.get("source", "internal"), "target_quantity": target,
                "setup_standard_hours": fields["setup_hours"],
                "processing_standard_hours": None if fields["unit_hours"] is None else fields["unit_hours"] * target,
                "nominal_standard_hours": None if fields["setup_hours"] is None else fields["setup_hours"] + fields["unit_hours"] * target,
                "external_cycle_days": fields.get("ext_days")})
    insert_row(conn, "Materials", plan["material"])
    conn.commit()
    deferred_ref = next(row["operation_ref"] for row in expected if row["batch_id"] == "DP-STAGE" and row["sequence"] == 2)
    with TransactionManager(conn).transaction():
        WorkbenchBatchMaterialService(conn).apply(case.ref("batch", "DP-STAGE"), {"removed_keys": [], "rows": [{
            "row_key": None, "material_ref": case.ref("material", plan["material"]["material_id"]),
            "required_quantity": 1, "available_quantity": 0, "operation_ref": deferred_ref, "arrivals": []}]})
    old_ids = [row[0] for row in conn.execute("SELECT id FROM BatchOperations WHERE batch_id=? ORDER BY seq,id",
                                          (plan["held_batch_id"],))]
    old_refs = [conn.execute("SELECT ref FROM WorkbenchPlanSourceRefs WHERE kind='operation' AND source_key=? AND active=1",
                           (str(value),)).fetchone()[0] for value in old_ids]
    held = conn.execute("SELECT s.start_time,s.end_time,s.machine_id,s.operator_id,r.ref FROM Schedule s "
        "JOIN BatchOperations o ON o.id=s.op_id JOIN WorkbenchPlanSourceRefs r ON r.kind='operation' "
        "AND r.source_key=CAST(o.id AS TEXT) AND r.active=1 WHERE o.op_code=? AND s.version=1",
        (plan["held_operation_code"],)).fetchone()
    if held is None or len(old_refs) != 3 or len(expected) != 17:
        raise RuntimeError("准备夹具的旧正式锚或新十七工序与批准参数不同。")
    settings = case.settings(*(row["batch_id"] for row in plan["batches"]), plan["held_batch_id"], **plan["settings"])
    scope_types = {}
    for op_id in operations + old_ids:
        row = conn.execute("SELECT o.op_type_id,t.name,r.ref FROM BatchOperations o "
            "LEFT JOIN OpTypes t ON t.op_type_id=o.op_type_id "
            "JOIN WorkbenchPlanSourceRefs r ON r.kind='operation' AND r.source_key=CAST(o.id AS TEXT) "
            "AND r.active=1 WHERE o.id=?", (op_id,)).fetchone()
        scope_types[row[2]] = {"id": row[0], "ref": case.ref("op_type", row[0]), "label": row[1]}
    return settings, {"new_batch_ids": [row["batch_id"] for row in plan["batches"]],
        "batch_refs": settings["batch_refs"], "operation_refs": [row["operation_ref"] for row in expected] + old_refs,
        "new_operation_refs": [row["operation_ref"] for row in expected], "expected_standard": expected,
        "unplanned_operation_ref": next(row["operation_ref"] for row in expected if row["batch_id"] == "DP-UNPLANNED"),
        "deferred_operation_ref": deferred_ref, "point_operation_ref": next(row["operation_ref"] for row in expected
            if row["batch_id"] == "DP-NORMAL" and row["sequence"] == 4),
        "held_operation": {"operation_ref": held[4], "start": held[0], "end": held[1],
                           "machine_ref": case.ref("machine", held[2]), "operator_ref": case.ref("operator", held[3])},
        "machine_refs": [case.ref("machine", row["machine_id"]) for row in plan["machines"]],
        "type_refs": {row["op_type_id"]: case.ref("op_type", row["op_type_id"]) for row in plan["types"]},
        "scope_type_identities": scope_types}


def preparation_fixtures(args):
    identity, client = _identity(args.instance)
    from core.infrastructure.database import get_connection
    from tests.workbench.run_compute_support import RunCase

    database = _environment(Path(identity["root"]))
    if database.resolve() != Path(identity["database_path"]).resolve():
        raise RuntimeError("准备需求验收夹具数据库与当前隔离实例不同。")
    evidence = ROOT / "evidence" / "decision-preparation-events" / Path(identity["root"]).name
    complete_formal = args.complete_formal_scope
    complete_piece = args.complete_piece_after_failure or complete_formal
    suffix = "-complete-formal-scope" if complete_formal else "-complete-piece" if complete_piece else ""
    manifest_path = evidence / ("preparation" + suffix + "-fixture.json")
    run_evidence = evidence / ("complete_formal_scope" if complete_formal else "complete_piece" if complete_piece else "mixed_preparation")
    manifest = json.loads(manifest_path.read_text(encoding="utf-8")) if manifest_path.exists() else None
    if manifest and Path(manifest["database_path"]).resolve() != database.resolve():
        raise RuntimeError("准备需求夹具记录属于其他数据库。")
    if manifest and manifest.get("state") == "generated":
        print(json.dumps(manifest, ensure_ascii=False), flush=True)
        return
    plan = _preparation_fixture_plan()
    if not complete_piece:
        _save(evidence / "fixture-plan.json", plan)
    conn = get_connection(str(database))
    try:
        retained = _retained_rows(conn)
        source_tables = ("Parts", "PartOperations", "Machines", "MachineOpTypes", "Operators", "OperatorMachine",
            "OperatorSkill", "OpTypes", "Suppliers", "WorkbenchSupplierOpTypes", "Materials", "BatchMaterials",
            "BatchMaterialArrivals", "BatchMaterialReviews", "BatchMaterialStages", "BatchExternalContexts",
            "WorkCalendar", "OperatorCalendar", "MachineDowntimes", "WorkbenchCalendarDefaults",
            "WorkbenchProductionReports")
        old_sources = {name: {tuple(row) for row in conn.execute("SELECT * FROM " + name)} for name in source_tables}
        keys = tuple(plan["runtime_config"])
        config = dict(conn.execute("SELECT config_key,config_value FROM ScheduleConfig WHERE config_key IN (?,?,?)", keys))
        if manifest is None:
            _save(evidence / ("runtime-before" + suffix + ".json"),
                {"base_url": identity["base_url"], "pid": identity["pid"],
                "database_path": str(database), "old_rows": {name: len(rows) for name, rows in retained.items()},
                "old_source_rows": {name: len(rows) for name, rows in old_sources.items()},
                "official_versions": [row[0] for row in conn.execute("SELECT version FROM ScheduleHistory ORDER BY version")],
                "runtime_config": config})
            if complete_piece:
                original = json.loads((evidence / "preparation-fixture.json").read_text(encoding="utf-8"))
                failed = client.document(API + "/runs/" + original["run_ref"])["data"]
                if failed["state"] != "failed" or failed["candidates"]:
                    raise RuntimeError("批准的分件续跑只能沿用已记录的原失败运行。")
                excluded = original["unplanned_operation_ref"]
                excluded_batch = RunCase(conn).ref("batch", "DP-UNPLANNED")
                settings = dict(original["settings"], batch_refs=[ref for ref in original["settings"]["batch_refs"]
                                                                if ref != excluded_batch])
                manifest = {key: value for key, value in original.items() if key not in
                    ("state", "run_ref", "candidate_refs", "settings", "retained_rows")}
                for field in ("operation_refs", "new_operation_refs"):
                    manifest[field] = [ref for ref in manifest[field] if ref != excluded]
                manifest["batch_refs"] = settings["batch_refs"]
                manifest["expected_standard"] = [row for row in manifest["expected_standard"] if row["operation_ref"] != excluded]
                manifest["scope_type_identities"] = {ref: row for ref, row in manifest["scope_type_identities"].items() if ref != excluded}
                manifest.update(settings=settings, state="seeded", database_path=str(database),
                    retained_rows={name: len(rows) for name, rows in retained.items()}, origin_failed_run_ref=original["run_ref"],
                    excluded_admission_batch_ref=excluded_batch, excluded_admission_operation_ref=excluded,
                    run_evidence_directory=str(run_evidence), new_data_created=False)
                if complete_formal:
                    static = json.loads((evidence / "complete-formal-scope-static-plan.json").read_text(encoding="utf-8"))
                    second = json.loads((evidence / "preparation-complete-piece-fixture.json").read_text(encoding="utf-8"))
                    second_failed = client.document(API + "/runs/" + second["run_ref"])["data"]
                    if second_failed["state"] != "failed" or second_failed["candidates"]:
                        raise RuntimeError("完整正式范围续跑必须保留已记录的第二次失败运行。")
                    settings = dict(static["settings"], end_date="2026-10-30")
                    scope_refs = set(static["operation_refs"])
                    placeholders = ",".join("?" for _ in scope_refs)
                    rows = conn.execute("SELECT r.ref,o.id,o.batch_id,o.seq,o.piece_id,o.op_type_id,o.source,"
                        "o.setup_hours,o.unit_hours,o.ext_days,b.quantity,t.name FROM BatchOperations o "
                        "JOIN Batches b ON b.batch_id=o.batch_id JOIN WorkbenchPlanSourceRefs r "
                        "ON r.kind='operation' AND r.source_key=CAST(o.id AS TEXT) AND r.active=1 "
                        "LEFT JOIN OpTypes t ON t.op_type_id=o.op_type_id WHERE r.ref IN (" + placeholders + ") "
                        "ORDER BY o.batch_id,o.seq,o.id", tuple(scope_refs)).fetchall()
                    expected, type_identities = [], {}
                    for row in rows:
                        target_quantity = 1 if row[4] is not None else row[10]
                        expected.append({"operation_ref": row[0], "op_id": row[1], "batch_id": row[2], "sequence": row[3],
                            "piece_id": row[4], "op_type_id": row[5], "source": row[6], "target_quantity": target_quantity,
                            "setup_standard_hours": row[7], "processing_standard_hours": None if row[8] is None else row[8] * target_quantity,
                            "nominal_standard_hours": None if row[7] is None else row[7] + row[8] * target_quantity,
                            "external_cycle_days": row[9]})
                        type_identities[row[0]] = {"id": row[5], "ref": RunCase(conn).ref("op_type", row[5]), "label": row[11]}
                    if len(rows) != 31 or len({row[0] for row in rows}) != 31:
                        raise RuntimeError("当前完整正式范围与批准的三十一工序静态计划不同。")
                    manifest.update(settings=settings, operation_refs=sorted(scope_refs), batch_refs=settings["batch_refs"],
                        expected_standard=expected, scope_type_identities=type_identities, full_formal_scope=True,
                        formal_operation_refs=static["formal_operation_refs"], admission_batch_ids=static["batch_ids"],
                        origin_failed_run_refs=[original["run_ref"], second["run_ref"]],
                        generation_horizon_reason="沿原夹具10月30日截止，给四候选完整性留足日历容量。")
                _save(evidence / ("fixture-plan" + suffix + ".json"), {"origin_failed_run_ref": original["run_ref"],
                    "settings": settings, "runtime_config": plan["runtime_config"], "operation_refs": manifest["operation_refs"],
                    "excluded_batch": "DP-UNPLANNED", "new_data_created": False})
            else:
                settings, entities = _seed_preparation_case(conn, plan)
                manifest = dict(entities, settings=settings, database_path=str(database), state="seeded",
                                retained_rows={name: len(rows) for name, rows in retained.items()})
            _save(manifest_path, manifest)
        accepted_file = run_evidence / "accepted-run.json"
        existing_ref = manifest.get("run_ref")
        if existing_ref is None and accepted_file.exists():
            existing_ref = json.loads(accepted_file.read_text(encoding="utf-8"))["run_ref"]
        try:
            if existing_ref is None:
                RunCase(conn).config(**plan["runtime_config"])
            run_ref, refs = _start_run(client, manifest["settings"], run_evidence, allow_partial=True, run_ref=existing_ref)
            manifest.update(run_ref=run_ref, candidate_refs=refs, state="computed")
            _save(manifest_path, manifest)
        finally:
            if accepted_file.exists():
                manifest["run_ref"] = json.loads(accepted_file.read_text(encoding="utf-8"))["run_ref"]
                _save(manifest_path, manifest)
            for key, value in config.items():
                conn.execute("UPDATE ScheduleConfig SET config_value=? WHERE config_key=?", (value, key))
            conn.commit()
        receipt = json.loads(conn.execute("SELECT result_json FROM WorkbenchRunReceipts WHERE run_ref=?", (run_ref,)).fetchone()[0])
        ledger = receipt.get("operation_effort")
        if not ledger or ledger.get("effort_basis") != "runtime_requirement_v1":
            raise RuntimeError("新运行收据没有保存共同运行需求账。")
        if {row["operation_ref"] for row in ledger["items"]} != set(manifest["operation_refs"]):
            raise RuntimeError("新需求账没有覆盖批准的二十道永久工序。")
        _save(evidence / ("saved-operation-effort" + suffix + ".json"), ledger)
        expected_arranged = set(manifest["operation_refs"]) - {
            manifest["unplanned_operation_ref"], manifest["deferred_operation_ref"]}
        actual = []
        for ref in refs:
            artifact = json.loads(conn.execute("SELECT artifact_json FROM WorkbenchRunCandidates WHERE candidate_ref=?", (ref,)).fetchone()[0])
            if artifact.get("analysis_contract_version") != "decision-preparation-v4":
                raise RuntimeError("新候选没有保存第4阶段需求分析合同版本。")
            workspace = client.document(API + "/candidates/" + ref + "/workspace")
            _save(run_evidence / (ref + "-workspace.json"), workspace)
            rows = {row["operation_ref"]: row for row in workspace["data"]["tasks"]}
            if set(rows) != expected_arranged:
                raise RuntimeError("真实准备夹具的已排、普通漏排和待料集合与批准参数不同。")
            held, point = rows[manifest["held_operation"]["operation_ref"]], rows[manifest["point_operation_ref"]]
            if held["start"].replace("T", " ") != manifest["held_operation"]["start"] or held["end"].replace("T", " ") != manifest["held_operation"]["end"]:
                raise RuntimeError("不重排窗口的旧正式锚没有原样保留。")
            if point["start"] != point["end"]:
                raise RuntimeError("准备验收的零工时工序没有产出真实点安排。")
            actual.append({"candidate_ref": ref, "task_count": len(rows), "held_operation_preserved": True, "point_retained": True})
        after = _retained_rows(conn)
        if any(not rows.issubset(after[name]) for name, rows in retained.items()):
            raise RuntimeError("新增准备夹具改变了原运行、正式历史或已保存候选。")
        if any(not rows.issubset({tuple(row) for row in conn.execute("SELECT * FROM " + name)}) for name, rows in old_sources.items()):
            raise RuntimeError("新增准备夹具改变了原资料、日历、报工或执行事实。")
        if dict(conn.execute("SELECT config_key,config_value FROM ScheduleConfig WHERE config_key IN (?,?,?)", keys)) != config:
            raise RuntimeError("准备验收生成参数没有恢复。")
        manifest.update(state="generated", base_url=identity["base_url"], pid=identity["pid"],
            old_sources_retained=True, actual_candidates=actual,
            current_batch_count=conn.execute("SELECT COUNT(*) FROM Batches").fetchone()[0],
            current_operation_count=conn.execute("SELECT COUNT(*) FROM BatchOperations").fetchone()[0])
        _save(manifest_path, manifest)
        print(json.dumps(manifest, ensure_ascii=False), flush=True)
    finally:
        conn.close()


def _preparation_expected_events(workspace, fixture, side):
    """Read the saved whole machine sequence before applying any window."""
    machines = {}
    for task in workspace["data"]["tasks"]:
        if task["source"] != "internal" or task["start"] == task["end"] or task["machine"] is None:
            continue
        machines.setdefault(task["machine"]["ref"], []).append(task)
    events = {}
    for machine, tasks in machines.items():
        tasks.sort(key=lambda row: (row["start"], row["end"], row["operation_ref"], row["row_ref"]))
        for previous, following in zip(tasks, tasks[1:]):
            a, b = (fixture["scope_type_identities"][row["operation_ref"]] for row in (previous, following))
            if a["id"] == b["id"]:
                continue
            events[(side, previous["row_ref"], following["row_ref"])] = {
                "machine_ref": machine, "previous_operation_ref": previous["operation_ref"],
                "next_operation_ref": following["operation_ref"], "previous_start": previous["start"],
                "previous_end": previous["end"], "next_start": following["start"], "next_end": following["end"],
                "previous_type_ref": a["ref"], "next_type_ref": b["ref"]}
    return events


def preparation_verify(args):
    """Compare HTTP projections with the original receipt and saved task rows."""
    identity, client = _identity(args.instance)
    evidence = ROOT / "evidence" / "decision-preparation-events" / Path(identity["root"]).name
    import math

    formal_path = evidence / "preparation-complete-formal-scope-fixture.json"
    complete_formal = formal_path.exists()
    complete_path = formal_path if complete_formal else evidence / "preparation-complete-piece-fixture.json"
    complete_piece = complete_path.exists()
    fixture = json.loads((complete_path if complete_piece else evidence / "preparation-fixture.json").read_text(encoding="utf-8"))
    if fixture["state"] != "generated":
        raise RuntimeError("真实准备需求夹具尚未完成。")
    target = evidence / "verification"
    selected, reference = fixture["candidate_refs"][-1], fixture["candidate_refs"][0]
    path = API + "/runs/" + fixture["run_ref"] + "/decision"
    query = {"comparison_mode": "candidates", "selected_candidate_ref": selected,
             "reference_candidate_ref": reference, "resource_view": "selected_batches",
             "range_start": "2026-10-13T00:00:00", "range_end": "2026-10-17T00:00:00"}
    suffix = "-complete-formal-scope" if complete_formal else "-complete-piece" if complete_piece else ""
    ledger = json.loads((evidence / ("saved-operation-effort" + suffix + ".json")).read_text(encoding="utf-8"))
    saved = {row["operation_ref"]: row for row in ledger["items"]}
    run_evidence = Path(fixture.get("run_evidence_directory", str(evidence / "mixed_preparation")))
    workspaces = {side: json.loads((run_evidence / (ref + "-workspace.json")).read_text(encoding="utf-8"))
                  for side, ref in (("before", reference), ("after", selected))}
    overview = client.document(path + "?" + urlencode(query))
    _save(target / "overview.json", overview)
    preparation = _detail_pages(client, path, query, target, "preparation")
    changeovers = _detail_pages(client, path, query, target, "changeovers")
    rows = preparation["data"]["items"]
    if {row["operation_ref"] for row in rows} != set(fixture["operation_refs"]):
        raise RuntimeError("准备需求没有完整返回批准的永久工序集合。")
    fields = ("source", "quantity_basis", "target_quantity", "execution_target_quantity", "execution_quantity_basis",
              "remaining_quantity", "setup_standard_hours", "processing_standard_hours", "required_standard_hours",
              "setup_state", "requirement_role", "external_cycle_days", "external_group_id", "external_merge_mode", "effort_basis")
    checked = 0
    for row in rows:
        original = saved[row["operation_ref"]]
        for side in ("before", "after"):
            for field in fields:
                if row[side][field] != original.get(field):
                    raise RuntimeError(f"准备需求与真实收据账不同：{row['operation_ref']} / {side} / {field}")
                checked += 1
            actual_tasks = [task for task in workspaces[side]["data"]["tasks"] if task["operation_ref"] == row["operation_ref"]]
            arrangements = row[side]["arrangements"]
            expected_identity = {(task["row_ref"], task["start"], task["end"]) for task in actual_tasks}
            if {(task["row_ref"], task["start"], task["end"]) for task in arrangements} != expected_identity:
                raise RuntimeError("准备需求的实际安排身份或时刻与原候选不同。")
    by_ref = {row["operation_ref"]: row for row in rows}
    for expected in fixture["expected_standard"]:
        row = by_ref[expected["operation_ref"]]["after"]
        if expected["source"] == "internal":
            standard = row["standard"]
            if standard["target_quantity"] != expected["target_quantity"] or standard["standard_total_hours"] is None or not math.isclose(
                    standard["standard_total_hours"], expected["nominal_standard_hours"], abs_tol=1e-8):
                raise RuntimeError("共同工序或分件的归档标准目标与真实输入不同。")
            if expected["operation_ref"] not in (fixture["deferred_operation_ref"], fixture["held_operation"]["operation_ref"]):
                for field in ("setup_standard_hours", "processing_standard_hours"):
                    if row[field] != expected[field]:
                        raise RuntimeError("真实交排槽需求与批准的工时公式不同。")
        elif row["external_cycle_days"] != 0.25 or row["allocation_state"] != "external":
            raise RuntimeError("真实外协周期没有独立保留。")
    if by_ref[fixture["point_operation_ref"]]["after"]["required_standard_hours"] != 0:
        raise RuntimeError("真实零工时准备需求没有保留原生零。")
    deferred = by_ref[fixture["deferred_operation_ref"]]["after"]
    held = by_ref[fixture["held_operation"]["operation_ref"]]["after"]
    if deferred["allocation_state"] != "deferred" or deferred["required_standard_hours"] is not None or deferred["standard"]["standard_total_hours"] != 0.5:
        raise RuntimeError("待料暂缓的运行缺项与归档标准没有分开呈现。")
    if held["allocation_state"] != "retained" or held["target_quantity"] is not None or held["execution_target_quantity"] != 20:
        raise RuntimeError("旧冻结需求目标与本次执行目标没有分开保留。")
    expected_metrics = {"pending_setup_hours": (None, 47.75, 17, 1, 18),
        "scheduled_setup_hours": (7.75, 7.75, 16, 0, 16), "unplanned_setup_hours": (40, 40, 1, 0, 1),
        "deferred_setup_hours": (None, 0, 0, 1, 1), "pending_processing_hours": (None, 17, 17, 1, 18),
        "scheduled_processing_hours": (17, 17, 16, 0, 16), "unplanned_processing_hours": (0, 0, 1, 0, 1),
        "deferred_processing_hours": (None, 0, 0, 1, 1)}
    if complete_piece:
        expected_metrics.update(pending_setup_hours=(None, 7.75, 16, 1, 17),
            pending_processing_hours=(None, 17, 16, 1, 17), unplanned_setup_hours=(0, 0, 0, 0, 0),
            unplanned_processing_hours=(0, 0, 0, 0, 0))
    if complete_formal:
        expected_metrics.update(pending_setup_hours=(None, 14.5, 28, 1, 29),
            pending_processing_hours=(None, 45.55, 28, 1, 29), scheduled_setup_hours=(14.5, 14.5, 28, 0, 28),
            scheduled_processing_hours=(45.55, 45.55, 28, 0, 28))
    for side in ("before", "after"):
        summary = preparation["data"]["summary"]
        for key, values in expected_metrics.items():
            metric = summary[side][key]
            for field, expected in zip(("value", "known_subtotal", "known_count", "unknown_count", "total_count"), values):
                actual = metric[field]
                if (actual is not None if expected is None else actual is None or not math.isclose(actual, expected, abs_tol=1e-8)):
                    raise RuntimeError("新排、漏排、暂缓的完整需求汇总不同：" + key)
        if summary["role_totals"][side]["retained"]["count"] != 1:
            raise RuntimeError("冻结保留种子没有独立于新排需求汇总。")
    for key in expected_metrics:
        delta = preparation["data"]["summary"]["deltas"][key]
        if delta["direction"] not in ("neutral", "equal"):
            raise RuntimeError("准备需求变化被错误解释为优化收益。")
    expected_events = {}
    for side, workspace in workspaces.items():
        expected_events.update(_preparation_expected_events(workspace, fixture, "reference" if side == "before" else "selected"))
    actual_events = changeovers["data"]["items"]
    keys = {(row["event_side"], row["previous_row_ref"], row["next_row_ref"]) for row in actual_events}
    if len(expected_events) != 6 or keys != set(expected_events) or len({row["event_ref"] for row in actual_events}) != len(actual_events):
        raise RuntimeError("真实工种切换与完整设备序列、点事件排除或稳定关系身份不同。")
    for row in actual_events:
        expected = expected_events[(row["event_side"], row["previous_row_ref"], row["next_row_ref"])]
        for field in ("machine_ref", "previous_operation_ref", "next_operation_ref", "previous_start", "previous_end", "next_start", "next_end"):
            if row[field] != expected[field]:
                raise RuntimeError("换型事件没有使用真实前后安排：" + field)
        if row["previous"]["type_ref"] != expected["previous_type_ref"] or row["next"]["type_ref"] != expected["next_type_ref"]:
            raise RuntimeError("换型事件没有使用保存的工种永久身份。")
    for side in ("before", "after"):
        if changeovers["data"]["summary"][side]["type_changeover_count"]["value"] != 3:
            raise RuntimeError("实际切换计数与三次工种身份转换不同。")
    normal_events = sorted((row for row in actual_events if row["event_side"] == "selected"
        and row["machine_ref"] == fixture["machine_refs"][0]), key=lambda row: row["next_start"])
    narrow_query = dict(query, machine_ref=fixture["machine_refs"][0], event_side="selected",
                        range_start=normal_events[0]["next_start"], range_end=normal_events[1]["next_start"])
    narrow = _detail_pages(client, path, narrow_query, target / "window-predecessor", "changeovers")
    if len(narrow["data"]["items"]) != 1 or narrow["data"]["items"][0]["event_ref"] != normal_events[0]["event_ref"] or not narrow["data"]["items"][0]["predecessor_outside_window"]:
        raise RuntimeError("窗口前驱、终点排除或稳定事件关系没有保留。")
    filtered = _detail_pages(client, path, dict(query, batch_ref=fixture["batch_refs"][0]), target / "batch-filter", "preparation")
    if filtered["data"]["summary"] != preparation["data"]["summary"] or len(filtered["data"]["items"]) != 5:
        raise RuntimeError("准备批次筛选改变完整需求汇总或没有缩小明细。")
    old_manifest = json.loads((ROOT / "evidence" / "decision-delivery-impact" / Path(identity["root"]).name /
                              "delivery-fixtures.json").read_text(encoding="utf-8"))
    old = old_manifest["fixtures"]["mixed_due"]
    old_query = dict(query, selected_candidate_ref=old["candidate_refs"][-1], reference_candidate_ref=old["candidate_refs"][0],
                     range_end="2026-10-21T00:00:00")
    legacy = _detail_pages(client, API + "/runs/" + old["run_ref"] + "/decision", old_query,
                           target / "legacy-multipage", "preparation")
    if len(legacy["data"]["items"]) != 42:
        raise RuntimeError("既有四十二工序归档没有提供真实多页准备明细。")
    for row in legacy["data"]["items"]:
        for side in ("before", "after"):
            if row[side]["effort_basis"] is not None or any(row[side][field] is not None for field in
                    ("setup_standard_hours", "processing_standard_hours", "required_standard_hours")):
                raise RuntimeError("旧归档缺少runtime账却回填了新的运行需求。")
            if row[side]["standard"]["standard_total_hours"] is None:
                raise RuntimeError("旧归档已知原标准没有独立保留。")
    result = {"state": "passed", "base_url": identity["base_url"], "pid": identity["pid"], "run_ref": fixture["run_ref"],
              "operation_count": len(rows), "receipt_fields_checked": checked, "expected_metrics": expected_metrics,
              "event_count": len(actual_events), "independent_events": list(expected_events.values()),
              "point_excluded_from_changeovers": True, "piece_targets": [1, 2], "external_cycle_days": 0.25,
              "retained_original_target_unknown_execution_target": 20, "narrow_window_predecessor_and_end_exclusion": True,
              "pages_20_50_same_snapshot": True, "batch_filter_preserves_full_summary": True,
              "legacy_multipage_operation_count": 42, "legacy_runtime_unknown_standard_known": True,
              "origin_failed_run_ref": fixture.get("origin_failed_run_ref"), "new_data_created": False if complete_piece else True,
              "origin_failed_run_refs": fixture.get("origin_failed_run_refs", []), "full_formal_scope": complete_formal,
              "ordinary_unplanned_source": "existing real SQLite worker tests and legacy ordinary_partial archive" if complete_piece else "this fixture",
              "actual_engine_rerun": False, "adoption_performed": False}
    _save(evidence / "preparation-verification.json", result)
    print(json.dumps(result, ensure_ascii=False), flush=True)


def _seed_delivery_case(conn, name):
    from core.infrastructure.transaction import TransactionManager
    from core.services.workbench.batch.materials import WorkbenchBatchMaterialService
    from tests.workbench.identity_metadata_support import insert_row
    from tests.workbench.run_compute_support import RunCase

    case = RunCase(conn)
    prefix = {"mixed_due": "DD-D", "ordinary_partial": "DD-P", "legal_stage": "DD-S"}[name]
    op_type, machine, operator = prefix + "-T", prefix + "-M", prefix + "-O"
    insert_row(conn, "OpTypes", dict(op_type_id=op_type, name="交付验收 " + name))
    insert_row(conn, "Machines", dict(machine_id=machine, name="交付验收专用设备 " + name, op_type_id=op_type))
    insert_row(conn, "Operators", dict(operator_id=operator, name="交付验收专用人员 " + name))
    insert_row(conn, "OperatorMachine", dict(operator_id=operator, machine_id=machine))
    batches, operations = [], []

    def add(code, due, hours, *, priority="normal", process="交付验收工序"):
        case.batch(code, quantity=1, part_name="交付影响验收零件", due_date=due, priority=priority)
        batches.append(code)
        for sequence, duration in enumerate(hours, 1):
            operations.append(case.operation(code, seq=sequence, op_type_id=op_type,
                op_type_name=process + " " + str(sequence), machine_id=machine, operator_id=operator,
                setup_hours=duration, unit_hours=0))

    if name == "mixed_due":
        for number in range(1, 22):
            due = "2026-10-12" if number <= 3 else "2026-10-13" if number <= 6 else "2026-10-20"
            if number >= 19:
                due = "9999-12-31" if number == 19 else None
            add(prefix + "-" + str(number).zfill(2), due, (0.3, 0.5),
                priority="urgent" if number in (5, 10) else "normal")
        settings = case.settings(*batches, start_date="2026-10-13", end_date="2026-10-20")
    elif name == "ordinary_partial":
        add(prefix + "-A", "2026-10-13", (6,), priority="urgent")
        add(prefix + "-B", "2026-10-13", (4, 4))
        case.config(dispatch_mode="sgs", sort_strategy="priority_first")
        settings = case.settings(*batches, start_date="2026-10-13", end_date="2026-10-13")
    else:
        add(prefix + "-01", "2026-10-13", (1, 1), process="分阶段待料工序")
        material = prefix + "-MATERIAL"
        insert_row(conn, "Materials", dict(material_id=material, name="交付验收后道物料", unit="件"))
        conn.commit()
        operation_ref = conn.execute("SELECT ref FROM WorkbenchPlanSourceRefs WHERE kind='operation' "
            "AND source_key=? AND active=1", (str(operations[-1]),)).fetchone()[0]
        with TransactionManager(conn).transaction():
            WorkbenchBatchMaterialService(conn).apply(case.ref("batch", batches[0]), {"removed_keys": [], "rows": [{
                "row_key": None, "material_ref": case.ref("material", material), "required_quantity": 1,
                "available_quantity": 0, "operation_ref": operation_ref, "arrivals": []}]})
        settings = case.settings(*batches, start_date="2026-10-13", end_date="2026-10-13", material_strategy="stage")
    conn.commit()
    return settings, {"batch_refs": [case.ref("batch", code) for code in batches], "batch_ids": batches,
        "operation_refs": [conn.execute("SELECT ref FROM WorkbenchPlanSourceRefs WHERE kind='operation' "
            "AND source_key=? AND active=1", (str(value),)).fetchone()[0] for value in operations],
        "machine_ref": case.ref("machine", machine), "operator_ref": case.ref("operator", operator)}


def _seed_resource_case(conn):
    from tests.workbench.identity_metadata_support import insert_row
    from tests.workbench.run_compute_support import RunCase

    case = RunCase(conn)
    old_machines = list(conn.execute("SELECT machine_id,name,op_type_id FROM Machines ORDER BY machine_id"))
    if len(old_machines) != 5:
        raise RuntimeError("资源现场夹具需要沿用已核对的五台原设备。")
    op_type = "DR-T"
    insert_row(conn, "OpTypes", dict(op_type_id=op_type, name="资源影响验收"))
    machines = [(row[0], row[1], row[2]) for row in old_machines]
    for number in range(1, 7):
        machine = "DR-M" + str(number).zfill(2)
        label = "资源验收设备 " + str(number)
        insert_row(conn, "Machines", dict(machine_id=machine, name=label, op_type_id=op_type))
        machines.append((machine, label, op_type))
    operators = ("DR-O1", "DR-O2")
    for operator, label in zip(operators, ("资源验收个人夜班", "资源验收轮班人员")):
        insert_row(conn, "Operators", dict(operator_id=operator, name=label))
        for machine, _, _ in machines:
            insert_row(conn, "OperatorMachine", dict(operator_id=operator, machine_id=machine))

    global_periods = [{"start": "22:00", "end": "00:00", "day_offset": 0},
                      {"start": "00:30", "end": "06:00", "day_offset": 1}]
    global_dates = ("2026-11-02", "2026-11-03", "2026-11-04")
    for day in global_dates:
        insert_row(conn, "WorkCalendar", dict(date=day, day_type="workday", shift_start="22:00", shift_end="06:00",
            shift_hours=7.5, efficiency=1, allow_normal="yes", allow_urgent="yes",
            periods_json=json.dumps(global_periods), remark="资源验收夜班；午夜休息半小时"))
    personal_periods = [{"start": "23:00", "end": "02:00", "day_offset": 0},
                        {"start": "02:30", "end": "07:00", "day_offset": 1}]
    insert_row(conn, "OperatorCalendar", dict(operator_id=operators[0], date=global_dates[0], day_type="workday",
        shift_start="23:00", shift_end="07:00", shift_hours=7.5, efficiency=0.5,
        allow_normal="yes", allow_urgent="yes", periods_json=json.dumps(personal_periods),
        remark="资源验收个人夜班；效率0.5"))
    profile = "DR-ROTATION"
    insert_row(conn, "WorkbenchShiftProfiles", dict(profile_id=profile, name="资源验收三日轮班",
        anchor_date=global_dates[0], cycle_days=3, status="active"))
    insert_row(conn, "WorkbenchOperatorProfiles", dict(operator_id=operators[1], shift_profile_id=profile,
        skills_declared=0))
    rotation_periods = {
        0: [{"start": "22:00", "end": "00:00", "day_offset": 0},
            {"start": "01:00", "end": "06:00", "day_offset": 1}],
        1: [],
        2: [{"start": "22:00", "end": "02:00", "day_offset": 0},
            {"start": "02:30", "end": "06:00", "day_offset": 1}],
    }
    for offset, periods in rotation_periods.items():
        insert_row(conn, "WorkbenchShiftPatternDays", dict(profile_id=profile, day_offset=offset,
            is_rest=int(not periods), shift_start="22:00" if periods else "00:00",
            shift_end="06:00" if periods else "00:00"))
        insert_row(conn, "WorkbenchShiftDayPeriods", dict(profile_id=profile, day_offset=offset,
            periods_json=json.dumps(periods)))
    override_periods = [{"start": "23:30", "end": "01:00", "day_offset": 0},
                        {"start": "01:30", "end": "05:30", "day_offset": 1}]
    insert_row(conn, "OperatorCalendar", dict(operator_id=operators[1], date=global_dates[1], day_type="workday",
        shift_start="23:30", shift_end="05:30", shift_hours=5.5, efficiency=1,
        allow_normal="yes", allow_urgent="yes", periods_json=json.dumps(override_periods),
        remark="资源验收个人日期覆盖轮班休息日"))
    insert_row(conn, "MachineDowntimes", dict(machine_id="DR-M01", scope_type="machine", scope_value="DR-M01",
        start_time="2026-11-02 23:00:00", end_time="2026-11-03 01:00:00",
        reason_code="maintenance", reason_detail="资源验收跨午夜停机", status="active"))

    batches, operations, assignments = [], [], []
    for index, (machine, label, machine_type) in enumerate(machines):
        batch = "DR-B" + str(index + 1).zfill(2)
        operator = operators[index % 2]
        case.batch(batch, quantity=1, part_name="资源影响验收零件", due_date="2026-11-04",
                   priority="urgent" if machine == "DR-M01" else "normal")
        op_id = case.operation(batch, op_type_id=machine_type, op_type_name="资源验收固定设备工序",
            machine_id=machine, operator_id=operator, setup_hours=5 if machine == "DR-M01" else 0.75,
            unit_hours=0)
        operation_ref = conn.execute("SELECT ref FROM WorkbenchPlanSourceRefs WHERE kind='operation' "
            "AND source_key=? AND active=1", (str(op_id),)).fetchone()[0]
        batches.append(batch)
        operations.append(operation_ref)
        assignments.append({"operation_ref": operation_ref, "machine_ref": case.ref("machine", machine),
                            "operator_ref": case.ref("operator", operator), "machine_id": machine,
                            "machine_label": label, "operator_id": operator})
    conn.commit()
    settings = case.settings(*batches, start_date=global_dates[0], end_date=global_dates[-1])
    return settings, {"batch_ids": batches, "batch_refs": [case.ref("batch", value) for value in batches],
        "operation_refs": operations, "assignments": assignments,
        "machine_refs": [case.ref("machine", value[0]) for value in machines],
        "operator_refs": [case.ref("operator", value) for value in operators],
        "global_dates": list(global_dates), "rotation_profile": profile,
        "schema_version": conn.execute("SELECT version FROM SchemaVersion ORDER BY version DESC LIMIT 1").fetchone()[0]}


def resource_fixtures(args):
    identity, client = _identity(args.instance)
    from core.infrastructure.database import get_connection
    from tests.workbench.run_compute_support import RunCase

    database = _environment(Path(identity["root"]))
    if database.resolve() != Path(identity["database_path"]).resolve():
        raise RuntimeError("资源验收夹具数据库与当前隔离实例不同。")
    evidence = ROOT / "evidence" / "decision-resource-impact" / Path(identity["root"]).name
    manifest_path = evidence / "resource-fixture.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8")) if manifest_path.exists() else None
    if manifest and Path(manifest["database_path"]).resolve() != database.resolve():
        raise RuntimeError("资源验收夹具记录属于其他数据库。")
    if manifest and manifest.get("state") == "generated":
        print(json.dumps(manifest, ensure_ascii=False), flush=True)
        return
    conn = get_connection(str(database))
    try:
        retained = _retained_rows(conn)
        source_tables = ("Machines", "Operators", "OperatorMachine", "OpTypes", "WorkCalendar", "OperatorCalendar",
                         "WorkbenchShiftProfiles", "WorkbenchOperatorProfiles", "WorkbenchShiftPatternDays",
                         "WorkbenchShiftDayPeriods", "MachineDowntimes", "WorkbenchCalendarDefaults")
        old_sources = {name: {tuple(row) for row in conn.execute("SELECT * FROM " + name)} for name in source_tables}
        keys = ("graph_candidate_weight_count", "dispatch_mode", "sort_strategy")
        config = dict(conn.execute("SELECT config_key,config_value FROM ScheduleConfig WHERE config_key IN (?,?,?)", keys))
        if manifest is None:
            settings, entities = _seed_resource_case(conn)
            manifest = dict(entities, settings=settings, database_path=str(database), state="seeded",
                            retained_rows={name: len(rows) for name, rows in retained.items()})
            _save(manifest_path, manifest)
        if manifest.get("run_ref"):
            run_ref, refs = manifest["run_ref"], manifest["candidate_refs"]
        else:
            try:
                RunCase(conn).config(graph_candidate_weight_count=3, dispatch_mode="sgs", sort_strategy="priority_first")
                run_ref, refs = _start_run(client, manifest["settings"], evidence / "calendar_resources", allow_partial=True)
                manifest.update(run_ref=run_ref, candidate_refs=refs, state="computed")
                _save(manifest_path, manifest)
            finally:
                for key, value in config.items():
                    conn.execute("UPDATE ScheduleConfig SET config_value=? WHERE config_key=?", (value, key))
                conn.commit()
        actual_by_candidate = []
        expected = {row["operation_ref"]: row for row in manifest["assignments"]}
        for candidate_ref in refs:
            workspace = client.document(API + "/candidates/" + candidate_ref + "/workspace")
            _save(evidence / "calendar_resources" / (candidate_ref + "-workspace.json"), workspace)
            tasks = workspace["data"]["tasks"]
            actual = {row["operation_ref"]: row for row in tasks}
            if set(actual) != set(expected):
                raise RuntimeError("真实资源夹具没有完整安排十一道固定设备工序。")
            if any(row["machine"]["ref"] != expected[operation_ref]["machine_ref"]
                   or row["operator"]["ref"] != expected[operation_ref]["operator_ref"]
                   for operation_ref, row in actual.items()):
                raise RuntimeError("真实引擎没有保留本次输入的固定设备与人员选择。")
            machine_refs = sorted({row["machine"]["ref"] for row in tasks})
            operator_refs = sorted({row["operator"]["ref"] for row in tasks})
            if len(machine_refs) != 11 or len(operator_refs) != 2:
                raise RuntimeError("真实候选的资源并集没有达到十一台设备与两名人员。")
            actual_by_candidate.append({"candidate_ref": candidate_ref, "task_count": len(tasks),
                                        "machine_refs": machine_refs, "operator_refs": operator_refs})
        after = _retained_rows(conn)
        if any(not rows.issubset(after[name]) for name, rows in retained.items()):
            raise RuntimeError("新增资源夹具改变了原排产、正式历史或交付夹具记录。")
        if any(not rows.issubset({tuple(row) for row in conn.execute("SELECT * FROM " + name)})
               for name, rows in old_sources.items()):
            raise RuntimeError("新增资源夹具改变了旧日期、旧资源字段或原日历资料。")
        restored = dict(conn.execute("SELECT config_key,config_value FROM ScheduleConfig WHERE config_key IN (?,?,?)", keys))
        if restored != config:
            raise RuntimeError("资源夹具生成后配置没有恢复原值。")
        manifest.update(state="generated", base_url=identity["base_url"], pid=identity["pid"],
            current_batch_count=len(after["Batches"]), current_operation_count=len(after["BatchOperations"]),
            actual_candidate_resources=actual_by_candidate, restored_configuration=restored,
            retained_source_rows={name: len(rows) for name, rows in old_sources.items()},
            evidence_directory=str(evidence))
        _save(manifest_path, manifest)
        print(json.dumps(manifest, ensure_ascii=False), flush=True)
    finally:
        conn.close()


def _resource_expected_windows(fixture):
    """Independent clock intervals for this fixture's saved three-day shifts."""
    from datetime import datetime

    def periods(rows):
        return [(datetime.fromisoformat(start), datetime.fromisoformat(end)) for start, end in rows]

    global_rows = [("2026-11-02T22:00:00", "2026-11-03T00:00:00"),
                   ("2026-11-03T00:30:00", "2026-11-03T06:00:00"),
                   ("2026-11-03T22:00:00", "2026-11-04T00:00:00"),
                   ("2026-11-04T00:30:00", "2026-11-04T06:00:00"),
                   ("2026-11-04T22:00:00", "2026-11-05T00:00:00")]
    stopped_rows = [("2026-11-02T22:00:00", "2026-11-02T23:00:00"),
                    ("2026-11-03T01:00:00", "2026-11-03T06:00:00")] + global_rows[2:]
    personal_rows = [("2026-11-02T23:00:00", "2026-11-03T02:00:00"),
                     ("2026-11-03T02:30:00", "2026-11-03T07:00:00")] + global_rows[2:]
    rotation_rows = [("2026-11-02T22:00:00", "2026-11-03T00:00:00"),
                     ("2026-11-03T01:00:00", "2026-11-03T06:00:00"),
                     ("2026-11-03T23:30:00", "2026-11-04T01:00:00"),
                     ("2026-11-04T01:30:00", "2026-11-04T05:30:00"),
                     ("2026-11-04T22:00:00", "2026-11-05T00:00:00")]
    result = {row["machine_ref"]: periods(stopped_rows if row["machine_id"] == "DR-M01" else global_rows)
              for row in fixture["assignments"]}
    result[fixture["operator_refs"][0]] = periods(personal_rows)
    result[fixture["operator_refs"][1]] = periods(rotation_rows)
    return result


def _resource_expected_measure(tasks, windows):
    from datetime import datetime

    def merged(rows):
        result = []
        for start, end in sorted(rows):
            if result and start <= result[-1][1]:
                result[-1] = (result[-1][0], max(end, result[-1][1]))
            else:
                result.append((start, end))
        return result

    def hours(rows):
        return sum((end - start).total_seconds() for start, end in rows) / 3600

    def inside(rows):
        return [(max(start, low), min(end, high)) for start, end in rows for low, high in windows
                if start < high and end > low]

    intervals = [(datetime.fromisoformat(row["start"]), datetime.fromisoformat(row["end"])) for row in tasks]
    occupied = hours(inside(merged(intervals)))
    load = hours(inside(intervals))
    capacity = hours(windows)
    return {"available_hours": capacity, "occupied_hours": occupied, "summed_load_hours": load,
            "outside_calendar_hours": hours(merged(intervals)) - occupied,
            "overlap_hours": load - occupied, "span_occupied_hours": hours(merged(intervals)),
            "span_summed_hours": hours(intervals), "span_overlap_hours": hours(intervals) - hours(merged(intervals)),
            "utilization_ratio": occupied / capacity, "demand_ratio": load / capacity}


def resource_verify(args):
    import math

    identity, client = _identity(args.instance)
    evidence = ROOT / "evidence" / "decision-resource-impact" / Path(identity["root"]).name
    fixture = json.loads((evidence / "resource-fixture.json").read_text(encoding="utf-8"))
    if fixture["state"] != "generated":
        raise RuntimeError("真实十一设备资源夹具尚未完成。")
    target = evidence / "verification"
    selected, reference = fixture["candidate_refs"][-1], fixture["candidate_refs"][0]
    path = API + "/runs/" + fixture["run_ref"] + "/decision"
    query = {"comparison_mode": "candidates", "selected_candidate_ref": selected,
             "reference_candidate_ref": reference, "resource_view": "selected_batches",
             "range_start": "2026-11-02T00:00:00", "range_end": "2026-11-05T00:00:00"}
    overview = client.document(path + "?" + urlencode(query))
    _save(target / "overview.json", overview)
    windows = _resource_expected_windows(fixture)
    capacities = {ref: sum((end - start).total_seconds() / 3600 for start, end in rows)
                  for ref, rows in windows.items()}
    expected_totals = {kind: sum(capacities[ref] for ref in fixture[kind + "_refs"])
                       for kind in ("machine", "operator")}
    for candidate in overview["data"]["candidates"]:
        for kind in ("machine", "operator"):
            metrics = candidate["metrics"]
            capacity = metrics[kind + "_available_hours"]
            ratio = metrics[kind + "_utilization_ratio"]
            if not math.isclose(capacity["value"], expected_totals[kind], abs_tol=1e-8):
                raise RuntimeError("全部候选共同容量与真实保存班段不同：" + kind)
            if set(capacity["covered_resource_refs"]) != set(fixture[kind + "_refs"]):
                raise RuntimeError("切换候选改变了共同资源集合：" + kind)
            if not math.isclose(ratio["denominator_hours"], expected_totals[kind], abs_tol=1e-8):
                raise RuntimeError("全部候选占用率没有使用同一资源容量分母：" + kind)
    documents, checked_cells = {}, 0
    for kind in ("machine", "operator"):
        document = _detail_pages(client, path, dict(query, resource_kind=kind, sort="label", order="asc"),
                                 target / kind, "resources")
        documents[kind] = document
        if {row["resource_ref"] for row in document["data"]["items"]} != set(fixture[kind + "_refs"]):
            raise RuntimeError("实际资源明细与十一设备/两名人员身份集合不同：" + kind)
        if document["data"]["rankings"]["total"] != len(fixture[kind + "_refs"]):
            raise RuntimeError("Top资源排行的总数没有覆盖真实共同资源集合。")
        if kind == "machine" and len(document["data"]["rankings"]["top_pressure"]) != 10:
            raise RuntimeError("十一台真实设备没有形成Top10和全部资源的区别。")
        for side, candidate_ref in (("before", reference), ("after", selected)):
            workspace = json.loads((evidence / "calendar_resources" / (candidate_ref + "-workspace.json")).read_text(encoding="utf-8"))
            all_tasks = workspace["data"]["tasks"]
            expected_rows = {}
            for row in document["data"]["items"]:
                ref = row["resource_ref"]
                tasks = [task for task in all_tasks if (task[kind] or {}).get("ref") == ref]
                expected = _resource_expected_measure(tasks, windows[ref])
                expected_rows[ref] = expected
                for field, value in expected.items():
                    actual = row[side][field]
                    if actual is None or not math.isclose(actual, value, abs_tol=1e-8):
                        raise RuntimeError(f"真实安排/班段独立计算与资源明细不同：{kind} / {side} / {field}")
                    checked_cells += 1
            summary = document["data"]["summary"][side]
            capacity = sum(row["available_hours"] for row in expected_rows.values())
            occupied = sum(row["occupied_hours"] for row in expected_rows.values())
            mean = occupied / capacity
            variance = sum(row["available_hours"] * (row["utilization_ratio"] - mean) ** 2
                           for row in expected_rows.values()) / capacity
            expected_cv = math.sqrt(variance) / mean
            if not math.isclose(summary["utilization_ratio"], mean, abs_tol=1e-8):
                raise RuntimeError("容量加权总占用率与真实全部资源分子/分母不同。")
            if not math.isclose(summary["balance_cv"], expected_cv, abs_tol=1e-8):
                raise RuntimeError("容量加权CV与真实全部资源占用率不同。")
        ref = next(row["machine_ref"] for row in fixture["assignments"] if row["machine_id"] == "DR-M01") if kind == "machine" else fixture["operator_refs"][1]
        days = _detail_pages(client, path, dict(query, resource_kind=kind, resource_ref=ref),
                             target / kind, "resource-days")
        if len(days["data"]["items"]) != 3:
            raise RuntimeError("三日班表没有返回三个真实自然日。")
        for side in ("before", "after"):
            row = next(item for item in document["data"]["items"] if item["resource_ref"] == ref)[side]
            for field in ("available_hours", "occupied_hours", "summed_load_hours", "outside_calendar_hours"):
                actual = sum(day[side][field] for day in days["data"]["items"])
                if not math.isclose(actual, row[field], abs_tol=1e-8):
                    raise RuntimeError("逐日数值之和与完整资源窗口不同：" + field)
        heat = next(row for row in document["data"]["heatmap"]["items"] if row["resource_ref"] == ref)
        if document["data"]["heatmap"]["grain"] != "day" or len(heat["cells"]) != len(days["data"]["items"]):
            raise RuntimeError("逐日热力图的粒度或时段数与单资源明细不同。")
        for cell, day in zip(heat["cells"], days["data"]["items"]):
            if any(cell[key] != day[key] for key in ("date", "start", "end")) or any(
                cell[side][field] != day[side][field] for side in ("before", "after")
                for field in ("available_hours", "occupied_hours", "summed_load_hours", "utilization_ratio")):
                raise RuntimeError("热力图实际数值与同资源逐日明细不同。")
    result = {"state": "passed", "base_url": identity["base_url"], "pid": identity["pid"],
              "database_path": identity["database_path"], "run_ref": fixture["run_ref"],
              "candidate_count": len(overview["data"]["candidates"]), "expected_capacities": expected_totals,
              "machine_count": 11, "operator_count": 2, "checked_independent_resource_cells": checked_cells,
              "checks": ["fixed_all_candidate_resource_cohort", "shared_window_denominators", "saved_global_personal_rotation_downtime",
                         "independent_interval_occupancy_demand_outside_overlap", "capacity_weighted_cv", "resource_day_heatmap_equality",
                         "top10_and_all", "same_snapshot_20_50_rows"], "engine_runs_reused": True, "adoption_repeated": False}
    _save(evidence / "resource-verification.json", result)
    print(json.dumps(result, ensure_ascii=False), flush=True)


def delivery_fixtures(args):
    identity, client = _identity(args.instance)
    from core.infrastructure.database import get_connection
    from tests.workbench.run_compute_support import RunCase

    database = _environment(Path(identity["root"]))
    if database.resolve() != Path(identity["database_path"]).resolve():
        raise RuntimeError("交付验收夹具数据库与当前隔离实例不同。")
    evidence = ROOT / "evidence" / "decision-delivery-impact" / Path(identity["root"]).name
    manifest_path = evidence / "delivery-fixtures.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8")) if manifest_path.exists() else {
        "database_path": str(database), "fixtures": {}}
    if Path(manifest["database_path"]).resolve() != database.resolve():
        raise RuntimeError("交付验收夹具记录属于其他数据库。")
    conn = get_connection(str(database))
    try:
        retained = _retained_rows(conn)
        keys = ("graph_candidate_weight_count", "dispatch_mode", "sort_strategy")
        config = dict(conn.execute("SELECT config_key,config_value FROM ScheduleConfig WHERE config_key IN (?,?,?)", keys))
        try:
            for name in args.scenario:
                fixture = manifest["fixtures"].get(name)
                if fixture and fixture.get("run_ref"):
                    continue
                if fixture is None:
                    settings, entities = _seed_delivery_case(conn, name)
                    fixture = dict(entities, settings=settings)
                    manifest["fixtures"][name] = fixture
                    _save(manifest_path, manifest)
                case = RunCase(conn)
                case.config(graph_candidate_weight_count=3, dispatch_mode="sgs",
                            sort_strategy="priority_first" if name == "ordinary_partial" else config["sort_strategy"])
                run_ref, refs = _start_run(client, fixture["settings"], evidence / name, allow_partial=True)
                fixture.update(run_ref=run_ref, candidate_refs=refs)
                _save(manifest_path, manifest)
        finally:
            for key, value in config.items():
                conn.execute("UPDATE ScheduleConfig SET config_value=? WHERE config_key=?", (value, key))
            conn.commit()
        after = _retained_rows(conn)
        if any(not rows.issubset(after[table]) for table, rows in retained.items()):
            raise RuntimeError("新增交付验收夹具时改变了原批次、工序、正式历史或候选归档。")
        restored = dict(conn.execute("SELECT config_key,config_value FROM ScheduleConfig WHERE config_key IN (?,?,?)", keys))
        if restored != config:
            raise RuntimeError("交付验收夹具完成后配置没有恢复原值。")
        complete = all(manifest["fixtures"].get(name, {}).get("run_ref") for name in
            ("mixed_due", "ordinary_partial", "legal_stage"))
        manifest.update(base_url=identity["base_url"], pid=identity["pid"], state="generated" if complete else "partially_generated",
            retained_rows={table: len(rows) for table, rows in retained.items()},
            added_batch_count=sum(len(row["batch_ids"]) for row in manifest["fixtures"].values()),
            added_operation_count=sum(len(row["operation_refs"]) for row in manifest["fixtures"].values()),
            current_batch_count=len(after["Batches"]),
            current_operation_count=len(after["BatchOperations"]), restored_configuration=restored,
            evidence_directory=str(evidence))
        _save(manifest_path, manifest)
        print(json.dumps(manifest, ensure_ascii=False), flush=True)
    finally:
        conn.close()


def serve(args):
    from web.bootstrap.entrypoint import _default_deps, configure_runtime_contract, create_app_with_mode
    from web.bootstrap.launcher_paths import current_runtime_owner
    from web.bootstrap.launcher_runtime_lock import acquire_runtime_lock
    from web.bootstrap.runtime_server import create_runtime_server
    from web.bootstrap.sample_http import SampleClient
    from web.bootstrap.workbench_run_lifecycle import release_runtime_after_jobs, stop_run_runtime
    from web.bootstrap.workbench_run_runtime import install_workbench_run_runtime
    from web.bootstrap.workbench_system_restore import install_workbench_system_restore_host

    previous = json.loads(Path(args.instance).read_text(encoding="utf-8")) if args.instance else None
    root = Path(previous["root"]).resolve() if previous else _new_root(args.root)
    database = _environment(root)
    if previous and Path(previous["database_path"]).resolve() != database:
        raise RuntimeError("原实例数据库路径与隔离目录不同。")
    evidence = ROOT / "evidence" / "decision-comparison-loop" / root.name
    payload = acquire_runtime_lock(str(root), db_path=str(database))
    deps = _default_deps("default")
    lock_state = {"payload": payload}
    atexit.register(release_runtime_after_jobs, deps, lock_state, str(root))
    runtime = app = server = None
    thread = None
    try:
        app = create_app_with_mode("default")
        if Path(app.config["DATABASE_PATH"]).resolve() != database or database == (ROOT / "db" / "aps.db").resolve():
            raise RuntimeError("应用数据库路径与本次隔离实例路径不同。")
        app.config["WORKBENCH_INSTANCE_LABEL"] = "决策验收 · 独立数据"
        settings = None if previous else _seed(database)
        runtime = install_workbench_run_runtime(app, runtime_lock=payload)
        lock_state["runtime"] = runtime
        if not runtime.ready:
            raise RuntimeError("隔离排产宿主未就绪：" + str(runtime.status))
        install_workbench_system_restore_host(app, runtime=runtime)
        server = create_runtime_server(app, "127.0.0.1", 0)
        publication = configure_runtime_contract(app, str(root), "127.0.0.1", server.server_port, current_runtime_owner(),
                                                 deps=deps, runtime_lock=payload)
        lock_state["contract"] = publication["payload"]
        lock_state["contract_cleanup_dir"] = str(Path(publication["path"]).parent)
        thread = threading.Thread(target=deps.serve_runtime_app, args=(app, "127.0.0.1", server.server_port),
                                  kwargs={"server": server}, name="decision-acceptance-http")
        thread.start()
        url = "http://127.0.0.1:" + str(server.server_port)
        client = SampleClient(url, request_prefix="decision-acceptance-")
        health = client.document("/system/health")
        if health["pid"] != os.getpid():
            raise RuntimeError("HTTP 服务 PID 与当前隔离宿主不同。")
        if previous:
            run_ref, refs = previous["run_ref"], previous["candidate_refs"]
            run = client.get(API + "/runs/" + run_ref)
            if run["state"] not in ("complete", "partial"):
                raise RuntimeError("原实例排产尚未产出可读取结果，请先核对原运行。")
        else:
            run_ref, refs = _start_run(client, settings, evidence)
        identity = {"base_url": url, "pid": os.getpid(), "root": str(root), "database_path": str(database),
                    "runtime_contract": str(root / "logs" / "aps_runtime.json"), "run_ref": run_ref,
                    "candidate_refs": refs, "evidence_directory": str(evidence),
                    "instance_file": str(root / "instance.json"), "batch_count": 6, "operation_count": 15}
        _save(root / "instance.json", identity)
        _save(evidence / "instance.json", identity)
        print(json.dumps(identity, ensure_ascii=False), flush=True)
        thread.join()
    except KeyboardInterrupt:
        pass
    finally:
        if server is not None:
            if thread is not None and thread.is_alive():
                server.shutdown()
                thread.join(timeout=15)
            server.server_close()
        if runtime is not None:
            stop_run_runtime(runtime)
        if app is not None:
            gate = app.extensions.get("workbench_request_lifecycle")
            if gate is not None:
                gate.shutdown(wait=True, timeout=15)


def _identity(instance):
    from web.bootstrap.sample_http import SampleClient

    identity = json.loads(Path(instance).read_text(encoding="utf-8"))
    contract = json.loads(Path(identity["runtime_contract"]).read_text(encoding="utf-8"))
    if Path(contract["db_path"]).resolve() != Path(identity["database_path"]).resolve() or contract["pid"] != identity["pid"]:
        raise RuntimeError("隔离实例身份文件与正常运行记录不一致。")
    client = SampleClient(identity["base_url"], request_prefix="decision-acceptance-")
    health = client.document("/system/health")
    if health["pid"] != identity["pid"] or health["status"] != "ok":
        raise RuntimeError("当前 HTTP 服务与隔离实例身份不一致。")
    return identity, client


def _fetch_export(client, path, evidence, fmt, *, section="overview"):
    with urlopen(client.base_url + path, timeout=120) as response:
        payload, headers = response.read(), dict(response.headers)
    target = evidence / ("decision-" + section + "." + fmt)
    target.write_bytes(payload)
    header_prefix = "decision-export" if section == "overview" else "decision-" + section
    _save(evidence / (header_prefix + "-headers-" + fmt + ".json"), {
        key: value for key, value in headers.items() if key.lower().startswith("x-workbench-")})
    return payload, headers


def _check_export_rows(data, rows, fmt):
    definitions = {item["label"]: item for item in data["metric_definitions"]}
    candidates = {item["candidate_ref"]: item for item in data["candidates"]}
    for row in rows:
        candidate = candidates[row["候选编号"]]
        definition = definitions[row["指标"]]
        key = definition["key"]
        metric, delta = candidate["metrics"][key], candidate["deltas"][key]
        expected = {"排产编号": data["run"]["run_ref"], "归档编号": data["capture_ref"],
                    "候选名称": candidate["label"], "查看基准名称": data["baseline"]["label"],
                    "查看基准编号": (data["baseline"]["candidate_ref"] if data["baseline"]["kind"] == "candidate"
                                     else data["baseline"]["baseline_ref"]),
                    "所选方案编号": data["scope"]["selected_candidate_ref"],
                    "推荐方案编号": data["recommendation"]["candidate_ref"], "单位": definition["unit"],
                    "缺项说明": (metric["reason"] or {}).get("message"),
                    "范围起点（含）": data["scope"]["range_start"], "范围终点（不含）": data["scope"]["range_end"],
                    "数值": metric["value"], "已知小计": metric["known_subtotal"],
                    "基准数值": delta["reference_value"], "差值（候选减基准）": delta["value"]}
        if definition["unit"] == "%":
            for field in ("数值", "已知小计", "基准数值", "差值（候选减基准）"):
                if expected[field] is not None:
                    expected[field] *= 100
        for field, value in expected.items():
            actual = row.get(field)
            if fmt == "xlsx" and type(value) in (int, float) and type(actual) not in (int, float):
                raise RuntimeError("XLSX 数字字段被写成文本：{} / {} / {}".format(row["候选编号"], key, field))
            if value is None:
                equal = actual == ("" if fmt == "csv" else None)
            elif type(value) in (int, float):
                # openpyxl writes numeric cells with sixteen significant digits.
                expected_number = float(format(value, ".16g")) if fmt == "xlsx" else float(value)
                equal = actual not in (None, "") and float(actual) == expected_number
            else:
                equal = actual == value
            if not equal:
                raise RuntimeError("导出与屏幕字段不同：{} / {} / {}".format(row["候选编号"], key, field))


def _detail_pages(client, path, query, evidence, section):
    documents, items, page = [], [], 1
    bound = dict(query)
    while True:
        document = client.document(path + "/" + section + "?" + urlencode(dict(bound, page=page, size=20)))
        documents.append(document)
        items.extend(document["data"]["items"])
        bound["snapshot_ref"] = document["meta"]["snapshot_ref"]
        if not document["data"]["page"]["has_more"]:
            break
        page += 1
    wide_pages, wide_items, page = [], [], 1
    while True:
        document = client.document(path + "/" + section + "?" + urlencode(dict(bound, page=page, size=50)))
        wide_pages.append(document)
        wide_items.extend(document["data"]["items"])
        if not document["data"]["page"]["has_more"]:
            break
        page += 1
    if len(items) != documents[0]["data"]["page"]["total"] or items != wide_items:
        raise RuntimeError("20/50 行分页返回的完整明细或顺序不同：" + section)
    snapshot = documents[0]["meta"]["snapshot_ref"]
    if any(item["meta"]["snapshot_ref"] != snapshot for item in documents + wide_pages):
        raise RuntimeError("翻页或切换每页数量改变了明细快照：" + section)
    all_rows = wide_pages[0]
    if len(wide_pages) > 1:
        # This is an explicit collection of unchanged public pages, not a GET.
        all_rows = dict(all_rows, projection_kind="complete_public_page_collection", source_pages=wide_pages,
                        data=dict(all_rows["data"], items=wide_items))
        if section == "resource-days" and wide_items:
            all_rows["data"]["display_scope"] = dict(all_rows["data"]["display_scope"],
                start=wide_items[0]["start"], end=wide_items[-1]["end"])
    _save(evidence / (section + "-pages.json"), documents)
    _save(evidence / (section + "-all.json"), all_rows)
    return all_rows


def _export_path(item, path):
    value = item
    for key in path.split("."):
        if value is None:
            return None
        value = value.get(key) if key == "segments" else value[key]
    return value


def _export_issue(value):
    if isinstance(value, (str, dict)):
        return value if isinstance(value, str) else value.get("message") or value.get("code")
    messages = [entry.get("message") or entry.get("code") if isinstance(entry, dict) else str(entry)
                for entry in value or []]
    return "；".join(messages) or None


def _detail_export_expected(data, item, section):
    baseline = data["baseline"]
    expected = dict(zip(DETAIL_EXPORT_COMMON, (data["run_ref"],
        baseline["candidate_ref"] if baseline["kind"] == "candidate" else baseline["baseline_ref"],
        data["candidate"]["candidate_ref"])))
    fields = BATCH_EXPORT_FIELDS if section == "batches" else OPERATION_EXPORT_FIELDS
    for header, path in fields.items():
        value = _export_path(item, path) if path is not None else "归档工序标准×归档执行目标量"
        if header.endswith("缺项"):
            value = _export_issue(value)
        elif path == "change_kinds":
            value = "、".join(EXPORT_DETAIL_LABELS[kind] for kind in value)
        elif path in ("deferred_operation_refs", "unplanned_operation_refs"):
            value = "、".join(value) or None
        elif path == "due_state" or path and path.endswith((".completion_state", ".risk")):
            value = EXPORT_DETAIL_LABELS[value]
        elif path == "effort.state":
            value = {"complete": "完整", "unavailable": "缺少资料", "not_applicable": "范围无适用对象"}[value]
        elif type(value) is bool:
            value = "是" if value else "否"
        elif header.endswith("多段安排"):
            value = value or None
        expected[header] = value
    return expected


def _check_detail_export_cell(actual, expected, fmt, header, cell_type):
    if expected is None:
        return actual == ("" if fmt == "csv" else None)
    if header.endswith("多段安排"):
        return isinstance(actual, str) and json.loads(actual) == expected and (fmt == "csv" or cell_type == "s")
    if type(expected) in (int, float):
        if fmt == "xlsx" and (type(actual) not in (int, float) or cell_type != "n"):
            return False
        number = float(format(expected, ".16g")) if fmt == "xlsx" else float(expected)
        return actual not in (None, "") and float(actual) == number
    if fmt == "xlsx":
        return actual == expected and cell_type == "s"
    numeric_identity = header.endswith("编号") and re.fullmatch(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?", expected)
    protected = expected.lstrip().startswith(("=", "+", "-", "@")) or numeric_identity
    return actual == ("'" + expected if protected else expected)


def _check_detail_export_table(document, rows, headers, fmt, cell_types):
    data, checked = document["data"], 0
    section = data["filters"]["section"]
    fields = BATCH_EXPORT_FIELDS if section == "batches" else OPERATION_EXPORT_FIELDS
    expected_headers = DETAIL_EXPORT_COMMON + tuple(fields)
    if tuple(headers) != expected_headers or len(rows) != len(data["items"]):
        raise RuntimeError("明细导出列或全部页行数与 API 不同：" + section)
    numeric_cells = 0
    for index, (item, row) in enumerate(zip(data["items"], rows)):
        expected = _detail_export_expected(data, item, section)
        for header, value in expected.items():
            cell_type = cell_types[index].get(header) if cell_types is not None else None
            if not _check_detail_export_cell(row.get(header), value, fmt, header, cell_type):
                raise RuntimeError(f"明细导出字段不同：{section} / 行{index + 1} / {header} / "
                                   f"实际 {row.get(header)!r} / 期待 {value!r}")
            checked += 1
            numeric_cells += type(value) in (int, float)
    return {"rows": len(rows), "columns": len(expected_headers), "checked_cells": checked,
            "numeric_cells": numeric_cells, "identity_order_matches_api": True}


def _check_export_download(content, fmt, documents):
    titles = {"overview": "方案总览", "batches": "批次影响", "operations": "工序变化"}
    checked = {}
    if fmt == "csv":
        section = next(iter(documents))
        reader = csv.DictReader(io.StringIO(content.decode("utf-8-sig")))
        rows = list(reader)
        checked[section] = _check_detail_export_table(documents[section], rows, reader.fieldnames, fmt, None)
        return checked
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    try:
        expected_sheets = [titles[section] for section in documents] + ["指标计算方式", "比较范围"]
        if workbook.sheetnames != expected_sheets:
            raise RuntimeError("XLSX 业务表和附表顺序不同：" + str(workbook.sheetnames))
        for section, document in documents.items():
            iterator = workbook[titles[section]].iter_rows()
            headers = [cell.value for cell in next(iterator)]
            cells = list(iterator)
            rows = [dict(zip(headers, (cell.value for cell in row))) for row in cells]
            types = [dict(zip(headers, (cell.data_type for cell in row))) for row in cells]
            if section == "overview":
                expected_count = len(document["data"]["candidates"]) * len(document["data"]["metric_definitions"])
                if len(rows) != expected_count:
                    raise RuntimeError("全部表中的方案总览行数与完整候选目录不同。")
                _check_export_rows(document["data"], rows, fmt)
                checked[section] = {"rows": len(rows), "columns": len(headers)}
            else:
                checked[section] = _check_detail_export_table(document, rows, headers, fmt, types)
        entries = {row[0]: row[1] if len(row) > 1 else None for row in workbook["比较范围"].iter_rows(values_only=True)}
        active = documents.get("operations") or documents.get("batches")
        if entries["排产编号"] != active["data"]["run_ref"] or entries["明细范围"] != "筛选结果全部页":
            raise RuntimeError("XLSX 比较范围与明细归档身份不同。")
        checked["sheet_names"] = workbook.sheetnames
    finally:
        workbook.close()
    return checked


def _verify_export_case(client, fixture, query, target, *, section, filter_section=None, identity=None):
    if section == "all":
        result, _ = current_all_export_case(client, identity, fixture, query, target, filter_section=filter_section)
        return result
    target.mkdir(parents=True, exist_ok=True)
    path = API + "/runs/" + fixture["run_ref"] + "/decision"
    active_section = filter_section or section
    common_keys = ("comparison_mode", "selected_candidate_ref", "reference_candidate_ref", "range_start", "range_end")
    common = {key: value for key, value in query.items() if key in common_keys}
    active = _detail_pages(client, path, query, target, active_section)
    documents = {active_section: active}
    if section == "all":
        overview = client.document(path + "?" + urlencode(common))
        _save(target / "overview.json", overview)
        documents = {"overview": overview}
        for name in ("batches", "operations"):
            other = dict(common, **({"batch_ref": query["batch_ref"]} if "batch_ref" in query else {}))
            documents[name] = active if name == active_section else _detail_pages(client, path, other, target, name)
            if documents[name]["data"]["scope"] != active["data"]["scope"]:
                raise RuntimeError("全部导出中的明细共同范围不同。")
    expected_count = sum(len(document["data"]["items"]) if name != "overview" else
        len(document["data"]["candidates"]) * len(document["data"]["metric_definitions"])
        for name, document in documents.items())
    export_query = dict(query, section=section, snapshot_ref=active["meta"]["snapshot_ref"])
    if filter_section:
        export_query["filter_section"] = filter_section
    _save(target / "request.json", export_query)
    formats = {}
    for fmt in (("xlsx",) if section == "all" else ("csv", "xlsx")):
        content, headers = _fetch_export(client, path + "/export?" + urlencode(dict(export_query, format=fmt)),
                                         target, fmt, section=section)
        if (headers.get("X-Workbench-Run-Ref") != fixture["run_ref"]
                or headers.get("X-Workbench-Snapshot-Ref") != export_query["snapshot_ref"]
                or int(headers["X-Workbench-Row-Count"]) != expected_count):
            raise RuntimeError("明细导出的运行、快照身份或业务行数合计不同。")
        formats[fmt] = _check_export_download(content, fmt, documents)
    result = {"state": "passed", "run_ref": fixture["run_ref"], "comparison_mode": query["comparison_mode"],
              "section": section, "filter_section": filter_section, "filters": active["data"]["filters"],
              "snapshot_ref": export_query["snapshot_ref"], "business_row_count": expected_count, "formats": formats}
    _save(target / "verification.json", result)
    return result


def delivery_export_verify(args):
    identity, client = _identity(args.instance)
    evidence = ROOT / "evidence" / "decision-delivery-impact" / Path(identity["root"]).name
    manifest = json.loads((evidence / "delivery-fixtures.json").read_text(encoding="utf-8"))
    legacy = json.loads((Path(identity["evidence_directory"]) / "repair-01" / "candidate-count-fixtures.json").read_text(encoding="utf-8"))
    if any(Path(item["database_path"]).resolve() != Path(identity["database_path"]).resolve() for item in (manifest, legacy)):
        raise RuntimeError("明细导出验收夹具与当前隔离数据库不同。")
    fixtures = manifest["fixtures"]
    mixed = fixtures["mixed_due"]
    old = next(item for item in legacy["fixtures"] if item["candidate_count"] == 8)

    def candidate_query(fixture):
        return {"comparison_mode": "candidates", "selected_candidate_ref": fixture["candidate_refs"][-1],
                "reference_candidate_ref": fixture["candidate_refs"][0]}

    candidates = candidate_query(mixed)
    old_query = {"comparison_mode": "candidates", "selected_candidate_ref": old["candidate_refs"][1],
                 "reference_candidate_ref": old["candidate_refs"][0]}
    old_resource = "73e0d4c2bb29e6de6e1d095e6b914932cc021601cc6cb4af"
    cases = []
    for section, sort in (("batches", "batch_id"), ("operations", "sequence")):
        cases.append(("mixed-" + section, mixed, dict(candidates, sort=sort, order="asc"), section, None))
    for name in ("ordinary_partial", "legal_stage"):
        cases.append((name + "-operations", fixtures[name], candidate_query(fixtures[name]), "operations", None))
    batch_filter = dict(candidates, change_kind="new_overdue", sort="delay_delta_hours", order="desc")
    resource_filter = dict(candidates, resource_ref=mixed["machine_ref"], change_kind="moved",
                           sort="start_delta_hours", order="desc")
    cross_filter = dict(candidates, batch_ref=mixed["batch_refs"][3], resource_ref=mixed["operator_ref"],
                        change_kind="moved", sort="end_delta_hours", order="desc")
    cases.extend([
        ("mixed-new-overdue", mixed, batch_filter, "batches", None),
        ("mixed-resource-moved", mixed, resource_filter, "operations", None),
        ("old8-resource-moved", old, dict(old_query, resource_ref=old_resource, change_kind="moved",
            sort="start_delta_hours", order="desc"), "operations", None),
        ("all-batches-filter", mixed, batch_filter, "all", "batches"),
        ("all-operations-filter", old, dict(old_query, resource_ref=old_resource, change_kind="moved",
            sort="start_delta_hours", order="desc"), "all", "operations"),
        ("all-batch-cross-filter", mixed, cross_filter, "all", "operations"),
        ("all-formal-filter", old, {"comparison_mode": "admission_official", "selected_candidate_ref": old["candidate_refs"][1],
            "resource_ref": old_resource, "sort": "end_delta_hours", "order": "desc"}, "all", "operations"),
    ])
    target = evidence / "exports"
    results = []
    current = None
    try:
        for current, fixture, query, section, filter_section in cases:
            result = _verify_export_case(client, fixture, query, target / current, section=section,
                                         filter_section=filter_section, identity=identity)
            results.append(dict(result, case=current))
    except Exception as error:
        _save(target / "export-verification.json", {"state": "failed", "case": current, "error": str(error), "cases": results})
        raise
    result = {"state": "passed", "base_url": identity["base_url"], "pid": identity["pid"],
        "database_path": identity["database_path"], "case_count": len(results), "cases": results,
        "business_row_count": sum(item["business_row_count"] * len(item["formats"]) for item in results),
        "evidence_directory": str(target)}
    _save(target / "export-verification.json", result)
    print(json.dumps({key: value for key, value in result.items() if key != "cases"}, ensure_ascii=False), flush=True)


def _resource_export_expected(data, row, section):
    item = dict(row)
    if section == "resource-days":
        item.update(resource_ref=data["resource"]["ref"], resource_kind=data["resource"]["kind"],
                    label=data["resource"]["label"])
    fields = RESOURCE_EXPORT_FIELDS if section == "resources" else RESOURCE_DAY_EXPORT_FIELDS
    public = dict(data, filter_window=data.get("pressure_scope") or data.get("display_scope") or {
        "start": data["scope"]["range_start"], "end": data["scope"]["range_end"]})
    expected = {}
    for header, path in fields.items():
        if path.startswith("="):
            value = path[1:]
        elif path.startswith("$"):
            value = _export_path(public, path[1:])
        elif path.startswith("deltas."):
            value = item["deltas"].get(path.split(".", 1)[1])
        else:
            value = _export_path(item, path)
        if header == "查看基准永久编号":
            value = value["candidate_ref"] if value["kind"] == "candidate" else value["baseline_ref"]
        elif header == "资源类型":
            value = {"machine": "设备", "operator": "人员"}[value]
        elif header == "资源视角":
            value = {"selected_batches": "本次所选批次", "full_candidate": "方案全部安排"}[value]
        elif header == "资源粒度":
            value = {"hour": "小时", "day": "天", "shift": "真实可排班段", "week": "周"}[value]
        elif header.endswith("覆盖状态"):
            value = {"complete": "完整", "partial": "部分", "unavailable": "缺少资料", "not_applicable": "范围无适用对象"}[value]
        elif header.endswith("缺项"):
            value = _export_issue(value)
        elif header.endswith("零容量证据原因"):
            value = value["message"] if value is not None else None
        elif header == "筛选时段高负荷区间":
            value = "；".join(period["start"] + " / " + period["end"] for period in value) or None
        elif header.endswith("冲突段数"):
            side = path.split(".", 1)[0]
            value = len(value) if item[side]["conflict"] is not None else None
        elif header.endswith("冲突工序数"):
            side = path.split(".", 1)[0]
            value = len({ref for conflict in value for ref in conflict["operation_refs"]}) if item[side]["conflict"] is not None else None
        elif path.startswith("$coverage.") and isinstance(value, list):
            value = len(value)
        elif type(value) is bool:
            value = "是" if value else "否"
        elif header.endswith("（%）") and value is not None:
            value *= 100
        expected[header] = value
    return expected


def _check_resource_export_table(documents, rows, headers, fmt, cell_types):
    section = documents[0]["data"]["filters"]["section"]
    fields = RESOURCE_EXPORT_FIELDS if section == "resources" else RESOURCE_DAY_EXPORT_FIELDS
    pairs = [(document["data"], item) for document in documents for item in document["data"]["items"]]
    if tuple(headers) != tuple(fields) or len(rows) != len(pairs):
        raise RuntimeError("资源导出列或全部页行数与 API 不同：" + section)
    checked = numeric = ratios = blanks = 0
    for index, ((data, item), row) in enumerate(zip(pairs, rows)):
        for header, value in _resource_export_expected(data, item, section).items():
            cell_type = cell_types[index].get(header) if cell_types is not None else None
            if not _check_detail_export_cell(row.get(header), value, fmt, header, cell_type):
                raise RuntimeError(f"资源导出字段不同：{section} / 行{index + 1} / {header} / "
                                   f"实际 {row.get(header)!r} / 期待 {value!r}")
            checked += 1
            numeric += type(value) in (int, float)
            ratios += value is not None and header.endswith("（%）")
            blanks += value is None
    return {"rows": len(rows), "columns": len(fields), "checked_cells": checked, "numeric_cells": numeric,
            "percentage_cells": ratios, "blank_cells": blanks, "identity_and_date_order_matches_api": True}


def _resource_archive_labels(identity, run_ref):
    import sqlite3

    from core.infrastructure.snapshot_connection import ddl_columns

    connection = sqlite3.connect(Path(identity["database_path"]).as_uri() + "?mode=ro", uri=True)
    try:
        connection.execute("PRAGMA query_only=ON")
        row = connection.execute("SELECT facts_json FROM WorkbenchRunJobs WHERE run_ref=?", (run_ref,)).fetchone()
        if row is None:
            raise RuntimeError("资源导出所需的排产归档事实缺失：" + run_ref)
        facts = json.loads(row[0])
    finally:
        connection.close()
    schema = {entry[1]: entry[3] for entry in facts["schema"] if entry[0] == "table"}
    tables = {name: [dict(zip(ddl_columns(schema[name], name, restricted=True), row)) for row in facts["tables"][name]]
              for name in ("BatchOperations", "Batches", "WorkbenchPlanSourceRefs", "WorkbenchEntityRefs")}
    operations = {row["id"]: row for row in tables["BatchOperations"]}
    batches = {row["batch_id"]: row for row in tables["Batches"]}
    batch_refs = {row["entity_key"]: row["ref"] for row in tables["WorkbenchEntityRefs"]
                  if row["kind"] == "batch" and row["active"] == 1}
    result = {}
    for ref in tables["WorkbenchPlanSourceRefs"]:
        if ref["kind"] != "operation" or ref["active"] != 1:
            continue
        operation = operations.get(int(ref["source_key"]), {})
        batch = batches.get(operation.get("batch_id"), {})
        result[ref["ref"]] = {"batch_ref": batch_refs.get(operation.get("batch_id")), "batch_label": batch.get("batch_id"),
            "sequence": operation.get("seq"), "process_label": operation.get("op_type_name")}
    return result


def _resource_coverage_expected(documents, archived_labels):
    resource = documents.get("resources", documents.get("resource-days"))[0]["data"]
    labels = {row["operation_ref"]: row for side in ("before", "after")
              for row in resource["coverage"][side]["missing_formal_arrangements"]}
    labels.update(archived_labels)
    if "operations" in documents:
        labels.update({row["operation_ref"]: row for row in documents["operations"]["data"]["items"]})
    baseline = resource["baseline"]
    common = [resource["run_ref"], baseline["candidate_ref"] if baseline["kind"] == "candidate" else baseline["baseline_ref"],
              resource["candidate"]["candidate_ref"]]
    rows = []
    for side, side_label in (("before", "基准"), ("after", "所选")):
        if side == "before" and baseline["state"] != "available":
            continue
        for field, category in (("missing_formal_operation_refs", "缺失正式"), ("unplanned_operation_refs", "其他未安排"),
                                ("deferred_operation_refs", "缺料未安排")):
            for ref in resource["coverage"][side][field]:
                label = labels.get(ref, {})
                rows.append(dict(zip(RESOURCE_COVERAGE_HEADERS, common + [side_label, category, ref,
                    label.get("batch_ref"), label.get("batch_label"), label.get("sequence"), label.get("process_label")])))
    return rows


def _resource_conflict_expected(documents):
    rows = []
    for document in documents.get("resource-days", []):
        data = document["data"]
        for item in data["items"]:
            day = _resource_export_expected(data, item, "resource-days")
            common = [day[header] for header in RESOURCE_EXPORT_COMMON]
            for side, side_label in (("before", "基准"), ("after", "所选")):
                tasks = {task["row_ref"]: task for task in item[side]["tasks"]}
                for conflict in item[side]["conflicts"]:
                    for row_ref in conflict["row_refs"]:
                        task = tasks[row_ref]
                        rows.append(dict(zip(RESOURCE_CONFLICT_HEADERS, common + [item["date"], side_label,
                            conflict["start"], conflict["end"], task["operation_ref"], row_ref, task.get("batch_ref"),
                            task.get("batch_label"), task.get("process_label")])))
    return rows


def _resource_metric_expected(definitions):
    return [dict(zip(RESOURCE_METRIC_HEADERS, [item["label"], RESOURCE_METRIC_LABELS["group"][item["group"]],
        item["unit"], item["precision"], RESOURCE_METRIC_LABELS["improvement"][item["improvement"]],
        RESOURCE_METRIC_LABELS["scope"][item["scope"]], RESOURCE_METRIC_LABELS["source"][item["source"]],
        "是" if item["common"] else "否", item["metric_version"]])) for item in definitions]


def _check_resource_record_table(sheet, headers, expected):
    iterator = sheet.iter_rows()
    actual_headers = tuple(cell.value for cell in next(iterator))
    if actual_headers != headers:
        raise RuntimeError("资源导出附表列与独立契约不同：" + sheet.title)
    actual = list(iterator)
    if len(actual) != len(expected):
        raise RuntimeError("资源导出附表行数与完整 API 清单不同：" + sheet.title)
    checked = numeric = blanks = 0
    longest = 0
    for index, (row, wanted) in enumerate(zip(actual, expected)):
        for header, cell in zip(headers, row):
            value = wanted[header]
            if not _check_detail_export_cell(cell.value, value, "xlsx", header, cell.data_type):
                raise RuntimeError(f"资源导出附表字段不同：{sheet.title} / 行{index + 1} / {header} / "
                                   f"实际 {cell.value!r} / 期待 {value!r}")
            checked += 1
            numeric += type(value) in (int, float)
            blanks += value is None
            longest = max(longest, len(value) if isinstance(value, str) else 0)
    return {"rows": len(actual), "columns": len(headers), "checked_cells": checked, "numeric_cells": numeric,
            "blank_cells": blanks, "max_text_cell_length": longest, "identity_and_order_matches_api": True}


def _check_resource_export_download(content, fmt, documents, active, metric_definitions, archived_labels):
    titles = {"overview": "方案总览", "batches": "批次影响", "operations": "工序变化",
              "resources": "资源影响", "resource-days": "资源逐日"}
    if "resource-days" in documents:
        titles["resource-days"] = {"day": "资源逐日", "hour": "资源小时", "shift": "资源班次", "week": "资源逐周"}[
            documents["resource-days"][0]["data"]["filters"]["grain"]]
    if fmt == "csv":
        reader = csv.DictReader(io.StringIO(content.decode("utf-8-sig")))
        if len(reader.fieldnames) != len(set(reader.fieldnames)):
            raise RuntimeError("资源 CSV 导出出现同名表头。")
        rows = list(reader)
        section = next(iter(documents))
        return {section: _check_resource_export_table(documents[section], rows, reader.fieldnames, fmt, None)}
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    checked = {}
    try:
        expected_sheets = [titles[section] for section in documents]
        if "resource-days" in documents:
            expected_sheets.append("资源冲突")
        expected_sheets.extend(("覆盖工序", "指标计算方式", "比较范围"))
        if workbook.sheetnames != expected_sheets:
            raise RuntimeError("资源 XLSX 业务表和附表顺序不同：" + str(workbook.sheetnames))
        for section, source in documents.items():
            iterator = workbook[titles[section]].iter_rows()
            headers = [cell.value for cell in next(iterator)]
            if len(headers) != len(set(headers)):
                raise RuntimeError("资源 XLSX 导出出现同名表头：" + titles[section])
            cells = list(iterator)
            rows = [dict(zip(headers, (cell.value for cell in row))) for row in cells]
            types = [dict(zip(headers, (cell.data_type for cell in row))) for row in cells]
            if section in ("resources", "resource-days"):
                checked[section] = _check_resource_export_table(source, rows, headers, fmt, types)
            elif section == "overview":
                count = len(source["data"]["candidates"]) * len(source["data"]["metric_definitions"])
                if len(rows) != count:
                    raise RuntimeError("资源全部表中的总览行数与完整候选目录不同。")
                _check_export_rows(source["data"], rows, fmt)
                checked[section] = {"rows": len(rows), "columns": len(headers)}
            else:
                checked[section] = _check_detail_export_table(source, rows, headers, fmt, types)
        if "resource-days" in documents:
            checked["resource-conflicts"] = _check_resource_record_table(workbook["资源冲突"], RESOURCE_CONFLICT_HEADERS,
                _resource_conflict_expected(documents))
        checked["coverage-operations"] = _check_resource_record_table(workbook["覆盖工序"], RESOURCE_COVERAGE_HEADERS,
            _resource_coverage_expected(documents, archived_labels))
        definitions = documents["overview"]["data"]["metric_definitions"] if "overview" in documents else metric_definitions
        checked["metric-definitions"] = _check_resource_record_table(workbook["指标计算方式"], RESOURCE_METRIC_HEADERS,
            _resource_metric_expected(definitions))
        entries = {row[0]: row[1] if len(row) > 1 else None for row in workbook["比较范围"].iter_rows(values_only=True)}
        scope = active["data"]["scope"]
        expected_scope = {"排产编号": scope["run_ref"], "读取编号": active["meta"]["snapshot_ref"],
            "范围起点（含）": scope["range_start"], "范围终点（不含）": scope["range_end"],
            "资源视角": "本次所选批次" if scope["resource_view"] == "selected_batches" else "方案全部安排",
            "共同设备数": len(scope["resource_refs"]["machine"]), "共同人员数": len(scope["resource_refs"]["operator"]),
            "筛选页签": active["data"].get("filters", {}).get("section"), "明细范围": "筛选结果全部页",
            "资源占用计算方式": "保存安排区间与归档可排窗口的交集；休息时段的保留跨度单列",
            "资源冲突计算方式": "归档可排窗口内同时占用同一资源的不同永久工序",
            "零容量工作证据": "安排跨度不证明实际工作；未保存工作片段的零容量时段留空，班表完整状态单列",
            "周占用比例计算方式": "所选周内班表内占用工时之和除以可排工时之和；日峰值和高负荷日保留",
            "小时筛选范围": "当前显示小时区间；完整比较窗口单列"}
        if any(entries.get(key) != value for key, value in expected_scope.items()):
            raise RuntimeError("资源 XLSX 比较范围、读取身份或资源视角不同。")
        filters = active["data"].get("filters")
        if filters is not None:
            exported_filters = dict(item.split("=", 1) for item in entries["筛选条件"].split("；"))
            if exported_filters != {key: str(value) for key, value in filters.items() if value is not None}:
                raise RuntimeError("资源 XLSX 筛选条件与明细快照不同。")
        checked["sheet_names"] = workbook.sheetnames
    finally:
        workbook.close()
    return checked


def _resource_export_document(client, fixture, query, target, section, cache):
    key = (fixture["run_ref"], section, tuple(sorted(query.items())))
    path = API + "/runs/" + fixture["run_ref"] + "/decision"
    if key not in cache:
        if section == "overview":
            cache[key] = client.document(path + "?" + urlencode(query))
        else:
            cache[key] = _detail_pages(client, path, query, target, section)
    _save(target / (section + "-all.json"), cache[key])
    return cache[key]


def _verify_resource_export_case(client, fixture, query, target, cache, metric_definitions, archived_labels, *, section,
                                 filter_section=None, identity=None):
    if section == "all":
        return current_all_export_case(client, identity, fixture, query, target, cache,
                                       filter_section=filter_section, archived_labels=archived_labels)
    target.mkdir(parents=True, exist_ok=True)
    path = API + "/runs/" + fixture["run_ref"] + "/decision"
    active_section = filter_section or ("overview" if section == "all" else section)
    active = _resource_export_document(client, fixture, query, target, active_section, cache)
    documents = {active_section: [active]}
    if section == "all":
        keys = ("comparison_mode", "reference_candidate_ref", "selected_candidate_ref", "range_start", "range_end", "resource_view")
        common = {key: value for key, value in query.items() if key in keys}
        overview = active if active_section == "overview" else _resource_export_document(client, fixture, common, target, "overview", cache)
        documents = {"overview": overview}
        for name in ("batches", "operations"):
            documents[name] = _resource_export_document(client, fixture, common, target, name, cache)
        if active_section == "resources":
            documents["resources"] = [active]
        elif active_section == "resource-days":
            resource_query = dict(common, resource_kind=query["resource_kind"], resource_ref=query["resource_ref"])
            documents["resources"] = [_resource_export_document(client, fixture, resource_query, target, "resources", cache)]
        else:
            documents["resources"] = [_resource_export_document(client, fixture, dict(common, resource_kind=kind),
                target / kind, "resources", cache) for kind in ("machine", "operator")]
        if active_section == "resource-days":
            documents["resource-days"] = [active]
    expected_count = 0
    for name, source in documents.items():
        for document in (source if isinstance(source, list) else [source]):
            if document["data"]["scope"] != active["data"]["scope"]:
                raise RuntimeError("资源全部导出中的共同窗口、视角或资源集合不同。")
            expected_count += (len(document["data"]["candidates"]) * len(document["data"]["metric_definitions"])
                               if name == "overview" else len(document["data"]["items"]))
    export_query = dict(query, section=section, snapshot_ref=active["meta"]["snapshot_ref"])
    export_query.update({key: value for key, value in active["data"].get("filters", {}).items()
                         if key != "section" and value is not None})
    if filter_section is not None:
        export_query["filter_section"] = filter_section
    _save(target / "request.json", export_query)
    formats, row_counts = {}, {}
    for fmt in (("xlsx",) if section == "all" else ("csv", "xlsx")):
        download_rows = expected_count + (len(_resource_coverage_expected(documents, archived_labels)) + len(_resource_conflict_expected(documents))
                                          if fmt == "xlsx" else 0)
        content, headers = _fetch_export(client, path + "/export?" + urlencode(dict(export_query, format=fmt)),
                                         target, fmt, section=section)
        if (headers.get("X-Workbench-Run-Ref") != fixture["run_ref"]
                or headers.get("X-Workbench-Snapshot-Ref") != export_query["snapshot_ref"]
                or int(headers["X-Workbench-Row-Count"]) != download_rows):
            raise RuntimeError("资源导出的运行、快照身份或业务行数合计不同。")
        formats[fmt] = _check_resource_export_download(content, fmt, documents, active, metric_definitions, archived_labels)
        row_counts[fmt] = download_rows
    result = {"state": "passed", "run_ref": fixture["run_ref"], "comparison_mode": query["comparison_mode"],
              "resource_view": active["data"]["scope"]["resource_view"], "section": section, "filter_section": filter_section,
              "filters": active["data"].get("filters"), "snapshot_ref": export_query["snapshot_ref"],
              "business_row_count": expected_count, "download_row_counts": row_counts, "formats": formats}
    _save(target / "verification.json", result)
    return result, active


def resource_export_verify(args):
    identity, client = _identity(args.instance)
    evidence = ROOT / "evidence" / "decision-resource-impact" / Path(identity["root"]).name
    fixture = json.loads((evidence / "resource-fixture.json").read_text(encoding="utf-8"))
    legacy = json.loads((Path(identity["evidence_directory"]) / "repair-01" / "candidate-count-fixtures.json").read_text(encoding="utf-8"))
    delivery = json.loads((ROOT / "evidence" / "decision-delivery-impact" / Path(identity["root"]).name / "delivery-fixtures.json").read_text(encoding="utf-8"))
    if any(Path(item["database_path"]).resolve() != Path(identity["database_path"]).resolve() for item in (fixture, legacy, delivery)):
        raise RuntimeError("资源导出夹具与当前隔离数据库不同。")
    if fixture["state"] != "generated":
        raise RuntimeError("真实资源夹具尚未完整生成，请先核对资源夹具记录。")
    old = next(item for item in legacy["fixtures"] if item["candidate_count"] == 8)
    mixed = delivery["fixtures"]["mixed_due"]
    candidates = {"comparison_mode": "candidates", "selected_candidate_ref": fixture["candidate_refs"][-1],
                  "reference_candidate_ref": fixture["candidate_refs"][0], "resource_view": "selected_batches"}
    long_days = dict(candidates, resource_view="full_candidate", resource_kind="operator", resource_ref=fixture["operator_refs"][1],
                     range_start="2026-10-13T00:00:00", range_end="2026-11-05T00:00:00")
    old_query = {"comparison_mode": "candidates", "selected_candidate_ref": old["candidate_refs"][1],
                 "reference_candidate_ref": old["candidate_refs"][0], "resource_view": "selected_batches", "resource_kind": "machine"}
    unused = dict(old_query, range_start="2026-10-13T08:00:00", range_end="2026-10-13T09:00:00", sort="label", order="asc")
    zero = dict(old_query, selected_candidate_ref=old["candidate_refs"][0], resource_ref="73e0d4c2bb29e6de6e1d095e6b914932cc021601cc6cb4af",
                range_start="2026-10-17T00:00:00", range_end="2026-10-19T00:00:00", pressure_filter="all")
    formal = {"comparison_mode": "admission_official", "selected_candidate_ref": mixed["candidate_refs"][-1],
              "resource_view": "full_candidate", "resource_kind": "machine", "resource_ref": "73e0d4c2bb29e6de6e1d095e6b914932cc021601cc6cb4af"}
    cases = [
        ("new-eleven-machines", fixture, dict(candidates, resource_kind="machine"), "resources", None),
        ("new-operator-days", fixture, long_days, "resource-days", None),
        ("old8-unused-positive", old, unused, "resources", None),
        ("old8-zero-capacity", old, zero, "resources", None),
        ("all-candidates-six-tables", fixture, candidates, "all", None),
        ("all-formal-full-filter", mixed, formal, "all", "resources"),
        ("all-operator-seven-tables", fixture, long_days, "all", "resource-days"),
    ]
    if args.case:
        cases = [case for case in cases if case[0] in args.case]
    target = Path(args.evidence_directory).resolve() if args.evidence_directory else evidence / "exports"
    cache, results, current = {}, [], None
    catalog = _resource_export_document(client, fixture, candidates, target / "catalog", "overview", cache)
    catalog_path = target / "catalog" / "overview-all.json"
    metric_definitions = catalog["data"]["metric_definitions"]
    _save(target / "metric-source.json", {"source": str(catalog_path), "metric_definitions": metric_definitions})
    labels_by_run = {}
    for source in (case[1] for case in cases):
        if source["run_ref"] not in labels_by_run:
            labels_by_run[source["run_ref"]] = _resource_archive_labels(identity, source["run_ref"])
            _save(target / ("coverage-labels-" + source["run_ref"] + ".json"), {"source": "WorkbenchRunJobs.facts_json",
                "run_ref": source["run_ref"], "operation_labels": labels_by_run[source["run_ref"]]})
    try:
        for current, source, query, section, filter_section in cases:
            result, document = _verify_resource_export_case(client, source, query, target / current, cache, metric_definitions,
                                                            labels_by_run[source["run_ref"]],
                                                            section=section, filter_section=filter_section, identity=identity)
            data = document["data"]
            if current == "new-eleven-machines" and (len(data["items"]) != 11 or len(data["rankings"]["top_pressure"]) != 10):
                raise RuntimeError("资源导出夹具的完整11资源与Top10排行数量不同。")
            if current == "new-operator-days" and len(data["items"]) != 23:
                raise RuntimeError("单人员逐日导出没有覆盖23天的完整窗口。")
            if current == "old8-unused-positive":
                row = next(item for item in data["items"] if item["resource_ref"] == "a004303dbdad560c9b7e6a711d77a5c530ea86acdab16f78")
                if any(row[side]["available_hours"] != 1 or row[side]["occupied_hours"] != 0 or row[side]["utilization_ratio"] != 0
                       for side in ("before", "after")):
                    raise RuntimeError("旧8候选的正容量未使用资源没有保留0占用率。")
            if current == "old8-zero-capacity":
                if len(data["items"]) != 1 or any(data["items"][0][side]["available_hours"] != 0
                    or data["items"][0][side]["utilization_ratio"] is not None or data["items"][0][side]["zero_capacity_task"] is not None
                    or data["items"][0][side]["actual_work_state"] != "not_recorded"
                    or data["items"][0][side]["zero_capacity_reason"]["code"] != "actual_work_fragments_not_recorded"
                    for side in ("before", "after")):
                    raise RuntimeError("跨周末的保留安排没有分别保留零容量、空比例和实际工作片段未知。")
            if current == "all-formal-full-filter":
                row = data["items"][0]
                if (len(data["items"]) != 1 or row["before"]["span_occupied_hours"] <= 0 or row["after"]["span_occupied_hours"] != 0
                        or data["coverage"]["after"]["covered_formal_operation_count"] != 0):
                    raise RuntimeError("完整候选导出叠加了候选未保存的正式安排。")
            results.append(dict(result, case=current))
    except Exception as error:
        _save(target / "export-verification.json", {"state": "failed", "case": current, "error": str(error), "cases": results,
            "base_url": identity["base_url"], "pid": identity["pid"]})
        raise
    result = {"state": "passed", "base_url": identity["base_url"], "pid": identity["pid"],
              "database_path": identity["database_path"], "case_count": len(results), "cases": results,
              "business_row_count": sum(sum(item["download_row_counts"].values()) for item in results),
              "evidence_directory": str(target)}
    _save(target / "export-verification.json", result)
    print(json.dumps({key: value for key, value in result.items() if key != "cases"}, ensure_ascii=False), flush=True)


def preparation_export_path(value, path):
    if value is None:
        return None
    if not path:
        return value
    first, _, rest = path.partition(".")
    if isinstance(value, list):
        return [preparation_export_path(item, path) for item in value]
    return preparation_export_path(value.get(first), rest)


def preparation_export_expected(data, item, section):
    fields = PREPARATION_EXPORT_FIELDS if section == "preparation" else CHANGEOVER_EXPORT_FIELDS
    expected = {}
    for header, path in fields.items():
        value = (preparation_export_path(data, path[1:]) if path and path.startswith("$")
                 else preparation_export_path(item, path) if path else None)
        if header == "查看基准永久编号":
            value = value["candidate_ref"] if value["kind"] == "candidate" else value["baseline_ref"]
        elif header.endswith("缺项") or header == "缺少资料":
            value = _export_issue(value)
        elif path and ".machines." in path:
            value = "、".join(machine or "" for machine in value) or None
        elif header.endswith("安排记录"):
            value = value or None
        elif type(value) is bool:
            value = "是" if value else "否"
        elif header == "计算方式":
            value = "同一设备上相邻两道工序的工种不同，后一道工序开始时间在所选范围内（含开始、不含结束）"
        elif header == "事件指标版本":
            value = "type_changeover_v1"
        expected[header] = value
    return expected


def preparation_export_check_table(document, rows, headers, fmt, types):
    data = document["data"]
    section = data["filters"]["section"]
    fields = PREPARATION_EXPORT_FIELDS if section == "preparation" else CHANGEOVER_EXPORT_FIELDS
    if tuple(headers) != tuple(fields) or len(headers) != len(set(headers)) or len(rows) != len(data["items"]):
        raise RuntimeError("准备或切换导出列及全部页行数与独立契约不同：" + section)
    checked = numeric = blanks = json_cells = 0
    for index, (item, row) in enumerate(zip(data["items"], rows)):
        for header, value in preparation_export_expected(data, item, section).items():
            cell_type = types[index].get(header) if types is not None else None
            if isinstance(value, list):
                actual = row.get(header)
                equal = isinstance(actual, str) and json.loads(actual) == value and (fmt == "csv" or cell_type == "s")
                json_cells += 1
            else:
                equal = _check_detail_export_cell(row.get(header), value, fmt, header, cell_type)
            if not equal:
                raise RuntimeError(f"准备或切换导出字段不同：{section} / 行{index + 1} / {header} / "
                                   f"实际 {row.get(header)!r} / 期待 {value!r}")
            checked += 1
            numeric += type(value) in (int, float)
            blanks += value is None
    return {"rows": len(rows), "columns": len(fields), "checked_cells": checked, "numeric_cells": numeric,
            "blank_cells": blanks, "json_cells": json_cells, "identity_and_order_matches_api": True}


def preparation_export_check_operations(document, rows, headers, fmt, types):
    original = DETAIL_EXPORT_COMMON + tuple(OPERATION_EXPORT_FIELDS)
    extra = {prefix + "所需工时：" + header: "runtime_effort." + side + "." + path
             for prefix, side in (("基准", "before"), ("所选", "after"))
             for header, path in PREPARATION_EXPORT_EFFORT_FIELDS.items()}
    if tuple(headers) != original + tuple(extra) or len(headers) != len(set(headers)):
        raise RuntimeError("工序导出的旧字段与运行需求附加字段不同或重名。")
    result = _check_detail_export_table(document, rows, original, fmt, types)
    blanks = 0
    for index, (item, row) in enumerate(zip(document["data"]["items"], rows)):
        for header, path in extra.items():
            value = preparation_export_path(item, path)
            if header.endswith("缺项"):
                value = _export_issue(value)
            cell_type = types[index].get(header) if types is not None else None
            if not _check_detail_export_cell(row.get(header), value, fmt, header, cell_type):
                raise RuntimeError(f"工序导出运行需求字段不同：行{index + 1} / {header} / "
                                   f"实际 {row.get(header)!r} / 期待 {value!r}")
            result["checked_cells"] += 1
            result["numeric_cells"] += type(value) in (int, float)
            blanks += value is None
    result.update(columns=len(headers), runtime_columns=len(extra), runtime_blank_cells=blanks)
    return result


def preparation_export_labels():
    labels = {key: dict(value) for key, value in RESOURCE_METRIC_LABELS.items()}
    labels["group"]["preparation"] = "准备与切换"
    labels["scope"]["fixed_changeover_window"] = "后一道工序在所选范围内开始的设备工种切换"
    labels["source"].update(runtime_operation_effort="计算时保存的准备和加工所需标准工时",
                            archived_changeover_events="同一设备上相邻工序的工种变化")
    labels["state"] = {"complete": "完整", "partial": "部分", "unavailable": "缺少资料", "not_applicable": "范围无适用对象",
        "completed": "计算完成", "failed": "失败", "skipped": "跳过", "queued": "待计算", "running": "计算中",
        "unknown": "缺少资料", "comparable": "可比较", "different_coverage": "覆盖范围不同",
        "better": "改善", "worse": "增加代价", "equal": "相同", "neutral": "中性变化",
        "available": "已保存", "absent": "排产时无正式计划", "selection_required": "需要选择基准"}
    return labels


def preparation_export_metric_expected(definitions):
    return _resource_metric_expected(definitions)


def preparation_export_overview_expected(data):
    headers = ("排产编号", "归档编号", "查看基准编号", "查看基准名称", "所选方案编号", "推荐方案编号",
        "范围起点（含）", "范围终点（不含）", "候选编号", "候选名称", "搜索状态", "应排范围状态", "交付覆盖状态",
        "指标", "单位", "指标版本", "指标范围", "指标计算方式", "覆盖状态", "数值", "已知小计", "已知数量", "未知数量",
        "总数量", "基准数值", "差值（候选减基准）", "差值状态", "变化方向", "缺项说明")
    scope, baseline = data["scope"], data["baseline"]
    baseline_ref = baseline["candidate_ref"] if baseline["kind"] == "candidate" else baseline["baseline_ref"]
    labels = preparation_export_labels()
    common = [data["run"]["run_ref"], data["capture_ref"], baseline_ref, baseline.get("label"),
        scope["selected_candidate_ref"], data["recommendation"]["candidate_ref"], scope["range_start"], scope["range_end"]]
    rows = []
    for candidate in data["candidates"]:
        identity = [candidate["candidate_ref"], candidate["label"], labels["state"][candidate["status"]],
            labels["state"][candidate["coverage"]["schedulable"]["state"]], labels["state"][candidate["coverage"]["delivery"]["state"]]]
        for definition in data["metric_definitions"]:
            metric, delta = candidate["metrics"][definition["key"]], candidate["deltas"][definition["key"]]
            numbers = [metric["value"], metric["known_subtotal"], delta["reference_value"], delta["value"]]
            if definition["unit"] == "%":
                numbers = [value * 100 if value is not None else None for value in numbers]
            value, subtotal, reference, difference = numbers
            values = common + identity + [definition["label"], definition["unit"], metric["metric_version"],
                labels["scope"][metric["scope"]], labels["source"][definition["source"]], labels["state"][metric["state"]],
                value, subtotal, metric["known_count"], metric["unknown_count"], metric["total_count"], reference,
                difference, labels["state"][delta["state"]], labels["state"][delta["direction"]], (metric["reason"] or {}).get("message")]
            rows.append(dict(zip(headers, values)))
    return headers, rows


def preparation_export_check_records(sheet, headers, expected):
    iterator = sheet.iter_rows()
    if tuple(cell.value for cell in next(iterator)) != headers:
        raise RuntimeError("准备与切换导出附表列与独立契约不同：" + sheet.title)
    actual = list(iterator)
    if len(actual) != len(expected):
        raise RuntimeError("准备与切换导出附表行数与完整 API 清单不同：" + sheet.title)
    checked = numeric = blanks = longest = 0
    for index, (row, wanted) in enumerate(zip(actual, expected)):
        for column, header in enumerate(headers):
            cell = row[column] if column < len(row) else None
            value, kind = (cell.value, cell.data_type) if cell is not None else (None, None)
            target = wanted[header]
            if not _check_detail_export_cell(value, target, "xlsx", header, kind):
                raise RuntimeError(f"准备与切换导出附表字段不同：{sheet.title} / 行{index + 1} / {header} / "
                                   f"实际 {value!r} / 期待 {target!r}")
            checked += 1
            numeric += type(target) in (int, float)
            blanks += target is None
            longest = max(longest, len(target) if isinstance(target, str) else 0)
    return {"rows": len(actual), "columns": len(headers), "checked_cells": checked, "numeric_cells": numeric,
            "blank_cells": blanks, "max_text_cell_length": longest, "identity_and_order_matches_api": True}


def preparation_export_check_download(content, fmt, documents, active, definitions, archived_labels):
    titles = {"overview": "方案总览", "batches": "批次影响", "operations": "工序变化", "resources": "资源影响",
              "resource-days": "资源逐日", "preparation": "准备需求", "changeovers": "工种切换"}
    if fmt == "csv":
        section = active["data"]["filters"]["section"]
        reader = csv.DictReader(io.StringIO(content.decode("utf-8-sig")))
        return {section: preparation_export_check_table(active, list(reader), reader.fieldnames, fmt, None)}
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    checked = {}
    try:
        expected_sheets = []
        for section in documents:
            expected_sheets.append(titles[section])
            if section == "resources" and "resource-days" not in documents:
                expected_sheets.append("覆盖工序")
            elif section == "resource-days":
                expected_sheets.extend(("资源冲突", "覆盖工序"))
        expected_sheets.extend(("指标计算方式", "比较范围"))
        if workbook.sheetnames != expected_sheets:
            raise RuntimeError("准备与切换 XLSX 表单顺序与独立契约不同：" + str(workbook.sheetnames))
        for section, document in documents.items():
            sheet = workbook[titles[section]]
            if section == "overview":
                headers, expected = preparation_export_overview_expected(document["data"])
                checked[section] = preparation_export_check_records(sheet, headers, expected)
                continue
            iterator = sheet.iter_rows()
            headers = [cell.value for cell in next(iterator)]
            cells = list(iterator)
            rows = [dict(zip(headers, (cell.value for cell in row))) for row in cells]
            types = [dict(zip(headers, (cell.data_type for cell in row))) for row in cells]
            if section in ("preparation", "changeovers"):
                checked[section] = preparation_export_check_table(document, rows, headers, fmt, types)
            elif section == "operations":
                checked[section] = preparation_export_check_operations(document, rows, headers, fmt, types)
            elif section in ("resources", "resource-days"):
                checked[section] = _check_resource_export_table(document, rows, headers, fmt, types)
            else:
                checked[section] = _check_detail_export_table(document, rows, headers, fmt, types)
        if "resources" in documents:
            checked["coverage-operations"] = preparation_export_check_records(workbook["覆盖工序"], RESOURCE_COVERAGE_HEADERS,
                _resource_coverage_expected(documents, archived_labels))
        if "resource-days" in documents:
            checked["resource-conflicts"] = preparation_export_check_records(workbook["资源冲突"], RESOURCE_CONFLICT_HEADERS,
                _resource_conflict_expected(documents))
        checked["metric-definitions"] = preparation_export_check_records(workbook["指标计算方式"], RESOURCE_METRIC_HEADERS,
            preparation_export_metric_expected(definitions))
        entries = {row[0]: row[1] if len(row) > 1 else None for row in workbook["比较范围"].iter_rows(values_only=True)}
        data, scope = active["data"], active["data"]["scope"]
        baseline = data["baseline"]
        wanted = {"排产编号": data["run_ref"], "读取编号": active["meta"]["snapshot_ref"],
            "查看基准编号": baseline["candidate_ref"] if baseline["kind"] == "candidate" else baseline["baseline_ref"],
            "所选方案编号": data["candidate"]["candidate_ref"], "范围起点（含）": scope["range_start"],
            "范围终点（不含）": scope["range_end"], "批次数": len(scope["batch_refs"]), "工序数": len(scope["operation_refs"]),
            "资源视角": "本次所选批次" if scope["resource_view"] == "selected_batches" else "方案全部安排",
            "共同设备数": len(scope["resource_refs"]["machine"]), "共同人员数": len(scope["resource_refs"]["operator"]),
            "筛选页签": data["filters"]["section"], "明细范围": "筛选结果全部页",
            "标准工时计算方式": "保存的工时定额乘以其对应的加工目标数量",
            "所需工时计算方式": "计算时保存的准备和加工所需标准工时；原加工目标与本次报工目标数量分别保留",
            "工种切换计算方式": "同一设备上相邻工序的工种不同，后一道工序开始时间在所选范围内（含开始、不含结束）",
            "不计工种切换的工序": "外协工序，以及保存记录中明确的零工时工序",
            "准备工时统计范围": "本次所选批次；新安排、其他未安排、缺料未安排分别汇总；原安排、已开工和已完工工序分别列出",
            "准备差值方向": "候选减基准，中性展示实际需求与安排覆盖变化"}
        if any(entries.get(key) != value for key, value in wanted.items()):
            raise RuntimeError("准备与切换 XLSX 的共同范围、来源或运行身份不同。")
        filters = dict(item.split("=", 1) for item in entries["筛选条件"].split("；"))
        if filters != {key: str(value) for key, value in data["filters"].items() if value is not None}:
            raise RuntimeError("准备与切换 XLSX 筛选条件与当前页签快照不同。")
        checked["sheet_names"] = workbook.sheetnames
    finally:
        workbook.close()
    return checked


def preparation_export_document(client, fixture, query, target, section, cache):
    key = (fixture["run_ref"], section, tuple(sorted(query.items())))
    if key not in cache:
        path = API + "/runs/" + fixture["run_ref"] + "/decision"
        cache[key] = (client.document(path + "?" + urlencode(query)) if section == "overview"
                      else _detail_pages(client, path, query, target, section))
    _save(target / (section + "-all.json"), cache[key])
    return cache[key]


def preparation_export_case(client, fixture, query, target, cache, definitions, archived_labels, *, section,
                            filter_section=None, identity=None):
    if section == "all":
        result, _ = current_all_export_case(client, identity, fixture, query, target, cache,
            filter_section=filter_section, archived_labels=archived_labels)
        return result
    target.mkdir(parents=True, exist_ok=True)
    active_section = filter_section or section
    active = preparation_export_document(client, fixture, query, target, active_section, cache)
    documents = {active_section: active}
    if section == "all":
        keys = ("comparison_mode", "reference_candidate_ref", "selected_candidate_ref", "range_start", "range_end", "resource_view")
        common = {key: value for key, value in query.items() if key in keys}
        documents = {"overview": preparation_export_document(client, fixture, common, target, "overview", cache)}
        for name in ("batches", "operations"):
            documents[name] = preparation_export_document(client, fixture, common, target, name, cache)
        documents["resources"] = [preparation_export_document(client, fixture, dict(common, resource_kind=kind),
            target / kind, "resources", cache) for kind in ("machine", "operator")]
        for name in ("preparation", "changeovers"):
            other = dict(common, **({"machine_ref": query["machine_ref"]} if query.get("machine_ref") else {}))
            documents[name] = active if name == active_section else preparation_export_document(client, fixture, other, target, name, cache)
    count = 0
    for name, source in documents.items():
        for document in source if isinstance(source, list) else [source]:
            if document["data"]["scope"] != active["data"]["scope"]:
                raise RuntimeError("准备与切换全部表中的共同范围不同。")
            count += (len(document["data"]["candidates"]) * len(document["data"]["metric_definitions"])
                      if name == "overview" else len(document["data"]["items"]))
    if "resources" in documents:
        count += len(_resource_coverage_expected(documents, archived_labels))
    export_query = dict(query, section=section, snapshot_ref=active["meta"]["snapshot_ref"])
    if filter_section:
        export_query["filter_section"] = filter_section
    _save(target / "request.json", export_query)
    formats = {}
    path = API + "/runs/" + fixture["run_ref"] + "/decision/export"
    for fmt in (("xlsx",) if section == "all" else ("csv", "xlsx")):
        content, headers = _fetch_export(client, path + "?" + urlencode(dict(export_query, format=fmt)), target, fmt, section=section)
        if (headers.get("X-Workbench-Run-Ref") != fixture["run_ref"]
                or headers.get("X-Workbench-Snapshot-Ref") != export_query["snapshot_ref"]
                or int(headers["X-Workbench-Row-Count"]) != count):
            raise RuntimeError("准备与切换导出的运行、快照身份或业务总行数不同。")
        formats[fmt] = preparation_export_check_download(content, fmt, documents, active, definitions, archived_labels)
    result = {"state": "passed", "run_ref": fixture["run_ref"], "section": section, "filter_section": filter_section,
        "filters": active["data"]["filters"], "scope": active["data"]["scope"], "snapshot_ref": export_query["snapshot_ref"],
        "business_row_count": count, "formats": formats, "pages_20_50_same_snapshot": True,
        "all_sections_scope_same": True, "active_filter_preserved": True}
    _save(target / "verification.json", result)
    return result


def preparation_export_verify(args):
    identity, client = _identity(args.instance)
    evidence = ROOT / "evidence" / "decision-preparation-events" / Path(identity["root"]).name
    fixture_path = (Path(args.fixture).resolve() if args.fixture else evidence / "preparation-complete-formal-scope-fixture.json")
    fixture = json.loads(fixture_path.read_text(encoding="utf-8"))
    if fixture["state"] != "generated" or Path(fixture["database_path"]).resolve() != Path(identity["database_path"]).resolve():
        raise RuntimeError("准备与切换导出所需的真实隔离夹具尚未完成或数据库不同。")
    target = Path(args.evidence_directory).resolve() if args.evidence_directory else evidence / "exports"
    query = {"comparison_mode": "candidates", "selected_candidate_ref": fixture["candidate_refs"][-1],
        "reference_candidate_ref": fixture["candidate_refs"][0], "resource_view": "selected_batches",
        "range_start": "2026-10-13T00:00:00", "range_end": "2026-10-17T00:00:00"}
    cache = {}
    overview = preparation_export_document(client, fixture, query, target / "query-source", "overview", cache)
    definitions = overview["data"]["metric_definitions"]
    if len(definitions) != 52 or any(candidate["metrics"]["scheduled_setup_hours"]["metric_version"] != "decision-preparation-v4"
                                      for candidate in overview["data"]["candidates"]):
        raise RuntimeError("准备与切换的共同指标目录或版本不同。")
    events = preparation_export_document(client, fixture, query, target / "query-source", "changeovers", cache)
    machine = fixture["machine_refs"][0]
    selected_events = sorted((row for row in events["data"]["items"] if row["event_side"] == "selected"
        and row["machine_ref"] == machine), key=lambda row: row["next_start"])
    if len(selected_events) < 2:
        raise RuntimeError("真实同设备切换样例缺少两次切换，无法验收窗口前驱与终点排除。")
    window = dict(query, machine_ref=machine, range_start=selected_events[0]["next_start"], range_end=selected_events[1]["next_start"])
    preparation = preparation_export_document(client, fixture, dict(query, sort="sequence", order="asc"),
        target / "new-runtime-preparation", "preparation", cache)
    if {row["operation_ref"] for row in preparation["data"]["items"]} != set(fixture["operation_refs"]):
        raise RuntimeError("准备导出没有覆盖批准的完整受理工序集合。")
    piece = next(row for row in preparation["data"]["items"] if row["piece_id"] is not None)
    cases = [
        ("new-runtime-preparation", dict(query, sort="sequence", order="asc"), "preparation", None),
        ("same-machine-window-changeovers", dict(window, event_side="all", sort="next_start", order="asc"), "changeovers", None),
        ("all-preparation-batch-filter", dict(query, batch_ref=piece["batch_ref"], allocation_state="all",
            sort="after_required_hours", order="desc"), "all", "preparation"),
        ("all-changeovers-selected-filter", dict(window, event_side="selected", sort="next_start", order="desc"), "all", "changeovers"),
    ]
    if args.case:
        cases = [case for case in cases if case[0] in args.case]
    archived_labels = _resource_archive_labels(identity, fixture["run_ref"])
    _save(target / "coverage-labels.json", archived_labels)
    results, current = [], None
    try:
        for current, current_query, section, filter_section in cases:
            results.append(dict(preparation_export_case(client, fixture, current_query, target / current, cache,
                definitions, archived_labels, section=section, filter_section=filter_section, identity=identity), case=current))
    except Exception as error:
        _save(target / "export-verification.json", {"state": "failed", "case": current, "error": str(error), "cases": results})
        raise
    legacy_path = evidence / "verification" / "legacy-multipage" / "preparation-all.json"
    legacy = {"source": str(legacy_path), "state": "not_saved"}
    if legacy_path.exists():
        document = json.loads(legacy_path.read_text(encoding="utf-8"))
        for item in document["data"]["items"]:
            for side in ("before", "after"):
                if any(item[side][field] is not None for field in
                       ("setup_standard_hours", "processing_standard_hours", "required_standard_hours")):
                    raise RuntimeError("复用的旧归档准备需求意外具有新 runtime 数值。")
                if item[side]["standard"]["standard_total_hours"] is None:
                    raise RuntimeError("复用的旧归档已知工序标准没有独立保留。")
        legacy.update(state="passed", rows=len(document["data"]["items"]), runtime_hours_remain_unknown=True,
                      archived_standard_remains_known=True, new_http_requests=False)
        _save(target / "legacy-source.json", document)
    result = {"state": "passed", "base_url": identity["base_url"], "pid": identity["pid"], "database_path": identity["database_path"],
        "run_ref": fixture["run_ref"], "fixture_source": str(fixture_path), "case_count": len(results),
        "download_count": sum(len(item["formats"]) for item in results),
        "business_row_count": sum(item["business_row_count"] * len(item["formats"]) for item in results),
        "metric_count": len(definitions), "preparation_columns": len(PREPARATION_EXPORT_FIELDS),
        "changeover_columns": len(CHANGEOVER_EXPORT_FIELDS), "legacy_api_reuse": legacy, "cases": results,
        "evidence_directory": str(target), "actual_engine_rerun": False, "adoption_performed": False}
    for name in ("checked_cells", "numeric_cells", "json_cells"):
        result[name] = sum(table.get(name, 0) for case in results for tables in case["formats"].values()
                           for table in tables.values() if isinstance(table, dict))
    _save(target / "export-verification.json", result)
    print(json.dumps({key: value for key, value in result.items() if key != "cases"}, ensure_ascii=False), flush=True)


def preparation_export_resource_days_verify(args):
    identity, client = _identity(args.instance)
    evidence = ROOT / "evidence" / "decision-preparation-events" / Path(identity["root"]).name
    fixture_path = Path(args.fixture).resolve() if args.fixture else evidence / "preparation-complete-formal-scope-fixture.json"
    fixture = json.loads(fixture_path.read_text(encoding="utf-8"))
    target = Path(args.evidence_directory).resolve() if args.evidence_directory else evidence / "exports"
    report_path = target / "export-verification.json"
    previous_bytes = report_path.read_bytes()
    previous = json.loads(previous_bytes)
    if (fixture["state"] != "generated" or previous["state"] != "passed" or previous["run_ref"] != fixture["run_ref"]
            or Path(fixture["database_path"]).resolve() != Path(identity["database_path"]).resolve()
            or Path(previous["database_path"]).resolve() != Path(identity["database_path"]).resolve()):
        raise RuntimeError("单人员逐日补证需要同一已完成夹具和四案例原始导出结果。")
    case_name = "all-operator-days-eleven-tables"
    current_target = target / case_name
    current_target.mkdir(parents=True, exist_ok=True)
    source_paths = {"overview": target / "all-preparation-batch-filter" / "overview-all.json",
        "batches": target / "all-preparation-batch-filter" / "batches-all.json",
        "operations": target / "all-preparation-batch-filter" / "operations-all.json",
        "preparation": target / "new-runtime-preparation" / "preparation-all.json",
        "changeovers": target / "all-preparation-batch-filter" / "changeovers-all.json"}
    saved = {name: json.loads(path.read_text(encoding="utf-8")) for name, path in source_paths.items()}
    normal = next(row["operation_ref"] for row in fixture["expected_standard"]
                  if row["batch_id"] == "DP-NORMAL" and row["sequence"] == 1)
    normal_row = next(row for row in saved["preparation"]["data"]["items"] if row["operation_ref"] == normal)
    operators = {(row["operator"] or {}).get("ref") for row in normal_row["after"]["arrangements"]}
    if len(operators) != 1 or None in operators:
        raise RuntimeError("DP-NORMAL 首道原安排没有唯一的 DP-O1 人员身份。")
    resource_ref = operators.pop()
    query = {"comparison_mode": "candidates", "selected_candidate_ref": fixture["candidate_refs"][-1],
        "reference_candidate_ref": fixture["candidate_refs"][0], "resource_view": "selected_batches",
        "range_start": "2026-10-13T00:00:00", "range_end": "2026-10-17T00:00:00",
        "resource_kind": "operator", "resource_ref": resource_ref}
    cache = {}
    active = preparation_export_document(client, fixture, query, current_target, "resource-days", cache)
    resource = preparation_export_document(client, fixture, query, current_target, "resources", cache)
    if (active["data"]["resource"]["ref"] != resource_ref or len(active["data"]["items"]) != 4
            or len(resource["data"]["items"]) != 1 or resource["data"]["items"][0]["resource_ref"] != resource_ref):
        raise RuntimeError("新增单人员逐日导出未保持 DP-O1 的单资源与四日期筛选。")
    documents = {"overview": saved["overview"], "batches": saved["batches"], "operations": saved["operations"],
                 "resources": [resource], "resource-days": [active], "preparation": saved["preparation"],
                 "changeovers": saved["changeovers"]}
    for name, sources in documents.items():
        for document in sources if isinstance(sources, list) else [sources]:
            if document["data"]["scope"] != active["data"]["scope"]:
                raise RuntimeError("新增单人员逐日导出与原共同受理范围不同：" + name)
    for name, document in saved.items():
        _save(current_target / (name + "-all.json"), document)
    _save(current_target / "reused-api-sources.json", {name: str(path) for name, path in source_paths.items()})
    labels = json.loads((target / "coverage-labels.json").read_text(encoding="utf-8"))
    current, _ = current_all_export_case(client, identity, fixture, query, current_target, cache,
        filter_section="resource-days", active=active, documents=documents, archived_labels=labels)
    current.update(case=case_name, original_cases_reused=True)
    count, checked = current["business_row_count"], current["formats"]["xlsx"]
    _save(current_target / "verification.json", current)
    archive_path = target / "export-verification-base-four.json"
    if not archive_path.exists():
        archive_path.write_bytes(previous_bytes)
    original_identity = {"base_url": previous["base_url"], "pid": previous["pid"]}
    cases = [row for row in previous["cases"] if row["case"] != case_name]
    for row in cases:
        row.setdefault("runtime_identity", original_identity)
    cases.append(current)
    previous.update(cases=cases, case_count=len(cases), download_count=sum(len(row["formats"]) for row in cases),
        business_row_count=sum(row["business_row_count"] * len(row["formats"]) for row in cases),
        final_read_instance=current["runtime_identity"], original_four_case_report=str(archive_path),
        latest_incremental_case=case_name, incremental_download_count=1)
    for name in ("checked_cells", "numeric_cells", "json_cells"):
        previous[name] = sum(table.get(name, 0) for row in cases for tables in row["formats"].values()
                             for table in tables.values() if isinstance(table, dict))
    previous["json_arrangement_cells"] = previous["json_cells"]
    _save(report_path, previous)
    print(json.dumps({"state": "passed", "case": case_name, "base_url": identity["base_url"], "pid": identity["pid"],
        "new_download_count": 1, "new_business_row_count": count, "sheet_names": checked["sheet_names"],
        "combined_case_count": len(cases), "combined_download_count": previous["download_count"],
        "combined_business_row_count": previous["business_row_count"], "evidence_directory": str(current_target)}, ensure_ascii=False), flush=True)


def recommendation_export_common(document):
    data, meta = document["data"], document["meta"]
    scope, baseline = data["scope"], data["baseline"]
    return dict(zip(RECOMMENDATION_EXPORT_COMMON, [data["run"]["run_ref"], data["capture_ref"], meta["snapshot_ref"], meta["as_of"],
        data["recommendation"]["candidate_ref"], scope["selected_candidate_ref"], scope["comparison_mode"],
        baseline["candidate_ref"] if baseline["kind"] == "candidate" else baseline["baseline_ref"], baseline["kind"],
        scope["range_start"], scope["range_end"], scope["resource_view"], data["recommendation"]["recommendation_version"],
        data["selection_basis"]["selection_policy_version"]]))


def recommendation_export_scalar_leaves(value, path):
    if isinstance(value, dict):
        result = {}
        for key in sorted(value):
            result.update(recommendation_export_scalar_leaves(value[key], path + "." + key))
        return result
    if isinstance(value, list):
        result = {}
        for index, item in enumerate(value):
            result.update(recommendation_export_scalar_leaves(item, path + "[" + str(index) + "]"))
        return result
    return {path: value}


def recommendation_export_definition(entry):
    values = {"指标标识": entry.get("metric_key") or entry.get("key"), "名称": entry["label"], "范围": entry["scope"],
        "来源": entry["source"], "指标版本": entry.get("metric_version"), "状态": entry["state"], "变化方向": entry.get("direction"), "单位": entry.get("unit"),
        "说明": entry["text"]}
    for header, key in (("基准值", "reference_value"), ("比较对象值", "selected_value"), ("差值", "value")):
        value = entry.get(key)
        values[header] = value * 100 if type(value) in (int, float) and entry.get("unit") == "%" else value
    return values


def recommendation_export_explanations(document):
    data = document["data"]
    recommendation, basis = data["recommendation"], data["selection_basis"]
    rows = [{"record_kind": "conclusion", "指标标识": "recommendation_type", "名称": "推荐结论", "范围": "full_admission",
             "状态": recommendation["state"], "比较对象值": recommendation["recommendation_type"], "说明": recommendation["headline"]}]
    for field in RECOMMENDATION_EXPORT_FIXED_FIELDS:
        for path, value in recommendation_export_scalar_leaves(basis[field], field).items():
            rows.append({"record_kind": "fixed_selection", "指标标识": path, "名称": field, "范围": "fixed_algorithm_selection",
                "来源": "saved_selection", "状态": "unavailable" if value is None else "complete", "比较对象值": value})
    for entry in basis["explanation"]["entries"]:
        rows.append({"record_kind": "selection_explanation", **recommendation_export_definition(entry)})
    for issue in basis["explanation"]["issues"]:
        rows.append({"record_kind": "missing_evidence", "指标标识": issue["code"], "名称": "推荐依据缺少资料",
                     "状态": "unavailable", "说明": issue["message"]})
    main = {row["metric_key"]: (kind, index + 1) for kind in ("benefits", "costs")
            for index, row in enumerate(recommendation["current_comparison"][kind])}
    for entry in recommendation["current_comparison"]["all_differences"]:
        category, order = main.get(entry["metric_key"], (None, None))
        rows.append({"record_kind": "difference", **recommendation_export_definition(entry), "主项类别": category, "主项顺序": order})
        for support in entry["supporting_refs"]:
            drill = support["drill"]
            record = {"record_kind": "support", "指标标识": entry["metric_key"], "名称": entry["label"],
                "范围": entry["scope"], "来源": entry["source"], "指标版本": entry["metric_version"], "相关明细类型": support["kind"], "相关明细永久编号": support["ref"],
                "相关明细候选编号": drill["candidate_ref"], "相关明细页签": drill.get("comparison_tab", "gantt"),
                "相关明细比较起点": drill["range_start"], "相关明细比较终点": drill["range_end"]}
            query = {"side": support.get("side", "selected"), **(drill.get("detail_query") or {key: drill[key] for key in
                ("batch_ref", "operation_ref", "selected_row_ref", "resource_ref", "resource_kind", "reference_view", "plan_ref") if key in drill})}
            if not query:
                rows.append(record)
            for field, value in query.items():
                rows.append({**record, "相关明细筛选字段": field, "相关明细筛选值": value})
    return rows


def recommendation_export_dispositions(identity, run_ref, source_path=None):
    if source_path is not None:
        document = json.loads(Path(source_path).read_text(encoding="utf-8"))
        if document["run_ref"] != run_ref:
            raise RuntimeError("推荐导出的原收据来源与排产编号不同。")
        receipt = document["result"]
        source = str(source_path)
    else:
        import sqlite3

        conn = sqlite3.connect(Path(identity["database_path"]).as_uri() + "?mode=ro", uri=True)
        try:
            conn.execute("PRAGMA query_only=ON")
            row = conn.execute("SELECT result_json FROM WorkbenchRunReceipts WHERE run_ref=?", (run_ref,)).fetchone()
            if row is None:
                raise RuntimeError("推荐导出所需的原始受理收据缺失。")
            receipt = json.loads(row[0])
        finally:
            conn.close()
        source = "WorkbenchRunReceipts.result_json (SQLite mode=ro)"
    return {"source": source, "run_ref": run_ref, "fields": ["status", "issues"],
        "operations": {row["operation_ref"]: {"status": row.get("status"), "issues": row.get("issues", [])}
                       for row in receipt.get("dispositions", [])}}


def recommendation_export_expected(document, candidates, dispositions):
    data, common = document["data"], recommendation_export_common(document)
    groups = {ref: group["group_index"] for group in data["completion_groups"] for ref in group["candidate_refs"]}
    complete = {ref: {row["batch_ref"] for row in facts["batches"]["data"]["items"]
                     if facts["available"] and row["after"]["completion_state"] == "complete"}
                for ref, facts in candidates.items()}
    fixed, reference = data["recommendation"]["candidate_ref"], data["baseline"].get("candidate_ref")
    formal = data["scope"]["comparison_mode"] == "admission_official"
    before_facts = candidates[fixed]
    before_complete = {row["batch_ref"] for row in before_facts["batches"]["data"]["items"]
        if data["baseline"]["state"] == "available" and row["before"]["completion_state"] == "complete"}
    shared = complete.get(fixed, set()) & (before_complete if formal else complete.get(reference, set()))
    completion, remaining = [], []
    contexts = [(candidate["candidate_ref"], candidate["label"], "candidate", "after", candidates[candidate["candidate_ref"]],
                 candidates[candidate["candidate_ref"]]["available"], complete[candidate["candidate_ref"]]) for candidate in data["candidates"]]
    if formal:
        contexts.append((None, data["baseline"].get("label"), "admission_official", "before", before_facts,
                         data["baseline"]["state"] == "available", before_complete))
    for ref, label, origin, side, facts, available, completed_refs in contexts:
        identity = {"安排组号": groups.get(ref), "安排来源类型": origin, "候选永久编号": ref,
                    "候选名称": label, "来源可用": available}
        by_batch = {row["batch_ref"]: row for row in facts["batches"]["data"]["items"]}
        for batch in data["scope"]["batch_refs"]:
            item = by_batch[batch]
            values = {header: preparation_export_path(item, path.replace("after.", side + "."))
                      for header, path in RECOMMENDATION_EXPORT_COMPLETION_FIELDS.items() if path}
            values["安排缺少资料"] = _export_issue(values["安排缺少资料"])
            category = item[side]["completion_state"] if available else "unknown"
            completion.append({"record_kind": "completion_set", **identity, **values,
                "安排分类": category, "双方工序都已排齐": batch in shared and category == "complete"})
        pending = [row for row in facts["preparation"]["data"]["items"] if available
                   and not row[side]["arrangements"] and row["batch_ref"] not in completed_refs]
        pending.sort(key=lambda row: (row["batch_ref"], row["sequence"] is None, row["sequence"] or 0, row["operation_ref"]))
        for item in pending:
            values = {header: preparation_export_path(item, path.replace("after.", side + "."))
                      for header, path in RECOMMENDATION_EXPORT_REMAINING_FIELDS.items() if path}
            for header in ("归档标准缺项", "需求分解缺项"):
                values[header] = _export_issue(values[header])
            disposition = dispositions["operations"].get(item["operation_ref"], {})
            remaining.append({"record_kind": "remaining_work", **identity, **values,
                "排产时安排状态": disposition.get("status"), "排产时说明": _export_issue(disposition.get("issues"))})
        if not available:
            remaining.append({"record_kind": "remaining_unavailable", **identity, "排产时说明": "这份方案未保存完整安排，无法确认未安排工序及其所需工时。"})
    tables = {"推荐说明": recommendation_export_explanations(document), "工序安排": completion, "剩余需求": remaining}
    return {title: [{**dict.fromkeys(RECOMMENDATION_EXPORT_HEADERS[title]), **common, **row} for row in rows]
            for title, rows in tables.items()}


def recommendation_export_cell_matches(actual, expected, fmt, header, kind):
    if type(expected) is bool:
        return actual == str(expected) if fmt == "csv" else actual is expected and kind == "b"
    return _check_detail_export_cell(actual, expected, fmt, header, kind)


def recommendation_export_record_identity(row, title):
    if title == "工序安排":
        fields = ("record_kind", "候选永久编号", "批次永久编号")
    elif title == "剩余需求":
        fields = ("record_kind", "候选永久编号", "工序永久编号")
    else:
        fields = ("record_kind", "指标标识", "相关明细类型", "相关明细永久编号", "相关明细候选编号", "相关明细筛选字段", "说明")
    return tuple(row.get(header) or None for header in fields)


def recommendation_export_check_table(rows, headers, fmt, types, expected, title, *, csv_union=False):
    wanted_headers = RECOMMENDATION_EXPORT_CSV_HEADERS if csv_union else RECOMMENDATION_EXPORT_HEADERS[title]
    if tuple(headers) != wanted_headers or len(headers) != len(set(headers)) or len(rows) != len(expected):
        raise RuntimeError("推荐导出列或全部事实行数与独立 API 清单不同：" + title)
    if title in ("工序安排", "剩余需求") and [recommendation_export_record_identity(row, title) for row in rows] != [
            recommendation_export_record_identity(row, title) for row in expected]:
        raise RuntimeError("工序安排或剩余需求的候选、批次或工序原序不同：" + title)
    if title == "推荐说明":
        if rows and rows[0].get("record_kind") != "conclusion":
            raise RuntimeError("推荐导出首行缺少固定推荐结论。")
        for kind in ("fixed_selection", "selection_explanation", "missing_evidence", "difference"):
            fields = ("指标标识", "说明") if kind == "missing_evidence" else ("指标标识",)
            actual_order = [tuple(row.get(field) or None for field in fields) for row in rows if row.get("record_kind") == kind]
            wanted_order = [tuple(row.get(field) or None for field in fields) for row in expected if row.get("record_kind") == kind]
            if actual_order != wanted_order:
                raise RuntimeError("推荐说明的公共有序数组不同：" + kind)
    buckets = {}
    for row in expected:
        buckets.setdefault(recommendation_export_record_identity(row, title), []).append(row)
    checked = numeric = boolean = blanks = longest = 0
    for index, row in enumerate(rows):
        key = recommendation_export_record_identity(row, title)
        pool = buckets.get(key, [])
        matching = None
        for wanted in pool:
            if all(recommendation_export_cell_matches(row.get(header), wanted.get(header), fmt, header,
                   types[index].get(header) if types is not None else None) for header in wanted_headers):
                matching = wanted
                break
        if matching is None:
            raise RuntimeError(f"推荐导出原始字段、来源或类型不同：{title} / 行{index + 1} / 身份 {key!r}")
        pool.remove(matching)
        for header in wanted_headers:
            value = matching.get(header)
            checked += 1
            numeric += type(value) in (int, float)
            boolean += type(value) is bool
            blanks += value is None
            longest = max(longest, len(value) if isinstance(value, str) else 0)
    if any(buckets.values()):
        raise RuntimeError("推荐导出缺少独立 API 事实记录：" + title)
    return {"rows": len(rows), "columns": len(wanted_headers), "checked_cells": checked, "numeric_cells": numeric,
        "boolean_cells": boolean, "blank_cells": blanks, "max_text_cell_length": longest,
        "complete_fact_membership_matches_api": True, "ordered_arrays_preserved": True, "dict_key_order_not_assumed": True}


def recommendation_export_check_scope(workbook, document):
    data, meta = document["data"], document["meta"]
    scope, baseline = data["scope"], data["baseline"]
    entries = {row[0]: row[1] if len(row) > 1 else None for row in workbook["比较范围"].iter_rows(values_only=True)}
    expected = {"排产编号": data["run"]["run_ref"], "归档编号": data["capture_ref"], "读取编号": meta["snapshot_ref"],
        "读取时间": meta["as_of"], "比较模式": "排产时的正式计划" if scope["comparison_mode"] == "admission_official" else "同次候选",
        "查看基准编号": baseline["candidate_ref"] if baseline["kind"] == "candidate" else baseline["baseline_ref"],
        "所选方案编号": scope["selected_candidate_ref"], "推荐方案编号": data["recommendation"]["candidate_ref"],
        "推荐版本": data["recommendation"]["recommendation_version"],
        "推荐记录格式版本": data["selection_basis"]["selection_policy_version"],
        "范围起点（含）": scope["range_start"], "范围终点（不含）": scope["range_end"],
        "批次数": len(scope["batch_refs"]), "工序数": len(scope["operation_refs"]),
        "资源视角": "本次所选批次" if scope["resource_view"] == "selected_batches" else "方案全部安排",
        "共同设备数": len(scope["resource_refs"]["machine"]), "共同人员数": len(scope["resource_refs"]["operator"]),
        "筛选页签": None, "筛选条件": "完整比较范围", "明细范围": "筛选结果全部页"}
    if any(entries.get(key) != value for key, value in expected.items()):
        raise RuntimeError("推荐导出范围、原推荐与当前查看身份或来源版本不同。")
    return {"checked_fields": len(expected), "snapshot_and_source_identities_match_api": True}


def recommendation_export_check_download(content, fmt, document, expected, documents=None, archived_labels=None):
    if fmt == "csv":
        reader = csv.DictReader(io.StringIO(content.decode("utf-8-sig")))
        actual = list(reader)
        grouped = {title: [] for title in RECOMMENDATION_EXPORT_HEADERS}
        rank, last = {title: index for index, title in enumerate(grouped)}, 0
        for row in actual:
            kind = row["record_kind"]
            title = "工序安排" if kind == "completion_set" else "剩余需求" if kind in ("remaining_work", "remaining_unavailable") else "推荐说明"
            if rank[title] < last:
                raise RuntimeError("推荐 CSV 的说明、完成和剩余事实分区顺序不同。")
            last = rank[title]
            grouped[title].append(row)
        return {title: recommendation_export_check_table(grouped[title], reader.fieldnames, fmt, None, expected[title], title,
                                                         csv_union=True) for title in grouped}
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    checked = {}
    try:
        business = ["方案总览", "批次影响", "工序变化", "资源影响", "覆盖工序", "准备需求", "工种切换"] if documents else []
        wanted_sheets = business + list(RECOMMENDATION_EXPORT_HEADERS) + ["指标计算方式", "比较范围"]
        if workbook.sheetnames != wanted_sheets:
            raise RuntimeError("推荐 XLSX 的五表或十二表顺序不同：" + str(workbook.sheetnames))
        titles = {"overview": "方案总览", "batches": "批次影响", "operations": "工序变化", "resources": "资源影响",
                  "preparation": "准备需求", "changeovers": "工种切换"}
        for name, source in (documents or {}).items():
            sheet = workbook[titles[name]]
            if name == "overview":
                headers, rows = preparation_export_overview_expected(source["data"])
                checked[name] = preparation_export_check_records(sheet, headers, rows)
                continue
            iterator = sheet.iter_rows()
            headers = [cell.value for cell in next(iterator)]
            cells = list(iterator)
            rows = [dict(zip(headers, (cell.value for cell in row))) for row in cells]
            types = [dict(zip(headers, (cell.data_type for cell in row))) for row in cells]
            if name in ("preparation", "changeovers"):
                checked[name] = preparation_export_check_table(source, rows, headers, fmt, types)
            elif name == "operations":
                checked[name] = preparation_export_check_operations(source, rows, headers, fmt, types)
            elif name == "resources":
                checked[name] = _check_resource_export_table(source, rows, headers, fmt, types)
            else:
                checked[name] = _check_detail_export_table(source, rows, headers, fmt, types)
        if documents:
            checked["coverage-operations"] = preparation_export_check_records(workbook["覆盖工序"], RESOURCE_COVERAGE_HEADERS,
                _resource_coverage_expected(documents, archived_labels))
        for title in RECOMMENDATION_EXPORT_HEADERS:
            iterator = workbook[title].iter_rows()
            headers = [cell.value for cell in next(iterator)]
            cells = list(iterator)
            rows = [dict(zip(headers, (cell.value for cell in row))) for row in cells]
            types = [dict(zip(headers, (cell.data_type for cell in row))) for row in cells]
            checked[title] = recommendation_export_check_table(rows, headers, fmt, types, expected[title], title)
        definitions = document["data"]["metric_definitions"]
        checked["metric-definitions"] = preparation_export_check_records(workbook["指标计算方式"], RESOURCE_METRIC_HEADERS,
            preparation_export_metric_expected(definitions))
        checked["comparison-scope"] = recommendation_export_check_scope(workbook, document)
        checked["sheet_names"] = workbook.sheetnames
    finally:
        workbook.close()
    return checked


def recommendation_export_document(client, fixture, query, target, section, cache):
    key = (fixture["run_ref"], section, tuple(sorted(query.items())))
    if key not in cache:
        path = API + "/runs/" + fixture["run_ref"] + "/decision"
        cache[key] = client.document(path + "?" + urlencode(query)) if section == "overview" else _detail_pages(client, path, query, target, section)
    _save(target / (section + "-all.json"), cache[key])
    return cache[key]


def recommendation_export_candidate_documents(client, fixture, query, document, target, cache):
    result = {}
    for candidate in document["data"]["candidates"]:
        ref = candidate["candidate_ref"]
        current = dict(query, selected_candidate_ref=ref)
        result[ref] = {"available": candidate["persisted_status"] == "completed"}
        for name in ("batches", "preparation"):
            public = recommendation_export_document(client, fixture, current, target / ref, name, cache)
            if public["data"]["candidate"]["candidate_ref"] != ref or public["data"]["run_ref"] != fixture["run_ref"]:
                raise RuntimeError("完成或剩余源资料的候选身份与本次原记录不同。")
            result[ref][name] = public
    return result


def recommendation_export_support_refs(document, candidates, *, require_reference_removed=False):
    data = document["data"]
    fixed = data["recommendation"]["candidate_ref"]
    source = candidates[fixed]["preparation"]["data"]["items"] if fixed in candidates else []
    by_side = {side: {row["operation_ref"]: {arrangement["row_ref"] for arrangement in row[side]["arrangements"]} for row in source}
               for side in ("before", "after")}
    row_refs = {ref for rows in by_side["after"].values() for ref in rows}
    refs = {"batch": set(data["scope"]["batch_refs"]), "operation": set(data["scope"]["operation_refs"]),
            "preparation": set(data["scope"]["operation_refs"]), "changeover": row_refs,
            "resource": {ref for rows in data["scope"]["resource_refs"].values() for ref in rows}, "collection": {None}}
    checked = reference_count = formal_count = removed_count = 0
    for entry in data["recommendation"]["current_comparison"]["all_differences"]:
        for support in entry["supporting_refs"]:
            if support["candidate_ref"] != fixed or support["ref"] not in refs[support["kind"]]:
                raise RuntimeError("推荐支撑事实没有指向保存推荐的永久身份。")
            drill = support["drill"]
            side = support.get("side", "selected")
            formal = side == "reference" and data["scope"]["comparison_mode"] == "admission_official"
            actual_candidate = data["baseline"].get("candidate_ref") if side == "reference" and not formal else fixed
            if drill.get("candidate_ref") != actual_candidate or drill.get("selected_candidate_ref") != fixed:
                raise RuntimeError("推荐支撑的固定推荐与实际安排来源侧候选不同。")
            if formal and (drill.get("reference_view") is not True or drill.get("plan_ref") != data["baseline"].get("plan_ref")):
                raise RuntimeError("推荐支撑的正式基准来源身份不同。")
            selected_row = drill.get("selected_row_ref")
            source_rows = by_side["before" if side == "reference" else "after"]
            if selected_row is not None and selected_row not in source_rows.get(drill.get("operation_ref"), set()):
                raise RuntimeError("推荐支撑的安排编号与声明来源侧的实际工序安排不同。")
            if (side == "reference" and support["kind"] == "operation"
                    and by_side["before"].get(support["ref"]) and not by_side["after"].get(support["ref"])
                    and selected_row in by_side["before"][support["ref"]]):
                removed_count += 1
            checked += 1
            reference_count += side == "reference"
            formal_count += formal
    if require_reference_removed and not removed_count:
        raise RuntimeError("普通部分结果未观察到来自比较基准实际安排的 removed 工序支撑。")
    return {"support_count": checked, "reference_side_support_count": reference_count, "formal_side_support_count": formal_count,
            "reference_side_removed_support_count": removed_count,
            "fixed_recommendation_refs_verified": True, "actual_source_side_refs_verified": True}


def recommendation_export_case(client, fixture, query, target, cache, dispositions, archived_labels=None, *, section="recommendation",
                               require_reference_removed=False, identity=None):
    if section == "all":
        result, _ = current_all_export_case(client, identity, fixture, query, target, cache,
            documents={"recommendation-dispositions": dispositions}, archived_labels=archived_labels)
        return result
    target.mkdir(parents=True, exist_ok=True)
    active = recommendation_export_document(client, fixture, query, target, "overview", cache)
    candidates = recommendation_export_candidate_documents(client, fixture, query, active, target / "candidate-sources", cache)
    expected = recommendation_export_expected(active, candidates, dispositions)
    _save(target / "expected-public-records.json", expected)
    source_refs = recommendation_export_support_refs(active, candidates, require_reference_removed=require_reference_removed)
    documents = None
    base_count = 0
    if section == "all":
        documents = {"overview": active, "batches": candidates[query["selected_candidate_ref"]]["batches"]}
        documents["operations"] = recommendation_export_document(client, fixture, query, target, "operations", cache)
        documents["resources"] = [recommendation_export_document(client, fixture, dict(query, resource_kind=kind),
            target / kind, "resources", cache) for kind in ("machine", "operator")]
        documents["preparation"] = candidates[query["selected_candidate_ref"]]["preparation"]
        documents["changeovers"] = recommendation_export_document(client, fixture, query, target, "changeovers", cache)
        for name, sources in documents.items():
            for document in sources if isinstance(sources, list) else [sources]:
                if document["data"]["scope"] != active["data"]["scope"]:
                    raise RuntimeError("推荐全部导出中的旧分区共同范围不同。")
                base_count += (len(document["data"]["candidates"]) * len(document["data"]["metric_definitions"])
                               if name == "overview" else len(document["data"]["items"]))
        base_count += len(_resource_coverage_expected(documents, archived_labels))
        for name in ("batches", "operations", "preparation", "changeovers"):
            _save(target / (name + "-all.json"), documents[name])
    count = base_count + sum(len(rows) for rows in expected.values())
    export_query = dict(query, section=section, snapshot_ref=active["meta"]["snapshot_ref"])
    _save(target / "request.json", export_query)
    formats = {}
    path = API + "/runs/" + fixture["run_ref"] + "/decision/export"
    for fmt in (("xlsx",) if section == "all" else ("csv", "xlsx")):
        content, headers = _fetch_export(client, path + "?" + urlencode(dict(export_query, format=fmt)), target, fmt, section=section)
        if (headers.get("X-Workbench-Run-Ref") != fixture["run_ref"]
                or headers.get("X-Workbench-Snapshot-Ref") != export_query["snapshot_ref"]
                or int(headers["X-Workbench-Row-Count"]) != count):
            raise RuntimeError("推荐导出的运行、快照绑定或全部业务行数不同。")
        formats[fmt] = recommendation_export_check_download(content, fmt, active, expected, documents, archived_labels)
    data = active["data"]
    result = {"state": "passed", "run_ref": fixture["run_ref"], "section": section, "scope": data["scope"],
        "snapshot_ref": active["meta"]["snapshot_ref"], "saved_recommendation_ref": data["recommendation"]["candidate_ref"],
        "viewed_candidate_ref": data["scope"]["selected_candidate_ref"], "reference_candidate_ref": data["baseline"].get("candidate_ref"),
        "comparison_mode": data["scope"]["comparison_mode"], "formal_reference_version": data["baseline"].get("version"),
        "recommendation_type": data["recommendation"]["recommendation_type"], "selection_reason": data["selection_basis"]["reason_code"],
        "selection_policy_version": data["selection_basis"]["selection_policy_version"],
        "business_row_count": count, "recommendation_table_counts": {name: len(rows) for name, rows in expected.items()},
        "source_refs": source_refs, "candidate_source_count": len(candidates), "pages_20_50_same_snapshot": True,
        "formats": formats, "adoption_performed": False, "actual_engine_rerun": False}
    _save(target / "verification.json", result)
    return result


def recommendation_export_verify(args):
    from datetime import date, timedelta

    identity, client = _identity(args.instance)
    evidence = ROOT / "evidence" / "decision-recommendation" / Path(identity["root"]).name
    balanced = json.loads((evidence / "balanced-fixture.json").read_text(encoding="utf-8"))
    legacy = json.loads((ROOT / "evidence" / "decision-delivery-impact" / Path(identity["root"]).name /
                         "delivery-fixtures.json").read_text(encoding="utf-8"))
    if (balanced["state"] != "run_complete" or any(Path(record["database_path"]).resolve() != Path(identity["database_path"]).resolve()
                                                  for record in (balanced, legacy))):
        raise RuntimeError("推荐导出夹具尚未完成或与当前隔离数据库不同。")
    def query(fixture):
        return {"comparison_mode": "candidates", "selected_candidate_ref": fixture["candidate_refs"][-1],
            "reference_candidate_ref": fixture["candidate_refs"][0], "resource_view": "selected_batches",
            "range_start": fixture["settings"]["start_date"] + "T00:00:00",
            "range_end": (date.fromisoformat(fixture["settings"]["end_date"]) + timedelta(days=1)).isoformat() + "T00:00:00"}
    balanced_query = dict(query(balanced), reference_candidate_ref=balanced["candidate_refs"][2], range_end="2026-10-17T00:00:00")
    partial, mixed = (legacy["fixtures"][name] for name in ("ordinary_partial", "mixed_due"))
    partial_query = query(partial)
    formal_query = {key: value for key, value in balanced_query.items() if key != "reference_candidate_ref"}
    formal_query["comparison_mode"] = "admission_official"
    cases = [("balanced-recommendation", balanced, balanced_query, "recommendation"),
        ("ordinary-partial-recommendation", partial, partial_query, "recommendation"),
        ("mixed-due-recommendation", mixed, query(mixed), "recommendation"),
        ("all-balanced-twelve-tables", balanced, formal_query, "all")]
    if args.case:
        cases = [case for case in cases if case[0] in args.case]
    target = Path(args.evidence_directory).resolve() if args.evidence_directory else evidence / "exports"
    cache, dispositions, labels = {}, {}, {}
    for _name, fixture, _query, section in cases:
        run_ref = fixture["run_ref"]
        if run_ref not in dispositions:
            source = evidence / "balanced-run" / "original-receipt.json" if run_ref == balanced["run_ref"] else None
            dispositions[run_ref] = recommendation_export_dispositions(identity, run_ref, source)
            _save(target / ("dispositions-" + run_ref + ".json"), dispositions[run_ref])
        if section == "all":
            labels[run_ref] = _resource_archive_labels(identity, run_ref)
            _save(target / ("coverage-labels-" + run_ref + ".json"), labels[run_ref])
    results, current = [], None
    try:
        for current, fixture, current_query, section in cases:
            result = recommendation_export_case(client, fixture, current_query, target / current, cache,
                dispositions[fixture["run_ref"]], labels.get(fixture["run_ref"]), section=section,
                require_reference_removed=current == "ordinary-partial-recommendation", identity=identity)
            result.update(case=current, runtime_identity={"base_url": identity["base_url"], "pid": identity["pid"]})
            _save(target / current / "verification.json", result)
            results.append(result)
    except Exception as error:
        _save(target / "export-verification.json", {"state": "failed", "case": current, "error": str(error), "cases": results,
            "base_url": identity["base_url"], "pid": identity["pid"]})
        raise
    summary = {"state": "passed", "base_url": identity["base_url"], "pid": identity["pid"], "database_path": identity["database_path"],
        "case_count": len(results), "download_count": sum(len(row["formats"]) for row in results),
        "business_row_count": sum(row["business_row_count"] * len(row["formats"]) for row in results),
        "recommendation_columns": {name: len(headers) for name, headers in RECOMMENDATION_EXPORT_HEADERS.items()},
        "csv_columns": len(RECOMMENDATION_EXPORT_CSV_HEADERS), "cases": results,
        "evidence_directory": str(target), "adoption_performed": False, "actual_engine_rerun": False}
    for name in ("checked_cells", "numeric_cells", "boolean_cells", "blank_cells", "json_cells"):
        summary[name] = sum(table.get(name, 0) for row in results for tables in row["formats"].values()
                            for table in tables.values() if isinstance(table, dict))
    _save(target / "export-verification.json", summary)
    print(json.dumps({key: value for key, value in summary.items() if key != "cases"}, ensure_ascii=False), flush=True)


# Independent process export contract from public DTO fields; no product row writer is used.
PROCESS_EXPORT_COMMON = ('排产编号', '归档编号', '读取编号', '读取时间', '保存推荐候选编号', '当前查看候选编号', '比较模式', '查看基准编号', '查看基准类型', '范围起点（含）', '范围终点（不含）', '资源视角', '推荐版本', '推荐记录格式版本', '过程版本', '查看过程候选编号', '查看过程候选名称', '求解来源候选编号', '过程来源类型', '过程来源可用', '过程证据版本', '过程证据状态', '主目标', '主指标key', '主指标', '主指标单位')
PROCESS_EXPORT_FIELD_PATHS = {'候选计算': {'record_kind': 'literal candidate_summary',
          '候选顺序类型': 'saved.candidate.kind',
          '保存状态': 'saved.candidate.persisted_status',
          '算法模式': 'saved.candidate.algo_mode',
          '策略': 'saved.candidate.used_strategy',
          '优化器配置类别': 'saved.candidate.algorithm_profile',
          '优化器Profile': 'saved.source.profile.profile',
          '接受准则': 'saved.source.profile.acceptance',
          '实际派工模式': 'saved.candidate.dispatch_mode',
          '配置派工规则': 'saved.candidate.configured_dispatch_rule',
          '实际采用派工规则': 'saved.candidate.adopted_dispatch_rule',
          '候选自身耗时（毫秒）': 'saved.candidate.own_elapsed_ms',
          '来源求解耗时（毫秒）': 'saved.candidate.source_solver_elapsed_ms',
          '停止原因': 'saved.candidate.stop_reason.code or null',
          '停止说明': 'saved.candidate.stop_reason.message or null',
          '优化器已启用': 'saved.source.profile.enabled',
          '优化器随机种子': 'saved.source.profile.seed',
          '配置求解预算（秒）': 'saved.source.profile.configured_time_budget_seconds',
          '实际求解预算（秒）': 'saved.source.profile.effective_time_budget_seconds',
          '已分配求解预算（毫秒）': 'saved.source.profile.assigned_time_budget_ms',
          '配置迭代上限': 'saved.source.profile.configured_max_iterations',
          '实际迭代上限': 'saved.source.profile.effective_max_iterations',
          '已应用系统限制': 'saved.source.profile.system_limit_applied',
          '系统限制原因': 'saved.source.profile.system_limit_reason',
          'Profile说明': 'saved.source.profile.message',
          '候选策略族': 'saved.source.profile.candidate_strategy_families',
          '配置邻域': 'saved.source.profile.configured_neighborhoods',
          '实际邻域': 'saved.source.profile.effective_neighborhoods',
          '候选自身评估数': 'saved.candidate.own_counts.evaluated',
          '候选自身接受数': 'saved.candidate.own_counts.accepted',
          '候选自身当前解接受数': 'saved.candidate.own_counts.current_accepted',
          '候选自身严格改善数': 'saved.candidate.own_counts.strict_improvements',
          '候选自身已记录拒绝数': 'saved.candidate.own_counts.rejected',
          '候选自身未验证数': 'saved.candidate.own_counts.unverified_trials',
          '候选自身实际排程调用数': 'saved.candidate.own_counts.decoder_invocations',
          '来源求解评估数': 'saved.candidate.solver_counts.evaluated',
          '来源求解接受数': 'saved.candidate.solver_counts.accepted',
          '来源求解当前解接受数': 'saved.candidate.solver_counts.current_accepted',
          '来源求解严格改善数': 'saved.candidate.solver_counts.strict_improvements',
          '来源求解已记录拒绝数': 'saved.candidate.solver_counts.rejected',
          '来源求解未验证数': 'saved.candidate.solver_counts.unverified_trials',
          '来源求解实际排程调用数': 'saved.candidate.solver_counts.decoder_invocations',
          '保存超期批次数': 'saved.candidate.final_metrics.overdue_count',
          '保存总拖期（小时）': 'saved.candidate.final_metrics.total_tardiness_hours',
          '保存加权拖期（小时）': 'saved.candidate.final_metrics.weighted_tardiness_hours',
          '保存工期（小时）': 'saved.candidate.final_metrics.makespan_hours',
          '保存算法工种切换次数': 'saved.candidate.final_metrics.changeover_count',
          '本次参与计算工序数': 'saved.candidate.completion.expected_operation_count',
          '本次计算范围缺少工序数': 'saved.candidate.completion.missing_operation_count',
          '本次计算范围外结果数': 'saved.candidate.completion.unexpected_result_count',
          '本次计算范围内工序排齐批次数': 'saved.candidate.completion.complete_batch_count',
          '本次计算范围内工序未排齐批次数': 'saved.candidate.completion.incomplete_batch_count',
          '保存代表尝试数': 'saved.summary.saved_attempt_count',
          '旧版最佳更新数（含同分）': 'saved.summary.legacy_best_updates',
          '原进步事件数': 'saved.progress.sampling.event_total_count',
          '保留进步点数': 'saved.progress.sampling.retained_point_count',
          '采样策略': 'saved.progress.sampling.policy',
          '分桶数': 'saved.progress.sampling.bucket_count',
          '桶宽（毫秒）': 'saved.progress.sampling.bucket_width_ms',
          '保留首次有效': 'saved.progress.coverage.endpoint_coverage.first_valid',
          '保留首次排齐本次参与计算工序': 'saved.progress.coverage.endpoint_coverage.first_expected_complete',
          '保留最终': 'saved.progress.coverage.endpoint_coverage.final',
          '覆盖说明': 'saved.progress.coverage.reason.message or null',
          '拒绝原因key': 'null in candidate_summary; reason key in rejection_reason',
          '拒绝原因数': 'null in candidate_summary; source.rejection_summary[key] in rejection_reason'},
 '进步事件': {'record_kind': 'literal progress_point',
          '事件序号': 'point.event_index',
          '累计求解耗时（毫秒）': 'point.elapsed_ms',
          '阶段': 'point.phase',
          '阶段访问序号': 'point.phase_visit',
          '连续段': 'point.segment',
          '首次有效': 'first_valid in point.flags if saved.candidate.evidence_version is not null else null',
          '首次排齐本次参与计算工序': 'first_expected_complete in point.flags if saved.candidate.evidence_version is not null else '
                    'null',
          '最终结果': 'final in point.flags if saved.candidate.evidence_version is not null else null',
          '业务主值': 'point.primary_value',
          '目标评价范围': 'point.objective_scope',
          '评价范围未排齐批次编号': 'point.evaluation_scope.incomplete_batch_ids',
          '已验证': 'point.verified',
          '本次参与计算工序数': 'point.completion.expected_operation_count',
          '本次计算范围缺少工序数': 'point.completion.missing_operation_count',
          '本次计算范围外结果数': 'point.completion.unexpected_result_count',
          '本次计算范围内工序排齐批次数': 'point.completion.complete_batch_count',
          '本次计算范围内工序未排齐批次数': 'point.completion.incomplete_batch_count',
          '评分分量0': 'point.score[0] if present else null',
          '评分分量1': 'point.score[1] if present else null',
          '评分分量2': 'point.score[2] if present else null',
          '评分分量3': 'point.score[3] if present else null',
          '评分分量4': 'point.score[4] if present else null',
          '评分分量5': 'point.score[5] if present else null',
          '评分分量0单位': 'saved.objective.score_fields item with index=0.unit if present else null',
          '评分分量1单位': 'saved.objective.score_fields item with index=1.unit if present else null',
          '评分分量2单位': 'saved.objective.score_fields item with index=2.unit if present else null',
          '评分分量3单位': 'saved.objective.score_fields item with index=3.unit if present else null',
          '评分分量4单位': 'saved.objective.score_fields item with index=4.unit if present else null',
          '评分分量5单位': 'saved.objective.score_fields item with index=5.unit if present else null'},
 '搜索阶段': {'record_kind': 'literal stage',
          '阶段': 'stage.phase',
          '阶段状态': 'stage.status',
          '阶段起点（毫秒）': 'stage.start_elapsed_ms',
          '阶段终点（毫秒）': 'stage.end_elapsed_ms',
          '阶段累计耗时（毫秒）': 'stage.elapsed_ms',
          '访问数': 'stage.visits',
          '评估数': 'stage.evaluated',
          '接受数': 'stage.accepted',
          '严格改善数': 'stage.strict_improvements',
          '已记录拒绝数': 'stage.rejected',
          '停止原因': 'stage.stop_reason'},
 '代表尝试': {'record_kind': 'literal representative_attempt',
          '保存尝试序号': 'attempt.saved_index',
          '代表排名': 'attempt.rank',
          '尝试来源': 'attempt.source_label',
          '尝试策略': 'attempt.strategy',
          '尝试派工模式': 'attempt.dispatch_mode',
          '尝试派工规则': 'attempt.dispatch_rule',
          '尝试保存状态': 'attempt.candidate_status',
          '尝试来源key': 'attempt.candidate_origin',
          '权重Profile': 'attempt.weight_profile_slug',
          '尝试事件序号': 'attempt.event_index',
          '尝试失败工序数': 'attempt.failed_ops',
          '尝试累计耗时（毫秒）': 'attempt.elapsed_ms',
          '尝试已接受': 'attempt.accepted',
          '尝试已验证': 'attempt.verified',
          '相对当时对象严格改善': 'attempt.strict_improved',
          '业务主值': 'attempt.primary_value',
          '目标评价范围': 'attempt.objective_scope',
          '当时比较对象': 'attempt.comparison.kind',
          '当时比较状态': 'attempt.comparison.state',
          '当时比较目标': 'attempt.comparison.objective_name',
          '当时本次评价范围': 'attempt.comparison.objective_scope',
          '当时参考评价范围': 'attempt.comparison.reference_scope',
          '保存超期批次数': 'attempt.metrics.overdue_count',
          '保存总拖期（小时）': 'attempt.metrics.total_tardiness_hours',
          '保存加权拖期（小时）': 'attempt.metrics.weighted_tardiness_hours',
          '保存工期（小时）': 'attempt.metrics.makespan_hours',
          '保存算法工种切换次数': 'attempt.metrics.changeover_count',
          '本次参与计算工序数': 'attempt.completion.expected_operation_count',
          '本次计算范围缺少工序数': 'attempt.completion.missing_operation_count',
          '本次计算范围外结果数': 'attempt.completion.unexpected_result_count',
          '本次计算范围内工序排齐批次数': 'attempt.completion.complete_batch_count',
          '本次计算范围内工序未排齐批次数': 'attempt.completion.incomplete_batch_count',
          '评分分量0': 'attempt.score[0] if present else null',
          '评分分量1': 'attempt.score[1] if present else null',
          '评分分量2': 'attempt.score[2] if present else null',
          '评分分量3': 'attempt.score[3] if present else null',
          '评分分量4': 'attempt.score[4] if present else null',
          '评分分量5': 'attempt.score[5] if present else null',
          '评分分量0单位': 'saved.objective.score_fields item with index=0.unit if present else null',
          '评分分量1单位': 'saved.objective.score_fields item with index=1.unit if present else null',
          '评分分量2单位': 'saved.objective.score_fields item with index=2.unit if present else null',
          '评分分量3单位': 'saved.objective.score_fields item with index=3.unit if present else null',
          '评分分量4单位': 'saved.objective.score_fields item with index=4.unit if present else null',
          '评分分量5单位': 'saved.objective.score_fields item with index=5.unit if present else null',
          '当时参考评分分量0': 'attempt.comparison.reference_score[0]',
          '当时参考评分分量1': 'attempt.comparison.reference_score[1]',
          '当时参考评分分量2': 'attempt.comparison.reference_score[2]',
          '当时参考评分分量3': 'attempt.comparison.reference_score[3]',
          '当时参考评分分量4': 'attempt.comparison.reference_score[4]',
          '当时参考评分分量5': 'attempt.comparison.reference_score[5]',
          '当时参考失败工序数': 'attempt.comparison.reference_metrics.failed_ops',
          '当时参考超期批次数': 'attempt.comparison.reference_metrics.overdue_count',
          '当时参考加权拖期（小时）': 'attempt.comparison.reference_metrics.weighted_tardiness_hours',
          '当时参考总拖期（小时）': 'attempt.comparison.reference_metrics.total_tardiness_hours',
          '当时参考工期（小时）': 'attempt.comparison.reference_metrics.makespan_hours',
          '当时参考算法工种切换次数': 'attempt.comparison.reference_metrics.changeover_count',
          '相对当时对象变化失败工序数': 'attempt.comparison.metric_deltas.failed_ops',
          '相对当时对象变化超期批次数': 'attempt.comparison.metric_deltas.overdue_count',
          '相对当时对象变化加权拖期（小时）': 'attempt.comparison.metric_deltas.weighted_tardiness_hours',
          '相对当时对象变化总拖期（小时）': 'attempt.comparison.metric_deltas.total_tardiness_hours',
          '相对当时对象变化工期（小时）': 'attempt.comparison.metric_deltas.makespan_hours',
          '相对当时对象变化算法工种切换次数': 'attempt.comparison.metric_deltas.changeover_count',
          '当时参考本次参与计算工序数': 'attempt.comparison.reference_completion.expected_operation_count',
          '当时参考本次计算范围缺少工序数': 'attempt.comparison.reference_completion.missing_operation_count',
          '当时参考本次计算范围外结果数': 'attempt.comparison.reference_completion.unexpected_result_count',
          '当时参考本次计算范围内工序排齐批次数': 'attempt.comparison.reference_completion.complete_batch_count',
          '当时参考本次计算范围内工序未排齐批次数': 'attempt.comparison.reference_completion.incomplete_batch_count'}}
PROCESS_EXPORT_HEADERS = {title: PROCESS_EXPORT_COMMON + tuple(fields) for title, fields in PROCESS_EXPORT_FIELD_PATHS.items()}
PROCESS_EXPORT_CSV_HEADERS = tuple(dict.fromkeys(header for headers in PROCESS_EXPORT_HEADERS.values() for header in headers))

def _process_public_pages(client, path, query, target, kind):
    """Keep raw HTTP pages and prove complete 20/50 row unions on one snapshot."""
    target.mkdir(parents=True, exist_ok=True)
    bound = dict(query, record_kind=kind)
    pages, combined = {}, {}
    first = None
    for size in (20, 50):
        current, items, number = [], [], 1
        while True:
            public = client.document(path + "/process?" + urlencode(dict(bound, page=number, size=size)))
            data, meta = public["data"], public["meta"]
            if first is None:
                first = public
                bound["snapshot_ref"] = meta["snapshot_ref"]
            if (meta["snapshot_ref"] != first["meta"]["snapshot_ref"] or meta["as_of"] != first["meta"]["as_of"]
                    or data["page"]["number"] != number or data["page"]["size"] != size
                    or data["page"]["total"] != first["data"]["page"]["total"]):
                raise RuntimeError("计算过程分页的快照、读取时间或页码不同。")
            if data["filters"] != first["data"]["filters"] or data["candidate"] != first["data"]["candidate"]:
                raise RuntimeError("计算过程翻页改变了锚点或筛选条件。")
            current.append(public)
            items.extend(data["items"])
            expected_more = number * size < data["page"]["total"]
            if len(data["items"]) != min(size, max(0, data["page"]["total"] - (number - 1) * size)) or data["page"]["has_more"] is not expected_more:
                raise RuntimeError("计算过程分页条数或下一页状态与完整总数不同。")
            if not data["page"]["has_more"]:
                break
            number += 1
        pages[size], combined[size] = current, items
        _save(target / (kind + "-pages-" + str(size) + ".json"), current)
    if combined[20] != combined[50] or len(combined[20]) != first["data"]["page"]["total"]:
        raise RuntimeError("计算过程20/50全页原始对象或顺序不同。")
    raw_points = first["data"]["progress"]["points"]
    if kind == "progress" and sorted(combined[20], key=lambda row: json.dumps(row, sort_keys=True)) != sorted(raw_points, key=lambda row: json.dumps(row, sort_keys=True)):
        raise RuntimeError("计算过程分页遗漏或增加了公开完整进步事件。")
    result = dict(pages[50][0])
    result["data"] = dict(result["data"], items=combined[50])
    result["read_evidence"] = {"kind": "public_GET_page_union", "total": len(combined[50]),
        "raw_pages_20": kind + "-pages-20.json", "raw_pages_50": kind + "-pages-50.json",
        "page_counts": {"20": len(pages[20]), "50": len(pages[50])}, "same_snapshot": True, "same_as_of": True,
        "complete_ordered_objects_equal": True}
    _save(target / (kind + "-readback.json"), result)
    return result


def process_export_candidate_documents(client, fixture, query, overview, target, *, anchor_ref=None):
    common_keys = ("comparison_mode", "selected_candidate_ref", "reference_candidate_ref", "range_start", "range_end", "resource_view")
    common = {key: value for key, value in query.items() if key in common_keys}
    refs = [anchor_ref] if anchor_ref is not None else [row["candidate_ref"] for row in overview["data"]["candidates"]]
    catalog = {row["candidate_ref"] for row in overview["data"]["candidates"]}
    if any(ref not in catalog for ref in refs):
        raise RuntimeError("导出的计算过程候选没有对应永久目录记录。")
    path = API + "/runs/" + fixture["run_ref"] + "/decision"
    records = []
    active_kind, order = query.get("record_kind", "attempts"), query.get("order", "asc")
    for ref in refs:
        documents = {}
        for kind in ("attempts", "progress"):
            sort = query.get("sort", "rank" if kind == "attempts" else "elapsed_ms") if active_kind == kind else ("rank" if kind == "attempts" else "elapsed_ms")
            current = dict(common, process_candidate_ref=ref, record_kind=kind, sort=sort, order=order)
            documents[kind] = _process_public_pages(client, path, current, target / ref, kind)
        attempts, progress = (documents[kind]["data"] for kind in ("attempts", "progress"))
        for key in ("candidate", "source", "summary", "objective", "stages", "scope", "baseline", "process_version"):
            if attempts[key] != progress[key]:
                raise RuntimeError("进步事件与代表尝试的保存过程来源不同：" + key)
        if (attempts["candidate"]["candidate_ref"] != ref or attempts["run_ref"] != fixture["run_ref"]
                or attempts["scope"] != overview["data"]["scope"]
                or any(value != overview["data"]["baseline"].get(key) for key, value in attempts["baseline"].items())):
            raise RuntimeError("计算过程的锚点、查看方案、基准或共同窗口不同。")
        if attempts["process_version"] != "decision-process-v1":
            raise RuntimeError("计算过程公共版本与本轮验收契约不同。")
        record = {key: attempts[key] for key in ("candidate", "source", "summary", "objective", "stages", "process_version")}
        record["attempts"] = attempts["items"]
        record["progress"] = dict(progress["progress"], points=progress["items"])
        record["public_readbacks"] = {kind: str((target / ref / (kind + "-readback.json")).resolve()) for kind in documents}
        if len(record["attempts"]) != attempts["summary"]["saved_attempt_count"] or len(record["progress"]["points"]) != attempts["summary"]["retained_point_count"]:
            raise RuntimeError("计算过程全部保存条数与摘要不同。")
        if len(record["attempts"]) > 12 or len(record["progress"]["points"]) > 200:
            raise RuntimeError("计算过程超过本轮保留上限。")
        records.append(record)
    _save(target / "public-process-candidates.json", records)
    return records


def _process_expected_path(root, path):
    value = root
    for key in path.split("."):
        if value is None:
            return None
        if not isinstance(value, dict):
            raise RuntimeError("计算过程预期字段结构不同：" + path)
        indexed = re.fullmatch(r"([^\[]+)\[([0-9]+)\]", key)
        if indexed is None:
            value = value.get(key)
        else:
            array = value.get(indexed.group(1))
            if array is None:
                return None
            if not isinstance(array, list):
                raise RuntimeError("计算过程预期评分结构不同：" + path)
            index = int(indexed.group(2))
            value = array[index] if index < len(array) else None
    return value


def _process_expected_grain(saved, title, item=None):
    root = {"saved": saved, "point": item, "stage": item, "attempt": item}
    kinds = {"候选计算": "candidate_summary", "进步事件": "progress_point", "搜索阶段": "stage", "代表尝试": "representative_attempt"}
    values = {}
    flags = {"首次有效": "first_valid", "首次排齐本次参与计算工序": "first_expected_complete", "最终结果": "final"}
    for header, path in PROCESS_EXPORT_FIELD_PATHS[title].items():
        if header == "record_kind":
            value = kinds[title]
        elif header in ("拒绝原因key", "拒绝原因数"):
            value = None
        elif header.startswith("评分分量"):
            index = int(re.search(r"评分分量([0-5])", header).group(1))
            if header.endswith("单位"):
                units = {field["index"]: field["unit"] for field in saved["objective"]["score_fields"]}
                value = units.get(index)
            else:
                score = item["score"]
                value = score[index] if index < len(score) else None
        elif title == "进步事件" and header in flags:
            value = flags[header] in item["flags"] if saved["candidate"]["evidence_version"] is not None else None
        else:
            value = _process_expected_path(root, (path[:-8] if path.endswith(" or null") else path))
        if type(value) is list:
            if not all(type(entry) is str for entry in value):
                raise RuntimeError("计算过程 Profile 列表含非文本值：" + header)
            value = json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":")) if header == "评价范围未排齐批次编号" else "；".join(value)
        values[header] = value
    return values


def process_export_expected(document, saved_candidates):
    base = recommendation_export_common(document)
    expected = {title: [] for title in PROCESS_EXPORT_HEADERS}
    for saved in saved_candidates:
        row, objective = saved["candidate"], saved["objective"]
        common_values = [saved["process_version"], row["candidate_ref"], row["label"], row["source_candidate_ref"], row["source_kind"],
            saved["source"]["available"], row["evidence_version"], row["evidence_state"], objective["name"],
            objective["primary_metric_key"], objective["primary_metric_label"], objective["unit"]]
        common = dict(base)
        common.update(dict(zip(PROCESS_EXPORT_COMMON[14:], common_values)))
        groups = {"候选计算": [_process_expected_grain(saved, "候选计算")],
            "进步事件": [_process_expected_grain(saved, "进步事件", point) for point in saved["progress"]["points"]],
            "搜索阶段": [_process_expected_grain(saved, "搜索阶段", stage) for stage in saved["stages"]],
            "代表尝试": [_process_expected_grain(saved, "代表尝试", attempt) for attempt in saved["attempts"]]}
        reasons = saved["source"]["rejection_summary"]
        if reasons is not None:
            for code in sorted(reasons):
                if type(reasons[code]) is not int or reasons[code] < 0:
                    raise RuntimeError("保存的拒绝原因计数不是非负原生整数。")
                groups["候选计算"].append({"record_kind": "rejection_reason", "拒绝原因key": code, "拒绝原因数": reasons[code]})
        for title, rows in groups.items():
            expected[title].extend({**dict.fromkeys(PROCESS_EXPORT_HEADERS[title]), **common, **values} for values in rows)
    return expected


def process_export_check_table(rows, headers, fmt, types, expected, title, *, csv_union=False):
    wanted_headers = PROCESS_EXPORT_CSV_HEADERS if csv_union else PROCESS_EXPORT_HEADERS[title]
    if tuple(headers) != wanted_headers or len(headers) != len(set(headers)) or len(rows) != len(expected):
        raise RuntimeError("计算过程导出列或完整事实行数不同：" + title)
    checked = numeric = boolean = blanks = 0
    for index, (actual, wanted) in enumerate(zip(rows, expected)):
        for header in wanted_headers:
            target = wanted.get(header)
            kind = types[index].get(header) if types is not None else None
            if not recommendation_export_cell_matches(actual.get(header), target, fmt, header, kind):
                raise RuntimeError(f"计算过程导出字段、单位或原生类型不同：{title} / 行{index + 1} / {header} / 实际 {actual.get(header)!r} / 期待 {target!r}")
            checked += 1
            numeric += type(target) in (int, float)
            boolean += type(target) is bool
            blanks += target is None
    return {"rows": len(rows), "columns": len(wanted_headers), "checked_cells": checked, "numeric_cells": numeric,
        "boolean_cells": boolean, "blank_cells": blanks, "ordered_records_match_public_API": True,
        "null_zero_false_preserved": True, "source_profile_reason_values_match": True}


_CURRENT_ALL_GLOBAL_KEYS = (
    "comparison_mode", "reference_candidate_ref", "selected_candidate_ref",
    "range_start", "range_end", "resource_view",
)
_CURRENT_ALL_FILTER_KEYS = {
    "batches": ("batch_ref", "change_kind", "sort", "order"),
    "operations": ("batch_ref", "resource_ref", "change_kind", "sort", "order"),
    "resources": ("resource_kind", "resource_ref", "pressure_filter", "sort", "order", "grain", "pressure_threshold", "display_start", "display_end"),
    "resource-days": ("resource_kind", "resource_ref", "grain", "pressure_threshold", "display_start", "display_end", "bucket_start", "bucket_end"),
    "preparation": ("batch_ref", "machine_ref", "allocation_state", "work_scope", "sort", "order"),
    "changeovers": ("machine_ref", "event_side", "sort", "order"),
    "process": ("process_candidate_ref", "record_kind", "sort", "order"),
}
_CURRENT_ALL_TITLES = {
    "overview": "方案总览", "batches": "批次影响", "operations": "工序变化",
    "resources": "资源影响", "resource-days": "资源逐日",
    "preparation": "准备需求", "changeovers": "工种切换",
}
_CURRENT_ALL_RECOMMENDATION_TITLES = ("推荐说明", "工序安排", "剩余需求")
_CURRENT_ALL_PROCESS_TITLES = ("候选计算", "进步事件", "搜索阶段", "代表尝试")


def _current_all_scope_query(scope):
    result = {key: scope[key] for key in (
        "comparison_mode", "selected_candidate_ref", "range_start", "range_end", "resource_view")}
    if scope["comparison_mode"] == "candidates" and scope["reference_state"] == "available":
        result["reference_candidate_ref"] = scope["reference"]["candidate_ref"]
    return {key: value for key, value in result.items() if value is not None}


def _current_all_requested_scope_matches(scope, query):
    actual = _current_all_scope_query(scope)
    for key in _CURRENT_ALL_GLOBAL_KEYS:
        if key not in query or query[key] is None:
            continue
        if key in ("range_start", "range_end"):
            if datetime.fromisoformat(query[key]) != datetime.fromisoformat(actual[key]):
                return False
        elif query[key] != actual.get(key):
            return False
    return True


def _current_all_default_filters(section, query, public=None):
    """Independent public defaults, used only to reuse an exact complete GET."""
    values = {"section": section}
    if section in ("batches", "operations"):
        values.update(batch_ref=None, resource_ref=None, change_kind="all",
                      sort="worsened" if section == "batches" else "sequence",
                      order="desc" if section == "batches" else "asc")
    elif section in ("resources", "resource-days"):
        values.update(resource_kind="machine", resource_ref=None, grain="day", pressure_threshold=0.9,
                      display_start=None, display_end=None, bucket_start=None, bucket_end=None)
        if section == "resources":
            values.update(pressure_filter="all", sort="after_peak_utilization", order="desc")
    elif section in ("preparation", "changeovers"):
        values.update(machine_ref=None, sort="sequence" if section == "preparation" else "next_start", order="asc")
        values.update({"batch_ref": None, "allocation_state": "all", "work_scope": "all"}
                      if section == "preparation" else {"event_side": "all"})
    else:
        raise ValueError("过程候选默认由公共 process GET 解析，不在此处重建。")
    for key in _CURRENT_ALL_FILTER_KEYS[section]:
        if key in query and query[key] is not None:
            values[key] = float(query[key]) if key == "pressure_threshold" else query[key]
    if section == "resources" and values["grain"] == "hour" and values["display_start"] is None and public is not None:
        display = public["heatmap"]["display_scope"]
        values.update(display_start=display["start"], display_end=display["end"])
    return values


def _current_all_sheet_rows(sheet):
    iterator = sheet.iter_rows()
    headers = [cell.value for cell in next(iterator)]
    if len(headers) != len(set(headers)):
        raise RuntimeError("完整 XLSX 导出出现同名表头：" + sheet.title)
    cells = list(iterator)
    rows = [dict(zip(headers, (cell.value for cell in row))) for row in cells]
    types = [dict(zip(headers, (cell.data_type for cell in row))) for row in cells]
    return rows, headers, types


def _current_all_check_scope(workbook, expected_document, filter_section, active_filters):
    """Check every ordered scope field, literal wording, and native XLSX types."""
    data, meta = expected_document["data"], expected_document["meta"]
    scope, baseline = data["scope"], data["baseline"]
    state_labels = preparation_export_labels()["state"]
    filter_orders = {
        "batches": ("section", "batch_ref", "resource_ref", "change_kind", "sort", "order"),
        "operations": ("section", "batch_ref", "resource_ref", "change_kind", "sort", "order"),
        "resources": ("section", "resource_kind", "resource_ref", "grain", "pressure_threshold", "display_start", "display_end",
                      "bucket_start", "bucket_end", "pressure_filter", "sort", "order"),
        "resource-days": ("section", "resource_kind", "resource_ref", "grain", "pressure_threshold", "display_start", "display_end", "bucket_start", "bucket_end"),
        "preparation": ("section", "machine_ref", "sort", "order", "batch_ref", "allocation_state", "work_scope"),
        "changeovers": ("section", "machine_ref", "sort", "order", "event_side"),
        "process": ("section", "process_candidate_ref", "record_kind", "sort", "order"),
    }
    if active_filters is not None and set(active_filters) != set(filter_orders[filter_section]):
        raise RuntimeError("完整导出的活动筛选字段集合不同。")
    description = ("完整比较范围" if active_filters is None else
                   "；".join(key + "=" + str(active_filters[key])
                            for key in filter_orders[filter_section] if active_filters[key] is not None))
    wanted = [
        ("排产编号", data["run"]["run_ref"]),
        ("归档编号", data["capture_ref"]),
        ("读取编号", meta["snapshot_ref"]),
        ("读取时间", meta["as_of"]),
        ("比较模式", "排产时的正式计划" if scope["comparison_mode"] == "admission_official" else "同次候选"),
        ("查看基准编号", baseline.get("candidate_ref") if baseline["kind"] == "candidate" else baseline.get("baseline_ref")),
        ("查看基准状态", state_labels.get(baseline["state"], baseline["state"])),
        ("正式基准版本", baseline.get("version")),
        ("正式来源候选编号", baseline.get("source_candidate_ref")),
        ("所选方案编号", scope["selected_candidate_ref"]),
        ("推荐方案编号", data["recommendation"]["candidate_ref"]),
        ("范围起点（含）", scope["range_start"]),
        ("推荐版本", data["recommendation"]["recommendation_version"]),
        ("推荐记录格式版本", data["selection_basis"]["selection_policy_version"]),
        ("范围终点（不含）", scope["range_end"]),
        ("批次数", len(scope["batch_refs"])),
        ("工序数", len(scope["operation_refs"])),
        ("差值方向", "候选减查看基准"),
        ("资源视角", "本次所选批次" if scope["resource_view"] == "selected_batches" else "方案全部安排"),
        ("共同设备数", len(scope["resource_refs"]["machine"])),
        ("共同人员数", len(scope["resource_refs"]["operator"])),
        ("资源占用计算方式", "保存安排区间与归档可排窗口的交集；休息时段的保留跨度单列"),
        ("资源冲突计算方式", "归档可排窗口内同时占用同一资源的不同永久工序"),
        ("零容量工作证据", "安排跨度不证明实际工作；未保存工作片段的零容量时段留空，班表完整状态单列"),
        ("周占用比例计算方式", "所选周内班表内占用工时之和除以可排工时之和；日峰值和高负荷日保留"),
        ("小时筛选范围", "当前显示小时区间；完整比较窗口单列"),
        ("筛选页签", filter_section),
        ("筛选条件", description),
        ("明细范围", "筛选结果全部页"),
        ("标准工时计算方式", "保存的工时定额乘以其对应的加工目标数量"),
        ("所需工时计算方式", "计算时保存的准备和加工所需标准工时；原加工目标与本次报工目标数量分别保留"),
        ("工种切换计算方式", "同一设备上相邻工序的工种不同，后一道工序开始时间在所选范围内（含开始、不含结束）"),
        ("不计工种切换的工序", "外协工序，以及保存记录中明确的零工时工序"),
        ("准备工时统计范围", "本次所选批次；新安排、其他未安排、缺料未安排分别汇总；原安排、已开工和已完工工序分别列出"),
        ("准备差值方向", "候选减基准，中性展示实际需求与安排覆盖变化"),
        ("推荐方案比较范围", "计算时推荐的方案与当前查看基准比较；交期差值只比较双方工序都已排齐的批次"),
        ("推荐工序安排范围", "本次所选批次的工序安排及未安排工序，包含所选时间范围之外及明细其他页的工序"),
        ("推荐依据来源", "计算时保存的目标数值、比较方案、指标相同时的选择顺序、交期上限和关键工序时间评价；旧记录未保存的项留空"),
        ("过程范围", "保存本方案的整个计算过程，包含当前时间范围之外的计算尝试"),
        ("过程工序统计范围", "过程中的工序数、排齐批次数和首次排齐标记，按本次参与计算工序统计；业务未安排工序见工序安排和剩余需求。"),
        ("过程耗时单位", "毫秒；候选自身处理耗时与来源方案实际计算耗时分别保存"),
        ("已保存代表尝试", "按计算时的目标数值排名；时间、是否留作继续搜索和相对当时比较方案的变化按原记录保存，未保存的项留空"),
        ("沿用方案的计算次数", "候选自身未搜索时计数为 0，来源方案的计算次数和编号另列"),
        ("过程次数计算方式", "评估次数、留作继续搜索次数、严格改善次数和放弃次数分别统计；放弃记录含重复结果，各项互有重合"),
    ]
    rows = list(workbook["比较范围"].iter_rows())
    if len(rows) != len(wanted):
        raise RuntimeError("完整导出的比较范围字段数不同：期待 " + str(len(wanted)))
    numeric = blanks = 0
    for cells, (key, value) in zip(rows, wanted):
        if not 1 <= len(cells) <= 2 or cells[0].value != key or cells[0].data_type != "s":
            raise RuntimeError("完整导出的比较范围字段、顺序或原生文本类型不同：" + key)
        cell = cells[1] if len(cells) == 2 else None
        actual, kind = (cell.value, cell.data_type) if cell is not None else (None, None)
        # Write-only XLSX omits an unformatted trailing blank cell. A saved
        # blank B cell must have native blank/numeric type; absent B is blank.
        valid = ((cell is None or actual is None and kind == "n") if value is None else
                 cell is not None and _check_detail_export_cell(actual, value, "xlsx", key, kind))
        if not valid:
            raise RuntimeError("完整导出的比较范围值、口径文案或原生类型不同：" + key)
        numeric += type(value) in (int, float)
        blanks += value is None
    return {"rows": len(wanted), "columns": 2, "checked_cells": 2 * len(wanted),
            "numeric_cells": numeric, "boolean_cells": 0, "blank_cells": blanks,
            "snapshot_and_source_identities_match_api": True, "active_filters_match_api": True,
            "all_scope_fields_and_literal_wording_match": True, "ordered_fields_match": True,
            "native_datatypes_match": True}


def current_all_export_case(client, identity, fixture, query, target, cache=None,
                            filter_section=None, active=None, documents=None, archived_labels=None):
    """One all-XLSX download; complete field/type checks for 16 or 18 sheets.

    Existing all callers dispatch here without renaming their historical cases.
    documents can contain already obtained complete public business responses;
    a response is reused only for the same global scope and exact normalized
    leaf filters. Optional recommendation-dispositions is the P5 receipt oracle.
    No source object is given fabricated snapshot/as_of metadata.
    """
    target = Path(target)
    target.mkdir(parents=True, exist_ok=True)
    cache = {} if cache is None else cache
    supplied = dict(documents or {})
    raw_query = dict(query)
    run_ref = fixture["run_ref"]
    requested_section = filter_section or raw_query.get("filter_section")
    if requested_section is None and active is not None:
        requested_section = active["data"].get("filters", {}).get("section")
    if requested_section is not None and requested_section not in _CURRENT_ALL_FILTER_KEYS:
        raise ValueError("all 导出的活动筛选页签不受支持：" + str(requested_section))
    raw_filter_keys = set(raw_query) & {key for keys in _CURRENT_ALL_FILTER_KEYS.values() for key in keys}
    if raw_filter_keys and requested_section is None:
        raise ValueError("带明细筛选的 all 导出需要明确活动页签。")
    if requested_section is not None and raw_filter_keys - set(_CURRENT_ALL_FILTER_KEYS[requested_section]):
        raise ValueError("all 导出的原叶条件含不适用于活动页签的字段。")
    if active is not None and not _current_all_requested_scope_matches(active["data"]["scope"], raw_query):
        raise RuntimeError("all 导出的原 query 与真实 active 共同范围不同。")
    common = (_current_all_scope_query(active["data"]["scope"]) if active is not None else
              {key: raw_query[key] for key in _CURRENT_ALL_GLOBAL_KEYS if raw_query.get(key) is not None})
    path = API + "/runs/" + run_ref + "/decision"
    source_uses = []
    overview_hint = supplied.get("overview")
    if overview_hint is not None and _current_all_requested_scope_matches(overview_hint["data"]["scope"], common):
        overview = overview_hint
        _save(target / "sources" / "overview-all.json", overview)
        source_uses.append({"section": "overview", "source": "supplied public GET response"})
    else:
        overview = preparation_export_document(client, fixture, common, target / "sources", "overview", cache)
        source_uses.append({"section": "overview", "source": "public GET /decision"})
    if overview["data"]["run"]["run_ref"] != run_ref:
        raise RuntimeError("all 的公共总览不属于指定真实运行。")
    scope = overview["data"]["scope"]
    common = _current_all_scope_query(scope)
    if active is not None and active["data"]["scope"] != scope:
        raise RuntimeError("all 的 active 与总览共同范围不同。")

    def read_business(section, requested, directory):
        candidates = supplied.get(section, [])
        if not isinstance(candidates, list):
            candidates = [candidates]
        if active is not None:
            candidates = [active] + candidates
        for document in candidates:
            data = document["data"]
            wanted_filters = _current_all_default_filters(section, requested, data)
            if (data.get("scope") == scope and data.get("filters") == wanted_filters
                    and len(data.get("items", [])) == data.get("page", {}).get("total")):
                _save(directory / (section + "-all.json"), document)
                source_uses.append({"section": section, "query": requested, "source": "reused complete public GET response"})
                return document
        document = preparation_export_document(client, fixture, requested, directory, section, cache)
        wanted_filters = _current_all_default_filters(section, requested, document["data"])
        if document["data"]["scope"] != scope or document["data"]["filters"] != wanted_filters:
            raise RuntimeError("all 的公共明细范围或实际叶筛选不同：" + section)
        if len(document["data"]["items"]) != document["data"]["page"]["total"]:
            raise RuntimeError("all 的公共明细尚未取得完整页：" + section)
        source_uses.append({"section": section, "query": requested, "source": "public GET /decision/" + section})
        return document

    leaf_query = dict(common)
    if requested_section is not None:
        leaf_query.update({key: raw_query[key] for key in _CURRENT_ALL_FILTER_KEYS[requested_section]
                           if raw_query.get(key) is not None})
        if active is None:
            if raw_query.get("snapshot_ref") is not None:
                leaf_query["snapshot_ref"] = raw_query["snapshot_ref"]
            if requested_section == "process":
                active = client.document(path + "/process?" + urlencode(dict(leaf_query, page=1, size=20)))
                _save(target / "sources" / "process-active.json", active)
                source_uses.append({"section": "process", "source": "public GET /decision/process", "query": leaf_query})
            else:
                active = read_business(requested_section, leaf_query, target / "sources" / "active")
        active_filters = active["data"]["filters"]
        if active_filters["section"] != requested_section or active["data"]["scope"] != scope:
            raise RuntimeError("all 的活动叶来源与声明页签或共同范围不同。")
        for key in _CURRENT_ALL_FILTER_KEYS[requested_section]:
            if raw_query.get(key) is not None and raw_query[key] != active_filters.get(key):
                raise RuntimeError("all 的原叶条件与 active GET 不同：" + key)
        leaf_query = dict(common, **{key: value for key, value in active_filters.items()
                                    if key in _CURRENT_ALL_FILTER_KEYS[requested_section] and value is not None})
    else:
        active = overview if active is None else active
        active_filters = None

    business = {"overview": overview}
    paired_batch = leaf_query.get("batch_ref") if requested_section in ("batches", "operations") else None
    for section in ("batches", "operations"):
        current = leaf_query if section == requested_section else dict(common, **({"batch_ref": paired_batch} if paired_batch is not None else {}))
        business[section] = read_business(section, current, target / "sources" / section)
    if requested_section in ("resources", "resource-days"):
        resource_query = leaf_query if requested_section == "resources" else dict(common, **{
            key: active_filters[key] for key in ("resource_kind", "resource_ref", "grain", "pressure_threshold", "display_start", "display_end")
            if active_filters[key] is not None})
        business["resources"] = [read_business("resources", resource_query, target / "sources" / "resources")]
    else:
        business["resources"] = [read_business("resources", dict(common, resource_kind=kind), target / "sources" / kind)
                                 for kind in ("machine", "operator")]
    if requested_section == "resource-days":
        business["resource-days"] = [read_business("resource-days", leaf_query, target / "sources" / "resource-days")]
    paired_machine = leaf_query.get("machine_ref") if requested_section in ("preparation", "changeovers") else None
    for section in ("preparation", "changeovers"):
        current = leaf_query if section == requested_section else dict(common, **({"machine_ref": paired_machine} if paired_machine is not None else {}))
        business[section] = read_business(section, current, target / "sources" / section)

    # Expected common cells only: public overview.data plus real active.meta.
    # This object is never labelled or saved as an actual API response.
    expected_document = {"data": overview["data"], "meta": dict(active["meta"])}
    candidate_sources = recommendation_export_candidate_documents(
        client, fixture, common, overview, target / "recommendation-sources", cache)
    dispositions = supplied.get("recommendation-dispositions")
    if dispositions is None:
        dispositions = recommendation_export_dispositions(identity, run_ref)
    if dispositions["run_ref"] != run_ref:
        raise RuntimeError("all 的保存受理状态 oracle 属于其他运行。")
    recommendation_expected = recommendation_export_expected(expected_document, candidate_sources, dispositions)
    support_refs = recommendation_export_support_refs(overview, candidate_sources)
    process_query = leaf_query if requested_section == "process" else common
    anchor_ref = active["data"]["candidate"]["candidate_ref"] if requested_section == "process" else None
    saved_process = process_export_candidate_documents(client, fixture, process_query, overview,
        target / "process-sources", anchor_ref=anchor_ref)
    process_expected = process_export_expected(expected_document, saved_process)
    if tuple(RECOMMENDATION_EXPORT_HEADERS) != _CURRENT_ALL_RECOMMENDATION_TITLES or tuple(PROCESS_EXPORT_HEADERS) != _CURRENT_ALL_PROCESS_TITLES:
        raise RuntimeError("all 的独立推荐/过程表顺序与16/18表合同不同。")

    labels = dict(archived_labels or {})
    if archived_labels is None:
        for sources in candidate_sources.values():
            labels.update({row["operation_ref"]: {key: row.get(key) for key in ("batch_ref", "batch_label", "sequence", "process_label")}
                           for row in sources["preparation"]["data"]["items"]})
    coverage = _resource_coverage_expected(business, labels)
    conflicts = _resource_conflict_expected(business)
    overview_headers, overview_rows = preparation_export_overview_expected(overview["data"])
    row_counts = {"方案总览": len(overview_rows), "批次影响": len(business["batches"]["data"]["items"]),
                  "工序变化": len(business["operations"]["data"]["items"]),
                  "资源影响": sum(len(document["data"]["items"]) for document in business["resources"])}
    if "resource-days" in business:
        resource_period_title = {"day": "资源逐日", "hour": "资源小时", "shift": "资源班次", "week": "资源逐周"}[
            business["resource-days"][0]["data"]["filters"]["grain"]]
        row_counts.update({resource_period_title: sum(len(document["data"]["items"]) for document in business["resource-days"]),
                           "资源冲突": len(conflicts)})
    row_counts.update({"覆盖工序": len(coverage), "准备需求": len(business["preparation"]["data"]["items"]),
                       "工种切换": len(business["changeovers"]["data"]["items"])})
    row_counts.update({title: len(recommendation_expected[title]) for title in _CURRENT_ALL_RECOMMENDATION_TITLES})
    row_counts.update({title: len(process_expected[title]) for title in _CURRENT_ALL_PROCESS_TITLES})
    business_count = sum(row_counts.values())  # 指标计算方式和比较范围不计入下载业务行数。
    wanted_sheets = list(row_counts) + ["指标计算方式", "比较范围"]
    if len(wanted_sheets) != (18 if requested_section == "resource-days" else 16):
        raise RuntimeError("all 的独立表名单数量不同。")
    export_query = dict(common, section="all", snapshot_ref=active["meta"]["snapshot_ref"])
    if requested_section is not None:
        export_query.update({key: value for key, value in active_filters.items()
                             if key in _CURRENT_ALL_FILTER_KEYS[requested_section] and value is not None})
        export_query["filter_section"] = requested_section
    _save(target / "request.json", export_query)
    _save(target / "expected-public-records.json", {"object_kind": "independent expected projection only",
        "common_meta_source": "unchanged real active.meta", "recommendation": recommendation_expected,
        "process": process_expected, "business_rows_by_sheet": row_counts})
    _save(target / "sources.json", {"business": source_uses, "recommendation_dispositions": {
        "source": dispositions["source"],
        "sqlite_mode": "ro" if "SQLite mode=ro" in dispositions["source"] else None,
        "query_only": 1 if "SQLite mode=ro" in dispositions["source"] else None,
        "reused": "recommendation-dispositions" in supplied,
        "run_ref": run_ref, "fields": dispositions["fields"]},
        "coverage_labels": "caller supplied archive oracle" if archived_labels is not None else "full-admission public preparation GET rows",
        "process": "public GET /decision/process, complete retained rows", "process_anchor_ref": anchor_ref})
    content, headers = _fetch_export(client, path + "/export?" + urlencode(dict(export_query, format="xlsx")),
                                     target, "xlsx", section="all")
    if (headers.get("X-Workbench-Run-Ref") != run_ref
            or headers.get("X-Workbench-Snapshot-Ref") != active["meta"]["snapshot_ref"]
            or int(headers["X-Workbench-Row-Count"]) != business_count):
        raise RuntimeError("all 的运行、真实活动快照或完整业务行数不同。")

    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    checked = {}
    try:
        if workbook.sheetnames != wanted_sheets:
            raise RuntimeError("all 的16/18表顺序不同：" + str(workbook.sheetnames))
        checked["overview"] = preparation_export_check_records(workbook["方案总览"], overview_headers, overview_rows)
        for section in ("batches", "operations", "resources", "resource-days", "preparation", "changeovers"):
            if section not in business:
                continue
            rows, columns, types = _current_all_sheet_rows(workbook[resource_period_title if section == "resource-days" else _CURRENT_ALL_TITLES[section]])
            if section in ("resources", "resource-days"):
                checked[section] = _check_resource_export_table(business[section], rows, columns, "xlsx", types)
            elif section in ("preparation", "changeovers"):
                checked[section] = preparation_export_check_table(business[section], rows, columns, "xlsx", types)
            elif section == "operations":
                checked[section] = preparation_export_check_operations(business[section], rows, columns, "xlsx", types)
            else:
                checked[section] = _check_detail_export_table(business[section], rows, columns, "xlsx", types)
        checked["coverage-operations"] = preparation_export_check_records(workbook["覆盖工序"], RESOURCE_COVERAGE_HEADERS, coverage)
        if "resource-days" in business:
            checked["resource-conflicts"] = preparation_export_check_records(workbook["资源冲突"], RESOURCE_CONFLICT_HEADERS, conflicts)
        for title in _CURRENT_ALL_RECOMMENDATION_TITLES:
            rows, columns, types = _current_all_sheet_rows(workbook[title])
            checked[title] = recommendation_export_check_table(rows, columns, "xlsx", types, recommendation_expected[title], title)
        for title in _CURRENT_ALL_PROCESS_TITLES:
            rows, columns, types = _current_all_sheet_rows(workbook[title])
            checked[title] = process_export_check_table(rows, columns, "xlsx", types, process_expected[title], title)
        checked["metric-definitions"] = preparation_export_check_records(workbook["指标计算方式"], RESOURCE_METRIC_HEADERS,
            preparation_export_metric_expected(overview["data"]["metric_definitions"]))
        checked["comparison-scope"] = _current_all_check_scope(workbook, expected_document, requested_section, active_filters)
        checked["sheet_names"] = workbook.sheetnames
    finally:
        workbook.close()
    result = {"state": "passed", "run_ref": run_ref, "section": "all", "filter_section": requested_section,
        "comparison_mode": scope["comparison_mode"], "resource_view": scope["resource_view"],
        "scope": scope, "filters": active_filters, "snapshot_ref": active["meta"]["snapshot_ref"], "as_of": active["meta"]["as_of"],
        "business_row_count": business_count, "download_row_counts": {"xlsx": business_count},
        "business_rows_by_sheet": row_counts, "formats": {"xlsx": checked}, "source_refs": support_refs,
        "candidate_source_count": len(candidate_sources), "process_source_count": len(saved_process),
        "process_anchor_ref": anchor_ref, "all_sections_scope_same": True, "active_filter_preserved": True,
        "common_meta_uses_real_active_snapshot": True, "download_count": 1,
        "saved_recommendation_ref": overview["data"]["recommendation"]["candidate_ref"],
        "viewed_candidate_ref": scope["selected_candidate_ref"],
        "reference_candidate_ref": overview["data"]["baseline"].get("candidate_ref"),
        "formal_reference_version": overview["data"]["baseline"].get("version"),
        "recommendation_type": overview["data"]["recommendation"]["recommendation_type"],
        "selection_reason": overview["data"]["selection_basis"]["reason_code"],
        "recommendation_table_counts": {title: len(recommendation_expected[title]) for title in _CURRENT_ALL_RECOMMENDATION_TITLES},
        "process_table_counts": {title: len(process_expected[title]) for title in _CURRENT_ALL_PROCESS_TITLES},
        "runtime_identity": {"base_url": identity["base_url"], "pid": identity["pid"]},
        "actual_engine_rerun": False, "adoption_performed": False, "evidence_directory": str(target)}
    _save(target / "verification.json", result)
    return result, active


def _current_export_metric_expected(definitions):
    return _resource_metric_expected(definitions)


def process_export_verify(args):
    identity, client = _identity(args.instance)
    target = Path(args.evidence_directory).resolve()
    target.mkdir(parents=True, exist_ok=True)
    fixture = {"run_ref": args.run_ref}
    query = {"comparison_mode": "candidates", "selected_candidate_ref": args.selected_candidate_ref,
             "reference_candidate_ref": args.reference_candidate_ref, "range_start": args.range_start,
             "range_end": args.range_end, "resource_view": args.resource_view}
    process_query = dict(query, process_candidate_ref=args.process_candidate_ref,
                         record_kind="progress", sort="elapsed_ms", order="asc")
    path = API + "/runs/" + args.run_ref + "/decision"
    if args.resume_process_capture:
        if not args.process_only:
            raise ValueError("原过程捕获续验仅支持 --process-only。")
        overview = json.loads((target / "overview.json").read_text(encoding="utf-8"))
        if not _current_all_requested_scope_matches(overview["data"]["scope"], query):
            raise ValueError("原过程捕获的共同范围与本次请求不同。")
    else:
        overview = client.document(path + "?" + urlencode(query))
        _save(target / "overview.json", overview)
    cases = []
    try:
        process_target = target / "process"
        if args.resume_process_capture:
            saved = json.loads((process_target / "sources" / "public-process-candidates.json").read_text(encoding="utf-8"))
            if len(saved) != 1 or saved[0]["candidate"]["candidate_ref"] != args.process_candidate_ref:
                raise ValueError("原过程捕获的候选与本次请求不同。")
        else:
            saved = process_export_candidate_documents(client, fixture, process_query, overview,
                process_target / "sources", anchor_ref=args.process_candidate_ref)
        active = json.loads(Path(saved[0]["public_readbacks"]["progress"]).read_text(encoding="utf-8"))
        expected_document = {"data": overview["data"], "meta": dict(active["meta"])}
        expected = process_export_expected(expected_document, saved)
        _save(process_target / "expected-public-records.json", {"object_kind": "independent expected projection only",
            "common_meta_source": "unchanged real process.meta", "records": expected})
        count = sum(len(records) for records in expected.values())
        export_query = dict(process_query, section="process", snapshot_ref=active["meta"]["snapshot_ref"])
        if args.resume_process_capture and json.loads((process_target / "request.json").read_text(encoding="utf-8")) != export_query:
            raise ValueError("原过程捕获的导出请求与本次请求不同。")
        _save(process_target / "request.json", export_query)
        formats = {}
        for fmt in ("csv", "xlsx"):
            if args.resume_process_capture and fmt == "csv":
                content = (process_target / "decision-process.csv").read_bytes()
                headers = json.loads((process_target / "decision-process-headers-csv.json").read_text(encoding="utf-8"))
            else:
                content, headers = _fetch_export(client, path + "/export?" + urlencode(dict(export_query, format=fmt)),
                                                 process_target, fmt, section="process")
            if (headers.get("X-Workbench-Run-Ref") != args.run_ref
                    or headers.get("X-Workbench-Snapshot-Ref") != active["meta"]["snapshot_ref"]
                    or int(headers["X-Workbench-Row-Count"]) != count):
                raise RuntimeError("过程导出的运行、快照身份或四个业务粒度行数不同。")
            if fmt == "csv":
                reader = csv.DictReader(io.StringIO(content.decode("utf-8-sig")))
                rows = list(reader)
                ordered = [row for records in expected.values() for row in records]
                formats[fmt] = {"process-long-table": process_export_check_table(
                    rows, reader.fieldnames, fmt, None, ordered, "计算过程", csv_union=True)}
            else:
                from openpyxl import load_workbook

                workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
                try:
                    wanted_sheets = list(PROCESS_EXPORT_HEADERS) + ["指标计算方式", "比较范围"]
                    if workbook.sheetnames != wanted_sheets:
                        raise RuntimeError("过程XLSX的四个业务粒度及附表顺序不同。")
                    checked = {}
                    for title, records in expected.items():
                        rows, columns, types = _current_all_sheet_rows(workbook[title])
                        checked[title] = process_export_check_table(rows, columns, fmt, types, records, title)
                    checked["metric-definitions"] = preparation_export_check_records(workbook["指标计算方式"],
                        RESOURCE_METRIC_HEADERS, _current_export_metric_expected(overview["data"]["metric_definitions"]))
                    checked["comparison-scope"] = _current_all_check_scope(workbook, expected_document,
                        "process", active["data"]["filters"])
                    checked["sheet_names"] = workbook.sheetnames
                    formats[fmt] = checked
                finally:
                    workbook.close()
        process_case = {"state": "passed", "section": "process", "run_ref": args.run_ref,
            "process_candidate_ref": args.process_candidate_ref, "snapshot_ref": active["meta"]["snapshot_ref"],
            "as_of": active["meta"]["as_of"], "scope": active["data"]["scope"], "filters": active["data"]["filters"],
            "business_row_count": count, "business_rows_by_sheet": {title: len(rows) for title, rows in expected.items()},
            "formats": formats, "download_count": 2, "new_download_count": 1 if args.resume_process_capture else 2,
            "original_csv_reused": args.resume_process_capture, "evidence_directory": str(process_target)}
        _save(process_target / "verification.json", process_case)
        cases.append(process_case)
        if not args.process_only:
            all_case, _ = current_all_export_case(client, identity, fixture, query, target / "all",
                                                 documents={"overview": overview})
            cases.append(all_case)
    except Exception as error:
        _save(target / "export-verification.json", {"state": "failed", "run_ref": args.run_ref,
            "error": str(error), "cases": cases})
        raise
    checked_cells = sum(table.get("checked_cells", 0) for case in cases for tables in case["formats"].values()
                        for table in tables.values() if isinstance(table, dict))
    result = {"state": "passed", "base_url": identity["base_url"], "pid": identity["pid"],
        "database_path": identity["database_path"], "run_ref": args.run_ref, "case_count": len(cases),
        "download_count": sum(case["download_count"] for case in cases),
        "process_only": args.process_only,
        "original_csv_reused": args.resume_process_capture,
        "new_download_count": sum(case.get("new_download_count", case["download_count"]) for case in cases),
        "business_row_count": sum(case["business_row_count"] * case["download_count"] for case in cases),
        "checked_cells": checked_cells, "cases": cases, "actual_engine_rerun": False, "adoption_performed": False,
        "evidence_directory": str(target)}
    _save(target / "export-verification.json", result)
    print(json.dumps({key: value for key, value in result.items() if key != "cases"}, ensure_ascii=False), flush=True)


def _check_arrangements(items, selected_tasks, reference_tasks):
    for side, tasks in (("after", selected_tasks), ("before", reference_tasks)):
        expected = {row["operation_ref"]: row for row in tasks}
        for item in items:
            value, task = item[side], expected.get(item["operation_ref"])
            if value is None or task is None:
                if value is not None or task is not None:
                    raise RuntimeError("工序明细的新增/移除状态与真实候选任务不同。")
                continue
            for key in ("start", "end", "source"):
                if value[key] != task[key]:
                    raise RuntimeError("工序明细与真实候选安排不同：" + key)
            for key in ("machine", "operator"):
                if (value[key] or {}).get("ref") != (task[key] or {}).get("ref"):
                    raise RuntimeError("工序明细与真实候选资源不同：" + key)


def _check_archive_after_current_edits(identity, client, path, query, overview, batches, operations, target):
    from core.infrastructure.database import get_connection

    conn = get_connection(identity["database_path"])
    batch = tuple(conn.execute("SELECT due_date,part_name FROM Batches WHERE batch_id='DD-D-01'").fetchone())
    setup = conn.execute("SELECT setup_hours FROM BatchOperations WHERE op_code='DD-D-01-1'").fetchone()[0]
    name = conn.execute("SELECT name FROM Machines WHERE machine_id='DD-D-M'").fetchone()[0]
    try:
        conn.execute("UPDATE Batches SET due_date='9999-12-31',part_name='现场修改后零件名' WHERE batch_id='DD-D-01'")
        conn.execute("UPDATE BatchOperations SET setup_hours=9 WHERE op_code='DD-D-01-1'")
        conn.execute("UPDATE Machines SET name='现场修改后设备名' WHERE machine_id='DD-D-M'")
        conn.commit()
        requests = (("overview", path, overview, {}),
                    ("batches", path + "/batches", batches, {"sort": "batch_id", "order": "asc", "size": 50}),
                    ("operations", path + "/operations", operations, {"sort": "sequence", "order": "asc", "size": 50}))
        checked = []
        for section, endpoint, document, filters in requests:
            changed = client.document(endpoint + "?" + urlencode(dict(query, **filters,
                snapshot_ref=document["meta"]["snapshot_ref"])))
            if changed["data"] != document["data"]:
                raise RuntimeError("当前资料修改改变了历史归档比较：" + section)
            checked.append(section)
    finally:
        conn.execute("UPDATE Batches SET due_date=?,part_name=? WHERE batch_id='DD-D-01'", batch)
        conn.execute("UPDATE BatchOperations SET setup_hours=? WHERE op_code='DD-D-01-1'", (setup,))
        conn.execute("UPDATE Machines SET name=? WHERE machine_id='DD-D-M'", (name,))
        conn.commit()
        restored = (tuple(conn.execute("SELECT due_date,part_name FROM Batches WHERE batch_id='DD-D-01'").fetchone()),
                    conn.execute("SELECT setup_hours FROM BatchOperations WHERE op_code='DD-D-01-1'").fetchone()[0],
                    conn.execute("SELECT name FROM Machines WHERE machine_id='DD-D-M'").fetchone()[0])
        conn.close()
        if restored != (batch, setup, name):
            raise RuntimeError("归档验证后新增 DD 记录的当前字段未恢复。")
    _save(target / "current-fields-archive-check.json", {"state": "passed", "sections": checked,
        "reused_original_snapshots": True, "current_fields_restored": True,
        "changed_fields": ["DD-D-01.due_date", "DD-D-01.part_name", "DD-D-01-1.setup_hours", "DD-D-M.name"]})


def delivery_verify(args):
    identity, client = _identity(args.instance)
    evidence = ROOT / "evidence" / "decision-delivery-impact" / Path(identity["root"]).name
    manifest = json.loads((evidence / "delivery-fixtures.json").read_text(encoding="utf-8"))
    if Path(manifest["database_path"]).resolve() != Path(identity["database_path"]).resolve():
        raise RuntimeError("交付验收夹具与当前隔离实例不同。")
    results = []
    for name in args.scenario:
        fixture = manifest["fixtures"][name]
        path = API + "/runs/" + fixture["run_ref"] + "/decision"
        selected, reference = fixture["candidate_refs"][-1], fixture["candidate_refs"][0]
        query = {"comparison_mode": "candidates", "selected_candidate_ref": selected,
                 "reference_candidate_ref": reference}
        target = evidence / name / "verification"
        overview = client.document(path + "?" + urlencode(query))
        _save(target / "overview.json", overview)
        batches = _detail_pages(client, path, dict(query, sort="batch_id", order="asc"), target, "batches")
        operations = _detail_pages(client, path, dict(query, sort="sequence", order="asc"), target, "operations")
        if {row["batch_ref"] for row in batches["data"]["items"]} != set(fixture["batch_refs"]):
            raise RuntimeError("逐批明细身份与真实夹具批次集合不同。")
        if {row["operation_ref"] for row in operations["data"]["items"]} != set(fixture["operation_refs"]):
            raise RuntimeError("工序明细身份与真实夹具全工序集合不同。")
        selected_workspace = client.document(API + "/candidates/" + selected + "/workspace")
        reference_workspace = client.document(API + "/candidates/" + reference + "/workspace")
        _save(target / "selected-workspace.json", selected_workspace)
        _save(target / "reference-workspace.json", reference_workspace)
        _check_arrangements(operations["data"]["items"], selected_workspace["data"]["tasks"], reference_workspace["data"]["tasks"])
        candidates = {row["candidate_ref"]: row for row in overview["data"]["candidates"]}
        after = candidates[selected]
        if after["coverage"]["search"]["state"] != "complete":
            raise RuntimeError("真实夹具的候选搜索覆盖应完整。")
        if name == "mixed_due":
            summary = batches["data"]["summary"]["after"]
            counts = {key: summary[key] for key in ("due_eligible_count", "excluded_due_count", "invalid_due_count",
                                                    "due_evaluation_count", "known_due_count", "unknown_due_count")}
            if counts != dict(due_eligible_count=18, excluded_due_count=1, invalid_due_count=2,
                              due_evaluation_count=20, known_due_count=18, unknown_due_count=2):
                raise RuntimeError("混合交期的三组数量与已知/未知范围不同：" + str(counts))
            rows = {row["batch_id"]: row for row in batches["data"]["items"]}
            if (rows["DD-D-19"]["after"]["due_state"] != "unspecified"
                    or rows["DD-D-19"]["after"]["risk"] != "not_applicable"
                    or rows["DD-D-19"]["after"]["delay_hours"] is not None
                    or rows["DD-D-20"]["after"]["due_issue"] != "due_date_missing"
                    or rows["DD-D-21"]["after"]["due_issue"] != "due_date_missing"
                    or after["metrics"]["total_tardiness_hours"]["value"] is not None
                    or after["metrics"]["total_tardiness_hours"]["known_subtotal"] <= 0):
                raise RuntimeError("未指定、非法、缺失交期或总拖期未知值表达不同。")
            _check_archive_after_current_edits(identity, client, path, query, overview, batches, operations, target)
        elif name == "ordinary_partial":
            before_tasks, after_tasks = reference_workspace["data"]["tasks"], selected_workspace["data"]["tasks"]
            if len(before_tasks) != 1 or len(after_tasks) != 2:
                raise RuntimeError("真实普通漏排夹具没有产生 1 道与 2 道的不同工作集合。")
            if (after["coverage"]["schedulable"]["state"] != "partial"
                    or after["deltas"]["makespan_hours"]["state"] != "different_coverage"
                    or after["deltas"]["makespan_hours"]["value"] is not None):
                raise RuntimeError("不同漏排集合仍显示可直接比较的完工跨度。")
        else:
            row = batches["data"]["items"][0]["after"]
            if (after["coverage"]["schedulable"]["state"] != "complete"
                    or after["coverage"]["schedulable"]["deferred_count"] != 1
                    or row["completion_state"] != "deferred" or row["schedule_complete"] is not False
                    or row["deferred_operation_count"] != 1 or row["ordinary_unscheduled_operation_count"] != 0
                    or row["planned_finish"] is not None or row["delay_hours"] is not None):
                raise RuntimeError("合法待料的可排覆盖与全链交付状态混合。")
        filtered = client.document(path + "/operations?" + urlencode(dict(query, resource_ref=fixture["machine_ref"],
            sort="sequence", order="asc", size=50)))
        if filtered["data"]["summary"] != operations["data"]["summary"]:
            raise RuntimeError("资源明细筛选改变了全范围交付汇总。")
        _save(target / "operations-resource-filter.json", filtered)
        results.append({"fixture": name, "run_ref": fixture["run_ref"], "state": "passed",
            "batch_count": len(batches["data"]["items"]), "operation_count": len(operations["data"]["items"]),
            "selected_candidate_ref": selected, "reference_candidate_ref": reference,
            "selected_task_count": len(selected_workspace["data"]["tasks"]),
            "reference_task_count": len(reference_workspace["data"]["tasks"])})
    result = {"state": "passed", "base_url": identity["base_url"], "pid": identity["pid"],
        "database_path": identity["database_path"], "fixtures": results, "evidence_directory": str(evidence)}
    _save(evidence / "delivery-verification.json", result)
    print(json.dumps(result, ensure_ascii=False), flush=True)


def verify(args):
    identity, client = _identity(args.instance)
    evidence = Path(identity["evidence_directory"])
    base = API + "/runs/" + identity["run_ref"] + "/decision"
    default = client.document(base)
    _save(evidence / "decision-default.json", default)
    data = default["data"]
    reference = args.reference_candidate_ref or identity["candidate_refs"][-1]
    selected = args.selected_candidate_ref or identity["candidate_refs"][0]
    if selected not in identity["candidate_refs"]:
        raise RuntimeError("所选候选引用与本次隔离排产不同。")
    if reference not in identity["candidate_refs"]:
        raise RuntimeError("基准候选引用与本次隔离排产不同。")
    query = {"reference_candidate_ref": reference, "selected_candidate_ref": selected}
    changed = client.document(base + "?" + urlencode(query))
    _save(evidence / "decision-changed-reference.json", changed)
    if changed["data"]["baseline"]["candidate_ref"] != reference:
        raise RuntimeError("切换比较基准之后，响应基准引用没有跟随本次选择。")
    if changed["data"]["selection_basis"] != data["selection_basis"]:
        raise RuntimeError("切换查看基准改变了保存的算法选择依据。")
    expected_rows = len(data["candidates"]) * len(data["metric_definitions"])
    export_query = dict(query, snapshot_ref=changed["meta"]["snapshot_ref"], section="overview")
    for fmt in ("csv", "xlsx"):
        content, headers = _fetch_export(client, base + "/export?" + urlencode(dict(export_query, format=fmt)), evidence, fmt)
        if (headers.get("X-Workbench-Run-Ref") != identity["run_ref"]
                or headers.get("X-Workbench-Snapshot-Ref") != export_query["snapshot_ref"]
                or int(headers["X-Workbench-Row-Count"]) != expected_rows):
            raise RuntimeError("总览导出的运行身份或行数与比较结果不同。")
        if fmt == "csv":
            rows = list(csv.DictReader(io.StringIO(content.decode("utf-8-sig"))))
            if len(rows) != expected_rows:
                raise RuntimeError("CSV 长表行数与同次全部候选的指标数不同。")
            _check_export_rows(changed["data"], rows, fmt)
        else:
            from openpyxl import load_workbook

            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            try:
                if workbook.sheetnames != ["方案总览", "指标计算方式", "比较范围"]:
                    raise RuntimeError("XLSX 工作表与总览导出约定不同。")
                iterator = workbook["方案总览"].iter_rows(values_only=True)
                headings = next(iterator)
                rows = [dict(zip(headings, row)) for row in iterator]
                if len(rows) != expected_rows:
                    raise RuntimeError("XLSX 长表行数与同次全部候选的指标数不同。")
                _check_export_rows(changed["data"], rows, fmt)
            finally:
                workbook.close()
    candidate_path = API + "/candidates/" + selected
    workspace = client.document(candidate_path + "/workspace")
    _save(evidence / "selected-candidate-workspace.json", workspace)
    result = {"state": "compared", "base_url": identity["base_url"], "database_path": identity["database_path"],
              "run_ref": identity["run_ref"], "candidate_count": len(data["candidates"]),
              "metric_count": len(data["metric_definitions"]), "export_row_count": expected_rows,
              "selected_candidate_ref": selected, "reference_candidate_ref": reference,
              "evidence_directory": str(evidence)}
    if args.skip_adoption:
        _save(evidence / "comparison-verification.json", result)
        print(json.dumps(result, ensure_ascii=False), flush=True)
        return
    if args.official_plan_ref:
        plan = {"plan_ref": args.official_plan_ref}
    else:
        preview = client.post(candidate_path + "/adopt-preview", {})["data"]
        if preview["validation"]["can_adopt"] is not True:
            raise RuntimeError("隔离候选采用预检未通过：" + str(preview["validation"]))
        adopted = client.command(candidate_path + "/adopt", preview["write_context"],
                                 {"confirm": True, "reason": "第 1 阶段隔离决策闭环验收", "declared_operator": "验收计划员"})
        _save(evidence / "adoption-receipt.json", adopted)
        plan = adopted["data"]["official_plan"]
    formal = client.document("/api/workbench/v1/plans/" + plan["plan_ref"] + "/workspace")
    _save(evidence / "adopted-official-workspace.json", formal)
    provenance = formal["data"]["provenance"]
    if provenance["source_ref"] != selected or provenance["run_ref"] != identity["run_ref"]:
        raise RuntimeError("正式计划采用来源与本次所选候选不同。")
    candidate_tasks = sorted((row["operation_ref"], row["start"], row["end"], row["quantity"],
        (row["machine"] or {}).get("ref"), (row["operator"] or {}).get("ref"), (row["supplier"] or {}).get("ref"))
        for row in workspace["data"]["tasks"])
    formal_tasks = sorted((row["operation_ref"], row["start"], row["end"], row["quantity"],
        row["machine_ref"], row["operator_ref"], row["supplier_ref"]) for row in formal["data"]["tasks"])
    if candidate_tasks != formal_tasks:
        raise RuntimeError("正式计划的工序、数量、时间或资源字段与所选候选不同。")
    result.update(state="passed", official_plan={"plan_ref": plan["plan_ref"], "version": formal["data"]["plan"]["version"]},
                  task_count=len(formal_tasks))
    _save(evidence / "verification.json", result)
    print(json.dumps(result, ensure_ascii=False), flush=True)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    commands = parser.add_subparsers(dest="command", required=True)
    host = commands.add_parser("serve", help="创建真实隔离实例并持续提供浏览器访问")
    host.add_argument("--root", default=str(DEFAULT_ROOT))
    host.add_argument("--instance", help="重启同一隔离数据库，沿用既有排产与正式采用结果")
    host.set_defaults(function=serve)
    check = commands.add_parser("verify", help="读取比较、导出并采用到指定隔离实例")
    check.add_argument("--instance", required=True)
    check.add_argument("--selected-candidate-ref", help="沿用浏览器实际选中的候选引用")
    check.add_argument("--reference-candidate-ref", help="沿用浏览器实际查看的候选基准")
    adoption = check.add_mutually_exclusive_group()
    adoption.add_argument("--skip-adoption", action="store_true", help="完成比较与导出后保留候选供浏览器采用")
    adoption.add_argument("--official-plan-ref", help="核对浏览器已采用的正式计划，复用实际采用结果")
    check.set_defaults(function=verify)
    resources = commands.add_parser("resource-fixtures", help="在同一隔离实例追加一次真实资源与夜班日历夹具")
    resources.add_argument("--instance", required=True)
    resources.set_defaults(function=resource_fixtures)
    resource_check = commands.add_parser("resource-verify", help="读取真实班段与安排，独立核对资源窗口、CV和逐日热力图")
    resource_check.add_argument("--instance", required=True)
    resource_check.set_defaults(function=resource_verify)
    preparation = commands.add_parser("preparation-fixtures", help="追加一次真实准备、分件、外协、冻结和漏排/暂缓夹具")
    preparation.add_argument("--instance", required=True)
    preparation_scope = preparation.add_mutually_exclusive_group()
    preparation_scope.add_argument("--complete-piece-after-failure", action="store_true",
        help="保留原失败20工序运行，用现有资料排除普通漏排批次，仅新受理一次合法19工序分件场景")
    preparation_scope.add_argument("--complete-formal-scope", action="store_true",
        help="保留前两次失败，用已核对31工序完整正式范围和原10月30日截止，仅受理一次新运行")
    preparation.set_defaults(function=preparation_fixtures)
    preparation_check = commands.add_parser("preparation-verify", help="逐收据账与完整设备序列核对准备需求、分件和窗口换型")
    preparation_check.add_argument("--instance", required=True)
    preparation_check.set_defaults(function=preparation_verify)
    fixtures = commands.add_parser("delivery-fixtures", help="在同一隔离实例追加交期、真实漏排和合法待料夹具")
    fixtures.add_argument("--instance", required=True)
    fixtures.add_argument("--scenario", nargs="+", choices=("mixed_due", "ordinary_partial", "legal_stage"),
                          default=("mixed_due", "ordinary_partial", "legal_stage"))
    fixtures.set_defaults(function=delivery_fixtures)
    delivery = commands.add_parser("delivery-verify", help="真实交期、漏排、待料与明细分页验证")
    delivery.add_argument("--instance", required=True)
    delivery.add_argument("--scenario", nargs="+", choices=("mixed_due", "ordinary_partial", "legal_stage"),
                          default=("mixed_due", "ordinary_partial", "legal_stage"))
    delivery.set_defaults(function=delivery_verify)
    exports = commands.add_parser("delivery-export-verify", help="读取既有真实夹具并逐字段验收明细导出")
    exports.add_argument("--instance", required=True)
    exports.set_defaults(function=delivery_export_verify)
    resource_exports = commands.add_parser("resource-export-verify", help="逐字段验收既有真实资源与单资源逐日导出")
    resource_exports.add_argument("--instance", required=True)
    resource_exports.add_argument("--evidence-directory", help="另存本次导出原始文件与结果，保留历史证据")
    resource_exports.add_argument("--case", nargs="+", choices=("new-eleven-machines", "new-operator-days",
        "old8-unused-positive", "old8-zero-capacity", "all-candidates-six-tables", "all-formal-full-filter",
        "all-operator-seven-tables"), help="只验收指定的既有案例")
    resource_exports.set_defaults(function=resource_export_verify)
    preparation_exports = commands.add_parser("preparation-export-verify", help="逐字段验收保存的准备需求与实际工种切换导出")
    preparation_exports.add_argument("--instance", required=True)
    preparation_exports.add_argument("--fixture", help="明确指定已生成的准备需求夹具记录")
    preparation_exports.add_argument("--evidence-directory", help="另存本次原始导出和结果")
    preparation_exports.add_argument("--case", nargs="+", choices=("new-runtime-preparation", "same-machine-window-changeovers",
        "all-preparation-batch-filter", "all-changeovers-selected-filter"), help="只验收指定的既定案例")
    preparation_exports.set_defaults(function=preparation_export_verify)
    preparation_days = commands.add_parser("preparation-resource-days-export-verify", help="复用四案例，仅补一份单人员逐日全部十八表XLSX")
    preparation_days.add_argument("--instance", required=True)
    preparation_days.add_argument("--fixture", help="明确指定已生成的准备需求夹具记录")
    preparation_days.add_argument("--evidence-directory", help="已有四案例原始导出的证据目录")
    preparation_days.set_defaults(function=preparation_export_resource_days_verify)
    recommendation_exports = commands.add_parser("recommendation-export-verify", help="独立核对三个代表运行的推荐导出及一份全部十六表XLSX")
    recommendation_exports.add_argument("--instance", required=True)
    recommendation_exports.add_argument("--evidence-directory", help="本次推荐原始导出和独立结果目录")
    recommendation_exports.add_argument("--case", nargs="+", choices=("balanced-recommendation", "ordinary-partial-recommendation",
        "mixed-due-recommendation", "all-balanced-twelve-tables"), help="只验收指定的既定案例")
    recommendation_exports.set_defaults(function=recommendation_export_verify)
    process_exports = commands.add_parser("process-export-verify", help="只读核对保存过程CSV、过程六表及全部十六表XLSX原文件")
    for argument in ("instance", "run-ref", "selected-candidate-ref", "reference-candidate-ref", "process-candidate-ref",
                     "range-start", "range-end", "evidence-directory"):
        process_exports.add_argument("--" + argument, required=True)
    process_exports.add_argument("--resource-view", choices=("selected_batches", "full_candidate"), default="selected_batches")
    process_exports.add_argument("--process-only", action="store_true", help="仅下载保存过程CSV与六表XLSX，复用既有全部导出证据")
    process_exports.add_argument("--resume-process-capture", action="store_true", help="复用同一公开快照已下载的原CSV，仅补XLSX并完成逐字段续验")
    process_exports.set_defaults(function=process_export_verify)
    args = parser.parse_args()
    args.function(args)


if __name__ == "__main__":
    main()
