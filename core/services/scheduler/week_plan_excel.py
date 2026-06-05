from __future__ import annotations

import time
from io import BytesIO
from typing import Any, Mapping, Optional, Sequence

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font

from core.services.common.excel_templates import build_xlsx_bytes


def build_week_plan_export_workbook(
    rows: Sequence[Mapping[str, Any]],
    *,
    plan_resolution: Optional[Mapping[str, Any]] = None,
    export_context: Optional[Mapping[str, Any]] = None,
) -> BytesIO:
    headers = ["日期", "批次号", "图号", "工序", "设备", "人员", "时段"]
    output = build_xlsx_bytes(
        headers,
        [[r.get(h, "") for h in headers] for r in rows],
        format_spec={
            "date_cols": [0],
            "text_cols": [1, 2, 4, 5, 6],
            "int_cols": [3],
            "column_widths": {0: 12, 1: 14, 2: 14, 3: 10, 4: 14, 5: 14, 6: 18},
        },
        sheet_title="周计划",
        sanitize_formula=True,
    )
    if not _needs_summary_sheet(plan_resolution or {}):
        return output
    return _with_plan_summary(output, plan_resolution=plan_resolution or {}, export_context=export_context or {})


def _needs_summary_sheet(plan_resolution: Mapping[str, Any]) -> bool:
    label = _plan_label(plan_resolution)
    return bool(plan_resolution.get("is_scenario_preview")) or (bool(label) and label != "正式采用方案")


def _plan_label(plan_resolution: Mapping[str, Any]) -> str:
    return _text(plan_resolution.get("user_label") or plan_resolution.get("selected_label") or plan_resolution.get("requested_label"))


def _with_plan_summary(
    output: BytesIO,
    *,
    plan_resolution: Mapping[str, Any],
    export_context: Mapping[str, Any],
) -> BytesIO:
    output.seek(0)
    wb = load_workbook(output)
    try:
        ws = wb.create_sheet("查询摘要", 0)
        for key, value in _plan_summary_rows(plan_resolution, export_context):
            ws.append([key, value])
        ws.freeze_panes = "A2"
        ws.column_dimensions["A"].width = 18
        ws.column_dimensions["B"].width = 42
        for row in ws.iter_rows():
            row[0].font = Font(bold=True)
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        result = BytesIO()
        wb.save(result)
        result.seek(0)
        return result
    finally:
        wb.close()


def _plan_summary_rows(
    plan_resolution: Mapping[str, Any],
    export_context: Mapping[str, Any],
) -> Sequence[Sequence[Any]]:
    week_start = _text(export_context.get("week_start"))
    week_end = _text(export_context.get("week_end"))
    plan_label = _plan_label(plan_resolution) or "正式采用方案"
    if not bool(plan_resolution.get("is_scenario_preview")):
        return [
            ["导出类型", "周计划导出"],
            ["方案", plan_label],
            ["提示", _plan_summary_hint(plan_resolution, plan_label)],
            ["排产版本", _text(export_context.get("version"))],
            ["周计划日期", f"{week_start} ～ {week_end}".strip()],
            ["导出时间", time.strftime("%Y-%m-%d %H:%M:%S")],
        ]
    scenario_label = _text(plan_resolution.get("scenario_display_name") or plan_resolution.get("scenario_name")) or "模拟预览（未命名）"
    return [
        ["导出类型", "模拟方案预览"],
        ["提示", "这是模拟方案预览，正式计划还没有改变。"],
        ["模拟方案", scenario_label],
        ["预览依据版本", _text(export_context.get("version"))],
        ["预览依据方案", _text(plan_resolution.get("selected_label") or plan_resolution.get("requested_label"))],
        ["周计划日期", f"{week_start} ～ {week_end}".strip()],
        ["导出时间", time.strftime("%Y-%m-%d %H:%M:%S")],
    ]


def _text(value: Any) -> str:
    return str(value or "").strip()


def _plan_summary_hint(plan_resolution: Mapping[str, Any], plan_label: str) -> str:
    readonly_reason = _text(plan_resolution.get("readonly_reason"))
    if readonly_reason:
        return readonly_reason
    if plan_label.startswith("历史正式方案"):
        return "这是历史版本，只能查看，不能当作当前可执行的正式计划。"
    return "这是查询时选择的排产方案。"
