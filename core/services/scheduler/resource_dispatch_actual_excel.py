from __future__ import annotations

import io
from datetime import date, datetime
from typing import Any, Dict, List, Sequence, cast

import openpyxl
from openpyxl import Workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.worksheet import Worksheet

from core.infrastructure.errors import AppError, ErrorCode, ValidationError
from core.services.common.excel_templates import _sanitize_export_cell

from .resource_dispatch_actual_records import (
    PAUSE_DETAIL_SHEET,
    PAUSE_HEADERS,
    TASK_FEEDBACK_SHEET,
    TASK_HEADERS,
    TaskRef,
    text,
)


def _excel_value(value: Any) -> Any:
    return _sanitize_export_cell(value)


def _normalize_cell_value(value: Any) -> Any:
    """把 openpyxl 读到的日期单元格归一成稳定字符串。

    ``data_only=True`` 会把日期读成 datetime/date 对象。这些对象一旦经接口下发再回传
    （预览→确认两步导入），序列化格式会改变，导致两次算出的幂等 token 对不上。在读取源头
    就把它们定型为 ``%Y-%m-%d %H:%M:%S`` / ``%Y-%m-%d`` 字符串，全链路类型一致。数字等其它
    类型保持原值，避免影响完成数量等字段“必须是整数”的校验语义。
    """
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    return value


def _append_row(ws: Worksheet, values: Sequence[Any]) -> None:
    ws.append([_excel_value(value) for value in values])


def _style_table(ws: Worksheet) -> None:
    ws.freeze_panes = "A2"
    for cell_obj in ws[1]:
        cell = cast(Cell, cell_obj)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell_obj in row:
            cell = cast(Cell, cell_obj)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for column_cells in ws.columns:
        first = cast(Cell, column_cells[0])
        col = first.column_letter
        width = max(len(str(cell.value or "")) for cell in column_cells) + 2
        ws.column_dimensions[col].width = min(max(width, 12), 34)


def build_actual_template_workbook(tasks: Sequence[TaskRef]) -> io.BytesIO:
    wb = Workbook()
    ws_task = cast(Worksheet, wb.active)
    ws_task.title = TASK_FEEDBACK_SHEET
    _append_row(ws_task, TASK_HEADERS)
    for task in tasks:
        _append_row(
            ws_task,
            [
                task.task_code,
                task.batch_id,
                task.op_name,
                task.planned_machine_label,
                task.planned_operator_label,
                "",
                "",
                "",
                "",
                "",
                "",
                "一般",
                "",
                "",
                "",
            ],
        )
    _style_table(ws_task)

    ws_pause = cast(Worksheet, wb.create_sheet(PAUSE_DETAIL_SHEET))
    _append_row(ws_pause, PAUSE_HEADERS)
    _style_table(ws_pause)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def read_actual_workbook_rows(file_bytes: bytes) -> List[Dict[str, Any]]:
    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    except Exception as exc:
        raise AppError(ErrorCode.EXCEL_READ_ERROR, "读取 Excel 文件失败，请确认文件未损坏且未被其他程序占用。", cause=exc) from exc
    try:
        rows: List[Dict[str, Any]] = []
        found_target_sheet = False
        for sheet_name in (TASK_FEEDBACK_SHEET, PAUSE_DETAIL_SHEET):
            if sheet_name not in wb.sheetnames:
                continue
            found_target_sheet = True
            ws = wb[sheet_name]
            values = list(ws.iter_rows(values_only=True))
            if not values:
                continue
            headers = [text(value) for value in values[0]]
            if "任务识别码" not in headers:
                raise ValidationError(f"“{sheet_name}”缺少任务识别码列，请使用下载的填写模板。", field="file")
            for row_number, raw in enumerate(values[1:], start=2):
                if all(not text(value) for value in raw):
                    continue
                item: Dict[str, Any] = {"sheet": sheet_name, "row_number": row_number}
                for index, header in enumerate(headers):
                    if not header:
                        continue
                    item[header] = _normalize_cell_value(raw[index]) if index < len(raw) else None
                rows.append(item)
        if not found_target_sheet:
            raise ValidationError("请使用下载的现场实际情况填写模板，模板里需要有“任务反馈”和“暂停明细”两个工作表。", field="file")
        return rows
    finally:
        wb.close()


__all__ = ["build_actual_template_workbook", "read_actual_workbook_rows"]
