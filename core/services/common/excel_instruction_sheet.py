"""把表描述协议落成工作簿里的第二张「填写说明」表和数据表的枚举下拉。

12 张导入表共用这一份写法。两派写出方式（openpyxl 只写工作簿与普通工作簿）都走同一个入口，
差别只有一处：只写表没有 `add_data_validation`，下拉只能直接挂到集合上。
"""

import os
from typing import Any, Dict, List

from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from core.models.workbench_table_descriptor import (
    INSTRUCTION_SHEET,
    enum_columns,
    instruction_rows,
)

#: 模板里没有数据行，下拉也要能用，所以往下多盖这么多行。
DROPDOWN_ROWS = 1000
_WIDTHS = (("A", 22), ("B", 14), ("C", 56), ("D", 48))


def append_instruction_sheet(wb, descriptor: Dict[str, Any]) -> Any:
    """在工作簿末尾追加填写说明表。返回这张表，由调用方按本家族的写法收尾。"""
    ws = wb.create_sheet(INSTRUCTION_SHEET)
    for letter, width in _WIDTHS:
        ws.column_dimensions[letter].width = width
    for row in instruction_rows(descriptor):
        ws.append(row)
    return ws


def add_enum_dropdowns(ws, descriptor: Dict[str, Any], row_count: int, *, write_only: bool) -> None:
    """给可填的枚举列挂下拉，覆盖已有数据行再往下留出可填区。"""
    last = max(row_count, 0) + 1 + DROPDOWN_ROWS
    for index, choices in enum_columns(descriptor):
        letter = get_column_letter(index)
        rule = DataValidation(type="list", formula1='"' + ",".join(choices) + '"', allow_blank=True)
        rule.add(letter + "2:" + letter + str(last))
        if write_only:
            # 只写表在保存前没有 add_data_validation，直接挂集合是 openpyxl 3.0.10 唯一可行的路径。
            ws.data_validations.dataValidation.append(rule)
        else:
            ws.add_data_validation(rule)


def close_write_only_sheets(sheets: List[Any]) -> None:
    """只写工作簿的收尾：每张表各有自己的临时写出器，都要关掉并删掉临时文件。"""
    for ws in sheets:
        if not ws.closed:
            ws.close()
        if ws._writer is not None and os.path.exists(ws._writer.out):
            ws._writer.cleanup()
