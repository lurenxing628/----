"""Bytes-only first-sheet XLSX, explicit formulas rejection and typed output text."""

import re
import zipfile
from io import BytesIO

import openpyxl
from openpyxl.cell.cell import TYPE_STRING
from openpyxl.comments import Comment
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from core.errors import ValidationError
from core.models.workbench_batch_file import HEADERS, MAX_BYTES, MAX_ROWS, TEMPLATE_FILENAME
from core.services.common.excel_templates import get_template_definition

TEMPLATE = get_template_definition(TEMPLATE_FILENAME)
if tuple(TEMPLATE["headers"]) != HEADERS:
    raise RuntimeError("批次文件合同与现有模板表头不一致，请检查当前模板定义。")


def read_batch_file(content):
    if not isinstance(content, bytes) or not content or len(content) > MAX_BYTES:
        raise ValidationError("请选择不超过10MB的有效批次XLSX文件。", field="file")
    workbook = None
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 64 * 1024 * 1024:
                raise ValidationError("Excel展开后超过64MB，请拆分文件。", field="file")
        workbook = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
        rows, reference_status = _read_first_sheet(workbook)
        warnings = ([{"code": "first_sheet_only", "message": "仅导入第一张工作表。"}] if len(workbook.worksheets) > 1 else [])
        if reference_status:
            warnings.append({"code": "reference_column_ignored", "message": "文件中的“状态”仅供参考，不导入；已有批次保留当前状态，新批次从待排开始。"})
        return rows, warnings
    except ValidationError:
        raise
    except Exception as exc:
        raise ValidationError("XLSX文件读取失败，请核对文件格式。", field="file") from exc
    finally:
        if workbook is not None:
            workbook.close()


def _read_first_sheet(workbook):
    if not workbook.worksheets:
        raise ValidationError("Excel没有数据工作表。", field="file")
    sheet = workbook.worksheets[0]
    sheet.reset_dimensions()
    iterator = sheet.iter_rows()
    headers = [cell.value for cell in next(iterator, ())]
    while headers and headers[-1] is None:
        headers.pop()
    reference_status = "状态" in headers
    expected = HEADERS + (("状态",) if reference_status else ())
    if len(headers) != len(expected) or set(headers) != set(expected):
        raise ValidationError("请使用批次信息模板的八列，或系统导出清单的九列（含只读状态）；不接受其他列或缺列。", field="headers")
    rows = []
    for line, cells in enumerate(iterator, 2):
        if not any(cell.value is not None for cell in cells):
            continue
        if len(rows) == MAX_ROWS:
            raise ValidationError("一次最多导入 5000 行；这次一行都没有写入。", field="file")
        rows.append(_read_data_row(line, cells, headers))
    return rows, reference_status


def _read_data_row(line, cells, headers):
    errors = []
    if any(cell.data_type in ("f", "e") for cell in cells):
        errors.append("不能导入公式或错误单元格，请提供实际值。")
    if any(cell.value is not None for cell in cells[len(headers):]):
        errors.append("数据行里有表头之外的多余列。")
    values = {key: cells[index].value if index < len(cells) else None for index, key in enumerate(headers) if key in HEADERS}
    return {"row": line, "values": values, "errors": errors}


def write_batch_file(rows, *, template=False):
    if len(rows) > MAX_ROWS:
        raise ValidationError("单次导出超过5000行，请缩小范围。", field="scope")
    workbook = openpyxl.Workbook()
    sheet = workbook.worksheets[0]
    sheet.title = "批次信息"
    sheet.freeze_panes = "A2"
    headers = HEADERS if template else HEADERS + ("状态",)
    sheet.append(headers)
    for index, _header in enumerate(headers, 1):
        cell = sheet.cell(1, index)
        cell.font = Font(bold=True)
        if template:
            example = TEMPLATE["sample_rows"][0][index - 1]
            cell.comment = Comment("示例：" + str(example if example is not None else "留空") + "。批次号、图号须为文本；日期不要填写时分。现有批次的空单元格不覆盖；新批次不会自动生成工序。", "APS")
        sheet.column_dimensions[get_column_letter(index)].width = (22 if index < 3 else 16) if index != 8 else 36
    for number, row in enumerate(rows, 2):
        for index, value in enumerate(row, 1):
            if isinstance(value, str) and (re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", value) or len(value) > 32767):
                raise ValidationError("这段文字里有 Excel 放不了的字符，或者太长；系统没有删除或截断内容。", field="file")
            cell = sheet.cell(number, index, value)
            if isinstance(value, str):
                cell.data_type = TYPE_STRING
                cell.number_format = "@"
    sheet.auto_filter.ref = "A1:" + get_column_letter(len(headers)) + str(max(1, len(rows) + 1))
    try:
        buffer = BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()
    finally:
        workbook.close()
