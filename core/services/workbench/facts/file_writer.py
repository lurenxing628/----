"""One-pass offline downloads; identifiers and formula-like text stay text."""

import codecs
import csv
from tempfile import SpooledTemporaryFile

from core.models.workbench_action_values import FileDownload
from core.models.workbench_resource_file import (
    ENUM_CODES,
    ENUM_LABELS,
    JSON_FIELDS,
    MULTI_CODES,
    NUMERIC_FIELDS,
    table_descriptor,
)
from core.models.workbench_table_descriptor import cell_notes
from core.services.common.excel_instruction_sheet import (
    add_enum_dropdowns,
    append_instruction_sheet,
    close_write_only_sheets,
)
from core.shared.strict_json import canonical_json
from core.shared.strict_parse import (
    ILLEGAL_XML_CHARS_RE as ILLEGAL,
)
from core.shared.strict_parse import (
    is_exact_finite_number,
)

from .file_codec import file_error
from .file_lists import encode_code_list
from .process_file_xml import preserve_carriage_returns
from .table_cells import number_text

XLSX_MAX_ROWS = 1048576
XLSX_MAX_CELL_CHARACTERS = 32767
# 下载文件名按界面上的叫法写，不用内部代号；工种还要分自制与外协。
_CATEGORY_NAMES = {"internal": "自制", "external": "外协"}


def check_capacity(count, fmt):
    if type(fmt) is not str or fmt not in ("csv", "xlsx"):
        raise file_error("只能下载 CSV 或 XLSX，没有开始下载。请重新选择格式。", field="format")
    if fmt == "xlsx" and count + 1 > XLSX_MAX_ROWS:
        raise file_error("行数超过 XLSX 能放的 1048576 行（含表头），请改用 CSV。", count + 1)


def _value(value, field, row, fmt, kind=None):
    if value is None:
        return r"\N"
    if field in MULTI_CODES:
        try:
            value = encode_code_list(value)
        except ValueError as exc:
            raise file_error("工种编号列表里有无效内容，没有导出。请到资料总览核对。", row, field) from exc
    elif field in JSON_FIELDS:
        value = canonical_json(value)
    elif field in NUMERIC_FIELDS:
        if not is_exact_finite_number(value):
            raise file_error("这个格子里存的数字不合法，没有导出，系统也没有替你改成 0。请到资料总览修正后重试。", row, field)
        return value
    elif type(value) is not str:
        raise file_error("这个格子里存的内容不是文字，没有导出。请到资料总览修正后重试。", row, field)
    if kind is not None and field in ENUM_CODES[kind]:
        value = ENUM_LABELS[field].get(value, value)
    return _escaped_text(value, field, row, fmt)


def _cell_text(value):
    """数字写成和页面一样的最短十进制文本：openpyxl 的数字格只留 16 位有效数字，
    3.3000000000000003 回导会变成 3.3，被当成改动；CSV 也不再写成 7.0、1e-07。"""
    return number_text(value) if type(value) in (int, float) else value


def _escaped_text(value, field, row, fmt):
    """Keep literal backslashes reversible and validate the selected file format."""
    if value.startswith("\\"):
        value = "\\" + value
    if fmt == "xlsx" and (ILLEGAL.search(value) or len(value) > XLSX_MAX_CELL_CHARACTERS):
        raise file_error("这个格子里有 XLSX 放不下的字符，或者超过 32767 个字，没有导出，内容也没有被截断。请改用 CSV。", row, field)
    if fmt == "csv" and "\x00" in value:
        raise file_error("这个格子里有 CSV 处理不了的空字符，没有导出，原内容也没有被删。请到资料总览修正后重试。", row, field)
    return value


def write_resource_file(kind, rows, fmt, *, template=False, category=None):
    check_capacity(0, fmt)
    descriptor = table_descriptor(kind)
    filename = ((_CATEGORY_NAMES[category] if category else "") + descriptor["file_stem"]
                + ("导入模板" if template else "清单"))
    if fmt == "csv":
        return _csv(kind, rows, filename, descriptor)
    return write_template_xlsx(
        rows, filename, template, descriptor,
        cell_value=lambda raw, field, row_no: _cell_text(_value(raw, field, row_no, "xlsx", kind)),
        min_width=24,
    )


def _csv(kind, rows, filename, descriptor):
    columns = descriptor["columns"]
    fields = [item["key"] for item in columns]
    count = 0
    with SpooledTemporaryFile(max_size=4 * 1024 * 1024, mode="w+b") as buffer:
        text = codecs.getwriter("utf-8-sig")(buffer)
        writer = csv.writer(text, lineterminator="\r\n")
        writer.writerow([item["label"] for item in columns])
        for count, row in enumerate(rows, 1):
            values = [_value(row[field], field, count + 1, "csv", kind) for field in fields]
            # 文字前加撇号防公式和自动转换；数字不加，回导照样认。
            writer.writerow(["'" + value if type(value) is str else _cell_text(value) for value in values])
        text.flush()
        buffer.seek(0)
        return FileDownload(filename + ".csv", "text/csv; charset=utf-8", buffer.read(), count)


def _headers(ws, template, descriptor):
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.comments import Comment
    from openpyxl.styles import Font

    columns = descriptor["columns"]
    cells, notes = [], cell_notes(descriptor)
    for index, item in enumerate(columns):
        cell = WriteOnlyCell(ws, value=item["label"])
        cell.font = Font(bold=True)
        if template:
            cell.comment = Comment(notes[index], "APS")
        cells.append(cell)
    return cells


def write_template_xlsx(rows, filename, template, descriptor, *, cell_value, min_width=18):
    import openpyxl
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.utils import get_column_letter

    columns = descriptor["columns"]
    fields = [item["key"] for item in columns]
    wb = openpyxl.Workbook(write_only=True)
    ws = wb.create_sheet(descriptor["sheet_name"])
    sheets, count = [ws], 0
    try:
        ws.freeze_panes = "A2"
        for index, column in enumerate(columns, 1):
            ws.column_dimensions[get_column_letter(index)].width = max(min_width, len(column["label"]) * 2 + 4)
        ws.append(_headers(ws, template, descriptor))
        for count, row in enumerate(rows, 1):
            check_capacity(count, "xlsx")
            cells = []
            for field in fields:
                value = cell_value(row[field], field, count + 1)
                cell = WriteOnlyCell(ws, value=value)
                if type(value) is str:
                    cell.data_type, cell.number_format = "s", "@"
                cells.append(cell)
            ws.append(cells)
        ws.auto_filter.ref = "A1:" + get_column_letter(len(fields)) + str(count + 1)
        add_enum_dropdowns(ws, descriptor, count, write_only=True)
        preserve_carriage_returns(ws)
        sheets.append(append_instruction_sheet(wb, descriptor))
        with SpooledTemporaryFile(max_size=4 * 1024 * 1024, mode="w+b") as buffer:
            wb.save(buffer)
            buffer.seek(0)
            return FileDownload(filename + ".xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                buffer.read(), count)
    finally:
        close_write_only_sheets(sheets)
        wb.close()
