"""One-pass offline downloads; identifiers and formula-like text stay text."""

import codecs
import csv
import math
import os
import re
from tempfile import SpooledTemporaryFile

import openpyxl
from openpyxl.cell import WriteOnlyCell
from openpyxl.comments import Comment
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from core.models.workbench_command import canonical_json
from core.models.workbench_resource_file import (
    INSTRUCTIONS,
    JSON_FIELDS,
    NUMERIC_FIELDS,
    READONLY,
    ResourceFileDownload,
    file_columns,
    public_columns,
)
from core.services.workbench.resource_file_codec import file_error

XLSX_MAX_ROWS = 1048576
XLSX_MAX_CELL_CHARACTERS = 32767
ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
# 下载文件名按界面上的叫法写，不用内部代号。
_KIND_NAMES = {"op_type": "工种", "machine": "设备", "operator": "人员", "supplier": "供应商"}
_CATEGORY_NAMES = {"internal": "自制", "external": "外协"}


def check_capacity(count, fmt):
    if type(fmt) is not str or fmt not in ("csv", "xlsx"):
        raise file_error("只能下载 CSV 或 XLSX，没有开始下载。请重新选择格式。", field="format")
    if fmt == "xlsx" and count + 1 > XLSX_MAX_ROWS:
        raise file_error("行数超过 XLSX 能放的 1048576 行（含表头），没有开始下载。请改用 CSV，内容不会被截断。", count + 1)


def _value(value, field, row, fmt):
    if value is None:
        return r"\N"
    if field in JSON_FIELDS:
        value = canonical_json(value)
    elif field in NUMERIC_FIELDS:
        if type(value) not in (int, float) or not math.isfinite(value):
            raise file_error("这个格子里存的数字不合法，没有导出，系统也没有替你改成 0。请到资料总览修正后重试。", row, field)
        return value
    elif type(value) is not str:
        raise file_error("这个格子里存的内容不是文字，没有导出，系统不猜着转换。请到资料总览修正后重试。", row, field)
    if value.startswith("\\"):
        value = "\\" + value
    if fmt == "xlsx" and (ILLEGAL.search(value) or len(value) > XLSX_MAX_CELL_CHARACTERS):
        raise file_error("这个格子里有 XLSX 放不下的字符，或者超过 32767 个字，没有导出，内容也没有被截断。请改用 CSV。", row, field)
    if fmt == "csv" and "\x00" in value:
        raise file_error("这个格子里有 CSV 处理不了的空字符，没有导出，原内容也没有被删。请到资料总览修正后重试。", row, field)
    return value


def write_resource_file(kind, rows, fmt, *, template=False, category=None):
    check_capacity(0, fmt)
    filename = (_CATEGORY_NAMES[category] if category else "") + _KIND_NAMES[kind] + ("导入模板" if template else "清单")
    if fmt == "csv":
        return _csv(kind, rows, filename)
    return _xlsx(kind, rows, filename, template, category)


def _csv(kind, rows, filename):
    count = 0
    with SpooledTemporaryFile(max_size=4 * 1024 * 1024, mode="w+b") as buffer:
        text = codecs.getwriter("utf-8-sig")(buffer)
        writer = csv.writer(text, lineterminator="\r\n")
        writer.writerow([item["label"] for item in public_columns(kind)])
        for count, row in enumerate(rows, 1):
            values = [_value(row[field], field, count + 1, "csv") for field in file_columns(kind)]
            writer.writerow(["'" + value if type(value) is str else value for value in values])
        text.flush()
        buffer.seek(0)
        return ResourceFileDownload(filename + ".csv", "text/csv; charset=utf-8", buffer.read(), count)


def _headers(kind, ws, template, category):
    cells = []
    for item in public_columns(kind):
        cell = WriteOnlyCell(ws, value=item["label"])
        cell.font = Font(bold=True)
        if template:
            key = item["key"]
            example = {"business_code": "000123", "label": "名称", "status": "active", "category": category or "原类别",
                       "default_days": "2.5", "default_merge_mode": "separate / merged", "op_type_code": "OT1",
                       "group_code": "G1", "shift_profile_code": "SHIFT1"}.get(key, "原值")
            if key in ("skill_codes", "op_type_codes", "explicit_op_type_codes"):
                example = '["OT1","OT2"]'
            prefix = "这是只读列，只能照原样填，不要改。" if key in READONLY[kind] else ""
            cell.comment = Comment(prefix + "示例：" + example + "。" + INSTRUCTIONS, "APS")
        cells.append(cell)
    return cells


def _xlsx(kind, rows, filename, template, category):
    wb = openpyxl.Workbook(write_only=True)
    ws = wb.create_sheet("资源")
    count = 0
    try:
        ws.freeze_panes = "A2"
        for index, column in enumerate(public_columns(kind), 1):
            ws.column_dimensions[get_column_letter(index)].width = max(24, len(column["label"]) * 2 + 4)
        ws.append(_headers(kind, ws, template, category))
        for count, row in enumerate(rows, 1):
            check_capacity(count, "xlsx")
            cells = []
            for field in file_columns(kind):
                value = _value(row[field], field, count + 1, "xlsx")
                cell = WriteOnlyCell(ws, value=value)
                if type(value) is str:
                    cell.data_type, cell.number_format = "s", "@"
                cells.append(cell)
            ws.append(cells)
        ws.auto_filter.ref = "A1:" + get_column_letter(len(file_columns(kind))) + str(count + 1)
        with SpooledTemporaryFile(max_size=4 * 1024 * 1024, mode="w+b") as buffer:
            wb.save(buffer)
            buffer.seek(0)
            return ResourceFileDownload(filename + ".xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                        buffer.read(), count)
    finally:
        if not ws.closed:
            ws.close()
        if ws._writer is not None and os.path.exists(ws._writer.out):
            ws._writer.cleanup()
        wb.close()
