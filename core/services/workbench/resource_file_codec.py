"""Strict CSV/XLSX decoding; blank cells omit, escaped nulls remain explicit."""

import csv
import json
import re
from io import BytesIO, StringIO

import openpyxl

from core.errors import ValidationError
from core.models.workbench_resource_file import (
    IMPORT_ROW_LIMIT,
    JSON_FIELDS,
    MULTI_CODES,
    NUMERIC_FIELDS,
    file_columns,
    public_columns,
)

NUMBER = re.compile(r"[+-]?(?:[0-9]+(?:\.[0-9]*)?|\.[0-9]+)(?:[eE][+-]?[0-9]+)?\Z")


def file_error(message, row=1, field="file"):
    return ValidationError(message, field=field, details={"row": row})


def _source_rows(content, fmt):
    if fmt == "csv":
        reader = None
        try:
            reader = csv.reader(StringIO(content.decode("utf-8-sig"), newline=""), strict=True)
            while True:
                number = reader.line_num + 1
                row = next(reader, None)
                if row is None:
                    return
                yield number, row, {}
        except (UnicodeDecodeError, csv.Error) as exc:
            raise file_error("这个 CSV 不是 UTF-8 编码，或者引号不成对，一行都没有导入。请另存为 UTF-8 编码后重新上传。", reader.line_num if reader else 1) from exc
    else:
        wb = None
        try:
            wb = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
            if len(wb.sheetnames) != 1 or len(wb.worksheets) != 1:
                raise file_error("文件里只能有一张数据表，没有导入。请删掉多余的工作表后重新上传。")
            ws = wb.worksheets[0]
            ws.reset_dimensions()
            for number, cells in enumerate(ws.iter_rows(), 1):
                errors = {i: "这个格子是公式或者显示为错误值，没有导入。请改成纯文本后重新上传。" for i, cell in enumerate(cells) if cell.data_type in ("f", "e")}
                yield number, [cell.value for cell in cells], errors
        except ValidationError:
            raise
        except Exception as exc:
            raise file_error("这个 XLSX 打不开，没有导入。请确认文件完整后重新上传。") from exc
        finally:
            if wb is not None:
                wb.close()


def _headers(kind, values):
    names = {column["label"]: column["key"] for column in public_columns(kind)}
    names.update({key: key for key in file_columns(kind)})
    fields = []
    for value in values:
        if type(value) is not str or value not in names:
            raise file_error("表头里有认不出的列或者空列，没有导入。请照模板里的列名填写。", field="headers")
        fields.append(names[value])
    if "business_code" not in fields or len(fields) != len(set(fields)):
        raise file_error("表头必须有编号这一列，而且不能有重复的列（中文名和英文名指同一列也算重复），没有导入。请照模板改好后重新上传。", field="headers")
    return fields


def _decode(value, field, fmt):
    if type(value) is str:
        if fmt == "csv" and value.startswith("'"):
            value = value[1:]
        if value == r"\N":
            return None
        if value.startswith("\\\\"):
            value = value[1:]
    if field in JSON_FIELDS:
        value = _decode_json(value, field)
    if field in NUMERIC_FIELDS and fmt == "csv":
        if type(value) is not str or NUMBER.fullmatch(value) is None:
            raise ValidationError("这个格子只能填数字，不能带单位、是或否、千分位逗号，没有导入。请改成纯数字。", field=field)
        value = float(value)
    return value


def _unique_pairs(pairs):
    value = {}
    for key, item in pairs:
        if key in value:
            raise ValueError("duplicate JSON key")
        value[key] = item
    return value


def _decode_json(value, field):
    if type(value) is not str:
        raise ValidationError("这个格子要填 JSON 文本，没有导入。多个编号请按 [\"A\",\"B\"] 这样填。", field=field)
    try:
        value = json.loads(value, object_pairs_hook=_unique_pairs,
                           parse_constant=lambda _: (_ for _ in ()).throw(ValueError("nonfinite")))
    except (ValueError, TypeError) as exc:
        raise ValidationError("这个格子里的 JSON 写得不对或者有重复的键，没有导入。多个编号请按 [\"A\",\"B\"] 这样填，不要只用逗号隔开。", field=field) from exc
    if field in MULTI_CODES:
        if type(value) is not list or any(type(code) is not str for code in value):
            raise ValidationError("这个格子要填一组用引号括起来的编号，没有导入。请按 [\"A\",\"B\"] 这样填。", field=field)
        if len(value) != len(set(value)):
            raise ValidationError("这个格子里的编号有重复，没有导入。请去掉重复的编号。", field=field)
    return value


def _parse(number, values, errors, fields, fmt):
    parsed, issues = {}, []
    if len(values) > len(fields) and any(v is not None and v != "" for v in values[len(fields):]):
        issues.append({"row": number, "field": "columns", "code": "invalid_input", "message": "这一行的列数比表头多，多出来的内容没有导入。请删掉多余的列。"})
    for index, field in enumerate(fields):
        value = values[index] if index < len(values) else None
        try:
            if index in errors:
                raise ValidationError(errors[index], field=field)
            if value is not None and value != "":
                parsed[field] = _decode(value, field, fmt)
        except ValidationError as exc:
            issues.append({"row": number, "field": field, "code": "invalid_input", "message": exc.message})
    return {"row": number, "values": parsed, "errors": issues}


def read_resource_file(kind, content, fmt):
    if type(content) is not bytes or fmt not in ("csv", "xlsx"):
        raise file_error("只能导入 CSV 或 XLSX 文件，没有导入。请重新选择文件。")
    source = _source_rows(content, fmt)
    try:
        header = next(source, None)
        if header is None or header[2]:
            raise file_error("文件第一行不是表头，没有导入。请照模板补上表头行。", field="headers")
        fields = _headers(kind, header[1])
        rows = []
        for number, values, errors in source:
            if not errors and all(value is None or value == "" for value in values):
                continue
            if len(rows) == IMPORT_ROW_LIMIT:
                raise file_error("一次最多导入 2000 行，这个文件超了，一行都没有导入。请拆成几个小文件分次上传。", number)
            rows.append(_parse(number, values, errors, fields, fmt))
        return rows
    finally:
        source.close()
