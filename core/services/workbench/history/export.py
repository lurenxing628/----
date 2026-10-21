"""Native files from the same archived history and condition projections."""

from tempfile import SpooledTemporaryFile

from core.models.workbench_action_values import FileDownload
from core.models.workbench_command import canonical_json
from core.models.workbench_run_candidate import reject
from core.services.common.excel_instruction_sheet import close_write_only_sheets
from core.services.workbench.facts.file_writer import ILLEGAL, XLSX_MAX_CELL_CHARACTERS, check_capacity
from core.services.workbench.facts.process_file_xml import preserve_carriage_returns

from ..run.decision_export import _append as append_excel_row
from ..run.decision_export import _csv as csv_content
from ..run.decision_export import _csv_cell as csv_cell

FILENAME = "排产历史趋势"
HISTORY_EXPORT_VERSION = 2
CSV_CONTENT_TYPE = "text/csv; charset=utf-8"
XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
SECTION_TITLES = {"entries": "历史条目", "conditions": "生产条件", "configurations": "算法配置"}
ENTRY_FIELDS = (
    ("历史条目编号", "entry_ref"), ("时间线", "timeline"), ("事件来源类型", "source_kind"),
    ("条目名称", "label"), ("条目状态", "state"), ("发生时间", "occurred_at"), ("时间依据", "time_basis"),
    ("运行编号", "identity.run_ref"), ("候选编号", "identity.candidate_ref"), ("试调编号", "identity.scenario_ref"),
    ("调整草案编号", "identity.draft_ref"), ("正式计划编号", "identity.plan_ref"), ("正式版本", "identity.plan_version"),
    ("旧正式记录编号", "identity.history_id"), ("采用回执编号", "identity.receipt_ref"), ("原始请求编号", "identity.request_key"),
    ("运行受理时间", "times.accepted_at"), ("计算开始时间", "times.started_at"), ("计算结束时间", "times.finished_at"),
    ("正式采用时间", "times.adopted_at"), ("提交时间（UTC）", "times.committed_at_utc"), ("旧正式记录时间", "times.schedule_time"),
    ("原正式基准编号", "baseline.plan_ref"), ("原正式基准版本", "baseline.version"),
    ("结果来源类型", "source.kind"), ("结果来源编号", "source.ref"), ("结果来源状态", "source.state"),
    ("结果来源缺项编码", "source.reason.code"), ("结果来源缺项说明", "source.reason.message"),
    ("生产范围状态", "scope.state"), ("批次数", "scope.batch_count"), ("工序数", "scope.operation_count"),
    ("批次总量", "scope.batch_quantity.value"), ("批次总量已知小计", "scope.batch_quantity.known_subtotal"),
    ("批次总量单位", "scope.batch_quantity.unit"), ("批次总量覆盖状态", "scope.batch_quantity.state"),
    ("批次总量指标范围", "scope.batch_quantity.scope"), ("批次总量指标版本", "scope.batch_quantity.metric_version"),
    ("批次总量已知批次数", "scope.batch_quantity.known_count"), ("批次总量未知批次数", "scope.batch_quantity.unknown_count"),
    ("批次总量总批次数", "scope.batch_quantity.total_count"),
    ("批次总量来源类型", "scope.batch_quantity.source.kind"), ("批次总量来源编号", "scope.batch_quantity.source.ref"),
    ("批次总量来源字段", "scope.batch_quantity.source.path"), ("批次总量缺项编码", "scope.batch_quantity.reason.code"),
    ("批次总量缺项说明", "scope.batch_quantity.reason.message"),
    ("需求工时（小时）", "scope.demand_hours"),
    ("范围缺项编码", "scope.reason.code"), ("范围缺项说明", "scope.reason.message"),
    ("结果读取方式", "reading_context.metric_basis"), ("资源窗口方式", "reading_context.resource_window.mode"),
    ("资源窗口起点（含）", "reading_context.resource_window.start"), ("资源窗口终点（不含）", "reading_context.resource_window.end"),
    ("资源窗口天数", "reading_context.resource_window.days"), ("资源窗口起点依据", "reading_context.resource_window.anchor"),
    ("资源窗口时间依据", "reading_context.resource_window.time_basis"), ("资源窗口边界", "reading_context.resource_window.boundary"),
    ("资源窗口状态", "reading_context.resource_window.state"),
    ("资源窗口缺项编码", "reading_context.resource_window.reason.code"),
    ("资源窗口缺项说明", "reading_context.resource_window.reason.message"),
)
METRIC_FIELDS = (
    ("覆盖状态", "state"), ("指标范围", "scope"), ("指标版本", "metric_version"),
    ("指标来源类型", "source.kind"), ("指标来源编号", "source.ref"), ("指标来源字段", "source.path"),
    ("已知数量", "known_count"), ("未知数量", "unknown_count"), ("总数量", "total_count"),
    ("分子工时（小时）", "numerator_hours"), ("分母工时（小时）", "denominator_hours"),
    ("班内未知任务数", "unknown_workload_count"),
    ("指标缺项编码", "reason.code"), ("指标缺项说明", "reason.message"),
)
ENTRY_REFERENCE_FIELDS = (
    ("批次总量覆盖批次永久编号", "scope.batch_quantity.covered_batch_refs", "批次"),
    ("范围批次永久编号", "scope.batch_refs", "批次"),
    ("范围工序永久编号", "scope.operation_refs", "工序"),
    ("窗口设备永久编号", "reading_context.resource_window.cohort.machine_refs", "设备"),
    ("窗口人员永久编号", "reading_context.resource_window.cohort.operator_refs", "人员"),
)
METRIC_REFERENCE_FIELDS = (
    ("指标覆盖批次永久编号", "covered_batch_refs", "批次"),
    ("指标覆盖工序永久编号", "covered_operation_refs", "工序"),
    ("指标覆盖资源永久编号", "covered_resource_refs", "资源"),
)
REFERENCE_HEADERS = ("引用字段", "引用名称", "引用对象类型", "引用原值类型", "引用数量", "引用顺序", "引用永久编号")
COMMON_HEADERS = ("历史格式版本", "导出格式版本", "读取编号", "读取时间")
ENTRY_HEADERS = COMMON_HEADERS + ("记录类型",) + tuple(label for label, _path in ENTRY_FIELDS) \
    + ("指标键", "指标", "单位", "数值", "已知小计") + tuple(label for label, _path in METRIC_FIELDS) + REFERENCE_HEADERS
PAIR_HEADERS = COMMON_HEADERS + ("时间线", "所选条目编号", "基准条目编号", "明确基准运行编号")
SIDE_HEADERS = ("值已保存", "值状态", "原值类型", "原值", "来源类型", "来源编号", "来源字段", "缺项编码", "缺项说明")
DETAIL_HEADERS = PAIR_HEADERS + ("记录类型", "条件类别", "类别名称", "实体编号", "实体名称", "字段", "对照状态") \
    + tuple(prefix + label for prefix in ("基准", "所选") for label in SIDE_HEADERS) + ("说明",)
DEFINITION_FIELDS = (("指标键", "key"), ("指标", "label"), ("分组", "group"), ("目录单位", "unit"),
                     ("显示精度", "precision"), ("改善方向", "improvement"), ("目录范围", "scope"),
                     ("目录来源", "source"), ("常用指标", "common"), ("目录指标版本", "metric_version"))
DEFINITION_HEADERS = tuple(label for label, _field in DEFINITION_FIELDS)
SCOPE_HEADERS = ("项目", "原值类型", "数值或说明")
ISSUE_HEADERS = COMMON_HEADERS + ("条目编号", "缺项类别", "指标或字段", "状态", "原因编码", "说明", "来源类型", "来源编号", "来源字段")


def read_history_export(conn, query=None):
    from .queries import read_history_all

    return read_history_all(conn, query)


def _at(document, path):
    for field in path.split("."):
        if document is None:
            return None
        document = document[field]
    return document


def _common(data, snapshot):
    return [data["history_version"], HISTORY_EXPORT_VERSION, snapshot["snapshot_ref"], snapshot["as_of"]]


def _entry_ref(entry):
    return entry["entry_ref"] if entry is not None else None


def _pair_common(data, snapshot):
    selected, reference = data["selection"]["entry"], data["selection"]["reference"]
    return _common(data, snapshot) + [data["filters"]["timeline"], _entry_ref(selected), _entry_ref(reference),
                                     data["filters"]["reference_run_ref"]]


def _display_number(value, unit):
    return value * 100 if value is not None and unit == "%" else value


def entry_rows(data, snapshot):
    definitions = {row["key"]: row for row in data["metric_definitions"]}
    common = _common(data, snapshot)
    for entry in data["entries"]:
        identity = [_native_value(_at(entry, path)) for _label, path in ENTRY_FIELDS]
        for key in data["filters"]["metric_keys"]:
            metric, definition = entry["metrics"][key], definitions[key]
            yield common + ["指标值"] + identity + [key, definition["label"], metric["unit"], _display_number(metric["value"], metric["unit"]),
                _display_number(metric["known_subtotal"], metric["unit"])] \
                + [_native_value(_at(metric, path)) for _label, path in METRIC_FIELDS] + [None] * len(REFERENCE_HEADERS)
        yield from entry_reference_rows(entry, common, data["filters"]["metric_keys"])


def entry_reference_rows(entry, common, metric_keys):
    """Preserve each collection and member without packing arbitrary scope into one cell."""
    # Reference rows link to their scalar history/metric row; they do not repeat
    # all history values for every member. Null and an explicit empty array each
    # keep one row, so an unused coverage axis is not changed into a known zero.
    identity = [_at(entry, path) if path in ("entry_ref", "timeline") else None for _label, path in ENTRY_FIELDS]
    metric_blanks = [None] * len(METRIC_FIELDS)
    scopes = [("范围引用", None, entry, ENTRY_REFERENCE_FIELDS),
              *(("指标覆盖引用", key, entry["metrics"][key], METRIC_REFERENCE_FIELDS) for key in metric_keys)]
    for record_type, metric_key, source, fields in scopes:
        prefix = common + [record_type] + identity + [metric_key, None, None, None, None] + metric_blanks
        for label, path, kind in fields:
            values = _at(source, path)
            value_type, count = ("null", None) if values is None else ("array", len(values))
            members = enumerate(values, 1) if values else ((None, None),)
            for index, reference in members:
                yield prefix + [path, label, kind, value_type, count, index, reference]


def _side_values(side):
    return [side["present"], side["state"], _value_type(side["value"]), _native_value(side["value"]),
            side["source"]["kind"], side["source"]["ref"], side["source"]["path"],
            _at(side, "reason.code"), _at(side, "reason.message")]


def detail_rows(data, snapshot, section):
    common, comparison = _pair_common(data, snapshot), data["comparison"]
    group = comparison[section]
    categories = {category["key"]: category["label"] for category in group["categories"]}
    pair = data["selection"]
    if pair["entry"] is None or pair["reference"] is None:
        conclusion = comparison["conclusion"]
        yield common + ["explanation", None, None, None, None, None, conclusion["state"]] \
            + [None] * (len(SIDE_HEADERS) * 2) + [_at(conclusion, "reason.message") or conclusion["message"]]
        return
    for row in group["rows"]:
        yield common + ["field", row["category"], categories[row["category"]], row["entity"]["key"], row["entity"]["label"],
                        row["field"], row["state"]] + _side_values(row["before"]) + _side_values(row["after"]) + [None]


def definition_rows(data):
    for definition in data["metric_definitions"]:
        yield [definition[field] for _label, field in DEFINITION_FIELDS]


def scope_rows(data, snapshot):
    selection, comparison = data["selection"], data["comparison"]
    fields = [("历史格式版本", data["history_version"]), ("导出格式版本", HISTORY_EXPORT_VERSION),
              ("读取编号", snapshot["snapshot_ref"]), ("读取时间", snapshot["as_of"]),
              ("明细范围", "筛选结果全部页"), ("所选条目编号", _entry_ref(selection["entry"])),
              ("基准条目编号", _entry_ref(selection["reference"])), ("运行时间线条目数", data["totals"]["run"]),
              ("正式采用时间线条目数", data["totals"]["official"]), ("当前筛选条目数", data["totals"]["filtered"]),
              ("比较结论状态", comparison["conclusion"]["state"]), ("比较结论", comparison["conclusion"]["message"]),
              ("比较原因编码", _at(comparison, "conclusion.reason.code")), ("比较原因说明", _at(comparison, "conclusion.reason.message")),
              ("生产条件状态", comparison["conditions"]["state"]), ("算法配置状态", comparison["configurations"]["state"]),
              ("生产条件和配置方向", "基准值 → 所选值"), ("百分比输出方式", "保存比例乘100；分子、分母工时仍按小时"),
              ("计算耗时单位", "秒，保留保存数值"), ("目录版本说明", "指标口径表列出当前目录；每条历史指标的实际版本和来源保留在历史条目表"),
              ("条目记录类型", "指标值、范围引用、指标覆盖引用；按历史条目编号和指标键关联"),
              ("引用集合说明", "引用原值类型为null时数量与成员留空；array数量0为空集合；成员按引用顺序逐行保留"),
              ("导出记录数说明", "响应记录数包含主分区的指标值、引用明细和空/null占位行；不含指标口径、比较范围、数据缺项三张附表")]
    for key, value in data["filters"].items():
        if key not in ("page", "size", "condition_page", "condition_size", "snapshot_ref"):
            fields.append(("筛选." + key, value))
    fields += [("资源窗口." + key, value) for key, value in data["resource_window"].items()]
    fields += [("比较." + key, comparison[key]) for key in ("objective", "metric_deltas")]
    for label, value in fields:
        yield [label, _value_type(value), _native_value(value)]


def issue_rows(data, snapshot):
    common = _common(data, snapshot)
    for entry in data["entries"]:
        ref = entry["entry_ref"]
        source = entry["source"]
        for reason in entry["issues"]:
            yield common + [ref, "entry", None, source["state"], reason["code"], reason["message"], source["kind"], source["ref"], None]
        quantity = entry["scope"]["batch_quantity"]
        for field, value, origin in (("scope", entry["scope"], source), ("batch_quantity", quantity, quantity["source"]),
                                     ("resource_window", entry["reading_context"]["resource_window"], source)):
            reason = value["reason"]
            if reason is not None:
                yield common + [ref, field, None, value["state"], reason["code"], reason["message"], origin["kind"], origin["ref"], origin.get("path")]
        for key in data["filters"]["metric_keys"]:
            metric = entry["metrics"][key]
            if metric["reason"] is not None or metric["state"] in ("partial", "unavailable"):
                reason, origin = metric["reason"], metric["source"]
                yield common + [ref, "metric", key, metric["state"], _at(reason, "code"), _at(reason, "message"),
                                origin["kind"], origin["ref"], origin["path"]]
    for section in ("conditions", "configurations"):
        for row in data["comparison"][section]["rows"]:
            for name, entry in (("before", data["selection"]["reference"]), ("after", data["selection"]["entry"])):
                side = row[name]
                if side["reason"] is not None:
                    reason, origin = side["reason"], side["source"]
                    yield common + [_entry_ref(entry), section + "." + name, row["field"], side["state"], reason["code"],
                                    reason["message"], origin["kind"], origin["ref"], origin["path"]]
    conclusion = data["comparison"]["conclusion"]
    if conclusion["reason"] is not None:
        yield common + [_entry_ref(data["selection"]["entry"]), "comparison", None, conclusion["state"],
                        conclusion["reason"]["code"], conclusion["reason"]["message"], None, None, None]


def write_history_export(data, snapshot, fmt, section="entries"):
    sections = tuple(SECTION_TITLES) if section == "all" else (section,)
    tables, count = {}, 0

    def counted(rows):
        nonlocal count
        for row in rows:
            count += 1
            yield row

    for name in sections:
        if name == "entries":
            headers, rows = ENTRY_HEADERS, entry_rows(data, snapshot)
        else:
            headers, rows = DETAIL_HEADERS, detail_rows(data, snapshot, name)
        tables[SECTION_TITLES[name]] = headers, counted(rows)
    if fmt == "csv":
        headers, rows = next(iter(tables.values()))
        content = csv_content(headers, _csv_rows(rows))
    else:
        tables["指标口径"] = DEFINITION_HEADERS, definition_rows(data)
        tables["比较范围"] = SCOPE_HEADERS, scope_rows(data, snapshot)
        tables["数据缺项"] = ISSUE_HEADERS, issue_rows(data, snapshot)
        content = _xlsx(tables)
    return FileDownload(FILENAME + "." + fmt, CSV_CONTENT_TYPE if fmt == "csv" else XLSX_CONTENT_TYPE, content, count)


def _native_value(value):
    return canonical_json(value) if isinstance(value, (dict, list, tuple)) else value


def _csv_rows(rows):
    # 原值列还会包含编号形状的文字；保存为文字，数字和布尔值继续保留原类型。
    for row in rows:
        yield [csv_cell(value, identity=True) if isinstance(value, str) else value for value in row]


def _value_type(value):
    if value is None:
        return "null"
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, int):
        return "integer"
    if isinstance(value, float):
        return "number"
    if isinstance(value, (list, tuple)):
        return "array"
    return "object" if isinstance(value, dict) else "string"


def _xlsx(tables):
    import openpyxl
    from openpyxl.utils import get_column_letter

    workbook = openpyxl.Workbook(write_only=True)
    try:
        for title, (headers, rows) in tables.items():
            sheet = workbook.create_sheet(title)
            sheet.freeze_panes = "A2"
            for index in range(1, len(headers) + 1):
                sheet.column_dimensions[get_column_letter(index)].width = 24
            append_excel_row(sheet, headers)
            for number, row in enumerate(rows, 1):
                check_capacity(number, "xlsx")
                for column, value in enumerate(row):
                    if isinstance(value, str) and (len(value) > XLSX_MAX_CELL_CHARACTERS or ILLEGAL.search(value)):
                        reject("export_format_limit", "XLSX 第" + str(number + 1) + "行的“" + headers[column]
                               + "”含超出单元格格式限制的文字，请选择 CSV 导出。", 422)
                append_excel_row(sheet, row)
            preserve_carriage_returns(sheet)
        with SpooledTemporaryFile(max_size=4 * 1024 * 1024, mode="w+b") as buffer:
            workbook.save(buffer)
            buffer.seek(0)
            return buffer.read()
    finally:
        close_write_only_sheets(workbook.worksheets)
        workbook.close()
