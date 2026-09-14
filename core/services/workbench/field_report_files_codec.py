"""Strict legacy ten-column and explicit-task thirteen-column XLSX contracts."""

import math
import re
from datetime import date, datetime
from io import BytesIO
from typing import NoReturn
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import TYPE_ERROR, TYPE_FORMULA, TYPE_STRING
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils.datetime import from_excel

from core.models.workbench_command import WorkbenchCommandRejected
from core.services.workbench.field_report_files_xml import check_package

HEADERS = ('报工编号', '批次号', '工序', '本次完成数量', '实际开工', '本次实际完工',
           '有效加工工时(h)', '实际设备', '实际人员', '备注')
FIELDS = ('report_no', 'batch_id', 'operation_label', 'completed_quantity', 'actual_start', 'actual_end',
          'effective_processing_hours', 'machine_label', 'operator_label', 'remark')
TASK_HEADERS = HEADERS + ('任务编号', '工序范围', '单件编号')
TASK_FIELDS = FIELDS + ('task_ref', 'operation_scope', 'piece_id')
ROW_LIMIT = 5000
BYTE_LIMIT = 8 * 1024 * 1024
MIME = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


def reject(message: str) -> NoReturn:
    raise WorkbenchCommandRejected('invalid_input', message, 422)


def _timestamp(value, epoch):
    if type(value) in (int, float):
        if not math.isfinite(value) or value < 0:
            reject('Excel 里的日期时间填得不对。')
        value = from_excel(value, epoch)
    if isinstance(value, datetime):
        if value.tzinfo is not None or value.second or value.microsecond:
            reject('实际时间只填到分钟，秒要写 00，例如 2026-09-13 08:30。')
        return value.isoformat(timespec='seconds')
    if isinstance(value, date):
        reject('实际时间不能只有日期，请填写时分。')
    if type(value) is not str or not re.fullmatch(r'\d{4}-\d\d-\d\d[ T]\d\d:\d\d(?::00)?', value.strip()):
        reject('实际时间请按 2026-09-13 08:30 这样填。')
    try:
        parsed = datetime.fromisoformat(value.strip())
    except ValueError as exc:
        raise WorkbenchCommandRejected('invalid_input', '实际时间不是有效的日期时间。', 422) from exc
    return parsed.isoformat(timespec='seconds')


def _value(cell, field, epoch):
    value = cell.value
    if cell.data_type in (TYPE_FORMULA, TYPE_ERROR) or type(value) is bool:
        reject('这一格不能填公式，也不能是 Excel 的错误值或「是/否」，请直接填实际值。')
    if value is None or type(value) is str and not value.strip():
        return None if field in FIELDS[3:7] else ''
    if field in ('actual_start', 'actual_end'):
        return _timestamp(value, epoch)
    if field in ('completed_quantity', 'effective_processing_hours'):
        return _numeric(value, field)
    if type(value) is not str:
        reject('编号、批次、工序、设备人员和备注都要填文字。')
    if len(value) > 32767:
        reject('这一格的文字超过了 Excel 单元格能放的长度。')
    return value.strip()


def _numeric(value, field):
    if type(value) not in (str, int, float) or type(value) is str and not re.fullmatch(r'\d+(?:\.\d+)?', value.strip()):
        reject('数量和工时要填 0 或大于 0 的数字。')
    number = float(value)
    if not math.isfinite(number) or number < 0:
        reject('数量和工时要填 0 或大于 0 的有效数字。')
    if field == 'completed_quantity':
        if not number.is_integer() or number > (1 << 53) - 1:
            reject('本次完成数量要填 0 或正整数。')
        return int(number)
    return number


def _column_order(first):
    headers = [cell.value.strip() if type(cell.value) is str else cell.value for cell in first]
    for version, expected, fields in ((1, HEADERS, FIELDS), (2, TASK_HEADERS, TASK_FIELDS)):
        if len(headers) == len(expected) and set(headers) == set(expected):
            return version, fields, [headers.index(label) for label in expected]
    reject('第一张工作表的表头必须正好是原来的 10 列，或者是带任务编号、工序范围、单件编号的 13 列。')


def _decode_row(cells, number, version, fields, order, epoch):
    row = {'row': number, 'format_version': version, 'values': {}, 'errors': []}
    for field, index in zip(fields, order):
        try:
            row['values'][field] = _value(cells[index], field, epoch) if index < len(cells) else None
        except (WorkbenchCommandRejected, ValueError, OverflowError) as exc:
            row['errors'].append({'row': number, 'field': field, 'message': str(exc)})
    return row


def decode_reports(content):
    if type(content) is not bytes or not content or len(content) > BYTE_LIMIT:
        reject('请选择不超过 8 MB 的 XLSX 文件。')
    columns = check_package(content)
    try:
        book = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
    except (BadZipFile, OSError, ValueError, KeyError) as exc:
        raise WorkbenchCommandRejected('invalid_input', '这个 XLSX 文件读不开，请确认文件没有损坏。', 422) from exc
    try:
        sheet = book.worksheets[0]
        # Do not trust a forged/undersized dimension to truncate physical rows.
        sheet.reset_dimensions()
        source = sheet.iter_rows()
        version, fields, order = _column_order(next(source, ()))
        if columns > len(fields):
            reject('有单元格超出了 10 列或 13 列的表头范围，系统不会覆盖或忽略原来的值。')
        result = []
        for number, cells in enumerate(source, 2):
            if number > ROW_LIMIT + 1:
                reject('最多只能导入 5000 行，这次一行都没有导入。')
            result.append(_decode_row(cells, number, version, fields, order, book.epoch))
        return result
    finally:
        book.close()


def _instructions(book):
    help_sheet = book.create_sheet('填写说明')
    help_sheet.append(['项目', '说明'])
    for row in [('数据范围', '只读第一张报工记录表，最多 5000 行；预检只看不保存。'),
                    ('逐次报工', '数量和有效工时为本次值，不是累计值；未知留空，0 为已知零。'),
                    ('编号', '请保留报工编号。任务编号、工序范围和单件编号由当前范围预填，不要改；不用手抄任务编号。'),
                    ('单件与共同工序', '工序范围为单件时，单件编号必须与预填值一致；共同工序的单件编号必须留空。原来的 10 列格式只支持批次和工序都唯一的情况；重名分不清会被拒绝。'),
                    ('目标数量', '实际甘特 CSV 里的目标数量按报工统计；计划应做数量、计划批次数量和依据另外列出，不知道就留空。'),
                    ('设备人员', '实际设备和人员要按现场核对后填写，不能拿计划里的当实际。'),
                    ('重复与更正', '重复导入不会累加，只补没填的；和已填内容对不上时请走更正并写原因。'),
                    ('时间', '按 2026-09-13 08:30 这样填；系统不会按开工到完工的时长自动算加工工时。')]:
        help_sheet.append(row)


def _style(page, is_report):
    page.freeze_panes = 'A2'
    page.auto_filter.ref = page.dimensions
    for cell in page[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='38625A')
    for row in page:
        for cell in row:
            if type(cell.value) is str:
                cell.data_type = TYPE_STRING
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    for column in page.columns:
        letter = column[0].column_letter
        page.column_dimensions[letter].width = 26 if is_report else 32
    if is_report:
        page.column_dimensions['J'].width = 48
    elif page.title == '填写说明':
        page.column_dimensions['B'].width = 90


def encode_reports(rows, *, template=False, summaries=(), metadata=(), format_version=1):
    if type(format_version) is not int or format_version not in (1, 2):
        reject('这个报工文件的格式版本不支持。')
    headers, fields = (HEADERS, FIELDS) if format_version == 1 else (TASK_HEADERS, TASK_FIELDS)
    book = Workbook()
    sheet = book.active
    if sheet is None:
        raise RuntimeError('New report workbook has no active worksheet.')
    sheet.title = '报工记录'
    sheet.append(headers)
    for number, row in enumerate(rows, 1):
        if number > ROW_LIMIT:
            reject('报工文件超过 5000 行上限，请缩小筛选范围；系统不会只导当前页。')
        sheet.append([row.get(field) for field in fields])
    if template:
        _instructions(book)
    for title, values in [('工序汇总', summaries), ('录入信息', metadata)]:
        values = list(values)
        if values:
            extra = book.create_sheet(title)
            for row in values:
                extra.append(row)
    for page in book:
        _style(page, page is sheet)
    output = BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


def encode_issues(rows):
    book = Workbook()
    sheet = book.active
    if sheet is None:
        raise RuntimeError('New issue workbook has no active worksheet.')
    sheet.title = '导入问题'
    sheet.append(['Excel 行号', '问题'])
    for row in rows:
        for issue in row['errors']:
            sheet.append([row['row'], issue['message']])
            sheet.cell(sheet.max_row, 2).data_type = TYPE_STRING
    sheet.column_dimensions['A'].width = 15
    sheet.column_dimensions['B'].width = 85
    sheet.freeze_panes = 'A2'
    for cells in sheet:
        for cell in cells:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    output = BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()
