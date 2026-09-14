"""Bounded CSV/XLSX byte sources, including all physical records and cells."""

import csv
from io import BytesIO, StringIO
from xml.etree import ElementTree
from zipfile import ZipFile

import openpyxl

from core.errors import ValidationError
from core.models.workbench_process_file import IMPORT_BYTE_LIMIT, XLSX_EXPANDED_BYTE_LIMIT, file_error
from core.services.workbench.process_file_xml import check_sheet_order

_WORKBOOK_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"


def check_bytes(content, file_format):
    if type(content) is not bytes:
        raise file_error("请上传 CSV 或 XLSX 文件。")
    if len(content) > IMPORT_BYTE_LIMIT:
        raise file_error("文件超过 16 MB 上限，没有读取，也不会只读前面一部分。请拆分文件后重新上传。")
    if file_format == "xlsx":
        _check_package(content)


def _check_package(content):
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if len({item.filename for item in entries}) != len(entries):
                raise file_error("这个 XLSX 内部有重复内容，认不出原始数据。请用 Excel 另存一份后重新上传。")
            if sum(item.file_size for item in entries) > XLSX_EXPANDED_BYTE_LIMIT:
                raise file_error("这个 XLSX 展开后超过 64 MB 上限，没有读取，也不会只读一部分。请拆分文件后重新上传。")
            if any(item.flag_bits & 1 for item in entries):
                raise file_error("加了密码的 XLSX 读不了。请去掉密码后重新上传。")
            if archive.testzip() is not None:
                raise file_error("这个 XLSX 的完整性检查没通过，文件可能已损坏。请用 Excel 另存一份后重新上传。")
            _check_workbook_content(archive)
    except ValidationError:
        raise
    except Exception as exc:
        raise file_error("这个 XLSX 打不开。请确认文件完整后重新上传。") from exc


def _check_workbook_content(archive):
    types = ElementTree.fromstring(archive.read("[Content_Types].xml"))
    workbooks = [item.attrib.get("ContentType", "") for item in types
                 if "spreadsheetml" in item.attrib.get("ContentType", "") and "main+xml" in item.attrib.get("ContentType", "")
                 or "macroEnabled.main+xml" in item.attrib.get("ContentType", "")]
    if workbooks != [_WORKBOOK_TYPE]:
        raise file_error("这个文件不是标准 XLSX 工作簿，只改扩展名或带宏的模板都不行。请用 Excel 另存为 xlsx 后重新上传。")
    for item in types:
        if item.attrib.get("ContentType", "").endswith("spreadsheetml.worksheet+xml"):
            with archive.open(item.attrib["PartName"].lstrip("/")) as stream:
                check_sheet_order(stream)


def csv_rows(content):
    reader = None
    try:
        reader = csv.reader(StringIO(content.decode("utf-8-sig"), newline=""), strict=True)
        while True:
            number = reader.line_num + 1
            values = next(reader, None)
            if values is None:
                return
            yield number, values, {}
    except (UnicodeDecodeError, csv.Error) as exc:
        raise file_error("CSV 必须是 UTF-8 编码，系统不猜编码，也不替你补引号。请另存为 UTF-8 后重新上传。", reader.line_num if reader else 1) from exc


def xlsx_rows(content):
    wb = None
    try:
        wb = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
        if len(wb.sheetnames) != 1 or len(wb.worksheets) != 1:
            raise file_error("文件只能有一张数据工作表，系统不会忽略其他表。请删掉多余的工作表后重新上传。")
        ws = wb.worksheets[0]
        ws.reset_dimensions()
        for number, cells in enumerate(ws.iter_rows(), 1):
            errors = {index: "不接受公式或 Excel 错误单元格，请填写实际值。"
                      for index, cell in enumerate(cells) if cell.data_type in ("f", "e")}
            yield number, [cell.value for cell in cells], errors
    except ValidationError:
        raise
    except Exception as exc:
        raise file_error("这个 XLSX 读不出来。请确认文件完整、单元格格式正常后重新上传。") from exc
    finally:
        if wb is not None:
            wb.close()
