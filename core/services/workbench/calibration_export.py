"""Complete filtered suggestions, with shared spreadsheet injection escaping."""

import csv
import io

import openpyxl

from core.models.workbench_calibration import MAX_EXPORT_BYTES
from core.models.workbench_command import WorkbenchCommandRejected, canonical_json
from core.services.common.excel_templates import _sanitize_export_cell
from core.services.report.exporters.xlsx import _append_write_only_row
from core.services.report.report_engine import ReportExport

COLUMNS = (("part_no", "图号"), ("part_name", "零件名称"), ("sequence", "工序号"),
           ("operation_label", "工序名称"), ("template_operation_ref", "模板工序编号"),
           ("template_revision", "模板版本"), ("template_snapshot", "模板数据版本"),
           ("old_unit_hours", "旧单件定额（小时）"), ("suggested_unit_hours", "建议单件定额（小时）"),
           ("deviation_percent", "偏差百分比"), ("deviation_basis", "偏差计算状态"),
           ("sample_count", "可用完工记录数"), ("candidate_count", "待核对完工记录数"),
           ("sample_refs", "完工记录编号"), ("sample_revisions", "完工记录版本"),
           ("exclusion_reasons", "剔除原因"), ("method_version", "计算方法"),
           ("generated_at", "生成时间"), ("as_of", "数据截至"), ("snapshot_ref", "数据版本编号"))


def _cell(value):
    if value is None:
        return "暂无数据"
    if isinstance(value, (dict, list)):
        value = canonical_json(value)
    if isinstance(value, str) and value and (value.lstrip()[:1] in ("=", "+", "-", "@") or value[0] in "\t\r\n"):
        return "'" + value
    return value


def _export_values(rows, scope, format_name):
    values = [[_cell(row[key]) for key, _ in COLUMNS] + [scope] for row in rows]
    if sum(len(str(value).encode("utf-8")) for row in values for value in row) > MAX_EXPORT_BYTES:
        raise WorkbenchCommandRejected("export_too_large", "完整导出超过 16 MB，请缩小范围后重试；系统没有截断内容。", 413)
    if format_name == "xlsx" and any(isinstance(value, str) and len(value) > 32767 for row in values for value in row):
        raise WorkbenchCommandRejected("export_too_large", "有单元格内容超出 XLSX 的上限，请改导 CSV；系统没有截断内容。", 413)
    return values


def _csv(values, headers, output):
    text = io.StringIO(newline="")
    writer = csv.writer(text)
    writer.writerow(headers)
    for row in values:
        writer.writerow([_sanitize_export_cell(value) for value in row])
    output.write(text.getvalue().encode("utf-8-sig"))
    return "text/csv;charset=utf-8"


def _xlsx(values, headers, output, metadata_rows):
    workbook = openpyxl.Workbook(write_only=True)
    try:
        sheet = workbook.create_sheet("校准建议")
        sheet.freeze_panes = "A2"
        _append_write_only_row(sheet, headers, is_header=True)
        for row in values:
            _append_write_only_row(sheet, row)
        metadata = workbook.create_sheet("范围与计算方式")
        for row in metadata_rows:
            _append_write_only_row(metadata, list(row))
        workbook.save(output)
    finally:
        workbook.close()
    return "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def export_calibration(data, rows, snapshot, format_name):
    if format_name not in ("csv", "xlsx"):
        raise WorkbenchCommandRejected("invalid_input", "只支持 CSV 或 XLSX 格式。", 400)
    if not rows:
        raise WorkbenchCommandRejected("empty_export", "当前筛选范围没有可导出的建议行。", 422)
    scope = canonical_json(data["scope"])
    values = _export_values(rows, scope, format_name)
    headers = [label for _, label in COLUMNS] + ["筛选范围"]
    output = io.BytesIO()
    if format_name == "csv":
        mime = _csv(values, headers, output)
    else:
        metadata_rows = (("数据来源", "生产库"), ("数据截至", snapshot["as_of"]),
                         ("数据版本编号", snapshot["snapshot_ref"]), ("筛选范围", scope),
                         ("来源限制", canonical_json(data["source_constraints"])),
                         ("采用与锁定", "请在工时校准页面预检并采用；采用后更新并锁定模板定额，已有批次不随之更改。"))
        mime = _xlsx(values, headers, output, metadata_rows)
    if output.tell() > MAX_EXPORT_BYTES:
        raise WorkbenchCommandRejected("export_too_large", "完整导出超过 16 MB，请缩小范围后重试；系统没有截断内容。", 413)
    output.seek(0)
    filename = "工时校准明细-" + snapshot["as_of"].replace(":", "") + "." + format_name
    return ReportExport(filename, mime, output, estimated_rows=len(rows))
