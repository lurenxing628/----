"""One-pass process file downloads; no import row cap or implicit row selection."""

import codecs
import csv
from collections.abc import Mapping
from tempfile import SpooledTemporaryFile

import openpyxl
from openpyxl.cell import WriteOnlyCell
from openpyxl.comments import Comment
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from core.models.workbench_process_file import (
    LABELS,
    ProcessFileDownload,
    file_columns,
    file_error,
    table_descriptor,
)
from core.models.workbench_table_descriptor import cell_notes
from core.services.common.excel_instruction_sheet import (
    add_enum_dropdowns,
    append_instruction_sheet,
    close_write_only_sheets,
)
from core.services.workbench.facts.file_writer import XLSX_MAX_ROWS
from core.services.workbench.facts.process_file_xml import preserve_carriage_returns

from .file_values import export_value


def _row_values(kind, row, number, file_format):
    fields = file_columns(kind)
    if not isinstance(row, Mapping) or any(type(key) is not str or key not in fields for key in row):
        raise file_error("导出内容里出现认不出的列，导出没有继续。请刷新重试；仍不行请联系维护人员。", number)
    # An empty string is a blank cell, just as it is in an uploaded template.
    return [export_value(row[field], field, number, file_format) if field in row and row[field] != "" else None
            for field in fields]


def write_csv(kind, rows):
    count = 0
    with SpooledTemporaryFile(max_size=4 * 1024 * 1024, mode="w+b") as buffer:
        text = codecs.getwriter("utf-8-sig")(buffer)
        writer = csv.writer(text, lineterminator="\r\n")
        writer.writerow([LABELS[field] for field in file_columns(kind)])
        for count, row in enumerate(rows, 1):
            values = _row_values(kind, row, count + 1, "csv")
            writer.writerow(["'" + value if value is not None else None for value in values])
        text.flush()
        buffer.seek(0)
        return ProcessFileDownload(_filename(kind, "csv"), "text/csv; charset=utf-8", buffer.read(), count)


def _filename(kind, file_format):
    return table_descriptor(kind)["file_stem"] + "." + file_format


def write_xlsx(kind, rows, template=False):
    descriptor = table_descriptor(kind)
    wb = openpyxl.Workbook(write_only=True)
    ws = wb.create_sheet(descriptor["sheet_name"])
    sheets, notes, count = [ws], cell_notes(descriptor), 0
    try:
        ws.freeze_panes = "A2"
        headers = []
        for index, field in enumerate(file_columns(kind), 1):
            cell = WriteOnlyCell(ws, value=LABELS[field])
            cell.font = Font(bold=True)
            if template:
                cell.comment = Comment(notes[index - 1], "APS")
            headers.append(cell)
            ws.column_dimensions[get_column_letter(index)].width = 44 if field in ("route_raw", "remark") else 24
        ws.append(headers)
        for count, row in enumerate(rows, 1):
            if count + 1 > XLSX_MAX_ROWS:
                raise file_error("行数超过 XLSX 单表 1048576 行上限（算上表头）。请改用 CSV 导出。", count + 1)
            cells = []
            for value in _row_values(kind, row, count + 1, "xlsx"):
                cell = WriteOnlyCell(ws, value=value)
                if value is not None:
                    cell.data_type, cell.number_format = "s", "@"
                cells.append(cell)
            ws.append(cells)
        ws.auto_filter.ref = "A1:" + get_column_letter(len(headers)) + str(count + 1)
        add_enum_dropdowns(ws, descriptor, count, write_only=True)
        # 这一步会关掉数据表并改写它的临时 XML，说明表要在它之后再建，各写各的。
        preserve_carriage_returns(ws)
        sheets.append(append_instruction_sheet(wb, descriptor))
        with SpooledTemporaryFile(max_size=4 * 1024 * 1024, mode="w+b") as buffer:
            wb.save(buffer)
            buffer.seek(0)
            return ProcessFileDownload(_filename(kind, "xlsx"),
                                       "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                       buffer.read(), count)
    finally:
        # Pinned openpyxl 3.0.10 needs explicit sheet cleanup when export aborts.
        close_write_only_sheets(sheets)
        wb.close()
