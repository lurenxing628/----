"""一次性离线下载。编号与看起来像公式的文字一律按文本写出。"""

import codecs
import csv
import re
from tempfile import SpooledTemporaryFile

import openpyxl
from openpyxl.cell import WriteOnlyCell
from openpyxl.comments import Comment
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from core.models.workbench_relation_file import (
    RelationFileDownload,
    file_columns,
    public_columns,
    relation_kind,
    table_descriptor,
)
from core.models.workbench_table_descriptor import cell_notes
from core.services.common.excel_instruction_sheet import (
    add_enum_dropdowns,
    append_instruction_sheet,
    close_write_only_sheets,
)

from .file_codec import file_error

XLSX_MAX_ROWS = 1048576
XLSX_MAX_CELL_CHARACTERS = 32767
ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def check_capacity(count, fmt):
    if type(fmt) is not str or fmt not in ("csv", "xlsx"):
        raise file_error("只能下载 CSV 或 XLSX，没有开始下载。请重新选择格式。", field="format")
    if fmt == "xlsx" and count + 1 > XLSX_MAX_ROWS:
        raise file_error("行数超过 XLSX 能放的 1048576 行（含表头），请改用 CSV。", count + 1)


def _value(value, field, row, fmt):
    if value is None:
        return r"\N"
    if type(value) is not str:
        raise file_error("这个格子里存的内容不是文字，没有导出。请到资料总览修正后重试。", row, field)
    if value.startswith("\\"):
        value = "\\" + value
    if fmt == "xlsx" and (ILLEGAL.search(value) or len(value) > XLSX_MAX_CELL_CHARACTERS):
        raise file_error("这个格子里有 XLSX 放不下的字符，或者超过 32767 个字，没有导出，内容也没有被截断。请改用 CSV。", row, field)
    if fmt == "csv" and "\x00" in value:
        raise file_error("这个格子里有 CSV 处理不了的空字符，没有导出，原内容也没有被删。请到资料总览修正后重试。", row, field)
    return value


def write_relation_file(kind, rows, fmt, *, template=False):
    relation_kind(kind)
    check_capacity(0, fmt)
    descriptor = table_descriptor(kind)
    filename = descriptor["file_stem"] + ("导入模板" if template else "清单")
    if fmt == "csv":
        return _csv(kind, rows, filename)
    return _xlsx(kind, rows, filename, template, descriptor)


def _csv(kind, rows, filename):
    count = 0
    with SpooledTemporaryFile(max_size=4 * 1024 * 1024, mode="w+b") as buffer:
        text = codecs.getwriter("utf-8-sig")(buffer)
        writer = csv.writer(text, lineterminator="\r\n")
        writer.writerow([item["label"] for item in public_columns(kind)])
        for count, row in enumerate(rows, 1):
            values = [_value(row[field], field, count + 1, "csv") for field in file_columns(kind)]
            writer.writerow(["'" + value if type(value) is str else value for value in values])
        text.flush()
        buffer.seek(0)
        return RelationFileDownload(filename + ".csv", "text/csv; charset=utf-8", buffer.read(), count)


def _headers(kind, ws, template, descriptor):
    cells, notes = [], cell_notes(descriptor)
    for index, item in enumerate(public_columns(kind)):
        cell = WriteOnlyCell(ws, value=item["label"])
        cell.font = Font(bold=True)
        if template:
            cell.comment = Comment(notes[index], "APS")
        cells.append(cell)
    return cells


def _xlsx(kind, rows, filename, template, descriptor):
    wb = openpyxl.Workbook(write_only=True)
    ws = wb.create_sheet(descriptor["sheet_name"])
    sheets, count = [ws], 0
    try:
        ws.freeze_panes = "A2"
        for index, column in enumerate(public_columns(kind), 1):
            ws.column_dimensions[get_column_letter(index)].width = max(18, len(column["label"]) * 2 + 4)
        ws.append(_headers(kind, ws, template, descriptor))
        for count, row in enumerate(rows, 1):
            check_capacity(count, "xlsx")
            cells = []
            for field in file_columns(kind):
                value = _value(row[field], field, count + 1, "xlsx")
                cell = WriteOnlyCell(ws, value=value)
                if type(value) is str:
                    cell.data_type, cell.number_format = "s", "@"
                cells.append(cell)
            ws.append(cells)
        ws.auto_filter.ref = "A1:" + get_column_letter(len(file_columns(kind))) + str(count + 1)
        add_enum_dropdowns(ws, descriptor, count, write_only=True)
        sheets.append(append_instruction_sheet(wb, descriptor))
        with SpooledTemporaryFile(max_size=4 * 1024 * 1024, mode="w+b") as buffer:
            wb.save(buffer)
            buffer.seek(0)
            return RelationFileDownload(
                filename + ".xlsx",
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                buffer.read(), count)
    finally:
        close_write_only_sheets(sheets)
        wb.close()
