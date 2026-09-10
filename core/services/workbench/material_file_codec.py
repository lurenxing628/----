"""Bytes-only material CSV/XLSX codec, using the repository's pinned openpyxl.

Defined headers are HEADERS or their COLUMNS aliases; business_code is required.
Missing columns and empty cells mean omission. The literal \\N means null (only
spec/unit/remark can be cleared). Double an initial backslash for literal text.
CSV is UTF-8 (optional BOM); exported text has one reversible apostrophe prefix
to prevent formula execution and automatic identifier/date coercion. XLSX text
is explicitly typed as text, never formulas. Numeric XLSX IDs are rejected by
the input contract, since any lost leading zeros cannot be recovered reliably.
created_at is an optional read-only assertion, not an importable timestamp.
"""

from __future__ import annotations

import codecs
import csv
import math
import os
import re
from io import BytesIO, StringIO
from tempfile import SpooledTemporaryFile

import openpyxl
from openpyxl.cell import WriteOnlyCell
from openpyxl.comments import Comment
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from core.errors import ValidationError
from core.models.workbench_material_file import COLUMNS, HEADER_FIELDS, HEADERS, IMPORT_ROW_LIMIT, MaterialFileDownload

_NUMBER = re.compile(r"[+-]?(?:[0-9]+(?:\.[0-9]*)?|\.[0-9]+)(?:[eE][+-]?[0-9]+)?\Z")
_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
_CLEARABLE = ("spec", "unit", "remark")
# Excel worksheet capacity includes the header. Neither limit belongs to CSV.
XLSX_MAX_ROWS = 1_048_576
XLSX_MAX_CELL_CHARACTERS = 32767
_SPOOL_BYTES = 4 * 1024 * 1024


def _file_error(message, row=1, field="file"):
    return ValidationError(message, field=field, details={"row": row})


def _headers(values):
    fields = []
    for value in values:
        if type(value) is not str or value not in HEADER_FIELDS:
            raise _file_error("表头含未知字段或空列，只能使用已定义的物料字段。", field="headers")
        fields.append(HEADER_FIELDS[value])
    if "business_code" not in fields:
        raise _file_error("缺少物料编号列。", field="headers")
    if len(set(fields)) != len(fields):
        raise _file_error("表头包含重复字段（含中英文同义列）。", field="headers")
    return fields


def _csv_rows(content):
    try:
        reader = csv.reader(StringIO(content.decode("utf-8-sig"), newline=""), strict=True)
        while True:
            number = reader.line_num + 1
            row = next(reader, None)
            if row is None:
                break
            yield number, row, {}
    except UnicodeDecodeError as exc:
        raise _file_error("CSV 必须使用 UTF-8 编码，未猜测或替换字符。") from exc
    except csv.Error as exc:
        raise _file_error("CSV 格式错误，请核对引号和分隔符。", reader.line_num) from exc


def _xlsx_rows(content):
    wb = None
    try:
        wb = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
        if len(wb.sheetnames) != 1 or len(wb.worksheets) != 1:
            raise _file_error("物料文件必须只有一张数据工作表，不能静默忽略其他表。")
        ws = wb.worksheets[0]
        # Do not trust a producer's cached dimensions to hide later rows/columns.
        ws.reset_dimensions()
        for number, cells in enumerate(ws.iter_rows(), 1):
            errors = {index: "不接受公式或 Excel 错误单元格，请提供实际值。"
                      for index, cell in enumerate(cells) if cell.data_type in ("f", "e")}
            yield number, [cell.value for cell in cells], errors
    except ValidationError:
        raise
    except Exception as exc:
        raise _file_error("XLSX 文件读取失败，请核对文件格式和内容。") from exc
    finally:
        if wb is not None:
            wb.close()


def _decode_value(value, field, file_format):
    if type(value) is str:
        if file_format == "csv" and value.startswith("'"):
            value = value[1:]
        if value == r"\N":
            return None
        if value.startswith("\\\\"):
            value = value[1:]
        if file_format == "csv" and field == "stock_qty":
            if _NUMBER.fullmatch(value) is None:
                raise ValidationError("库存数量必须是独立数值，不能混入单位、布尔值或分组逗号。", field="stock_qty")
            value = float(value)
    return value


def _parse_row(number, values, cell_errors, fields, file_format):
    parsed, errors = {}, []
    if len(values) > len(fields) and any(v is not None and v != "" for v in values[len(fields):]):
        errors.append({"row": number, "field": "columns", "code": "invalid_input", "message": "数据行有未声明的多余列。"})
    for index, field in enumerate(fields):
        value = values[index] if index < len(values) else None
        if index in cell_errors:
            errors.append({"row": number, "field": field, "code": "invalid_input", "message": cell_errors[index]})
        elif value is not None and value != "":
            try:
                parsed[field] = _decode_value(value, field, file_format)
            except ValidationError as exc:
                errors.append({"row": number, "field": field, "code": "invalid_input", "message": exc.message})
    return {"row": number, "values": parsed, "errors": errors}


def read_material_file(content: bytes, file_format: str):
    if type(content) is not bytes or file_format not in ("csv", "xlsx"):
        raise _file_error("必须提供 CSV/XLSX 的原始文件字节。")
    source = _csv_rows(content) if file_format == "csv" else _xlsx_rows(content)
    try:
        header = next(source, None)
        if header is None or header[2]:
            raise _file_error("文件缺少有效的物料表头。", field="headers")
        fields = _headers(header[1])
        result = []
        for number, values, errors in source:
            if not errors and all(v is None or v == "" for v in values):
                continue
            if len(result) == IMPORT_ROW_LIMIT:
                raise _file_error("单次最多导入 2000 行，未截断或导入前半部分。", number)
            result.append(_parse_row(number, values, errors, fields, file_format))
        return result
    finally:
        source.close()


def _export_value(value, field, number, file_format):
    if field == "stock_qty":
        if value is not None and (type(value) not in (int, float) or not math.isfinite(value) or value < 0):
            raise _file_error("原始库存不是有效的非负数，未转换成零。", number, field)
        return value
    if value is None:
        return r"\N" if field in _CLEARABLE else None
    if type(value) is not str:
        raise _file_error("原始文字字段不是文字，未猜测转换类型。", number, field)
    value = "\\" + value if value.startswith("\\") else value
    _check_export_text(value, field, number, file_format)
    return value


def _check_export_text(value, field, number, file_format):
    if file_format == "csv":
        if "\x00" in value:
            raise _file_error("CSV 读取工具链不支持 NUL 字符，未删除或替换原文。", number, field)
    elif _ILLEGAL.search(value) or len(value) > XLSX_MAX_CELL_CHARACTERS:
        raise _file_error("文字含 XLSX 不支持的控制字符或超过单元格 32767 字符容量，未截断内容。", number, field)


def check_export_capacity(row_count, file_format):
    if file_format not in ("csv", "xlsx"):
        raise _file_error("物料导出仅支持 CSV 和 XLSX。", field="format")
    if file_format == "xlsx" and row_count + 1 > XLSX_MAX_ROWS:
        raise _file_error("XLSX 超过单张工作表 1048576 行容量（含表头），未截断；可导出 CSV。", row_count + 1)


def write_material_file(rows, file_format, *, template=False):
    """Consume rows exactly once; only returned file bytes require full materialization."""
    check_export_capacity(0, file_format)
    filename = "materials-template" if template else "materials"
    if file_format == "csv":
        return _write_csv(rows, filename)
    return _write_xlsx(rows, filename, template)


def _write_csv(rows, filename):
    count = 0
    with SpooledTemporaryFile(max_size=_SPOOL_BYTES, mode="w+b") as buffer:
        text = codecs.getwriter("utf-8-sig")(buffer)
        writer = csv.writer(text, lineterminator="\r\n")
        writer.writerow(HEADERS)
        for count, row in enumerate(rows, 1):
            values = [_export_value(row[field], field, count + 1, "csv") for field in COLUMNS]
            writer.writerow(["'" + value if type(value) is str else value for value in values])
        text.flush()
        buffer.seek(0)
        content = buffer.read()
    return MaterialFileDownload(filename + ".csv", "text/csv; charset=utf-8", content, count)


def _xlsx_headers(ws, template):
    examples = ("000123", "45# 圆钢", "D25", "kg", "12.375", "active / inactive", "采购备注", "只读原始时间")
    headers = []
    for index, value in enumerate(HEADERS):
        cell = WriteOnlyCell(ws, value=value)
        cell.font = Font(bold=True)
        if template:
            cell.comment = Comment("示例：" + examples[index] + "。编号必须为文本；缺列/空格子不改；\\N 仅清规格、单位或备注；创建时间只读。", "APS")
        headers.append(cell)
    return headers


def _xlsx_row(ws, row, number):
    cells = []
    for field in COLUMNS:
        value = _export_value(row[field], field, number, "xlsx")
        cell = WriteOnlyCell(ws, value=value)
        if type(value) is str:
            cell.data_type = "s"
            cell.number_format = "@"
        elif field == "stock_qty":
            cell.number_format = "0.###############"
        cells.append(cell)
    return cells


def _write_xlsx(rows, filename, template):
    wb = openpyxl.Workbook(write_only=True)
    ws = wb.create_sheet("物料")
    count = 0
    try:
        ws.freeze_panes = "A2"
        for index, width in enumerate((22, 28, 28, 14, 18, 18, 40, 24), 1):
            ws.column_dimensions[get_column_letter(index)].width = width
        ws.append(_xlsx_headers(ws, template))
        for count, row in enumerate(rows, 1):
            check_export_capacity(count, "xlsx")
            ws.append(_xlsx_row(ws, row, count + 1))
        ws.auto_filter.ref = "A1:H" + str(count + 1)
        with SpooledTemporaryFile(max_size=_SPOOL_BYTES, mode="w+b") as buffer:
            wb.save(buffer)
            buffer.seek(0)
            content = buffer.read()
        return MaterialFileDownload(filename + ".xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                    content, count)
    finally:
        # Pinned openpyxl 3.0.10 Workbook.close() alone leaves an aborted sheet's temp XML behind.
        if not ws.closed:
            ws.close()
        writer = ws._writer
        if writer is not None and os.path.exists(writer.out):
            writer.cleanup()
        wb.close()
