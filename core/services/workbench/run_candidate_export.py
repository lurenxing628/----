"""Export the exact admitted candidate workspace; no query or second interpretation."""

import codecs
import csv
import os
from tempfile import SpooledTemporaryFile

import openpyxl
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from core.models.workbench_command import canonical_json
from core.models.workbench_resource_file import ResourceFileDownload

from .process_file_xml import preserve_carriage_returns
from .resource_file_writer import _value, check_capacity

HEADERS = ("记录类型", "排产编号", "候选方案编号", "候选方案名称", "候选方案状态", "候选方案完整性", "排产时间",
           "数据版本编号", "读取时间", "范围起点（含）", "范围终点（不含）", "行编号", "工序编号", "批次编号",
           "批次名称", "零件名称", "工序顺序", "工序名称", "数量", "交期", "安排开始", "安排结束",
           "任务来源", "排产时锁定", "设备名称", "人员名称", "供应商名称", "排产时报工状态", "未安排原因", "数据缺项")
# 存的是英文代号，写进导出格子时换成用户看得懂的说法；代号本身留在接口里不动。
RECORD_KINDS = {"task": "工序安排", "unplanned_operation": "未安排工序", "candidate": "空范围"}
CANDIDATE_STATUS = {"completed": "已排完", "partial": "部分排完", "failed": "没排成", "skipped": "本次跳过"}
COMPLETENESS = {"complete": "完整", "partial": "不完整", "no_result": "没有结果", "unknown": "暂无数据"}
TASK_SOURCES = {"internal": "自制", "external": "外协"}
REPORT_STATES = {"unreported": "待报工", "started": "已开工", "partial": "部分完成",
                 "paused": "已暂停", "exception": "异常", "complete": "已完工"}


def _text(texts, value):
    return None if value is None else texts[value]


def export_rows(data, snapshot):
    candidate, generation = data["candidate"], data["generation"]
    common = [data["candidate"]["run_ref"], candidate["candidate_ref"], candidate["label"],
              _text(CANDIDATE_STATUS, candidate["status"]), _text(COMPLETENESS, candidate["completeness"]),
              generation["accepted_at"], snapshot["snapshot_ref"], snapshot["as_of"],
              data["time_scope"]["range_start"], data["time_scope"]["range_end"]]
    records = [("task", row) for row in data["tasks"]] + [("unplanned_operation", row) for row in data["unplanned_operations"] or []]
    if not records:
        records = [("candidate", {})]
    for kind, row in records:
        values = [RECORD_KINDS[kind]] + common + [row.get(key) for key in ("row_ref", "operation_ref", "batch_ref", "batch_label",
            "part_label", "sequence", "process_label", "quantity", "due_date", "start", "end")]
        values.append(_text(TASK_SOURCES, row.get("source")))
        locked = row.get("locked")
        values.append(None if locked is None else "是" if locked else "否")
        values.extend((row.get(key) or {}).get("label") for key in ("machine", "operator", "supplier"))
        values.extend([_text(REPORT_STATES, (row.get("execution_at_generation") or {}).get("execution_state")),
                       (row.get("reason") or {}).get("message"),
                       canonical_json(row.get("data_gaps", []) + data["data_gaps"] + candidate["data_gaps"]
                                      + generation["data_gaps"])])
        yield values


def _values(values, number, fmt):
    return [_value(value, "default_days" if type(value) in (int, float) else "text", number, fmt) for value in values]


def write_run_candidate_export(data, snapshot, fmt):
    count = max(1, len(data["tasks"]) + len(data["unplanned_operations"] or []))
    check_capacity(count, fmt)
    content = _csv(data, snapshot) if fmt == "csv" else _xlsx(data, snapshot)
    mime = "text/csv; charset=utf-8" if fmt == "csv" else "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    return ResourceFileDownload("候选范围导出." + fmt, mime, content, count)


def _csv(data, snapshot):
    with SpooledTemporaryFile(max_size=4 * 1024 * 1024, mode="w+b") as buffer:
        text = codecs.getwriter("utf-8-sig")(buffer)
        writer = csv.writer(text, lineterminator="\r\n")
        writer.writerow(HEADERS)
        for number, row in enumerate(export_rows(data, snapshot), 2):
            writer.writerow(["'" + value if type(value) is str else value for value in _values(row, number, "csv")])
        text.flush()
        buffer.seek(0)
        return buffer.read()


def _xlsx(data, snapshot):
    wb = openpyxl.Workbook(write_only=True)
    ws = wb.create_sheet("候选任务")
    try:
        ws.freeze_panes = "A2"
        for index in range(1, len(HEADERS) + 1):
            ws.column_dimensions[get_column_letter(index)].width = 24 if index not in (2, 3, 8, 12, 13, 14) else 52
        headers = [WriteOnlyCell(ws, value=value) for value in HEADERS]
        for cell in headers:
            cell.font = Font(bold=True)
        ws.append(headers)
        for number, row in enumerate(export_rows(data, snapshot), 2):
            cells = []
            for value in _values(row, number, "xlsx"):
                cell = WriteOnlyCell(ws, value=value)
                if type(value) is str:
                    cell.data_type, cell.number_format = "s", "@"
                cells.append(cell)
            ws.append(cells)
        preserve_carriage_returns(ws)
        with SpooledTemporaryFile(max_size=4 * 1024 * 1024, mode="w+b") as buffer:
            wb.save(buffer)
            buffer.seek(0)
            return buffer.read()
    finally:
        if not ws.closed:
            ws.close()
        if ws._writer is not None and os.path.exists(ws._writer.out):
            ws._writer.cleanup()
        wb.close()
