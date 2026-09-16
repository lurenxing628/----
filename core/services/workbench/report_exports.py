"""Download the entire snapshot cohort, using the existing safe XLSX renderer."""

from __future__ import annotations

import csv
import io

import openpyxl
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment

from core.models.workbench_command import WorkbenchCommandRejected, canonical_json
from core.services.common.excel_templates import _sanitize_export_cell
from core.services.report.exporters.xlsx import SUMMARY_IDENTITY_LABELS, _append_write_only_row
from core.services.report.report_engine import ReportExport

from .execution_ledger_projection import COMPLETION_BASIS_TEXT, DATA_QUALITY_TEXT

# 导出标题与列名都是用户直接看到的文字，统一走词表，不再暴露专题代号。
TOPIC_TITLES = {"delivery": "工序完成情况", "records": "报工记录", "machines": "设备工时",
                "people": "人员工时", "quality": "数据完整性",
                "overdue": "超期批次", "utilization": "资源负荷", "downtime": "停机影响"}
# 这两行写的是记录编号，必须原样往返，不能被 Excel 当成公式。
IDENTITY_LABELS = SUMMARY_IDENTITY_LABELS
# 存的是英文代号，写进导出格子时换成用户看得懂的说法；代号本身留在接口里不动。
CELL_TEXTS = {
    "data_quality": DATA_QUALITY_TEXT,
    "completion_basis": COMPLETION_BASIS_TEXT,
    "recorded_at_time_basis": {"factory_local": "现场记录时间", "legacy_storage": "历史系统导入的原始时间"},
    "source": {"manual": "手工填写", "excel": "Excel 导入"},
}


def _value(value):
    if value is None:
        return "未知"
    if isinstance(value, dict) or isinstance(value, list) and any(isinstance(item, (dict, list)) for item in value):
        return canonical_json(value)
    if isinstance(value, list):
        return "；".join(str(item) for item in value)
    if isinstance(value, bool):
        return "是" if value else "否"
    return value


def _cell(column_key, value):
    texts = CELL_TEXTS.get(column_key)
    if texts is None or value is None:
        return _value(value)
    return texts[value]


def metadata(data, snapshot):
    return [["数据来源", data["provenance"]], ["计划", data["plan"]["display_name"]],
            [IDENTITY_LABELS[0], data["plan"]["plan_ref"]], ["数据截至", snapshot["as_of"]],
            [IDENTITY_LABELS[1], snapshot["snapshot_ref"]], ["筛选范围", canonical_json(data["scope"])],
            ["统计说明与待补资料", "；".join(data["data_gaps"])]]


def ensure_export_size(engine, count):
    decision = engine._build_export_decision(count)
    if decision.mode == "reject_need_async":
        raise WorkbenchCommandRejected("export_too_large", "这次筛出来的条数超过一次能导出的上限，没有开始下载。请缩小筛选范围后再点「导出」。", 413)
    return decision


def _export_values(columns, rows, format_name):
    values = [[_cell(column["key"], row.get(column["key"])) for column in columns] for row in rows]
    if format_name == "xlsx" and any(isinstance(value, str) and len(value) > 32767 for row in values for value in row):
        raise WorkbenchCommandRejected("export_too_large", "单元格内容超过 XLSX 上限，请改用 CSV 导出。", 413)
    return values


def _append_xlsx_metadata(sheet, data, snapshot):
    for label, value in metadata(data, snapshot):
        if label not in IDENTITY_LABELS:
            _append_write_only_row(sheet, [label, value])
            continue
        # References must round-trip exactly without ever becoming formulas.
        cells = [WriteOnlyCell(sheet, value=label), WriteOnlyCell(sheet, value=value)]
        cells[1].data_type = "s"
        for cell in cells:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        sheet.append(cells)


def export_table(engine, data, rows, snapshot, format_name):
    if format_name not in ("csv", "xlsx"):
        raise WorkbenchCommandRejected("invalid_input", "只能导出 CSV 或 XLSX，没有开始下载。请重新选择格式。", 400)
    if not rows:
        raise WorkbenchCommandRejected("empty_export", "当前筛选范围里没有可导出的结果，没有开始下载。请放宽筛选条件后重试。", 422)
    decision = ensure_export_size(engine, len(rows))
    columns = data["columns"]
    values = _export_values(columns, rows, format_name)
    filename = TOPIC_TITLES[data["topic"]] + "-" + snapshot["as_of"].replace(":", "")
    if format_name == "csv":
        output = io.StringIO(newline="")
        writer = csv.writer(output)
        writer.writerow([column["label"] for column in columns] + ["数据截至", IDENTITY_LABELS[1], "筛选范围", "数据来源", "统计说明与待补资料"])
        for row in values:
            writer.writerow([_sanitize_export_cell(value) for value in row + [snapshot["as_of"], snapshot["snapshot_ref"], canonical_json(data["scope"]), data["provenance"], "；".join(data["data_gaps"])]])
        return ReportExport(filename + ".csv", "text/csv;charset=utf-8", io.BytesIO(output.getvalue().encode("utf-8-sig")), estimated_rows=len(rows))
    workbook = openpyxl.Workbook(write_only=True)
    output = io.BytesIO()
    try:
        sheet = workbook.create_sheet("范围与计算方式")
        _append_xlsx_metadata(sheet, data, snapshot)
        for key, value in data["summary"].items():
            _append_write_only_row(sheet, [key, _value(value)])
        sheet = workbook.create_sheet("范围全部结果")
        _append_write_only_row(sheet, [column["label"] for column in columns], is_header=True)
        for row in values:
            _append_write_only_row(sheet, row)
        workbook.save(output)
        output.seek(0)
    finally:
        workbook.close()
    return ReportExport(filename + ".xlsx", engine.XLSX_CONTENT_TYPE, output, mode=decision.mode, estimated_rows=len(rows))
