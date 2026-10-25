"""同次比较目录、数值和差值直接写入离线总览。"""

import codecs
import csv
import json
import re
from tempfile import SpooledTemporaryFile

from core.infrastructure.read_evidence import read_evidence_scope
from core.models.workbench_action_values import FileDownload
from core.models.workbench_run_candidate import reject
from core.services.common.excel_instruction_sheet import close_write_only_sheets
from core.services.workbench.facts.candidate_store import CandidateStore
from core.services.workbench.facts.decision_metrics import metric_definitions
from core.services.workbench.facts.table_cells import number_text

from .decision import decision_overview, decision_summary
from .decision_preparation_export import (
    CHANGEOVER_HEADERS,
    PREPARATION_HEADERS,
    RUNTIME_OPERATION_HEADERS,
    changeover_rows,
    preparation_rows,
    runtime_operation_values,
)
from .decision_preparation_queries import normalize_preparation_filters, preparation_details
from .decision_process import normalize_process_filters, process_details, process_export_records
from .decision_process_export import process_csv, process_export_tables
from .decision_queries import decision_details, normalize_filters
from .decision_recommendation_export import (
    recommendation_csv,
    recommendation_export_records,
    recommendation_export_tables,
)
from .decision_resource_export import (
    COVERAGE_OPERATION_HEADERS,
    RESOURCE_CONFLICT_HEADERS,
    RESOURCE_DAY_HEADERS,
    RESOURCE_HEADERS,
    coverage_operation_rows,
    resource_conflict_rows,
    resource_day_rows,
    resource_rows,
)
from .decision_resources import normalize_resource_filters, resource_details
from .decision_source import DecisionSource

HEADERS = ("排产编号", "归档编号", "查看基准编号", "查看基准名称", "所选方案编号", "推荐方案编号",
           "范围起点（含）", "范围终点（不含）", "候选编号", "候选名称", "搜索状态", "应排范围状态",
           "交付覆盖状态", "指标", "单位", "指标版本", "指标范围", "指标计算方式", "覆盖状态", "数值",
           "已知小计", "已知数量", "未知数量", "总数量", "基准数值", "差值（候选减基准）", "差值状态", "变化方向", "缺项说明")
BATCH_HEADERS = ("排产编号", "查看基准永久编号", "所选方案编号", "批次永久编号", "批次编号", "零件编号", "零件名称", "数量", "优先级",
                 "交期", "交期状态", "交期缺项", "基准工序安排", "所选工序安排", "基准交付状态", "所选交付状态",
                 "基准预计完工", "所选预计完工", "基准已排部分预计结束", "所选已排部分预计结束", "结束变化（小时）",
                 "基准拖期（小时）", "所选拖期（小时）", "拖期变化（小时）", "变化类型", "工序总数",
                 "基准已排工序数", "所选已排工序数", "基准未排工序数", "所选未排工序数", "基准缺料未安排工序数", "所选缺料未安排工序数",
                 "基准其他未安排工序数", "所选其他未安排工序数", "所选缺料未安排工序编号", "所选其他未安排工序编号", "基准缺项", "所选缺项")
BATCH_HEADERS += ("当前待处理类别", "已确认变化工序数", "变化待确认工序数", "工序变化覆盖状态")
OPERATION_HEADERS = ("排产编号", "查看基准永久编号", "所选方案编号", "批次永久编号", "批次编号", "工序永久编号", "零件名称", "工序号",
                     "工序名称", "单件编号", "数量", "基准开始", "所选开始", "开始变化（小时）", "基准结束", "所选结束", "结束变化（小时）",
                     "基准设备编号", "基准设备", "所选设备编号", "所选设备", "设备变化", "基准人员编号", "基准人员", "所选人员编号", "所选人员", "人员变化",
                     "基准来源", "所选来源", "来源变化", "变化类型", "标准准备工时", "标准单件工时", "执行目标量", "剩余数量", "标准总工时",
                     "工时计算方式", "工时状态", "排产时保护", "基准锁定", "所选锁定", "排产时执行状态", "排产时已完数量", "排产时首次实开", "排产时确认完工",
                     "基准多段安排", "所选多段安排", "基准缺项", "所选缺项")
OPERATION_HEADERS += RUNTIME_OPERATION_HEADERS
OPERATION_HEADERS += ("基准排产状态", "所选排产状态", "排产状态变化", "基准未排说明", "所选未排说明", "已确认变化字段",
                      "基准供方编号", "基准供方", "所选供方编号", "所选供方", "供方变化")
_STATE = {"complete": "完整", "partial": "部分", "unavailable": "缺少资料", "not_applicable": "范围无适用对象",
          "completed": "计算完成", "failed": "失败", "skipped": "跳过", "queued": "待计算", "running": "计算中",
          "unknown": "缺少资料", "no_result": "尚无结果", "comparable": "可比较", "different_coverage": "覆盖范围不同",
          "better": "改善", "worse": "增加代价", "equal": "相同", "neutral": "中性变化"}
_STATE.update(available="已保存", absent="排产时无正式计划", selection_required="需要选择基准")
_DETAIL_STATE = {"specified": "有交期", "unspecified": "未设置交期", "invalid": "交期无效", "complete": "工序已排齐",
                 "deferred": "因缺料暂未安排", "incomplete": "有工序未安排", "unknown": "缺少资料", "overdue": "拖期", "on_time": "按期",
                 "not_applicable": "未设置交期", "earlier": "提前结束", "later": "延后结束", "new_overdue": "新增拖期",
                 "relieved_overdue": "解除拖期", "completion_improved": "工序已排齐", "completion_worsened": "有工序未安排",
                 "unchanged": "相同", "added": "新增安排", "removed": "移除安排", "moved": "时间调整", "machine_changed": "设备调整",
                 "operator_changed": "人员调整", "source_changed": "来源调整"}
_DETAIL_STATE.update(operation_changed="工序安排变化", status_changed="排产状态变化", supplier_changed="供方调整",
                     arranged="已安排", missing="未安排", unavailable="尚无结果", unplanned="其他未安排")
_SCOPE = {"full_admission": "整次排产范围", "full_candidate": "完整候选安排", "completed_batches_only": "工序已排齐的批次",
          "candidate_search": "本候选计算", "saved_schedule_resources": "算法保存的资源范围",
          "saved_algorithm_batches": "算法评价批次集合"}
_SCOPE["fixed_resource_window"] = "相同设备人员与班表范围"
_SCOPE["fixed_changeover_window"] = "后一道工序在所选范围内开始的设备工种切换"
_SOURCE = {"admission": "排产时归档", "candidate_tasks": "保存的候选安排", "archived_delivery": "保存的批次工序安排和交期",
           "saved_algorithm": "排产时算法统计", "candidate_comparison": "按同一道工序比较安排"}
_SOURCE["archived_resource_calendar"] = "排产时归档班表与保存安排"
_SOURCE.update(runtime_operation_effort="计算时保存的准备和加工所需标准工时", archived_changeover_events="同一设备上相邻工序的工种变化")
_GROUP = {"completeness": "工序安排", "delivery": "交付", "time": "时间", "preparation": "准备与切换",
          "resources": "资源", "stability": "调整"}
_IMPROVEMENT = {"lower": "越低越好", "higher": "越高越好", "neutral": "按实际需求判断"}
_NUMERIC_TEXT = re.compile(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?")


def _label(value):
    return _STATE.get(value, value)


def _display_number(value, definition):
    # 占用率在算法记录中是比例；总览文件按目录的百分数单位给出数字。
    return value * 100 if value is not None and definition["unit"] == "%" else value


def overview_rows(data):
    scope, baseline = data["scope"], data["baseline"]
    common = [data["run"]["run_ref"], data["capture_ref"], _baseline_ref(baseline), baseline.get("label"),
              scope["selected_candidate_ref"], data["recommendation"]["candidate_ref"], scope["range_start"], scope["range_end"]]
    for candidate in data["candidates"]:
        coverage = candidate["coverage"]
        identity = [candidate["candidate_ref"], candidate["label"], _label(candidate["status"]),
                    _label(coverage["schedulable"]["state"]), _label(coverage["delivery"]["state"])]
        for definition in data["metric_definitions"]:
            metric, delta = candidate["metrics"][definition["key"]], candidate["deltas"][definition["key"]]
            value, subtotal, reference, difference = [_display_number(number, definition) for number in
                (metric["value"], metric["known_subtotal"], delta["reference_value"], delta["value"])]
            yield common + identity + [definition["label"], definition["unit"], metric["metric_version"], _SCOPE[metric["scope"]],
                _SOURCE[definition["source"]], _label(metric["state"]), value, subtotal, metric["known_count"],
                metric["unknown_count"], metric["total_count"], reference, difference, _label(delta["state"]),
                _label(delta["direction"]), (metric["reason"] or {}).get("message")]


def _baseline_ref(baseline):
    return baseline.get("candidate_ref") if baseline["kind"] == "candidate" else baseline.get("baseline_ref")


def _detail_label(value: str) -> str:
    return _DETAIL_STATE.get(value, value)


def _issues(value):
    if isinstance(value, str):
        return value
    if isinstance(value, dict):
        return value.get("message") or value.get("code")
    return "；".join(str(item.get("message") or item.get("code")) if isinstance(item, dict) else str(item) for item in value or ()) or None


def batch_rows(data):
    common = [data["run_ref"], _baseline_ref(data["baseline"]), data["candidate"]["candidate_ref"]]
    for row in data["items"]:
        before, after = row["before"], row["after"]
        yield common + [row["batch_ref"], row["batch_id"], row["part_no"], row["part_label"], row["quantity"], row["priority"],
            row["due_date"], _detail_label(row["due_state"]), _issues(row["due_issue"]),
            _detail_label(before["completion_state"]), _detail_label(after["completion_state"]), _detail_label(before["risk"]), _detail_label(after["risk"]),
            before["planned_finish"], after["planned_finish"], before["partial_planned_finish"], after["partial_planned_finish"], row["finish_delta_hours"],
            before["delay_hours"], after["delay_hours"], row["delay_delta_hours"], "、".join(_detail_label(kind) for kind in row["change_kinds"]),
            after["operation_count"], before["scheduled_operation_count"], after["scheduled_operation_count"],
            before["unscheduled_operation_count"], after["unscheduled_operation_count"], before["deferred_operation_count"], after["deferred_operation_count"],
            before["ordinary_unscheduled_operation_count"], after["ordinary_unscheduled_operation_count"],
            "、".join(row["deferred_operation_refs"] or ()) or None, "、".join(row["unplanned_operation_refs"] or ()) or None,
            _issues(before["issues"]), _issues(after["issues"]),
            "、".join(_detail_label(kind) for kind in row["attention_kinds"]) or None,
            row["operation_change_count"], row["operation_unknown_count"], _label(row["operation_changes_state"])]


def _resource(side, kind, field):
    return ((side or {}).get(kind) or {}).get(field)


def _boolean(value):
    return None if value is None else "是" if value else "否"


def _segments(side):
    segments = (side or {}).get("segments")
    if not segments:
        return None
    # 多段事实单列保存，独立时间和资源差值继续为空。
    return json.dumps(segments, ensure_ascii=False, separators=(",", ":"))


def operation_rows(data):
    common = [data["run_ref"], _baseline_ref(data["baseline"]), data["candidate"]["candidate_ref"]]
    for row in data["items"]:
        before, after = row["before"] or {}, row["after"] or {}
        effort, constraints = row["effort"], row["constraints"]
        yield common + [row["batch_ref"], row["batch_label"], row["operation_ref"], row["part_label"], row["sequence"], row["process_label"], row["piece_id"], row["quantity"],
            before.get("start"), after.get("start"), row["start_delta_hours"], before.get("end"), after.get("end"), row["end_delta_hours"],
            _resource(before, "machine", "ref"), _resource(before, "machine", "label"), _resource(after, "machine", "ref"), _resource(after, "machine", "label"), _boolean(row["machine_changed"]),
            _resource(before, "operator", "ref"), _resource(before, "operator", "label"), _resource(after, "operator", "ref"), _resource(after, "operator", "label"), _boolean(row["operator_changed"]),
            before.get("source"), after.get("source"), _boolean(row["source_changed"]), "、".join(_detail_label(kind) for kind in row["change_kinds"]),
            effort["setup_standard_hours"], effort["unit_standard_hours"], effort["target_quantity"], effort["remaining_quantity"], effort["standard_total_hours"],
            "归档工序标准×归档执行目标量", _label(effort["state"]), _boolean(constraints["protected_at_admission"]),
            _boolean(constraints["reference_locked"]), _boolean(constraints["selected_locked"]), constraints["execution_state"],
            constraints["known_completed_quantity"], constraints["first_actual_start"], constraints["confirmed_finish"],
            _segments(before), _segments(after), _issues(before.get("data_gaps")), _issues(after.get("data_gaps"))] + runtime_operation_values(row) + [
            _detail_label(row["before_state"]["status"]), _detail_label(row["after_state"]["status"]), _boolean(row["status_changed"]),
            _issues(row["before_state"]["reason"]), _issues(row["after_state"]["reason"]),
            "、".join({"start": "开工", "end": "完工", "machine": "设备", "operator": "人员", "supplier": "供方",
                      "source": "来源", "status": "排产状态"}[field] for field in row["changed_fields"]) or None,
            _resource(before, "supplier", "ref"), _resource(before, "supplier", "label"),
            _resource(after, "supplier", "ref"), _resource(after, "supplier", "label"), _boolean(row["supplier_changed"])]


def read_decision_export(conn, run_ref, section, *, filter_section=None, **query):
    selected = query.get("selected_candidate_ref")
    details = ("batches", "operations", "resources", "resource-days", "preparation", "changeovers", "process")
    detail_section = filter_section or (section if section in details else None)
    if filter_section is not None and (filter_section not in details or section == "overview" or section in details and filter_section != section):
        reject("invalid_input", "导出筛选页签与当前明细不一致，请重新导出。", 422)
    filters = {key: query.pop(key) for key in ("batch_ref", "resource_ref", "change_kind", "delivery_filter", "selected_batch_ref",
        "resource_kind", "pressure_filter", "grain", "pressure_threshold", "display_start", "display_end", "bucket_start", "bucket_end",
        "machine_ref", "allocation_state", "work_scope", "event_side", "process_candidate_ref", "objective", "record_kind", "sort", "order") if key in query}
    query.pop("page", None)
    query.pop("size", None)
    query.pop("heatmap_page", None)
    query.pop("heatmap_size", None)
    if filters and detail_section is None:
        reject("invalid_input", "请指定明细筛选页签后导出筛选结果。", 422)
    if detail_section is not None and detail_section != "process" and selected is None:
        reject("invalid_input", "明细导出需要指定所选候选方案。", 422)
    store = CandidateStore(conn)
    with store.snapshot(), read_evidence_scope(conn):
        source = DecisionSource(store, run_ref, **query)
        active_filters = _export_filters(detail_section, filters) if detail_section is not None else None
        selected_process = None
        if detail_section == "process":
            selected_process = process_details(source, active_filters, all_rows=True)
            active_filters = selected_process["filters"]
        snapshot_scope = {"kind": "run-decision", **source.scope}
        if active_filters is not None:
            snapshot_scope["filters"] = active_filters
        data = decision_overview(source) if section in ("overview", "all") else (
            decision_summary(source) if section == "recommendation" else _export_context(source))
        if section in ("recommendation", "all"):
            data["export_recommendation"] = recommendation_export_records(source, data)
        if section in ("process", "all"):
            data["export_process"] = process_export_records(
                source,
                active_filters if detail_section == "process" else None,
                include_all_objectives=section == "all",
                selected=selected_process,
            )
        for name in ("batches", "operations"):
            if section not in (name, "all"):
                continue
            batch = filters.get("batch_ref") if detail_section in ("batches", "operations") else None
            current_filters = active_filters if name == detail_section else normalize_filters(name, batch_ref=batch)
            data["export_" + name] = decision_details(source, name, current_filters, all_rows=True)
        if section in ("resources", "all"):
            if active_filters is not None and detail_section in ("resources", "resource-days"):
                current_filters = active_filters if detail_section == "resources" else normalize_resource_filters(
                    "resources", **{key: active_filters[key] for key in ("resource_kind", "resource_ref", "grain", "pressure_threshold", "display_start", "display_end")})
                resource_documents = [resource_details(source, current_filters, all_rows=True)]
            else:
                resource_documents = [resource_details(source, normalize_resource_filters("resources", resource_kind=kind), all_rows=True)
                                      for kind in ("machine", "operator")]
            data["export_resources"] = resource_documents
        if section == "resource-days" or section == "all" and detail_section == "resource-days":
            data["export_resource_days"] = resource_details(source, active_filters, all_rows=True)
        if section in ("resources", "resource-days", "all"):
            resource_document = data["export_resources"][0] if "export_resources" in data else data["export_resource_days"]
            data["export_coverage_operations"] = list(coverage_operation_rows(source, resource_document["coverage"]))
        for name in ("preparation", "changeovers"):
            if section not in (name, "all"):
                continue
            machine = filters.get("machine_ref") if detail_section in ("preparation", "changeovers") else None
            current_filters = active_filters if name == detail_section else normalize_preparation_filters(name, machine_ref=machine)
            data["export_" + name] = preparation_details(source, current_filters, all_rows=True)
        data["export_filter_section"] = detail_section
        data["export_filters"] = active_filters
        return data, snapshot_scope, source.state()


def _export_filters(section, filters):
    allowed = {"batches": {"batch_ref", "change_kind", "sort", "order", "delivery_filter", "selected_batch_ref"},
               "operations": {"batch_ref", "resource_ref", "change_kind", "sort", "order"},
               "resources": {"resource_kind", "resource_ref", "pressure_filter", "sort", "order", "grain", "pressure_threshold", "display_start", "display_end"},
               "resource-days": {"resource_kind", "resource_ref", "grain", "pressure_threshold", "display_start", "display_end", "bucket_start", "bucket_end"},
               "preparation": {"batch_ref", "machine_ref", "allocation_state", "work_scope", "sort", "order"},
               "changeovers": {"machine_ref", "event_side", "sort", "order"},
               "process": {"process_candidate_ref", "objective", "record_kind", "sort", "order"}}[section]
    if set(filters) - allowed:
        reject("invalid_input", "导出条件不适用于当前页签，请重新导出。", 422)
    if section in ("preparation", "changeovers"):
        return normalize_preparation_filters(section, **filters)
    if section == "process":
        return normalize_process_filters(**filters)
    return normalize_resource_filters(section, **filters) if section in ("resources", "resource-days") else normalize_filters(section, **filters)


def _export_context(source):
    return {"run": source.run, "scope": source.scope, "capture_ref": source.run["run_ref"],
            "baseline": source.baseline, "selection_basis": source.basis,
            "metric_definitions": metric_definitions(),
            "recommendation": {"candidate_ref": source.basis["selected_candidate_ref"],
                               "recommendation_version": "decision-recommendation-v1"}}


def write_run_decision_export(data, snapshot, fmt, section="overview"):
    if fmt == "csv" and section in ("recommendation", "process"):
        csv_table = recommendation_csv(data, snapshot) if section == "recommendation" else process_csv(data, snapshot)
        return FileDownload("排产方案比较.csv", "text/csv; charset=utf-8", csv_bytes(*csv_table), len(csv_table[1]))
    tables = {}
    if section in ("overview", "all"):
        tables["方案总览"] = (HEADERS, list(overview_rows(data)))
    if section in ("batches", "all"):
        tables["批次影响"] = (BATCH_HEADERS, list(batch_rows(data["export_batches"])))
    if section in ("operations", "all"):
        tables["工序变化"] = (OPERATION_HEADERS, list(operation_rows(data["export_operations"])))
    if section in ("resources", "all"):
        tables["资源影响"] = (RESOURCE_HEADERS, list(resource_rows(data["export_resources"])))
    if section == "resource-days" or section == "all" and "export_resource_days" in data:
        period_document = data["export_resource_days"]
        period_title = {"day": "资源逐日", "hour": "资源小时", "shift": "资源班次", "week": "资源逐周"}[period_document["filters"]["grain"]]
        tables[period_title] = (RESOURCE_DAY_HEADERS, list(resource_day_rows(period_document)))
        if fmt == "xlsx":
            tables["资源冲突"] = (RESOURCE_CONFLICT_HEADERS, list(resource_conflict_rows(data["export_resource_days"])))
    if fmt == "xlsx" and section in ("resources", "resource-days", "all"):
        tables["覆盖工序"] = (COVERAGE_OPERATION_HEADERS, data["export_coverage_operations"])
    if section in ("preparation", "all"):
        tables["准备需求"] = (PREPARATION_HEADERS, list(preparation_rows(data["export_preparation"])))
    if section in ("changeovers", "all"):
        tables["工种切换"] = (CHANGEOVER_HEADERS, list(changeover_rows(data["export_changeovers"])))
    if section in ("recommendation", "all"):
        tables.update(recommendation_export_tables(data, snapshot))
    if section in ("process", "all"):
        tables.update(process_export_tables(data, snapshot))
    content = csv_bytes(*next(iter(tables.values()))) if fmt == "csv" else _xlsx(data, snapshot, tables)
    mime = "text/csv; charset=utf-8" if fmt == "csv" else "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    return FileDownload("排产方案比较." + fmt, mime, content, sum(len(rows) for _, rows in tables.values()))


def csv_bytes(headers, rows):
    with SpooledTemporaryFile(max_size=4 * 1024 * 1024, mode="w+b") as buffer:
        text = codecs.getwriter("utf-8-sig")(buffer)
        writer = csv.writer(text, lineterminator="\r\n")
        writer.writerow(headers)
        for row in rows:
            writer.writerow([csv_cell(value, identity=headers[index].endswith("编号")) for index, value in enumerate(row)])
        text.flush()
        buffer.seek(0)
        return buffer.read()


def csv_cell(value, *, identity):
    if value is None:
        return ""
    if type(value) in (int, float):
        return number_text(value)
    if type(value) is str and (value.lstrip().startswith(("=", "+", "-", "@"))
                              or identity and _NUMERIC_TEXT.fullmatch(value)):
        return "'" + value
    return value


def append_sheet_row(sheet, row):
    from openpyxl.cell import WriteOnlyCell

    cells = []
    for value in row:
        cell = WriteOnlyCell(sheet, value=value)
        if type(value) is str:
            cell.data_type, cell.number_format = "s", "@"
        cells.append(cell)
    sheet.append(cells)


def _xlsx(data, snapshot, tables):
    import openpyxl
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook(write_only=True)
    try:
        for title, (headers, rows) in tables.items():
            sheet = wb.create_sheet(title)
            sheet.freeze_panes = "D2"
            for index in range(1, len(headers) + 1):
                sheet.column_dimensions[get_column_letter(index)].width = 24
            append_sheet_row(sheet, headers)
            for row in rows:
                append_sheet_row(sheet, row)
        sheet = wb.create_sheet("指标计算方式")
        append_sheet_row(sheet, ("指标", "分组", "单位", "精度", "改善方向", "范围", "来源", "常用指标", "指标版本"))
        for item in data["metric_definitions"]:
            append_sheet_row(sheet, [item["label"], _GROUP[item["group"]], item["unit"], item["precision"],
                           _IMPROVEMENT[item["improvement"]], _SCOPE[item["scope"]], _SOURCE[item["source"]]]
                    + ["是" if item["common"] else "否", item["metric_version"]])
        sheet = wb.create_sheet("比较范围")
        scope = data["scope"]
        entries = [("排产编号", data["run"]["run_ref"]), ("归档编号", data["capture_ref"]),
                   ("读取编号", snapshot["snapshot_ref"]), ("读取时间", snapshot["as_of"]),
                   ("比较模式", "排产时的正式计划" if scope["comparison_mode"] == "admission_official" else "同次候选"),
                   ("查看基准编号", _baseline_ref(data["baseline"])), ("查看基准状态", _label(data["baseline"]["state"])),
                   ("正式基准版本", data["baseline"].get("version")), ("正式来源候选编号", data["baseline"].get("source_candidate_ref")),
                   ("所选方案编号", scope["selected_candidate_ref"]),
                   ("推荐方案编号", data["recommendation"]["candidate_ref"]), ("范围起点（含）", scope["range_start"]),
                   ("推荐版本", data["recommendation"]["recommendation_version"]), ("推荐记录格式版本", data["selection_basis"]["selection_evidence_version"]),
                   ("范围终点（不含）", scope["range_end"]), ("批次数", len(scope["batch_refs"])),
                   ("工序数", len(scope["operation_refs"])), ("差值方向", "候选减查看基准"),
                   ("资源视角", "本次所选批次" if scope["resource_view"] == "selected_batches" else "方案全部安排"),
                   ("共同设备数", len(scope["resource_refs"]["machine"])), ("共同人员数", len(scope["resource_refs"]["operator"])),
                   ("资源占用计算方式", "保存安排区间与归档可排窗口的交集；休息时段的保留跨度单列"),
                   ("资源冲突计算方式", "归档可排窗口内同时占用同一资源的不同永久工序"),
                   ("零容量工作证据", "安排跨度不证明实际工作；未保存工作片段的零容量时段留空，班表完整状态单列"),
                   ("周占用比例计算方式", "所选周内班表内占用工时之和除以可排工时之和；日峰值和高负荷日保留"),
                   ("小时筛选范围", "当前显示小时区间；完整比较窗口单列"),
                   ("筛选页签", data.get("export_filter_section")), ("筛选条件", _filter_description(data.get("export_filters"))),
                   ("明细范围", "筛选结果全部页"), ("标准工时计算方式", "保存的工时定额乘以其对应的加工目标数量"),
                   ("所需工时计算方式", "计算时保存的准备和加工所需标准工时；原加工目标与本次报工目标数量分别保留"),
                   ("工种切换计算方式", "同一设备上相邻工序的工种不同，后一道工序开始时间在所选范围内（含开始、不含结束）"),
                   ("不计工种切换的工序", "外协工序，以及保存记录中明确的零工时工序"),
                   ("准备工时统计范围", "本次所选批次；新安排、其他未安排、缺料未安排分别汇总；原安排、已开工和已完工工序分别列出"),
                   ("准备差值方向", "候选减基准，中性展示实际需求与安排覆盖变化"),
                   ("推荐方案比较范围", "计算时推荐的方案与当前查看基准比较；交期差值只比较双方工序都已排齐的批次"),
                   ("推荐工序安排范围", "本次所选批次的工序安排及未安排工序，包含所选时间范围之外及明细其他页的工序"),
                   ("推荐依据来源", "计算时保存的目标数值、比较方案、指标相同时的选择顺序、交期上限和关键工序时间评价；旧记录未保存的项留空")]
        if "export_process" in data:
            entries += [("过程范围", "保存本方案的整个计算过程，包含当前时间范围之外的计算尝试"),
                        ("过程工序统计范围", "过程中的工序数、排齐批次数和首次排齐标记，按本次参与计算工序统计；业务未安排工序见工序安排和剩余需求。"),
                        ("过程耗时单位", "毫秒；候选自身处理耗时与来源方案实际计算耗时分别保存"),
                        ("已保存代表尝试", "按计算时的目标数值排名；时间、是否留作继续搜索和相对当时比较方案的变化按原记录保存，未保存的项留空"),
                        ("沿用方案的计算次数", "候选自身未搜索时计数为 0，来源方案的计算次数和编号另列"),
                        ("过程次数计算方式", "评估次数、留作继续搜索次数、严格改善次数和放弃次数分别统计；放弃记录含重复结果，各项互有重合")]
        for row in entries:
            append_sheet_row(sheet, row)
        with SpooledTemporaryFile(max_size=4 * 1024 * 1024, mode="w+b") as buffer:
            wb.save(buffer)
            buffer.seek(0)
            return buffer.read()
    finally:
        close_write_only_sheets(wb.worksheets)
        wb.close()


def _filter_description(filters):
    if filters is None:
        return "完整比较范围"
    return "；".join(key + "=" + str(value) for key, value in filters.items() if value is not None)
