"""File preview/confirm use AJ's real domain with zero-write preview enforcement."""

from io import BytesIO

from openpyxl import load_workbook

from core.services.workbench.field_report_files_codec import encode_reports
from tests.workbench.test_execution_ledger_support import all_rows
from tests.workbench.test_field_workspace_support import BASE, _ledger_fixture, success
from tests.workbench.test_field_workspace_support import field_api as _field_api


def rows(quantity=3):
    return [{'report_no': 'BG-EXCEL-1', 'batch_id': 'B1', 'operation_label': '1 Turning',
        'completed_quantity': quantity, 'actual_start': '2026-09-01T08:00:00', 'actual_end': '2026-09-01T10:00:00',
        'effective_processing_hours': 1.5, 'machine_label': 'Lathe', 'operator_label': 'Operator', 'remark': '现场核对'}]


def test_real_preview_original_bytes_replay_and_no_duplicate(field_api):
    api = field_api
    content = encode_reports(rows())
    before = all_rows(api.case.conn)
    preview = success(api.upload(content))
    assert preview['data']['can_confirm'] and preview['data']['summary']['changed'] == 1
    assert all_rows(api.case.conn) == before
    body = api.confirm_body(preview)
    result = success(api.client.post(BASE + '/files/confirm', json=body))
    assert result['result'] == 'committed' and api.task()['execution']['known_completed_quantity'] == 3
    repeated = success(api.client.post(BASE + '/files/confirm', json={**body, 'write_token': 'expired'}))
    assert repeated['replayed'] and repeated['receipt_ref'] == result['receipt_ref']
    duplicate = success(api.upload(content))
    assert duplicate['data']['summary']['unchanged'] == 1
    success(api.client.post(BASE + '/files/confirm', json=api.confirm_body(duplicate)))
    assert api.task()['execution']['known_completed_quantity'] == 3


def test_file_known_conflict_atomic_and_unknown_supplement(field_api):
    api = field_api
    source = rows(); source[0].pop('completed_quantity'); source[0].pop('effective_processing_hours')
    preview = success(api.upload(encode_reports(source)))
    success(api.client.post(BASE + '/files/confirm', json=api.confirm_body(preview)))
    full = success(api.upload(encode_reports(rows())))
    assert full['data']['can_confirm']
    success(api.client.post(BASE + '/files/confirm', json=api.confirm_body(full)))
    before = all_rows(api.case.conn)
    conflict = success(api.upload(encode_reports(rows(4))))
    assert not conflict['data']['can_confirm'] and conflict['data']['summary']['rejected'] == 1
    assert all_rows(api.case.conn) == before


def test_export_whole_scope_and_stale_bytes_confirm(field_api):
    api = field_api
    ids = [api.case.op_id] + [api.case.op('X' + str(index), seq=index) for index in range(2, 24)]
    api.case.plan(2, ids)
    for op in ids:
        api.case.command('create', api.case.task(2, op), api.case.values(1))
    first = api.read(size=10)
    response = api.client.get(BASE + '/files/export', query_string={**first['data']['scope'], 'snapshot_ref': first['meta']['snapshot_ref'], 'page': 1, 'size': 10})
    assert response.status_code == 200
    book = load_workbook(BytesIO(response.data)); assert book.active.max_row == 24; book.close()
    preview = success(api.upload(encode_reports(rows()), first))
    api.case.conn.execute("UPDATE Batches SET quantity=11 WHERE batch_id='B1'"); api.case.conn.commit()
    result = api.client.post(BASE + '/files/confirm', json=api.confirm_body(preview))
    assert result.status_code == 409


def test_file_batch_overreport_and_missing_snapshot(field_api):
    api = field_api
    values = rows(6); second = dict(values[0], report_no='BG-EXCEL-2', actual_start='2026-09-01T11:00:00', actual_end='2026-09-01T13:00:00'); values.append(second)
    before = all_rows(api.case.conn)
    preview = success(api.upload(encode_reports(values)))
    assert not preview['data']['can_confirm'] and preview['data']['rows'][1]['errors']
    assert all_rows(api.case.conn) == before
    assert api.client.get(BASE + '/files/export').status_code == 400


def test_import_replay_after_context_expiry_and_conflicting_preview(field_api):
    api = field_api
    preview = success(api.upload(encode_reports(rows())))
    body = api.confirm_body(preview)
    first = success(api.client.post(BASE + '/files/confirm', json=body))
    from web.routes.workbench.resource_action_context import EXTENSION
    api.app.extensions[EXTENSION].clear()
    replay = success(api.client.post(BASE + '/files/confirm', json={**body, 'write_token': 'expired'}))
    assert replay['receipt_ref'] == first['receipt_ref'] and replay['replayed']
    other = api.client.post(BASE + '/files/confirm', json={**body, 'input': {'preview_ref': 'x' * 32}})
    assert other.status_code == 409 and other.get_json()['error']['code'] == 'request_key_conflict'


def test_unknown_no_number_repeat_and_problem_download(field_api):
    api = field_api
    value = rows(); value[0]['report_no'] = ''
    content = encode_reports(value)
    preview = success(api.upload(content))
    success(api.client.post(BASE + '/files/confirm', json=api.confirm_body(preview)))
    duplicate = success(api.upload(content))
    assert duplicate['data']['summary']['unchanged'] == 1
    values = rows(); values[0]['batch_id'] = 'not-found'
    rejected = success(api.upload(encode_reports(values)))
    response = api.client.get(BASE + '/files/errors', query_string={'preview_ref': rejected['data']['preview_ref']})
    assert response.status_code == 200
    book = load_workbook(BytesIO(response.data)); assert book.active['A2'].value == 2; book.close()
