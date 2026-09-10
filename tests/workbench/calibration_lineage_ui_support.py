"""CE real template writes and execution ledger, served only from a pytest database."""

import csv
import hashlib
import io
import json
import sqlite3
import threading
from contextlib import contextmanager
from datetime import datetime, timedelta
from pathlib import Path

import openpyxl
import pytest
from flask import Blueprint, Flask, g, jsonify, request
from werkzeug.serving import make_server

from core.infrastructure.transaction import TransactionManager
from core.services.workbench.calibration_export import COLUMNS
from tests.workbench.template_lineage_support import completed, create, origin
from web.routes.workbench.calibration import register_calibration_routes


@pytest.fixture(name="calibration_case")
def calibration_case(lineage_case):
    return lineage_case


def template_completed(case, template_id, hours, *, prefix, version):
    ids = []
    for index in range(len(hours)):
        code = prefix + "-" + str(index).zfill(3)
        case.batch_service.create(code, "P1", 10)
        with TransactionManager(case.conn).transaction():
            ids.append(case.lineage_writer.copy_template(code, template_id))
    case.plan(version, ids)
    for index, (op_id, value) in enumerate(zip(ids, hours)):
        end = datetime(2026, 9, 9, 11) + timedelta(minutes=index)
        case.command("create", case.task(version, op_id), case.values(10,
            effective_processing_hours=value * 10, actual_start=(end - timedelta(hours=value * 10 + 1)).isoformat(), actual_end=end.isoformat()))
    return ids


def prepare(case):
    case.command("create", case.task(1, case.op_id), case.values(10, effective_processing_hours=0))
    case.conn.execute("UPDATE PartOperations SET op_type_name='Turning-01'")
    case.conn.commit()
    old_ids, _ = completed(case, [8, 8], prefix="OLD", version=2)
    case.conn.execute("UPDATE PartOperations SET unit_hours=2")
    case.conn.commit()
    ids, reports = completed(case, [0, 1, 3, 4, 90, 6, 7, 8, 9], prefix="LIVE", version=3)
    case.command("correct", reports[4]["report_ref"], {"original_revision_ref": reports[4]["revision_ref"],
        "effective_processing_hours": 1000, "actual_start": (datetime(2026, 9, 9, 10, 4) - timedelta(hours=1001)).isoformat(),
        "reason": "核对原始报工单，更正有效加工小时"})
    with TransactionManager(case.conn).transaction():
        case.conn.execute("UPDATE BatchOperations SET unit_hours=9 WHERE id=?", (ids[5],))
        case.lineage_writer.withdraw(origin(case, ids[6])["operation_ref"], "核对来源后撤回")
    for index, patch in ((7, {"completed_quantity": 5}), (8, {"effective_processing_hours": None})):
        case.command("correct", reports[index]["report_ref"], {"original_revision_ref": reports[index]["revision_ref"],
            "reason": "复核现场记录", **patch})
    unfinished = create(case, "LIVE-UNFINISHED")
    case.batch_service.create("UNBOUND-UNKNOWN", "P1", 10)
    with TransactionManager(case.conn).transaction():
        unknown = case.lineage_writer.copy_instance("UNBOUND-UNKNOWN", case.op_id)
    case.plan(4, [unknown])
    case.command("create", case.task(4, unknown), case.values(None, effective_processing_hours=None))
    for seq in range(2, 24):
        case.conn.execute("""INSERT INTO PartOperations(part_no,seq,op_type_id,op_type_name,source,unit_hours)
            VALUES ('P1',?,'T1',?,?,?)""", (seq, "Turning-" + str(seq).zfill(2),
                "external" if seq == 23 else "internal", 0 if seq == 2 else None if seq == 3 else 1))
    case.conn.commit()
    templates = {row["seq"]: row["id"] for row in case.conn.execute("SELECT seq,id FROM PartOperations")}
    other_ids = template_completed(case, templates[2], [20] * 5, prefix="OTHER", version=5)
    recent_ids = template_completed(case, templates[4], list(range(1, 26)), prefix="RECENT", version=6)
    return {"template_ref": origin(case, ids[0])["template_operation_ref"],
            "selected_refs": [origin(case, value)["operation_ref"] for value in ids[:5]],
            "excluded_refs": [origin(case, value)["operation_ref"] for value in old_ids + ids[5:] + [unfinished]],
            "other_refs": [origin(case, value)["operation_ref"] for value in other_ids],
            "recent_refs": [origin(case, value)["operation_ref"] for value in recent_ids],
            "report_ref": reports[4]["report_ref"], "correction_sample_ref": origin(case, ids[4])["operation_ref"],
            "unknown_ref": case.lineage_repo.instance(unknown)["operation_ref"], "rows": 23}


def export_rows(content, format_name, scope, snapshot, as_of):
    if format_name == "csv":
        assert content.startswith(b"\xef\xbb\xbf")
        rows = list(csv.reader(io.StringIO(content.decode("utf-8-sig"))))
    else:
        book = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=False)
        try:
            assert book.sheetnames == ["校准建议", "范围与口径"]
            assert all(cell.data_type != "f" for sheet in book for row in sheet for cell in row)
            rows = list(book["校准建议"].iter_rows(values_only=True))
            metadata = dict(book["范围与口径"].iter_rows(values_only=True))
            assert metadata["范围快照"].lstrip("'") == snapshot
            assert metadata["数据截至"] == as_of and json.loads(metadata["筛选范围"]) == scope
        finally:
            book.close()
    return rows


def verify_cell(value, actual, key):
    value = value[1:] if isinstance(value, str) and value.startswith("'") else value
    if isinstance(actual, (dict, list)):
        assert json.loads(value) == actual, key
    elif actual is None:
        assert value == "未知", key
    elif isinstance(actual, (int, float)):
        assert float(value) == actual, key
    else:
        assert value == actual, key


def verify_export(content, format_name, expected, scope, snapshot, as_of):
    rows = export_rows(content, format_name, scope, snapshot, as_of)
    assert list(rows[0]) == [label for _, label in COLUMNS] + ["筛选范围"]
    assert len(rows) == len(expected) + 1
    for values, item in zip(rows[1:], expected):
        for value, (key, _) in zip(values, COLUMNS):
            verify_cell(value, item[key], key)
        assert json.loads(values[-1]) == scope and item["snapshot_ref"] == snapshot and item["as_of"] == as_of


@contextmanager
def serve(case, output, config):
    database = Path(case.conn.execute("PRAGMA database_list").fetchone()[2]).resolve()
    assert database.name == "template-lineage.sqlite"
    app = Flask(__name__)
    app.config.update(TESTING=True, SECRET_KEY="calibration-lineage-isolated-only")
    bp = Blueprint("calibration_lineage_ui", __name__)
    register_calibration_routes(bp)
    app.register_blueprint(bp)
    journal, violations, mutations = [], [], []
    digest = lambda: hashlib.sha256(database.read_bytes()).hexdigest()

    @app.before_request
    def open_database():
        if not request.path.startswith("/api/workbench/v1/calibration"):
            return
        g.before_hash, g.statements = digest(), []
        g.db = sqlite3.connect(database.as_uri() + "?mode=ro", uri=True)
        g.db.row_factory = sqlite3.Row
        g.db.execute("PRAGMA query_only=ON")
        g.db.set_trace_callback(g.statements.append)

    @app.after_request
    def verify_read(response):
        if hasattr(g, "db"):
            writes = [sql for sql in g.statements if sql.lstrip().split()[0].upper() in
                      {"INSERT", "UPDATE", "DELETE", "REPLACE", "CREATE", "ALTER", "DROP"}]
            row = {"path": request.path, "method": request.method, "status": response.status_code,
                   "query": dict(request.args), "writes": writes, "unchanged": digest() == g.before_hash}
            journal.append(row)
            if writes or not row["unchanged"]:
                violations.append(row)
        return response

    @app.teardown_request
    def close_database(_error):
        if hasattr(g, "db"):
            g.db.close()

    @app.post("/__calibration_lineage_fixture__/recreate")
    def recreate_template():
        assert not mutations
        conn = sqlite3.connect(str(database))
        try:
            conn.execute("DELETE FROM PartOperations WHERE seq=1")
            conn.execute("""INSERT INTO PartOperations(part_no,seq,op_type_id,op_type_name,source,unit_hours)
                VALUES ('P1',1,'T1','Turning-01','internal',2)""")
            conn.commit()
            mutations.append("delete-recreate-template-1")
        finally:
            conn.close()
        return jsonify(ok=True)

    server = make_server("127.0.0.1", 0, app, threaded=False)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        yield {**config, "api_origin": "http://127.0.0.1:" + str(server.server_port)}
    finally:
        server.shutdown()
        thread.join(timeout=10)
        server.server_close()
        evidence = {"database": str(database), "journal": journal, "violations": violations,
                    "mutations": mutations, "server_stopped": not thread.is_alive()}
        (output / "calibration-lineage-server.json").write_text(json.dumps(evidence, ensure_ascii=False, indent=2), encoding="utf-8")
        assert not violations and not thread.is_alive()
