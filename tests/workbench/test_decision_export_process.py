"""Real spawned XLSX writers preserve native files and propagate terminal failures."""

import csv
import io
import os
from concurrent.futures.process import BrokenProcessPool
from copy import deepcopy

import pytest
from openpyxl import load_workbook

from core.models.workbench_action_values import FileDownload
from core.models.workbench_command import WorkbenchCommandRejected
from core.services.workbench.decision_file_export import write_decision_xlsx
from core.services.workbench.history.export import write_history_export
from core.services.workbench.history.queries import read_history_all
from core.services.workbench.run.decision_export import read_decision_export, write_run_decision_export
from tests.workbench.decision_history_support import exported_history_references, four_source_case


@pytest.fixture(scope="module")
def materialized_exports(tmp_path_factory):
    """Use the existing small four-source archive; only materialized values leave it."""
    with four_source_case(tmp_path_factory.mktemp("decision-export-process")) as case:
        refs = case.history_refs
        run, _scope, _state = read_decision_export(case.conn, refs["run_ref"], "overview",
            selected_candidate_ref=refs["candidate_ref"])
        history, _scope, _state = read_history_all(case.conn, {
            "timeline": "official", "selected_entry_ref": "adoption:" + refs["trial_receipt_ref"],
            "reference_run_ref": refs["run_ref"],
            "metric_keys": ["expected_operation_count", "overdue_count", "makespan_hours", "machine_utilization_ratio"]})
    snapshot = {"snapshot_ref": "s" * 32, "as_of": "2026-10-17T12:00:00"}
    return {"run": (run, snapshot, "overview"), "history": (history, snapshot, "all")}


def workbook_values_and_types(content):
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    try:
        return [(sheet.title, [tuple((cell.value, type(cell.value), cell.data_type) for cell in row)
                               for row in sheet.iter_rows()]) for sheet in workbook.worksheets]
    finally:
        workbook.close()


@pytest.mark.parametrize("kind", ["run", "history"])
def test_spawned_xlsx_matches_direct_writer_values_types_and_metadata(materialized_exports, kind):
    source, snapshot, section = materialized_exports[kind]
    data = deepcopy(source)
    label = "=原方案001"
    if kind == "run":
        data["candidates"][0]["label"] = label
        direct = write_run_decision_export(data, snapshot, "xlsx", section)
    else:
        data["entries"][0]["label"] = label
        direct = write_history_export(data, snapshot, "xlsx", section)
    spawned = write_decision_xlsx(kind, data, snapshot, section)
    assert isinstance(spawned, FileDownload)
    assert (spawned.filename, spawned.mime_type, spawned.row_count) == (direct.filename, direct.mime_type, direct.row_count)
    expected, actual = workbook_values_and_types(direct.content), workbook_values_and_types(spawned.content)
    assert actual == expected
    assert any(value == label and value_type is str and cell_type == "s"
               for _title, rows in actual for row in rows for value, value_type, cell_type in row)


def test_spawned_xlsx_restores_writer_rejection_code_message_and_status(materialized_exports):
    source, snapshot, section = materialized_exports["history"]
    data = deepcopy(source)
    data["entries"][0]["label"] = "超" * 32768
    with pytest.raises(WorkbenchCommandRejected) as direct:
        write_history_export(data, snapshot, "xlsx", section)
    with pytest.raises(WorkbenchCommandRejected) as spawned:
        write_decision_xlsx("history", data, snapshot, section)
    expected = direct.value
    assert expected.code == "export_format_limit" and expected.status == 422
    assert (spawned.value.code, str(spawned.value), spawned.value.status, spawned.value.operation_refs) == (
        expected.code, str(expected), expected.status, expected.operation_refs)
    assert spawned.value.committed is expected.committed is False


@pytest.mark.parametrize("operation_count", [643, 883, 20000])
def test_history_files_keep_large_scope_and_metric_reference_collections(materialized_exports, operation_count):
    source, snapshot, _section = materialized_exports["history"]
    data = deepcopy(source)
    entry = next(item for item in data["entries"] if item["scope"]["operation_refs"] is not None)
    data["entries"] = [entry]
    keys = data["filters"]["metric_keys"] = ["overdue_count", "makespan_hours", "machine_utilization_ratio"]
    operations = [f"{index:048x}" for index in range(1, operation_count + 1)]
    batch_ref = entry["scope"]["batch_refs"][0]
    entry["scope"].update(operation_refs=operations, operation_count=operation_count)
    entry["scope"]["batch_quantity"]["covered_batch_refs"] = [batch_ref]
    cohort = entry["reading_context"]["resource_window"]["cohort"]
    resources = [f"{index:048x}" for index in (900002, 900001)]
    cohort.update(machine_refs=resources, operator_refs=[])
    for key in keys:
        entry["metrics"][key].update(covered_batch_refs=None, covered_operation_refs=None, covered_resource_refs=None)
    entry["metrics"]["overdue_count"]["covered_batch_refs"] = [batch_ref]
    entry["metrics"]["makespan_hours"]["covered_operation_refs"] = operations
    entry["metrics"]["machine_utilization_ratio"]["covered_resource_refs"] = resources
    expected = {(entry["entry_ref"], None, path): value for path, value in (
        ("scope.batch_quantity.covered_batch_refs", [batch_ref]), ("scope.batch_refs", entry["scope"]["batch_refs"]),
        ("scope.operation_refs", operations), ("reading_context.resource_window.cohort.machine_refs", resources),
        ("reading_context.resource_window.cohort.operator_refs", []))}
    expected.update({(entry["entry_ref"], key, field): entry["metrics"][key][field] for key in keys
                     for field in ("covered_batch_refs", "covered_operation_refs", "covered_resource_refs")})
    csv_file = write_history_export(data, snapshot, "csv", "entries")
    csv_rows = list(csv.reader(io.StringIO(csv_file.content.decode("utf-8-sig"))))
    headers, rows = csv_rows[0], csv_rows[1:]
    assert exported_history_references(headers, rows) == expected
    assert csv_file.row_count == len(rows)
    assert all(row[headers.index("导出格式版本")] == "2" for row in rows)
    assert all(len(value) <= 32767 for row in rows for value in row)
    xlsx_file = write_history_export(data, snapshot, "xlsx", "entries")
    workbook = load_workbook(io.BytesIO(xlsx_file.content), read_only=True, data_only=False)
    try:
        native = list(workbook["历史条目"].iter_rows(max_col=len(headers), values_only=True))
        assert tuple(headers) == native[0]
        assert exported_history_references(headers, native[1:]) == expected
        assert xlsx_file.row_count == csv_file.row_count == len(native) - 1
        assert all(row[headers.index("导出格式版本")] == 2 for row in native[1:])
        assert all(len(value) <= 32767 for row in native[1:] for value in row if isinstance(value, str))
    finally:
        workbook.close()


class ExitDuringWorkerUnpickle:
    def __reduce__(self):
        return os._exit, (7,)


def test_spawned_xlsx_reports_worker_exit_as_broken_process_pool():
    # The parent pickles this normally; only child deserialization exits.
    with pytest.raises(BrokenProcessPool):
        write_decision_xlsx("history", {"exit_in_worker": ExitDuringWorkerUnpickle()},
                            {"snapshot_ref": "s" * 32, "as_of": "2026-10-17T12:00:00"}, "all")
