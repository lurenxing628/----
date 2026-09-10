"""R1-B isolated real HTTP fixtures and edits to downloaded XLSX bytes."""

from io import BytesIO

import pytest
from flask import Blueprint
from openpyxl import load_workbook

from tests.workbench.execution_ledger_support import ledger_case as ledger_case
from tests.workbench.field_workspace_support import BASE, success
from tests.workbench.field_workspace_support import field_api as field_api
from web.routes.workbench.actual_gantt import register_actual_gantt_routes

ACTUAL = '/api/workbench/v1/actual-gantt'
PIECES = ('分件甲', '分件乙', '0')


def seed_pieces(api):
    case = api.case
    ids = [case.op_id] + [case.op('R1B-' + piece, piece=piece) for piece in PIECES]
    case.plan(2, ids)
    case.conn.execute("UPDATE Machines SET name='中文车床' WHERE machine_id='M1'")
    case.conn.execute("UPDATE Operators SET name='中文人员' WHERE operator_id='O1'")
    case.conn.commit()
    bp = Blueprint('round1_actual_files', __name__)
    register_actual_gantt_routes(bp)
    api.app.register_blueprint(bp)
    return api


@pytest.fixture(name='piece_file_api')
def piece_file_api(field_api):
    return seed_pieces(field_api)


def download(api, kind='template', reading=None):
    reading = reading or api.read()
    response = api.client.get(BASE + '/files/' + kind, query_string={
        **reading['data']['scope'], 'snapshot_ref': reading['meta']['snapshot_ref']})
    assert response.status_code == 200, response.get_data(as_text=True)
    return response.data


def fill_template(content, changes):
    book = load_workbook(BytesIO(content))
    sheet = book.active
    assert sheet is not None
    headers = {cell.value: index for index, cell in enumerate(sheet[1], 1)}
    for row in range(2, sheet.max_row + 1):
        piece = sheet.cell(row, headers['单件编号']).value
        for name, value in changes.get(piece, {}).items():
            sheet.cell(row, headers[name]).value = value
    output = BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


def report_values(quantity=0):
    return {'本次完成数量': quantity, '实际开工': '2026-09-01T08:00:00',
            '本次实际完工': '2026-09-01T09:00:00', '有效加工工时(h)': 0,
            '实际设备': '中文车床', '实际人员': '中文人员', '备注': '中文现场报工'}


def confirm(api, content):
    preview = success(api.upload(content))
    assert preview['data']['can_confirm'], preview
    body = api.confirm_body(preview)
    return preview, body, success(api.client.post(BASE + '/files/confirm', json=body))


def by_piece(api):
    return {task['piece_id']: task for task in api.read()['data']['tasks']}
