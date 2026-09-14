"""Full server scope and lossless CSV/XLSX round trips, with no plan identities."""

import csv
import json
from io import BytesIO, StringIO

import openpyxl
import pytest

from tests.workbench.run_candidate_support import BASE, api, compute, edit_capture, read, retained
from tests.workbench.run_candidate_support import candidate_case as _candidate_case


def decode(response, fmt):
    if fmt == "csv":
        rows = list(csv.reader(StringIO(response.data.decode("utf-8-sig"))))
        return rows[0], [[value[1:] if value.startswith("'") else value for value in row] for row in rows[1:]]
    wb = openpyxl.load_workbook(BytesIO(response.data), read_only=True, data_only=False)
    try:
        rows = list(wb.active.iter_rows(values_only=True))
        return list(rows[0]), [list(row) for row in rows[1:]]
    finally:
        wb.close()


@pytest.mark.parametrize("fmt", ["csv", "xlsx"])
def test_fullscope_export_matches_workspace_and_retains_partial_rows(candidate_case, fmt):
    case = candidate_case
    case.batch("B2", ready_status="no")
    case.operation("B2")
    case.operation(seq=2)
    case.conn.commit()
    _, refs = compute(case, case.settings("B1", "B2"))
    client, _ = api(case)
    path = "/candidates/" + refs[0]
    with retained(case.conn):
        result = read(client, path, sort="start", order="desc")
        data, token = result["data"], result["meta"]["snapshot_ref"]
        response = client.get(BASE + path + "/export", query_string={"format": fmt, "snapshot_ref": token, "sort": "start", "order": "desc"})
        assert response.status_code == 200, response.get_data(as_text=True) if response.status_code != 200 else ""
        headers, rows = decode(response, fmt)
        assert response.headers["X-Workbench-Task-Count"] == "2"
        assert response.headers["X-Workbench-Row-Count"] == "3"
        assert response.headers["X-Workbench-Snapshot-Ref"] == token
        assert [row[headers.index("行编号")] for row in rows[:2]] == [row["row_ref"] for row in data["tasks"]]
        assert [row[headers.index("安排开始")] for row in rows[:2]] == [row["start"] for row in data["tasks"]]
        assert rows[-1][0] == "未安排工序" and rows[-1][headers.index("行编号")] == r"\N"
        assert all(row[headers.index("候选方案状态")] == "部分排完" for row in rows)
        assert not any(value in headers for value in ("version", "plan_ref", "task_ref"))


@pytest.mark.parametrize("query", [{"format": "csv"}, {"format": "pdf", "snapshot_ref": "x"},
                                   {"format": "csv", "snapshot_ref": "x", "page": 1}])
def test_export_requires_format_scope_snapshot_and_no_client_paging(candidate_case, query):
    client, _ = api(candidate_case)
    response = client.get(BASE + "/candidates/" + "a" * 48 + "/export", query_string=query)
    assert response.status_code == 400


def test_export_cannot_reuse_catalog_or_another_range_snapshot(candidate_case):
    case = candidate_case
    run_ref, refs = compute(case)
    client, _ = api(case)
    tokens = [read(client, "/runs/" + run_ref + "/candidates")["meta"]["snapshot_ref"],
              read(client, "/candidates/" + refs[1])["meta"]["snapshot_ref"],
              read(client, "/candidates/" + refs[0], sort="end")["meta"]["snapshot_ref"]]
    for token in tokens:
        response = client.get(BASE + "/candidates/" + refs[0] + "/export", query_string={"format": "csv", "snapshot_ref": token})
        assert response.status_code == 409 and response.get_json()["error"]["code"] == "snapshot_stale"


@pytest.mark.parametrize("fmt", ["csv", "xlsx"])
def test_export_formula_like_labels_and_control_characters_are_lossless(candidate_case, fmt):
    case = candidate_case
    _, refs = compute(case)
    label = '=SUM(1,2)\r\n"quoted" \\N'

    def change(facts):
        from core.services.workbench.run_candidate_facts import _columns
        sql = next(row[3] for row in facts["schema"] if row[0] == "table" and row[1] == "Machines")
        facts["tables"]["Machines"][0][_columns(sql, "Machines").index("name")] = label

    edit_capture(case, "facts_json", change)
    client, _ = api(case)
    result = read(client, "/candidates/" + refs[0])
    response = client.get(BASE + "/candidates/" + refs[0] + "/export", query_string={"format": fmt, "snapshot_ref": result["meta"]["snapshot_ref"]})
    assert response.status_code == 200
    headers, rows = decode(response, fmt)
    assert rows[0][headers.index("设备名称")] == label
    if fmt == "xlsx":
        wb = openpyxl.load_workbook(BytesIO(response.data), read_only=True, data_only=False)
        try:
            assert all(cell.data_type != "f" for row in wb.active.iter_rows() for cell in row)
        finally:
            wb.close()


def test_http_envelope_is_bounded_without_truncation(candidate_case, monkeypatch):
    from web.routes.workbench import run_candidates
    case = candidate_case
    _, refs = compute(case)
    client, _ = api(case)
    first = read(client, "/candidates/" + refs[0])
    monkeypatch.setattr(run_candidates, "MAX_RESPONSE_BYTES", 100)
    with retained(case.conn):
        response = client.get(BASE + "/candidates/" + refs[0])
        assert response.status_code == 413 and response.get_json()["error"]["code"] == "candidate_capacity_exceeded"
        download = client.get(BASE + "/candidates/" + refs[0] + "/export", query_string={"format": "csv", "snapshot_ref": first["meta"]["snapshot_ref"]})
        assert download.status_code == 413


def test_raw_generation_inputs_cannot_be_nested_private_payloads(candidate_case):
    case = candidate_case
    _, refs = compute(case)
    edit_capture(case, "normalized_input_json", lambda value: value.update(ready_check={"facts_json": "SECRET"}))
    client, _ = api(case)
    response = client.get(BASE + "/candidates/" + refs[0])
    assert "SECRET" not in json.dumps(response.get_json())
    assert response.status_code in (200, 500)
