"""Version headers are exact; physical XML bounds cannot be normalized away."""

from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import load_workbook

from core.models.workbench_command import WorkbenchCommandRejected
from core.services.workbench.field_report_files_codec import HEADERS, TASK_HEADERS, decode_reports, encode_reports


def changed_book(content, mutate):
    book = load_workbook(BytesIO(content))
    sheet = book.active
    assert sheet is not None
    mutate(sheet)
    output = BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


def test_v2_header_keeps_ten_column_prefix_null_and_string_zero():
    content = encode_reports([{'task_ref': 'a' * 48, 'piece_id': '0', 'operation_scope': '单件',
                              'completed_quantity': 0, 'effective_processing_hours': None}], format_version=2)
    book = load_workbook(BytesIO(content))
    sheet = book.active
    assert sheet is not None
    assert tuple(cell.value for cell in sheet[1]) == TASK_HEADERS
    assert TASK_HEADERS[:10] == HEADERS and sheet.max_column == 13
    book.close()
    row = decode_reports(content)[0]
    assert row['format_version'] == 2 and row['errors'] == []
    assert row['values']['piece_id'] == '0' and row['values']['completed_quantity'] == 0
    assert row['values']['effective_processing_hours'] is None


@pytest.mark.parametrize('column,value', [('K1', '任务标识(v2)'), ('L1', '任务编号'), ('N1', '额外字段')])
def test_unrecognized_or_duplicate_headers_are_not_guessed(column, value):
    content = encode_reports([], format_version=2)
    modified = changed_book(content, lambda sheet: sheet.__setitem__(column, value))
    with pytest.raises(WorkbenchCommandRejected):
        decode_reports(modified)


@pytest.mark.parametrize('version,column', [(1, 'K2'), (2, 'N2')])
def test_extra_physical_columns_rejected_even_if_dimension_lies(version, column):
    content = changed_book(encode_reports([{}], format_version=version),
                           lambda sheet: sheet.__setitem__(column, '不许忽略'))
    output = BytesIO()
    with ZipFile(BytesIO(content)) as source, ZipFile(output, 'w', ZIP_DEFLATED) as target:
        for item in source.infolist():
            value = source.read(item.filename)
            if item.filename == 'xl/worksheets/sheet1.xml':
                value = value.replace(('ref="A1:' + column + '"').encode(), b'ref="A1:A1"')
            target.writestr(item, value)
    with pytest.raises(WorkbenchCommandRejected):
        decode_reports(output.getvalue())


def test_reordered_exact_headers_are_still_verifiable():
    content = encode_reports([{'task_ref': 'b' * 48, 'piece_id': '中文件', 'operation_scope': '单件'}], format_version=2)
    def reorder(sheet):
        for row in (1, 2):
            first, last = sheet.cell(row, 1).value, sheet.cell(row, 13).value
            sheet.cell(row, 1).value, sheet.cell(row, 13).value = last, first
    row = decode_reports(changed_book(content, reorder))[0]
    assert row['errors'] == [] and row['values']['piece_id'] == '中文件'
