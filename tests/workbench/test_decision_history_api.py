"""Archived result identity, semantic comparisons and native complete history files."""

import csv
import io
import json
from collections import Counter
from copy import deepcopy
from datetime import datetime, timedelta
from types import SimpleNamespace
from typing import Any, Dict

import pytest
from openpyxl import load_workbook

from tests.workbench.decision_history_support import (
    BASE,
    api,
    clone_run,
    compare_conditions,
    copy_case,
    exported_history_references,
    four_source_case,
    read,
)
from tests.workbench.run_candidate_support import candidate_case as _candidate_case  # noqa: F401
from tests.workbench.run_candidate_support import compute_two_objectives, retained


@pytest.fixture(scope="module")
def history_template(tmp_path_factory):
    with four_source_case(tmp_path_factory.mktemp("decision-history-template")) as case:
        yield case


@pytest.fixture
def history_case(history_template, tmp_path):
    case: Any = copy_case(history_template, tmp_path)
    with case.app.app_context():
        yield case
    case.conn.close()


def test_history_event_keeps_all_metrics_but_enriches_only_its_candidate(candidate_case, monkeypatch):
    from core.infrastructure.transaction import TransactionManager
    from core.services.workbench.facts.candidate_store import CandidateStore
    from core.services.workbench.history.reading import HistoryReader
    from core.services.workbench.run.decision import decision_overview
    from core.services.workbench.run.decision_resource_engine import DecisionResources
    from core.services.workbench.run.decision_source import DecisionSource

    case = candidate_case
    run_ref, candidates = compute_two_objectives(case)
    candidate_ref = candidates[-1]
    start, end = "2026-09-09T00:00:00", "2026-09-10T00:00:00"
    with TransactionManager(case.conn).transaction():
        source = DecisionSource(CandidateStore(case.conn), run_ref, selected_candidate_ref=candidate_ref,
            range_start=start, range_end=end, resource_view="full_candidate")
        expected = decision_overview(source)
        selected = next(row for row in expected["candidates"] if row["candidate_ref"] == candidate_ref)
        calls, original = [], DecisionResources.metrics

        def metrics(engine, key):
            calls.append(key)
            return original(engine, key)

        monkeypatch.setattr(DecisionResources, "metrics", metrics)
        actual = HistoryReader(case.conn).decision(run_ref, start, end, candidate_ref=candidate_ref)
        assert actual["resource_cohort"] == expected["resource_cohort"]
        assert actual["candidates"][0]["metrics"] == selected["metrics"]
        assert actual["candidates"][0]["coverage"] == selected["coverage"]
        assert calls == [candidate_ref]


def test_three_saved_sources_preserve_events_and_result_identity_on_two_timelines(history_case):
    from core.infrastructure.transaction import TransactionManager
    from core.models.workbench_run_job import new_run_ref
    from core.services.workbench.facts.decision_read_model import build_failure_read_models
    from core.services.workbench.facts.decision_read_models import decode_model
    from core.services.workbench.history.catalog import read_history_catalog
    from data.repositories.workbench_run_repo import WorkbenchRunRepository

    case, refs = history_case, history_case.history_refs
    queued = case.accept(key="history-queued-run-0001")["run_ref"]
    failed = case.accept(key="history-failed-run-0001")["run_ref"]
    repo = WorkbenchRunRepository(case.conn)
    repo.claim(failed, new_run_ref(), "2026-09-10T12:00:01")
    failure = {"state": "failed", "result_persisted": False, "candidates": [],
        "error": {"code": "fixture_saved_failure", "message": "Saved worker failure"}}
    repo.finish(failed, "failed", failure, "2026-09-10T12:00:02")
    admission = decode_model(repo.admission_read_model(failed, ("decision",)), ("decision",))["decision"]
    models = build_failure_read_models(repo.head(failed), admission, failure)
    repo.insert_decision_read_models(failed, models["run_model"], models["candidate_models"])
    case.conn.commit()
    client, statements = api(case)
    with retained(case.conn):
        runs = read(client, timeline="run", state="complete")["data"]
        official = read(client, timeline="official", order="asc")["data"]
        assert runs["totals"] == {"run": 3, "official": 3, "filtered": 1}
        assert official["totals"] == {"run": 3, "official": 3, "filtered": 3}
        assert len(runs["entries"]) == 1 and runs["entries"][0]["identity"]["run_ref"] == refs["run_ref"]
        assert runs["entries"][0]["identity"]["candidate_ref"] == refs["candidate_ref"]
        by_ref = {row["entry_ref"]: row for row in official["entries"]}
        for prefix in ("candidate", "second_candidate"):
            event = by_ref["adoption:" + refs[prefix + "_receipt_ref"]]
            assert event["source_kind"] == "candidate_adoption"
            assert event["identity"]["run_ref"] == refs["run_ref"]
            assert event["identity"]["candidate_ref"] == refs["candidate_ref"]
            assert event["identity"]["plan_ref"] == refs[prefix + "_plan_ref"]
            assert event["identity"]["plan_version"] == refs[prefix + "_version"]
        trial = by_ref["adoption:" + refs["trial_receipt_ref"]]
        assert trial["source_kind"] == "trial_adoption"
        assert trial["identity"]["scenario_ref"] == refs["scenario_ref"]
        assert trial["identity"]["draft_ref"] == refs["draft_ref"]
        assert trial["identity"]["run_ref"] is None and trial["identity"]["candidate_ref"] is None
        assert trial["identity"]["plan_ref"] == refs["trial_plan_ref"]
        trial_gantt = next(row for row in trial["navigation"] if row["view"] == "gantt")
        assert trial_gantt["enabled"] and trial_gantt["context"] == {"plan_ref": refs["trial_plan_ref"]}
        unfinished = read(client, timeline="run", state="queued", selected_entry_ref="run:" + queued)["data"]["selection"]["entry"]
        assert unfinished["identity"]["run_ref"] == queued and unfinished["source"]["state"] == "unavailable"
        run_link = next(row for row in unfinished["navigation"] if row["view"] == "run")
        assert run_link["enabled"] and run_link["context"] == {"run_ref": queued}
        for state, ref in (("queued", queued), ("failed", failed)):
            entry = read(client, timeline="run", state=state, selected_entry_ref="run:" + ref,
                metric_keys=["expected_operation_count", "scheduled_operation_count", "machine_available_hours"])["data"]["selection"]["entry"]
            expected = entry["metrics"]["expected_operation_count"]
            assert expected["value"] == entry["scope"]["operation_count"] == 1
            assert expected["state"] == "complete" and expected["unknown_count"] == 0
            assert expected["covered_operation_refs"] == entry["scope"]["operation_refs"]
            assert entry["metrics"]["scheduled_operation_count"]["value"] is None
            assert entry["metrics"]["machine_available_hours"]["value"] is None
        assert set(by_ref) == {"adoption:" + refs[key] for key in ("candidate_receipt_ref", "second_candidate_receipt_ref", "trial_receipt_ref")}
        for row in runs["entries"] + official["entries"]:
            assert set(row["times"]) == {"accepted_at", "started_at", "finished_at", "adopted_at", "committed_at_utc", "schedule_time"}
        assert client.get(BASE, query_string={"timeline": "invalid"}).status_code == 422
    assert not any(sql.lstrip().upper().startswith(("INSERT", "UPDATE", "DELETE", "CREATE", "ALTER", "DROP")) for sql in statements)
    case.conn.execute("DELETE FROM ScheduleHistory WHERE version=?", (refs["second_candidate_version"],))
    case.conn.execute("UPDATE WorkbenchCommandReceipts SET committed_at_utc=? WHERE receipt_ref=?",
                      (b"bad-time", refs["second_candidate_receipt_ref"]))
    case.conn.commit()
    with TransactionManager(case.conn).transaction():
        descriptors, _ = read_history_catalog(case.conn)
    degraded = next(row for row in descriptors if row["entry_ref"] == "adoption:" + refs["second_candidate_receipt_ref"])
    assert degraded["occurred_at"] is None
    assert any(gap["field"] == "committed_at_utc" and gap["code"] == "invalid_stored_time"
               for gap in degraded["evidence_gaps"])
    official_after_blob = read(client, timeline="official")["data"]
    degraded_entry = next(
        row for row in official_after_blob["entries"]
        if row["entry_ref"] == "adoption:" + refs["second_candidate_receipt_ref"]
    )
    assert degraded_entry["times"]["committed_at_utc"] is None
    assert {"code": "invalid_stored_time", "message": "原记录的此项无法确认。"} in degraded_entry["issues"]


def test_saved_production_conditions_and_metric_gates_keep_algorithm_experiments_separate(history_case, monkeypatch):
    from core.infrastructure.read_evidence import read_evidence_scope
    from core.infrastructure.snapshot_connection import ddl_columns
    from core.infrastructure.transaction import TransactionManager
    from core.services.workbench.facts.candidate_facts import GenerationFacts
    from core.services.workbench.facts.run_input_readonly import candidate_read_snapshot
    from core.services.workbench.history import conditions as condition_module
    from core.services.workbench.history.catalog import read_history_catalog
    from core.services.workbench.history.conditions import (
        _compare,
        _materials,
        _Projection,
        compare_prepared_history_conditions,
    )
    from core.services.workbench.history.conditions_source import ConditionSource
    from core.services.workbench.history.metric_gate import gate_history_metric
    from core.services.workbench.history.projection import missing_packet
    from core.services.workbench.history.reading import HistoryReader
    from core.shared.strict_json import canonical_json

    case, refs = history_case, history_case.history_refs
    experiment, _ = clone_run(case, "2026-09-11T12:00:00", change_artifact=lambda value: value.update(time_budget_seconds=30))
    changed, _ = clone_run(case, "2026-09-12T12:00:00", change_input=lambda value: value.update(ready_check=False))
    for day in range(13, 25):
        clone_run(case, "2026-09-" + str(day) + "T12:00:00")
    client, _ = api(case)
    query = {"comparison": "same_conditions", "reference_run_ref": refs["run_ref"],
             "selected_entry_ref": "run:" + experiment, "metric_keys": ["overdue_count", "makespan_hours"]}
    with retained(case.conn):
        reference_only = read(client, comparison="same_conditions", reference_run_ref=refs["run_ref"],
                              metric_keys=query["metric_keys"])["data"]
        assert reference_only["selection"]["entry"] is None
        assert reference_only["selection"]["reference"]["identity"]["run_ref"] == refs["run_ref"]
        assert reference_only["entries"] and reference_only["comparison"]["conditions"]["state"] == "unknown"
        assert all(row["state"] == "unavailable" and row["value"] is None
                   for row in reference_only["comparison"]["metric_deltas"].values())
        counts = {key: Counter() for key in ("source", "production", "experiments")}
        calendar_calls = Counter()
        source_init, production, experiments = ConditionSource.__init__, condition_module._production, condition_module._experiments
        calendar = condition_module._calendar

        def count_source(self, packet):
            counts["source"][(packet["kind"], packet["ref"], packet["entry_ref"])] += 1
            return source_init(self, packet)

        def count_production(source):
            counts["production"][(source.kind, source.ref)] += 1
            return production(source)

        def count_calendar(projection, resources, known, window):
            source = projection.source
            calendar_calls[(source.kind, source.ref, window["start"], window["end"])] += 1
            return calendar(projection, resources, known, window)

        def count_experiments(source, packet):
            counts["experiments"][(source.kind, source.ref)] += 1
            return experiments(source, packet)

        with monkeypatch.context() as checks:
            checks.setattr(ConditionSource, "__init__", count_source)
            checks.setattr(condition_module, "_production", count_production)
            checks.setattr(condition_module, "_calendar", count_calendar)
            checks.setattr(condition_module, "_experiments", count_experiments)
            equal = read(client, **query)["data"]
            assert len(equal["trend"]["points"]) == 15
            assert sum(counts["source"].values()) == sum(counts["production"].values()) == sum(counts["experiments"].values()) == 15
            assert all(count == 1 for values in counts.values() for count in values.values())
            next_request = read(client, **query)["data"]
            assert next_request == equal
            assert all(count == 2 for values in counts.values() for count in values.values())
        assert equal["comparison"]["conditions"]["state"] == "equal"
        assert equal["comparison"]["configurations"]["state"] == "different"
        assert equal["comparison"]["objective"]["state"] == "equal"
        assert equal["selection"]["reference"]["identity"]["run_ref"] == refs["run_ref"]
        delta = equal["comparison"]["metric_deltas"]["overdue_count"]
        assert delta["state"] == "comparable" and delta["value"] == 0 and delta["direction"] == "equal"
        different = read(client, **dict(query, selected_entry_ref="run:" + changed))["data"]
        assert different["comparison"]["conditions"]["state"] == "different"
        assert different["comparison"]["conditions"]["summary"]["different"] > 0
        states = [row["state"] for row in different["comparison"]["conditions"]["rows"]]
        assert states[0] == "different"
        assert states == sorted(states, key={"different": 0, "unknown": 1, "equal": 2}.__getitem__)
        assert different["comparison"]["conclusion"]["state"] == "different_conditions"
        before = equal["selection"]["reference"]["metrics"]["makespan_hours"]
        after = equal["selection"]["entry"]["metrics"]["makespan_hours"]
        definition = next(row for row in equal["metric_definitions"] if row["key"] == "makespan_hours")
        for field, value, state in (("metric_version", "another-saved-version", "different_version"),
                                   ("unit", "another-saved-unit", "different_version"),
                                   ("scope", "another-saved-scope", "different_coverage"),
                                   ("covered_operation_refs", ["e" * 48], "different_coverage")):
            altered = deepcopy(after)
            altered[field] = value
            gated = gate_history_metric(definition, before, altered, production_state="equal", objective_state="equal")
            assert gated["state"] == state and gated["value"] is None
        assert len(before["covered_operation_refs"]) == len(after["covered_operation_refs"]) == 1
        objective = gate_history_metric(definition, before, after, production_state="equal", objective_state="different")
        assert objective["state"] == "different_objective" and objective["value"] is None
        with TransactionManager(case.conn).transaction(), candidate_read_snapshot(case.conn), read_evidence_scope(case.conn):
            descriptors, _ = read_history_catalog(case.conn)
            descriptor = next(row for row in descriptors if row["entry_ref"] == "run:" + refs["run_ref"])
            packet = HistoryReader(case.conn, catalog=descriptors).read(descriptor)
        window = {"start": "2026-09-09T08:00:00", "end": "2026-09-09T16:00:00"}
        for values in counts.values():
            values.clear()
        calendar_calls.clear()
        with monkeypatch.context() as checks, TransactionManager(case.conn).transaction():
            checks.setattr(ConditionSource, "__init__", count_source)
            checks.setattr(condition_module, "_production", count_production)
            checks.setattr(condition_module, "_calendar", count_calendar)
            checks.setattr(condition_module, "_experiments", count_experiments)
            reader = HistoryReader(case.conn)
            source = reader.condition_source(packet)
            assert not counts["production"] and not calendar_calls
            side = reader.condition_side(packet, window)
            alias = dict(packet, entry_ref="adoption:same-source")
            assert reader.condition_source(alias) is source
            assert reader.condition_side(alias, dict(window, cohort={"machine_refs": []})) is side
            saved_rows, saved_families = deepcopy(side["conditions"].rows), deepcopy(side["conditions"].families)
            compare_prepared_history_conditions(side, side)
            assert side["conditions"].rows == saved_rows and side["conditions"].families == saved_families
            other_window = dict(window, end="2026-09-10T16:00:00")
            other_side = reader.condition_side(packet, other_window)
            assert other_side is not side and other_side["configurations"] is side["configurations"]
            assert sum(counts["source"].values()) == sum(counts["experiments"].values()) == sum(counts["production"].values()) == 1
            assert sum(calendar_calls.values()) == 2 and set(calendar_calls.values()) == {1}
            assert side["conditions"].rows == saved_rows and side["conditions"].families == saved_families
            assert any(key[0] == "calendar_downtime" and "2026-09-10" in key[1] for key in other_side["conditions"].rows)
            assert not any(key[0] == "calendar_downtime" and "2026-09-10" in key[1] for key in side["conditions"].rows)
            without_result = dict(packet, candidate=None, raw_tasks=None, entry_ref="run:without-result")
            assert reader.condition_source(without_result) is not source
            issues = []
            for code in ("audit_missing", "audit_unproven"):
                issue = {"code": code, "message": "Saved adoption audit gap"}
                unavailable = dict(packet, entry_ref="adoption:" + code, source_issue=issue)
                isolated = reader.condition_source(unavailable)
                assert isolated is not source and not isolated.available and isolated.issue == issue
                issues.append(isolated)
            assert issues[0] is not issues[1]
            assert reader.condition_source(missing_packet()) is not reader.condition_source(
                missing_packet("selection_required", "请选择实际历史条目。"))
        for source, expected_state in (("internal", "unknown"), ("external", "equal")):
            hours = deepcopy(packet)
            operation = next(iter(hours["facts"].tables["BatchOperations"].values()))
            operation.update(source=source, setup_hours=None, unit_hours=None, ext_days=2)
            compared = compare_conditions(hours, hours, window)
            rows = [row for row in compared["conditions"]["rows"] if row["category"] == "requirements"
                    and row["field"] in ("setup_hours", "unit_hours")]
            assert len(rows) == 2 and all(row["state"] == expected_state for row in rows)
            if source == "external":
                assert all(row["before"]["state"] == "known" and row["before"]["value"] is None for row in rows)
                operation.pop("setup_hours")
                operation.pop("unit_hours")
                missing = compare_conditions(hours, hours, window)
                assert all(row["state"] == "unknown" for row in missing["conditions"]["rows"]
                           if row["category"] == "requirements" and row["field"] in ("setup_hours", "unit_hours"))
        saved_selection = packet["receipt"]["selection"]
        assert saved_selection["selection_policy_version"] == "business_objectives_v1"
        primary = saved_selection["primary_objective"]
        saved_objective = compare_conditions(packet, packet, window)["objective"]["before"]
        assert saved_objective["source"]["path"] == "receipt.selection"
        assert saved_objective["value"] == {"name": primary, "field_order": saved_selection["objective_field_order"][primary]}
        old_objective = deepcopy(packet)
        old_objective["receipt"].pop("selection", None)
        old_objective["candidate"]["artifact"]["summary"] = {}
        objective = compare_conditions(packet, old_objective, window)["objective"]
        assert objective["state"] == "equal" and len(objective["after"]["value"]["field_order"]) == 6
        old_objective["receipt"]["summary"]["algo"].pop("best_score_schema", None)
        assert compare_conditions(packet, old_objective, window)["objective"]["state"] == "unknown"
        metadata = deepcopy(packet)
        facts = json.loads(metadata["capture"]["facts_text"])
        for kind, name, _table, ddl in facts["schema"]:
            if kind != "table" or name not in ("Batches", "BatchOperations", "SchemaVersion", "ScheduleConfig"):
                continue
            columns = ddl_columns(ddl, name, restricted=True)
            for row in facts["tables"][name]:
                for field in ("created_at", "updated_at"):
                    if field in columns:
                        row[columns.index(field)] = "2026-09-12T23:59:59"
                if name == "SchemaVersion":
                    row[columns.index("version")] += 1
                if name == "ScheduleConfig" and row[columns.index("config_key")] == "graph_candidate_weight_count":
                    row[columns.index("config_value")] = "7"
        metadata["capture"]["facts_text"] = canonical_json(facts)
        metadata["facts"] = GenerationFacts(metadata["capture"])
        metadata.update(ref="7" * 48, entry_ref="run:" + "7" * 48)
        compared = compare_conditions(packet, metadata, window)
        assert compared["conditions"]["state"] == "equal"
        assert compared["configurations"]["state"] == "different"
        missing_ledger = deepcopy(packet)
        missing_ledger["receipt"].pop("operation_effort", None)
        assert compare_conditions(packet, missing_ledger, window)["conditions"]["state"] == "equal"
        mapping = {packet["capture"]["input"]["batch_refs"][0]: "e" * 48,
                   packet["raw_tasks"][0]["operation_ref"]: "f" * 48}

        def replace_members(value: Any) -> Any:
            if isinstance(value, dict):
                return {mapping.get(key, key): replace_members(item) for key, item in value.items()}
            if isinstance(value, list):
                return [replace_members(item) for item in value]
            return mapping.get(value, value) if isinstance(value, str) else value

        members: Dict[str, Any] = replace_members({key: value for key, value in deepcopy(packet).items() if key != "facts"})
        members["capture"]["facts_text"] = canonical_json(replace_members(json.loads(packet["capture"]["facts_text"])))
        members["facts"] = GenerationFacts(members["capture"])
        compared = compare_conditions(packet, members, window)
        assert compared["conditions"]["state"] == "different"
        scope = [row for row in compared["conditions"]["rows"] if row["category"] == "scope"]
        for field in ("batch_member", "operation_member"):
            previous = {row["entity"]["key"] for row in scope if row["field"] == field and row["before"]["value"] is True}
            current = {row["entity"]["key"] for row in scope if row["field"] == field and row["after"]["value"] is True}
            assert len(previous) == len(current) == 1 and previous != current
        arrival_tables = {"BatchMaterials": [{"id": 1, "batch_id": "B0", "material_id": "M0",
            "required_qty": 1, "available_qty": 1, "ready_status": "yes"}],
            "BatchMaterialReviews": [], "BatchMaterialStages": [],
            "BatchMaterialArrivals": [{"requirement_id": 1, "arrival_date": "2026-09-10", "quantity": quantity}
                                     for quantity in (0.1, 0.2, 0.3)]}
        arrival_tables["BatchMaterialArrivals"].append({"requirement_id": 99, "arrival_date": None,
                                                       "quantity": "Unconsumed malformed quantity"})
        stub = SimpleNamespace(table=lambda name: arrival_tables.get(name), original_table=lambda name: [],
            entity_refs={("material", "M0"): "a" * 48},
            source=lambda path, origin="environment": {"kind": "run", "ref": refs["run_ref"], "path": path})
        original, reversed_rows = _Projection(stub), _Projection(stub)
        _materials(original, {"c" * 48: {"batch_id": "B0"}}, {})
        arrival_tables["BatchMaterialArrivals"].reverse()
        _materials(reversed_rows, {"c" * 48: {"batch_id": "B0"}}, {})
        arrivals = _compare(original, reversed_rows, (("materials_outsource", "物料与外协"),))
        quantity = next(row for row in arrivals["rows"] if row["field"] == "quantity")
        assert quantity["before"]["value"] == quantity["after"]["value"] == 0.6
        assert arrivals["state"] == "equal"


def test_shared_resource_window_uses_saved_capacity_and_keeps_missing_capacity_unknown(history_case):
    from core.services.workbench.facts.trial_scenario_archive import load_saved_scenario
    from core.services.workbench.history.conditions_source import ConditionSource
    from core.services.workbench.history.metrics import _metric, _path, _trial, metric_definitions
    from core.services.workbench.trial.projection import comparison as trial_comparison
    from tests.workbench.trial_support import change, create
    from tests.workbench.trial_support import service as trial_service

    case, refs = history_case, history_case.history_refs
    period_index = [row[1] for row in case.conn.execute("PRAGMA table_info(WorkbenchCalendarDefaults)")].index("periods_json")

    def half_shift(facts):
        facts["tables"]["WorkbenchCalendarDefaults"][0][period_index] = '[{"start":"08:00","end":"12:00","day_offset":0}]'

    changed, _ = clone_run(case, "2026-09-11T12:00:00", change_facts=half_shift)
    client, _ = api(case)
    query = {"reference_run_ref": refs["run_ref"], "selected_entry_ref": "run:" + changed,
             "comparison": "same_conditions", "range_start": "2026-09-09T08:00:00", "range_end": "2026-09-09T16:00:00",
             "metric_keys": ["machine_available_hours", "machine_occupied_hours", "machine_utilization_ratio",
                             "operator_available_hours", "operator_occupied_hours", "operator_utilization_ratio", "makespan_hours"]}
    with retained(case.conn):
        data = read(client, **query)["data"]
        before, after = data["selection"]["reference"], data["selection"]["entry"]
        assert data["resource_window"]["mode"] == "absolute" and data["resource_window"]["boundary"] == "half_open"
        for entry in (before, after):
            window = entry["reading_context"]["resource_window"]
            assert window["start"] == query["range_start"] and window["end"] == query["range_end"]
            assert window["time_basis"] == "factory_local"
        assert before["reading_context"]["resource_window"]["cohort"] == after["reading_context"]["resource_window"]["cohort"]
        for kind in ("machine", "operator"):
            count = len(before["reading_context"]["resource_window"]["cohort"][kind + "_refs"])
            assert count > 0
            for entry, hours in ((before, 8), (after, 4)):
                metrics = entry["metrics"]
                assert metrics[kind + "_available_hours"]["value"] == count * hours
                assert metrics[kind + "_occupied_hours"]["value"] == 3
                ratio = metrics[kind + "_utilization_ratio"]
                assert ratio["numerator_hours"] == 3 and ratio["denominator_hours"] == count * hours
                assert ratio["value"] == pytest.approx(3 / (count * hours))
        assert data["comparison"]["conditions"]["state"] == "different"
        assert any(point["range_changed"] for point in data["trend"]["points"])
        trial = read(client, **dict(query, timeline="official", selected_entry_ref="adoption:" + refs["trial_receipt_ref"]))["data"]["selection"]["entry"]
        assert trial["metrics"]["makespan_hours"]["value"] == 3
        for kind in ("machine", "operator"):
            count = len(trial["reading_context"]["resource_window"]["cohort"][kind + "_refs"])
            ratio = trial["metrics"][kind + "_utilization_ratio"]
            assert ratio["numerator_hours"] == 3 and ratio["denominator_hours"] == count * 4
            assert ratio["value"] == pytest.approx(3 / (count * 4))
            assert ratio["source"]["path"].startswith("head.admission.facts")

    case.conn.execute("INSERT INTO OpTypes(op_type_id,name) VALUES ('T2','Drilling')")
    for machine in ("M1", "M2"):
        case.conn.execute("INSERT INTO MachineOpTypes(machine_id,op_type_id) VALUES (?,'T2')", (machine,))
    case.conn.execute("UPDATE BatchOperations SET unit_hours=? WHERE id=?", (1 / 3, case.op_id))
    second = case.operation(seq=2, op_type_id="T2", op_type_name="Drilling", unit_hours=1 / 3)
    case.plan(5, [case.op_id, second], end="2026-09-09T09:00:00")
    case.conn.execute("UPDATE Schedule SET start_time='2026-09-09T09:00:00',end_time='2026-09-09T10:00:00' WHERE version=5 AND op_id=?", (second,))
    case.conn.commit()
    draft = create(case, {"base": {"plan_ref": case.plan_ref(5)}}, key="history-windowed-switch-create")
    draft = change(case, draft, task=0, start="2026-09-09T08:00:00", key="history-windowed-switch-first")["data"]
    draft = change(case, draft, task=1, start="2026-09-09T09:00:00", key="history-windowed-switch-second")["data"]
    saved = trial_service(case.conn).save(draft["draft_ref"], {"name": "Saved windowed switches"},
        draft["write_context"]["write_token"], "history-windowed-switch-save")["data"]
    saved, head, rows = load_saved_scenario(case.conn, saved["scenario_ref"])
    packet: Dict[str, Any] = {"kind": "trial", "ref": saved["scenario_ref"], "entry_ref": "scenario:" + saved["scenario_ref"],
              "saved": saved, "head": head, "rows": rows}
    reader = SimpleNamespace(condition_source=ConditionSource)
    for start, end, expected in (("2026-09-09T08:00:00", "2026-09-09T10:00:00", 1),
                                 ("2026-09-10T00:00:00", "2026-09-11T00:00:00", 0)):
        metrics, _, _ = _trial(packet, {"start": start, "end": end, "cohort": {}}, reader=reader)
        count = metrics["type_changeover_count"]
        assert count["value"] == expected and count["state"] == "complete"
        assert count["scope"] == "fixed_changeover_window"

    # Pure adapter probes reuse verified permanent rows; no saved record or receipt is rewritten.
    partial = deepcopy(packet)
    missing_task = partial["saved"]["tasks"].pop()
    missing_row = partial["rows"].pop()
    missing_ref = missing_row["operation_ref"]
    partial["saved"]["unplanned_operations"] = [{key: missing_task[key]
        for key in ("operation_ref", "batch_ref", "sequence", "piece_id", "predecessor_refs")}]
    partial["saved"]["unplanned_operations"][0]["status"] = "skipped"
    source = partial["head"]["admission"]["source"]
    definitions = {row["key"]: row for row in metric_definitions()}
    for deferred_ids, deferred_count, ordinary_count, expected_count in (([second], 1, 0, 1), ([], 0, 1, 2)):
        source.update(deferred_operation_ids=deferred_ids, deferred_operation_issues=[])
        if ordinary_count:
            partial["saved"]["unplanned_operations"] = []
            ledger = partial["saved"]["operation_effort"]
            ledger["items"] = [row for row in ledger["items"] if row["operation_ref"] != missing_ref]
        partial["saved"]["comparison"] = trial_comparison(partial["saved"]["tasks"], partial["saved"]["validation"],
            partial["rows"], partial["head"]["admission"])
        metrics, _, _ = _trial(partial, {"start": "2026-09-09T08:00:00", "end": "2026-09-09T10:00:00", "cohort": {}},
                               reader=reader)
        assert metrics["deferred_operation_count"]["value"] == deferred_count
        assert metrics["unplanned_operation_count"]["value"] == ordinary_count
        assert metrics["expected_operation_count"]["value"] == expected_count
        assert missing_ref in metrics["deferred_operation_count"]["covered_operation_refs"]
        conditions = ConditionSource(partial)
        assert conditions.operations[missing_ref]["id"] == second
        index = next(index for index, row in enumerate(partial["head"]["admission"]["facts"]["tables"]["BatchOperations"])
                     if row["id"] == second)
        assert conditions.source("operations." + missing_ref, origin="original")["path"] == (
            f"head.admission.facts.tables.BatchOperations[{index}]")
        for key in ("expected_operation_count", "deferred_operation_count", "unplanned_operation_count"):
            item = _metric(metrics[key], definitions[key], partial, _path(partial, definitions[key]))
            assert item["source"]["path"] == "head.admission"
        scheduled = _metric(metrics["scheduled_operation_count"], definitions["scheduled_operation_count"], partial,
                            _path(partial, definitions["scheduled_operation_count"]))
        assert scheduled["source"]["path"] == "saved.tasks"
        if ordinary_count:
            demand = metrics["unplanned_setup_hours"]
            assert demand["value"] is None and demand["state"] == "unavailable"
            assert demand["unknown_count"] == demand["total_count"] == 1


def test_filtered_twenty_and_fifty_pages_export_all_native_rows_and_reject_stale_snapshot(history_case):
    case = history_case
    expected = []
    for index in range(51):
        at = (datetime(2026, 9, 13, 8) + timedelta(minutes=index)).isoformat()
        expected.append("run:" + clone_run(case, at)[0])
    outside = "run:" + clone_run(case, "2026-09-14T08:00:00")[0]
    client, _ = api(case)
    query = {"timeline": "run", "date_from": "2026-09-13", "date_to": "2026-09-13", "state": "complete", "order": "asc",
             "metric_keys": ["overdue_count", "makespan_hours", "machine_utilization_ratio"]}
    with retained(case.conn):
        initial = read(client, **query)
        assert initial["data"]["page"]["total"] == 51
        query["snapshot_ref"] = initial["meta"]["snapshot_ref"]
        entries = {}
        for size, lengths in ((20, [20, 20, 11]), (50, [50, 1])):
            rows = []
            for page, length in enumerate(lengths, 1):
                result = read(client, **dict(query, page=page, size=size))["data"]
                assert len(result["entries"]) == length and result["page"]["total"] == 51
                rows.extend(row["entry_ref"] for row in result["entries"])
                for row in result["entries"]:
                    if row["entry_ref"] in entries:
                        assert row["label"] == entries[row["entry_ref"]]["label"]
                entries.update((row["entry_ref"], row) for row in result["entries"])
            assert rows == expected and outside not in rows and len(set(rows)) == 51
        assert entries[expected[0]]["label"] == "第2次运行"
        assert entries[expected[-1]]["label"] == "第52次运行"
        export_query = dict(query, metric_keys=",".join(query["metric_keys"]), page=2, size=20)
        csv_response = client.get(BASE + "/export", query_string=dict(export_query, format="csv", section="entries"))
        assert csv_response.status_code == 200, csv_response.get_data(as_text=True)
        csv_rows = list(csv.reader(io.StringIO(csv_response.data.decode("utf-8-sig"))))
        header, exported = csv_rows[0], csv_rows[1:]
        assert len(header) == len(set(header)) and len(exported) > 51 * 3
        assert int(csv_response.headers["X-Workbench-Row-Count"]) == len(exported)
        assert csv_response.headers["X-Workbench-Export-Version"] == "2"
        csv_references = exported_history_references(header, exported)
        xlsx_response = client.get(BASE + "/export", query_string=dict(export_query, format="xlsx", section="all"))
        assert xlsx_response.status_code == 200, xlsx_response.get_data(as_text=True)
        workbook = load_workbook(io.BytesIO(xlsx_response.data), read_only=True, data_only=True)
        try:
            native = list(workbook["历史条目"].iter_rows(max_col=len(header), values_only=True))
            assert tuple(header) == native[0] and len(native) == 1 + len(exported)
            assert len(set(native[0])) == len(native[0])
            assert xlsx_response.headers["X-Workbench-Export-Version"] == "2"
            native_references = exported_history_references(header, native[1:])
            assert native_references == csv_references
            primary_count = len(native) - 1 + sum(sum(1 for _row in workbook[title].values) - 1
                                                for title in ("生产条件", "算法配置"))
            assert int(xlsx_response.headers["X-Workbench-Row-Count"]) == primary_count
            for value in (0,):
                assert any(type(cell) in (int, float) and cell == value for row in native[1:] for cell in row)
            assert any(type(cell) is float and cell > 0 for row in native[1:] for cell in row)
            assert any(cell is None for row in native[1:] for cell in row)
            entry_index, key_index = header.index("历史条目编号"), header.index("指标键")
            record_index = header.index("记录类型")
            csv_by_key = {(row[entry_index], row[key_index]): row for row in exported if row[record_index] == "指标值"}
            native_by_key = {(row[entry_index], row[key_index]): row for row in native[1:] if row[record_index] == "指标值"}
            assert set(csv_by_key) == set(native_by_key) == {(ref, key) for ref in expected for key in query["metric_keys"]}
            for ref, entry in entries.items():
                quantity = entry["scope"]["batch_quantity"]
                for path, value in (("scope.batch_quantity.covered_batch_refs", quantity["covered_batch_refs"]),
                                    ("scope.batch_refs", entry["scope"]["batch_refs"]),
                                    ("scope.operation_refs", entry["scope"]["operation_refs"]),
                                    ("reading_context.resource_window.cohort.machine_refs", entry["reading_context"]["resource_window"]["cohort"]["machine_refs"]),
                                    ("reading_context.resource_window.cohort.operator_refs", entry["reading_context"]["resource_window"]["cohort"]["operator_refs"])):
                    assert csv_references[ref, None, path] == value
                for key in query["metric_keys"]:
                    source_metric = entry["metrics"][key]
                    csv_row, native_row = csv_by_key[ref, key], native_by_key[ref, key]
                    for label, value in (("条目名称", entry["label"]), ("运行编号", entry["identity"]["run_ref"]), ("候选编号", entry["identity"]["candidate_ref"]),
                                         ("指标版本", source_metric["metric_version"]), ("指标范围", source_metric["scope"]),
                                         ("指标来源编号", source_metric["source"]["ref"]),
                                         ("批次总量单位", quantity["unit"]), ("批次总量覆盖状态", quantity["state"]),
                                         ("批次总量指标范围", quantity["scope"]), ("批次总量指标版本", quantity["metric_version"]),
                                         ("批次总量来源类型", quantity["source"]["kind"]), ("批次总量来源编号", quantity["source"]["ref"]),
                                         ("批次总量来源字段", quantity["source"]["path"])):
                        assert csv_row[header.index(label)] == native_row[header.index(label)] == value
                    for label, field in (("批次总量", "value"), ("批次总量已知小计", "known_subtotal"),
                                         ("批次总量已知批次数", "known_count"), ("批次总量未知批次数", "unknown_count"),
                                         ("批次总量总批次数", "total_count")):
                        number = native_row[header.index(label)]
                        assert type(number) in (int, float) and number == quantity[field]
                        assert float(csv_row[header.index(label)]) == quantity[field]
                    assert native_row[header.index("导出格式版本")] == 2
                    assert csv_row[header.index("导出格式版本")] == "2"
                    for field in ("covered_batch_refs", "covered_operation_refs", "covered_resource_refs"):
                        assert csv_references[ref, key, field] == source_metric[field]
                    for label in ("批次总量缺项编码", "批次总量缺项说明"):
                        assert native_row[header.index(label)] is None and csv_row[header.index(label)] == ""
                    value = source_metric["value"]
                    displayed = value * 100 if value is not None and source_metric["unit"] == "%" else value
                    number = native_row[header.index("数值")]
                    if displayed is None:
                        assert number is None and csv_row[header.index("数值")] == ""
                    else:
                        assert type(number) in (int, float) and number == pytest.approx(displayed)
                        assert float(csv_row[header.index("数值")]) == pytest.approx(displayed)
        finally:
            workbook.close()
    clone_run(case, "2026-09-13T09:00:00")
    stale = client.get(BASE + "/export", query_string=dict(export_query, format="csv", section="entries"))
    assert stale.status_code == 409 and stale.get_json()["error"]["code"] == "snapshot_stale"
