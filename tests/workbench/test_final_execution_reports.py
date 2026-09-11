"""Five topics and four catalog exports through the complete current main entry."""

import json
import shutil
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from urllib.parse import urlencode

import openpyxl
import pytest

from core.services.report.exporters import xlsx
from core.services.report.report_engine import ReportEngine
from core.services.workbench.report_exports import export_table
from tests.workbench.final_execution_cases import final_e_runtime as final_e_runtime
from tests.workbench.final_execution_support import restart_preserved, serving
from tests.workbench.live_environment import write_json
from tests.workbench.reports_review_browser_oracle import rows_from_download, verify, verify_snapshot


def test_full_main_reports_all_download_bytes_scope_sql_and_real_restart(final_e_runtime):
    parent, built, runtime = final_e_runtime
    with serving(parent, built, runtime, "reports") as host:
        print("FINAL_REPORTS_ROOT " + str(host.root), flush=True)
        initial = host.state()
        host.browser("final_execution_reports.cjs")
        report = json.loads((host.root / "final-reports-initial.json").read_text(encoding="utf-8"))
        assert all(row["passed"] for row in report["cases"])
        assert len(report["downloads"]) == 56 and host.state() == initial
        ready = host.ready
        assert ready is not None
        directory = Path(ready["root"]) / "sessions" / ready["session"]
        host.stop()
        write_json(host.root / "ei-browser.json", report)
        for source, target in (("http.jsonl", "ei-http.jsonl"), ("streams.jsonl", "ei-streams.jsonl")):
            shutil.copyfile(str(directory / source), str(host.root / target))
        proof = verify(host.root)
        assert proof["wire_bytes_identical"] and proof["database_unchanged"]
        host.start(reuse=True)
        restarted_state = host.state()
        restart_audit = restart_preserved(initial, restarted_state)
        host.browser("final_execution_reports.cjs", "restart")
        restarted = json.loads((host.root / "final-reports-restart.json").read_text(encoding="utf-8"))
        assert len(restarted["cases"]) == 4 and all(row["passed"] for row in restarted["cases"])
        assert host.state() == restarted_state
        write_json(host.root / "final-reports-proof.json", {"initial_cases": report["cases"], "restart_cases": restarted["cases"],
            "wire_sql_proof": proof, "strict_changed_tables": [], "old_rows_preserved": True,
            "private_full_build": ready["assets"]["build_id"], "restart_audit": restart_audit})


@pytest.mark.parametrize("format_name", ["csv", "xlsx"])
@pytest.mark.parametrize("token", ["plain-ref", "-FU964sBK0YQFoS6hBSnxbrKOPzxW-9J", "=ref", "+ref", "@ref", "'-ref"])
def test_report_snapshot_oracle_checks_exact_format_and_rejects_wrong_identity(schema_conn, tmp_path, format_name, token):
    data = {"topic": "records", "provenance": "original-source", "data_gaps": ["unknown"],
            "plan": {"display_name": "original-plan", "plan_ref": "plan-ref"}, "scope": {"query": "B1"},
            "summary": {"rows": 2}, "columns": [{"key": "remark", "label": "备注"}]}
    exported = export_table(ReportEngine(schema_conn), data, [{"remark": "=1+1"}, {"remark": "-note"}],
                            {"snapshot_ref": token, "as_of": "2026-09-11T15:44:11"}, format_name)
    path = tmp_path / ("snapshot." + format_name)
    try:
        path.write_bytes(exported.data.read())
    finally:
        exported.data.close()
    download = {"path": str(path), "url": "http://127.0.0.1/export?" + urlencode({"snapshot_ref": token}),
                "headers": {"x-workbench-snapshot": token}}
    sheets, metadata = rows_from_download(download)
    expected = "'" + token if format_name == "csv" and token[0] in "=+-@" else token
    assert metadata["范围快照"] == expected
    assert [row[0] for row in sheets["范围全部结果"][1:]] == ["'=1+1", "'-note"]
    verify_snapshot(download, sheets, metadata)
    for wrong in {"wrong-ref", "'" + expected, expected[1:] if expected.startswith("'") else "-" + expected}:
        with pytest.raises(AssertionError):
            verify_snapshot(download, sheets, {**metadata, "范围快照": wrong})
    with pytest.raises(AssertionError):
        verify_snapshot({**download, "headers": {"x-workbench-snapshot": "wrong-ref"}}, sheets, metadata)
    with pytest.raises(AssertionError):
        verify_snapshot({**download, "url": "http://127.0.0.1/export?snapshot_ref=wrong-ref"}, sheets, metadata)
    if format_name == "csv":
        altered = deepcopy(sheets)
        altered["范围全部结果"][2][-4] = "wrong-ref"
        with pytest.raises(AssertionError):
            verify_snapshot(download, altered, metadata)
    else:
        workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=False)
        try:
            cell = next(row[1] for row in workbook["范围与口径"].iter_rows() if row[0].value == "范围快照")
            assert cell.value == token and cell.data_type == "s"
            assert all(cell.data_type != "f" for sheet in workbook for row in sheet.iter_rows() for cell in row)
        finally:
            workbook.close()
    write_json(tmp_path / "snapshot-oracle-proof.json", {"format": format_name, "raw_snapshot": token,
               "exported_snapshot": expected, "headers_exact": True, "wrong_identities_rejected": True,
               "formula_protection_preserved": True})


@pytest.mark.parametrize("write_only", (False, True), ids=("normal", "write-only"))
@pytest.mark.parametrize("prefix", ("-", "="), ids=("minus", "equals"))
def test_report_xlsx_reference_metadata_preserves_literal_text(tmp_path, write_only, prefix):
    references = {"计划引用": prefix + "plan-ref", "范围快照": prefix + "snapshot-ref"}
    formula = "=1+1"
    options = {"summary_rows": [list(item) for item in references.items()] + [["备注", formula]],
               "write_only": write_only}
    exporters = (
        ("overdue", lambda: xlsx.export_overdue_xlsx([{"batch_id": formula}], **options), "超期清单", "B2"),
        ("utilization", lambda: xlsx.export_utilization_xlsx([{"machine_id": formula}], [], **options), "设备负荷", "A2"),
        ("downtime", lambda: xlsx.export_downtime_impact_xlsx([{"machine_id": formula}], **options), "停机影响", "A2"),
        ("review", lambda: xlsx.export_execution_review_xlsx([{"batch_id_label": formula}], **options), "计划和现场实际", "A2"),
    )
    for name, export, sheet_name, address in exporters:
        stream = export()
        try:
            raw = stream.read()
        finally:
            stream.close()
        (tmp_path / (name + ".xlsx")).write_bytes(raw)
        workbook = openpyxl.load_workbook(BytesIO(raw), read_only=True, data_only=False)
        try:
            summary = {row[0].value: row[1] for row in workbook["查询摘要"].iter_rows()
                       if row and row[0].value in set(references) | {"备注"}}
            ordinary = workbook[sheet_name][address]
            formulas = [cell.coordinate for sheet in workbook for row in sheet.iter_rows()
                        for cell in row if cell.data_type == "f"]
            evidence = {"exporter": name, "write_only": write_only, "prefix": prefix, "expected": references,
                        "summary": {label: {"value": cell.value, "data_type": cell.data_type}
                                    for label, cell in summary.items()},
                        "ordinary_data": {"value": ordinary.value, "data_type": ordinary.data_type}, "formulas": formulas}
            write_json(tmp_path / (name + ".json"), evidence)
            for label, expected in references.items():
                assert summary[label].value.encode("utf-8") == expected.encode("utf-8"), (name, label)
                assert summary[label].data_type == "s"
            assert summary["备注"].value == "'" + formula
            assert summary["备注"].data_type == "s"
            assert ordinary.value == "'" + formula and ordinary.data_type == "s"
            assert formulas == []
        finally:
            workbook.close()
