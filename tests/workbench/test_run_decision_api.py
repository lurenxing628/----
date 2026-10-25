"""Real archived candidates keep comparison identity, scope and exports aligned."""

import csv
import io
import json
from copy import deepcopy

import pytest
from openpyxl import load_workbook

from core.services.workbench.facts.candidate_read_source import read_candidate_source
from core.services.workbench.facts.candidate_store import CandidateStore
from core.services.workbench.run.decision import CORE_BASELINE_COVERAGE, CORE_CANDIDATE_COVERAGE, CORE_METRIC_KEYS
from core.services.workbench.run.decision_export import write_run_decision_export
from core.services.workbench.run.worker import WorkbenchRunWorker
from scripts.verify_decision_workbench import _check_export_rows
from tests.workbench.run_candidate_support import BASE, api, compute
from tests.workbench.run_candidate_support import candidate_case as _candidate_case  # noqa: F401
from tests.workbench.test_material_stage_release import add_requirement
from web.routes.workbench.navigation_boot import WorkbenchNavigationInvalid, read_navigation


def decision(client, run_ref, **query):
    response = client.get(BASE + "/runs/" + run_ref + "/decision", query_string=query)
    assert response.status_code == 200, response.get_json()
    assert response.headers["Cache-Control"] == "no-store"
    return response.get_json()


def decision_summary(client, run_ref, **query):
    response = client.get(BASE + "/runs/" + run_ref + "/decision/summary", query_string=query)
    assert response.status_code == 200, response.get_json()
    assert response.headers["Cache-Control"] == "no-store"
    return response.get_json()


def _assert_summary_matches_full(summary, full):
    shared = ("run", "scope", "capture_ref", "catalog_complete", "metric_definitions", "selection_basis",
              "recommendation", "difference_summary", "resource_cohort")
    assert set(summary) == set(shared) | {"candidates", "baseline"}
    for key in shared:
        assert summary[key] == full[key], key
    assert len(summary["candidates"]) == len(full["candidates"])
    for core, complete in zip(summary["candidates"], full["candidates"]):
        expected = {**complete,
                    "metrics": {key: complete["metrics"][key] for key in CORE_METRIC_KEYS},
                    "deltas": {key: complete["deltas"][key] for key in CORE_METRIC_KEYS},
                    "coverage": {key: complete["coverage"][key] for key in CORE_CANDIDATE_COVERAGE}}
        assert core == expected
    baseline = full["baseline"]
    if baseline["state"] == "selection_required":
        metrics, coverage = {}, {"schedulable": baseline["coverage"]["schedulable"]}
    else:
        metrics = {key: baseline["metrics"][key] for key in CORE_METRIC_KEYS}
        coverage = {key: baseline["coverage"][key] for key in CORE_BASELINE_COVERAGE}
    assert summary["baseline"] == {**baseline, "metrics": metrics, "coverage": coverage}


def test_all_real_candidates_share_reference_and_preserve_saved_selection(candidate_case):
    case = candidate_case
    case.batch("B2", due_date="2026-09-09", priority="urgent")
    case.operation("B2", unit_hours=2)
    case.conn.commit()
    run_ref, refs = compute(case, case.settings("B1", "B2"))
    client, _ = api(case)
    initial = decision(client, run_ref)["data"]
    compared = decision(client, run_ref, reference_candidate_ref=refs[-1], selected_candidate_ref=refs[0])["data"]
    summary = decision_summary(client, run_ref, reference_candidate_ref=refs[-1], selected_candidate_ref=refs[0])["data"]
    _assert_summary_matches_full(summary, compared)
    recommended_ref = compared["recommendation"]["candidate_ref"]
    same_reference = decision(client, run_ref,
        reference_candidate_ref=recommended_ref, selected_candidate_ref=refs[0])["data"]
    _assert_summary_matches_full(decision_summary(client, run_ref,
        reference_candidate_ref=recommended_ref, selected_candidate_ref=refs[0])["data"], same_reference)
    assert {row["candidate_ref"] for row in compared["candidates"]} == set(refs)
    assert 1 <= len(compared["candidates"]) <= 3
    assert compared["scope"]["reference"]["candidate_ref"] == refs[-1]
    assert compared["scope"]["selected_candidate_ref"] == refs[0]
    assert compared["scope"]["batch_refs"] == initial["scope"]["batch_refs"]
    assert compared["scope"]["operation_refs"] == initial["scope"]["operation_refs"]
    assert compared["selection_basis"] == initial["selection_basis"]
    assert compared["recommendation"]["candidate_ref"] == initial["recommendation"]["candidate_ref"]
    basis = compared["selection_basis"]
    assert basis["generation_policy_version"] == "unified_search_v1"
    assert basis["selection_policy_version"] == basis["selection_evidence_version"] == "business_objectives_v1"
    assert basis["decision_evidence"] is basis["health_reference_candidate_ref"] is basis["health_comparison"] is None
    assert basis["target_objectives"] == ["min_overdue", "min_tardiness", "min_changeover"]
    assert set(basis["objective_winners"]) == set(basis["target_objectives"])
    assert all(winner in refs for winner in basis["objective_winners"].values())
    assert basis["objective_winners"][basis["primary_objective"]] == basis["selected_candidate_ref"]
    for row in compared["candidates"]:
        assert row["kind"] == "business_objective"
        assert "algorithm_baseline" not in row["selection_role"]
    for entry in basis["explanation"]["entries"]:
        assert all(entry[key] is None or type(entry[key]) in (int, float)
                   for key in ("reference_value", "selected_value", "value", "limit", "tolerance"))
    recommendation = compared["recommendation"]
    assert recommendation["recommendation_version"] == "decision-recommendation-v1"
    assert recommendation["recommendation_type"] == "complete"
    assert recommendation["viewed_candidate_ref"] == refs[0]
    assert recommendation["candidate_ref"] == basis["selected_candidate_ref"]
    comparison = recommendation["current_comparison"]
    assert comparison["candidate_ref"] == recommendation["candidate_ref"]
    assert comparison["viewed_candidate_ref"] == refs[0]
    assert comparison["reference"]["candidate_ref"] == refs[-1]
    assert comparison["state"] == "available"
    assert len(comparison["all_differences"]) == len(compared["metric_definitions"]) == 52
    assert len(comparison["benefits"]) <= 3 and len(comparison["costs"]) <= 3
    assert recommendation["axes"]["preflight"] == {"state": "not_checked", "source": "existing_adoption_preview"}
    saved_recommended = next(row for row in compared["candidates"] if row["candidate_ref"] == recommendation["candidate_ref"])
    definitions = {row["key"]: row for row in compared["metric_definitions"]}
    for item in comparison["all_differences"]:
        assert item["metric_version"] == definitions[item["metric_key"]]["metric_version"]
        assert len(item["sample_compared_batch_refs"]) <= 3 and len(item["supporting_refs"]) <= 3
        if item["metric_key"] not in ("overdue_count", "total_tardiness_hours"):
            expected = saved_recommended["deltas"][item["metric_key"]]
            fields = ("value", "reference_value", "selected_value", "state", "direction")
            assert {key: item[key] for key in fields} == {key: expected[key] for key in fields}
        for support in item["supporting_refs"]:
            assert support["candidate_ref"] == recommendation["candidate_ref"]
            target = (comparison["reference"]["candidate_ref"] if support.get("side", "selected") == "reference"
                      and compared["scope"]["comparison_mode"] == "candidates" else recommendation["candidate_ref"])
            assert support["drill"]["candidate_ref"] == target
            assert support["drill"]["selected_candidate_ref"] == recommendation["candidate_ref"]
            assert support["drill"]["run_ref"] == run_ref
            if support["kind"] == "operation":
                assert support["drill"]["operation_ref"] == support["ref"]
                assert "comparison_tab" not in support["drill"] and "detail_query" not in support["drill"]
    reference = next(row for row in compared["candidates"] if row["candidate_ref"] == refs[-1])
    assert all(value["value"] == 0 for value in reference["deltas"].values() if value["state"] == "comparable")
    for row in compared["candidates"]:
        delta = row["deltas"]["makespan_hours"]
        assert delta["value"] == row["metrics"]["makespan_hours"]["value"] - reference["metrics"]["makespan_hours"]["value"]
        assert row["coverage"]["schedulable"]["missing_count"] == 0


def test_legacy_saved_selection_evidence_keeps_permanent_anchors_and_view_identity(candidate_case, monkeypatch):
    from tests.workbench.legacy_run_archive_support import install_complete_legacy_archive

    # Explicit historical saved evidence, independent of current selection algorithms.
    baseline = {"candidate_key": "baseline", "score": [0, 0, 0, 0.75, 0],
        "failed_ops": 0, "primary_score": 0, "overdue_count": 0, "total_tardiness_hours": 0}
    saved_evidence = {"version": "candidate-selection-evidence-v1", "state": "complete",
        "selected": deepcopy(baseline),
        "anchors": {"raw_score_best": deepcopy(baseline), "health_reference": deepcopy(baseline)},
        "tiebreak": {"rule": "full_score_baseline_sequence_v1", "raw_score_best_candidate_key": "baseline",
            "runner_up_candidate_key": "graph_w1_of_3", "decisive_step": "baseline_preference", "decisive_score_index": None,
            "primary_tied_candidate_keys": ["baseline", "graph_w1_of_3", "graph_w2_of_3", "graph_w3_of_3"],
            "full_score_tied_candidate_keys": ["baseline", "graph_w1_of_3", "graph_w2_of_3", "graph_w3_of_3"]},
        "balanced_checks": {"eligible": False, "checked": False, "reason_code": "critical_health_not_better",
            "candidate_key": "graph_w1_of_3", "reference_candidate_key": "baseline", "health_reference_candidate_key": "baseline",
            "health_state": "same", "checks": [], "accepted": False, "override_applied": False}}
    frozen = deepcopy(saved_evidence)
    install_complete_legacy_archive(monkeypatch, selection_snapshot={"decision_evidence": saved_evidence})
    run_ref, refs = compute(candidate_case)
    client, _ = api(candidate_case)
    initial = decision(client, run_ref)["data"]
    compared = decision(client, run_ref, reference_candidate_ref=refs[-1], selected_candidate_ref=refs[-2])["data"]
    _assert_summary_matches_full(decision_summary(client, run_ref,
        reference_candidate_ref=refs[-1], selected_candidate_ref=refs[-2])["data"], compared)
    assert len(compared["candidates"]) == 4
    assert compared["scope"]["selected_candidate_ref"] == refs[-2]
    assert compared["scope"]["reference"]["candidate_ref"] == refs[-1]
    assert compared["selection_basis"] == initial["selection_basis"]
    basis, evidence = compared["selection_basis"], compared["selection_basis"]["decision_evidence"]
    assert basis["selection_evidence_version"] == evidence["version"] == "candidate-selection-evidence-v1"
    assert evidence["selected"]["candidate_ref"] == basis["selected_candidate_ref"] == refs[0]
    assert evidence["anchors"]["raw_score_best"]["candidate_ref"] == basis["raw_score_best_candidate_ref"] == refs[0]
    assert evidence["anchors"]["health_reference"]["candidate_ref"] == basis["health_reference_candidate_ref"] == refs[0]
    assert evidence["tiebreak"]["rule"] == "full_score_baseline_sequence_v1"
    assert evidence["tiebreak"]["runner_up_candidate_ref"] == refs[1]
    assert "candidate_key" not in json.dumps(evidence)
    assert saved_evidence == frozen
    assert compared["recommendation"]["candidate_ref"] == initial["recommendation"]["candidate_ref"] == refs[0]
    assert compared["recommendation"]["viewed_candidate_ref"] == refs[-2]
    for entry in basis["explanation"]["entries"]:
        assert all(entry[key] is None or type(entry[key]) in (int, float)
                   for key in ("reference_value", "selected_value", "value", "limit", "tolerance"))


def test_candidate_scope_and_window_errors_are_rejected_before_comparison(candidate_case):
    case = candidate_case
    first_ref, first_candidates = compute(case)
    accepted = case.accept(key="decision-second-run-0001")
    client, _ = api(case)
    pending = decision(client, accepted["run_ref"])["data"]
    pending_document = decision_summary(client, accepted["run_ref"])
    pending_summary = pending_document["data"]
    pending_snapshot = pending_document["meta"]["snapshot_ref"]
    _assert_summary_matches_full(pending_summary, pending)
    assert pending_summary["recommendation"]["state"] == "unavailable"
    assert pending_summary["baseline"]["state"] == "selection_required"
    assert pending_summary["baseline"]["metrics"] == {} and pending_summary["candidates"] == []
    explanation = pending_summary["selection_basis"]["explanation"]
    assert explanation["headline"] == "本次排产尚未保存推荐方案。"
    assert explanation["issues"] == []
    second = WorkbenchRunWorker(case.conn).execute(accepted["run_ref"])
    second_candidate = second["candidates"][0]["candidate_ref"]
    completed_path = BASE + "/runs/" + second["run_ref"] + "/decision"
    stale = client.get(completed_path, query_string={"snapshot_ref": pending_snapshot})
    assert stale.status_code == 409 and stale.get_json()["error"]["code"] == "snapshot_stale"
    refreshed = decision_summary(client, second["run_ref"])
    refreshed_snapshot = refreshed["meta"]["snapshot_ref"]
    full = decision(client, second["run_ref"], snapshot_ref=refreshed_snapshot)
    assert full["meta"]["snapshot_ref"] == refreshed_snapshot
    assert full["meta"]["as_of"] == refreshed["meta"]["as_of"]
    _assert_summary_matches_full(refreshed["data"], full["data"])
    path = BASE + "/runs/" + first_ref + "/decision"
    for query in ({"reference_candidate_ref": second_candidate}, {"selected_candidate_ref": second_candidate},
                  {"range_start": "2026-09-09T08:00:00"},
                  {"range_start": "2026-09-10T08:00:00", "range_end": "2026-09-09T08:00:00"}):
        for endpoint in (path, path + "/summary"):
            response = client.get(endpoint, query_string=query)
            assert response.status_code == 422, response.get_json()
            assert response.get_json()["ok"] is False
    valid = decision(client, first_ref, reference_candidate_ref=first_candidates[0],
                     range_start="2026-09-09T00:00:00", range_end="2026-09-12T00:00:00")["data"]
    assert valid["scope"]["range_start"] == "2026-09-09T00:00:00"
    assert valid["scope"]["range_end"] == "2026-09-12T00:00:00"


def test_overview_exports_require_the_screen_reference_snapshot(candidate_case, monkeypatch):
    # Fixed historical directory preserves distinct viewer/reference export risks.
    from tests.workbench.legacy_run_archive_support import install_complete_legacy_archive
    install_complete_legacy_archive(monkeypatch)
    run_ref, refs = compute(candidate_case)
    client, _ = api(candidate_case)

    document = decision(client, run_ref, reference_candidate_ref=refs[-1], selected_candidate_ref=refs[0])
    summary = decision_summary(client, run_ref, reference_candidate_ref=refs[-1], selected_candidate_ref=refs[0])
    data, snapshot = document["data"], summary["meta"]["snapshot_ref"]
    _assert_summary_matches_full(summary["data"], data)
    query = {"reference_candidate_ref": refs[-1], "selected_candidate_ref": refs[0], "snapshot_ref": snapshot}
    path = BASE + "/runs/" + run_ref + "/decision/export"
    count = len(data["candidates"]) * len(data["metric_definitions"])
    for fmt in ("csv", "xlsx"):
        response = client.get(path, query_string=dict(query, format=fmt, section="overview"))
        assert response.status_code == 200, response.get_data(as_text=True) if fmt == "csv" else response.status_code
        assert response.headers["X-Workbench-Run-Ref"] == run_ref
        assert response.headers["X-Workbench-Snapshot-Ref"] == snapshot
        assert int(response.headers["X-Workbench-Row-Count"]) == count
        if fmt == "csv":
            rows = list(csv.DictReader(io.StringIO(response.data.decode("utf-8-sig"))))
            assert len(rows) == count
            _check_export_rows(data, rows, fmt)
            assert all(ref in response.data.decode("utf-8-sig") for ref in refs)
        else:
            book = load_workbook(io.BytesIO(response.data), read_only=True, data_only=True)
            try:
                assert book.sheetnames == ["方案总览", "指标计算方式", "比较范围"]
                iterator = book["方案总览"].iter_rows(values_only=True)
                headers = next(iterator)
                rows = [dict(zip(headers, row)) for row in iterator]
                assert len(rows) == count
                _check_export_rows(data, rows, fmt)
            finally:
                book.close()
    boundaries = deepcopy(data)
    numeric_ref = "0" * 47 + "1"
    boundaries["candidates"][0].update(candidate_ref=numeric_ref, label="=1+1")
    boundaries["candidates"][1]["label"] = r"方案'甲\N"
    for fmt in ("csv", "xlsx"):
        download = write_run_decision_export(boundaries, document["meta"], fmt)
        if fmt == "csv":
            rows = list(csv.DictReader(io.StringIO(download.content.decode("utf-8-sig"))))
            assert rows[0]["候选编号"] == "'" + numeric_ref and rows[0]["候选名称"] == "'=1+1"
            assert rows[len(data["metric_definitions"])]["候选名称"] == r"方案'甲\N"
        else:
            book = load_workbook(io.BytesIO(download.content), read_only=True, data_only=False)
            try:
                rows = list(book["方案总览"].iter_rows())
                identity, label, number = rows[1][8], rows[1][9], rows[1][19]
                assert identity.value == numeric_ref and identity.data_type == "s"
                assert label.value == "=1+1" and label.data_type == "s"
                assert type(number.value) in (int, float) and number.data_type == "n"
                assert rows[1 + len(data["metric_definitions"])][9].value == r"方案'甲\N"
            finally:
                book.close()
    stale = client.get(path, query_string=dict(query, reference_candidate_ref=refs[0], format="csv"))
    assert stale.status_code == 409 and stale.get_json()["error"]["code"] == "snapshot_stale"


def test_stage_material_wait_has_complete_schedulable_scope_and_deferred_delivery(candidate_case):
    case = candidate_case
    second = case.operation(seq=2)
    add_requirement(case, second, [])
    run_ref, _ = compute(case, case.settings(material_strategy="stage"))
    client, _ = api(case)
    data = decision(client, run_ref)["data"]
    _assert_summary_matches_full(decision_summary(client, run_ref)["data"], data)
    for row in data["candidates"]:
        coverage = row["coverage"]
        assert coverage["schedulable"]["state"] == "complete"
        assert coverage["schedulable"]["scheduled_count"] == coverage["schedulable"]["expected_count"] == 1
        assert coverage["schedulable"]["missing_count"] == 0 and coverage["schedulable"]["deferred_count"] == 1
        assert coverage["delivery"]["state"] == "partial" and coverage["delivery"]["deferred_batch_count"] == 1
        assert row["metrics"]["unplanned_operation_count"]["value"] == 0
        assert row["metrics"]["deferred_operation_count"]["value"] == 1
    recommendation = data["recommendation"]
    assert recommendation["recommendation_type"] == "schedulable_scope"
    scope = recommendation["completion_scope"]
    assert scope["state"] == "partial" and scope["deferred_batch_count"] == 1
    assert scope["ordinary_incomplete_batch_count"] == scope["unplanned_operation_count"] == 0
    assert scope["deferred_operation_count"] == 1
    assert scope["remaining_work"]["operation_count"] == scope["remaining_work"]["deferred_operation_count"] == 1
    assert scope["drill"]["candidate_ref"] == recommendation["candidate_ref"]
    assert scope["drill"]["detail_query"]["work_scope"] == "remaining"
    path = BASE + "/runs/" + run_ref + "/decision/preparation"
    query = {"selected_candidate_ref": recommendation["candidate_ref"]}
    full = client.get(path, query_string=query).get_json()["data"]
    response = client.get(path, query_string=dict(query, work_scope="remaining"))
    assert response.status_code == 200, response.get_json()
    remaining = response.get_json()["data"]
    assert len(full["items"]) == 2 and remaining["page"]["total"] == len(remaining["items"]) == 1
    assert remaining["summary"] == full["summary"]
    row = remaining["items"][0]
    assert row["operation_ref"] == case.conn.execute(
        "SELECT ref FROM WorkbenchPlanSourceRefs WHERE kind='operation' AND source_key=? AND active=1", (str(second),),
    ).fetchone()[0]
    candidate_ref = recommendation["candidate_ref"]
    source = read_candidate_source(CandidateStore(case.conn), candidate_ref)
    stored = source.disposition[row["operation_ref"]]
    assert stored["status"] == "skipped" and {issue["code"] for issue in stored["issues"]} == {"material_not_ready"}
    workspace_response = client.get(BASE + "/candidates/" + candidate_ref)
    assert workspace_response.status_code == 200, workspace_response.get_json()
    unplanned = workspace_response.get_json()["data"]["unplanned_operations"]
    assert len(unplanned) == 1 and unplanned[0]["operation_ref"] == row["operation_ref"]
    assert unplanned[0]["status"] == "skipped" and unplanned[0]["reason"]["code"] == "input_operation_skipped"
    assert unplanned[0]["operation_effort"]["allocation_state"] == "deferred"
    assert unplanned[0]["reason"]["message"].startswith("缺料未安排")
    assert all(issue["message"] in unplanned[0]["reason"]["message"] for issue in stored["issues"])
    assert row["after"]["allocation_state"] == "deferred" and row["after"]["required_standard_hours"] is None
    assert row["after"]["standard"]["standard_total_hours"] == 0.75
    for allocation, count in (("deferred", 1), ("unplanned", 0), ("new_schedule", 0)):
        filtered = client.get(path, query_string=dict(query, work_scope="remaining", allocation_state=allocation,
                                                      batch_ref=case.ref("batch", "B1"), machine_ref=case.ref("machine", "M1")))
        assert filtered.status_code == 200, filtered.get_json()
        assert filtered.get_json()["data"]["page"]["total"] == len(filtered.get_json()["data"]["items"]) == count


def test_explicit_comparison_navigation_preserves_identity_and_complete_window():
    context = {"run_ref": "a" * 48, "candidate_ref": "b" * 48, "reference_candidate_ref": "c" * 48,
               "comparison_mode": "candidates", "range_start": "2026-09-09T00:00:00", "range_end": "2026-09-12T00:00:00"}
    for view in ("analysis", "gantt"):
        navigation = {"version": 1, "view": view, "context": context}
        assert read_navigation(view, {"nav": json.dumps(navigation)}) == navigation
    formal = {key: value for key, value in context.items() if key != "reference_candidate_ref"}
    formal.update(comparison_mode="admission_official", reference_view=True, plan_ref="d" * 48,
                  operation_ref="e" * 48, comparison_tab="operations", metric_view="all")
    navigation = {"version": 1, "view": "gantt", "context": formal}
    assert read_navigation("gantt", {"nav": json.dumps(navigation)}) == navigation
    resources = dict(formal, comparison_tab="resources", resource_view="full_candidate", resource_kind="operator", resource_ref="f" * 48)
    for view in ("analysis", "gantt"):
        navigation = {"version": 1, "view": view, "context": resources}
        assert read_navigation(view, {"nav": json.dumps(navigation)}) == navigation
        for tab in ("preparation", "changeovers"):
            scoped = {"version": 1, "view": view, "context": dict(resources, comparison_tab=tab, resource_kind="machine")}
            assert read_navigation(view, {"nav": json.dumps(scoped)}) == scoped
    reading = dict(context, selected_batch_ref="1" * 48, range_start="2026-09-09T00:00:00.123456")
    for tab in ("delivery", "differences", "process", "metrics", "completion"):
        navigation = {"version": 1, "view": "analysis", "context": dict(reading, comparison_tab=tab)}
        assert read_navigation("analysis", {"nav": json.dumps(navigation)}) == navigation
    mixed = dict(context, plan_ref="d" * 48)
    incomplete = {key: value for key, value in context.items() if key != "range_end"}
    for invalid in (mixed, incomplete, dict(formal, reference_candidate_ref="c" * 48),
                    {key: value for key, value in formal.items() if key != "plan_ref"},
                    {key: value for key, value in formal.items() if key != "candidate_ref"},
                    dict(formal, reference_view="true"), dict(formal, comparison_tab="invalid"),
                    dict(formal, metric_view="unknown"), dict(formal, operation_ref="invalid"), dict(reading, selected_batch_ref="invalid"),
                    dict(resources, resource_view="invalid"), dict(resources, resource_kind="invalid"), dict(resources, resource_ref="invalid")):
        with pytest.raises(WorkbenchNavigationInvalid):
            read_navigation("analysis", {"nav": json.dumps({"version": 1, "view": "analysis", "context": invalid})})
    history_state = {"query": {"timeline": "official", "comparison": "same_conditions", "reference_run_ref": "a" * 48,
        "metric_keys": ["overdue_count", "makespan_hours"], "date_from": "2026-09-09", "date_to": "2026-09-12",
        "range_start": "2026-09-09T08:00:00.123456", "range_end": "2026-09-12T16:00:00.654321",
        "resource_window_days": 3, "state": "committed", "sort": "occurred_at", "order": "asc", "page": 2, "size": 50,
        "condition_page": 2, "condition_size": 50, "selected_entry_ref": "adoption:" + "9" * 32},
        "selected_entry_ref": "adoption:" + "9" * 32, "reference_entry_ref": "run:" + "a" * 48,
        "metric_key": "makespan_hours", "detail_queries": {"conditions": {"state": "different", "family": "calendar"}}}
    target = {"source": "run_history", "history_topic": "decision_history", "decision_history_context": history_state,
              "history_query": {"state": "complete", "page": 3, "size": 50}}
    for view in ("analysis", "gantt"):
        navigation = {"version": 1, "view": view, "context": target}
        assert read_navigation(view, {"nav": json.dumps(navigation)}) == navigation
    for view, source in (("analysis", context), ("gantt", {"plan_ref": "d" * 48}),
                         ("run", {"run_ref": "a" * 48}), ("trial", {"scenario_ref": "b" * 48})):
        navigation = {"version": 1, "view": view, "context": dict(source, return_decision_history_context=history_state)}
        assert read_navigation(view, {"nav": json.dumps(navigation)}) == navigation
    for patch in ({"timeline": "invalid"}, {"comparison": "invalid"}, {"reference_run_ref": "invalid"},
                  {"metric_keys": ["overdue_count", "overdue_count"]}, {"range_end": None}, {"state": "invalid"},
                  {"sort": "invalid"}, {"order": "invalid"}, {"page": 0}, {"size": 25}):
        invalid_state = dict(history_state, query=dict(history_state["query"], **patch))
        invalid = dict(target, decision_history_context=invalid_state)
        with pytest.raises(WorkbenchNavigationInvalid):
            read_navigation("analysis", {"nav": json.dumps({"version": 1, "view": "analysis", "context": invalid})})
    for patch in ({"selected_entry_ref": "invalid"}, {"reference_entry_ref": "invalid"},
                  {"metric_key": "elapsed_seconds"}, {"detail_queries": []}):
        invalid = dict(target, decision_history_context=dict(history_state, **patch))
        with pytest.raises(WorkbenchNavigationInvalid):
            read_navigation("analysis", {"nav": json.dumps({"version": 1, "view": "analysis", "context": invalid})})


@pytest.mark.parametrize("status", ("queued", "running", "failed", "cancelled"))
def test_unsaved_recommendation_does_not_invent_missing_saved_selection_evidence(status):
    from core.services.workbench.facts.decision_selection import selection_basis

    candidate = {"candidate_ref": "a" * 48, "status": status,
                 "artifact": {"candidate_key": "unfinished", "label": "Candidate"}}
    basis = selection_basis(None, [candidate])
    assert basis["selected_candidate_ref"] is None and basis["evidence_state"] == "unavailable"
    assert basis["explanation"]["headline"] == "本次排产尚未保存推荐方案。"
    assert basis["explanation"]["issues"] == []
    legacy = selection_basis({"selection": {"selected_candidate_ref": candidate["candidate_ref"]}},
                             [dict(candidate, status="completed")])["explanation"]
    assert legacy["headline"] == "推荐方案已保存；这条旧记录未保存具体选择依据。"
    assert {issue["code"] for issue in legacy["issues"]} == {"selection_detail_missing", "health_detail_missing"}


_NAVIGATION_MALFORMED = "页面地址里的定位或范围不对，页面没有打开；系统没有恢复上次选择，也没有换记录或放宽范围。请从侧栏重新进入。"
_NAVIGATION_NON_FINITE = "页面地址里的数字不是有效数字，页面没有打开。请从侧栏重新进入。"


@pytest.mark.parametrize("raw,message", [
    ('{"version": 1, "version": 1, "view": "analysis", "context": {}}',
     "页面地址里的定位内容有重复项，页面没有打开；系统没有替你挑一个继续。请从侧栏重新进入。"),
    ('{"version": 1, "view": "analysis", "context": {"range_start": NaN}}', _NAVIGATION_NON_FINITE),
    ('{"version": 1, "view": "analysis", "context": {"range_start": -Infinity}}', _NAVIGATION_NON_FINITE),
    ('{"version": 1e999, "view": "analysis", "context": {}}', _NAVIGATION_NON_FINITE),
    ('{"version": 1, "view": "analysis", "context": {}', _NAVIGATION_MALFORMED),
], ids=("duplicate", "nan", "infinity", "overflow", "syntax"))
def test_navigation_json_rejects_duplicate_keys_and_non_finite_numbers(raw, message):
    with pytest.raises(WorkbenchNavigationInvalid) as raised:
        read_navigation("analysis", {"nav": raw})
    assert str(raised.value) == message
