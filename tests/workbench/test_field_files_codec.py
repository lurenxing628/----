"""Ten-column bytes, numeric/null/date preservation, and strict bounds."""

from datetime import datetime
from io import BytesIO

import pytest
from openpyxl import load_workbook

from core.models.workbench_command import WorkbenchCommandRejected
from core.services.workbench.field_report_files_codec import HEADERS, decode_reports, encode_reports


def test_ten_columns_null_zero_and_excel_datetime():
    content = encode_reports([{'report_no': 'BG1', 'batch_id': 'B1', 'operation_label': '1 Turning', 'completed_quantity': 0,
        'actual_start': datetime(2026, 9, 1, 8), 'effective_processing_hours': None}])
    row = decode_reports(content)[0]
    assert row['errors'] == []
    assert row['values']['completed_quantity'] == 0 and row['values']['effective_processing_hours'] is None
    assert row['values']['actual_start'] == '2026-09-01T08:00:00'
    book = load_workbook(BytesIO(content))
    assert tuple(cell.value for cell in book.active[1]) == HEADERS and book.active.max_column == 10
    book.close()


def test_formula_is_text_on_export_but_formula_input_rejected():
    content = encode_reports([{'remark': '=1+1'}])
    assert decode_reports(content)[0]['values']['remark'] == '=1+1'
    book = load_workbook(BytesIO(content)); book.active['D2'] = '=1+1'
    out = BytesIO(); book.save(out); book.close()
    assert decode_reports(out.getvalue())[0]['errors']


def test_capacity_exact_and_extra_column():
    content = encode_reports([{'report_no': 'BG' + str(index)} for index in range(5000)])
    assert len(decode_reports(content)) == 5000
    with pytest.raises(WorkbenchCommandRejected, match='5000'):
        encode_reports([{} for _ in range(5001)])
    book = load_workbook(BytesIO(encode_reports([{}]))); book.active['K2'] = 'hidden extra'
    out = BytesIO(); book.save(out); book.close()
    with pytest.raises(WorkbenchCommandRejected, match='单元格超出表头范围'):
        decode_reports(out.getvalue())


def test_invalid_date_and_boolean_remain_errors():
    content = encode_reports([{'actual_start': '2026-02-30T10:00', 'completed_quantity': True}])
    assert len(decode_reports(content)[0]['errors']) == 2


def test_merged_cells_are_rejected_not_normalized():
    book = load_workbook(BytesIO(encode_reports([{'report_no': 'BG'}])))
    book.active.merge_cells('D2:E2')
    output = BytesIO(); book.save(output); book.close()
    with pytest.raises(WorkbenchCommandRejected, match='合并'):
        decode_reports(output.getvalue())
