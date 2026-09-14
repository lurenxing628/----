"""Independently compare browser downloads with original wire bytes and public facts."""

import csv
import hashlib
import json
from io import BytesIO, StringIO
from pathlib import Path
from urllib.parse import parse_qs, urlsplit

from openpyxl import load_workbook

from core.services.workbench.report_exports import CELL_TEXTS

FIELD_HEADERS = ['报工编号', '批次号', '工序', '本次完成数量', '实际开工', '本次实际完工',
                 '有效加工工时(h)', '实际设备', '实际人员', '备注', '任务编号', '工序范围', '单件编号']
CALIB_FIELDS = ["part_no", "part_name", "sequence", "operation_label", "template_operation_ref", "template_revision",
                "template_snapshot", "old_unit_hours", "suggested_unit_hours", "deviation_percent", "deviation_basis",
                "sample_count", "candidate_count", "sample_refs", "sample_revisions", "exclusion_reasons", "method_version",
                "generated_at", "as_of", "snapshot_ref"]
CALIB_HEADERS = ["图号", "零件名称", "工序号", "工序名称", "模板工序编号", "模板版本", "模板数据版本", "旧单件定额（小时）", "建议单件定额（小时）",
                 "偏差百分比", "偏差计算状态", "可用完工记录数", "待核对完工记录数", "完工记录编号", "完工记录版本", "剔除原因", "计算方法", "生成时间", "数据截至", "数据版本编号", "筛选范围"]


def _unescape(value):
    return value[1:] if isinstance(value, str) and value.startswith("'") else value


def _scalar(value, expected):
    value = _unescape(value)
    if expected is None or expected == "":
        assert value is None or value == ""
    elif isinstance(expected, (int, float)):
        assert value is not None and value != "" and float(value) == expected
    else:
        assert str(value) == str(expected), (value, expected)


def _readings(report, snapshot, key):
    return [row["payload"] for row in report["responses"] if row["status"] == 200
            and row["payload"].get("meta", {}).get("snapshot_ref") == snapshot and key in row["payload"].get("data", {})]


def _calibration(content, name, query, report):
    snapshot = query["snapshot_ref"][0]
    source = _readings(report, snapshot, "items")[-1]
    if name.endswith("csv"):
        assert content.startswith(b"\xef\xbb\xbf")
        rows = list(csv.reader(StringIO(content.decode("utf-8-sig"))))
    else:
        book = load_workbook(BytesIO(content), read_only=True, data_only=True)
        try:
            assert book.sheetnames == ["校准建议", "范围与计算方式"]
            rows = list(book["校准建议"].iter_rows(values_only=True))
            meta = {}
            for row in book["范围与计算方式"].iter_rows(values_only=True):
                assert len(row) == 2 and isinstance(row[0], str)
                meta[row[0]] = row[1]
            assert _unescape(meta["数据版本编号"]) == snapshot
            assert meta["数据截至"] == source["meta"]["as_of"]
        finally:
            book.close()
    assert list(rows[0]) == CALIB_HEADERS
    expected = source["data"]["items"]
    assert len(rows) - 1 == source["data"]["summary"]["total"] == len(expected) == 2
    for values, item in zip(rows[1:], expected):
        for value, key in zip(values, CALIB_FIELDS):
            value = _unescape(value)
            if item[key] is None:
                assert value == "暂无数据", key
            elif isinstance(item[key], (list, dict)):
                assert isinstance(value, str)
                assert json.loads(value) == item[key], key
            else:
                _scalar(value, item[key])
        scope_value = values[-1]
        assert isinstance(scope_value, str)
        exported_scope = json.loads(scope_value)
        assert exported_scope["sort"] == query["sort"][0]
        assert exported_scope["direction"] == query["direction"][0]
        assert _unescape(values[-2]) == snapshot
    return {"rows": len(expected), "source_snapshot": snapshot, "all_suggestion_fields_equal": True}


def _field(content, name, query, report):
    book = load_workbook(BytesIO(content), read_only=True, data_only=True)
    try:
        first = book.active
        assert first is not None
        rows = list(first.iter_rows(values_only=True))
        if name == "rejected-rows":
            assert [row[0] for row in rows[1:]] == list(range(2, 12))
            assert all(any(cell for cell in row[1:]) for row in rows[1:])
            return {"rows": 10, "physical_rows": list(range(2, 12))}
        assert list(rows[0]) == FIELD_HEADERS
        if name == "template":
            assert len(rows) > 10
            assert all(row[10] and row[11] in ("共同工序", "单件") for row in rows[1:])
            assert any(row[12] == "0" for row in rows[1:])
            return {"rows": len(rows) - 1, "columns": FIELD_HEADERS}
        assert name == "saved-records"
        assert book.sheetnames == ["报工记录", "工序汇总", "录入信息"]
        source = _readings(report, query["snapshot_ref"][0], "tasks")[-1]
        tasks = source["data"]["tasks"]
        assert len(tasks) == source["data"]["page"]["total"] == 33
        original = {row["report_no"]: (task, row) for task in tasks for row in task["execution"]["reports"]}
        assert len(rows) - 1 == len(original) == 9
        for row in rows[1:]:
            task, record = original[row[0]]
            assert row[1] == task["batch_id"] and row[10] == record["recorded_against_task_ref"]
            for index, field in [(3, "completed_quantity"), (4, "actual_start"), (5, "actual_end"),
                                 (6, "effective_processing_hours"), (9, "remark")]:
                _scalar(row[index], record[field])
        summaries = list(book["工序汇总"].iter_rows(values_only=True))
        assert len(summaries) == len(tasks) + 1
        for row, task in zip(summaries[1:], tasks):
            assert row[0] == task["batch_id"] and row[1] == task["operation_label"]
            for index, key in [(2, "target_quantity"), (3, "known_completed_quantity"), (4, "remaining_quantity")]:
                _scalar(row[index], task["execution"][key])
        meta = list(book["录入信息"].iter_rows(values_only=True))
        assert len(meta) - 1 == len(original)
        for row in meta[1:]:
            record = original[row[0]][1]
            assert row[1] == record["recorded_at"]
            assert row[3] == sum(revision["action"] in ("supplement", "correct") for revision in record["correction_history"])
        return {"rows": len(original), "summary_tasks": len(tasks), "sheets": book.sheetnames, "all_original_reports_equal": True}
    finally:
        book.close()


def _actual(content, query, report):
    rows = list(csv.DictReader(StringIO(content.decode("utf-8-sig"))))
    assert len(rows) == 2
    source = report["actual_final"]["data"]
    task = report["completed_task"]
    item = next(item for item in source["items"] if item["task"]["task_ref"] == task["task_ref"])
    originals = {record["report_ref"]: record for record in item["execution"]["reports"]}
    assert {row["报工编号"] for row in rows} == set(originals)
    for row in rows:
        record = originals[row["报工编号"]]
        assert row["计划编号"] == row["录入依据计划"] == task["plan_ref"]
        assert row["任务编号"] == row["录入依据任务"] == task["task_ref"]
        assert row["工序编号"] == task["operation_ref"]
        assert row["数据版本编号"] == query["snapshot_ref"][0]
        assert json.loads(row["本地筛选"])["local_query"] == query["local_query"][0]
        assert row["计划应做数量"] == row["计划批次数量"] == ""
        assert row["计划数量缺失原因"] == "plan_target_not_recorded"
        for key, field in [("本次数量", "completed_quantity"), ("本次开工", "actual_start"),
                           ("本次结束", "actual_end"), ("有效加工工时（小时）", "effective_processing_hours")]:
            _scalar(row[key], record[field])
    return {"rows": 2, "operations": 1, "same_original_plan_task_reports": True}


def _review(content, query, report):
    rows = list(csv.DictReader(StringIO(content.decode("utf-8-sig"))))
    expected = {row["operation_ref"]: row for response in report["export_source"] for row in response["data"]["rows"]}
    columns = report["export_source"][0]["data"]["columns"]
    assert len(expected) == len(rows) == 66 and {row["工序编号"] for row in rows} == set(expected)
    for row in rows:
        original = expected[row["工序编号"]]
        for column in columns:
            value, source = row[column["label"]], original[column["key"]]
            if source is None:
                assert value == "未知", column
            elif column["key"] in CELL_TEXTS:
                assert value == CELL_TEXTS[column["key"]][source], (column, value, source)
            elif isinstance(source, (list, dict)):
                assert json.loads(_unescape(value)) == source
            else:
                _scalar(value, source)
        assert row["数据版本编号"] == query["snapshot_ref"][0]
    return {"rows": 66, "all_pages_and_columns_equal": True, "source_snapshot": query["snapshot_ref"][0]}


def verify_downloads(root, report):
    root = Path(root).resolve()
    streams = [json.loads(line) for path in (root / "sessions").glob("*/streams.jsonl")
               for line in path.read_text(encoding="utf-8").splitlines()]
    checked = []
    for download in report["downloads"]:
        path = Path(download["path"]).resolve()
        assert root in path.parents
        content = path.read_bytes()
        digest = hashlib.sha256(content).hexdigest()
        assert digest == download["sha256"] and len(content) == download["bytes"]
        matches = [row for row in streams if row["sha256"] == digest and row["bytes"] == len(content)]
        assert matches, download["name"]
        wire = matches[0]
        assert Path(wire["path"]).read_bytes() == content
        query = parse_qs(urlsplit(wire["url"]).query)
        if download["name"].startswith("calibration-"):
            proof = _calibration(content, download["name"], query, report)
        elif download["name"] == "actual-visible-scope":
            proof = _actual(content, query, report)
        elif download["name"] == "review-complete-csv":
            proof = _review(content, query, report)
        else:
            proof = _field(content, download["name"], query, report)
        checked.append({"name": download["name"], "path": str(path), "sha256": digest,
                        "original_wire_bytes_equal": True, **proof})
    assert checked
    return checked
