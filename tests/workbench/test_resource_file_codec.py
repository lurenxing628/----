"""Byte-level strictness, explicit empties, read-only types and XLSX capacities."""

from datetime import datetime
from io import BytesIO
from zipfile import ZipFile

import openpyxl
import pytest

from core.errors import ValidationError
from core.services.workbench import resource_file_writer
from core.services.workbench.resource_file_codec import read_resource_file
from core.services.workbench.resource_file_writer import check_capacity
from core.services.workbench.resource_files import WorkbenchResourceFileService
from tests.workbench.test_resource_file_support import (
    confirm,
    exported,
    file_bytes,
    resource_database,
    scope,
    snapshot,
)
from tests.workbench.test_resource_file_support import existing_raw as raw


@pytest.mark.parametrize("fmt", ("csv", "xlsx"))
@pytest.mark.parametrize("value", (0.0, 2.0, 4.0))
def test_integer_valued_readonly_hours_roundtrip_is_unchanged(resource_conn, fmt, value):
    resource_conn.execute("UPDATE OpTypes SET default_hours=? WHERE op_type_id='OT1'", (value,))
    resource_conn.commit()
    download = exported(resource_conn, "op_type", fmt, category="internal")
    preview = WorkbenchResourceFileService(resource_conn, "op_type").preview_import(download.content, file_format=fmt, scope=scope("internal"))
    assert preview.as_dict()["summary"]["rejected"] == 0
    assert confirm(resource_conn, "op_type", preview, download.content, fmt)["result"] == "unchanged"
    assert raw(resource_conn, "op_type", "OT1")["default_hours"] == value


@pytest.mark.parametrize("fmt", ("csv", "xlsx"))
@pytest.mark.parametrize("value", (0.0, 4.0))
def test_supplier_integer_days_roundtrip_does_not_rewrite_legacy_zero(resource_conn, fmt, value):
    resource_conn.execute("UPDATE Suppliers SET default_days=? WHERE supplier_id='S1'", (value,))
    resource_conn.commit()
    download = exported(resource_conn, "supplier", fmt)
    service = WorkbenchResourceFileService(resource_conn, "supplier")
    preview = service.preview_import(download.content, file_format=fmt, scope={})
    assert preview.as_dict()["summary"]["unchanged"] == download.row_count
    row = preview.as_dict()["rows"][0]
    assert not row["changes"] and row["input"]["fields"] == {}
    assert confirm(resource_conn, "supplier", preview, download.content, fmt)["result"] == "unchanged"
    assert raw(resource_conn, "supplier", "S1")["default_days"] == value


@pytest.mark.parametrize("field,left,right,expected", [
    ("default_hours", 0, 0.0, True), ("default_days", 4, 4.0, True),
    ("default_hours", True, 1.0, False), ("default_days", False, 0.0, False),
    ("default_hours", float("inf"), float("inf"), False),
    ("business_code", 4, 4.0, False), ("skill_codes", [4], [4.0], False),
    ("skill_details", [{"level": 4}], [{"level": 4.0}], False),
])
def test_numeric_equivalence_is_finite_field_scoped_and_excludes_bool(field, left, right, expected):
    from core.services.workbench.resource_file_input import same_value
    assert same_value(left, right, field) is expected


def test_bare_empty_skills_declares_but_exported_undeclared_assertion_does_not(resource_conn):
    service = WorkbenchResourceFileService(resource_conn, "operator")
    download = exported(resource_conn, "operator", "xlsx")
    preview = service.preview_import(download.content, file_format="xlsx", scope={})
    assert preview.as_dict()["summary"]["rejected"] == 0
    assert confirm(resource_conn, "operator", preview, download.content, "xlsx")["result"] == "unchanged"
    content = file_bytes([["EMPTY", "[]"]], headers=("business_code", "skill_codes"))
    preview = service.preview_import(content, file_format="csv", scope={})
    row = preview.as_dict()["rows"][0]
    assert row["result"] == "update" and row["requires_confirmation"]
    assert row["changes"]["skills_declared"] == {"before": False, "after": True}
    confirm(resource_conn, "operator", preview, content)
    assert resource_conn.execute("SELECT skills_declared FROM WorkbenchOperatorProfiles WHERE operator_id='EMPTY'").fetchone()[0] == 1
    assert resource_conn.execute("SELECT COUNT(*) FROM OperatorMachine WHERE operator_id='EMPTY'").fetchone()[0] == 0


@pytest.mark.parametrize("value", (True, 123, datetime(2026, 1, 2), "=1+1"))
@pytest.mark.parametrize("field", ("business_code", "label", "op_type_code"))
def test_xlsx_text_fields_never_guessed_or_evaluated(resource_conn, value, field):
    fields = {"business_code": "NEW", "label": "new", "status": "active", field: value}
    content = file_bytes([list(fields.values())], "xlsx", tuple(fields))
    preview = WorkbenchResourceFileService(resource_conn, "machine").preview_import(content, file_format="xlsx", scope={})
    assert preview.as_dict()["summary"]["rejected"] == 1


@pytest.mark.parametrize("fmt", ("csv", "xlsx"))
def test_formula_like_text_leading_zero_null_and_slashes_are_reversible(resource_conn, fmt):
    resource_conn.execute("INSERT INTO OpTypes(op_type_id,name,category,remark) VALUES ('000001','=1+1','internal',?)", (r"\N",))
    resource_conn.commit()
    download = exported(resource_conn, "op_type", fmt, category="internal")
    preview = WorkbenchResourceFileService(resource_conn, "op_type").preview_import(download.content, file_format=fmt, scope=scope("internal"))
    assert preview.as_dict()["summary"]["rejected"] == 0
    assert confirm(resource_conn, "op_type", preview, download.content, fmt)["result"] == "unchanged"
    assert raw(resource_conn, "op_type", "000001")["name"] == "=1+1"
    if fmt == "xlsx":
        with ZipFile(BytesIO(download.content)) as archive:
            assert archive.testzip() is None
            assert b"<f>" not in archive.read("xl/worksheets/sheet1.xml")


def test_xlsx_template_has_only_headers_and_json_array_explanations(resource_conn):
    for kind, category in (("op_type", "external"), ("operator", None), ("machine", None), ("supplier", None)):
        result = WorkbenchResourceFileService.template(kind, category=category)
        wb = openpyxl.load_workbook(BytesIO(result.content))
        try:
            ws = wb.worksheets[0]
            assert ws.max_row == 1 and result.row_count == 0
            assert all(cell.comment and "只读" in cell.comment.text and "JSON" in cell.comment.text for cell in ws[1])
            assert any(cell.comment and '["OT1","OT2"]' in cell.comment.text for cell in ws[1])
        finally:
            wb.close()


@pytest.mark.parametrize("fmt", ("csv", "xlsx"))
def test_duplicate_headers_unheaded_cells_and_multiple_sheets_fail(resource_conn, fmt):
    with pytest.raises(ValidationError):
        read_resource_file("machine", file_bytes([["M1", "M1"]], fmt, ("business_code", "编号")), fmt)
    rows = read_resource_file("machine", file_bytes([["M1", "extra"]], fmt, ("business_code",)), fmt)
    assert rows[0]["errors"][0]["field"] == "columns"
    if fmt == "xlsx":
        wb = openpyxl.Workbook()
        wb.worksheets[0].append(["business_code"])
        wb.create_sheet("another")
        buffer = BytesIO()
        wb.save(buffer)
        wb.close()
        with pytest.raises(ValidationError):
            read_resource_file("machine", buffer.getvalue(), fmt)


def test_missing_reference_metadata_is_not_repaired(resource_conn):
    resource_conn.execute("DELETE FROM WorkbenchEntityRefs WHERE kind='machine'")
    resource_conn.commit()
    before, changes = snapshot(resource_conn), resource_conn.total_changes
    from core.models.workbench_command import WorkbenchCommandRejected
    with pytest.raises(WorkbenchCommandRejected) as error:
        exported(resource_conn, "machine", "csv")
    assert error.value.code == "storage_failure"
    assert snapshot(resource_conn) == before and resource_conn.total_changes == changes


def test_xlsx_capacity_and_abort_cleanup_do_not_truncate(resource_conn, monkeypatch):
    from openpyxl.worksheet._writer import ALL_TEMP_FILES
    check_capacity(1048575, "xlsx")
    with pytest.raises(ValidationError):
        check_capacity(1048576, "xlsx")
    check_capacity(2000000, "csv")
    resource_conn.execute("UPDATE OpTypes SET remark=?", ("x" * 32768,))
    resource_conn.commit()
    before = list(ALL_TEMP_FILES)
    with pytest.raises(ValidationError, match="32767"):
        exported(resource_conn, "op_type", "xlsx", category="internal")
    assert list(ALL_TEMP_FILES) == before
    assert exported(resource_conn, "op_type", "csv", category="internal").row_count == 2
    monkeypatch.setattr(resource_file_writer, "XLSX_MAX_ROWS", 1)
    with pytest.raises(ValidationError):
        exported(resource_conn, "machine", "xlsx")
    assert list(ALL_TEMP_FILES) == before
