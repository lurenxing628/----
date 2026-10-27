"""Field route evidence against a real, isolated database and AJ's real ledger."""

from io import BytesIO

from flask import Blueprint
from openpyxl import load_workbook

from data.repositories.operator_machine_repo import OperatorMachineRepository
from data.repositories.workbench_execution_report_repo import WorkbenchExecutionReportRepository, new_ref
from tests.workbench.execution_ledger_support import all_rows
from tests.workbench.field_workspace_support import BASE, FieldAPI, _ledger_fixture, success  # noqa: F401
from tests.workbench.field_workspace_support import field_api as _field_api  # noqa: F401
from web.routes.workbench.reports import register_report_routes


def test_unreported_null_zero_partial_finish_and_readonly(field_api):
    api = field_api
    before = all_rows(api.case.conn)
    task = api.task()
    assert task['execution']['execution_state'] == 'unreported'
    assert task['execution']['reports'] == []
    assert all_rows(api.case.conn) == before
    assert not any(sql.lstrip().split()[0].upper() in ('INSERT', 'UPDATE', 'DELETE', 'CREATE') for sql in api.app.ak_statements)
    api.create({'actual_start': '2026-09-01T08:00:00'})
    p = api.task()['execution']
    assert p['execution_state'] == 'started' and p['remaining_quantity'] is None
    record = p['reports'][0]
    body = api.body(record['write_context'], {'original_revision_ref': record['revision_ref'], 'reason': '核对现场记录',
        **api.case.values(0, actual_start='2026-09-01T08:00:00', actual_end='2026-09-01T10:00:00', effective_processing_hours=0)})
    success(api.client.post(BASE + '/reports/' + record['report_ref'] + '/supplement', json=body))
    assert api.task()['execution']['reports'][0]['completed_quantity'] == 0
    api.create(api.case.values(4, actual_start='2026-09-01T11:00:00', actual_end='2026-09-01T13:00:00'))
    assert api.task()['execution']['execution_state'] == 'partial'
    api.create(api.case.values(6, actual_start='2026-09-01T14:00:00', actual_end='2026-09-01T16:00:00'))
    p = api.task()['execution']
    assert p['execution_state'] == 'complete' and p['known_completed_quantity'] == 10 and p['remaining_quantity'] == 0
    api.case.conn.execute("INSERT INTO Machines(machine_id,name,op_type_id) VALUES ('M2','Lathe','T1'),('M3','M1','T1')")
    api.case.conn.execute("INSERT INTO Operators(operator_id,name) VALUES ('O2','Operator'),('O3','O1')")
    api.case.conn.commit()
    reading = api.read()
    before = all_rows(api.case.conn)
    exported = api.client.get(BASE + '/files/export', query_string={**reading['data']['scope'],
        'snapshot_ref': reading['meta']['snapshot_ref']})
    assert exported.status_code == 200
    book = load_workbook(BytesIO(exported.data))
    assert [(row[7].value, row[8].value) for row in list(book.worksheets[0])[1:]] == [('M1', 'O1')] * 3
    checked = success(api.upload(exported.data, reading))['data']
    assert checked['can_confirm'] and checked['summary']['unchanged'] == 3 and checked['summary']['rejected'] == 0
    assert all_rows(api.case.conn) == before
    book.worksheets[0].cell(2, 8).value = 'Lathe'
    output = BytesIO()
    book.save(output)
    book.close()
    checked = success(api.upload(output.getvalue(), reading))['data']
    assert not checked['can_confirm'] and checked['summary']['rejected'] == 1
    assert all_rows(api.case.conn) == before


def test_supplement_keeps_known_facts_through_correction(ledger_case):
    case = ledger_case
    api = FieldAPI(case)
    reports = Blueprint("report_hours", __name__)
    register_report_routes(reports)
    api.app.register_blueprint(reports)
    task = api.task()
    api.create(case.values(10, effective_processing_hours=None, actual_operator_ref=None), task)
    case.conn.execute("INSERT INTO Machines(machine_id,name,op_type_id) VALUES ('M2','New planned machine','T1')")
    case.plan(2, [case.op_id])
    case.conn.execute("UPDATE Schedule SET machine_id='M2' WHERE version=2")
    case.conn.commit()
    case.conn.execute("INSERT INTO OpTypes(op_type_id,name) VALUES ('T2','New craft')")
    case.conn.execute("UPDATE Machines SET op_type_id='T2' WHERE machine_id='M1'")
    OperatorMachineRepository(case.conn).remove("O1", "M1")
    case.conn.commit()
    task = api.task()
    assert task["planned_machine_ref"] == case.ref("machine", "M2")
    assert len(api.read(resource_type="machine", resource_ref=case.ref("machine", "M1"))["data"]["tasks"]) == 1
    assert len(api.read(query="Lathe")["data"]["tasks"]) == 1
    summary = api.read()["data"]["summary"]
    assert summary["effective_processing_hours"] is None and summary["unknown_hour_reports"] == 1
    record = task["execution"]["reports"][0]
    before = all_rows(case.conn)
    response = api.client.post(BASE + "/reports/" + record["report_ref"] + "/supplement", json=api.body(
        record["write_context"], {"original_revision_ref": record["revision_ref"], "reason": "不能用补齐覆盖原开工时间",
                                  "actual_start": "2026-09-09T09:00:00", "effective_processing_hours": 0.5}))
    assert response.status_code == 409 and all_rows(case.conn) == before
    success(api.client.post(BASE + "/reports/" + record["report_ref"] + "/supplement", json=api.body(
        record["write_context"], {"original_revision_ref": record["revision_ref"], "reason": "只补原记录缺少的工时",
                                  "effective_processing_hours": 1.5})))
    summary = api.read()["data"]["summary"]
    assert summary["effective_processing_hours"] == 1.5 and summary["unknown_hour_reports"] == 0
    report = success(api.client.get("/api/workbench/v1/analytics", query_string={"topic": "records"}))["data"]
    assert report["summary"]["production_reports"] == 1 and "records" not in report["summary"]
    assert report["summary"]["effective_processing_hours"] == 1.5 and report["summary"]["unknown_hour_events"] == 0
    assert report["resources"]["machines"][0]["effective_processing_hours"] == 1.5
    assert report["resources"]["people"][0]["effective_processing_hours"] == 1.5
    execution = api.task()["execution"]
    record = execution["reports"][0]
    assert (record["actual_start"], record["actual_end"], record["completed_quantity"]) == (
        "2026-09-09T08:00:00", "2026-09-09T10:00:00", 10)
    assert record["actual_machine_ref"] == case.ref("machine", "M1") and record["actual_operator_ref"] is None
    assert execution["known_completed_quantity"] == 10 and execution["unknown_record_count"] == 0
    assert execution["data_quality"] == "incomplete"
    before = all_rows(case.conn)
    response = api.client.post(BASE + "/reports/" + record["report_ref"] + "/correct", json=api.body(
        record["write_context"], {"original_revision_ref": record["revision_ref"], "reason": "未知人员仍需当前资格",
                                "actual_operator_ref": case.ref("operator", "O1")}))
    assert response.status_code == 409 and all_rows(case.conn) == before
    # Maintaining hours/remark on an earlier plan's report preserves historical machine/operator actuals.
    original = case.ledger.load([record["operation_ref"]])["reports"][record["operation_ref"]][record["report_ref"]][-1]
    damaged = dict(original, revision_ref=new_ref(), sequence=3, previous_revision_ref=original["revision_ref"],
                   action="correct", reason="old accepted correction", values=dict(original["values"],
                   effective_processing_hours=0.5))
    request_key = case.conn.execute("SELECT request_key FROM WorkbenchProductionReportRevisions WHERE revision_ref=?",
                                   (original["revision_ref"],)).fetchone()[0]
    case.conn.execute("BEGIN")
    WorkbenchExecutionReportRepository(case.conn).append(damaged, request_key=request_key)
    case.conn.commit()
    assert api.read()["data"]["summary"]["effective_processing_hours"] == 0.5
    record = api.task()["execution"]["reports"][0]
    success(api.client.post(BASE + "/reports/" + record["report_ref"] + "/correct", json=api.body(
        record["write_context"], {"original_revision_ref": record["revision_ref"], "reason": "恢复核对后的有效工时",
                                "effective_processing_hours": 1.5})))
    execution = api.task()["execution"]
    assert execution["known_completed_quantity"] == 10 and execution["unknown_record_count"] == 0
    assert execution["reports"][0]["actual_start"] == "2026-09-09T08:00:00"
    assert api.read()["data"]["summary"]["effective_processing_hours"] == 1.5


def test_file_confirmation_retains_source_after_read_expiry_and_rechecks_write_expiry(field_api, monkeypatch):
    from core.services.workbench.commands import WorkbenchCommandService
    from web import public_token_registry
    from web.routes.workbench.execution_files import PREVIEW_NAMESPACE
    from web.routes.workbench.read_context import _SCOPE

    api = field_api
    api.create()
    reading = api.read()
    exported = api.client.get(BASE + '/files/export', query_string={**reading['data']['scope'],
        'snapshot_ref': reading['meta']['snapshot_ref']})
    assert exported.status_code == 200
    preview = success(api.upload(exported.data, reading))
    assert preview['data']['can_confirm']
    with api.app.app_context():
        public_token_registry._registry()[_SCOPE]['tokens'][reading['meta']['snapshot_ref']]['expires_at'] = 0
    confirmed = success(api.client.post(BASE + '/files/confirm', json=api.confirm_body(preview)))
    assert confirmed['result'] == 'unchanged'

    preview = success(api.upload(exported.data))
    ref, before = preview['data']['preview_ref'], all_rows(api.case.conn)
    execute = WorkbenchCommandService.execute

    def expired_after_prefetch(service, **kwargs):
        # The file has been prepared; the write-lock guard must still reject expiry.
        public_token_registry._registry()[PREVIEW_NAMESPACE]['tokens'][ref]['expires_at'] = 0
        return execute(service, **kwargs)

    with monkeypatch.context() as patch:
        patch.setattr(WorkbenchCommandService, 'execute', expired_after_prefetch)
        response = api.client.post(BASE + '/files/confirm', json=api.confirm_body(preview))
    assert response.status_code == 409 and response.get_json()['error']['code'] == 'stale_write'
    assert all_rows(api.case.conn) == before


def test_report_and_task_contexts_locate_shared_operation_facts_once(field_api, record_property):
    from core.models.content import content_text
    from core.services.workbench.facts.execution_projection import attach_context
    from web.routes.workbench.write_context import issue_write_context

    class SharedFacts(str):
        equalities = 0
        __hash__ = str.__hash__

        def __eq__(self, other):
            SharedFacts.equalities += 1
            return str.__eq__(self, other)

    api = field_api
    api.create()
    task = api.task()
    with api.case.ledger.read_snapshot():
        facts = api.case.ledger.load([task['operation_ref']])
        projection = api.case.ledger.project_loaded(facts, contexts=False)[0]
        snapshot = SharedFacts(content_text(api.case.ledger.fact_snapshot(facts, task['operation_ref'])))
    # The common attachment owner serves every report and the task in either ledger path.
    with api.app.app_context():
        first = attach_context(projection, issue_write_context, snapshot)
        same_content = SharedFacts(snapshot)
        SharedFacts.equalities = 0
        reopened = attach_context(projection, issue_write_context, same_content)
    assert SharedFacts.equalities == 1
    assert first.write_context['write_token'] == reopened.write_context['write_token']
    assert [row.write_context['write_token'] for row in first.reports] == [row.write_context['write_token'] for row in reopened.reports]
    record_property('shared_facts_equalities', SharedFacts.equalities)


def test_file_preview_scope_rejects_duplicate_keys_and_non_finite_numbers(field_api):
    for raw in ('{"query": "a", "query": "b"}', '{"query": NaN}', '{"query": Infinity}', '{"query": 1e999}', '{"query": '):
        response = field_api.client.post(BASE + '/files/preview', data={
            'file': (BytesIO(b'never parsed'), 'field.xlsx'), 'scope': raw, 'snapshot_ref': 'x'})
        assert response.status_code == 400, response.get_data(as_text=True)
        error = response.get_json()['error']
        assert (error['code'], error['message']) == (
            'invalid_input', '文件导入范围读不出来，报工还没有写入。请刷新页面后重新点「开始预检」。'), raw
