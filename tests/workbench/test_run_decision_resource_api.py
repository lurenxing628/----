"""Archived resource comparisons retain fixed cohorts, windows and exports."""

import csv
import io
from copy import deepcopy

import pytest
from openpyxl import load_workbook

from core.infrastructure.read_evidence import read_evidence_scope
from core.services.scheduler.calendar.service import CalendarService
from core.services.workbench.facts import resource_projection
from core.services.workbench.facts.candidate_baseline import AdmissionBaseline
from core.services.workbench.facts.candidate_store import CandidateStore
from core.services.workbench.facts.resource_projection import ResourceTaskIndex
from core.services.workbench.run.decision_export import read_decision_export, write_run_decision_export
from core.services.workbench.run.decision_resource_export import coverage_operation_rows
from core.services.workbench.run.decision_source import DecisionSource
from core.services.workbench.run.worker import WorkbenchRunWorker
from tests.workbench.run_candidate_adoption_support import INTENT, rewrite_candidate, service
from tests.workbench.run_candidate_support import BASE, api, compute
from tests.workbench.run_candidate_support import candidate_case as _candidate_case  # noqa: F401
from tests.workbench.test_run_decision_api import decision
from tests.workbench.test_run_decision_impact_api import operation_ref

DAY = "2026-09-09"
WINDOW = {"range_start": DAY + "T08:00:00", "range_end": DAY + "T16:00:00"}


def resource_document(client, run_ref, section="resources", **query):
    response = client.get(BASE + "/runs/" + run_ref + "/decision/" + section, query_string=query)
    assert response.status_code == 200, response.get_json()
    assert response.headers["Cache-Control"] == "no-store"
    return response.get_json()


def add_pair(case, machine="M2", operator="O2", label=r"备用'甲\N"):
    case.conn.execute("INSERT INTO Machines(machine_id,name,op_type_id) VALUES (?,?,'T1')", (machine, label))
    case.conn.execute("INSERT INTO Operators(operator_id,name) VALUES (?,?)", (operator, label))
    case.conn.execute("INSERT INTO OperatorMachine(operator_id,machine_id) VALUES (?,?)", (operator, machine))


def legacy_resource_archive(case, monkeypatch, settings=None, *, edits=None):
    """Create a named legacy reading fixture before deriving its immutable models."""
    from tests.workbench import decision_history_support
    from tests.workbench.legacy_run_archive_support import install_complete_legacy_archive

    with monkeypatch.context() as archived:
        install_complete_legacy_archive(archived)
        original_run, original_refs = compute(case, settings)
    existing_history = getattr(case, "history_refs", None)
    case.history_refs = {"run_ref": original_run, "candidate_ref": original_refs[0]}
    derive = decision_history_support._save_cloned_read_models

    def derive_fixture(conn, run_ref):
        refs = [row[0] for row in conn.execute(
            "SELECT candidate_ref FROM WorkbenchRunCandidates WHERE run_ref=? ORDER BY sequence", (run_ref,))]
        for index, edit in (edits or {}).items():
            rewrite_candidate(case, refs[index], edit)
        derive(conn, run_ref)

    try:
        with monkeypatch.context() as cloned:
            cloned.setattr(decision_history_support, "_save_cloned_read_models", derive_fixture)
            run_ref, _ = decision_history_support.clone_run(case, "2026-09-10T12:00:00.000000")
    finally:
        if existing_history is None:
            del case.history_refs
        else:
            case.history_refs = existing_history
    refs = [row[0] for row in case.conn.execute(
        "SELECT candidate_ref FROM WorkbenchRunCandidates WHERE run_ref=? ORDER BY sequence", (run_ref,))]
    return run_ref, refs


def reassigned_run(case, monkeypatch):
    add_pair(case)
    case.conn.commit()
    return legacy_resource_archive(case, monkeypatch,
        edits={-1: lambda row: row.update(machine_id="M2", operator_id="O2")})


def row_by_ref(document, resource_ref):
    return next(row for row in document["data"]["items"] if row["resource_ref"] == resource_ref)


def test_candidate_switch_keeps_unused_resources_and_archived_capacity(candidate_case, monkeypatch):
    case = candidate_case
    run_ref, refs = reassigned_run(case, monkeypatch=monkeypatch)
    client, _ = api(case)
    query = {**WINDOW, "selected_candidate_ref": refs[0], "reference_candidate_ref": refs[0]}
    documents = {}
    for kind, used, unused in (("machine", "M1", "M2"), ("operator", "O1", "O2")):
        document = resource_document(client, run_ref, **dict(query, resource_kind=kind))
        documents[kind] = document
        resource_refs = {case.ref(kind, used), case.ref(kind, unused)}
        assert set(document["data"]["scope"]["resource_refs"][kind]) == resource_refs
        assert {row["resource_ref"] for row in document["data"]["items"]} == resource_refs
        idle = row_by_ref(document, case.ref(kind, unused))["after"]
        assert idle["available_hours"] == 8
        assert idle["occupied_hours"] == idle["summed_load_hours"] == idle["utilization_ratio"] == 0
        assert row_by_ref(document, case.ref(kind, used))["after"]["occupied_hours"] == 0.75
        summary = document["data"]["summary"]["after"]
        assert summary["resource_count"] == 2 and summary["used_resource_count"] == 1
        assert summary["available_hours"] == 16 and summary["occupied_hours"] == 0.75
        assert summary["utilization_ratio"] == pytest.approx(0.75 / 16)
        assert summary["balance_cv"] == pytest.approx(1)
        switched = resource_document(client, run_ref, **dict(query, selected_candidate_ref=refs[-1],
                                                           reference_candidate_ref=refs[-1], resource_kind=kind))
        assert switched["data"]["scope"]["resource_refs"] == document["data"]["scope"]["resource_refs"]
        assert row_by_ref(switched, case.ref(kind, used))["after"]["occupied_hours"] == 0
        assert row_by_ref(switched, case.ref(kind, unused))["after"]["occupied_hours"] == 0.75
    window_reads, calendar_reads = [], []
    original_window, original_calendar = ResourceTaskIndex.window, resource_projection._ResourceCalendar.__init__

    def record_window(index, resource, calendar, **values):
        window_reads.append(resource["ref"])
        return original_window(index, resource, calendar, **values)

    def record_calendar(value, calendar, low, high, **kwargs):
        calendar_reads.append(calendar["resource_ref"])
        original_calendar(value, calendar, low, high, **kwargs)

    def unexpected_execution(_baseline):
        pytest.fail("普通覆盖汇总只需要真实正式成员。")

    with monkeypatch.context() as scoped:
        scoped.setattr(ResourceTaskIndex, "window", record_window)
        scoped.setattr(resource_projection._ResourceCalendar, "__init__", record_calendar)
        scoped.setattr(AdmissionBaseline, "_execution", unexpected_execution)
        overview = decision(client, run_ref, **query)["data"]
    cohort_refs = {case.ref(kind, code) for kind, codes in (("machine", ("M1", "M2")), ("operator", ("O1", "O2"))) for code in codes}
    assert set(calendar_reads) == cohort_refs and len(calendar_reads) == len(cohort_refs)
    assert all(window_reads.count(ref) == len(refs) for ref in cohort_refs)
    assert len(window_reads) == len(refs) * len(cohort_refs)
    chosen = next(row for row in overview["candidates"] if row["candidate_ref"] == refs[0])
    for kind in ("machine", "operator"):
        assert chosen["metrics"][kind + "_utilization_ratio"]["value"] == pytest.approx(documents[kind]["data"]["summary"]["after"]["utilization_ratio"])
    assert chosen["metrics"]["machine_util_avg"]["value"] == 1
    window_reads.clear()
    calendar_reads.clear()
    store = CandidateStore(case.conn)
    with store.snapshot(), read_evidence_scope(case.conn):
        source = DecisionSource(store, run_ref, selected_candidate_ref=refs[-1], reference_candidate_ref=refs[0], **WINDOW)
        engine = source.resources
        with monkeypatch.context() as scoped:
            scoped.setattr(ResourceTaskIndex, "window", record_window)
            scoped.setattr(resource_projection._ResourceCalendar, "__init__", record_calendar)
            engine.prepare_keys([])
            assert window_reads == calendar_reads == [] and engine.indexes == engine.tasks == {}
            engine.prepare_keys([refs[-1], "reference", refs[0]])
            pair_metrics = engine.metrics(refs[-1])
            reference_metrics = engine.metrics("reference")
            assert engine.metrics(refs[0]) == reference_metrics
        selected = next(row for row in overview["candidates"] if row["candidate_ref"] == refs[-1])
        assert pair_metrics == {key: selected["metrics"][key] for key in pair_metrics}
        assert set(engine.indexes) == {refs[0], refs[-1]}
        assert source.scope["resource_refs"] == overview["scope"]["resource_refs"]
    assert set(calendar_reads) == cohort_refs and len(calendar_reads) == len(cohort_refs)
    assert all(window_reads.count(ref) == 2 for ref in cohort_refs)
    assert len(window_reads) == 2 * len(cohort_refs)
    for kind, idle in (("machine", "M2"), ("operator", "O2")):
        absent = resource_document(client, run_ref, **{**WINDOW, "comparison_mode": "admission_official",
                                                     "selected_candidate_ref": refs[0], "resource_kind": kind})
        assert absent["data"]["baseline"]["state"] == "absent"
        before = absent["data"]["summary"]["before"]
        assert before["available_hours"] == before["known_available_hours"] == 16
        assert before["known_capacity_count"] == 2
        for field in ("occupied_hours", "summed_load_hours", "utilization_ratio", "high_pressure_count"):
            assert before[field] is None and absent["data"]["summary"]["deltas"][field] is None
        for row in absent["data"]["items"]:
            assert row["before"]["available_hours"] == 8 and row["before"]["high_pressure"] is None
            assert all(row["deltas"][field] is None for field in ("occupied_hours", "summed_load_hours", "utilization_ratio", "peak_utilization"))
        assert row_by_ref(absent, case.ref(kind, idle))["after"]["occupied_hours"] == 0
    case.conn.execute("UPDATE WorkbenchCalendarDefaults SET periods_json='[{\"start\":\"08:00\",\"end\":\"12:00\",\"day_offset\":0}]'")
    case.conn.execute("UPDATE Machines SET name='Current lathe' WHERE machine_id='M1'")
    case.conn.commit()
    previous = documents["machine"]
    reread = resource_document(client, run_ref, **dict(query, resource_kind="machine", snapshot_ref=previous["meta"]["snapshot_ref"]))
    assert reread["data"] == previous["data"]


def test_partial_night_window_splits_days_and_preserves_overlap_identity(candidate_case, monkeypatch):
    case = candidate_case
    second, boundary = case.operation(seq=2), case.operation(seq=3)
    case.conn.commit()
    calendar = CalendarService(case.conn)
    calendar._admin.repo.upsert_many([
        calendar._admin._build_work_calendar(
            date_value=DAY, shift_start="22:00", shift_end="06:00", efficiency=1, allow_normal="yes", allow_urgent="yes"
        )
    ])
    calendar._admin.operator_calendar_repo.upsert_many([
        calendar._admin._build_operator_calendar(
            operator_id="O1", date_value=DAY, shift_start="22:00", shift_end="06:00", efficiency=0.5,
            allow_normal="yes", allow_urgent="yes"
        )
    ])
    case.conn.commit()
    times = {case.op_id: ("2026-09-09T23:00:00", "2026-09-10T01:00:00"),
             second: ("2026-09-10T00:00:00", "2026-09-10T02:00:00"),
             boundary: ("2026-09-10T01:30:00", "2026-09-10T02:30:00")}
    run_ref, refs = legacy_resource_archive(case, monkeypatch, case.settings(end_date="2026-09-10"),
        edits={0: lambda row: row.update(start_time=times[row["op_id"]][0], end_time=times[row["op_id"]][1])})
    client, _ = api(case)
    query = {"selected_candidate_ref": refs[0], "reference_candidate_ref": refs[0],
             "range_start": "2026-09-09T23:30:00", "range_end": "2026-09-10T01:30:00"}
    expected_refs = {operation_ref(case, case.op_id), operation_ref(case, second)}
    for kind, key in (("machine", "M1"), ("operator", "O1")):
        ref = case.ref(kind, key)
        document = resource_document(client, run_ref, **dict(query, resource_kind=kind))
        after = row_by_ref(document, ref)["after"]
        assert after["available_hours"] == after["occupied_hours"] == after["span_occupied_hours"] == 2
        assert after["summed_load_hours"] == after["span_summed_hours"] == 3
        assert after["overlap_hours"] == after["span_overlap_hours"] == 1
        assert row_by_ref(document, ref)["deltas"]["overlap_span_hours"] == 0
        assert after["outside_calendar_hours"] == 0
        assert after["utilization_ratio"] == 1 and after["demand_ratio"] == 1.5
        assert after["conflict"] is True and after["high_pressure"] is True
        days = resource_document(client, run_ref, "resource-days", **dict(query, resource_kind=kind, resource_ref=ref))["data"]
        assert [row["date"] for row in days["items"]] == [DAY, "2026-09-10"]
        first, next_day = days["items"]
        assert first["start"] == query["range_start"] and first["end"] == "2026-09-10T00:00:00"
        assert next_day["start"] == "2026-09-10T00:00:00" and next_day["end"] == query["range_end"]
        assert first["after"]["occupied_hours"] == first["after"]["summed_load_hours"] == 0.5
        assert next_day["after"]["occupied_hours"] == 1.5 and next_day["after"]["summed_load_hours"] == 2.5
        assert first["after"]["conflicts"] == []
        conflict = next_day["after"]["conflicts"][0]
        assert (conflict["start"], conflict["end"]) == ("2026-09-10T00:00:00", "2026-09-10T01:00:00")
        assert set(conflict["operation_refs"]) == expected_refs and len(conflict["row_refs"]) == 2
        assert conflict["drill"]["candidate_ref"] == refs[0] and conflict["drill"]["resource_ref"] == ref
        task_refs = {task["operation_ref"] for day in days["items"] for task in day["after"]["tasks"]}
        assert task_refs == expected_refs
        assert operation_ref(case, boundary) not in task_refs
        shifts = resource_document(client, run_ref, **dict(query, resource_kind=kind, grain="shift"))["data"]
        heatmap = next(row for row in shifts["heatmap"]["items"] if row["resource_ref"] == ref)
        assert len(heatmap["cells"]) == 1
        shift = heatmap["cells"][0]
        assert (shift["start"], shift["end"]) == (query["range_start"], query["range_end"])
        assert shift["after"]["occupied_hours"] == 2 and shift["after"]["summed_load_hours"] == 3
        assert heatmap["continuations"] == [{"group_key": "2026-09-10", "bucket_key": shift["bucket_key"], "date": DAY,
            "start": query["range_start"], "end": query["range_end"], "continuation_start": "2026-09-10T00:00:00",
            "continuation_end": query["range_end"], "basis": "continuation_of_archived_work_period"}]
        selected_shift = resource_document(client, run_ref, "resource-days", **dict(query, resource_kind=kind, resource_ref=ref,
            grain="shift", bucket_start=shift["start"], bucket_end=shift["end"]))["data"]
        assert len(selected_shift["items"]) == 1
        assert len(selected_shift["items"][0]["after"]["tasks"]) == 2
        assert {row["row_ref"] for row in selected_shift["items"][0]["after"]["tasks"]} == {
            row["row_ref"] for day in days["items"] for row in day["after"]["tasks"]}
    reference_query = dict(query, reference_candidate_ref=refs[-1], resource_kind="machine", resource_ref=case.ref("machine", "M1"))
    compared = resource_document(client, run_ref, **reference_query)
    assert compared["data"]["items"][0]["deltas"]["overlap_span_hours"] == 1
    export_path = BASE + "/runs/" + run_ref + "/decision/export"
    response = client.get(export_path, query_string=dict(reference_query, snapshot_ref=compared["meta"]["snapshot_ref"],
                                                        section="resources", format="csv"))
    assert response.status_code == 200, response.get_json()
    exported = list(csv.DictReader(io.StringIO(response.data.decode("utf-8-sig"))))
    assert len(exported) == 1 and float(exported[0]["变化保留安排重叠时段（小时）"]) == 1
    full_shift_query = {"selected_candidate_ref": refs[0], "reference_candidate_ref": refs[0], "resource_kind": "machine",
                        "resource_ref": case.ref("machine", "M1"), "grain": "shift",
                        "range_start": DAY + "T22:00:00", "range_end": "2026-09-10T06:00:00"}
    full_shift = resource_document(client, run_ref, "resource-days", **full_shift_query)
    shift_after = full_shift["data"]["items"][0]["after"]
    assert shift_after["peak_utilization"] == 3.5 / 8 and shift_after["daily_peak_utilization"] == 0.5
    for fmt in ("csv", "xlsx"):
        response = client.get(export_path, query_string={**full_shift_query, "snapshot_ref": full_shift["meta"]["snapshot_ref"],
                                                        "section": "resource-days", "format": fmt})
        assert response.status_code == 200, response.get_json()
        if fmt == "csv":
            table = list(csv.reader(io.StringIO(response.data.decode("utf-8-sig"))))
        else:
            book = load_workbook(io.BytesIO(response.data), read_only=True, data_only=False)
            try:
                table = list(book["资源班次"].iter_rows(values_only=True))
            finally:
                book.close()
        headers, values = table[0], table[1]
        assert len(headers) == len(set(headers)) == len(values)
        assert float(values[headers.index("所选时段占用比例（%）")]) == 43.75
        assert float(values[headers.index("所选日峰占用比例（%）")]) == 50

    midnight_query = {**full_shift_query, "range_start": "2026-09-10T00:00:00"}
    midnight = resource_document(client, run_ref, **midnight_query)["data"]
    heatmap = next(row for row in midnight["heatmap"]["items"] if row["resource_ref"] == midnight_query["resource_ref"])
    assert heatmap["cells"] == []
    assert heatmap["continuations"] == [{
        "group_key": "2026-09-10", "bucket_key": "shift:" + midnight_query["range_start"] + "/" + midnight_query["range_end"],
        "date": DAY, "start": midnight_query["range_start"], "end": midnight_query["range_end"],
        "continuation_start": midnight_query["range_start"], "continuation_end": midnight_query["range_end"],
        "basis": "continuation_of_archived_work_period"}]
    midnight_detail_query = dict(midnight_query, bucket_start=midnight_query["range_start"], bucket_end=midnight_query["range_end"])
    midnight_detail = resource_document(client, run_ref, "resource-days", **midnight_detail_query)
    assert len(midnight_detail["data"]["items"]) == 1
    tail = midnight_detail["data"]["items"][0]
    assert tail["date"] == tail["group_key"] == DAY
    assert (tail["start"], tail["end"]) == (midnight_query["range_start"], midnight_query["range_end"])
    assert tail["after"]["available_hours"] == 6
    assert {task["operation_ref"] for task in tail["after"]["tasks"]} == {
        operation_ref(case, case.op_id), operation_ref(case, second), operation_ref(case, boundary)}
    response = client.get(export_path, query_string={**midnight_detail_query,
        "snapshot_ref": midnight_detail["meta"]["snapshot_ref"], "section": "resource-days", "format": "csv"})
    assert response.status_code == 200, response.get_json()
    exported = list(csv.DictReader(io.StringIO(response.data.decode("utf-8-sig"))))
    assert len(exported) == 1 and exported[0]["日期"] == DAY
    assert (exported[0]["时段起点（含）"], exported[0]["时段终点（不含）"]) == (
        midnight_query["range_start"], midnight_query["range_end"])
    assert float(exported[0]["所选可排总时长（小时）"]) == 6


def test_unknown_operator_capacity_and_known_zero_capacity_remain_distinct(candidate_case, monkeypatch):
    case = candidate_case
    add_pair(case)
    add_pair(case, "M3", "O3")
    case.conn.commit()
    calendar = CalendarService(case.conn)
    calendar._admin.operator_calendar_repo.upsert_many([
        calendar._admin._build_operator_calendar(
            operator_id="O1", date_value=DAY, shift_start="08:00", shift_end="16:00", efficiency=1,
            allow_normal="yes", allow_urgent="yes"
        ),
        calendar._admin._build_operator_calendar(
            operator_id="O2", date_value=DAY, shift_start="08:00", shift_end="16:00", efficiency=1,
            allow_normal="no", allow_urgent="no"
        ),
        calendar._admin._build_operator_calendar(
            operator_id="O3", date_value=DAY, shift_start="08:00", shift_end="16:00", efficiency=1,
            allow_normal="yes", allow_urgent="yes"
        ),
    ])
    case.conn.execute("UPDATE OperatorCalendar SET efficiency=NULL WHERE operator_id IN ('O1','O3') AND date=?", (DAY,))
    case.conn.commit()
    run_ref, refs = legacy_resource_archive(case, monkeypatch,
        edits={-1: lambda row: row.update(machine_id="M2", operator_id="O2"),
               -2: lambda row: row.update(machine_id="M3", operator_id="O3")})
    client, _ = api(case)
    query = {**WINDOW, "selected_candidate_ref": refs[0], "reference_candidate_ref": refs[0]}
    operators = resource_document(client, run_ref, **dict(query, resource_kind="operator"))
    missing = row_by_ref(operators, case.ref("operator", "O1"))["after"]
    assert missing["available_hours"] is missing["utilization_ratio"] is None
    assert missing["span_occupied_hours"] == missing["span_summed_hours"] == 0.75
    assert missing["unknown_days"] == 1 and missing["zero_capacity_task"] is None
    summary = operators["data"]["summary"]["after"]
    assert summary["known_capacity_count"] == 1 and summary["unknown_capacity_count"] == 2
    assert summary["available_hours"] is summary["utilization_ratio"] is None
    assert summary["zero_capacity_task_count"] is None and summary["known_zero_capacity_task_count"] == 0
    idle_unknown = row_by_ref(operators, case.ref("operator", "O3"))["after"]
    assert idle_unknown["available_hours"] is None and idle_unknown["high_pressure"] is False
    overview = decision(client, run_ref, **query)["data"]
    chosen = next(row for row in overview["candidates"] if row["candidate_ref"] == refs[0])
    for metric, field, expected in (("operator_available_hours", "available_hours", {case.ref("operator", "O2")}),
                                    ("operator_high_pressure_count", "high_pressure", {case.ref("operator", "O2"), case.ref("operator", "O3")}),
                                    ("operator_zero_capacity_task_count", "zero_capacity_task", {case.ref("operator", "O2"), case.ref("operator", "O3")})):
        value = chosen["metrics"][metric]
        assert set(value["covered_resource_refs"]) == expected
        assert value["known_count"] == len(expected) and value["unknown_count"] == 3 - len(expected)
        assert expected == {row["resource_ref"] for row in operators["data"]["items"] if row["after"][field] is not None}
    machines = resource_document(client, run_ref, **dict(query, resource_kind="machine"))
    assert row_by_ref(machines, case.ref("machine", "M1"))["after"]["available_hours"] == 8
    unknown = resource_document(client, run_ref, **dict(query, resource_kind="operator", pressure_filter="unknown"))
    assert {row["resource_ref"] for row in unknown["data"]["items"]} == {case.ref("operator", "O1"), case.ref("operator", "O3")}
    alternate = dict(query, selected_candidate_ref=refs[-1], resource_kind="operator", pressure_filter="zero_capacity")
    zero = resource_document(client, run_ref, **alternate)
    assert zero["data"]["items"] == []
    unfiltered = resource_document(client, run_ref, **dict(alternate, pressure_filter="all"))
    actual = row_by_ref(unfiltered, case.ref("operator", "O2"))["after"]
    assert actual["available_hours"] == actual["occupied_hours"] == 0
    assert actual["outside_calendar_hours"] == 0.75 and actual["unknown_days"] == 0
    assert actual["utilization_ratio"] is None and actual["zero_capacity_task"] is None
    assert actual["calendar_state"] == "complete" and actual["zero_capacity_reason"]["code"] == "actual_work_fragments_not_recorded"


def test_hour_pressure_tracks_display_page_and_exports_its_exact_filter_window(candidate_case, monkeypatch):
    case = candidate_case
    run_ref, refs = reassigned_run(case, monkeypatch=monkeypatch)
    client, _ = api(case)
    query = {"selected_candidate_ref": refs[0], "reference_candidate_ref": refs[0], "resource_kind": "machine", "grain": "hour",
             "range_start": DAY + "T00:00:00", "range_end": "2026-09-11T00:00:00", "heatmap_size": 7,
             "pressure_filter": "high_pressure", "pressure_threshold": 0.7}
    first = resource_document(client, run_ref, **query)
    assert first["data"]["items"] == []
    second = resource_document(client, run_ref, **dict(query, heatmap_page=2))
    data = second["data"]
    assert [row["resource_ref"] for row in data["items"]] == [case.ref("machine", "M1")]
    assert data["filters"]["display_start"] == DAY + "T07:00:00"
    assert data["filters"]["display_end"] == DAY + "T14:00:00"
    assert data["scope"]["range_start"] == query["range_start"] and data["scope"]["range_end"] == query["range_end"]
    row = data["items"][0]
    assert row["after"]["available_hours"] == 16 and row["after"]["occupied_hours"] == 0.75
    assert row["filter_summary"]["high_pressure_period_count"] == 1
    assert row["filter_summary"]["period_peak_utilization"] == 0.75
    assert row["filter_summary"]["daily_peak_utilization"] == pytest.approx(0.75 / 6)
    stale = client.get(BASE + "/runs/" + run_ref + "/decision/resources", query_string=dict(query, heatmap_page=2,
        snapshot_ref=first["meta"]["snapshot_ref"]))
    assert stale.status_code == 409
    effective = {key: value for key, value in data["filters"].items() if key != "section" and value is not None}
    export_query = {**query, **effective, "snapshot_ref": second["meta"]["snapshot_ref"], "section": "resources", "format": "csv"}
    response = client.get(BASE + "/runs/" + run_ref + "/decision/export", query_string=export_query)
    assert response.status_code == 200, response.get_json()
    exported = list(csv.DictReader(io.StringIO(response.data.decode("utf-8-sig"))))
    assert [row["资源永久编号"] for row in exported] == [case.ref("machine", "M1")]
    assert exported[0]["筛选时段起点（含）"] == data["filters"]["display_start"]
    assert exported[0]["筛选时段终点（不含）"] == data["filters"]["display_end"]
    assert exported[0]["范围起点（含）"] == query["range_start"]
    assert exported[0]["范围终点（不含）"] == query["range_end"]
    assert exported[0]["资源指标版本"] == "decision-resource-v4"
    assert float(exported[0]["筛选时段峰占用比例（%）"]) == 75
    assert int(exported[0]["筛选时段高负荷数"]) == 1


def test_full_candidate_view_exposes_uncovered_official_pressure(candidate_case):
    case = candidate_case
    add_pair(case)
    case.batch("B2", quantity=1)
    second_id = case.operation("B2", unit_hours=6, machine_id="M2", operator_id="O2")
    case.conn.commit()
    _, old_refs = compute(case, case.settings("B1", "B2"))
    adopter = service(case.conn)
    preview = adopter.preview(old_refs[0])
    assert preview["validation"]["can_adopt"] is True, preview
    adopted = adopter.adopt(old_refs[0], preview["write_context"]["write_token"], "resource-formal-isolated-adopt", INTENT)
    assert adopted["ok"] is True
    accepted = case.accept(key="resource-b1-only-run", settings=case.settings("B1"))
    result = WorkbenchRunWorker(case.conn).execute(accepted["run_ref"])
    run_ref, selected = result["run_ref"], result["candidates"][0]["candidate_ref"]
    client, _ = api(case)
    query = {**WINDOW, "comparison_mode": "admission_official", "selected_candidate_ref": selected}
    contribution = resource_document(client, run_ref, **dict(query, resource_view="selected_batches"))
    complete = resource_document(client, run_ref, **dict(query, resource_view="full_candidate"))
    assert contribution["data"]["scope"]["resource_refs"]["machine"] == [case.ref("machine", "M1")]
    assert set(complete["data"]["scope"]["resource_refs"]["machine"]) == {case.ref("machine", "M1"), case.ref("machine", "M2")}
    other = row_by_ref(complete, case.ref("machine", "M2"))
    assert other["before"]["occupied_hours"] == 6 and other["after"]["occupied_hours"] == 0
    coverage = complete["data"]["coverage"]["after"]
    assert coverage == contribution["data"]["coverage"]["after"]
    assert coverage["formal_state"] == "available" and coverage["formal_operation_count"] == 2
    assert coverage["covered_formal_operation_count"] == 1 and coverage["formal_coverage_complete"] is False
    assert coverage["missing_formal_operation_refs"] == [operation_ref(case, second_id)]
    assert {row["operation_ref"] for row in coverage["missing_formal_arrangements"]} == {operation_ref(case, second_id)}
    assert coverage["unplanned_operation_refs"] == coverage["deferred_operation_refs"] == []
    assert other["drill"]["reference"]["reference_view"] is True
    assert other["drill"]["reference"]["plan_ref"] == adopted["data"]["official_plan"]["plan_ref"]


def test_partial_worker_resource_load_keeps_ordinary_unplanned_demand_visible(candidate_case):
    case = candidate_case
    case.conn.execute("UPDATE Batches SET quantity=1,due_date=? WHERE batch_id='B1'", (DAY,))
    case.conn.execute("UPDATE BatchOperations SET unit_hours=5 WHERE id=?", (case.op_id,))
    case.batch("B2", quantity=1, due_date="2026-09-10")
    second_id = case.operation("B2", unit_hours=6)
    case.config(dispatch_mode="batch_order", sort_strategy="due_date_first")
    run_ref, refs = compute(case, case.settings("B1", "B2", end_date=DAY))
    client, _ = api(case)
    query = {**WINDOW, "selected_candidate_ref": refs[0], "reference_candidate_ref": refs[0]}
    document = resource_document(client, run_ref, **query)
    machine = row_by_ref(document, case.ref("machine", "M1"))["after"]
    assert machine["available_hours"] == 8 and machine["occupied_hours"] == machine["summed_load_hours"] == 5
    assert machine["utilization_ratio"] == 0.625
    coverage = document["data"]["coverage"]["after"]
    assert coverage["formal_state"] == "absent"
    assert coverage["unplanned_operation_refs"] == [operation_ref(case, second_id)]
    assert coverage["deferred_operation_refs"] == []
    full = resource_document(client, run_ref, **dict(query, resource_view="full_candidate"))
    assert full["data"]["coverage"]["after"] == coverage
    formal_query = {key: value for key, value in query.items() if key != "reference_candidate_ref"}
    formal_query["comparison_mode"] = "admission_official"
    formal = resource_document(client, run_ref, **formal_query)
    assert formal["data"]["baseline"]["state"] == "absent"
    assert set(formal["data"]["coverage"]["before"]["unplanned_operation_refs"]) == {
        operation_ref(case, case.op_id), operation_ref(case, second_id)}
    export_path = BASE + "/runs/" + run_ref + "/decision/export"
    response = client.get(export_path, query_string=dict(formal_query, snapshot_ref=formal["meta"]["snapshot_ref"],
                                                        section="resources", format="xlsx"))
    assert response.status_code == 200, response.get_json()
    assert int(response.headers["X-Workbench-Row-Count"]) == 2
    book = load_workbook(io.BytesIO(response.data), read_only=True, data_only=False)
    try:
        iterator = book["覆盖工序"].iter_rows(values_only=True)
        headers = next(iterator)
        exported = [dict(zip(headers, row)) for row in iterator]
        assert len(exported) == 1
        row = exported[0]
        assert (row["比较侧"], row["分类"], row["工序永久编号"]) == ("所选", "其他未安排", operation_ref(case, second_id))
        assert row["排产编号"] == run_ref and row["所选方案编号"] == refs[0]
        assert row["批次永久编号"] == case.ref("batch", "B2") and row["批次编号"] == "B2"
        iterator = book["资源影响"].iter_rows()
        headers = [cell.value for cell in next(iterator)]
        resources = [dict(zip(headers, row)) for row in iterator]
        assert len(resources) == 1
        resource = resources[0]
        assert resource["所选其他未安排工序数"].value == 1 and resource["所选其他未安排工序数"].data_type == "n"
        assert type(resource["所选其他未安排工序数"].value) is int
        assert resource["范围起点（含）"].value == formal_query["range_start"] and resource["范围终点（不含）"].value == formal_query["range_end"]
        assert resource["所选方案编号"].value == refs[0]
        scope = {row[0]: row[1] if len(row) > 1 else None for row in book["比较范围"].iter_rows(values_only=True)}
        assert scope["比较模式"] == "排产时的正式计划" and scope["查看基准状态"] == "排产时无正式计划"
        assert scope["查看基准编号"] == formal["data"]["baseline"]["baseline_ref"] and scope["所选方案编号"] == refs[0]
    finally:
        book.close()


def test_resource_paging_heatmap_snapshot_and_full_filtered_exports(candidate_case, monkeypatch):
    import json
    from collections import Counter

    from core.services.workbench.run import decision_resources as resource_runtime
    from core.services.workbench.run.decision_resources import DecisionResources

    case = candidate_case
    batch_codes = ["B1"]
    for index in range(2, 26):
        machine, operator, batch = f"M{index:02d}", f"O{index:02d}", f"B{index:02d}"
        add_pair(case, machine, operator, r"设备'甲\N" + str(index))
        case.batch(batch)
        case.operation(batch, machine_id=machine, operator_id=operator)
        batch_codes.append(batch)
    case.conn.commit()
    run_ref, refs = legacy_resource_archive(case, monkeypatch, case.settings(*batch_codes))
    client, _ = api(case)
    query = {"selected_candidate_ref": refs[0], "reference_candidate_ref": refs[0], "resource_kind": "machine",
             "resource_view": "selected_batches", "range_start": DAY + "T00:00:00", "range_end": "2026-10-04T00:00:00",
             "sort": "label", "order": "asc"}
    calendar_refs, measurement_refs, calculation_keys, full_cells, drill_calls = [], [], [], [], []
    parsed_calendar_times, day_facts_calls, day_statistics_calls = [], [], []
    single_windows, bulk_windows = [], []
    measurements = []
    original_calendar = resource_projection._ResourceCalendar.__init__
    original_measurement = resource_projection._ResourceMeasurement.__init__
    original_values = resource_projection._ResourceMeasurement._values
    original_cell = resource_projection._ResourceMeasurement.cell
    original_drill = DecisionResources.drill
    original_parse = resource_runtime.instant
    original_facts = ResourceTaskIndex.period_facts
    original_statistics = resource_projection._ResourceMeasurement.period_statistics
    original_single = resource_projection.ResourceUtilizationMetrics.window
    original_bulk = resource_projection.ResourceUtilizationMetrics.windows_at_boundaries

    def record_calendar(value, calendar, low, high, **kwargs):
        calendar_refs.append(calendar["resource_ref"])
        original_calendar(value, calendar, low, high, **kwargs)

    def record_parse(value):
        parsed_calendar_times.append(value)
        return original_parse(value)

    def record_facts(index, resource, calendar, **kwargs):
        day_facts_calls.append(resource["ref"])
        return original_facts(index, resource, calendar, **kwargs)

    def record_statistics(value, low, high, calendar_days, grain, group, *, whole=False):
        day_statistics_calls.append((value.resource["ref"], low, high))
        return original_statistics(value, low, high, calendar_days, grain, group, whole=whole)

    def record_single(metrics, low, high, **kwargs):
        single_windows.append((id(metrics), low, high))
        return original_single(metrics, low, high, **kwargs)

    def record_bulk(metrics, boundaries, **kwargs):
        for index, raw in enumerate(original_bulk(metrics, boundaries, **kwargs)):
            bulk_windows.append((id(metrics), boundaries[index], boundaries[index + 1]))
            yield raw

    def record_measurement(value, resource, *args, **kwargs):
        measurement_refs.append(resource["ref"])
        original_measurement(value, resource, *args, **kwargs)
        measurements.append(value)

    def record_values(value, low, high, calendar_days, *, potential_work=None, prepared=None):
        result = original_values(value, low, high, calendar_days, potential_work=potential_work, prepared=prepared)
        effective_work = (result.raw["span_occupied_hours"] > 0 or not value.work_complete
                          if potential_work is None else potential_work)
        calculation_keys.append((value.resource["ref"], low, high, value._day_key(calendar_days), effective_work))
        return result

    def record_cell(value, low, high, calendar_days, *, whole=False):
        full_cells.append((value.resource["ref"], whole))
        return original_cell(value, low, high, calendar_days, whole=whole)

    def record_drill(engine, key, resource, **kwargs):
        drill_calls.append((key, resource["ref"]))
        return original_drill(engine, key, resource, **kwargs)

    with monkeypatch.context() as scoped:
        scoped.setattr(resource_projection._ResourceCalendar, "__init__", record_calendar)
        scoped.setattr(resource_projection._ResourceMeasurement, "__init__", record_measurement)
        scoped.setattr(resource_projection._ResourceMeasurement, "_values", record_values)
        scoped.setattr(resource_projection._ResourceMeasurement, "cell", record_cell)
        scoped.setattr(DecisionResources, "drill", record_drill)
        scoped.setattr(resource_runtime, "instant", record_parse)
        scoped.setattr(ResourceTaskIndex, "period_facts", record_facts)
        scoped.setattr(resource_projection._ResourceMeasurement, "period_statistics", record_statistics)
        scoped.setattr(resource_projection.ResourceUtilizationMetrics, "window", record_single)
        scoped.setattr(resource_projection.ResourceUtilizationMetrics, "windows_at_boundaries", record_bulk)
        first = resource_document(client, run_ref, **dict(query, heatmap_size=7))
    cohort_refs = set(first["data"]["scope"]["resource_refs"]["machine"])
    assert len(cohort_refs) == 25
    assert Counter(calendar_refs) == Counter(measurement_refs) == Counter({ref: 1 for ref in cohort_refs})
    assert len(calculation_keys) == 25 * (1 + 25) + 20 * 7 and len(set(calculation_keys)) == 650
    assert sum(len(key[3]) == 1 for key in calculation_keys) == 25 * 25 + 20 * 7
    assert len(single_windows) == 25 + 20 * 7 and len(bulk_windows) == 25 * 25
    assert len(single_windows + bulk_windows) == 790 and len(set(single_windows + bulk_windows)) == 650
    assert Counter(Counter(calculation_keys).values()) == {1: 510, 2: 140}
    duplicates = {key for key, count in Counter(calculation_keys).items() if count == 2}
    assert {(key[0], key[1], key[2]) for key in duplicates} == {
        (row["resource_ref"], resource_projection.instant(cell["start"]), resource_projection.instant(cell["end"]))
        for row in first["data"]["heatmap"]["items"] for cell in row["cells"]}
    retained_contexts = sum(len(value._value_cache) for value in measurements)
    retained_variants = sum(sum(len(variants) for variants in value._value_cache.values()) for value in measurements)
    assert retained_contexts == retained_variants == 25 + 20 * 7
    assert all(not value.retain_statistics_values and all(value._value_cache.values()) for value in measurements)
    assert Counter(whole for _ref, whole in full_cells) == {True: 25, False: 20 * 7}
    assert len(parsed_calendar_times) == len(set(parsed_calendar_times))
    assert day_facts_calls == [] and len(day_statistics_calls) == 20 * 7
    output_rows = (first["data"]["items"] + first["data"]["rankings"]["top_pressure"] +
                   first["data"]["rankings"]["top_increase"])
    output_refs = {row["resource_ref"] for row in output_rows}
    assert Counter(drill_calls) == Counter({(key, ref): 1 for key in (refs[0], "reference") for ref in output_refs})
    print("U5_RESOURCE_FACT_COUNTS=" + json.dumps({
        "resources": len(cohort_refs), "natural_days": 25, "visible_resources": 20, "visible_buckets": 7,
        "normalized_calendars": len(calendar_refs), "measurements": len(measurement_refs),
        "numeric_calculations": len(calculation_keys), "unique_numeric_keys": len(set(calculation_keys)),
        "complete_window_rows": 25, "complete_period_rows": len(full_cells) - 25,
        "drill_resources": len(output_refs), "drill_locations": len(drill_calls),
        "calendar_time_parses": len(parsed_calendar_times), "unique_calendar_time_strings": len(set(parsed_calendar_times)),
        "single_metrics_windows": len(single_windows), "bulk_metrics_windows": len(bulk_windows),
        "retained_value_contexts": retained_contexts, "retained_value_variants": retained_variants,
        "intentional_visible_recalculations": len(duplicates),
        "full_day_period_facts_calls": len(day_facts_calls), "visible_period_statistics": len(day_statistics_calls),
        "public_period_sides": 2 * 20 * 7}, sort_keys=True))
    token = first["meta"]["snapshot_ref"]
    page = first["data"]["page"]
    assert (page["number"], page["size"], page["total"], page["has_more"]) == (1, 20, 25, True)
    second = resource_document(client, run_ref, **dict(query, page=2, snapshot_ref=token, heatmap_size=7))
    rows = first["data"]["items"] + second["data"]["items"]
    assert len(rows) == len({row["resource_ref"] for row in rows}) == 25
    assert rows == sorted(rows, key=lambda row: (row["label"], row["resource_ref"]))
    assert len(first["data"]["items"]) == 20 and len(second["data"]["items"]) == 5
    for row in output_rows:
        for key in ("selected", "reference"):
            assert row["drill"][key]["resource_ref"] == row["resource_ref"]
            assert row["drill"][key]["candidate_ref"] == refs[0]
    other_dates = resource_document(client, run_ref, **dict(query, heatmap_page=2, heatmap_size=7, snapshot_ref=token))
    assert other_dates["data"]["summary"] == first["data"]["summary"]
    assert {row["date"] for row in other_dates["data"]["heatmap"]["buckets"]}.isdisjoint(
        row["date"] for row in first["data"]["heatmap"]["buckets"])
    assert len(first["data"]["heatmap"]["items"]) == 20
    assert all(len(row["cells"]) == 7 for row in first["data"]["heatmap"]["items"])
    path = BASE + "/runs/" + run_ref + "/decision/resources"
    for change in ({"range_end": "2026-10-03T00:00:00"}, {"resource_view": "full_candidate"}, {"pressure_filter": "high_pressure"}):
        response = client.get(path, query_string=dict(query, snapshot_ref=token, **change))
        assert response.status_code == 409 and response.get_json()["error"]["code"] == "snapshot_stale"
    export_path = BASE + "/runs/" + run_ref + "/decision/export"
    for fmt in ("csv", "xlsx"):
        response = client.get(export_path, query_string=dict(query, snapshot_ref=token, section="resources", format=fmt))
        assert response.status_code == 200, response.get_json()
        assert int(response.headers["X-Workbench-Row-Count"]) == 25
        if fmt == "csv":
            reader = csv.DictReader(io.StringIO(response.data.decode("utf-8-sig")))
            assert len(reader.fieldnames) == len(set(reader.fieldnames))
            exported = list(reader)
        else:
            book = load_workbook(io.BytesIO(response.data), read_only=True, data_only=False)
            try:
                assert book.sheetnames == ["资源影响", "覆盖工序", "指标计算方式", "比较范围"]
                coverage_rows = list(book["覆盖工序"].iter_rows(values_only=True))
                assert len(coverage_rows) == 1 and "工序永久编号" in coverage_rows[0]
                iterator = book["资源影响"].iter_rows()
                headers = [cell.value for cell in next(iterator)]
                assert len(headers) == len(set(headers))
                cells = [dict(zip(headers, row)) for row in iterator]
                exported = [{name: cell.value for name, cell in row.items()} for row in cells]
                assert all(row["所选占用比例（%）"].data_type == "n" for row in cells)
                for header in ("所选缺失正式工序数", "所选其他未安排工序数", "所选缺料未安排工序数"):
                    assert all(row[header].data_type == "n" and type(row[header].value) is int and row[header].value == 0 for row in cells)
                definitions = book["指标计算方式"].iter_rows(values_only=True)
                definition_headers = next(definitions)
                assert {dict(zip(definition_headers, row))["常用指标"] for row in definitions} == {"是", "否"}
            finally:
                book.close()
        assert [row["资源永久编号"] for row in exported] == [row["resource_ref"] for row in rows]
        for actual, expected in zip(exported, rows):
            assert float(actual["所选占用比例（%）"]) == pytest.approx(expected["after"]["utilization_ratio"] * 100)
            assert float(actual["所选日峰占用比例（%）"]) == pytest.approx(expected["after"]["daily_peak_utilization"] * 100)
            assert "所选时段占用比例（%）" not in actual
            assert actual["资源名称"] == expected["label"]
    combined = client.get(export_path, query_string=dict(query, snapshot_ref=token, section="all", filter_section="resources", format="xlsx"))
    assert combined.status_code == 200, combined.get_json()
    book = load_workbook(io.BytesIO(combined.data), read_only=True, data_only=True)
    try:
        assert book.sheetnames == ["方案总览", "批次影响", "工序变化", "资源影响", "覆盖工序", "准备需求", "工种切换",
                                  "推荐说明", "工序安排", "剩余需求", "候选计算", "进步事件", "搜索阶段", "代表尝试", "指标计算方式", "比较范围"]
        assert len(list(book["资源影响"].iter_rows(values_only=True))) - 1 == 25
    finally:
        book.close()
    actual, _scope, _state = read_decision_export(case.conn, run_ref, "resources", **query)
    assert len(actual["export_resources"][0]["items"]) == 25
    assert all({"before", "after", "filter_summary", "drill"}.issubset(row)
               for row in actual["export_resources"][0]["items"])
    boundary = deepcopy(actual)
    coverage = boundary["export_resources"][0]["coverage"]
    categories = (("missing_formal_operation_refs", "缺失正式"), ("unplanned_operation_refs", "其他未安排"),
                  ("deferred_operation_refs", "缺料未安排"))
    coverage_refs = {field: [f"{number:048x}" for number in range(index * 800 + 1, (index + 1) * 800 + 1)]
                     for index, (field, _category) in enumerate(categories)}
    assert all(len("、".join(values)) > 32767 for values in coverage_refs.values())
    coverage["after"].update(coverage_refs, formal_operation_count=800, covered_formal_operation_count=0, formal_coverage_complete=False)
    source_query = {key: value for key, value in query.items() if key not in ("resource_kind", "sort", "order")}
    store = CandidateStore(case.conn)
    with store.snapshot(), read_evidence_scope(case.conn):
        source = DecisionSource(store, run_ref, **source_query)
        boundary["export_coverage_operations"] = list(coverage_operation_rows(source, coverage))
        engine = source.resources
        resource = engine.resource(case.ref("machine", "M1"), "machine")
        calendar = engine.resource_calendar(resource)
        parsed_calendar_times.clear()
        with monkeypatch.context() as scoped:
            scoped.setattr(resource_runtime, "instant", record_parse)
            clock = DAY + "T08:00:00"
            precise_clock = clock + ".000000"
            assert engine._calendar_time(clock) == engine._calendar_time(precise_clock)
            assert engine._calendar_time(precise_clock) is engine._calendar_times[precise_clock]
            assert parsed_calendar_times == [precise_clock]
            for _attempt in range(2):
                with pytest.raises(ValueError, match="Invalid local plan/calendar time"):
                    engine._calendar_time("invalid-time")
            assert parsed_calendar_times == [precise_clock, "invalid-time", "invalid-time"]
            assert "invalid-time" not in engine._calendar_times
        detailed = engine.periods(resource, grain="hour", start=DAY + "T08:00:00", end=DAY + "T09:00:00", details=True)[0]
        assert detailed["before"]["tasks"] is not detailed["after"]["tasks"]
        assert detailed["before"]["tasks"][0]["drill"] is not detailed["after"]["tasks"][0]["drill"]
        assert detailed["before"]["tasks"][0]["drill"]["candidate_ref"] == refs[0]
        assert detailed["after"]["tasks"][0]["drill"]["candidate_ref"] == refs[0]
        detailed["after"]["tasks"][0]["drill"]["test_location"] = "selected"
        assert "test_location" not in detailed["before"]["tasks"][0]["drill"]
        original_index = engine.index(refs[0])
        engine.configure(0.9)
        assert engine.index(refs[0]) is original_index and engine.resource_calendar(resource) is calendar
        assert detailed["after"]["high_pressure"] is False
        engine.configure(0.7)
        assert engine.index(refs[0]) is not original_index and engine.resource_calendar(resource) is calendar
        refreshed = engine.periods(resource, grain="hour", start=DAY + "T08:00:00", end=DAY + "T09:00:00")[0]
        assert refreshed["after"]["occupied_hours"] == detailed["after"]["occupied_hours"] == 0.75
        assert refreshed["after"]["pressure_threshold"] == 0.7 and refreshed["after"]["high_pressure"] is True
    download = write_run_decision_export(boundary, first["meta"], "xlsx", "resources")
    assert download.row_count == 25 + 3 * 800
    book = load_workbook(io.BytesIO(download.content), read_only=True, data_only=False)
    try:
        assert book.sheetnames == ["资源影响", "覆盖工序", "指标计算方式", "比较范围"]
        iterator = book["资源影响"].iter_rows()
        headers = [cell.value for cell in next(iterator)]
        resource_cells = [dict(zip(headers, row)) for row in iterator]
        assert len(resource_cells) == 25
        assert resource_cells[0]["正式归档工序数"].value == 800
        assert resource_cells[0]["所选覆盖正式工序数"].value == 0 and resource_cells[0]["所选完整覆盖正式范围"].value == "否"
        for header in ("所选缺失正式工序数", "所选其他未安排工序数", "所选缺料未安排工序数"):
            assert all(row[header].data_type == "n" and type(row[header].value) is int and row[header].value == 800 for row in resource_cells)
        iterator = book["覆盖工序"].iter_rows()
        headers = [cell.value for cell in next(iterator)]
        coverage_cells = [dict(zip(headers, row)) for row in iterator]
        assert len(coverage_cells) == 3 * 800
        assert {row["比较侧"].value for row in coverage_cells} == {"所选"}
        for field, category in categories:
            group = [row for row in coverage_cells if row["分类"].value == category]
            assert [row["工序永久编号"].value for row in group] == coverage_refs[field]
            assert all(row["工序永久编号"].data_type == "s" for row in group)
        assert coverage_cells[-1]["工序永久编号"].value == coverage_refs["deferred_operation_refs"][-1]
    finally:
        book.close()
    day_query = {key: value for key, value in query.items() if key not in ("sort", "order")}
    day_query["resource_ref"] = case.ref("machine", "M1")
    calls = []
    original_window = ResourceTaskIndex.window

    def record_window(index, resource, calendar):
        calls.append(resource["ref"])
        return original_window(index, resource, calendar)

    with monkeypatch.context() as scoped:
        scoped.setattr(ResourceTaskIndex, "window", record_window)
        initial_days = resource_document(client, run_ref, "resource-days", **day_query)
    assert calls == [day_query["resource_ref"]]
    day_token = initial_days["meta"]["snapshot_ref"]
    next_days = resource_document(client, run_ref, "resource-days", **dict(day_query, page=2, snapshot_ref=day_token))
    days = initial_days["data"]["items"] + next_days["data"]["items"]
    assert len(days) == 25 and [row["date"] for row in days] == sorted(row["date"] for row in days)
    wide_days = resource_document(client, run_ref, "resource-days", **dict(day_query, size=50, snapshot_ref=day_token))
    assert wide_days["data"]["items"] == days
    response = client.get(export_path, query_string=dict(day_query, snapshot_ref=day_token, section="resource-days", format="csv"))
    assert response.status_code == 200, response.get_json()
    exported_days = list(csv.DictReader(io.StringIO(response.data.decode("utf-8-sig"))))
    assert int(response.headers["X-Workbench-Row-Count"]) == len(days)
    assert [row["日期"] for row in exported_days] == [row["date"] for row in days]
    assert [row["时段起点（含）"] for row in exported_days] == [row["start"] for row in days]
    assert [row["时段终点（不含）"] for row in exported_days] == [row["end"] for row in days]
    assert {row["范围起点（含）"] for row in exported_days} == {day_query["range_start"]}
    assert {row["范围终点（不含）"] for row in exported_days} == {day_query["range_end"]}
    daily_workbook = client.get(export_path, query_string=dict(day_query, snapshot_ref=day_token, section="resource-days", format="xlsx"))
    assert daily_workbook.status_code == 200, daily_workbook.get_json()
    assert int(daily_workbook.headers["X-Workbench-Row-Count"]) == len(days)
    book = load_workbook(io.BytesIO(daily_workbook.data), read_only=True, data_only=False)
    try:
        assert book.sheetnames == ["资源逐日", "资源冲突", "覆盖工序", "指标计算方式", "比较范围"]
        assert len(list(book["覆盖工序"].iter_rows(values_only=True))) == 1
        conflicts = list(book["资源冲突"].iter_rows(values_only=True))
        assert len(conflicts) == 1 and "安排永久编号" in conflicts[0]
        iterator = book["资源逐日"].iter_rows()
        headers = [cell.value for cell in next(iterator)]
        day_cells = [dict(zip(headers, row)) for row in iterator]
        for header in ("基准冲突工序数", "所选冲突工序数"):
            assert all(row[header].data_type == "n" and type(row[header].value) is int and row[header].value == 0 for row in day_cells)
    finally:
        book.close()
    actual_days, _scope, _state = read_decision_export(case.conn, run_ref, "resource-days", **day_query)
    conflict_boundary = deepcopy(actual_days)
    day_document = conflict_boundary["export_resource_days"]
    chosen_day = day_document["items"][0]
    operation_refs = coverage_refs["missing_formal_operation_refs"][:3]
    row_refs = [f"{number:048x}" for number in range(3001, 3004)]
    conflict = {"start": DAY + "T08:15:00", "end": DAY + "T08:30:00", "operation_refs": operation_refs, "row_refs": row_refs}
    for side in ("before", "after"):
        original_task = chosen_day[side]["tasks"][0]
        tasks = [{**original_task, "operation_ref": operation_ref, "row_ref": row_ref}
                 for operation_ref, row_ref in zip(operation_refs, row_refs)]
        chosen_day[side].update(tasks=tasks, conflicts=[deepcopy(conflict)], conflict=True)
    download = write_run_decision_export(conflict_boundary, initial_days["meta"], "xlsx", "resource-days")
    assert download.row_count == len(days) + 2 * 3
    book = load_workbook(io.BytesIO(download.content), read_only=True, data_only=False)
    try:
        assert book.sheetnames == ["资源逐日", "资源冲突", "覆盖工序", "指标计算方式", "比较范围"]
        iterator = book["资源逐日"].iter_rows()
        headers = [cell.value for cell in next(iterator)]
        first_day = dict(zip(headers, next(iterator)))
        for header in ("基准冲突工序数", "所选冲突工序数"):
            assert first_day[header].value == 3 and type(first_day[header].value) is int and first_day[header].data_type == "n"
        assert first_day["基准冲突段数"].value == first_day["所选冲突段数"].value == 1
        iterator = book["资源冲突"].iter_rows()
        headers = [cell.value for cell in next(iterator)]
        conflict_cells = [dict(zip(headers, row)) for row in iterator]
        assert len(conflict_cells) == 6
        assert {row["比较侧"].value for row in conflict_cells} == {"基准", "所选"}
        for side_label in ("基准", "所选"):
            group = [row for row in conflict_cells if row["比较侧"].value == side_label]
            assert [(row["工序永久编号"].value, row["安排永久编号"].value) for row in group] == list(zip(operation_refs, row_refs))
            assert all(row["工序永久编号"].data_type == row["安排永久编号"].data_type == "s" for row in group)
            assert all(row["冲突起点（含）"].value == conflict["start"] and row["冲突终点（不含）"].value == conflict["end"] for row in group)
    finally:
        book.close()
    combined_days = client.get(export_path, query_string=dict(day_query, snapshot_ref=day_token, section="all",
                                                             filter_section="resource-days", format="xlsx"))
    assert combined_days.status_code == 200, combined_days.get_json()
    book = load_workbook(io.BytesIO(combined_days.data), read_only=True, data_only=True)
    try:
        assert book.sheetnames == ["方案总览", "批次影响", "工序变化", "资源影响", "资源逐日", "资源冲突", "覆盖工序", "准备需求", "工种切换",
                                  "推荐说明", "工序安排", "剩余需求", "候选计算", "进步事件", "搜索阶段", "代表尝试", "指标计算方式", "比较范围"]
        assert len(list(book["资源影响"].iter_rows(values_only=True))) - 1 == 1
        iterator = book["资源逐日"].iter_rows(values_only=True)
        headers = next(iterator)
        exported_days = [dict(zip(headers, row)) for row in iterator]
        assert [row["日期"] for row in exported_days] == [row["date"] for row in days]
        assert {row["资源永久编号"] for row in exported_days} == {day_query["resource_ref"]}
        assert exported_days[0]["所选占用比例（%）"] == pytest.approx(days[0]["after"]["utilization_ratio"] * 100)
    finally:
        book.close()


def test_resource_workspace_shares_source_restores_view_and_keeps_resource_failures(candidate_case, monkeypatch):
    from core.models.workbench_command import WorkbenchCommandRejected
    from core.services.workbench.run import decision_resource_workspace as workspace
    from tests.workbench.test_run_decision_api import decision_summary
    from web.routes.workbench.read_context import bind_read_snapshot

    case = candidate_case
    add_pair(case)
    case.conn.commit()
    run_ref, refs = compute(case)
    client, statements = api(case)
    path = BASE + "/runs/" + run_ref + "/decision/resource-workspace"
    query = {**WINDOW, "selected_candidate_ref": refs[-1], "reference_candidate_ref": refs[0]}
    constructed, summary_sources, resource_sources = [], [], []
    original_source, original_summary, original_resource = workspace.DecisionSource, workspace.decision_summary, workspace.resource_details

    def source_once(*args, **kwargs):
        source = original_source(*args, **kwargs)
        constructed.append(source)
        return source

    def summary_from(source):
        summary_sources.append(source)
        return original_summary(source)

    def resources_from(source, *args, **kwargs):
        resource_sources.append(source)
        return original_resource(source, *args, **kwargs)

    with monkeypatch.context() as scoped:
        scoped.setattr(workspace, "DecisionSource", source_once)
        scoped.setattr(workspace, "decision_summary", summary_from)
        scoped.setattr(workspace, "resource_details", resources_from)
        combined = client.get(path, query_string=query)
    assert combined.status_code == 200 and combined.headers["Cache-Control"] == "no-store"
    result = combined.get_json()
    assert constructed == summary_sources == resource_sources and len(constructed) == 1
    resource = result["resource_result"]
    assert result["ok"] is resource["ok"] is True
    assert result["meta"]["as_of"] == resource["meta"]["as_of"]
    assert result["meta"]["snapshot_ref"] != resource["meta"]["snapshot_ref"]
    assert "snapshot_ref" not in result["resource_query"]
    independent_summary = decision_summary(client, run_ref, **query)
    independent_resources = resource_document(client, run_ref, **query, heatmap_size=7)
    assert result["data"] == independent_summary["data"]
    assert resource["data"] == independent_resources["data"]
    assert resource["schema_version"] == independent_resources["schema_version"]
    assert resource["warnings"] == independent_resources["warnings"]
    threshold = client.get(path, query_string={**query, "pressure_threshold": 0.5}).get_json()
    assert threshold["data"] == independent_summary["data"]
    assert threshold["resource_result"]["data"] == resource_document(
        client, run_ref, **query, pressure_threshold=0.5, heatmap_size=7)["data"]

    no_selected = client.get(path, query_string={**WINDOW, "reference_candidate_ref": refs[0]}).get_json()
    initial_summary = decision_summary(client, run_ref, **WINDOW, reference_candidate_ref=refs[0])["data"]
    assert no_selected["data"] == initial_summary
    assert no_selected["resource_result"]["data"]["candidate"]["candidate_ref"] == initial_summary["scope"]["selected_candidate_ref"]
    token = resource["meta"]["snapshot_ref"]
    paging = client.get(path, query_string={**query, "resource_snapshot_ref": token, "page": 2}).get_json()
    assert paging["ok"] and paging["resource_result"]["ok"]
    assert paging["resource_query"]["snapshot_ref"] == token
    assert paging["resource_result"]["meta"]["snapshot_ref"] == token
    assert paging["resource_result"]["data"]["page"]["number"] == 2
    assert paging["meta"]["as_of"] == paging["resource_result"]["meta"]["as_of"]
    store = CandidateStore(case.conn)
    with store.snapshot(), read_evidence_scope(case.conn):
        source = original_source(store, run_ref, **query)
        state, scope = source.state(), {"kind": "run-decision", **source.scope}
    old_summary = bind_read_snapshot(scope, state, as_of=DAY + "T08:01:00")
    old_resource = bind_read_snapshot({**scope, "filters": resource["data"]["filters"]}, state, as_of=DAY + "T08:02:00")
    only_summary = client.get(path, query_string={**query, "snapshot_ref": old_summary["snapshot_ref"]}).get_json()
    assert only_summary["ok"] and only_summary["resource_result"]["ok"]
    assert only_summary["meta"]["as_of"] == only_summary["resource_result"]["meta"]["as_of"] == old_summary["as_of"]
    both_old = client.get(path, query_string={**query, "snapshot_ref": old_summary["snapshot_ref"],
                                             "resource_snapshot_ref": old_resource["snapshot_ref"]}).get_json()
    assert both_old["ok"] and both_old["resource_result"]["ok"]
    assert both_old["meta"]["snapshot_ref"] == old_summary["snapshot_ref"]
    assert both_old["resource_result"]["meta"]["snapshot_ref"] == old_resource["snapshot_ref"]
    assert both_old["meta"]["as_of"] == old_summary["as_of"]
    assert both_old["resource_result"]["meta"]["as_of"] == old_resource["as_of"]
    export = client.get(BASE + "/runs/" + run_ref + "/decision/export",
        query_string={**query, "snapshot_ref": token, "heatmap_size": 7, "section": "resources", "format": "csv"})
    assert export.status_code == 200, export.get_json()
    assert export.headers["X-Workbench-Snapshot-Ref"] == token
    exported_resources = list(csv.DictReader(io.StringIO(export.data.decode("utf-8-sig"))))
    assert {row["资源永久编号"] for row in exported_resources} == set(resource["data"]["scope"]["resource_refs"]["machine"])

    invalid_ref = "e" * 48
    restored = client.get(path, query_string={**query, "resource_ref": invalid_ref, "resource_snapshot_ref": "expired-view", "page": 9})
    assert restored.status_code == 200
    restored = restored.get_json()
    assert restored["resource_result"]["ok"]
    assert "resource_ref" not in restored["resource_query"] and "snapshot_ref" not in restored["resource_query"]
    assert restored["resource_query"]["page"] == 1
    assert client.get(BASE + "/runs/" + run_ref + "/decision/resources", query_string={**query, "resource_ref": invalid_ref}).status_code == 404
    display = {"grain": "hour", "display_start": DAY + "T07:00:00", "display_end": DAY + "T17:00:00",
               "resource_snapshot_ref": "expired-view", "heatmap_page": 9}
    restored = client.get(path, query_string={**query, **display}).get_json()
    assert restored["resource_result"]["ok"] and restored["resource_query"]["grain"] == "hour"
    assert not {"display_start", "display_end", "snapshot_ref"}.intersection(restored["resource_query"])
    assert restored["resource_query"]["heatmap_page"] == 1
    ordinary_display = {key: value for key, value in display.items() if key != "resource_snapshot_ref"}
    assert client.get(BASE + "/runs/" + run_ref + "/decision/resources", query_string={**query, **ordinary_display}).status_code == 422
    for invalid in ({"pressure_threshold": 0}, {"resource_ref": "malformed"}, {"page": "01"},
                    {"heatmap_size": 15}, {"display_start": DAY + "T09:00:00"},
                    {"grain": "hour", "display_start": "invalid-time", "display_end": "invalid-time"}, {"unknown": "field"}):
        rejected = client.get(path, query_string={**query, **invalid})
        ordinary = client.get(BASE + "/runs/" + run_ref + "/decision/resources", query_string={**query, **invalid})
        assert rejected.status_code == ordinary.status_code
        assert rejected.get_json()["error"]["code"] == ordinary.get_json()["error"]["code"] == "invalid_input"
        assert rejected.get_json()["error"]["message"] == ordinary.get_json()["error"]["message"]
    duplicate = client.get(path, query_string=list(query.items()) + [("grain", "day"), ("grain", "week")])
    assert duplicate.status_code == 422
    stale = client.get(path, query_string={**query, "resource_snapshot_ref": "expired-view"}).get_json()
    assert stale["ok"] and stale["data"] == result["data"]
    assert not stale["resource_result"]["ok"]
    assert stale["resource_result"]["status"] == 409 and stale["resource_result"]["error"]["code"] == "snapshot_stale"
    stale_summary = client.get(path, query_string={**query, "snapshot_ref": "expired-summary"})
    assert stale_summary.status_code == 409 and not stale_summary.get_json()["ok"]

    def known_failure(*_args, **_kwargs):
        raise WorkbenchCommandRejected("resource_evidence_invalid", "原资源读取失败", 500)

    with monkeypatch.context() as scoped:
        scoped.setattr(workspace, "resource_details", known_failure)
        known = client.get(path, query_string=query)
    assert known.status_code == 200 and known.get_json()["data"] == result["data"]
    failure = known.get_json()["resource_result"]
    assert failure["status"] == 500 and failure["error"]["code"] == "resource_evidence_invalid"
    assert failure["error"]["message"] == "原资源读取失败" and failure["error"]["retryable"] is True

    def unexpected(*_args, **_kwargs):
        raise RuntimeError("unanticipated resource reader failure")

    with monkeypatch.context() as scoped:
        scoped.setattr(workspace, "resource_details", unexpected)
        unknown = client.get(path, query_string=query)
    assert unknown.status_code == 500 and unknown.get_json()["error"]["code"] == "storage_failure"
    assert not unknown.get_json()["ok"] and "resource_result" not in unknown.get_json()
    pending = case.accept(key="resource-workspace-pending")
    no_result = client.get(BASE + "/runs/" + pending["run_ref"] + "/decision/resource-workspace").get_json()
    assert no_result["ok"] and no_result["data"]["candidates"] == [] and no_result["resource_result"] is None
    assert no_result["resource_query"]["page"] == 1
    # All client connections were query_only; writes above belong to the test's
    # explicit saved arrangement and pending-run setup, never to these readers.
    assert not any(sql.lstrip().upper().startswith(("INSERT", "UPDATE", "DELETE", "REPLACE")) for sql in statements)
