"""Offline CSV/XLSX import/export with strict types and atomic preview confirmation."""

import csv
from dataclasses import replace
from datetime import datetime
from io import BytesIO, StringIO
from unittest.mock import patch

import openpyxl
import pytest

from core.errors import ValidationError
from core.infrastructure.transaction import TransactionManager
from core.models.workbench_command import WorkbenchCommandRejected, WorkbenchCommandUncertain, canonical_json
from core.models.workbench_material_file import HEADERS
from core.services.material.material_service import MaterialService
from core.services.workbench import material_file_codec
from core.services.workbench.commands import WorkbenchCommandService
from core.services.workbench.material_file_codec import check_export_capacity, write_material_file
from core.services.workbench.material_files import WorkbenchMaterialFileService
from core.services.workbench.materials import WorkbenchMaterialService
from tests.workbench.identity_metadata_support import business_snapshot, connect_temp, copy_to_temp, table_rows
from tests.workbench.material_file_support import (
    KEY,
    OnePassRows,
    codec_rows,
    confirm_import,
    exercise_scale_export,
    export_file,
    file_bytes,
    seed_many,
    small_dimensions,
    verify_download,
)
from tests.workbench.material_support import identity_for, material_database, material_row, stored_state


@pytest.mark.parametrize("fmt", ("csv", "xlsx"))
def test_upsert_full_preview_and_atomic_confirm_preserve_hidden_and_batch_columns(material_conn, fmt):
    conn = material_conn
    conn.execute("ALTER TABLE Materials ADD COLUMN legacy_hidden TEXT DEFAULT 'retained'")
    conn.execute("UPDATE Materials SET status = 'legacy HOLD ', stock_qty = NULL WHERE material_id = 'MAT1'")
    conn.commit()
    before, changes = business_snapshot(conn), conn.total_changes
    old, identity = material_row(conn), identity_for(conn)
    content = file_bytes([["MAT1", "改名", r"\N", None, 0], ["000007", "圆钢，特殊\"字符", "D25", "kg", 1.25]], fmt,
                         headers=HEADERS[:5])
    service = WorkbenchMaterialFileService(conn)
    preview = service.preview_import(content, file_format=fmt, scope={})
    rows = preview.as_dict()["rows"]
    assert [row["result"] for row in rows] == ["update", "new"]
    assert rows[0]["expected"]["material"]["stock_qty"] is None
    assert rows[0]["expected"]["material"]["legacy_hidden"] == "retained"
    assert rows[0]["expected"]["identity"]["ref"] == identity.ref
    assert rows[0]["requires_confirmation"] and rows[0]["changes"]["stock_qty"] == {"before": None, "after": 0.0}
    assert conn.total_changes == changes and business_snapshot(conn) == before and not conn.in_transaction
    result = confirm_import(conn, preview, content, fmt)
    assert result["result"] == "committed"
    assert material_row(conn) == {**old, "name": "改名", "spec": None, "stock_qty": 0.0}
    assert material_row(conn, "000007")["name"] == "圆钢，特殊\"字符"
    assert identity_for(conn) == replace(identity, revision=identity.revision + 1)
    after = business_snapshot(conn)
    assert {k: v for k, v in after.items() if k != "Materials"} == {k: v for k, v in before.items() if k != "Materials"}
    assert all(set(row) == {"row", "result", "entity_ref", "business_code"} for row in result["data"]["rows"])


@pytest.mark.parametrize("fmt", ("csv", "xlsx"))
def test_complete_export_roundtrip_retains_unknown_status_null_and_literal_text(material_conn, fmt):
    material_conn.execute("UPDATE Materials SET status = 'legacy HOLD ', stock_qty = NULL, spec = ?, unit = NULL, remark = ? WHERE material_id = 'MAT1'",
                          (r"\N", "literal '\\note,\n第二行"))
    material_conn.commit()
    exported = export_file(material_conn, fmt, scope={})
    before, identity = stored_state(material_conn), identity_for(material_conn)
    service = WorkbenchMaterialFileService(material_conn)
    preview = service.preview_import(exported.content, file_format=fmt, scope={})
    assert preview.as_dict()["summary"]["unchanged"] == 1
    assert preview.as_dict()["rows"][0]["input"] == {"fields": {}}
    result = confirm_import(material_conn, preview, exported.content, fmt)
    assert result["result"] == "unchanged" and identity_for(material_conn) == identity
    assert stored_state(material_conn)[:2] == before[:2]


@pytest.mark.parametrize("fmt", ("csv", "xlsx"))
def test_omitted_column_and_blank_cell_are_not_clear_or_default(material_conn, fmt):
    content = file_bytes([["MAT1", None, None]], fmt, headers=("物料编号", "库存数量", "状态"))
    service = WorkbenchMaterialFileService(material_conn)
    preview = service.preview_import(content, file_format=fmt, scope={})
    assert preview.as_dict()["rows"][0]["input"] == {"fields": {}}
    before = material_row(material_conn)
    assert confirm_import(material_conn, preview, content, fmt)["result"] == "unchanged"
    assert material_row(material_conn) == before


@pytest.mark.parametrize("fmt", ("csv", "xlsx"))
@pytest.mark.parametrize("header,value", [("状态", r"\N"), ("库存数量", r"\N"), ("名称", "   "),
                                          ("状态", "low_stock"), ("状态", "other unknown")])
def test_invalid_changes_rejected_with_exact_row_field_and_no_partial_import(material_conn, fmt, header, value):
    content = file_bytes([["MAT2", "new", None], ["MAT1", "would change", value]], fmt,
                         headers=("物料编号", "名称", header) if header != "名称" else ("物料编号", "规格", "名称"))
    preview = WorkbenchMaterialFileService(material_conn).preview_import(content, file_format=fmt, scope={})
    assert preview.as_dict()["rows"][1]["result"] == "rejected"
    assert preview.as_dict()["rows"][1]["errors"][0]["row"] == 3
    before = stored_state(material_conn)
    with patch.object(WorkbenchMaterialService, "apply", side_effect=AssertionError("full preflight required")):
        with pytest.raises(WorkbenchCommandRejected) as error:
            confirm_import(material_conn, preview, content, fmt)
    assert error.value.code == "constraint_conflict" and stored_state(material_conn) == before


@pytest.mark.parametrize("value", (True, False, 123, 1.25, datetime(2026, 1, 2)))
@pytest.mark.parametrize("header", ("物料编号", "名称", "规格", "单位", "备注"))
def test_xlsx_never_guesses_types_for_text_columns(material_conn, value, header):
    values = {"物料编号": "MAT2", "名称": "valid", header: value}
    content = file_bytes([list(values.values())], "xlsx", headers=tuple(values))
    row = WorkbenchMaterialFileService(material_conn).preview_import(content, file_format="xlsx", scope={}).as_dict()["rows"][0]
    assert row["result"] == "rejected" and row["errors"][0]["row"] == 2


@pytest.mark.parametrize("fmt,value", [("csv", x) for x in ("NaN", "Infinity", "-Infinity", "1e9999", "true", "false", "1,200 kg", -1)]
                         + [("xlsx", x) for x in (True, False, "1.5", -1, datetime(2026, 1, 2))])
def test_invalid_stock_never_turns_into_inventory(material_conn, fmt, value):
    content = file_bytes([["MAT2", "valid", value]], fmt, headers=("物料编号", "名称", "库存数量"))
    row = WorkbenchMaterialFileService(material_conn).preview_import(content, file_format=fmt, scope={}).as_dict()["rows"][0]
    assert row["result"] == "rejected" and "stock_qty" in row["errors"][0]["field"]


@pytest.mark.parametrize("fmt", ("csv", "xlsx"))
@pytest.mark.parametrize("headers", [("名称",), ("物料编号", "low_stock"), ("物料编号", "business_code"),
                                     ("物料编号", "material_id"), ("物料编号", ""), ("物料编号", "未知字段")])
def test_unknown_duplicate_or_missing_headers_fail_precisely(material_conn, fmt, headers):
    with pytest.raises(ValidationError) as error:
        WorkbenchMaterialFileService(material_conn).preview_import(file_bytes([], fmt, headers), file_format=fmt, scope={})
    assert error.value.details["row"] == 1 and error.value.field == "headers"


@pytest.mark.parametrize("fmt", ("csv", "xlsx"))
def test_duplicate_trimmed_codes_report_every_source_line_including_blank_gap(material_conn, fmt):
    content = file_bytes([[" MAT2 ", "first"], [], ["MAT2", "second"]], fmt, headers=HEADERS[:2])
    rows = WorkbenchMaterialFileService(material_conn).preview_import(content, file_format=fmt, scope={}).as_dict()["rows"]
    assert [row["row"] for row in rows] == [2, 4]
    assert all(row["result"] == "rejected" and row["errors"][-1]["code"] == "duplicate_entry" for row in rows)
    assert all("2, 4" in row["errors"][-1]["message"] for row in rows)


def test_csv_multiline_record_keeps_physical_source_line_and_chinese(material_conn):
    content = file_bytes([["MAT2", "中文,\"引号\"\n跨行"], ["MAT3", "valid", "extra"]], "csv", headers=HEADERS[:2])
    rows = WorkbenchMaterialFileService(material_conn).preview_import(content, file_format="csv", scope={}).as_dict()["rows"]
    assert rows[0]["input"]["label"] == "中文,\"引号\"\n跨行"
    assert rows[1]["errors"][0]["row"] == 4 and rows[1]["errors"][0]["field"] == "columns"


def test_xlsx_formula_and_extra_sheet_are_not_silently_evaluated_or_ignored(material_conn):
    content = file_bytes([["MAT2", "=1+1"]], "xlsx", headers=HEADERS[:2], formulas=True)
    rows = WorkbenchMaterialFileService(material_conn).preview_import(content, file_format="xlsx", scope={}).as_dict()["rows"]
    assert rows[0]["errors"][0]["field"] == "label" and rows[0]["result"] == "rejected"
    wb = openpyxl.load_workbook(BytesIO(content))
    wb.create_sheet("extra")
    output = BytesIO()
    wb.save(output)
    wb.close()
    with pytest.raises(ValidationError, match="一张"):
        WorkbenchMaterialFileService(material_conn).preview_import(output.getvalue(), file_format="xlsx", scope={})


@pytest.mark.parametrize("fmt", ("csv", "xlsx"))
def test_2000_import_rows_commit_and_2001_fail_without_truncation(schema_conn, fmt):
    rows = [[f"ROW{n:05d}", f"圆钢{n:05d}"] for n in range(2001)]
    service = WorkbenchMaterialFileService(schema_conn)
    content = file_bytes(rows[:2000], fmt, headers=HEADERS[:2])
    preview = service.preview_import(content, file_format=fmt, scope={})
    assert preview.as_dict()["summary"]["new"] == 2000
    oversized = file_bytes(rows, fmt, headers=HEADERS[:2])
    if fmt == "xlsx":
        oversized = small_dimensions(oversized)
    with pytest.raises(ValidationError, match="2000") as error:
        service.preview_import(oversized, file_format=fmt, scope={})
    assert error.value.details["row"] == 2002 and not table_rows(schema_conn, "Materials")
    result = confirm_import(schema_conn, preview, content, fmt)
    assert len(result["data"]["rows"]) == 2000 and len(table_rows(schema_conn, "Materials")) == 2000


@pytest.mark.parametrize("fmt", ("csv", "xlsx"))
def test_empty_table_roundtrip_and_template_do_not_create_demo_rows(schema_conn, fmt):
    service = WorkbenchMaterialFileService(schema_conn)
    for content in (export_file(schema_conn, fmt, scope={}).content, service.template(fmt).content):
        preview = service.preview_import(content, file_format=fmt, scope={})
        assert preview.as_dict()["rows"] == []
        result = confirm_import(schema_conn, preview, content, fmt, key=KEY + str(len(content)))
        assert result["result"] == "unchanged" and not table_rows(schema_conn, "Materials")


@pytest.mark.parametrize("change", ("content", "mode", "scope", "preview", "format"))
def test_any_approved_input_hash_or_scope_change_is_stale(material_conn, change):
    content = file_bytes([["MAT2", "new"]], "csv", headers=HEADERS[:2])
    service = WorkbenchMaterialFileService(material_conn)
    preview = service.preview_import(content, file_format="csv", scope={})
    mode, scope, fmt = "upsert", {}, "csv"
    if change == "content":
        content += b"\n"
    elif change == "mode":
        mode = "replace"
    elif change == "scope":
        scope = {"query": "different"}
    elif change == "format":
        fmt = "xlsx"
    else:
        body = preview.as_dict()
        body["rows"][0]["input"]["label"] = "changed"
        preview = replace(preview, document=canonical_json(body))
    before = stored_state(material_conn)
    with pytest.raises(WorkbenchCommandRejected) as error:
        with TransactionManager(material_conn).transaction(begin_immediate=True):
            service.confirm_import(preview, content, file_format=fmt, mode=mode, scope=scope)
    assert error.value.code == "stale_write" and stored_state(material_conn) == before


@pytest.mark.parametrize("change", ("create", "update", "recreate", "reference_same_count", "absent_aba"))
def test_concurrent_same_code_or_full_fact_changes_reject_entire_import(material_conn, tmp_path, change):
    path = tmp_path / "concurrent-materials.db"
    copy_to_temp(material_conn, path)
    with connect_temp(path) as conn:
        content = file_bytes([["MAT2", "new"], ["MAT1", "renamed"]], "csv", headers=HEADERS[:2])
        preview = WorkbenchMaterialFileService(conn).preview_import(content, file_format="csv", scope={})
        with connect_temp(path) as concurrent:
            domain = MaterialService(concurrent)
            if change in ("create", "absent_aba"):
                domain.create("MAT2", "created elsewhere")
                if change == "absent_aba":
                    domain.delete("MAT2")
            elif change == "update":
                domain.update("MAT1", remark="concurrent hidden edit")
            elif change == "recreate":
                concurrent.execute("DELETE FROM BatchMaterials WHERE material_id = 'MAT1'")
                concurrent.commit()
                domain.delete("MAT1")
                domain.create("MAT1", "rebuilt")
            else:
                concurrent.execute("UPDATE BatchMaterials SET required_qty = 99 WHERE material_id = 'MAT1'")
                concurrent.commit()
        before = stored_state(conn)
        with pytest.raises(WorkbenchCommandRejected) as error:
            confirm_import(conn, preview, content, "csv")
        assert error.value.code == "stale_write" and stored_state(conn) == before


def test_atomic_savepoint_survives_late_failure_even_when_outer_catches(material_conn):
    content = file_bytes([["MAT2", "new"], ["MAT1", "renamed"]], "csv", headers=HEADERS[:2])
    service = WorkbenchMaterialFileService(material_conn)
    preview = service.preview_import(content, file_format="csv", scope={})
    original = service.adapter.apply

    def fail_second(action, normalized, identity):
        outcome = original(action, normalized, identity)
        if action == "update":
            raise RuntimeError("fixture later domain failure")
        return outcome

    before = stored_state(material_conn)
    with patch.object(service.adapter, "apply", side_effect=fail_second):
        with TransactionManager(material_conn).transaction(begin_immediate=True):
            with pytest.raises(RuntimeError, match="fixture later"):
                service.confirm_import(preview, content, file_format="csv", scope={})
            assert stored_state(material_conn) == before
    assert stored_state(material_conn) == before


def test_receipt_failure_rolls_back_import_and_original_intent_replays_once(material_conn, monkeypatch):
    content = file_bytes([["MAT2", "new"], ["MAT1", "renamed"]], "csv", headers=HEADERS[:2])
    preview = WorkbenchMaterialFileService(material_conn).preview_import(content, file_format="csv", scope={})
    command = WorkbenchCommandService(material_conn)
    before = stored_state(material_conn)
    with monkeypatch.context() as local:
        local.setattr(command.repo, "insert", lambda **_: (_ for _ in ()).throw(OSError("fixture receipt failure")))
        with pytest.raises(WorkbenchCommandUncertain):
            confirm_import(material_conn, preview, content, "csv", command=command)
    assert stored_state(material_conn) == before
    first = confirm_import(material_conn, preview, content, "csv")
    MaterialService(material_conn).update("MAT2", name="later fact")
    before = stored_state(material_conn)
    second = confirm_import(material_conn, preview, content, "csv", guard=lambda: pytest.fail("expired preview must not run"))
    assert second == {**first, "replayed": True} and stored_state(material_conn) == before
    other = WorkbenchMaterialFileService(material_conn).preview_import(content, file_format="csv", scope={})
    with pytest.raises(WorkbenchCommandRejected) as error:
        confirm_import(material_conn, other, content, "csv")
    assert error.value.code == "request_key_conflict"


@pytest.mark.parametrize("fmt", ("csv", "xlsx"))
def test_export_all_filtered_and_explicit_selection_not_only_visible_page(material_conn, fmt):
    seed_many(material_conn, 240)
    filtered = export_file(material_conn, fmt, scope={"query": "BULK", "sort": "business_code", "direction": "desc"})
    assert filtered.row_count == 240 and filtered.filename == "物料清单." + fmt
    refs = [identity_for(material_conn, code).ref for code in ("BULK00230", "BULK00001", "MAT1")]
    selected = export_file(material_conn, fmt, selected_refs=refs)
    assert selected.row_count == 3
    if fmt == "csv":
        parsed = list(csv.reader(StringIO(selected.content.decode("utf-8-sig"))))
        assert selected.mime_type == "text/csv; charset=utf-8"
        assert parsed[0] == list(HEADERS)
        assert [row[0] for row in parsed[1:]] == ["'BULK00230", "'BULK00001", "'MAT1"]
        assert parsed[3][4] == "12.375" and parsed[3][3] == "'kg"
    else:
        wb = openpyxl.load_workbook(BytesIO(selected.content), data_only=False)
        assert selected.mime_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        assert list(wb.active.values)[0] == HEADERS and wb.active.max_row == 4
        assert [wb.active.cell(n, 1).value for n in (2, 3, 4)] == ["BULK00230", "BULK00001", "MAT1"]
        assert wb.active["E4"].value == 12.375 and wb.active["D4"].value == "kg"
        wb.close()


@pytest.mark.parametrize("fmt", ("csv", "xlsx"))
def test_text_codes_dates_formula_like_text_survive_export_without_execution(schema_conn, fmt):
    codes = ["000123", "2026-01-02", "=1+1", "+cmd", "@SUM(A1)", "-2", "'literal"]
    for code in codes:
        MaterialService(schema_conn).create(code, "=2+2", remark="中文\n换行,引号\"", stock_qty=0)
    refs = [identity_for(schema_conn, code).ref for code in codes]
    result = export_file(schema_conn, fmt, selected_refs=refs)
    if fmt == "csv":
        rows = list(csv.reader(StringIO(result.content.decode("utf-8-sig"))))[1:]
        assert [row[0] for row in rows] == ["'" + code for code in codes]
        assert all(row[1] == "'=2+2" for row in rows)
    else:
        wb = openpyxl.load_workbook(BytesIO(result.content), data_only=False)
        assert [row[0].value for row in list(wb.active.rows)[1:]] == codes
        assert all(cell.data_type != "f" for row in wb.active for cell in row)
        assert all(wb.active.cell(n, 1).data_type == "s" for n in range(2, 9))
        wb.close()
    preview = WorkbenchMaterialFileService(schema_conn).preview_import(result.content, file_format=fmt, scope={})
    assert preview.as_dict()["summary"]["unchanged"] == len(codes)


def test_export_requires_validated_transaction_exact_scope_and_current_ref(material_conn):
    service = WorkbenchMaterialFileService(material_conn)
    with pytest.raises(RuntimeError, match="快照事务"):
        service.export("csv", scope={})
    for selection in ({}, {"scope": {}, "selected_refs": []}, {"scope": {"page": 1}}, {"selected_refs": ["MAT1"]}):
        with pytest.raises(ValidationError):
            export_file(material_conn, "csv", **selection)
    MaterialService(material_conn).create("MAT2", "old")
    old = identity_for(material_conn, "MAT2")
    MaterialService(material_conn).delete("MAT2")
    MaterialService(material_conn).create("MAT2", "rebuilt")
    before = stored_state(material_conn)
    with pytest.raises(WorkbenchCommandRejected):
        export_file(material_conn, "xlsx", selected_refs=[identity_for(material_conn).ref, old.ref])
    assert stored_state(material_conn) == before


@pytest.mark.parametrize("fmt", ("csv", "xlsx"))
@pytest.mark.parametrize("selection", ("all", "filtered", "selected"))
def test_10000_material_export_rows_and_tail_are_complete(material_conn, fmt, selection):
    exercise_scale_export(material_conn, fmt, selection)


def test_empty_selection_never_expands_to_all(schema_conn):
    seed_many(schema_conn, 2001)
    assert export_file(schema_conn, "xlsx", selected_refs=[]).row_count == 0


def test_bad_formats_corrupt_files_and_legacy_keys(material_conn):
    service = WorkbenchMaterialFileService(material_conn)
    for content, fmt, mode in ((b"", "csv", "upsert"), (b"bad zip", "xlsx", "upsert"), (b"\xff", "csv", "upsert"),
                               (b"", "xls", "upsert"), (b"", "csv", "replace"), (b"", "csv", "append"),
                               ("not bytes", "csv", "upsert")):
        with pytest.raises(ValidationError):
            service.preview_import(content, file_format=fmt, scope={}, mode=mode)
    material_conn.execute("INSERT INTO Materials (material_id, name) VALUES (' MAT1 ', 'legacy')")
    material_conn.commit()
    preview = service.preview_import(file_bytes([[" MAT1 ", "unsafe"]], "csv", HEADERS[:2]), file_format="csv", scope={})
    assert preview.as_dict()["rows"][0]["result"] == "rejected"
    with pytest.raises(RuntimeError, match="外层"):
        service.confirm_import(preview, b"", file_format="csv", scope={})


@pytest.mark.parametrize("fmt", ("csv", "xlsx"))
def test_export_reads_only_and_missing_metadata_is_not_repaired(material_conn, fmt):
    before, changes = stored_state(material_conn), material_conn.total_changes
    export_file(material_conn, fmt, scope={})
    assert stored_state(material_conn) == before and material_conn.total_changes == changes
    assert not material_conn.in_transaction
    material_conn.execute("DELETE FROM WorkbenchEntityRefs WHERE kind = 'material'")
    material_conn.commit()
    before, changes = stored_state(material_conn), material_conn.total_changes
    with pytest.raises(WorkbenchCommandRejected) as error:
        export_file(material_conn, fmt, scope={})
    assert error.value.code == "storage_failure"
    assert stored_state(material_conn) == before and material_conn.total_changes == changes


def test_malformed_csv_and_unheaded_xlsx_cells_report_source_location(material_conn):
    service = WorkbenchMaterialFileService(material_conn)
    with pytest.raises(ValidationError) as error:
        service.preview_import(b'business_code,label\nMAT2,"unclosed', file_format="csv", scope={})
    assert error.value.details["row"] == 2
    content = file_bytes([["MAT2", "label", "undeclared"]], "xlsx", headers=HEADERS[:2])
    row = service.preview_import(content, file_format="xlsx", scope={}).as_dict()["rows"][0]
    assert row["errors"][0]["row"] == 2 and row["errors"][0]["field"] == "columns"


@pytest.mark.parametrize("fmt", ("csv", "xlsx"))
def test_export_consumes_one_pass_iterator_once_and_preserves_last_row(material_conn, fmt):
    rows = [{**material_row(material_conn), "material_id": f"G{n:05d}"} for n in range(3001)]
    source = OnePassRows(codec_rows(rows))
    download = write_material_file(source, fmt)
    verify_download(download, rows, fmt)
    assert source.iterations == 1


def test_xlsx_capacity_includes_header_but_does_not_limit_csv(material_conn, monkeypatch):
    from openpyxl.worksheet._writer import ALL_TEMP_FILES

    assert material_file_codec.XLSX_MAX_ROWS == 1048576
    check_export_capacity(1048575, "xlsx")
    with pytest.raises(ValidationError, match="1048576"):
        check_export_capacity(1048576, "xlsx")
    check_export_capacity(2000000, "csv")
    row = material_row(material_conn)
    monkeypatch.setattr(material_file_codec, "XLSX_MAX_ROWS", 3)
    verify_download(write_material_file(codec_rows([row] * 2), "xlsx"), [row] * 2, "xlsx")
    before = list(ALL_TEMP_FILES)
    with pytest.raises(ValidationError):
        write_material_file(codec_rows([row] * 3), "xlsx")
    assert list(ALL_TEMP_FILES) == before
    verify_download(write_material_file(codec_rows([row] * 4), "csv"), [row] * 4, "csv")


@pytest.mark.parametrize("text", ("x" * 32768, "\\" + "x" * 32766, "embedded\x01text"))
def test_xlsx_cell_capacity_never_truncates_and_is_not_a_csv_limit(material_conn, text):
    row = {**material_row(material_conn), "remark": text}
    with pytest.raises(ValidationError, match="XLSX"):
        write_material_file(codec_rows([row]), "xlsx")
    verify_download(write_material_file(codec_rows([row]), "csv"), [row], "csv")


def test_csv_rejects_runtime_unsupported_nul_without_silent_deletion(material_conn):
    row = {**material_row(material_conn), "remark": "original\x00text"}
    with pytest.raises(ValidationError, match="CSV 打不开的隐藏字符"):
        write_material_file(codec_rows([row]), "csv")


@pytest.mark.parametrize("fmt", ("csv", "xlsx"))
def test_export_keeps_one_read_snapshot_while_another_connection_updates_tail(material_conn, tmp_path, fmt):
    seed_many(material_conn, 1000)
    path = tmp_path / "read-snapshot.db"
    copy_to_temp(material_conn, path)
    with connect_temp(path) as reader, connect_temp(path) as writer:
        writer.execute("PRAGMA journal_mode=WAL")
        service = WorkbenchMaterialFileService(reader)
        expected = [dict(row) for row in reader.execute("SELECT * FROM Materials WHERE material_id LIKE 'BULK%' ORDER BY material_id")]
        batch_before = table_rows(reader, "BatchMaterials")
        original = service._export_rows

        def concurrent_rows(rows):
            for index, row in enumerate(original(rows)):
                if index == 10:
                    MaterialService(writer).update("BULK00999", name="new tail", status="inactive")
                    MaterialService(writer).create("BULK_NEW", "late insert")
                yield row

        with patch.object(service, "_export_rows", side_effect=concurrent_rows):
            with TransactionManager(reader).transaction():
                download = service.export(fmt, scope={"query": "BULK", "status": "active"})
        verify_download(download, expected, fmt)
        assert material_row(reader, "BULK00999")["name"] == "new tail"
        assert material_row(reader, "BULK_NEW") is not None
        assert table_rows(reader, "BatchMaterials") == batch_before
