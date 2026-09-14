"""Byte-level plan exports remain pinned to the admitted workspace snapshot."""

import codecs
import csv
import io
import json
from types import SimpleNamespace
from zipfile import ZipFile

import openpyxl
import pytest

import web.public_token_registry as tokens
from core.services.workbench.plan_export import HEADERS
from tests.workbench.plan_catalog_support import history
from tests.workbench.plan_read_support import NIGHT_END, NIGHT_START, add_tasks, assert_error, plan_read_api
from tests.workbench.test_plan_workspace_projections import prepare


def read_rows(response, fmt):
    assert response.status_code == 200, response.get_data(as_text=True) if response.is_json else response.status_code
    assert response.headers['Cache-Control'] == 'no-store'
    assert response.headers['Content-Disposition'].startswith('attachment;')
    if fmt == 'csv':
        assert response.data.startswith(codecs.BOM_UTF8)
        assert response.mimetype == 'text/csv'
        raw = list(csv.reader(io.StringIO(response.data.decode('utf-8-sig'), newline='')))
        assert tuple(raw[0]) == HEADERS
        return [raw[0]] + [[value[1:] if value.startswith("'") else value for value in row] for row in raw[1:]]
    assert response.mimetype == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    assert response.data.startswith(b'PK')
    with ZipFile(io.BytesIO(response.data)) as archive:
        assert archive.testzip() is None
        assert b'<f>' not in archive.read('xl/worksheets/sheet1.xml')
    wb = openpyxl.load_workbook(io.BytesIO(response.data), read_only=True, data_only=False)
    try:
        assert wb.sheetnames == ['计划任务']
        raw = list(wb.active.iter_rows(values_only=True))
        assert tuple(raw[0]) == HEADERS
        return raw
    finally:
        wb.close()


@pytest.mark.parametrize('fmt', ['csv', 'xlsx'])
@pytest.mark.parametrize('role,scenario,kind', [('adopted', None, '正式计划'), ('baseline_best', None, '候选方案'),
                                               ('critical_best', None, '候选方案'), ('adopted', 'PRIVATE-ACTIVE', '试调方案')])
def test_exact_workspace_tasks_and_source_in_file(plan_api, fmt, role, scenario, kind):
    prepare(plan_api)
    ref = plan_api.ref(role=role, scenario_id=scenario)
    workspace = plan_api.read('/' + ref + '/workspace')
    data, snapshot = workspace['data'], workspace['meta']
    before = plan_api.state()
    response = plan_api.get('/' + ref + '/export', format=fmt, snapshot_ref=snapshot['snapshot_ref'])
    rows = read_rows(response, fmt)
    assert len(rows) - 1 == data['task_count'] == int(response.headers['X-Workbench-Row-Count'])
    assert response.headers['X-Workbench-Snapshot-Ref'] == snapshot['snapshot_ref']
    assert response.headers['X-Workbench-Plan-Ref'] == ref
    for values, task in zip(rows[1:], data['tasks']):
        assert list(values[0:4]) == [ref, '3', kind, data['plan']['display_name']]
        assert list(values[6:8]) == [snapshot['snapshot_ref'], snapshot['as_of']]
        assert values[11] == task['task_ref'] and values[12] == task['operation_ref']
        assert values[22] == task['start'] and values[23] == task['end']
    assert plan_api.state() == before
    assert not any(sql.lstrip().split()[0].upper() not in {'SELECT', 'WITH', 'BEGIN', 'COMMIT'}
                   for sql in plan_api.statements if not sql.startswith('-- PRAGMA '))


@pytest.mark.parametrize('fmt', ['csv', 'xlsx'])
def test_lossless_formula_text_null_zero_date_and_int64(plan_api, fmt):
    prepare(plan_api)
    label = '=SUM(1,2)\r\n中文\t"quote"'
    version = (1 << 63) - 1
    with plan_api.db() as conn:
        history(conn, version, op_id=1)
        conn.execute('UPDATE BatchOperations SET seq=?,op_type_name=? WHERE id=1', (version, label))
        conn.execute("UPDATE Batches SET due_date='2026-09-09'")
    ref = plan_api.ref(version)
    snapshot = plan_api.read('/' + ref + '/workspace')['meta']['snapshot_ref']
    response = plan_api.get('/' + ref + '/export', format=fmt, snapshot_ref=snapshot)
    row = read_rows(response, fmt)[1]
    assert row[1] == row[14] == str(version)
    assert row[15] == label
    assert row[16] == row[17] == row[20] == r'\N'
    assert row[25] == '2026-09-09T09:00:00' and row[27] == '2026-09-09'
    assert row[29] == row[30] == ('0.0' if fmt == 'csv' else 0)
    if fmt == 'csv':
        raw = list(csv.reader(io.StringIO(response.data.decode('utf-8-sig'), newline='')))[1]
        assert raw[15] == "'" + label and raw[1] == "'" + str(version)
    with plan_api.db() as conn:
        conn.execute("UPDATE Batches SET due_date=NULL")
        conn.execute("UPDATE BatchOperations SET op_type_name=?", (r'\N',))
    snapshot = plan_api.read('/' + ref + '/workspace')['meta']['snapshot_ref']
    row = read_rows(plan_api.get('/' + ref + '/export', format=fmt, snapshot_ref=snapshot), fmt)[1]
    assert row[15] == r'\\N' and row[27] == row[29] == row[30] == r'\N'


@pytest.mark.parametrize('fmt', ['csv', 'xlsx'])
def test_scope_exports_whole_overlapping_tasks_and_empty_header_only(plan_api, fmt):
    prepare(plan_api)
    ref = plan_api.ref()
    scope = {'range_start': '2026-09-10T00:00:00', 'range_end': '2026-09-10T01:00:00'}
    snapshot = plan_api.read('/' + ref + '/workspace', **scope)['meta']['snapshot_ref']
    row = read_rows(plan_api.get('/' + ref + '/export', format=fmt, snapshot_ref=snapshot, **scope), fmt)[1]
    assert row[8] == scope['range_start'] and row[9] == scope['range_end']
    assert row[22] == NIGHT_START and row[23] == NIGHT_END
    assert_error(plan_api.get('/' + ref + '/export', format=fmt, snapshot_ref=snapshot), 'snapshot_stale')
    scope = {'range_start': NIGHT_END, 'range_end': '2026-09-10T08:00:00'}
    snapshot = plan_api.read('/' + ref + '/workspace', **scope)['meta']['snapshot_ref']
    response = plan_api.get('/' + ref + '/export', format=fmt, snapshot_ref=snapshot, **scope)
    assert len(read_rows(response, fmt)) == 1
    assert response.headers['X-Workbench-Row-Count'] == '0'


@pytest.mark.parametrize('query', [{}, {'format': 'csv'}, {'format': 'pdf', 'snapshot_ref': 'x'},
                                  {'format': 'csv', 'snapshot_ref': ''}, {'format': 'csv', 'snapshot_ref': 'x', 'query': 'part'}])
def test_missing_snapshot_and_unsupported_export_parameters_fail_before_read(plan_api, query):
    ref = plan_api.ref()
    assert_error(plan_api.get('/' + ref + '/export', **query), 'invalid_input', 400)
    assert plan_api.statements == []


def test_export_wrong_identity_catalog_snapshot_and_expiry_rejected(plan_api, monkeypatch):
    now = [1000000.0]
    monkeypatch.setattr(tokens, 'time', SimpleNamespace(time=lambda: now[0]))
    ref = plan_api.ref()
    snapshot = plan_api.read('/' + ref + '/workspace')['meta']['snapshot_ref']
    assert_error(plan_api.get('/' + plan_api.ref(role='baseline_best') + '/export', format='csv', snapshot_ref=snapshot), 'snapshot_stale')
    assert_error(plan_api.get('/' + 'b' * 48 + '/export', format='csv', snapshot_ref=snapshot), 'entity_not_found', 404)
    catalog = plan_api.read()['meta']['snapshot_ref']
    assert_error(plan_api.get('/' + ref + '/export', format='csv', snapshot_ref=catalog), 'snapshot_stale')
    now[0] += 901
    assert_error(plan_api.get('/' + ref + '/export', format='csv', snapshot_ref=snapshot), 'snapshot_stale')


def test_exports_do_not_issue_replacement_tokens_or_write_receipts(plan_api, monkeypatch):
    ref = plan_api.ref()
    snapshot = plan_api.read('/' + ref + '/workspace')['meta']['snapshot_ref']

    def forbidden(*args, **kwargs):
        pytest.fail('Pinned export cannot allocate a new token')

    monkeypatch.setattr('web.routes.workbench.read_context.issue_public_token', forbidden)
    before = plan_api.state()
    assert plan_api.get('/' + ref + '/export', format='csv', snapshot_ref=snapshot).status_code == 200
    assert plan_api.state() == before


def test_10000_tasks_export_complete_and_real_8mib_aggregate_limit(plan_api):
    with plan_api.db() as conn:
        add_tasks(conn, 9999)
    ref = plan_api.ref()
    workspace = plan_api.read('/' + ref + '/workspace')
    assert workspace['data']['task_count'] == 10000
    for fmt in ('csv', 'xlsx'):
        response = plan_api.get('/' + ref + '/export', format=fmt, snapshot_ref=workspace['meta']['snapshot_ref'])
        rows = read_rows(response, fmt)
        assert len(rows) == 10001
        assert {row[11] for row in rows[1:]} == {task['task_ref'] for task in workspace['data']['tasks']}
        print('PLAN_EXPORT_SCALE ' + json.dumps({'format': fmt, 'tasks': 10000, 'bytes': len(response.data)}))
    with plan_api.db() as conn:
        conn.execute('UPDATE BatchOperations SET op_type_name=?', ('X' * 700,))
    assert_error(plan_api.get('/' + ref + '/workspace'), 'query_too_large', 413)
    assert_error(plan_api.get('/' + ref + '/export', format='csv', snapshot_ref=workspace['meta']['snapshot_ref']), 'query_too_large', 413)
    with plan_api.db() as conn:
        conn.execute("UPDATE BatchOperations SET op_type_name='Operation'")
        add_tasks(conn, 1)
    assert_error(plan_api.get('/' + ref + '/workspace'), 'query_too_large', 413)
