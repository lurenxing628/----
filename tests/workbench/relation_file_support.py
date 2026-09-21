"""可操作设备关系文件的临时库夹具与独立取值口径。

断言用的是直接查 OperatorMachine 表的独立口径，不调生产的读取方法来验证生产的写入结果。
"""

import csv
from io import BytesIO, StringIO

import openpyxl
import pytest

from core.models.workbench_table_descriptor import INSTRUCTION_SHEET
from tests.workbench.identity_metadata_support import insert_row
from tests.workbench.resource_entity_support import resource_database  # noqa: F401

OPERATORS = (("RO1", "关系人员一"), ("RO2", "关系人员二"))
MACHINES = (("RM1", "关系设备一"), ("RM2", "关系设备二"))


@pytest.fixture(name="relation_env")
def relation_database(resource_conn):
    """在资源种子之上加一组没有任何关联的人员与设备，让关系用例从空白起步。"""
    for code, name in OPERATORS:
        insert_row(resource_conn, "Operators", {"operator_id": code, "name": name, "status": "active"})
    for code, name in MACHINES:
        insert_row(resource_conn, "Machines", {"machine_id": code, "name": name, "op_type_id": "OT1"})
    resource_conn.commit()
    return resource_conn


def links_of(conn, operator_code):
    rows = conn.execute("SELECT machine_id, skill_level, is_primary FROM OperatorMachine WHERE operator_id=?",
                        (operator_code,)).fetchall()
    return {row["machine_id"]: (row["skill_level"], row["is_primary"]) for row in rows}


def file_bytes(rows, fmt="csv", headers=("工号", "设备编号", "技能等级", "主操设备")):
    if fmt == "csv":
        target = StringIO(newline="")
        writer = csv.writer(target)
        writer.writerow(headers)
        writer.writerows(rows)
        return target.getvalue().encode("utf-8-sig")
    wb = openpyxl.Workbook()
    ws = wb.worksheets[0]
    ws.append(list(headers))
    for row in rows:
        ws.append(list(row))
    output = BytesIO()
    wb.save(output)
    wb.close()
    return output.getvalue()


def decode(download, fmt):
    if fmt == "csv":
        rows = list(csv.reader(StringIO(download.content.decode("utf-8-sig"))))
        return rows[0], [[value[1:] if value.startswith("'") else value for value in row] for row in rows[1:]]
    wb = openpyxl.load_workbook(BytesIO(download.content), read_only=True, data_only=False)
    try:
        # 模板与导出文件固定两张表：数据表在前，「填写说明」在后。
        assert wb.sheetnames[1:] == [INSTRUCTION_SHEET], wb.sheetnames
        source = wb.worksheets[0].iter_rows()
        headers = [cell.value for cell in next(source)]
        rows = []
        for cells in source:
            assert all(cell.data_type not in ("f", "e") for cell in cells)
            rows.append([cell.value for cell in cells])
        return headers, rows
    finally:
        wb.close()
