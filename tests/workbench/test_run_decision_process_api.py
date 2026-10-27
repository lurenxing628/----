"""Shared-run HTTP processes and public search summaries preserve saved evidence."""

import csv
import io
import json
import sys
from copy import deepcopy

import pytest
from openpyxl import load_workbook

from core.services.scheduler.contracts.optimizer_public_safety import project_attempt_metrics, project_public_metrics
from core.services.scheduler.run.optimizer.shared_process import UNIFIED_PROCESS_VERSION
from core.services.scheduler.summary.optimizer_public_attempts import COMPARISON_METRIC_KEYS, project_attempts
from core.services.scheduler.summary.optimizer_public_summary import project_public_algo_summary
from core.services.workbench.run.decision_process_export import UNIFIED_PROCESS_CSV_HEADERS
from tests.workbench.run_candidate_support import BASE, api, compute, retained
from tests.workbench.run_candidate_support import candidate_case as _candidate_case  # noqa: F401
from tests.workbench.test_run_decision_api import decision_summary
from tests.workbench.test_run_decision_impact_api import detail

PROCESS_SHEETS = ("共同计算", "目标预算", "进步事件", "搜索动作", "搜索阶段", "评价范围", "范围成员", "真实桥接")


def expect_error(client, path, query, status, code):
    response = client.get(path, query_string=query)
    assert response.status_code == status, response.get_json()
    assert response.get_json()["error"]["code"] == code


def process_download(client, run_ref, query, token, *, fmt, section="process"):
    response = client.get(BASE + "/runs/" + run_ref + "/decision/export", query_string={
        **query, "snapshot_ref": token, "section": section, "format": fmt})
    assert response.status_code == 200, response.get_json()
    assert response.headers["X-Workbench-Snapshot-Ref"] == token
    assert response.headers["X-Workbench-Run-Ref"] == run_ref
    if fmt == "csv":
        reader = csv.DictReader(io.StringIO(response.data.decode("utf-8-sig")))
        rows = list(reader)
        assert reader.fieldnames == list(UNIFIED_PROCESS_CSV_HEADERS)
    else:
        book = load_workbook(io.BytesIO(response.data), read_only=True, data_only=False)
        try:
            if section == "process":
                assert book.sheetnames == list(PROCESS_SHEETS) + ["指标计算方式", "比较范围"]
            assert all(book.sheetnames.count(title) == 1 for title in PROCESS_SHEETS)
            rows = []
            for title in PROCESS_SHEETS:
                cells = iter(book[title].iter_rows())
                headers = [cell.value for cell in next(cells)]
                for cells_row in cells:
                    for cell in cells_row:
                        if type(cell.value) is str:
                            assert cell.data_type == "s" and cell.number_format == "@"
                        elif type(cell.value) is bool:
                            assert cell.data_type == "b"
                        elif type(cell.value) in (int, float):
                            assert cell.data_type == "n"
                    rows.append(dict(zip(headers, [cell.value for cell in cells_row])))
        finally:
            book.close()
    if section == "process":
        assert int(response.headers["X-Workbench-Row-Count"]) == len(rows)
    return rows


def exported_value(row, label, value, fmt):
    actual = row[label]
    if type(value) in (int, float):
        assert float(actual) == pytest.approx(value, rel=1e-12, abs=1e-9), label
    elif fmt == "csv":
        expected = "" if value is None else str(value)
        if type(value) is str and (value.lstrip().startswith(("=", "+", "-", "@"))
                                  or label.endswith("编号") and value.isdecimal()):
            expected = "'" + value
        assert actual == expected, (label, actual, expected)
    else:
        assert actual == (None if value == "" else value), label


def check_scope_exports(rows, scopes, fmt):
    groups = [row for row in rows if row["record_kind"] == "score_scope"]
    members = [row for row in rows if row["record_kind"] == "scope_member"]
    assert len(groups) == len(scopes)
    expected_members = [(ref, kind, index, member) for ref, scope in scopes.items()
                        for kind, field in (("scheduled_operation", "scheduled_operation_refs"),
                                            ("complete_batch", "complete_batch_refs"))
                        for index, member in enumerate(scope[field], 1)]
    assert len(members) == len(expected_members)
    for row, (ref, kind, index, member) in zip(members, expected_members):
        for label, value in (("评价范围编号", ref), ("成员类型", kind), ("成员顺序", index), ("成员永久编号", member)):
            exported_value(row, label, value, fmt)
    for row, (ref, scope) in zip(groups, scopes.items()):
        for label, value in (("评价范围编号", ref), ("目标评价范围", scope["objective_scope"]),
                             ("评分分量已知", scope["components_known"]), ("本次工序已排齐", scope["is_complete"]),
                             ("已排工序成员数", len(scope["scheduled_operation_refs"])),
                             ("排齐批次成员数", len(scope["complete_batch_refs"]))):
            exported_value(row, label, value, fmt)


def test_real_shared_process_queries_and_exports_preserve_one_snapshot(candidate_case):
    case = candidate_case
    case.config(algo_mode="improve", dispatch_mode="sgs", dispatch_rule="slack", time_budget_seconds=1, objective="min_overdue")
    run_ref, refs = compute(case)
    assert 1 <= len(refs) <= 3
    raw = json.loads(case.conn.execute("SELECT process_json FROM WorkbenchRunDecisionReadModels WHERE run_ref=?", (run_ref,)).fetchone()[0])
    artifacts = {row[0]: json.loads(row[1]) for row in case.conn.execute(
        "SELECT candidate_ref,artifact_json FROM WorkbenchRunCandidates WHERE run_ref=? ORDER BY sequence", (run_ref,))}
    assert raw["version"] == UNIFIED_PROCESS_VERSION
    assert raw["primary_objective"] == "min_overdue" and raw["target_objectives"] == ["min_overdue", "min_tardiness", "min_changeover"]
    assert raw["budget"]["configured_seconds"] == 1
    events = {row["event_id"]: row for row in raw["events"]}
    client, statements = api(case)
    query = {"selected_candidate_ref": refs[-1], "reference_candidate_ref": refs[0]}
    path = BASE + "/runs/" + run_ref + "/decision/process"
    units = {"min_overdue": ("overdue_count", "批次"), "min_tardiness": ("total_tardiness_hours", "小时"),
             "min_changeover": ("changeover_count", "次")}
    with retained(case.conn):
        overview = decision_summary(client, run_ref, **query)
        statements.clear()
        initial = detail(client, run_ref, "process", **query)
        data = initial["data"]
        reads = [sql.upper() for sql in statements if sql.lstrip().upper().startswith("SELECT")]
        assert not any(any(column in sql for column in ("FACTS_JSON", "EXECUTION_JSON", "RECEIPT_JSON", "ARTIFACT_JSON", "PAYLOAD_JSON")) for sql in reads)
        assert data["process_version"] == UNIFIED_PROCESS_VERSION
        assert data["scope"]["selected_candidate_ref"] == refs[-1] and data["scope"]["reference"]["candidate_ref"] == refs[0]
        assert data["candidate"]["candidate_ref"] == data["filters"]["process_candidate_ref"] == refs[-1]
        assert data["filters"]["objective"] == data["selected_objective"]["name"] == raw["primary_objective"]
        assert data["counts"] == raw["counts"] and data["budget"] == raw["budget"]
        assert data["summary"]["elapsed_ms"] == raw["elapsed_ms"] and data["stop_reason"] == raw["stop_reason"]
        assert set(data["summary"]) == {"elapsed_ms", "first_complete_elapsed_ms", "first_improvement_elapsed_ms", "last_improvement_elapsed_ms"}
        for field in ("first_improvement_elapsed_ms", "last_improvement_elapsed_ms"):
            assert data["summary"][field] == raw["targets"][raw["primary_objective"]][field]
        complete_times = [event["elapsed_ms"] for event in raw["events"] if event["event_type"] == "archive_observation"
                          and event["verified"] is True and raw["scopes"][event["scope_ref"]]["is_complete"] is True]
        assert data["summary"]["first_complete_elapsed_ms"] == min(complete_times)
        assert data["scopes"] == raw["scopes"] and data["bridges"] == raw["bridges"]
        explicit = detail(client, run_ref, "process", **dict(query, objective=raw["primary_objective"], snapshot_ref=initial["meta"]["snapshot_ref"]))
        assert all(explicit["meta"][key] == initial["meta"][key] for key in ("snapshot_ref", "as_of"))
        assert explicit["data"] == data
        assert {candidate["candidate_ref"] for candidate in data["candidates"]} == set(refs)
        for ref in refs:
            anchored = detail(client, run_ref, "process", **dict(query, process_candidate_ref=ref))["data"]
            assert anchored["scope"] == data["scope"] and anchored["baseline"] == data["baseline"]
            assert anchored["summary"] == data["summary"] and anchored["counts"] == data["counts"]
            assert anchored["candidate"]["candidate_ref"] == ref
            for field in ("representative_objectives", "origin", "first_found_elapsed_ms", "search_evidence_ref"):
                assert anchored["candidate"][field] == artifacts[ref][field]
            evidence = anchored["candidate"]["search_evidence_ref"]
            assert evidence["run_ref"] == run_ref and anchored["source_event"] == events[evidence["source_event_id"]]
        for objective in raw["target_objectives"]:
            current = detail(client, run_ref, "process", **dict(query, objective=objective, record_kind="progress", size=50))["data"]
            descriptor = current["selected_objective"]
            metric, unit = units[objective]
            assert (descriptor["primary_metric_key"], descriptor["unit"]) == (metric, unit)
            assert current["scope"] == data["scope"]
            assert current["summary"] == {**data["summary"], **{field: raw["targets"][objective][field]
                for field in ("first_improvement_elapsed_ms", "last_improvement_elapsed_ms")}}
            assert current["actions"] == [action for action in raw["actions"] if action["search_objective"] == objective]
            assert current["progress"]["sampling"] == raw["targets"][objective]["sampling"]
            assert 1 <= len(current["progress"]["points"]) <= 200
            for point, saved in zip(current["progress"]["points"], raw["targets"][objective]["points"]):
                assert all(point[key] == value for key, value in saved.items())
                assert point["phase"] == events[point["source_event_id"]]["phase"]
                assert point["verified"] == events[point["source_event_id"]]["verified"]
                assert point["primary_value"] == point["score"][point["field_order"].index(metric)]
            event_order = detail(client, run_ref, "process", **dict(query, objective=objective, record_kind="progress", sort="event_index", size=50))["data"]
            assert [point["source_event_id"] for point in event_order["items"]] == sorted(point["source_event_id"] for point in current["progress"]["points"])[:50]
        actions = detail(client, run_ref, "process", **dict(query, record_kind="attempts", sort="rank", order="desc", size=50))["data"]
        assert [action["action_id"] for action in actions["items"]] == sorted((action["action_id"] for action in data["actions"]), reverse=True)[:50]
        assert all("rank" not in action and "primary_value" not in action for action in actions["actions"])
        progress_query = dict(query, process_candidate_ref=refs[0], objective="min_tardiness", record_kind="progress", sort="elapsed_ms", order="desc")
        first = detail(client, run_ref, "process", **progress_query)
        token = first["meta"]["snapshot_ref"]
        assert first["data"]["scope"] == data["scope"]
        second = detail(client, run_ref, "process", **dict(progress_query, page=2, snapshot_ref=token))
        wide = detail(client, run_ref, "process", **dict(progress_query, size=50, snapshot_ref=token))
        assert all(second["meta"][key] == wide["meta"][key] == first["meta"][key] for key in ("snapshot_ref", "as_of"))
        assert first["data"]["items"] + second["data"]["items"] == wide["data"]["items"][:40]
        total = first["data"]["page"]["total"]
        empty = detail(client, run_ref, "process", **dict(progress_query, page=(total + 19) // 20 + 1, snapshot_ref=token))["data"]
        assert empty["items"] == [] and empty["page"]["total"] == total and empty["page"]["has_more"] is False
        assert first["data"]["page"]["size"] == 20 and wide["data"]["page"]["size"] == 50
        for changed in ({"objective": "min_overdue"}, {"order": "asc"}, {"sort": "event_index"}):
            expect_error(client, path, dict(progress_query, snapshot_ref=token, **changed), 409, "snapshot_stale")
        for invalid in ({"objective": "min_weighted_tardiness"}, {"record_kind": "raw"}, {"record_kind": "attempts", "sort": "event_index"},
                        {"record_kind": "attempts", "sort": "primary_value"}, {"page": 0}, {"size": 27}, {"order": "random"}):
            expect_error(client, path, dict(query, **invalid), 422, "invalid_input")
        expect_error(client, path, dict(query, process_candidate_ref="bad"), 400, "invalid_input")
        expect_error(client, path, dict(query, process_candidate_ref="f" * 48), 404, "entity_not_found")
        expect_error(client, BASE + "/runs/" + "e" * 48 + "/decision/process", {}, 404, "entity_not_found")
        progress = first["data"]
        expected_points = sorted(sorted(progress["progress"]["points"], key=lambda point: point["source_event_id"]), key=lambda point: point["elapsed_ms"], reverse=True)
        expected_actions = sorted(progress["actions"], key=lambda action: action["action_id"], reverse=True)
        for fmt in ("csv", "xlsx"):
            rows = process_download(client, run_ref, progress_query, token, fmt=fmt)
            summary = [row for row in rows if row["record_kind"] == "run_process_summary"]
            assert len(summary) == 1
            for label, value in (("过程版本", UNIFIED_PROCESS_VERSION), ("共同总耗时（毫秒）", raw["elapsed_ms"]),
                                 ("首次完整解（毫秒）", data["summary"]["first_complete_elapsed_ms"]),
                                 ("共同实际排程调用数", raw["counts"]["actual_decoder_calls"]),
                                 ("查看过程候选编号", progress["candidate"]["candidate_ref"]), ("查看过程来源事件编号", progress["source_event"]["event_id"])):
                exported_value(summary[0], label, value, fmt)
            for row in rows:
                for label, value in (("读取编号", token), ("排产编号", run_ref), ("归档编号", run_ref),
                                     ("保存推荐候选编号", overview["data"]["recommendation"]["candidate_ref"]),
                                     ("当前查看候选编号", refs[-1]), ("查看基准编号", refs[0])):
                    exported_value(row, label, value, fmt)
            budgets = [row for row in rows if row["record_kind"] == "objective_budget"]
            assert len(budgets) == 1 and budgets[0]["优化方向"] == "min_tardiness"
            exported_value(budgets[0], "首次真实评分改善（毫秒）", progress["summary"]["first_improvement_elapsed_ms"], fmt)
            exported_value(budgets[0], "最后真实评分改善（毫秒）", progress["summary"]["last_improvement_elapsed_ms"], fmt)
            points = [row for row in rows if row["record_kind"] == "progress_point"]
            assert len(points) == len(expected_points)
            for row, point in zip(points, expected_points):
                for label, value in (("累计共同耗时（毫秒）", point["elapsed_ms"]), ("来源事件编号", point["source_event_id"]),
                                     ("评价范围编号", point["scope_ref"]), ("连续段", point["segment"]), ("业务主值", point["primary_value"]),
                                     ("已验证", point["verified"]), ("主指标单位", "小时")):
                    exported_value(row, label, value, fmt)
                fields = {field["key"]: field for field in progress["selected_objective"]["score_fields"]}
                for index, (key, value) in enumerate(zip(point["field_order"], point["score"])):
                    exported_value(row, "评分分量" + str(index), value, fmt)
                    exported_value(row, "评分分量" + str(index) + "指标key", key, fmt)
                    exported_value(row, "评分分量" + str(index) + "单位", fields[key]["unit"], fmt)
            exported_actions = [row for row in rows if row["record_kind"] == "search_action"]
            assert len(exported_actions) == len(expected_actions)
            for row, action in zip(exported_actions, expected_actions):
                for label, value in (("动作编号", action["action_id"]), ("动作实际耗时（毫秒）", action["elapsed_ms"]),
                                     ("动作新增决策数", action["new_decisions"]), ("动作复用决策数", action["reused_decisions"])):
                    exported_value(row, label, value, fmt)
            check_scope_exports(rows, raw["scopes"], fmt)
        all_rows = process_download(client, run_ref, {**progress_query, "filter_section": "process"}, token, fmt="xlsx", section="all")
        all_summary = [row for row in all_rows if row["record_kind"] == "run_process_summary"]
        assert len(all_summary) == 1
        assert all_summary[0]["查看过程候选编号"] == progress["candidate"]["candidate_ref"]
        assert all_summary[0]["查看过程来源事件编号"] == progress["source_event"]["event_id"]
        assert all(row["读取编号"] == token and row["当前查看候选编号"] == query["selected_candidate_ref"]
                   and row["查看基准编号"] == query["reference_candidate_ref"] for row in all_rows)
        current_points = [row for row in all_rows if row["record_kind"] == "progress_point" and row["优化方向"] == "min_tardiness"]
        assert [row["累计共同耗时（毫秒）"] for row in current_points] == [point["elapsed_ms"] for point in expected_points]
        assert [row["优化方向"] for row in all_rows if row["record_kind"] == "objective_budget"] == raw["target_objectives"]
        for row in all_rows:
            if row["record_kind"] == "objective_budget":
                target = raw["targets"][row["优化方向"]]
                exported_value(row, "首次真实评分改善（毫秒）", target["first_improvement_elapsed_ms"], "xlsx")
                exported_value(row, "最后真实评分改善（毫秒）", target["last_improvement_elapsed_ms"], "xlsx")
        assert sum(row["record_kind"] == "search_action" for row in all_rows) == len(raw["actions"])
        check_scope_exports(all_rows, raw["scopes"], "xlsx")
        assert decision_summary(client, run_ref, **query)["data"]["selection_basis"] == overview["data"]["selection_basis"]


def test_public_attempt_projection_drops_numbers_beyond_the_float_range_without_raising():
    huge = 10 ** 400
    completion = {"objective_scope": "all_batches", "incomplete_work_weight": huge}
    for project in (project_attempt_metrics, project_public_metrics):
        assert project({"completion": completion})["completion"] == {"objective_scope": "all_batches"}
    public_attempts, _diagnostics = project_attempts([{"score": [huge, 1.5], "elapsed_ms": huge, "metrics": {"completion": completion}}])
    assert public_attempts == [{"score": [None, 1.5], "elapsed_ms": None, "metrics": {"completion": {"objective_scope": "all_batches"}}}]


def test_public_algo_summary_keeps_sentinels_graph_origins_and_comparisons_idempotent():
    # Explicit unknown components are ordering sentinels; completed-subset zero is real.
    sentinel = sys.float_info.max
    attempts = [
        {"tag": "local:unknown_scope", "score": [2, sentinel, sentinel],
         "metrics": {"total_tardiness_hours": sentinel, "completion": {"objective_scope": "unknown"}}},
        {"tag": "local:unknown_policy", "score": [1, sentinel, sentinel],
         "metrics": {"total_tardiness_hours": 0, "completion": {"objective_score_policy": "unknown_all_components"}}},
        {"tag": "local:completed_subset", "score": [1, 0, 0],
         "metrics": {"total_tardiness_hours": 0, "completion": {"objective_scope": "completed_batches_only",
                     "objective_score_policy": "completed_batches_only", "objective_defined": False,
                     "expected_operation_count": 3, "missing_operation_count": 1, "unexpected_result_count": 0,
                     "complete_batch_count": 1, "incomplete_batch_count": 1}}},
        {"tag": "local:all_batches", "score": [0, 5.5, 1.25], "metrics": {"total_tardiness_hours": 5.5,
         "completion": {"objective_scope": "all_batches", "expected_operation_count": 31, "missing_operation_count": 0,
                        "unexpected_result_count": 0, "complete_batch_count": 11, "incomplete_batch_count": 0}}},
    ]
    trace = [{"elapsed_ms": index, **deepcopy(attempt)} for index, attempt in enumerate(attempts, 1)]
    ordinary, _diagnostics = project_public_algo_summary({"attempts": attempts, "improvement_trace": trace})
    expected = [[2, None, None], [1, None, None], [1, 0, 0], [0, 5.5, 1.25]]
    assert [item["score"] for item in ordinary["attempts"]] == expected
    assert [item["score"] for item in ordinary["improvement_trace"]] == expected

    origins = ("graph_ready_v2_generated", "graph_ready_v2_repaired", "graph_ready_v2_iterated_greedy")
    graph_attempts = [{"candidate_origin": origin, "weight_profile_slug": "profile-" + str(index),
                       "tag": "graph_ready:profile|sgs:slack", "score": [0, 1, 1.5], "metrics": {"overdue_count": 1}}
                      for index, origin in enumerate(origins, 1)]
    public_attempts, _diagnostics = project_attempts(graph_attempts)
    assert project_attempts(public_attempts)[0] == public_attempts
    summary, _diagnostics = project_public_algo_summary({"attempts": public_attempts})
    assert summary["attempts"] == public_attempts

    # Freeze the recorded object and decisions; comparison is independent of representative rank.
    metrics = {"failed_ops": 0, "overdue_count": 1, "weighted_tardiness_hours": 1.5, "total_tardiness_hours": 1.5,
               "makespan_hours": 2.5, "changeover_count": 0}
    complete = {"expected_operation_count": 1, "missing_operation_count": 0, "unexpected_result_count": 0,
                "complete_batch_count": 1, "incomplete_batch_count": 0}
    comparator = {"kind": "incumbent", "state": "comparable", "objective_name": "min_overdue", "objective_scope": "all_batches",
                  "reference_scope": "all_batches", "reference_score": [0, 1, 1.5, 1.5, 2.5, 0],
                  "reference_metrics": metrics, "reference_completion": complete, "metric_deltas": dict.fromkeys(COMPARISON_METRIC_KEYS, 0)}
    for attempt in graph_attempts:
        attempt.update(score=[0, 1, 1.5, 1.5, 2.5, 0], failed_ops=0,
                       metrics={**{key: value for key, value in metrics.items() if key != "failed_ops"},
                                "completion": {**complete, "objective_scope": "all_batches"}})
    first = {**deepcopy(comparator), "kind": "none", "state": "no_reference", "reference_scope": None,
             "reference_score": None, "reference_metrics": None, "reference_completion": None, "metric_deltas": None}
    graph_attempts[0].update(elapsed_ms=10, event_index=1, accepted=True, verified=True, strict_improved=None, comparison=first)
    graph_attempts[1].update(elapsed_ms=32, event_index=7, accepted=True, verified=True, strict_improved=False, comparison=deepcopy(comparator))
    graph_attempts[1]["comparison"]["reference_completion"]["incomplete_batch_ids"] = ["PRIVATE-BATCH"]
    graph_attempts[1]["comparison"]["reference_metrics"]["machine_id"] = "PRIVATE-MACHINE"
    graph_attempts[2].update(elapsed_ms=19, event_index=4, accepted=True, verified=True, strict_improved=True, comparison=deepcopy(comparator))
    graph_attempts[2]["comparison"]["reference_metrics"]["total_tardiness_hours"] = 2.5
    graph_attempts[2]["comparison"]["reference_score"][3] = 2.5
    graph_attempts[2]["comparison"]["metric_deltas"]["total_tardiness_hours"] = -1.0
    ordinary, _diagnostics = project_public_algo_summary({"attempts": graph_attempts})
    repeated, _diagnostics = project_public_algo_summary(ordinary)
    assert ordinary == repeated
    assert ordinary["attempts"][0]["comparison"] == first
    assert ordinary["attempts"][1]["comparison"] == comparator
    assert ordinary["attempts"][2]["comparison"]["metric_deltas"]["total_tardiness_hours"] == -1.0
    assert "PRIVATE" not in json.dumps(ordinary)
