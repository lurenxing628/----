"""Temporary SQLite fixtures and independent byte oracles for resource files."""

import csv
import json
import time
from io import BytesIO, StringIO
from uuid import uuid4

import openpyxl

from core.infrastructure.transaction import TransactionManager
from core.services.workbench.commands import WorkbenchCommandService
from core.services.workbench.resource_bulk import WorkbenchResourceBulkService
from core.services.workbench.resource_files import WorkbenchResourceFileService
from data.repositories.workbench_identity_repo import WorkbenchIdentityRepository
from tests.workbench.identity_metadata_support import business_snapshot
from tests.workbench.resource_entity_support import resource_database  # noqa: F401

KINDS = ("op_type", "machine", "operator", "supplier")
SCOPES = (("op_type", "internal"), ("op_type", "external"), ("machine", None), ("operator", None), ("supplier", None))
TABLES = {"op_type": ("OpTypes", "op_type_id"), "machine": ("Machines", "machine_id"),
          "operator": ("Operators", "operator_id"), "supplier": ("Suppliers", "supplier_id")}


def scope(category=None):
    return {"category": category} if category else {}


def ref(conn, kind, code):
    identity = WorkbenchIdentityRepository(conn).find_active(kind, code)
    assert identity is not None
    return identity.ref


def raw(conn, kind, code):
    table, key = TABLES[kind]
    row = conn.execute(f"SELECT * FROM {table} WHERE {key}=?", (code,)).fetchone()
    return dict(row) if row else None


def existing_raw(conn, kind, code):
    row = raw(conn, kind, code)
    assert row is not None
    return row


def snapshot(conn):
    return list(conn.iterdump())


def file_bytes(rows, fmt="csv", headers=("business_code", "label", "status")):
    if fmt == "csv":
        target = StringIO(newline="")
        writer = csv.writer(target)
        writer.writerow(headers)
        writer.writerows(rows)
        return target.getvalue().encode("utf-8-sig")
    wb = openpyxl.Workbook()
    ws = wb.worksheets[0]
    ws.append(list(headers))
    for row in rows:
        ws.append(row)
    output = BytesIO()
    wb.save(output)
    wb.close()
    return output.getvalue()


def decode(download, fmt, *, check_text=True):
    if fmt == "csv":
        rows = list(csv.reader(StringIO(download.content.decode("utf-8-sig"))))
        return rows[0], [[v[1:] if v.startswith("'") else v for v in row] for row in rows[1:]]
    wb = openpyxl.load_workbook(BytesIO(download.content), read_only=True, data_only=False)
    try:
        assert len(wb.worksheets) == 1
        source = wb.worksheets[0].iter_rows()
        headers = [cell.value for cell in next(source)]
        rows = []
        for cells in source:
            assert all(cell.data_type != "f" for cell in cells)
            if check_text:
                assert cells[0].data_type == "s"
            rows.append([cell.value for cell in cells])
        return headers, rows
    finally:
        wb.close()


def confirm(conn, kind, preview, content=None, fmt="csv", *, command=None, key=None, guard=None):
    request = preview.as_dict()["request"]
    action = preview.as_dict()["operation"]
    def mutate(_):
        if action.endswith("bulk_delete"):
            return WorkbenchResourceBulkService(conn, kind).confirm_delete(preview, request["refs"], scope=request["scope"])
        return WorkbenchResourceFileService(conn, kind).confirm_import(preview, content, file_format=fmt, scope=request["scope"])
    return (command or WorkbenchCommandService(conn)).execute(
        request_key=key or "resource-file-" + uuid4().hex, action=action, context_ref=preview.digest,
        normalized_input={"preview_ref": preview.digest}, guard=guard or (lambda: None), mutate=mutate)


def exported(conn, kind, fmt, *, category=None, selected_refs=None, query=""):
    with TransactionManager(conn).transaction():
        return WorkbenchResourceFileService(conn, kind).export(fmt, scope={**scope(category), "query": query}, selected_refs=selected_refs)


def seed_many(conn, kind, count=10000):
    table, key = TABLES[kind]
    fields = [key, "name"] + (["category"] if kind == "op_type" else ["status"])
    conn.executemany(f"INSERT INTO {table} ({','.join(fields)}) VALUES (?,?,?)",
                     [(f"FILE{i:05d}", f"file row {i}", "internal" if kind == "op_type" else "active") for i in range(count)])
    conn.commit()
    return [row[0] for row in conn.execute("SELECT ref FROM WorkbenchEntityRefs WHERE kind=? AND active=1 AND entity_key LIKE 'FILE%' ORDER BY entity_key", (kind,))]


def measure(label, function):
    start = time.monotonic()
    result = function()
    print("RESOURCE_PERF", label, "seconds=", round(time.monotonic() - start, 3))
    return result


def assert_business_equal(conn, expected):
    assert business_snapshot(conn) == expected
