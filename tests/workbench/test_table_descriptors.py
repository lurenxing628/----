"""12 张导入表的表描述协议合同：形状一致、说明表同源、表头与下拉都从描述里来。

这里不复述任何一条业务规则的文字，只锁住"各家族说的和文件里写出来的是同一份"。
"""

from io import BytesIO

import openpyxl
import pytest

from core.models.workbench_batch_file import table_descriptor as batch_descriptor
from core.models.workbench_calendar_file import table_descriptor as calendar_descriptor
from core.models.workbench_material_file import table_descriptor as material_descriptor
from core.models.workbench_process_file import table_descriptor as process_descriptor
from core.models.workbench_relation_file import table_descriptor as relation_descriptor
from core.models.workbench_resource_file import table_descriptor as resource_descriptor
from core.models.workbench_table_descriptor import (
    INSTRUCTION_SHEET,
    cell_notes,
    check_table_descriptor,
    enum_columns,
    extra_sheet_notice,
    instruction_rows,
    writable_columns,
)
from core.services.workbench.batch.file_codec import write_batch_file
from core.services.workbench.execution.field_report_files_codec import encode_reports
from core.services.workbench.execution.field_report_files_codec import (
    table_descriptor as report_descriptor,
)
from core.services.workbench.facts.file_writer import write_resource_file
from core.services.workbench.material.file_codec import write_material_file
from core.services.workbench.process.file_codec import encode_process_file
from core.services.workbench.resource.calendar_files.file_writer import write_calendar_file
from core.services.workbench.resource.relation_files.file_writer import write_relation_file
from tests.workbench.table_template_support import check_instruction_sheet, sheet_rows


def _resource(kind):
    return resource_descriptor(kind), write_resource_file(kind, [], "xlsx", template=True).content


def _process(kind):
    return process_descriptor(kind), encode_process_file(kind, [], "xlsx", template=True).content


def _calendar(kind):
    return calendar_descriptor(kind), write_calendar_file(kind, [], "xlsx", template=True).content


def _report(version):
    return report_descriptor(version), encode_reports([], format_version=version)


#: 工作台全部可导入的表。每一项给出表描述和这张表的 XLSX 模板字节。
TABLES = {
    "op_type": lambda: _resource("op_type"),
    "machine": lambda: _resource("machine"),
    "operator": lambda: _resource("operator"),
    "supplier": lambda: _resource("supplier"),
    "material": lambda: (material_descriptor(), write_material_file([], "xlsx", template=True).content),
    "route": lambda: _process("route"),
    "hours": lambda: _process("hours"),
    "batch": lambda: (batch_descriptor(), write_batch_file([], template=True)),
    "field_report_v1": lambda: _report(1),
    "field_report_v2": lambda: _report(2),
    "operator_machine": lambda: (relation_descriptor("operator_machine"),
                                 write_relation_file("operator_machine", [], "xlsx", template=True).content),
    "work_calendar": lambda: _calendar("work_calendar"),
    "operator_calendar": lambda: _calendar("operator_calendar"),
}


def test_every_importable_table_is_registered_here():
    """新增一张可导入的表就要登记，不然模板统一会漏掉它。"""
    # 报工有两种格式版本，算同一张表的两副列目录，所以清单是 13 项对 12 张表。
    assert len(TABLES) == 13
    assert {name for name in TABLES if name.startswith("field_report")} == {"field_report_v1", "field_report_v2"}


@pytest.mark.parametrize("name", sorted(TABLES))
def test_descriptor_matches_the_protocol(name):
    descriptor = TABLES[name]()[0]
    check_table_descriptor(descriptor)
    assert descriptor["table_id"] == name
    assert len(cell_notes(descriptor)) == len(descriptor["columns"])


@pytest.mark.parametrize("name", sorted(TABLES))
def test_template_headers_come_from_the_descriptor(name):
    descriptor, content = TABLES[name]()
    wb = openpyxl.load_workbook(BytesIO(content))
    try:
        data = wb[descriptor["sheet_name"]]
        assert [cell.value for cell in data[1]] == [column["label"] for column in descriptor["columns"]]
        assert data.max_row == 1, "模板的数据表只留表头，示例行在说明表里"
    finally:
        wb.close()


@pytest.mark.parametrize("name", sorted(TABLES))
def test_template_carries_the_generated_instruction_sheet(name):
    descriptor, content = TABLES[name]()
    wb = openpyxl.load_workbook(BytesIO(content))
    try:
        check_instruction_sheet(wb, descriptor)
        rows = sheet_rows(wb, INSTRUCTION_SHEET)
        # 说明表里的示例行数和宽度必须和可填列对得上，用户能照着抄。
        width = len(writable_columns(descriptor))
        assert rows[-len(descriptor["sample_rows"]) - 1] == [column["label"] for column in writable_columns(descriptor)]
        assert all(len(row) <= width for row in rows[-len(descriptor["sample_rows"]):])
    finally:
        wb.close()


@pytest.mark.parametrize("name", sorted(TABLES))
def test_enum_columns_get_a_dropdown_in_the_data_sheet(name):
    descriptor, content = TABLES[name]()
    wb = openpyxl.load_workbook(BytesIO(content))
    try:
        data = wb[descriptor["sheet_name"]]
        rules = {str(rule.sqref): rule.formula1 for rule in data.data_validations.dataValidation}
        assert len(rules) == len(enum_columns(descriptor))
        for choices in (item[1] for item in enum_columns(descriptor)):
            assert any(formula == '"' + ",".join(choices) + '"' for formula in rules.values()), choices
    finally:
        wb.close()


def _batch_accepts(key, choice):
    """批次的中文取值要先过归一矩阵才落到内部代号，所以这一列只能按行为验，不能比集合。"""
    from core.models.enums import BATCH_PRIORITY_VALUES, READY_STATUS_VALUES, BatchPriority, ReadyStatus
    from core.services.common.normalization_matrix import (
        normalize_batch_priority_value,
        normalize_ready_status_value,
    )

    if key == "priority":
        return normalize_batch_priority_value(
            choice, default=BatchPriority.NORMAL.value, unknown_policy="passthrough") in BATCH_PRIORITY_VALUES
    return normalize_ready_status_value(
        choice, default=ReadyStatus.YES.value, unknown_policy="passthrough") in READY_STATUS_VALUES


def _accepts(name, key, choice):
    """模板下拉给的取值，解析器必须认。否则用户照着下拉选也会被拒。

    取值一律从**解析器那一侧**拿，不能从描述符那一侧拿：下拉选项本来就是按描述符生成的，
    拿描述符去比描述符是自反的，改错了也全绿（2026-09-21 试过一次，确实抓不到）。
    也不要在这里另抄一份元组——生产侧收紧取值时，那只会证明"测试里抄的那份没变"。
    """
    from core.models.workbench_material import STATUS_VALUES as material_status
    from core.models.workbench_process_file import SOURCE_VALUES
    from core.models.workbench_resource_input import _CATEGORY, _MERGE_MODE, _STATUS
    from core.models.workbench_supplier import STATUS_VALUES as supplier_status
    from core.services.workbench.facts.file_codec import _decode

    if name == "batch":
        return _batch_accepts(key, choice)
    if name in ("machine", "operator", "supplier", "op_type"):
        choice = _decode(choice, key, "xlsx")
    if name in ("machine", "operator"):
        return choice in _STATUS[name]
    if name == "supplier":
        return choice in supplier_status
    if name == "op_type":
        return choice in (_CATEGORY if key == "category" else _MERGE_MODE)
    if name == "material":
        from core.services.workbench.material.file_codec import _decode_value

        return _decode_value(choice, key, "xlsx") in material_status
    if name in ("route", "hours"):
        return choice in SOURCE_VALUES
    if name == "operator_machine":
        from core.models.workbench_relation_file import ENUMS as relation_enums

        return choice in relation_enums[key]
    from core.models.workbench_calendar_file import ENUMS as calendar_enums

    return choice in calendar_enums[key]


@pytest.mark.parametrize("name", sorted(TABLES))
def test_dropdown_choices_are_all_accepted_by_the_parser(name):
    descriptor = TABLES[name]()[0]
    by_index = {index: column for index, column in enumerate(descriptor["columns"], 1)}
    for index, choices in enum_columns(descriptor):
        key = by_index[index]["key"]
        for choice in choices:
            assert _accepts(name, key, choice), (name, key, choice)


def test_instruction_sheet_is_not_written_twice_for_the_same_table():
    """同一份说明只能有一个来源：生成器给什么，文件里就是什么。"""
    descriptor = calendar_descriptor("work_calendar")
    assert instruction_rows(descriptor)[1] == ["列名", "必填", "能填什么", "会报错的情况"]
    assert instruction_rows(descriptor) == instruction_rows(calendar_descriptor("work_calendar"))


@pytest.mark.parametrize("count,expected", ((0, 0), (1, 0), (2, 1), (5, 1)))
def test_extra_sheet_notice_only_fires_for_more_than_one_sheet(count, expected):
    assert len(extra_sheet_notice(count)) == expected


def test_no_table_declares_a_limit_the_transport_layer_would_reject():
    """哪张表都不能声明一个比传输层上限还大的 byte_limit。

    协议本身只校验 byte_limit 是正整数，没有上界。某张表要是声明了 32MB，用户传 20MB 的文件
    会先被 Flask 按 MAX_CONTENT_LENGTH 挡成 413，读取器那句"不超过 32MB"的提示永远走不到——
    表面上是上限放宽了，实际是给了个兑现不了的承诺。
    """
    from core.models.workbench_table_catalog import all_descriptors
    from web.bootstrap.app_config import Config

    over = [(descriptor["table_id"], descriptor["byte_limit"]) for descriptor in all_descriptors()
            if int(descriptor["byte_limit"]) > Config.EXCEL_MAX_UPLOAD_BYTES]
    assert not over, f"这些表声明的上限超过了传输层的 {Config.EXCEL_MAX_UPLOAD_BYTES} 字节：{over}"


def test_the_process_import_limit_comes_from_the_protocol():
    """工艺表原来自己写了一份 16MB，注释说 Matches Config，可两边都是独立字面量。"""
    from core.models.workbench_process_file import IMPORT_BYTE_LIMIT
    from core.models.workbench_table_descriptor import DEFAULT_IMPORT_BYTE_LIMIT
    from web.bootstrap.app_config import Config

    assert IMPORT_BYTE_LIMIT == DEFAULT_IMPORT_BYTE_LIMIT == Config.EXCEL_MAX_UPLOAD_BYTES
