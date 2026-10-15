"""Saved runtime requirements and actual type events retain their API provenance."""

import csv
import io
import json
from copy import deepcopy
from datetime import datetime

import pytest
from openpyxl import load_workbook

from core.services.workbench.facts.candidate_store import CandidateStore
from core.services.workbench.facts.operation_effort_evidence import archived_standard_effort
from core.services.workbench.facts.resource_intervals import task_interval
from core.services.workbench.run.decision_source import DecisionSource
from core.services.workbench.run.worker import WorkbenchRunWorker
from tests.workbench.decision_history_support import clone_run
from tests.workbench.legacy_run_archive_support import install_different_work_archive
from tests.workbench.run_candidate_adoption_support import INTENT, service
from tests.workbench.run_candidate_support import BASE, api, compute
from tests.workbench.run_candidate_support import candidate_case as _candidate_case  # noqa: F401
from tests.workbench.run_jobs_support import service as run_service
from tests.workbench.test_material_stage_release import add_requirement
from tests.workbench.test_run_decision_api import decision
from tests.workbench.test_run_decision_impact_api import detail, operation_ref
from tests.workbench.test_run_decision_resource_api import add_pair

DAY = "2026-09-09"
WINDOW = {"range_start": DAY + "T08:00:00", "range_end": DAY + "T16:00:00"}


def run(case, key, settings=None):
    input_ref = case.preflight(settings)
    service = run_service(case.conn)
    preview = service.preview(input_ref)["write_context"]
    assert preview["capabilities"]["scheduling.run"] is True, preview["blocked_reasons"]
    accepted = service.accept(input_ref, preview["write_token"], key)
    result = WorkbenchRunWorker(case.conn).execute(accepted["run_ref"])
    return result["run_ref"], [row["candidate_ref"] for row in result["candidates"]]


def receipt(case, run_ref):
    return json.loads(case.conn.execute("SELECT result_json FROM WorkbenchRunReceipts WHERE run_ref=?", (run_ref,)).fetchone()[0])


def by_operation(document):
    return {row["operation_ref"]: row for row in document["data"]["items"]}


def adopt(case, candidate_ref, key):
    adopter = service(case.conn)
    preview = adopter.preview(candidate_ref)
    assert preview["validation"]["can_adopt"] is True, preview
    return adopter.adopt(candidate_ref, preview["write_context"]["write_token"], key, INTENT)["data"]["official_plan"]


def test_real_worker_saves_one_requirement_used_by_all_candidates_and_workspaces(candidate_case):
    case = candidate_case
    case.conn.execute("UPDATE Batches SET quantity=6 WHERE batch_id='B1'")
    case.conn.execute("UPDATE BatchOperations SET setup_hours=2,unit_hours=0.5 WHERE id=?", (case.op_id,))
    case.conn.commit()
    run_ref, refs = compute(case)
    saved = receipt(case, run_ref)["operation_effort"]
    assert saved["effort_basis"] == "runtime_requirement_v1" and len(saved["items"]) == 1
    item = saved["items"][0]
    ref = operation_ref(case, case.op_id)
    assert item["operation_ref"] == ref and item["quantity_basis"] == "batch" and item["target_quantity"] == 6
    assert (item["setup_standard_hours"], item["processing_standard_hours"], item["required_standard_hours"]) == (2, 3, 5)
    assert item["requirement_role"] == "new_schedule"
    client, _ = api(case)
    query = dict(WINDOW, selected_candidate_ref=refs[0], reference_candidate_ref=refs[0])
    with CandidateStore(case.conn).snapshot():
        source = DecisionSource(CandidateStore(case.conn), run_ref, **query)
        for method in (source.preparation.groups, source.preparation.rows, source.preparation.summary,
                       source.changeovers.projection, source.changeovers.metrics):
            assert method("reference") is method(refs[0])
    overview = decision(client, run_ref, **query)["data"]
    assert len(overview["metric_definitions"]) == 52
    definitions = {row["key"]: row for row in overview["metric_definitions"]}
    assert definitions["scheduled_setup_hours"]["label"] == "新安排准备标准工时"
    assert definitions["scheduled_setup_hours"]["improvement"] == "neutral"
    assert definitions["type_changeover_count"]["label"] == "工种切换次数"
    assert definitions["changeover_count"]["label"] == "算法工种切换"
    for candidate in overview["candidates"]:
        assert candidate["metrics"]["scheduled_setup_hours"]["value"] == 2
        delta = candidate["deltas"]["scheduled_setup_hours"]
        assert delta["value"] == 0 and delta["direction"] == "equal"
    preparation = detail(client, run_ref, "preparation", **query)
    projected = preparation["data"]["items"][0]
    assert projected["operation_ref"] == ref and projected["after"]["allocation_state"] == "new_schedule"
    for field in ("setup_standard_hours", "processing_standard_hours", "required_standard_hours", "target_quantity"):
        assert projected["after"][field] == item[field]
    assert projected["after"]["standard"]["standard_total_hours"] == 5
    assert projected["after"]["evidence_source"]["run_ref"] == run_ref
    operation = detail(client, run_ref, "operations", **query)["data"]["items"][0]
    assert operation["effort"]["standard_total_hours"] == 5
    assert operation["runtime_effort"]["after"]["required_standard_hours"] == 5
    workspace = client.get(BASE + "/candidates/" + refs[0] + "/workspace")
    assert workspace.status_code == 200, workspace.get_json()
    task = workspace.get_json()["data"]["tasks"][0]
    assert (datetime.fromisoformat(task["end"]) - datetime.fromisoformat(task["start"])).total_seconds() / 3600 == 5
    assert task["operation_effort"]["required_standard_hours"] == 5
    case.conn.execute("UPDATE BatchOperations SET setup_hours=9,unit_hours=2 WHERE id=?", (case.op_id,))
    case.conn.commit()
    reread = detail(client, run_ref, "preparation", **dict(query, snapshot_ref=preparation["meta"]["snapshot_ref"]))
    assert reread["data"] == preparation["data"]


def test_unplanned_and_material_deferred_work_remain_neutral_and_old_runtime_stays_missing(candidate_case, monkeypatch):
    case = candidate_case
    case.conn.execute("UPDATE Batches SET quantity=1,due_date=? WHERE batch_id='B1'", (DAY,))
    case.conn.execute("UPDATE BatchOperations SET setup_hours=2,unit_hours=3 WHERE id=?", (case.op_id,))
    case.batch("B2", quantity=1, due_date="2026-09-10")
    missing_id = case.operation("B2", setup_hours=3, unit_hours=4)
    case.config(dispatch_mode="batch_order", sort_strategy="due_date_first")
    with monkeypatch.context() as archive_patch:
        install_different_work_archive(archive_patch)
        run_ref, refs = compute(case, case.settings("B1", "B2", end_date=DAY))
    client, _ = api(case)
    catalog = decision(client, run_ref)["data"]["candidates"]
    reference = next(row["candidate_ref"] for row in catalog if row["kind"] == "critical_chain")
    query = dict(WINDOW, selected_candidate_ref=refs[0], reference_candidate_ref=reference)
    document = detail(client, run_ref, "preparation", **query)
    row = by_operation(document)[operation_ref(case, missing_id)]
    assert row["after"]["allocation_state"] == "unplanned"
    assert row["after"]["setup_standard_hours"] == 3
    summary = document["data"]["summary"]
    assert summary["after"]["pending_setup_hours"]["value"] == 5
    assert summary["after"]["scheduled_setup_hours"]["value"] == 2
    assert summary["after"]["unplanned_setup_hours"]["value"] == 3
    delta = summary["deltas"]["scheduled_setup_hours"]
    assert delta["value"] == -1 and delta["state"] == "different_coverage" and delta["direction"] == "neutral"
    assert row["drill"]["selected"]["operation_ref"] == operation_ref(case, missing_id)
    # A historical run without the ledger retains independently saved standards.
    case.history_refs = {"run_ref": run_ref, "candidate_ref": refs[0]}
    accepted = case.conn.execute("SELECT accepted_at FROM WorkbenchRunJobs WHERE run_ref=?", (run_ref,)).fetchone()[0]
    historical_run, _ = clone_run(case, accepted,
        change_artifact=lambda artifact: artifact.update(analysis_contract_version="decision-resource-v3"),
        change_receipt=lambda result: result.pop("operation_effort"))
    historical_refs = [row[0] for row in case.conn.execute(
        "SELECT candidate_ref FROM WorkbenchRunCandidates WHERE run_ref=? ORDER BY sequence", (historical_run,))]
    historical_query = dict(query, selected_candidate_ref=historical_refs[refs.index(query["selected_candidate_ref"])],
                            reference_candidate_ref=historical_refs[refs.index(reference)])
    historical = detail(client, historical_run, "preparation", **historical_query)
    legacy = by_operation(historical)[operation_ref(case, case.op_id)]["after"]
    assert legacy["required_standard_hours"] is legacy["setup_standard_hours"] is None
    assert legacy["standard"]["setup_standard_hours"] == 2 and legacy["standard"]["standard_total_hours"] == 5
    assert legacy["execution_target_quantity"] == 1 and legacy["data_gaps"]
    stage_id = case.operation(seq=2, setup_hours=1, unit_hours=0.25)
    case.conn.commit()
    _stage_original_run, stage_original_refs = run(case, "prep-stage-before-material-run")
    adopt(case, stage_original_refs[0], "prep-stage-before-material-adopt")
    add_requirement(case, stage_id, [])
    stage_run, stage_refs = run(case, "prep-material-stage-real", case.settings(material_strategy="stage"))
    stage_query = dict(WINDOW, selected_candidate_ref=stage_refs[0], comparison_mode="admission_official")
    staged = detail(client, stage_run, "preparation", **stage_query)
    deferred = by_operation(staged)[operation_ref(case, stage_id)]["after"]
    assert deferred["allocation_state"] == "deferred"
    assert staged["data"]["summary"]["after"]["deferred_setup_hours"]["value"] is None
    assert deferred["standard"]["setup_standard_hours"] == 1
    stage_delta = staged["data"]["summary"]["deltas"]["scheduled_setup_hours"]
    assert stage_delta["value"] == -1 and stage_delta["state"] == "different_coverage" and stage_delta["direction"] == "neutral"


def test_formal_requirements_keep_last_source_missingness_and_retained_targets(candidate_case):
    case = candidate_case
    case.conn.execute("UPDATE Batches SET quantity=6 WHERE batch_id='B1'")
    case.conn.execute("UPDATE BatchOperations SET setup_hours=2,unit_hours=0.5 WHERE id=?", (case.op_id,))
    case.batch("B2", quantity=6)
    second_id = case.operation("B2", setup_hours=1, unit_hours=0.5)
    case.conn.commit()
    original_run, original_refs = compute(case, case.settings("B1", "B2"))
    old_plan = adopt(case, original_refs[0], "prep-original-real-adopt")
    # This authentic legacy partial-version shape leaves B2 at its actual v1
    # adoption. B1's newer manual source contains no saved runtime or standard.
    legacy_version = old_plan["version"] + 1
    case.conn.execute("INSERT INTO Schedule(op_id,machine_id,operator_id,start_time,end_time,lock_status,version) "
                      "SELECT op_id,machine_id,operator_id,start_time,end_time,lock_status,? FROM Schedule WHERE version=? AND op_id=?",
                      (legacy_version, old_plan["version"], case.op_id))
    case.conn.execute("INSERT INTO ScheduleHistory(version,strategy,batch_count,op_count,result_status,result_summary) "
                      "VALUES(?,'manual',1,1,'success','{}')", (legacy_version,))
    case.conn.execute("UPDATE Batches SET quantity=20 WHERE batch_id='B2'")
    case.conn.execute("UPDATE BatchOperations SET unit_hours=0.15 WHERE id=?", (second_id,))
    case.config(freeze_window_enabled="yes", freeze_window_days=5)
    final_run, final_refs = run(case, "prep-last-source-frozen-run", case.settings("B1", "B2"))
    client, _ = api(case)
    query = {**WINDOW, "range_end": "2026-09-12T16:00:00", "comparison_mode": "admission_official", "selected_candidate_ref": final_refs[0]}
    document = detail(client, final_run, "preparation", **query)
    rows = by_operation(document)
    first, second = rows[operation_ref(case, case.op_id)], rows[operation_ref(case, second_id)]
    assert document["data"]["baseline"]["version"] == legacy_version
    assert first["before"]["required_standard_hours"] is first["before"]["setup_standard_hours"] is None
    assert first["before"]["standard"]["standard_total_hours"] is None
    assert first["before"]["standard"]["state"] == "unavailable"
    assert second["before"]["required_standard_hours"] == 4
    assert second["before"]["target_quantity"] == 6
    assert second["before"]["execution_target_quantity"] == second["before"]["remaining_quantity"] == 20
    assert second["before"]["standard"]["unit_standard_hours"] == 0.5
    assert second["before"]["evidence_source"]["run_ref"] == original_run
    assert second["before"]["evidence_source"]["plan_ref"] == old_plan["plan_ref"]
    assert first["after"]["allocation_state"] == "retained" and second["after"]["allocation_state"] == "new_schedule"
    assert first["after"]["required_standard_hours"] is first["after"]["setup_standard_hours"] is None
    assert first["after"]["target_quantity"] is None and first["after"]["execution_target_quantity"] == 6
    assert first["after"]["data_gaps"]
    assert second["after"]["required_standard_hours"] == 4 and second["after"]["target_quantity"] == 20
    assert second["after"]["execution_target_quantity"] == second["after"]["remaining_quantity"] == 20
    assert second["after"]["standard"]["unit_standard_hours"] == 0.15 and second["after"]["standard"]["standard_total_hours"] == 4
    assert document["data"]["summary"]["role_totals"]["after"]["retained"]["count"] == 1
    artifact = json.loads(case.conn.execute("SELECT artifact_json FROM WorkbenchRunCandidates WHERE candidate_ref=?", (final_refs[0],)).fetchone()[0])
    trace = next(row for row in artifact["results"] if row["op_id"] == case.op_id)
    if "op_type_id" in trace and trace["op_type_id"] is None:
        response = client.get(BASE + "/candidates/" + final_refs[0] + "/workspace")
        assert response.status_code == 200, response.get_json()
        task = next(row for row in response.get_json()["data"]["tasks"] if row["operation_ref"] == first["operation_ref"])
        assert task["process_type_ref"] is task["process_type_label"] is None
    workspace = detail(client, final_run, "reference-workspace", **query)
    work = {row["operation_ref"]: row["operation_effort"] for row in workspace["data"]["tasks"]}
    assert work[second["operation_ref"]]["required_standard_hours"] == 4
    assert work[first["operation_ref"]]["required_standard_hours"] is None


def test_preparation_pages_filters_snapshot_and_all_filtered_exports(candidate_case):
    case = candidate_case
    add_pair(case)
    case.conn.execute("UPDATE BatchOperations SET setup_hours=0.2,unit_hours=0.05 WHERE id=?", (case.op_id,))
    ids = [case.op_id]
    for sequence in range(2, 28):
        machine, operator = ("M1", "O1") if sequence <= 25 else ("M2", "O2")
        ids.append(case.operation(seq=sequence, machine_id=machine, operator_id=operator, setup_hours=0.2, unit_hours=0.05))
    case.conn.commit()
    run_ref, refs = compute(case)
    missing_ref = operation_ref(case, ids[24])
    case.history_refs = {"run_ref": run_ref, "candidate_ref": refs[0]}
    accepted = case.conn.execute("SELECT accepted_at FROM WorkbenchRunJobs WHERE run_ref=?", (run_ref,)).fetchone()[0]

    def change_receipt(result):
        result["operation_effort"]["items"] = [row for row in result["operation_effort"]["items"] if row["operation_ref"] != missing_ref]

    run_ref, _ = clone_run(case, accepted, change_receipt=change_receipt)
    refs = [row[0] for row in case.conn.execute(
        "SELECT candidate_ref FROM WorkbenchRunCandidates WHERE run_ref=? ORDER BY sequence", (run_ref,))]
    assert receipt(case, run_ref)["selection"]["generation_policy_version"] == "unified_search_v1"
    client, _ = api(case)
    query = dict(WINDOW, selected_candidate_ref=refs[0], reference_candidate_ref=refs[0], machine_ref=case.ref("machine", "M1"),
                 sort="after_setup_hours", order="asc")
    initial = detail(client, run_ref, "preparation", **query)
    token = initial["meta"]["snapshot_ref"]
    page = initial["data"]["page"]
    assert (page["number"], page["size"], page["total"], page["has_more"]) == (1, 20, 25, True)
    following = detail(client, run_ref, "preparation", **dict(query, page=2, snapshot_ref=token))
    wide = detail(client, run_ref, "preparation", **dict(query, size=50, snapshot_ref=token))
    rows = initial["data"]["items"] + following["data"]["items"]
    assert rows == wide["data"]["items"] and len(rows) == 25
    assert [row["operation_ref"] for row in rows[:-1]] == sorted(row["operation_ref"] for row in rows[:-1])
    assert rows[-1]["operation_ref"] == missing_ref and rows[-1]["after"]["setup_standard_hours"] is None
    assert following["data"]["summary"] == initial["data"]["summary"]
    complete_query = {key: value for key, value in query.items() if key != "machine_ref"}
    unfiltered = detail(client, run_ref, "preparation", **complete_query)
    assert unfiltered["data"]["summary"] == initial["data"]["summary"]
    machines = initial["data"]["filter_options"]["machines"]
    assert {row["ref"] for row in machines} == {case.ref("machine", "M1"), case.ref("machine", "M2")}
    path = BASE + "/runs/" + run_ref + "/decision/preparation"
    for change in ({"machine_ref": case.ref("machine", "M2")}, {"allocation_state": "unplanned"}):
        stale = client.get(path, query_string=dict(query, snapshot_ref=token, **change))
        assert stale.status_code == 409 and stale.get_json()["error"]["code"] == "snapshot_stale"
    export_path = BASE + "/runs/" + run_ref + "/decision/export"
    for fmt in ("csv", "xlsx"):
        response = client.get(export_path, query_string=dict(query, snapshot_ref=token, section="preparation", format=fmt))
        assert response.status_code == 200, response.get_json()
        assert int(response.headers["X-Workbench-Row-Count"]) == 25
        if fmt == "csv":
            exported = list(csv.DictReader(io.StringIO(response.data.decode("utf-8-sig"))))
        else:
            book = load_workbook(io.BytesIO(response.data), read_only=True, data_only=False)
            try:
                iterator = book["准备需求"].iter_rows()
                headers = [cell.value for cell in next(iterator)]
                cells = [dict(zip(headers, row)) for row in iterator]
                exported = [{name: cell.value for name, cell in row.items()} for row in cells]
                for actual, expected in zip(cells, rows):
                    for header, field in (("所选准备标准工时", "setup_standard_hours"), ("所选加工标准工时", "processing_standard_hours"),
                                          ("所选所需标准工时", "required_standard_hours")):
                        assert actual[header].value == pytest.approx(expected["after"][field]) if expected["after"][field] is not None else actual[header].value is None
                        if expected["after"][field] is not None:
                            assert actual[header].data_type == "n" and type(actual[header].value) in (int, float)
                    assert actual["所选执行目标量"].value == 3 and actual["所选执行目标量"].data_type == "n"
                    assert actual["所选归档标准总工时"].value == pytest.approx(0.35)
            finally:
                book.close()
        assert [row["工序永久编号"] for row in exported] == [row["operation_ref"] for row in rows]
        for actual, expected in zip(exported, rows):
            for header, field in (("所选准备标准工时", "setup_standard_hours"), ("所选加工标准工时", "processing_standard_hours"),
                                  ("所选所需标准工时", "required_standard_hours")):
                value = expected["after"][field]
                if value is None:
                    assert actual[header] == ("" if fmt == "csv" else None)
                else:
                    assert float(actual[header]) == pytest.approx(value)
    response = client.get(export_path, query_string=dict(query, snapshot_ref=token, section="all", filter_section="preparation", format="xlsx"))
    assert response.status_code == 200, response.get_json()
    book = load_workbook(io.BytesIO(response.data), read_only=True, data_only=True)
    try:
        assert book.sheetnames == ["方案总览", "批次影响", "工序变化", "资源影响", "覆盖工序", "准备需求", "工种切换",
                                  "推荐说明", "工序安排", "剩余需求", "共同计算", "目标预算", "进步事件", "搜索动作", "搜索阶段",
                                  "评价范围", "范围成员", "真实桥接", "指标计算方式", "比较范围"]
        assert len(list(book["准备需求"].iter_rows(values_only=True))) - 1 == 25
        assert int(response.headers["X-Workbench-Row-Count"]) == sum(len(list(book[name].iter_rows(values_only=True))) - 1 for name in book.sheetnames[:-2])
    finally:
        book.close()


def test_changeover_api_keeps_type_identity_points_predecessors_and_partial_overlap(candidate_case, monkeypatch):
    case = candidate_case
    for code in ("T2", "T3"):
        case.conn.execute("INSERT INTO OpTypes(op_type_id,name) VALUES (?,?)", (code, "精车" if code == "T2" else "磨削"))
        case.conn.execute("INSERT INTO MachineOpTypes(machine_id,op_type_id) VALUES ('M1',?)", (code,))
    second_id = case.operation(seq=2)
    third_id = case.operation(seq=3, op_type_id="T2", op_type_name="精车")
    point_id = case.operation(seq=4, op_type_id="T3", op_type_name="磨削", setup_hours=0, unit_hours=0)
    case.conn.commit()
    run_ref, refs = compute(case)
    first_ref, second_ref, third_ref, point_ref = [operation_ref(case, item) for item in (case.op_id, second_id, third_id, point_id)]
    original_tasks = DecisionSource.analysis_tasks
    overlapping = [False]

    def copied_arrangements(source, candidate_ref):
        rows = original_tasks(source, candidate_ref)
        if candidate_ref != refs[0]:
            return rows
        copied = deepcopy(rows)
        times = {first_ref: ("08:00:00", "09:00:00", "车削"), second_ref: ("09:00:00", "10:00:00", "精车"),
                 third_ref: ("09:30:00" if overlapping[0] else "10:00:00", "11:00:00", "精车"),
                 point_ref: ("09:30:00", "09:30:00", "磨削")}
        for row in copied:
            start, end, label = times[row["operation_ref"]]
            row.update(start=DAY + "T" + start, end=DAY + "T" + end, process_type_label=label)
        source._task_intervals[candidate_ref] = {row["row_ref"]: task_interval(row) for row in copied}
        return copied

    monkeypatch.setattr(DecisionSource, "analysis_tasks", copied_arrangements)
    client, _ = api(case)
    query = {"selected_candidate_ref": refs[0], "reference_candidate_ref": refs[-1], "event_side": "selected",
             "range_start": DAY + "T08:00:00", "range_end": DAY + "T12:00:00"}
    document = detail(client, run_ref, "changeovers", **query)
    count = document["data"]["summary"]["after"]["type_changeover_count"]
    assert count["value"] == 1 and count["state"] == "complete"
    assert document["data"]["summary"]["after"]["changed_machine_count"]["value"] == 1
    event, = document["data"]["items"]
    assert (event["previous_operation_ref"], event["next_operation_ref"]) == (second_ref, third_ref)
    assert event["previous"]["type_ref"] == case.ref("op_type", "T1") and event["next"]["type_ref"] == case.ref("op_type", "T2")
    assert event["previous"]["type_label"] == event["next"]["type_label"] == "精车"
    export_path = BASE + "/runs/" + run_ref + "/decision/export"
    for fmt in ("csv", "xlsx"):
        response = client.get(export_path, query_string=dict(query, snapshot_ref=document["meta"]["snapshot_ref"], section="changeovers", format=fmt))
        assert response.status_code == 200, response.get_json()
        assert int(response.headers["X-Workbench-Row-Count"]) == 1
        if fmt == "csv":
            exported = list(csv.DictReader(io.StringIO(response.data.decode("utf-8-sig"))))
        else:
            book = load_workbook(io.BytesIO(response.data), read_only=True, data_only=False)
            try:
                iterator = book["工种切换"].iter_rows()
                headers = [cell.value for cell in next(iterator)]
                cells = [dict(zip(headers, row)) for row in iterator]
                exported = [{name: cell.value for name, cell in row.items()} for row in cells]
                assert cells[0]["设备序列位置"].data_type == "n"
                assert cells[0]["事件编号"].data_type == cells[0]["前序安排永久编号"].data_type == "s"
            finally:
                book.close()
        assert len(exported) == 1
        assert (exported[0]["前序工序永久编号"], exported[0]["后序工序永久编号"]) == (second_ref, third_ref)
        assert (exported[0]["前序安排永久编号"], exported[0]["后序安排永久编号"]) == (event["previous_row_ref"], event["next_row_ref"])
        assert exported[0]["事件编号"] == event["event_ref"] and exported[0]["事件侧"] == "selected"
    for first, last in (("08:00:00", "09:00:00"), ("09:00:00", "10:00:00")):
        empty = detail(client, run_ref, "changeovers", **dict(query, range_start=DAY + "T" + first, range_end=DAY + "T" + last))
        assert empty["data"]["summary"]["after"]["type_changeover_count"]["value"] == 0 and empty["data"]["items"] == []
    narrow_query = dict(query, range_start=DAY + "T10:00:00", range_end=DAY + "T11:00:00")
    narrow = detail(client, run_ref, "changeovers", **narrow_query)
    witness, = narrow["data"]["items"]
    assert witness["event_ref"] == event["event_ref"] and witness["predecessor_outside_window"] is True
    assert witness["drill"]["previous"]["operation_ref"] == second_ref and witness["drill"]["next"]["operation_ref"] == third_ref
    assert witness["drill"]["next"]["candidate_ref"] == refs[0]
    repeated = detail(client, run_ref, "changeovers", **dict(narrow_query, snapshot_ref=narrow["meta"]["snapshot_ref"], size=50))
    assert repeated["data"]["items"] == narrow["data"]["items"]
    stale = client.get(BASE + "/runs/" + run_ref + "/decision/changeovers",
                       query_string=dict(narrow_query, snapshot_ref=narrow["meta"]["snapshot_ref"], event_side="reference"))
    assert stale.status_code == 409 and stale.get_json()["error"]["code"] == "snapshot_stale"
    overlapping[0] = True
    partial = detail(client, run_ref, "changeovers", **query)
    count = partial["data"]["summary"]["after"]["type_changeover_count"]
    assert count["value"] is None and count["known_subtotal"] == 1 and count["state"] == "partial"
    assert partial["data"]["items"][0]["overlap"] is True


def test_archived_standard_effort_keeps_large_integer_display_and_exact_total():
    setup, unit = (1 << 53) + 1, (1 << 53) + 3
    operation = {"source": "internal", "setup_hours": setup, "unit_hours": unit}
    execution = {"target_quantity": 3}
    effort = archived_standard_effort(operation, execution)
    assert effort["setup_standard_hours"] == str(setup) and effort["unit_standard_hours"] == str(unit)
    assert type(effort["standard_total_hours"]) is int and effort["standard_total_hours"] == setup + unit * 3
    assert effort["state"] == "complete" and effort["reason"] is None

    missing = archived_standard_effort(dict(operation, unit_hours=None), execution)
    assert missing["setup_standard_hours"] == str(setup) and missing["unit_standard_hours"] is None
    assert missing["standard_total_hours"] is None and missing["state"] == "unavailable"
    assert missing["reason"]["code"] == "standard_effort_missing"
