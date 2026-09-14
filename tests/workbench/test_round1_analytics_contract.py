"""Round-one report boundaries preserve source rows, unknowns and full cohorts."""

import csv
import io
from copy import deepcopy
from datetime import datetime, timezone

import openpyxl
import pytest

from core.models.workbench_command import WorkbenchCommandRejected
from core.models.workbench_report import ReportPage, ReportScope
from core.models.workbench_report_catalog import ReportCatalogScope
from core.services.workbench.report_catalog import catalog_workspace
from core.services.workbench.review_legacy import legacy_review
from tests.workbench.plan_read_support import assert_error
from tests.workbench.report_api_support import event
from tests.workbench.report_api_support import report_api as _report_api


@pytest.mark.parametrize("start,end", [
    (None, "2026-09-02"), ("2026-09-02", None), ("", "2026-09-02"),
    ("2026-09-02", ""), ("2026-02-30", "2026-03-01"),
    ("2026-09-03", "2026-09-02"), ("2026-9-2", "2026-09-03"),
    (True, "2026-09-02"), ("2026-09-02", 20260903),
])
def test_finish_date_pair_is_not_dropped_or_coerced(start, end):
    with pytest.raises(WorkbenchCommandRejected) as caught:
        ReportScope(plan_finish_date_from=start, plan_finish_date_to=end)
    assert caught.value.code == "invalid_input" and caught.value.status == 400


def test_optional_and_same_day_finish_scope_keep_original_identity():
    assert ReportScope().scope()["plan_finish_date_to"] is None
    scope = ReportScope(plan_ref="a" * 48, plan_finish_date_from="2026-09-02", plan_finish_date_to="2026-09-02")
    assert scope.scope()["plan_ref"] == "a" * 48
    assert scope.plan_finish_date_from == scope.plan_finish_date_to == "2026-09-02"


@pytest.mark.parametrize("topic,sort", [
    ("delivery", "batch_label"), ("records", "event_time"), ("quality", "batch_label"),
    ("machines", "resource_label"), ("people", "resource_label"),
])
def test_default_sort_is_used_only_for_an_absent_parameter(report_api, topic, sort):
    data = report_api.read(topic=topic)["data"]
    assert data["page"]["sort"] == [{"field": sort, "direction": "asc"}]
    assert_error(report_api.get(topic=topic, sort=""), "invalid_input", 400)
    assert_error(report_api.get(topic=topic, sort="unknown"), "invalid_input", 400)


@pytest.mark.parametrize("kwargs", [{"number": True}, {"size": True}, {"number": 1.5}, {"size": "20"}])
def test_pagination_keeps_exact_integer_contract(kwargs):
    with pytest.raises(WorkbenchCommandRejected) as caught:
        ReportPage(**kwargs)
    assert caught.value.code == "invalid_input" and caught.value.status == 400


def _legacy_event(kind, time, task="task-a", **extra):
    return {"recorded_against_task_ref": task, "event_type": kind, "event_time": time,
            "reason_code": "other", "severity": "low", "impact_minutes": None,
            "handling_status": "new", "suggest_reschedule": 0,
            "actual_machine_ref": None, "actual_operator_ref": None, **extra}


@pytest.mark.parametrize("invalid", [None, "broken", "2026-09-02T08:30:00+08:00", True, 8,
    datetime(2026, 9, 2, 8, 30, tzinfo=timezone.utc)])
@pytest.mark.parametrize("kind", ["pause", "resume", "exception"])
def test_invalid_legacy_time_keeps_unknown_metrics_and_all_source_rows(kind, invalid):
    projection = {"legacy_facts": [
        _legacy_event("pause", "2026-09-02T08:00:00"),
        _legacy_event(kind, invalid),
        _legacy_event("exception", "2026-09-02T09:00:00"),
    ], "data_gaps": []}
    before = deepcopy(projection)
    result = legacy_review(projection)
    assert all(value is None for value in result.values())
    assert projection == before and len(projection["legacy_facts"]) == 3


@pytest.mark.parametrize("gap", ["invalid_legacy_sequence", "legacy_identity_unresolved"])
def test_legacy_invalid_gap_is_not_reinterpreted_as_valid_history(gap):
    projection = {"legacy_facts": [_legacy_event("exception", "2026-09-02T08:00:00")],
                  "data_gaps": [{"code": gap}]}
    before = deepcopy(projection)
    assert all(value is None for value in legacy_review(projection).values())
    assert projection == before


def test_legacy_duration_groups_tasks_and_selects_latest_valid_exception():
    projection = {"legacy_facts": [
        _legacy_event("pause", "2026-09-01T23:50:00"),
        _legacy_event("resume", "2026-09-02T00:20:00"),
        _legacy_event("pause", "2026-09-02T08:00:00", "task-b"),
        _legacy_event("finish", "2026-09-02T08:10:00", "task-b"),
        _legacy_event("exception", "2026-09-02T10:00:00", severity="high", impact_minutes=0),
        _legacy_event("exception", "2026-09-02T09:00:00", impact_minutes=5),
    ], "data_gaps": []}
    before = deepcopy(projection)
    result = legacy_review(projection)
    assert result["pause_duration_minutes"] == 40
    assert result["exception_reason"] == "其他" and result["exception_severity"] == "严重"
    assert result["exception_impact_minutes"] == 0
    assert projection == before


def test_open_or_reversed_pause_is_unknown_not_zero():
    rows = [_legacy_event("pause", "2026-09-02T08:00:00")]
    assert legacy_review({"legacy_facts": rows, "data_gaps": []})["pause_duration_minutes"] is None
    rows.append(_legacy_event("resume", "2026-09-02T07:59:00"))
    assert legacy_review({"legacy_facts": rows, "data_gaps": []})["pause_duration_minutes"] is None
    assert legacy_review({"legacy_facts": [], "data_gaps": []})["pause_duration_minutes"] is None


def test_catalog_resource_pagination_preserves_unknowns_and_original_export_rows():
    machines = [{"machine_id": str(index), "machine_name": f"Machine {index:02d}",
                 "hours": index, "task_count": 1, "capacity_hours": None if index == 0 else 8,
                 "utilization": None if index == 0 else 0 if index == 1 else .5,
                 "source_marker": index} for index in range(12)]
    facts = {"scope": ReportCatalogScope("utilization"), "plan": {"plan_ref": "a" * 48},
             "report": {"machines": machines, "operators": [], "report_degraded": True},
             "identities": {"machine": {str(index): {"ref": format(index, "048x")} for index in range(12)}}}
    before = deepcopy(facts)
    data, ordered, selected = catalog_workspace(None, facts, {}, ReportPage(size=10, sort="utilization_percent", direction="desc"))
    assert data["page"]["total"] == 12 and len(data["rows"]) == 10
    assert data["summary"] == {"rows": 12}
    assert len(ordered) == len(selected) == 12
    assert ordered[-2]["utilization_percent"] == 0 and ordered[-1]["utilization_percent"] is None
    assert ordered[-1]["capacity_hours"] is None
    assert [row["source_marker"] for row in selected] == list(range(2, 12)) + [1, 0]
    assert all(row["resource_kind"] == "machine" for row in selected)
    assert data["plan"] == facts["plan"] and len(data["data_gaps"]) == 2
    assert data["scope"]["selection"] == "schedule_window_overlap"
    assert facts == before


def test_catalog_overdue_keeps_invalid_and_unscheduled_rows_in_export_selection():
    raw = [
        {"batch_id": "late", "due_date": "2026-09-01", "finish_time": "2026-09-02T10:00:00"},
        {"batch_id": "due-bad", "due_date": "bad", "finish_time": None},
        {"batch_id": "time-bad", "due_date": "2026-09-01", "finish_time": "bad", "schedule_row_count": 1},
        {"batch_id": "unscheduled", "due_date": "2026-09-01", "finish_time": None},
        {"batch_id": "on-time", "due_date": "2026-09-03", "finish_time": "2026-09-02T10:00:00"},
    ]
    identities = {row["batch_id"]: {"ref": format(index, "048x")} for index, row in enumerate(raw)}
    facts = {"scope": ReportCatalogScope("overdue"), "plan": {"plan_ref": "a" * 48},
             "raw": raw, "report": None, "identities": {"batch": identities}}
    before = deepcopy(facts)
    data, ordered, selected = catalog_workspace(None, facts, {"as_of": "2026-09-10T12:00:00"}, ReportPage())
    assert data["summary"] == {"rows": 4}
    assert [row["batch_label"] for row in ordered] == [row["batch_id"] for row in selected]
    assert {row["bucket"] for row in selected} == {
        "scheduled_overdue", "due_date_invalid", "schedule_time_invalid", "unscheduled_overdue",
    }
    assert ordered[0]["batch_label"] == "due-bad" and ordered[0]["delay_hours"] is None
    assert data["scope"]["selection"] == "all_plan_batches" and facts == before


@pytest.mark.parametrize("format_name", ["csv", "xlsx"])
def test_invalid_future_fact_stays_in_cohort_charts_and_download(report_api, format_name):
    with report_api.db() as conn:
        event(conn, 3, "finish", "2099-09-02T09:00:00", quantity_done=1)
    before = report_api.state()
    reading = report_api.read(size=10)
    data, token = reading["data"], reading["meta"]["snapshot_ref"]
    assert data["summary"]["operations"] == data["summary"]["due"] == 23
    assert data["summary"]["confirmed_due"] == 2
    assert data["summary"]["completion_rate"] == 2 / 23
    assert sum(row["count"] for row in data["charts"]["finish"]) == 2
    assert sum(row["count"] for row in data["charts"]["aging"]) == 21
    assert data["charts"]["trend"][-1]["planned"] == 23
    assert data["charts"]["trend"][-1]["actual"] == 2
    records = report_api.read(topic="records", snapshot_ref=token)["data"]
    assert records["charts"] == data["charts"] and records["summary"] == data["summary"]
    invalid = next(row for row in records["rows"] if row["event_time_raw"] == "2099-09-02 09:00:00")
    assert invalid["data_quality"] == "invalid" and invalid["effective_processing_hours"] is None
    response = report_api.get("/export", size=10, snapshot_ref=token, format=format_name)
    assert response.status_code == 200 and response.headers["X-Workbench-Row-Count"] == "23"
    if format_name == "csv":
        rows = list(csv.reader(io.StringIO(response.data.decode("utf-8-sig"))))
    else:
        book = openpyxl.load_workbook(io.BytesIO(response.data), read_only=True)
        try:
            rows = list(book["范围全部结果"].iter_rows(values_only=True))
            exported_snapshots = [row[1] for row in book["范围与计算方式"].iter_rows(values_only=True) if row[0] == "数据版本编号"]
            assert exported_snapshots == [token]
        finally:
            book.close()
    assert len(rows) == 24
    assert report_api.state() == before
