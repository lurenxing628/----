"""Physical reader contracts that domain codecs must preserve."""

from functools import partial
from io import BytesIO

import openpyxl
import pytest

from core.errors import ValidationError
from core.services.workbench.facts.file_codec import file_error, read_resource_file
from core.services.workbench.facts.file_source import source_rows
from core.services.workbench.material.file_codec import read_material_file
from core.services.workbench.process.file_codec import decode_process_file
from core.services.workbench.resource.calendar_files.file_codec import read_calendar_file
from core.services.workbench.resource.relation_files.file_codec import read_relation_file


def test_csv_row_numbers_identify_multiline_record_start():
    content = b'code,label\r\nA,"first\r\nsecond"\r\nB,third\r\n'
    assert list(source_rows(content, "csv", {}, error=file_error)) == [
        (1, ["code", "label"], {}),
        (2, ["A", "first\r\nsecond"], {}),
        (4, ["B", "third"], {}),
    ]


@pytest.mark.parametrize("read", [
    partial(read_resource_file, "machine"),
    partial(read_calendar_file, "work_calendar"),
    partial(read_relation_file, "operator_machine"),
    read_material_file,
    partial(decode_process_file, "route"),
], ids=["resource", "calendar", "relation", "material", "process"])
def test_codec_header_rejection_closes_the_open_workbook(monkeypatch, read):
    book = openpyxl.Workbook()
    book.active.append(["unknown header"])
    buffer = BytesIO()
    book.save(buffer)
    book.close()
    load = openpyxl.load_workbook
    closed = []

    def track_load(*args, **kwargs):
        opened = load(*args, **kwargs)
        close = opened.close

        def track_close():
            closed.append(opened)
            close()

        opened.close = track_close
        return opened

    monkeypatch.setattr(openpyxl, "load_workbook", track_load)
    with pytest.raises(ValidationError) as exc:
        read(buffer.getvalue(), "xlsx")
    assert exc.value.field == "headers"
    assert len(closed) == 1
