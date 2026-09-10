"""Real SQLite/routes/projection + isolated Chrome 109 widgets, not a main build."""

import csv
import hashlib
import io
import json
import os
import subprocess
from pathlib import Path

import openpyxl

from tests.workbench.calibration_support import calibration_api as api_fixture
from tests.workbench.calibration_support import calibration_case as case_fixture
from tests.workbench.calibration_support import ledger_fixture
from tests.workbench.calibration_widgets_support import serve
from tests.workbench.test_live_browser import runtime_tools

ROOT = Path(__file__).resolve().parents[2]
HERE = Path(__file__).resolve().parent
EXPORT_FIELDS = ["part_no", "part_name", "sequence", "operation_label", "template_operation_ref", "template_revision",
                 "template_snapshot", "old_unit_hours", "suggested_unit_hours", "deviation_percent", "deviation_basis",
                 "sample_count", "candidate_count", "sample_refs", "sample_revisions", "exclusion_reasons", "method_version",
                 "generated_at", "as_of", "snapshot_ref"]
HEADERS = ["图号", "零件名称", "工序号", "工序名称", "模板工序引用", "模板修订", "模板快照", "旧单件定额h", "建议单件定额h",
           "偏差百分比", "偏差计算状态", "有效样本数", "待核对实例数", "样本引用", "样本修订", "剔除原因", "计算方法", "生成时间", "数据截至", "范围快照", "筛选范围"]


def unescape(value):
    if isinstance(value, str) and value.startswith("'"):
        return value[1:]
    return value


def verify_download(download):
    content = Path(download["path"]).read_bytes()
    assert hashlib.sha256(content).hexdigest() == download["sha256"]
    if download["format"] == "csv":
        assert content.startswith(b"\xef\xbb\xbf")
        rows = list(csv.reader(io.StringIO(content.decode("utf-8-sig"))))
    else:
        book = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        try:
            assert book.sheetnames == ["校准建议", "范围与口径"]
            rows = list(book["校准建议"].iter_rows(values_only=True))
            metadata = dict(book["范围与口径"].iter_rows(values_only=True))
            assert unescape(metadata["范围快照"]) == download["snapshot"]
            assert metadata["数据截至"] == download["as_of"]
            assert json.loads(metadata["筛选范围"]) == download["scope"]
        finally:
            book.close()
    assert list(rows[0]) == HEADERS
    assert len(rows) - 1 == download["headers_count"] == len(download["expected_rows"])
    for values, expected in zip(rows[1:], download["expected_rows"]):
        for value, key in zip(values, EXPORT_FIELDS):
            value = unescape(value)
            expected_value = expected[key]
            if isinstance(expected_value, (list, dict)):
                assert json.loads(value) == expected_value, key
            elif expected_value is None:
                assert value == "未知", key
            elif isinstance(expected_value, (int, float)):
                assert float(value) == expected_value, key
            else:
                assert str(value) == expected_value, key
        assert json.loads(values[-1]) == download["scope"]
        assert expected["sample_count"] == 0 and expected["suggested_unit_hours"] is None
        assert unescape(values[-2]) == download["snapshot"]


def test_calibration_widgets(calibration_api, tmp_path):
    output = Path(os.environ.get("CALIBRATION_UI_OUTPUT", str(tmp_path / "browser-evidence"))).resolve()
    assert output != ROOT and ROOT not in output.parents
    output.mkdir(parents=True, exist_ok=True)
    node, browser, modules = runtime_tools()
    print("CALIBRATION_UI_ARTIFACTS " + str(output), flush=True)
    with serve(calibration_api, output) as config:
        result = subprocess.run([node, str(HERE / "calibration_widgets_probe.cjs"), str(output)],
            input=json.dumps(config, ensure_ascii=False), text=True, capture_output=True, timeout=300,
            env=dict(os.environ, NODE_PATH=modules, WORKBENCH_BROWSER=browser))
    assert result.returncode == 0, result.stdout + result.stderr + "\n" + str(output)
    report = json.loads((output / "calibration-ui-result.json").read_text(encoding="utf-8"))
    assert report["browser"].startswith("109.") and report["compile_global_build"] is False
    assert len(report["cases"]) == 4 and len(report["downloads"]) == 8
    assert report["external"] == report["errors"] == []
    assert report["unknown_permission_disabled"] and report["contract_rejections"] >= 6
    assert report["export_header_rejections"] == 2
    for source in report["sources"]:
        assert hashlib.sha256((ROOT / source["path"]).read_bytes()).hexdigest() == source["sha256"]
    for download in report["downloads"]:
        verify_download(download)
        assert download["page"] == 2 and download["headers_count"] == 22
    for case in report["cases"]:
        assert case["zero_distinct_null"] and case["selected_source_retained"] and case["real_409"]
        assert case["geometry"]["scroll_width"] <= case["viewport"]["width"]
    server = json.loads((output / "calibration-server-result.json").read_text(encoding="utf-8"))
    assert server["violations"] == [] and server["server_stopped"]
    assert len(server["explicit_fixture_mutations"]) == 4
    assert all(row["method"] == "GET" and row["database_unchanged"] and row["writes"] == [] for row in server["journal"])
    assert sum(row["status"] == 409 for row in server["journal"]) == 4
    print("CALIBRATION_UI_VERIFIED " + json.dumps({"cases": 4, "downloads": 8, "server_gets": len(server["journal"]), "output": str(output)}), flush=True)
