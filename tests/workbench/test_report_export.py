"""Filtered report exports preserve complete cohorts and CSV/XLSX payloads."""

import csv
import io

import openpyxl
import pytest

from tests.workbench.plan_read_support import assert_error
from tests.workbench.report_api_support import report_api as _report_api


@pytest.mark.parametrize("topic", ["delivery", "records", "machines", "people", "quality"])
@pytest.mark.parametrize("format_name", ["csv", "xlsx"])
def test_entire_filtered_cohort_export_and_bytes(report_api, topic, format_name):
    first = report_api.read(topic=topic, size=10)
    snapshot = first["meta"]["snapshot_ref"]
    before = report_api.state()
    response = report_api.get("/export", topic=topic, size=10, snapshot_ref=snapshot, format=format_name)
    assert response.status_code == 200, response.get_data(as_text=True) if response.status_code != 200 else ""
    assert response.headers["X-Workbench-Snapshot"] == snapshot
    assert response.headers["X-Workbench-As-Of"] == first["meta"]["as_of"]
    if format_name == "csv":
        rows = list(csv.reader(io.StringIO(response.data.decode("utf-8-sig"))))
        assert len(rows[0]) == len(set(rows[0])), "CSV headers must preserve distinct row and scope fields"
    else:
        assert response.data[:2] == b"PK"
        workbook = openpyxl.load_workbook(io.BytesIO(response.data), read_only=True)
        rows = list(workbook["范围全部结果"].values)
        meta = dict(workbook["范围与计算方式"].values)
        assert meta["数据版本编号"] == snapshot and meta["数据截至"] == first["meta"]["as_of"]
        workbook.close()
    assert len(rows) - 1 == first["data"]["page"]["total"]
    assert report_api.state() == before


def test_csv_scope_notes_do_not_overwrite_operation_gaps(report_api):
    first = report_api.read(topic="quality", size=50)
    response = report_api.get("/export", topic="quality", size=50,
                              snapshot_ref=first["meta"]["snapshot_ref"], format="csv")
    assert response.status_code == 200
    rows = list(csv.DictReader(io.StringIO(response.data.decode("utf-8-sig"))))
    by_ref = {row["工序编号"]: row for row in rows}
    for operation in first["data"]["rows"]:
        exported = by_ref[operation["operation_ref"]]
        assert exported["数据缺口"] == "；".join(operation["data_gaps"])
        assert exported["统计说明与待补资料"] == "；".join(first["data"]["data_gaps"])


def test_export_stale_and_range_changes_rejected(report_api):
    first = report_api.read(query="OP-02")
    token = first["meta"]["snapshot_ref"]
    assert_error(report_api.get("/export", snapshot_ref=token), "snapshot_stale")
    with report_api.db() as conn:
        conn.execute("UPDATE Schedule SET end_time='2026-09-02 09:30:00' WHERE version=3")
    assert_error(report_api.get("/export", snapshot_ref=token, query="OP-02"), "snapshot_stale")


@pytest.mark.parametrize("kind", ["overdue", "utilization", "downtime", "official-review"])
def test_real_catalog_xlsx_roundtrip(report_api, kind):
    path = "/api/workbench/v1/reports/" + kind
    response = report_api.client.get(path)
    assert response.status_code == 200, response.get_data(as_text=True)
    data = response.get_json()
    file = report_api.client.get(path + "/export", query_string={"snapshot_ref": data["meta"]["snapshot_ref"], "format": "xlsx"})
    assert file.status_code == 200, file.get_data(as_text=True) if file.status_code != 200 else ""
    workbook = openpyxl.load_workbook(io.BytesIO(file.data), read_only=True)
    assert workbook.sheetnames
    assert file.headers["X-Workbench-Snapshot"] == data["meta"]["snapshot_ref"]
    workbook.close()


def test_catalog_rejects_implicit_finish_date_to_window_conversion(report_api):
    response = report_api.client.get("/api/workbench/v1/reports/utilization", query_string={"plan_finish_date_from": "2026-09-02", "plan_finish_date_to": "2026-09-02"})
    assert_error(response, "invalid_input", 400)
