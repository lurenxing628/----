"""XLSX metadata identities are literal text; CSV and other escaping stay intact."""

import csv
import io
import xml.etree.ElementTree as ET
from zipfile import ZipFile

import openpyxl
import pytest

from core.services.report.report_engine import ReportEngine
from core.services.workbench.report_exports import export_table

IDENTITIES = ["-R1G标识", "+R1G标识", "=1+1", "@R1G标识", "中文标识"]
NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}


def _export_fixture(identity):
    data = {"topic": "records", "provenance": "正式计划与现场事实", "data_gaps": ["未知不是零"],
            "plan": {"display_name": "中文正式计划", "plan_ref": identity},
            "scope": {"kind": "execution_analysis", "plan_ref": identity}, "summary": {"rows": 2},
            "columns": [{"key": "label", "label": "工序"}, {"key": "remark", "label": "备注"},
                        {"key": "hours", "label": "工时"}]}
    rows = [{"label": "中文工序", "remark": "=1+1", "hours": None},
            {"label": "零工时", "remark": "-保留原转义", "hours": 0}]
    snapshot = {"snapshot_ref": identity, "as_of": "2026-09-10T12:00:00"}
    return data, rows, snapshot


@pytest.mark.parametrize("identity", IDENTITIES)
@pytest.mark.parametrize("stream", [False, True])
def test_xlsx_identity_bytes_and_dtype_are_exact_text(schema_conn, tmp_path, monkeypatch, identity, stream):
    engine = ReportEngine(schema_conn)
    if stream:
        monkeypatch.setattr(engine, "EXPORT_DIRECT_MAX_ROWS", 1)
    data, rows, snapshot = _export_fixture(identity)
    exported = export_table(engine, data, rows, snapshot, "xlsx")
    payload = exported.data.read()
    (tmp_path / "identities.xlsx").write_bytes(payload)
    assert exported.mode == ("stream" if stream else "direct")
    assert exported.estimated_rows == 2
    book = openpyxl.load_workbook(io.BytesIO(payload), read_only=True, data_only=False)
    try:
        cells = list(book["范围与口径"].iter_rows())
        identities = [row[1] for row in cells if row[0].value in ("计划引用", "范围快照")]
        assert len(identities) == 2
        assert all(cell.value == identity and cell.data_type == "s" for cell in identities)
        assert cells[1][1].value == "中文正式计划" and cells[1][1].data_type == "s"
        table = list(book["范围全部结果"].iter_rows())
        assert [cell.value for cell in table[1]] == ["中文工序", "'=1+1", "未知"]
        assert [cell.value for cell in table[2]] == ["零工时", "'-保留原转义", 0]
        assert table[1][1].data_type == table[2][1].data_type == "s"
        assert table[2][2].data_type == "n"
    finally:
        book.close()
        exported.data.close()
    with ZipFile(io.BytesIO(payload)) as archive:
        metadata_xml = archive.read("xl/worksheets/sheet1.xml")
        (tmp_path / "metadata-sheet.xml").write_bytes(metadata_xml)
        root = ET.fromstring(metadata_xml)
        for coordinate in ("B3", "B5"):
            cell = root.find(".//m:c[@r='" + coordinate + "']", NS)
            assert cell is not None and cell.get("t") == "inlineStr"
            text = cell.find("m:is/m:t", NS)
            assert text is not None and text.text == identity
        for path in ("xl/worksheets/sheet1.xml", "xl/worksheets/sheet2.xml"):
            assert ET.fromstring(archive.read(path)).findall(".//m:f", NS) == []


@pytest.mark.parametrize("identity", IDENTITIES)
def test_csv_keeps_original_prefix_and_unknown_value_contract(schema_conn, identity):
    data, rows, snapshot = _export_fixture(identity)
    exported = export_table(ReportEngine(schema_conn), data, rows, snapshot, "csv")
    try:
        payload = exported.data.read()
    finally:
        exported.data.close()
    records = list(csv.DictReader(io.StringIO(payload.decode("utf-8-sig"))))
    assert len(records) == exported.estimated_rows == 2
    expected_identity = "'" + identity if identity[0] in "-+=@" else identity
    assert all(row["范围快照"] == expected_identity for row in records)
    assert records[0]["备注"] == "'=1+1" and records[1]["备注"] == "'-保留原转义"
    assert records[0]["工时"] == "未知" and records[1]["工时"] == "0"
    assert records[0]["工序"] == "中文工序"
