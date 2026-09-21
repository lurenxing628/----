"""日历文件测试的独立取值口径：直接读 WorkCalendar 表，不调生产读取来验证生产写入。"""

import csv
from io import BytesIO, StringIO

import openpyxl

from core.models.workbench_table_descriptor import INSTRUCTION_SHEET

HEADERS = ("日期", "类型", "可排工时（小时）", "效率（%）", "允许普通件", "允许急件", "备注")


def file_bytes(rows, fmt="csv", headers=HEADERS):
    if fmt == "csv":
        target = StringIO(newline="")
        writer = csv.writer(target)
        writer.writerow(headers)
        writer.writerows(rows)
        return target.getvalue().encode("utf-8-sig")
    wb = openpyxl.Workbook()
    ws = wb.worksheets[0]
    ws.append(list(headers))
    for row in rows:
        ws.append(list(row))
    output = BytesIO()
    wb.save(output)
    wb.close()
    return output.getvalue()


def stored_day(conn, day):
    row = conn.execute("SELECT * FROM WorkCalendar WHERE date=?", (day,)).fetchone()
    return dict(row) if row else None


def stored_days(conn):
    return {row["date"]: dict(row) for row in conn.execute("SELECT * FROM WorkCalendar ORDER BY date")}


def decode(download, fmt):
    if fmt == "csv":
        rows = list(csv.reader(StringIO(download.content.decode("utf-8-sig"))))
        return rows[0], [[value[1:] if value.startswith("'") else value for value in row] for row in rows[1:]]
    wb = openpyxl.load_workbook(BytesIO(download.content), read_only=True, data_only=False)
    try:
        # 模板与导出文件固定两张表：数据表在前，「填写说明」在后。
        assert wb.sheetnames[1:] == [INSTRUCTION_SHEET], wb.sheetnames
        source = wb.worksheets[0].iter_rows()
        headers = [cell.value for cell in next(source)]
        rows = []
        for cells in source:
            assert all(cell.data_type not in ("f", "e") for cell in cells)
            rows.append([cell.value for cell in cells])
        return headers, rows
    finally:
        wb.close()
