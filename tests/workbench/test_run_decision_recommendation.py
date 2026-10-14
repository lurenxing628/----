"""Saved recommendation evidence and its complete business export stay aligned."""

import csv
import io
import json
from copy import deepcopy
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from core.services.workbench.facts.resource_intervals import task_interval
from core.services.workbench.run.decision_metrics import metric_definitions, operation_changes
from core.services.workbench.run.decision_recommendation import RecommendationSupport
from core.services.workbench.run.decision_recommendation_export import (
    COMPLETION_HEADERS,
    RECOMMENDATION_CSV_HEADERS,
    RECOMMENDATION_HEADERS,
    REMAINING_HEADERS,
    _definition_values,
    _explanations,
)
from core.services.workbench.run.decision_selection_explanation import selection_explanation
from core.services.workbench.run.decision_tradeoffs import aggregate_tradeoffs, tradeoff_entry
from tests.workbench.run_candidate_support import BASE, api, compute, corrupt_update, edit_artifact
from tests.workbench.run_candidate_support import candidate_case as _candidate_case  # noqa: F401
from tests.workbench.test_run_decision_api import _assert_summary_matches_full, decision, decision_summary
from tests.workbench.test_run_decision_impact_api import operation_ref


def test_tradeoffs_prioritize_business_groups_and_keep_all_differences_and_neutral_work():
    definitions = metric_definitions()
    values = {
        "unplanned_operation_count": (-1, "better"), "overdue_count": (-2, "better"),
        "machine_high_pressure_count": (-1, "better"), "makespan_hours": (-1000000, "better"),
        "total_tardiness_hours": (2, "worse"), "operator_high_pressure_count": (1, "worse"),
        "type_changeover_count": (1, "worse"), "changeover_count": (1000, "worse"),
        "changed_operation_count": (1000000, "worse"),
    }
    candidate_ref, run_ref = "b" * 48, "a" * 48
    supports = [{"kind": "operation", "ref": str(index).zfill(48), "candidate_ref": candidate_ref,
                 "drill": {"run_ref": run_ref, "candidate_ref": candidate_ref, "operation_ref": str(index).zfill(48)}}
                for index in range(7)]
    compared = [str(index).zfill(48) for index in range(6)]
    entries = []
    for definition in definitions:
        value, direction = values.get(definition["key"], (0, "equal"))
        delta = {"value": value, "state": "comparable", "direction": direction,
                 "reference_value": 10, "selected_value": 10 + value}
        if definition["key"] == "scheduled_setup_hours":
            delta.update(value=-2, selected_value=8, state="different_coverage", direction="neutral")
        entries.append(tradeoff_entry(definition, delta, scope=definition["scope"], source=definition["source"],
                                      metric_version="decision-preparation-v4", supporting_refs=supports,
                                      compared_batch_refs=compared + compared))
    original = deepcopy(entries)
    result = aggregate_tradeoffs(entries, definitions)
    assert len(result["all_differences"]) == len(definitions) == 52
    assert result["all_differences"] == original and entries == original
    assert [item["metric_key"] for item in result["benefits"]] == [
        "unplanned_operation_count", "overdue_count", "machine_high_pressure_count"]
    assert [item["metric_key"] for item in result["costs"]] == [
        "total_tardiness_hours", "operator_high_pressure_count", "type_changeover_count"]
    reduced = next(item for item in result["all_differences"] if item["metric_key"] == "scheduled_setup_hours")
    assert reduced["value"] == -2 and reduced["direction"] == "neutral"
    assert "少安排的准备需求" in reduced["text"] and "节省" not in reduced["text"]
    assert reduced["compared_batch_count"] == 6 and reduced["sample_compared_batch_refs"] == compared[:3]
    assert reduced["supporting_refs"] == supports[:3]
    assert all(item["metric_key"] != "scheduled_setup_hours" for item in result["benefits"] + result["costs"])
    catalog = {definition["key"]: definition for definition in definitions}
    assert (catalog["machine_high_pressure_count"]["label"], catalog["machine_high_pressure_count"]["unit"]) == ("高负荷设备", "台")
    assert (catalog["operator_high_pressure_count"]["label"], catalog["operator_high_pressure_count"]["unit"]) == ("高负荷人员", "人")
    for key, value, wording in (("makespan_hours", -0.001, "不到 0.01 小时"),
                                ("machine_utilization_ratio", 0.002, "0.2 个百分点")):
        item = tradeoff_entry(catalog[key], {"value": value, "reference_value": 1, "selected_value": 1 + value,
                                           "state": "comparable", "direction": "neutral"},
                              scope=catalog[key]["scope"], source=catalog[key]["source"], metric_version="decision-preparation-v4")
        assert item["value"] == value and wording in item["text"]
        assert "相同" not in item["text"]
    basis = {
        "decision_evidence": {"selected": {"score": [0, 1, 5]},
            "anchors": {"raw_score_best": {"score": [0, 1, 5]}},
            "tiebreak": {"primary_tied_candidate_refs": [candidate_ref, "c" * 48],
                         "full_score_tied_candidate_refs": [candidate_ref], "decisive_step": "secondary_score",
                         "decisive_score_index": 2, "runner_up_candidate_ref": "c" * 48,
                         "raw_score_best_candidate_ref": candidate_ref},
            "balanced_checks": {"checked": False, "reason_code": "score_only"}},
        "health_decision_evidence": {"state": "unavailable", "reason_code": "critical_path_unavailable", "components": []},
        "reason_code": "score_only_raw_score_best", "selection_policy": "score_only", "objective_name": "min_changeover",
        "objective_field_order": ["failed_ops", "changeover_count", "makespan_hours"],
        "selected_candidate_ref": candidate_ref, "raw_score_best_candidate_ref": candidate_ref,
    }
    explanation = selection_explanation(basis, {}, [])
    explanation_rows = {item["key"]: item for item in explanation["entries"]}
    assert explanation_rows["score_1"]["label"] == "换型次数"
    assert explanation_rows["score_1"]["selected_value"] == 1
    assert explanation_rows["score_1"]["unit"] == "次"
    assert "推荐方案的换型次数 1 次" in explanation["headline"]
    assert "继续比较“安排总时长”" in explanation["headline"]
    assert "主要指标相同的方案有 2 份" in explanation_rows["tiebreak"]["text"]
    assert "各项目标指标都相同" not in explanation_rows["tiebreak"]["text"]
    assert explanation_rows["health_unavailable"]["state"] == "unavailable"
    assert "未保存关键工序列表" in explanation_rows["health_unavailable"]["text"]
    checked_basis = deepcopy(basis)
    checked_basis.update(selection_policy="balanced", reason_code="balanced_raw_score_best")
    checked_basis["decision_evidence"]["balanced_checks"] = {
        "checked": True, "candidate_ref": "c" * 48, "reference_candidate_ref": candidate_ref,
        "accepted": False, "override_applied": False,
        "checks": [{"key": "total_tardiness_hours", "reference_value": 10, "candidate_value": 12,
                    "limit": 11, "tolerance": 0.1, "passed": False}],
    }
    recorded = deepcopy(checked_basis["decision_evidence"])
    checked = selection_explanation(checked_basis, {}, [
        {"candidate_ref": "c" * 48, "artifact": {"label": "实际检查候选"}},
        {"candidate_ref": candidate_ref, "artifact": {"label": "目标评分最优候选"}},
    ])
    tolerance = next(item for item in checked["entries"] if item["key"] == "tolerance_total_tardiness_hours")
    assert tolerance["unit"] == "小时" and tolerance["tolerance_unit"] == "%" and tolerance["tolerance"] == 0.1
    assert tolerance["reference_value"] == 10 and tolerance["selected_value"] == 12 and tolerance["limit"] == 11
    assert "实际检查候选" in tolerance["text"] and "目标评分最优候选" in tolerance["text"]
    assert "c" * 48 in tolerance["text"] and candidate_ref in tolerance["text"]
    exported = _definition_values(tolerance)
    assert exported["容差"] == 10 and exported["容差单位"] == "%" and exported["单位"] == "小时"
    assert exported["基准值"] == 10 and exported["比较对象值"] == 12 and exported["实际限制"] == 11
    assert checked_basis["decision_evidence"] == recorded
    assert "总拖期 12 小时" in checked["headline"] and "允许上限 11 小时" in checked["headline"]
    tied = deepcopy(basis)
    tied.update(objective_name="min_overdue", objective_field_order=["failed_ops", "overdue_count", "total_tardiness_hours"],
                selection_policy="balanced", reason_code="balanced_raw_score_best")
    tied["decision_evidence"]["selected"].update(score=[0, 2, 5], overdue_count=2, total_tardiness_hours=5)
    tied["decision_evidence"]["tiebreak"].update(full_score_tied_candidate_refs=[candidate_ref, "c" * 48],
                                               decisive_step="baseline_preference", decisive_score_index=None)
    tied["decision_evidence"]["balanced_checks"] = {"checked": False, "reason_code": "critical_health_not_better"}
    tied_rows = [{"candidate_ref": ref, "status": "completed", "artifact": {"label": "原算法方案" if ref == candidate_ref else "优化方案"}}
                 for ref in (candidate_ref, "c" * 48)]
    tied_explanation = selection_explanation(tied, {}, tied_rows)
    assert tied_explanation["headline"] == "本次优化未减少超期批次和累计超期时长，推荐原算法方案；本次计算范围内：预计超期批次 2 批、总拖期 5 小时。"
    cost_basis = deepcopy(checked_basis)
    cost_basis["reason_code"] = "balanced_critical_health_better"
    cost_basis["decision_evidence"]["selected"]["score"][1] = 3
    cost_basis["decision_evidence"]["balanced_checks"].update(accepted=True, override_applied=True)
    check = cost_basis["decision_evidence"]["balanced_checks"]["checks"][0]
    check.update(reference_value=284.61, candidate_value=301.85, limit=313.071, passed=True)
    raw_delta = check["candidate_value"] - check["reference_value"]
    assert raw_delta != 17.24
    cost_basis.update(health_reference_candidate_ref="c" * 48, raw_score_best_candidate_ref="c" * 48)
    cost_basis["decision_evidence"]["balanced_checks"].update(candidate_ref=candidate_ref, reference_candidate_ref="c" * 48)
    cost_basis["health_decision_evidence"] = {"state": "complete", "components": [
        {"key": key, "state": "complete", "contribution": 1, "improvement": "lower",
         "reference_value": before, "candidate_value": after, "delta_hours": delta, "tolerance_hours": tolerance}
        for key, before, after, delta, tolerance in (
            ("critical_chain_finish_hours_delta", 76.235, 30.33, -45.905, 3.81175),
            ("critical_chain_wait_hours_delta", 46.735, 0.83, -45.905, 2.33675),
            ("top_impact_ops_avg_start_hours_delta", 18.04842105263158, 16.462105263157895, -1.5863157894736843, 0.01))]}
    cost_candidates = [{"candidate_ref": "c" * 48, "artifact": {"label": "原算法"}}]
    recorded_cost = deepcopy(cost_basis)
    cost_explanation = selection_explanation(cost_basis, {"health": {"state": "better", "score": 3}}, cost_candidates)
    assert "本次计算范围内" in cost_explanation["headline"]
    assert "相对原算法，关键工序总时长减少 45.905 小时" in cost_explanation["headline"]
    assert "关键工序间等待减少 45.905 小时" in cost_explanation["headline"]
    assert "总拖期增加 17.24 小时" in cost_explanation["headline"]
    assert str(raw_delta) not in cost_explanation["headline"]
    cost = next(item for item in cost_explanation["entries"] if item["key"] == "tolerance_total_tardiness_hours")
    assert cost["reference_value"] == 284.61 and cost["selected_value"] == 301.85 and cost["value"] == raw_delta
    assert _definition_values(cost)["差值"] == raw_delta and cost_basis == recorded_cost


def test_operation_supports_keep_known_changes_and_removed_reference_arrangements():
    def ref(number):
        return format(number, "048x")

    run_ref, fixed_ref, viewed_ref, reference_ref, plan_ref, batch_ref = [ref(number) for number in range(1, 7)]
    operation_refs = {name: ref(index + 100) for index, name in enumerate(
        ("both_none", "added", "removed", "old_multiple", "new_multiple", "machine_changed", "time_unknown"))}

    def task(name, row, *, start="2026-09-09T08:00:00", end="2026-09-09T09:00:00", machine=7):
        return {"row_ref": ref(row), "operation_ref": operation_refs[name], "batch_ref": batch_ref,
                "start": start, "end": end, "source": "internal", "machine": {"ref": ref(machine), "label": "设备"},
                "operator": {"ref": ref(9), "label": "人员"}, "supplier": None, "locked": False,
                "data_gaps": [], "interval_comparable": True}

    before = {value: [] for value in operation_refs.values()}
    after = {value: [] for value in operation_refs.values()}
    after[operation_refs["added"]] = [task("added", 201)]
    before[operation_refs["removed"]] = [task("removed", 202, start="2026-09-08T20:00:00", end="2026-09-08T21:00:00")]
    before[operation_refs["old_multiple"]] = [task("old_multiple", 203, end="2026-09-09T08:30:00"),
                                              task("old_multiple", 204, start="2026-09-09T08:30:00")]
    after[operation_refs["old_multiple"]] = [task("old_multiple", 205)]
    before[operation_refs["new_multiple"]] = [task("new_multiple", 206)]
    after[operation_refs["new_multiple"]] = [task("new_multiple", 207, end="2026-09-09T08:30:00"),
                                             task("new_multiple", 208, start="2026-09-09T08:30:00")]
    before[operation_refs["machine_changed"]] = [task("machine_changed", 209)]
    after[operation_refs["machine_changed"]] = [task("machine_changed", 210, machine=8)]
    unknown = task("time_unknown", 211, start=None, end=None)
    unknown.update(interval_comparable=False, data_gaps=[{"field": "start", "code": "source_missing"}])
    before[operation_refs["time_unknown"]] = [unknown]
    after[operation_refs["time_unknown"]] = [deepcopy(unknown)]
    groups = {"reference": before, fixed_ref: after}
    intervals = {key: {row["row_ref"]: task_interval(row) for rows in grouped.values() for row in rows}
                 for key, grouped in groups.items()}
    source = SimpleNamespace(
        run={"run_ref": run_ref}, selected_ref=viewed_ref, reference_ref=reference_ref, comparison_mode="candidates",
        resource_view="selected_batches", baseline={"state": "available", "kind": "candidate", "candidate_ref": reference_ref},
        scope={"operation_refs": list(operation_refs.values()), "batch_refs": [batch_ref],
               "range_start": "2026-09-09T08:00:00", "range_end": "2026-09-09T12:00:00"},
        preparation=SimpleNamespace(groups=groups.__getitem__, labels={value: {"batch_ref": batch_ref} for value in operation_refs.values()}),
        task_intervals=intervals.__getitem__,
    )
    supports = RecommendationSupport(source, fixed_ref)
    candidate_supports = supports.operations("changed_operation_count")
    assert [item["ref"] for item in candidate_supports] == [operation_refs[name] for name in ("added", "removed", "machine_changed")]
    assert all(item["candidate_ref"] == fixed_ref and item["drill"]["selected_candidate_ref"] == fixed_ref for item in candidate_supports)
    assert [item["side"] for item in candidate_supports] == ["selected", "reference", "selected"]
    assert [item["ref"] for item in supports.operations("machine_change_count")] == [operation_refs["machine_changed"]]
    removed = candidate_supports[1]
    assert removed["drill"]["candidate_ref"] == reference_ref and removed["candidate_ref"] != removed["drill"]["candidate_ref"]
    assert removed["drill"]["selected_row_ref"] == ref(202) and removed["drill"]["operation_ref"] == operation_refs["removed"]
    assert removed["drill"]["range_start"] == "2026-09-08T20:00:00" and removed["drill"]["drill_range_adjusted"] is True
    assert "comparison_tab" not in removed["drill"] and "reference_view" not in removed["drill"]
    before_tasks = [row for rows in before.values() for row in rows]
    after_tasks = [row for rows in after.values() for row in rows]
    metrics = operation_changes(after_tasks, before_tasks, list(operation_refs.values()), True)
    assert metrics["changed_operation_count"]["value"] is None
    assert metrics["changed_operation_count"]["known_subtotal"] == 3 and metrics["changed_operation_count"]["unknown_count"] == 4
    assert metrics["machine_change_count"]["value"] is None
    assert metrics["machine_change_count"]["known_subtotal"] == 1 and metrics["machine_change_count"]["unknown_count"] == 2
    basis = {"selected_candidate_ref": fixed_ref, "raw_score_best_candidate_ref": fixed_ref,
             "health_reference_candidate_ref": reference_ref, "selection_policy": "score_only", "selection_policy_version": "candidate-selection-v1",
             "reason_code": "score_only_raw_score_best", "objective_name": "min_changeover", "objective_field_order": [],
             "selected_score": [], "overdue_tolerance_count": None, "tardiness_tolerance_ratio": None,
             "decision_evidence": None, "health_decision_evidence": None, "explanation": {"entries": [], "issues": []}}
    definition = next(row for row in metric_definitions() if row["key"] == "changed_operation_count")

    def export_supports(items):
        entry = tradeoff_entry(definition, {"value": None, "reference_value": None, "selected_value": None,
                                           "state": "unavailable", "direction": "neutral"},
                              scope=definition["scope"], source=definition["source"], metric_version="decision-preparation-v4", supporting_refs=items)
        data = {"selection_basis": basis, "recommendation": {"state": "available", "recommendation_type": "partial", "headline": "synthetic",
                "current_comparison": {"benefits": [], "costs": [], "all_differences": [entry]}}}
        return [row for row in _explanations(data) if row["record_kind"] == "support"]

    candidate_export = export_supports(candidate_supports)
    removed_export = [row for row in candidate_export if row["相关明细永久编号"] == operation_refs["removed"]]
    assert removed_export and all(row["相关明细候选编号"] == reference_ref for row in removed_export)
    assert {row["相关明细筛选字段"]: row["相关明细筛选值"] for row in removed_export}["side"] == "reference"
    source.comparison_mode, source.reference_ref = "admission_official", None
    source.baseline = {"state": "available", "kind": "admission_official", "plan_ref": plan_ref}
    formal_supports = supports.operations("changed_operation_count")
    formal_removed = formal_supports[1]
    assert formal_removed["side"] == "reference" and formal_removed["drill"]["candidate_ref"] == fixed_ref
    assert formal_removed["drill"]["reference_view"] is True and formal_removed["drill"]["plan_ref"] == plan_ref
    assert formal_removed["drill"]["selected_row_ref"] == ref(202)
    formal_export = export_supports(formal_supports)
    formal_removed_export = [row for row in formal_export if row["相关明细永久编号"] == operation_refs["removed"]]
    assert all(row["相关明细候选编号"] == fixed_ref for row in formal_removed_export)
    filters = {row["相关明细筛选字段"]: row["相关明细筛选值"] for row in formal_removed_export}
    assert filters["side"] == "reference" and filters["reference_view"] is True and filters["plan_ref"] == plan_ref
    source.changeovers = SimpleNamespace(projection=lambda _key: {"events": []}, tasks=lambda _key: [])
    event_definition = next(row for row in metric_definitions() if row["key"] == "type_changeover_count")
    zero_events = supports.entries(event_definition, {"value": 0, "state": "comparable", "direction": "equal"})
    assert len(zero_events) == 1 and zero_events[0]["kind"] == "collection"
    assert zero_events[0]["drill"]["comparison_tab"] == "changeovers" and zero_events[0]["drill"]["detail_query"] == {}
    source.baseline["state"] = "absent"
    assert supports.operations("changed_operation_count") == supports.operations("machine_change_count") == []
    Path("/tmp/decision-recommendation-removed-support.json").write_text(json.dumps({
        "synthetic": True, "scenario": "seven archived-operation arrangements; support DTOs produced by RecommendationSupport",
        "run_ref": run_ref, "fixed_candidate_ref": fixed_ref, "viewed_candidate_ref": viewed_ref,
        "reference_candidate_ref": reference_ref, "plan_ref": plan_ref, "batch_ref": batch_ref,
        "range_start": source.scope["range_start"], "range_end": source.scope["range_end"],
        "candidate_supports": candidate_supports, "formal_supports": formal_supports,
        "candidate_export_supports": candidate_export, "formal_export_supports": formal_export, "zero_changeover_supports": zero_events,
        "metrics": metrics,
    }, ensure_ascii=False), encoding="utf-8")


def test_legacy_missing_decision_evidence_keeps_saved_selection_and_unspecified_due(candidate_case):
    case = candidate_case
    case.conn.execute("UPDATE Batches SET due_date='9999-12-31' WHERE batch_id='B1'")
    case.conn.commit()
    run_ref, refs = compute(case)
    client, _ = api(case)
    original = decision(client, run_ref)["data"]
    receipt = json.loads(case.conn.execute(
        "SELECT result_json FROM WorkbenchRunReceipts WHERE run_ref=?", (run_ref,)).fetchone()[0])
    receipt["selection"].pop("decision_evidence", None)
    corrupt_update(case.conn, "WorkbenchRunReceipts", "UPDATE WorkbenchRunReceipts SET result_json=? WHERE run_ref=?",
                   (json.dumps(receipt, ensure_ascii=False), run_ref))
    for ref in refs:
        edit_artifact(case, ref, lambda artifact: (artifact.get("health") or {}).pop("decision_evidence", None))
    legacy = decision(client, run_ref, reference_candidate_ref=refs[-1], selected_candidate_ref=refs[1])["data"]
    _assert_summary_matches_full(decision_summary(client, run_ref,
        reference_candidate_ref=refs[-1], selected_candidate_ref=refs[1])["data"], legacy)
    before, after = original["selection_basis"], legacy["selection_basis"]
    for key in ("selected_candidate_ref", "raw_score_best_candidate_ref", "health_reference_candidate_ref",
                "selection_policy", "selection_policy_version", "reason_code", "objective_name", "selected_score",
                "overdue_tolerance_count", "tardiness_tolerance_ratio"):
        assert after[key] == before[key]
    assert after["evidence_state"] == "partial"
    assert after["decision_evidence"] is after["health_decision_evidence"] is None
    assert after["explanation"]["issues"]
    recommendation = legacy["recommendation"]
    assert recommendation["candidate_ref"] == before["selected_candidate_ref"]
    assert recommendation["recommendation_type"] == "complete"
    assert recommendation["viewed_candidate_ref"] == refs[1]
    common = recommendation["current_comparison"]["common_completed_delivery"]
    assert common["batch_count"] == 1 and common["state"] == "not_applicable"
    assert common["delay_delta_hours"] is common["overdue_delta_count"] is None
    for side in ("before", "after"):
        assert common[side]["excluded_due_count"] == 1
        assert common[side]["total_tardiness_hours"] is common[side]["overdue_count"] is None


def test_recommendation_exports_include_complete_remaining_work_and_preserve_view_snapshot(candidate_case, monkeypatch):
    case = candidate_case
    case.conn.execute("UPDATE Batches SET quantity=1,due_date='2026-09-09' WHERE batch_id='B1'")
    case.conn.execute("UPDATE BatchOperations SET setup_hours=2,unit_hours=3 WHERE id=?", (case.op_id,))
    operation_ids = {"B1": case.op_id}
    for number in range(2, 7):
        case.batch("B" + str(number), quantity=1, due_date="2026-09-10")
        operation_ids["B" + str(number)] = case.operation("B" + str(number), setup_hours=3, unit_hours=3)
    case.config(dispatch_mode="batch_order", sort_strategy="due_date_first")
    run_ref, refs = compute(case, case.settings(*("B" + str(number) for number in range(1, 7)), end_date="2026-09-09"))
    from itertools import count

    from web import public_token_registry

    token_numbers = count(1)
    monkeypatch.setattr(public_token_registry.secrets, "token_urlsafe", lambda _size: "-" + format(next(token_numbers), "031x"))
    client, _ = api(case)
    document = decision_summary(client, run_ref, reference_candidate_ref=refs[-1], selected_candidate_ref=refs[-2])
    data, snapshot = document["data"], document["meta"]["snapshot_ref"]
    recommendation = data["recommendation"]
    assert "本次计算范围内已排齐的批次中" in data["selection_basis"]["explanation"]["headline"]
    assert recommendation["recommendation_type"] == "partial"
    assert recommendation["candidate_ref"] == data["selection_basis"]["selected_candidate_ref"]
    assert recommendation["viewed_candidate_ref"] == refs[-2]
    scope = recommendation["completion_scope"]
    assert scope["complete_batch_count"] == 1 and scope["ordinary_incomplete_batch_count"] == 5
    assert scope["unplanned_operation_count"] == scope["remaining_work"]["operation_count"] == 5
    assert len(scope["sample_batch_refs"]["incomplete"]) == 3
    query = {"reference_candidate_ref": refs[-1], "selected_candidate_ref": refs[-2], "snapshot_ref": snapshot}
    path = BASE + "/runs/" + run_ref + "/decision/export"
    csv_response = client.get(path, query_string=dict(query, section="recommendation", format="csv"))
    assert csv_response.status_code == 200, csv_response.get_data(as_text=True)
    rows = list(csv.DictReader(io.StringIO(csv_response.data.decode("utf-8-sig"))))
    assert rows and tuple(rows[0]) == RECOMMENDATION_CSV_HEADERS and len(rows[0]) == 97
    identity = {"排产编号": run_ref, "归档编号": data["capture_ref"], "读取编号": snapshot,
                "保存推荐候选编号": recommendation["candidate_ref"], "当前查看候选编号": refs[-2],
                "查看基准编号": refs[-1], "比较模式": "candidates", "查看基准类型": "candidate",
                "范围起点（含）": data["scope"]["range_start"], "范围终点（不含）": data["scope"]["range_end"],
                "资源视角": data["scope"]["resource_view"], "推荐版本": "decision-recommendation-v1",
                "推荐记录格式版本": "candidate-selection-evidence-v1"}
    for row in rows:
        assert {field: row[field] for field in identity} == {**identity, "读取编号": "'" + snapshot}
    differences = {row["指标标识"]: row for row in rows if row["record_kind"] == "difference"}
    assert len(differences) == 52
    definitions = {entry["key"]: entry for entry in data["metric_definitions"]}
    for entry in recommendation["current_comparison"]["all_differences"]:
        exported = differences[entry["metric_key"]]
        assert entry["metric_version"] == definitions[entry["metric_key"]]["metric_version"]
        assert exported["指标版本"] == entry["metric_version"]
        assert exported["范围"] == entry["scope"] and exported["来源"] == entry["source"]
        assert exported["状态"] == entry["state"] and exported["变化方向"] == entry["direction"]
        assert exported["说明"] == entry["text"]
        for field, key in (("基准值", "reference_value"), ("比较对象值", "selected_value"), ("差值", "value")):
            if entry[key] is None:
                assert exported[field] == ""
            else:
                assert float(exported[field]) == pytest.approx(entry[key] * (100 if entry["unit"] == "%" else 1))
    supports = [row for row in rows if row["record_kind"] == "support"]
    support_targets = {(entry["metric_key"], support["ref"] or ""): support["drill"]["candidate_ref"]
                       for entry in recommendation["current_comparison"]["all_differences"] for support in entry["supporting_refs"]}
    assert supports and all(row["相关明细候选编号"] == support_targets[row["指标标识"], row["相关明细永久编号"]] for row in supports)
    completion_rows = [row for row in rows if row["record_kind"] == "completion_set"]
    assert len(completion_rows) == 24
    for ref in refs:
        candidate_rows = [row for row in completion_rows if row["候选永久编号"] == ref]
        assert len(candidate_rows) == 6 and {row["批次永久编号"] for row in candidate_rows} == set(data["scope"]["batch_refs"])
        assert sum(row["安排分类"] == "complete" for row in candidate_rows) == 1
        assert all(row["预计完工"] == row["拖期小时"] == "" for row in candidate_rows if row["安排分类"] == "incomplete")
    remaining = [row for row in rows if row["record_kind"] == "remaining_work"]
    assert len(remaining) == 20 and {row["候选永久编号"] for row in remaining} == set(refs)
    for row in remaining:
        code = row["批次编号"]
        assert row["工序永久编号"] == operation_ref(case, operation_ids[code])
        assert float(row["需求目标量"]) == float(row["执行目标量"]) == 1
        assert float(row["准备标准工时"]) == (2 if code == "B1" else 3)
        assert float(row["加工标准工时"]) == 3
        assert float(row["所需标准工时"]) == (5 if code == "B1" else 6)
        assert row["工序安排状态"] == "unplanned" and row["需求来源类型"] == "run"
        assert row["需求来源排产编号"] == run_ref and row["需求来源候选编号"] == ""
        assert row["需求来源正式版本"] == ""
    assert csv_response.headers["X-Workbench-Snapshot-Ref"] == snapshot
    assert csv_response.headers["X-Workbench-Run-Ref"] == run_ref
    assert int(csv_response.headers["X-Workbench-Row-Count"]) == len(rows)
    xlsx_response = client.get(path, query_string=dict(query, section="recommendation", format="xlsx"))
    assert xlsx_response.status_code == 200
    book = load_workbook(io.BytesIO(xlsx_response.data), read_only=True, data_only=False)
    try:
        assert book.sheetnames == ["推荐说明", "工序安排", "剩余需求", "指标计算方式", "比较范围"]
        header_lists = {"推荐说明": RECOMMENDATION_HEADERS, "工序安排": COMPLETION_HEADERS, "剩余需求": REMAINING_HEADERS}
        assert [len(headers) for headers in header_lists.values()] == [40, 36, 59]
        explanation_rows = book["推荐说明"].iter_rows(min_row=2)
        index = {field: i for i, field in enumerate(RECOMMENDATION_HEADERS)}
        for values in explanation_rows:
            if values[index["record_kind"]].value == "difference":
                key = values[index["指标标识"]].value
                assert values[index["指标版本"]].value == definitions[key]["metric_version"]
                assert values[index["指标版本"]].data_type == "s"
        for title, headers in header_lists.items():
            assert tuple(cell.value for cell in next(book[title].iter_rows())) == headers
            cells = next(book[title].iter_rows(min_row=2))
            assert cells[headers.index("读取编号")].value == snapshot
            assert cells[headers.index("读取编号")].data_type == "s"
        remaining_rows = list(book["剩余需求"].iter_rows(max_col=len(REMAINING_HEADERS)))
        assert len(remaining_rows) == 21
        indexed = {cell.value: index for index, cell in enumerate(remaining_rows[0])}
        for values in remaining_rows[1:]:
            cells = {field: values[index] for field, index in indexed.items()}
            assert cells["工序永久编号"].data_type == "s" and cells["批次编号"].data_type == "s"
            assert cells["需求目标量"].value == cells["执行目标量"].value == 1
            assert cells["需求目标量"].data_type == cells["执行目标量"].data_type == "n"
            code = cells["批次编号"].value
            assert cells["准备标准工时"].value == (2 if code == "B1" else 3)
            assert cells["加工标准工时"].value == 3
            assert cells["所需标准工时"].value == (5 if code == "B1" else 6)
            assert cells["所需标准工时"].data_type == "n"
            assert cells["需求来源候选编号"].value is None
            assert cells["需求来源正式版本"].value is None
        assert int(xlsx_response.headers["X-Workbench-Row-Count"]) == sum(
            sum(1 for _ in book[name].iter_rows()) - 1 for name in ("推荐说明", "工序安排", "剩余需求"))
    finally:
        book.close()
    all_response = client.get(path, query_string=dict(query, section="all", format="xlsx"))
    assert all_response.status_code == 200
    book = load_workbook(io.BytesIO(all_response.data), read_only=True, data_only=True)
    try:
        assert book.sheetnames == ["方案总览", "批次影响", "工序变化", "资源影响", "覆盖工序", "准备需求", "工种切换",
                                   "推荐说明", "工序安排", "剩余需求", "候选计算", "进步事件", "搜索阶段", "代表尝试", "指标计算方式", "比较范围"]
    finally:
        book.close()
    day_query = {"reference_candidate_ref": refs[-1], "selected_candidate_ref": refs[-2], "resource_kind": "machine",
                 "resource_ref": data["scope"]["resource_refs"]["machine"][0]}
    days = client.get(BASE + "/runs/" + run_ref + "/decision/resource-days", query_string=day_query)
    assert days.status_code == 200, days.get_json()
    day_snapshot = days.get_json()["meta"]["snapshot_ref"]
    all_days = client.get(path, query_string=dict(day_query, snapshot_ref=day_snapshot, section="all",
                                                 filter_section="resource-days", format="xlsx"))
    assert all_days.status_code == 200, all_days.get_json()
    book = load_workbook(io.BytesIO(all_days.data), read_only=True, data_only=True)
    try:
        assert book.sheetnames == ["方案总览", "批次影响", "工序变化", "资源影响", "资源逐日", "资源冲突", "覆盖工序",
                                   "准备需求", "工种切换", "推荐说明", "工序安排", "剩余需求", "候选计算", "进步事件", "搜索阶段", "代表尝试", "指标计算方式", "比较范围"]
    finally:
        book.close()
    _assert_summary_matches_full(data, decision(client, run_ref,
        reference_candidate_ref=refs[-1], selected_candidate_ref=refs[-2])["data"])
    stale = client.get(path, query_string=dict(query, reference_candidate_ref=refs[0], section="recommendation", format="csv"))
    assert stale.status_code == 409 and stale.get_json()["error"]["code"] == "snapshot_stale"
