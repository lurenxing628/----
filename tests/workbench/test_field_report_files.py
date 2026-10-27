"""Strict XLSX numeric precision and XML character validation for field report files."""

import re
from io import BytesIO
from typing import Dict
from zipfile import ZIP_DEFLATED, ZipFile

import pytest

from core.models.workbench_command import WorkbenchCommandRejected
from core.models.workbench_execution_input import normalize_report_input
from core.models.workbench_field_report_file import FIELDS
from core.services.workbench.execution.field_report_files_codec import decode_reports, encode_reports


def _sample_row(**overrides: object) -> Dict[str, object]:
    base: Dict[str, object] = dict.fromkeys(FIELDS, "")
    base.update(
        report_no="R-0001",
        batch_id="B1",
        operation_label="1 Turning",
        completed_quantity=1,
        actual_start="2026-09-13 08:30:00",
        actual_end="2026-09-13 10:30:00",
        effective_processing_hours=2.0,
        machine_label="M1",
        operator_label="O1",
        remark="normal",
        task_ref="a" * 48,
        operation_scope="共同",
        piece_id="",
    )
    base.update(overrides)
    return base


def _replace_d2_with_raw_number(xlsx_bytes: bytes, raw_literal: str) -> bytes:
    source = BytesIO(xlsx_bytes)
    target = BytesIO()
    with ZipFile(source, "r") as zin, ZipFile(target, "w", ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                xml_text = data.decode("utf-8")
                xml_text, count = re.subn(
                    r'<c r="D2"[^>]*>.*?</c>',
                    f'<c r="D2" t="n"><v>{raw_literal}</v></c>',
                    xml_text,
                     count=1,
                )
                assert count == 1
                data = xml_text.encode("utf-8")
            zout.writestr(item, data)
    return target.getvalue()


@pytest.mark.parametrize(
    "raw_literal",
    [
        "1.0000000000000001",
        "0.99999999999999999",
        "9007199254740990.5",
        "1e1000000000",
    ],
)
def test_decode_reports_rejects_non_integral_raw_xlsx_numbers_masked_by_float(raw_literal):
    xlsx = encode_reports([_sample_row()], summaries=[], metadata=[])
    patched = _replace_d2_with_raw_number(xlsx, raw_literal)
    rows = decode_reports(patched)
    assert len(rows) == 1
    assert rows[0]["errors"] == [
        {
            "row": 2,
            "field": "completed_quantity",
            "message": "本次完成数量要填 0 或正整数。",
        }
    ]


@pytest.mark.parametrize(("raw_qty", "expected_qty"), [(1, 100), (0.5, 50)])
def test_decode_reports_scales_percent_formatted_numeric_cells(raw_qty, expected_qty):
    from openpyxl import load_workbook

    xlsx = encode_reports([_sample_row()], summaries=[], metadata=[])
    book = load_workbook(BytesIO(xlsx))
    sheet = book.worksheets[0]
    sheet["D2"].value = raw_qty
    sheet["D2"].number_format = "0%"
    sheet["G2"].value = 0.07
    sheet["G2"].number_format = "0%"
    out = BytesIO()
    book.save(out)
    book.close()

    rows = decode_reports(out.getvalue())
    assert len(rows) == 1
    assert rows[0]["errors"] == []
    assert rows[0]["values"]["completed_quantity"] == expected_qty
    assert rows[0]["values"]["effective_processing_hours"] == 7.0


def _percent_xlsx_with_raw_d2_and_g2(raw_d2: str, raw_g2: str) -> bytes:
    from openpyxl import load_workbook

    xlsx = encode_reports([_sample_row()], summaries=[], metadata=[])
    book = load_workbook(BytesIO(xlsx))
    sheet = book.worksheets[0]
    sheet["D2"].value = 1
    sheet["D2"].number_format = "0%"
    sheet["G2"].value = 0.07
    sheet["G2"].number_format = "0%"
    out = BytesIO()
    book.save(out)
    book.close()

    source = BytesIO(out.getvalue())
    target = BytesIO()
    with ZipFile(source, "r") as zin, ZipFile(target, "w", ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                xml_text = data.decode("utf-8")
                xml_text, c1 = re.subn(
                    r'(<c r="D2"[^>]*><v>).*?(</v></c>)',
                    rf"\g<1>{raw_d2}\g<2>",
                    xml_text,
                    count=1,
                )
                xml_text, c2 = re.subn(
                    r'(<c r="G2"[^>]*><v>).*?(</v></c>)',
                    rf"\g<1>{raw_g2}\g<2>",
                    xml_text,
                    count=1,
                )
                assert c1 == 1 and c2 == 1
                data = xml_text.encode("utf-8")
            zout.writestr(item, data)
    return target.getvalue()


@pytest.mark.parametrize(
    ("raw_d2", "raw_g2"),
    [
        ("1.000000000000000000000000000001", "1e-1000000000"),
        ("1e1000000000", "1e1000000000"),
        ("1e-1000000000", "1e-1000000000"),
    ],
)
def test_decode_reports_rejects_percent_cells_beyond_default_decimal_context(raw_d2, raw_g2):
    patched = _percent_xlsx_with_raw_d2_and_g2(raw_d2, raw_g2)
    rows = decode_reports(patched)
    assert len(rows) == 1
    assert rows[0]["errors"] == [
        {
            "row": 2,
            "field": "completed_quantity",
            "message": "本次完成数量要填 0 或正整数。",
        },
        {
            "row": 2,
            "field": "effective_processing_hours",
            "message": "数量和工时要填 0 或大于 0 的有效数字。",
        },
    ]


@pytest.mark.parametrize(
    "cell_patch",
    [
        (r'<c r="D2"[^>]*>.*?</c>', '<c r="D2" t="n"><v>abc</v></c>'),
        (r'<c r="J2"[^>]*>.*?</c>', '<c r="J2" t="s"><v>999</v></c>'),
    ],
)
def test_decode_reports_rejects_corrupted_numeric_or_dangling_shared_string_cell(cell_patch):
    pattern, replacement = cell_patch
    xlsx = encode_reports([_sample_row()], summaries=[], metadata=[])
    source = BytesIO(xlsx)
    target = BytesIO()
    with ZipFile(source, "r") as zin, ZipFile(target, "w", ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                xml_text = data.decode("utf-8")
                xml_text, count = re.subn(pattern, replacement, xml_text, count=1)
                assert count == 1
                data = xml_text.encode("utf-8")
            zout.writestr(item, data)
    with pytest.raises(WorkbenchCommandRejected) as exc_info:
        decode_reports(target.getvalue())
    assert exc_info.value.code == "invalid_input"
    assert exc_info.value.status == 422


@pytest.mark.parametrize("bad_char", ["\x00", "\x07", "\x1f", "\ufffe", "\uffff"])
def test_report_input_and_xlsx_codec_reject_illegal_xml_characters(bad_char):
    with pytest.raises(WorkbenchCommandRejected):
        normalize_report_input(
            "create",
            {
                "completed_quantity": 1,
                "remark": f"bad{bad_char}text",
            },
        )
    with pytest.raises(WorkbenchCommandRejected):
        encode_reports([_sample_row(remark=f"bad{bad_char}text")], summaries=[], metadata=[])
    with pytest.raises(WorkbenchCommandRejected):
        encode_reports([_sample_row()], summaries=[[f"bad{bad_char}summary"]], metadata=[])
    with pytest.raises(WorkbenchCommandRejected):
        encode_reports([_sample_row()], summaries=[], metadata=[[f"bad{bad_char}meta"]])


def test_encode_reports_rejects_oversized_auxiliary_sheet_cells():
    oversized = "a" * 32768
    with pytest.raises(WorkbenchCommandRejected):
        encode_reports([_sample_row()], summaries=[[oversized]], metadata=[])
    with pytest.raises(WorkbenchCommandRejected):
        encode_reports([_sample_row()], summaries=[], metadata=[[oversized]])


def test_decode_reports_rejects_shared_strings_with_illegal_xml_char():
    xlsx = encode_reports([_sample_row()], summaries=[], metadata=[])
    source = BytesIO(xlsx)
    target = BytesIO()
    with ZipFile(source, "r") as zin, ZipFile(target, "w", ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            zout.writestr(item, zin.read(item.filename))
        zout.writestr(
            "xl/sharedStrings.xml",
            (
                '<?xml version="1.0" encoding="UTF-8"?>'
                '<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
                "<si><t>bad\ufffetext</t></si></sst>"
            ).encode("utf-8"),
        )
    with pytest.raises(WorkbenchCommandRejected) as exc_info:
        decode_reports(target.getvalue())
    assert exc_info.value.code == "invalid_input"
    assert exc_info.value.status == 422


def test_decode_reports_rejects_alternate_workbook_part_in_content_types():
    xlsx = encode_reports([_sample_row()], summaries=[], metadata=[])
    source = BytesIO(xlsx)
    target = BytesIO()
    with ZipFile(source, "r") as zin, ZipFile(target, "w", ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                xml_text = data.decode("utf-8").replace(
                    'PartName="/xl/workbook.xml"',
                    'PartName="/xl/altWorkbook.xml"',
                )
                data = xml_text.encode("utf-8")
            zout.writestr(item, data)
        zout.writestr("xl/altWorkbook.xml", zin.read("xl/workbook.xml"))
    with pytest.raises(WorkbenchCommandRejected) as exc_info:
        decode_reports(target.getvalue())
    assert exc_info.value.code == "invalid_input"
    assert exc_info.value.status == 422


@pytest.mark.parametrize(
    "styles_xml",
    [
        (
            '<?xml version="1.0" encoding="UTF-8"?>'
            '<styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
            '<fills count="1"><fill/></fills>'
            "</styleSheet>"
        ),
        (
            '<?xml version="1.0" encoding="UTF-8"?>'
            '<styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
            '<cellStyleXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/></cellStyleXfs>'
            '<cellStyles count="1"><cellStyle name="Normal" xfId="99" builtinId="0"/></cellStyles>'
            "</styleSheet>"
        ),
    ],
)
def test_decode_reports_rejects_malformed_stylesheet_xml(styles_xml):
    xlsx = encode_reports([_sample_row()], summaries=[], metadata=[])
    source = BytesIO(xlsx)
    target = BytesIO()
    with ZipFile(source, "r") as zin, ZipFile(target, "w", ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = styles_xml.encode("utf-8") if item.filename == "xl/styles.xml" else zin.read(item.filename)
            zout.writestr(item, data)
    with pytest.raises(WorkbenchCommandRejected) as exc_info:
        decode_reports(target.getvalue())
    assert exc_info.value.code == "invalid_input"
    assert exc_info.value.status == 422


@pytest.mark.parametrize("raw_zero", ["-0", "-0.0", "0e999999999999999999"])
def test_decode_reports_accepts_negative_and_extreme_exponent_raw_zeros(raw_zero):
    patched = _percent_xlsx_with_raw_d2_and_g2(raw_zero, raw_zero)
    rows = decode_reports(patched)
    assert len(rows) == 1
    assert rows[0]["errors"] == []
    assert rows[0]["values"]["completed_quantity"] == 0
    assert rows[0]["values"]["effective_processing_hours"] == 0.0


def test_decode_reports_rejects_duplicate_sheets_element_in_workbook():
    xlsx = encode_reports([_sample_row()], summaries=[], metadata=[])
    source = BytesIO(xlsx)
    target = BytesIO()
    with ZipFile(source, "r") as zin, ZipFile(target, "w", ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/workbook.xml":
                xml_text = data.decode("utf-8").replace(
                    "</sheets>",
                    '</sheets><sheets><sheet name="Alt" sheetId="99" r:id="rId2"/></sheets>',
                )
                data = xml_text.encode("utf-8")
            zout.writestr(item, data)
    with pytest.raises(WorkbenchCommandRejected) as exc_info:
        decode_reports(target.getvalue())
    assert exc_info.value.code == "invalid_input"
    assert exc_info.value.status == 422
