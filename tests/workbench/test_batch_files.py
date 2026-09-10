"""Real first-sheet XLSX bytes, three modes, range exports and atomic receipts."""

import json
import zipfile
from io import BytesIO

import openpyxl
import pytest

from core.models.workbench_batch_file import HEADERS
from core.services.workbench.batch_file_codec import write_batch_file
from tests.workbench.batch_support import BASE, assert_error, batch_database, body, detail, list_data, ref_for, state
from tests.workbench.legacy_batch_lineage_copy_support import assert_dashboard_copy_items

_batch_fixture = batch_database


def uploaded(client, rows, mode="overwrite", content=None):
    content = content or write_batch_file(rows, template=True)
    return client.post(BASE + "/import-preview", data={"file": (BytesIO(content), "batch.xlsx"), "mode": mode,
        "scope": json.dumps({}), "snapshot_ref": list_data(client)["meta"]["snapshot_ref"]}, content_type="multipart/form-data")


def confirm(client, document, key="batch-file-confirm-00001"):
    return client.post(BASE + "/import-confirm", json=body(document["write_context"], {"preview_ref": document["preview_ref"]}, key))


def data_rows(response):
    assert response.status_code == 200, response.get_json()
    workbook = openpyxl.load_workbook(BytesIO(response.data), data_only=False)
    try:
        return list(workbook.worksheets[0].iter_rows(values_only=True))
    finally:
        workbook.close()


@pytest.mark.parametrize("mode", ("overwrite", "append"))
def test_modes_preserve_existing_empty_columns_and_no_automatic_operations(batch_client, mode):
    client = batch_client
    before = state(client)
    response = uploaded(client, [["FREE-001", "P1", 5, None, None, None, None, "changed"],
                                 ["NEW-XLSX", "P1", 3, "2028-02-29", "急件", "未齐套", None, "import"]], mode)
    assert response.status_code == 200, response.get_json()
    document = response.get_json()["data"]
    assert document["can_confirm"] and state(client) == before
    response = confirm(client, document)
    assert response.status_code == 200, response.get_json()
    record = detail(client)["data"]
    assert record["fields"]["remark"] == ("changed" if mode == "overwrite" else "keep-hidden")
    assert record["fields"]["ready_status"] == "no" and record["fields"]["priority"] == "normal"
    new = detail(client, ref_for(client, key="NEW-XLSX"))["data"]
    assert new["operations"] == [] and new["fields"]["due_date"] == "2028-02-29"
    after = state(client)
    assert after[0] == before[0]
    assert_dashboard_copy_items(before[1], after[1], 1)
    for table in before[1]:
        if table not in ("Batches", "WorkbenchEntityRefs", "WorkbenchCommandReceipts", "WorkbenchDashboardItems"):
            assert after[1][table] == before[1][table], table
    saved = state(client)
    replay = confirm(client, document)
    assert replay.get_json()["replayed"] and state(client) == saved


def test_replace_shows_entire_deletion_set_and_blocks_planned_fact(batch_client):
    client = batch_client
    before = state(client)
    response = uploaded(client, [["NEW-XLSX", "P1", 3, None, "普通", "齐套", None, None]], "replace")
    document = response.get_json()["data"]
    assert document["can_confirm"] is False and document["write_context"] is None
    assert {row["before"]["business_code"] for row in document["deleted"]} == {"B1", "FREE-001"}
    assert any(row["errors"] for row in document["deleted"]) and state(client) == before


def test_replace_unreferenced_is_atomic_and_new_identity(batch_client):
    client, conn = batch_client, batch_client.batch_conn
    conn.execute("DELETE FROM Schedule")
    conn.execute("DELETE FROM BatchMaterials")
    conn.commit()
    old_ref = ref_for(client)
    document = uploaded(client, [["FREE-001", "P1", 3, None, "普通", "齐套", None, None]], "replace").get_json()["data"]
    assert document["can_confirm"]
    response = confirm(client, document)
    assert response.status_code == 200, response.get_json()
    assert list_data(client)["data"]["page"]["total"] == 1 and ref_for(client) != old_ref
    assert not detail(client)["data"]["operations"]


@pytest.mark.parametrize("row", [[1, "P1", 3, None, None, None, None, None], ["NEW", "P1", 0, None, None, None, None, None],
                                 ["NEW", "P1", 2, "2026-02-29", None, None, None, None], ["NEW", "missing", 2, None, None, None, None, None]])
def test_invalid_file_rows_never_write(batch_client, row):
    before = state(batch_client)
    response = uploaded(batch_client, [row])
    assert response.status_code == 200, response.get_json()
    assert response.get_json()["data"]["can_confirm"] is False
    assert state(batch_client) == before


def test_template_bytes_and_selected_export_preserve_text_and_scope(batch_client):
    client = batch_client
    before = state(client)
    template = client.get(BASE + "/template")
    assert tuple(data_rows(template)[0]) == HEADERS
    assert state(client) == before
    conn = client.batch_conn
    conn.execute("UPDATE Batches SET remark='=FORMULA_NOT_EXECUTED' WHERE batch_id='FREE-001'")
    conn.commit()
    result = list_data(client, size=1)
    response = client.post(BASE + "/export-preview", json={"selection": "selected", "refs": [ref_for(client)],
                          "scope": {"size": 1, "snapshot_ref": result["meta"]["snapshot_ref"]}})
    assert response.status_code == 200, response.get_json()
    export = client.get(BASE + "/export", query_string={"export_ref": response.get_json()["data"]["export_ref"]})
    rows = data_rows(export)
    assert len(rows) == 2 and rows[1][0] == "FREE-001" and rows[1][7] == "=FORMULA_NOT_EXECUTED"
    workbook = openpyxl.load_workbook(BytesIO(export.data), data_only=False)
    sheet = workbook.active
    assert sheet is not None
    assert sheet["H2"].data_type == "s"
    workbook.close()
    assert export.headers["Cache-Control"] == "no-store" and export.headers["X-Workbench-Row-Count"] == "1"


def test_file_confirm_stale_rolls_back_and_duplicate_rejected(batch_client):
    client = batch_client
    rows = [["NEW", "P1", 2, None, None, None, None, None]]
    duplicate = uploaded(client, rows * 2).get_json()["data"]
    assert not duplicate["can_confirm"]
    document = uploaded(client, rows).get_json()["data"]
    client.batch_conn.execute("UPDATE Materials SET stock_qty=0")
    client.batch_conn.commit()
    before = state(client)
    assert_error(confirm(client, document), "stale_write")
    assert state(client) == before


def test_first_sheet_only_formulas_and_bool_quantities(batch_client):
    client = batch_client
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(HEADERS)
    sheet.append(["FIRST", "P1", 3, None, None, None, None, None])
    second = workbook.create_sheet("ignored")
    second.append(["not a batch sheet"])
    content = BytesIO()
    workbook.save(content)
    document = uploaded(client, [], content=content.getvalue()).get_json()["data"]
    assert document["can_confirm"] and document["count"] == 1
    assert document["warnings"][0]["code"] == "first_sheet_only"
    for value in (True, "=1+1"):
        sheet["C2"] = value
        content = BytesIO()
        workbook.save(content)
        before = state(client)
        document = uploaded(client, [], content=content.getvalue()).get_json()["data"]
        assert not document["can_confirm"] and state(client) == before
    workbook.close()


def test_corrupt_xlsx_is_known_rejection_not_unknown_commit(batch_client):
    before = state(batch_client)
    assert_error(uploaded(batch_client, [], content=b"invalid file"), "invalid_input")
    assert state(batch_client) == before
    content = BytesIO()
    with zipfile.ZipFile(content, "w") as archive:
        archive.writestr("[Content_Types].xml", b"<invalid xml")
    assert_error(uploaded(batch_client, [], content=content.getvalue()), "invalid_input")
    assert state(batch_client) == before
