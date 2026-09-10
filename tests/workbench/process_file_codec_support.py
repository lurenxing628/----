"""DB-free process file fixtures, package mutations and one-pass measurements."""

import csv
import time
import tracemalloc
from io import BytesIO, StringIO
from xml.etree import ElementTree
from zipfile import ZIP_DEFLATED, ZipFile

import openpyxl

NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
SHEET = "xl/worksheets/sheet1.xml"


def file_bytes(headers, rows, file_format, *, text=False):
    if file_format == "csv":
        buffer = StringIO(newline="")
        writer = csv.writer(buffer, lineterminator="\r\n")
        writer.writerow(headers)
        writer.writerows(rows)
        return buffer.getvalue().encode("utf-8-sig")
    wb = openpyxl.Workbook()
    try:
        ws = wb.active
        ws.append(headers)
        for row in rows:
            ws.append(row)
        if text:
            for row in ws.iter_rows():
                for cell in row:
                    if type(cell.value) is str:
                        cell.data_type = "s"
        buffer = BytesIO()
        wb.save(buffer)
        return buffer.getvalue()
    finally:
        wb.close()


def mutate_xml(content, mutation, member=SHEET):
    buffer = BytesIO()
    with ZipFile(BytesIO(content)) as source, ZipFile(buffer, "w", ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == member:
                root = ElementTree.fromstring(data)
                mutation(root)
                data = ElementTree.tostring(root, encoding="utf-8", xml_declaration=True)
            target.writestr(item, data)
    return buffer.getvalue()


def number_cell(content, coordinate, value):
    def mutation(root):
        cell = root.find(".//s:c[@r='" + coordinate + "']", NS)
        cell.clear()
        cell.attrib.update({"r": coordinate, "t": "n"})
        ElementTree.SubElement(cell, "{" + NS["s"] + "}v").text = value
    return mutate_xml(content, mutation)


def source_cells(content, file_format):
    if file_format == "csv":
        reader = csv.reader(StringIO(content.decode("utf-8-sig"), newline=""))
        yield from reader
        return
    wb = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=False)
    try:
        yield from wb.worksheets[0].iter_rows(values_only=True)
    finally:
        wb.close()


def measure(label, operation):
    tracemalloc.start()
    started = time.perf_counter()
    try:
        result = operation()
        elapsed = time.perf_counter() - started
        _, peak = tracemalloc.get_traced_memory()
    finally:
        tracemalloc.stop()
    print(f"process_codec {label} seconds={elapsed:.3f} peak_mib={peak / 1024 ** 2:.3f}")
    assert elapsed < 45, (label, elapsed)
    assert peak < 96 * 1024 ** 2, (label, peak)
    return result


class SinglePassRows:
    def __init__(self, kind, count):
        self.kind, self.count, self.iterations = kind, count, 0

    def __iter__(self):
        self.iterations += 1
        assert self.iterations == 1
        for index in range(self.count):
            if self.kind == "route":
                yield {"business_code": f"图{index:05d}", "label": "零件", "route_raw": "10铣20检", "remark": "首行\n末行"}
            else:
                yield {"business_code": f"图{index:05d}", "sequence": index + 1, "source": "internal",
                       "op_type_name": "铣", "setup_hours": 0, "unit_hours": 0.12345678901234567,
                       "external_days": None, "group_start": None, "group_end": None, "group_total_days": None}
