"""Populate a browser-downloaded 13-column workbook, without touching any database."""

import argparse
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from tests.workbench.test_round1_field_piece_files_support import fill_template, report_values


def selected_piece_file(content):
    book = load_workbook(BytesIO(content))
    sheet = book.active
    assert sheet is not None
    headers = {cell.value: index for index, cell in enumerate(sheet[1], 1)}
    assert len(headers) == 13
    for row in range(sheet.max_row, 1, -1):
        if sheet.cell(row, headers['单件编号']).value not in ('分件甲', '0'):
            sheet.delete_rows(row)
    assert sheet.max_row == 3
    output = BytesIO()
    book.save(output)
    book.close()
    return fill_template(output.getvalue(), {"分件甲": report_values(), "0": report_values()})


def rejected_file(content):
    book = load_workbook(BytesIO(content))
    sheet = book.active
    assert sheet is not None
    headers = {cell.value: index for index, cell in enumerate(sheet[1], 1)}
    assert len(headers) == 13 and sheet.max_row >= 11
    if sheet.max_row > 11:
        sheet.delete_rows(12, sheet.max_row - 11)
    quantity = headers['本次完成数量']
    for row in range(2, 12):
        sheet.cell(row, quantity, -1)
    output = BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("template", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--mode", choices=("pieces", "rejected"), default="pieces")
    args = parser.parse_args()
    source, output = args.template.resolve(), args.output.resolve()
    repo = Path(__file__).resolve().parents[2]
    if repo in source.parents or repo in output.parents or source.parent != output.parent:
        raise ValueError("Only sibling private downloaded files may be used")
    transform = selected_piece_file if args.mode == "pieces" else rejected_file
    output.write_bytes(transform(source.read_bytes()))


if __name__ == "__main__":
    main()
