"""CSV compatibility, preserved report provenance, and scoped refactor invariants."""

import ast
import csv
import json
from io import BytesIO, StringIO
from pathlib import Path

from openpyxl import load_workbook
from radon.complexity import cc_visit

from core.services.workbench.actual_gantt_export import HEADERS
from tests.workbench.execution_ledger_support import all_rows
from tests.workbench.execution_ledger_support import ledger_case as ledger_case
from tests.workbench.field_workspace_support import BASE, success
from tests.workbench.field_workspace_support import field_api as field_api
from tests.workbench.round1_field_piece_files_support import (
    ACTUAL,
    PIECES,
    by_piece,
    confirm,
    download,
    fill_template,
    report_values,
)
from tests.workbench.round1_field_piece_files_support import (
    piece_file_api as piece_file_api,
)

LEGACY_CSV_PREFIX = ('计划引用', '任务引用', '工序引用', '批次', '工序', '计划开工', '计划完工', '计划设备', '计划人员',
    '目标数量', '整道状态', '完成依据', '已知完成数量', '未知记录数', '整道实际完工', '数据质量',
    '报工引用', '报工单号', '录入依据计划', '录入依据任务', '本次开工', '本次结束', '本次数量', '有效加工小时',
    '实际设备', '实际人员', '备注', '登记时间', '剩余数量', '剩余计划开工', '剩余计划完工', '剩余设备', '剩余人员',
    '旧事实', '数据缺项', '数据截至', '快照引用', '服务端范围', '本地筛选', '计划事件类型', '计划时长秒', '计划占用资源')
PLAN_CSV_FIELDS = ('单件编号', '计划应做数量', '计划批次数量', '计划数量依据', '计划数量缺失原因')
SOURCE_FIELDS = ('piece_id', 'quantity', 'batch_quantity', 'quantity_basis', 'quantity_reason')


def csv_export(api, **query):
    reading = success(api.client.get(ACTUAL, query_string={'plan_ref': api.case.plan_ref(2)}))
    response = api.client.get(ACTUAL + '/export', query_string={'plan_ref': api.case.plan_ref(2),
        'snapshot_ref': reading['meta']['snapshot_ref'], 'format': 'csv', **query})
    assert response.status_code == 200, response.get_data(as_text=True)
    assert response.data.startswith(b'\xef\xbb\xbf')
    return reading, list(csv.reader(StringIO(response.data.decode('utf-8-sig'))))


def test_csv_old_prefix_exact_positions_values_and_plan_fields(piece_file_api):
    api = piece_file_api
    confirm(api, fill_template(download(api), {PIECES[0]: report_values()}))
    before = all_rows(api.case.conn)
    reading, rows = csv_export(api)
    assert tuple(HEADERS) == LEGACY_CSV_PREFIX + PLAN_CSV_FIELDS
    assert tuple(rows[0][:42]) == LEGACY_CSV_PREFIX
    by_task = {item['task']['task_ref']: item for item in reading['data']['items']}
    for row in rows[1:]:
        item = by_task[row[1]]
        task, execution = item['task'], item['execution']
        assert row[:5] == [task['plan_ref'], task['task_ref'], task['operation_ref'], task['batch_id'], '1 Turning']
        assert row[9] == str(execution['target_quantity'])
        assert row[42:] == ['' if task[key] is None else str(task[key]) for key in SOURCE_FIELDS]
        assert row[35:39] == [reading['meta']['as_of'], reading['meta']['snapshot_ref'],
            json.dumps(reading['data']['scope'], ensure_ascii=False, sort_keys=True, separators=(',', ':')), '{}']
        if task['piece_id'] == PIECES[0]:
            report = execution['reports'][0]
            assert row[16:20] == [report['report_ref'], report['report_no'],
                                   report['recorded_against_plan_ref'], report['recorded_against_task_ref']]
            assert row[22:24] == ['0', '0.0']
        else:
            assert row[16:28] == [''] * 12
    assert all_rows(api.case.conn) == before


def test_new_plan_export_reimport_keeps_original_refs_and_correction_history(piece_file_api):
    api = piece_file_api
    confirm(api, fill_template(download(api), {PIECES[0]: report_values()}))
    original = by_piece(api)[PIECES[0]]['execution']['reports'][0]
    correction = api.body(original['write_context'], {'original_revision_ref': original['revision_ref'],
        'effective_processing_hours': 0.25, 'reason': '中文原始工时复核'})
    success(api.client.post(BASE + '/reports/' + original['report_ref'] + '/correct', json=correction))
    ids = [row[0] for row in api.case.conn.execute('SELECT id FROM BatchOperations')]
    api.case.plan(3, ids)
    before = all_rows(api.case.conn)
    exported = download(api, 'export')
    book = load_workbook(BytesIO(exported))
    sheet = book.active
    assert sheet is not None
    record = next(row for row in sheet.iter_rows(min_row=2, values_only=True) if row[0] == original['report_no'])
    assert record[6] == 0.25 and record[10] != original['recorded_against_task_ref']
    metadata = list(book['录入信息'].iter_rows(min_row=2, values_only=True))
    assert metadata[0][3] == 1
    book.close()
    assert all_rows(api.case.conn) == before
    _, _, result = confirm(api, exported)
    assert result['result'] == 'unchanged'
    final = by_piece(api)[PIECES[0]]['execution']['reports'][0]
    for key in ('report_ref', 'report_no', 'operation_ref', 'recorded_against_plan_ref', 'recorded_against_task_ref'):
        assert final[key] == original[key]
    assert len(final['correction_history']) == 2
    assert final['correction_history'][1]['reason'] == '中文原始工时复核'
    assert final['correction_history'][0]['after']['effective_processing_hours'] == 0
    assert final['correction_history'][1]['before']['effective_processing_hours'] == 0
    for table in ('WorkbenchProductionReports', 'WorkbenchProductionReportRevisions', 'OperationExecutionEvents'):
        assert all_rows(api.case.conn)[table] == before[table]


def test_owned_sources_keep_python38_grammar_and_complexity_limit():
    root = Path(__file__).resolve().parents[2]
    files = ['field_report_files.py', 'field_report_files_codec.py', 'field_report_files_xml.py',
             'field_report_files_identity.py', 'actual_gantt_export.py', 'actual_gantt_scope.py', 'field_workspace.py']
    for name in files:
        source = (root / 'core/services/workbench' / name).read_text(encoding='utf-8')
        ast.parse(source, feature_version=8)
        blocks = cc_visit(source)
        for block in blocks:
            assert block.complexity <= 15, (name, block.name, block.complexity)
            for method in getattr(block, 'methods', []):
                assert method.complexity <= 15, (name, method.name, method.complexity)
