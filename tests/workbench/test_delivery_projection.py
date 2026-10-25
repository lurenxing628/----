"""Shared due and whole-chain contracts, including adopted and staged results."""

import csv
import json
from datetime import datetime
from io import BytesIO, StringIO

import pytest
from openpyxl import load_workbook

from core.infrastructure.transaction import TransactionManager
from core.models.workbench_batch_material import normalize_material_changes
from core.models.workbench_plan_scope import PlanReadScope
from core.models.workbench_run_candidate import RunCandidateReadScope
from core.services.workbench.batch.materials import WorkbenchBatchMaterialService
from core.services.workbench.dashboard.service import WorkbenchDashboardService
from core.services.workbench.facts.delivery_projection import (
    delivery_batch_delta,
    delivery_summary,
    project_delivery_batch,
    task_intervals,
)
from core.services.workbench.plan.export import HEADERS, write_plan_export
from core.services.workbench.plan.queries import WorkbenchPlanQueryService
from core.services.workbench.run.candidate_analysis import read_candidate_analysis
from core.services.workbench.run.candidates import WorkbenchRunCandidateQueryService
from tests.workbench.dashboard_support import dashboard_case as _dashboard_case  # noqa: F401
from tests.workbench.run_candidate_adoption_support import INTENT, service
from tests.workbench.run_candidate_support import candidate_case as _candidate_case  # noqa: F401
from tests.workbench.run_candidate_support import compute


def batch_delivery(due="2026-10-13", *, finish="2026-10-13T12:00:00", missing=False, deferred=False, part_name="零件"):
    batch = {"batch_id": "B1", "part_no": "P1", "part_name": part_name, "due_date": due,
             "quantity": 3, "priority": "urgent"}
    rows = [{"schedule_id": 1, "op_id": 1, "batch_id": "B1", "start_time": "2026-10-13T08:00:00", "end_time": finish}]
    operations = [{"op_id": 1}, {"op_id": 2}] if missing else [{"op_id": 1}]
    return project_delivery_batch(batch, "batch-ref", rows, operations, task_intervals(rows), False, [],
                                  deferred_operation_ids={2} if deferred else set())


@pytest.mark.parametrize("due,state,risk,issue", [
    ("2026-10-13", "specified", "on_time", None),
    ("9999-12-31", "unspecified", "not_applicable", None),
    (None, "invalid", "unknown", "due_date_missing"),
    ("", "invalid", "unknown", "due_date_missing"),
    ("2026-02-30", "invalid", "unknown", "due_date_invalid"),
])
def test_due_state_is_independent_of_whole_chain_completion(due, state, risk, issue):
    row = batch_delivery(due, part_name=None)
    assert row["due_state"] == state and row["risk"] == risk and row["due_issue"] == issue
    assert row["schedule_complete"] and row["completion_state"] == row["completeness"] == "complete"
    assert row["planned_finish"] == "2026-10-13T12:00:00"
    assert row["delay_hours"] == (0 if state == "specified" else None)
    assert "part_label_missing" in row["issues"]
    if issue:
        assert issue in row["issues"]


def test_due_coverage_excludes_unspecified_and_keeps_invalid_in_unknown():
    rows = [batch_delivery(), batch_delivery(finish="2026-10-14T00:00:00.000001"),
            batch_delivery(missing=True, deferred=True), batch_delivery(None), batch_delivery("9999-12-31")]
    for index, row in enumerate(rows):
        row["batch_ref"] = str(index)
    summary = delivery_summary(rows)
    assert summary["state"] == "partial"
    assert (summary["due_eligible_count"], summary["excluded_due_count"], summary["invalid_due_count"]) == (3, 1, 1)
    assert (summary["due_evaluation_count"], summary["known_due_count"], summary["unknown_due_count"]) == (4, 2, 2)
    assert summary["overdue_count"] is summary["total_tardiness_hours"] is None
    assert summary["known_overdue_count"] == 1 and summary["known_tardiness_hours"] > 0
    assert summary["deferred_batch_count"] == 1
    assert summary["unknown_batch_refs"] == ["2", "3"]
    not_applicable = delivery_summary([rows[-1]])
    assert not_applicable["state"] == "not_applicable"
    assert not_applicable["overdue_count"] is not_applicable["total_tardiness_hours"] is not_applicable["known_tardiness_hours"] is None
    assert delivery_summary([rows[3]])["state"] == "unavailable"


def test_dashboard_unspecified_due_is_excluded_from_evaluation_gaps(dashboard_case):
    case = dashboard_case
    case.conn.execute("UPDATE Batches SET due_date='9999-12-31' WHERE batch_id='DB1'")
    case.conn.commit()
    _, data = case.read()
    summary = data["categories"]["delivery"]
    assert summary["state"] == "not_applicable" and summary["risk_count"] is None
    assert summary["unknown_count"] == summary["assessed_count"] == summary["due_eligible_count"] == 0
    assert summary["excluded_due_count"] == 1 and summary["evaluation_gaps"] == []


def test_legal_material_deferral_is_distinct_from_ordinary_omission():
    deferred = batch_delivery(missing=True, deferred=True)
    omitted = batch_delivery(missing=True)
    assert deferred["completion_state"] == "deferred"
    assert (deferred["deferred_operation_count"], deferred["ordinary_unscheduled_operation_count"]) == (1, 0)
    assert omitted["completion_state"] == "incomplete"
    assert (omitted["deferred_operation_count"], omitted["ordinary_unscheduled_operation_count"]) == (0, 1)
    for row in (deferred, omitted):
        assert not row["schedule_complete"] and row["planned_finish"] is row["delay_hours"] is None
        assert row["partial_planned_finish"] == "2026-10-13T12:00:00"
        assert row["risk"] == "unknown"


def test_delivery_difference_uses_only_whole_batch_finishes_and_unrounded_tardiness():
    before = batch_delivery(finish="2026-10-14T00:00:00.000001")
    after = batch_delivery()
    change = delivery_batch_delta(before, after)
    assert change["delay_delta_hours"] < 0 and change["finish_delta_hours"] < 0
    assert change["change_kinds"] == ["earlier", "relieved_overdue"]
    partial = delivery_batch_delta(after, batch_delivery(missing=True))
    assert partial["finish_delta_hours"] is partial["delay_delta_hours"] is None
    assert partial["change_kinds"] == ["completion_worsened"]


def test_unspecified_candidate_analysis_and_adopted_plan_share_projection(candidate_case):
    case = candidate_case
    case.conn.execute("UPDATE Batches SET due_date='9999-12-31',part_name='零件'")
    case.conn.commit()
    _, refs = compute(case)
    candidate = refs[0]
    saved = case.conn.execute("SELECT artifact_json FROM WorkbenchRunCandidates WHERE candidate_ref=?", (candidate,)).fetchone()[0]
    workspace, _ = WorkbenchRunCandidateQueryService(case.conn).workspace(RunCandidateReadScope(candidate))
    assert workspace["delivery_risks"]["items"][0]["risk"] == "not_applicable"
    assert workspace["delivery_risks"]["summary"]["state"] == "not_applicable"
    analysis, _ = read_candidate_analysis(case.conn, candidate)
    for key in ("overdue_count", "total_tardiness_hours"):
        assert analysis["metrics"][key]["state"] == "not_applicable" and analysis["metrics"][key]["value"] is None
    assert json.loads(saved) == json.loads(case.conn.execute("SELECT artifact_json FROM WorkbenchRunCandidates WHERE candidate_ref=?", (candidate,)).fetchone()[0])
    adopter = service(case.conn)
    preview = adopter.preview(candidate)
    assert preview["validation"]["can_adopt"]
    result = adopter.adopt(candidate, preview["write_context"]["write_token"], "delivery-unspecified-adopt", INTENT)
    plan_ref = result["data"]["official_plan"]["plan_ref"]
    reader = WorkbenchPlanQueryService(case.conn)
    with reader.read_snapshot():
        plan, _ = reader.workspace(PlanReadScope(plan_ref))
    projection = plan["projections"]["delivery_risks"]
    assert projection["state"] == projection["summary"]["state"] == "not_applicable"
    assert projection["items"][0]["due_state"] == "unspecified"
    assert projection["items"][0]["schedule_complete"]
    assert projection["summary"]["due_eligible_count"] == 0 and projection["summary"]["excluded_due_count"] == 1
    item = projection["items"][0]
    snapshot = {"snapshot_ref": "unspecified-plan-export", "as_of": "2026-09-10T12:00:00"}
    exported = write_plan_export(plan, snapshot, "csv")
    rows = list(csv.DictReader(StringIO(exported.content.decode("utf-8-sig"))))
    assert len(rows) == exported.row_count == plan["task_count"]
    assert tuple(rows[0]) == HEADERS
    for row in rows:
        assert row["交付风险"] == "'未设置交期"
        assert row["预计完工"] == "'" + item["planned_finish"]
        assert row["交期"] == "'9999-12-31"
        assert all(row[key] == r"'\N" for key in ("交期截止（不含）", "超期小时", "超期天数"))
    workbook = load_workbook(BytesIO(write_plan_export(plan, snapshot, "xlsx").content), read_only=True)
    try:
        values = list(workbook.worksheets[0].iter_rows(values_only=True))
        assert values[0] == HEADERS and len(values) == plan["task_count"] + 1
        for value in values[1:]:
            row = dict(zip(HEADERS, value))
            assert row["交付风险"] == "未设置交期"
            assert row["预计完工"] == item["planned_finish"]
            assert row["交期"] == "9999-12-31"
            assert all(row[key] == r"\N" for key in ("交期截止（不含）", "超期小时", "超期天数"))
    finally:
        workbook.close()


def test_stage_adoption_keeps_deferred_operation_evidence(candidate_case):
    case = candidate_case
    second = case.operation(seq=2)
    case.batch("B2")
    third = case.operation(batch_id="B2", op_code="B2-1")
    case.conn.execute("INSERT INTO Materials(material_id,name,unit) VALUES ('STEEL','钢材','件')")
    case.conn.commit()
    with TransactionManager(case.conn).transaction():
        for batch_id, operation_id in (("B1", second), ("B2", third)):
            operation_ref = case.conn.execute("SELECT ref FROM WorkbenchPlanSourceRefs WHERE kind='operation' AND source_key=? AND active=1", (str(operation_id),)).fetchone()[0]
            WorkbenchBatchMaterialService(case.conn).apply_normalized(case.ref("batch", batch_id), normalize_material_changes({"removed_keys": [], "rows": [{
                "row_key": None, "material_ref": case.ref("material", "STEEL"), "required_quantity": 3,
                "available_quantity": 0, "operation_ref": operation_ref, "arrivals": []}]}))
    _, refs = compute(case, case.settings(material_strategy="stage", batch_refs=[case.ref("batch", "B1"), case.ref("batch", "B2")]))
    assert refs
    workspace, _ = WorkbenchRunCandidateQueryService(case.conn).workspace(RunCandidateReadScope(refs[0]))
    assert len(workspace["delivery_risks"]["items"]) == 2
    assert all(batch["completion_state"] == "deferred" and batch["ordinary_unscheduled_operation_count"] == 0
               for batch in workspace["delivery_risks"]["items"])
    adopter = service(case.conn)
    preview = adopter.preview(refs[0])
    assert preview["validation"]["can_adopt"], preview
    result = adopter.adopt(refs[0], preview["write_context"]["write_token"], "delivery-stage-adoption", INTENT)
    plan = result["data"]["official_plan"]
    audit = json.loads(case.conn.execute("SELECT result_summary FROM ScheduleHistory WHERE version=?", (plan["version"],)).fetchone()[0])
    assert audit["deferred_operation_ids"] == [second, third]
    reader = WorkbenchPlanQueryService(case.conn)
    with reader.read_snapshot():
        data, _ = reader.workspace(PlanReadScope(plan["plan_ref"]))
    items = data["projections"]["delivery_risks"]["items"]
    assert {batch["batch_id"] for batch in items} == {"B1", "B2"}
    assert all(batch["completion_state"] == "deferred" and batch["risk"] == "unknown" for batch in items)
    assert all(batch["planned_finish"] is batch["delay_hours"] is None for batch in items)
    assert next(batch for batch in items if batch["batch_id"] == "B2")["task_count"] == 0
    dashboard = WorkbenchDashboardService(case.conn, clock=lambda: datetime(2026, 9, 10, 12))
    with dashboard.read_snapshot():
        data = dashboard.read()
    assert data["categories"]["delivery"]["unknown_count"] == 2
    assert {row["source_ref"] for row in data["categories"]["delivery"]["evaluation_gaps"]} == {
        batch["batch_ref"] for batch in items}
    # Older stage adoptions carry source identities, before explicit pending IDs.
    audit.pop("deferred_operation_ids")
    case.conn.execute("UPDATE ScheduleHistory SET result_summary=? WHERE version=?", (json.dumps(audit), plan["version"]))
    case.conn.commit()
    with reader.read_snapshot():
        data, _ = reader.workspace(PlanReadScope(plan["plan_ref"]))
    assert len(data["projections"]["delivery_risks"]["items"]) == 2
    assert all(batch["completion_state"] == "deferred" for batch in data["projections"]["delivery_risks"]["items"])
