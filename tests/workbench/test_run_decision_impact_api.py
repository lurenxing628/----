"""Delivery details keep archived identities, filters and downloads aligned."""

import csv
import io
import json
from datetime import datetime, timedelta

import pytest
from openpyxl import load_workbook

from core.services.workbench.facts.candidate_store import CandidateStore
from core.services.workbench.facts.decision_metrics import project_official_core_metrics
from core.services.workbench.run.decision_queries import read_decision_details
from core.services.workbench.run.decision_source import DecisionSource
from core.services.workbench.run.worker import WorkbenchRunWorker
from tests.workbench.decision_history_support import clone_run
from tests.workbench.legacy_run_archive_support import install_different_work_archive
from tests.workbench.run_candidate_adoption_support import INTENT, service
from tests.workbench.run_candidate_support import BASE, api, compute, corrupt_update
from tests.workbench.run_candidate_support import candidate_case as _candidate_case  # noqa: F401
from tests.workbench.test_material_stage_release import add_requirement
from tests.workbench.test_run_decision_api import _assert_summary_matches_full, decision, decision_summary


def detail(client, run_ref, section, **query):
    response = client.get(BASE + "/runs/" + run_ref + "/decision/" + section, query_string=query)
    assert response.status_code == 200, response.get_json()
    assert response.headers["Cache-Control"] == "no-store"
    return response.get_json()


def operation_ref(case, op_id):
    return case.conn.execute(
        "SELECT ref FROM WorkbenchPlanSourceRefs WHERE kind='operation' AND source_key=? AND active=1",
        (str(op_id),),
    ).fetchone()[0]


def many_batches(case):
    codes = ["B1", "000001"] + [f"B{index:02d}" for index in range(3, 28)]
    for code in codes[1:]:
        due = "9999-12-31" if code == "B25" else None if code == "B26" else "2026-09-25"
        case.batch(code, due_date=due, part_name=r"零件'甲\N")
        case.operation(code, unit_hours=0.05)
    case.conn.execute("UPDATE Batches SET part_name=? WHERE batch_id='B1'", (r"零件'甲\N",))
    case.conn.execute("UPDATE Parts SET part_name=? WHERE part_no='P1'", (r"零件'甲\N",))
    case.conn.execute("UPDATE BatchOperations SET unit_hours=0.05 WHERE id=?", (case.op_id,))
    case.conn.commit()
    run_ref, refs = compute(case, case.settings(*codes))
    # Save a historical copy with the bad date before deriving its read models.
    case.history_refs = {"run_ref": run_ref, "candidate_ref": refs[0]}
    accepted = case.conn.execute("SELECT accepted_at FROM WorkbenchRunJobs WHERE run_ref=?", (run_ref,)).fetchone()[0]
    columns = [row[1] for row in case.conn.execute("PRAGMA table_info(Batches)")]

    def change_facts(facts):
        for row in facts["tables"]["Batches"]:
            if row[columns.index("batch_id")] == "B27":
                row[columns.index("due_date")] = "2026-99-99"

    run_ref, _ = clone_run(case, accepted, change_facts=change_facts)
    refs = [row[0] for row in case.conn.execute(
        "SELECT candidate_ref FROM WorkbenchRunCandidates WHERE run_ref=? ORDER BY sequence", (run_ref,))]
    return run_ref, refs, codes


def test_mixed_due_values_keep_nulls_last_and_permanent_ties_across_pages(candidate_case):
    case = candidate_case
    run_ref, refs, codes = many_batches(case)
    client, _ = api(case)
    query = {"selected_candidate_ref": refs[-1], "reference_candidate_ref": refs[-1],
             "sort": "after_delay_hours", "order": "desc"}
    with CandidateStore(case.conn).snapshot():
        batch_source = DecisionSource(
            CandidateStore(case.conn), run_ref,
            selected_candidate_ref=refs[-1], reference_candidate_ref=refs[-1],
        )
        _ = batch_source.batch_impacts
        from core.services.workbench.run.decision_impact import filter_options
        _ = filter_options(batch_source, "batches")
        _ = filter_options(batch_source, "operations")
        assert "facts" not in batch_source.__dict__
        _ = batch_source.operation_impacts
        assert not batch_source._raw_tasks and not batch_source._projected
    first = detail(client, run_ref, "batches", **query)
    assert first["data"]["candidate"]["candidate_ref"] == refs[-1]
    page = first["data"]["page"]
    assert (page["number"], page["size"], page["total"], page["has_more"]) == (1, 20, 27, True)
    second = detail(client, run_ref, "batches", **dict(query, page=2, snapshot_ref=first["meta"]["snapshot_ref"]))
    wide = detail(client, run_ref, "batches", **dict(query, size=50, snapshot_ref=first["meta"]["snapshot_ref"]))
    rows = first["data"]["items"] + second["data"]["items"]
    assert rows == wide["data"]["items"]
    assert len(rows) == len({row["batch_ref"] for row in rows}) == len(codes)
    assert wide["data"]["page"]["has_more"] is False
    values = [row["after"]["delay_hours"] for row in rows]
    assert values == [0] * 24 + [None] * 3
    for cohort in (rows[:24], rows[24:]):
        assert [row["batch_ref"] for row in cohort] == sorted(row["batch_ref"] for row in cohort)
    by_id = {row["batch_id"]: row for row in rows}
    attention = detail(client, run_ref, "batches", selected_candidate_ref=refs[-1], reference_candidate_ref=refs[-1],
        delivery_filter="attention", selected_batch_ref=case.ref("batch", "B1"), sort="current_priority")["data"]
    assert attention["items"] == [] and attention["page"]["total"] == 0
    assert attention["selected_batch"] == by_id["B1"]
    assert attention["selected_batch"]["after"]["is_overdue"] is False
    overview = decision(client, run_ref, selected_candidate_ref=refs[-1], reference_candidate_ref=refs[-1])["data"]
    candidate = next(row for row in overview["candidates"] if row["candidate_ref"] == refs[-1])
    changes = {key: candidate["metrics"][key] for key in ("changed_operation_count", "machine_change_count")}
    assert all(data["summary"]["operation_changes"] == changes for data in
               (first["data"], second["data"], wide["data"], attention))
    assert all(value["total_count"] == 27 and value["scope"] == "full_admission" for value in changes.values())
    assert by_id["B25"]["due_state"] == "unspecified" and by_id["B25"]["due_issue"] is None
    assert by_id["B26"]["due_state"] == "invalid" and by_id["B26"]["due_issue"] == "due_date_missing"
    assert by_id["B27"]["due_state"] == "invalid" and by_id["B27"]["due_issue"] == "due_date_invalid"
    summary = wide["data"]["summary"]["after"]
    assert (summary["due_eligible_count"], summary["excluded_due_count"], summary["invalid_due_count"]) == (24, 1, 2)
    assert summary["known_due_count"] == 24 and summary["unknown_due_count"] == 2
    assert summary["total_tardiness_hours"] is None and summary["known_tardiness_hours"] == 0
    operations = detail(client, run_ref, "operations", selected_candidate_ref=refs[-1],
                        reference_candidate_ref=refs[-1], sort="sequence", order="desc", size=50)["data"]
    operation_refs = [row["operation_ref"] for row in operations["items"]]
    assert operation_refs == sorted(operation_refs) and len(operation_refs) == 27
    assert operations["summary"]["operation_changes"] == changes
    path = BASE + "/runs/" + run_ref + "/decision/batches"
    for invalid in ({"page": 0}, {"size": 30}):
        rejected = client.get(path, query_string=dict(query, **invalid))
        assert rejected.status_code == 422, rejected.get_json()


@pytest.mark.parametrize("mode", ["stage", "ordinary_missing"])
def test_operation_filters_preserve_identity_and_distinguish_material_deferral(candidate_case, mode, monkeypatch):
    case = candidate_case
    if mode == "stage":
        second_id = case.operation(seq=2)
        add_requirement(case, second_id, [])
        run_ref, refs = compute(case, case.settings(material_strategy="stage"))
        selected, reference = refs[0], refs[0]
        batch = "B1"
    else:
        case.conn.execute("UPDATE Batches SET quantity=1,due_date='2026-09-09' WHERE batch_id='B1'")
        case.conn.execute("UPDATE BatchOperations SET unit_hours=5 WHERE id=?", (case.op_id,))
        case.batch("B2", quantity=1, due_date="2026-09-10")
        second_id = case.operation("B2", unit_hours=6)
        case.config(dispatch_mode="batch_order", sort_strategy="due_date_first")
        with monkeypatch.context() as archive_patch:
            install_different_work_archive(archive_patch)
            run_ref, refs = compute(case, case.settings("B1", "B2", end_date="2026-09-09"))
        selected = refs[0]
        client, _ = api(case)
        catalog = decision(client, run_ref)["data"]["candidates"]
        reference = next(row["candidate_ref"] for row in catalog if row["kind"] == "critical_chain")
        batch = "B2"
    if mode == "stage":
        client, _ = api(case)
    query = {"selected_candidate_ref": selected, "reference_candidate_ref": reference}
    batch_ref, second_ref = case.ref("batch", batch), operation_ref(case, second_id)
    batches = detail(client, run_ref, "batches", **dict(query, batch_ref=batch_ref))["data"]
    assert len(batches["items"]) == 1
    item = batches["items"][0]
    assert item["after"]["schedule_complete"] is False and item["after"]["planned_finish"] is None
    expected_state = "deferred" if mode == "stage" else "incomplete"
    assert item["after"]["completion_state"] == expected_state
    assert item["deferred_operation_refs"] == ([second_ref] if mode == "stage" else [])
    assert item["unplanned_operation_refs"] == ([] if mode == "stage" else [second_ref])
    assert item["after"]["deferred_operation_count"] == (1 if mode == "stage" else 0)
    assert item["after"]["ordinary_unscheduled_operation_count"] == (0 if mode == "stage" else 1)
    operations = detail(client, run_ref, "operations", **dict(query, batch_ref=batch_ref))["data"]
    missing = next(row for row in operations["items"] if row["operation_ref"] == second_ref)
    assert missing["after"] is None
    assert missing["after_state"]["status"] == ("deferred" if mode == "stage" else "missing")
    assert missing["status_changed"] is (mode != "stage")
    assert missing["drill"]["selected"]["operation_ref"] == second_ref
    assert missing["drill"]["selected"]["candidate_ref"] == selected
    assert missing["effort"]["target_quantity"] == (3 if mode == "stage" else 1)
    machine_ref = case.ref("machine", "M1")
    filtered = detail(client, run_ref, "operations", **dict(query, resource_ref=machine_ref))["data"]
    overview = decision(client, run_ref, **query)["data"]
    candidate = next(row for row in overview["candidates"] if row["candidate_ref"] == selected)
    changes = {key: candidate["metrics"][key] for key in ("changed_operation_count", "machine_change_count")}
    assert batches["summary"]["operation_changes"] == operations["summary"]["operation_changes"] == filtered["summary"]["operation_changes"] == changes
    assert all(value["total_count"] == 2 and value["scope"] == "full_admission" for value in changes.values())
    if mode == "stage":
        assert {row["operation_ref"] for row in filtered["items"]} == {operation_ref(case, case.op_id)}
        unknown = detail(client, run_ref, "operations", **dict(query, batch_ref=batch_ref, change_kind="unknown"))["data"]
        assert [row["operation_ref"] for row in unknown["items"]] == [second_ref]
        assert missing["before_state"]["status"] == "deferred"
        assert missing["changed_fields"] == [] and missing["change_kinds"] == ["unknown"]
        current, _ = read_decision_details(case.conn, run_ref, "batches", **query, delivery_filter="attention", sort="current_priority")
        assert [row["batch_ref"] for row in current["items"]] == [batch_ref]
        assert current["items"][0]["attention_kinds"] == ["deferred"]
        assert current["items"][0]["operation_unknown_count"] == 1
        assert missing["after_state"]["reason"]["message"].startswith("缺料未安排")
    else:
        assert {row["operation_ref"] for row in filtered["items"]} == {operation_ref(case, case.op_id), second_ref}
        changed = detail(client, run_ref, "operations", **dict(query, change_kind="changed"))["data"]
        assert {row["operation_ref"] for row in changed["items"]} == {operation_ref(case, case.op_id), second_ref}
        assert {kind for row in changed["items"] for kind in row["change_kinds"]} == {"added", "removed", "status_changed"}
        assert missing["before_state"]["status"] == "arranged" and "status" in missing["changed_fields"]


def test_admission_official_details_and_gantt_keep_archived_facts_after_live_changes(candidate_case):
    case = candidate_case
    case.conn.execute("UPDATE Batches SET part_name='Archived batch part' WHERE batch_id='B1'")
    case.conn.commit()
    _, old_refs = compute(case)
    adopter = service(case.conn)
    preview = adopter.preview(old_refs[0])
    assert preview["validation"]["can_adopt"] is True, preview
    adopted = adopter.adopt(old_refs[0], preview["write_context"]["write_token"], "impact-first-real-adopt", INTENT)
    official = adopted["data"]["official_plan"]
    old = case.conn.execute("SELECT start_time,end_time FROM Schedule WHERE version=? AND op_id=?",
                            (official["version"], case.op_id)).fetchone()
    case.conn.execute("UPDATE BatchOperations SET unit_hours=0.5 WHERE id=?", (case.op_id,))
    case.conn.execute("UPDATE WorkbenchCalendarDefaults SET periods_json=? WHERE singleton=1",
                     ('[{"start":"07:15","end":"16:00","day_offset":0}]',))
    case.conn.commit()
    accepted = case.accept(key="impact-second-real-run")
    result = WorkbenchRunWorker(case.conn).execute(accepted["run_ref"])
    run_ref, selected = result["run_ref"], result["candidates"][0]["candidate_ref"]
    client, _ = api(case)
    query = {"comparison_mode": "admission_official", "selected_candidate_ref": selected}
    before = detail(client, run_ref, "operations", **query)
    batches = detail(client, run_ref, "batches", **query)
    workspace = detail(client, run_ref, "reference-workspace", **query)
    row = before["data"]["items"][0]
    assert row["before"]["start"] == datetime.fromisoformat(old[0]).isoformat()
    assert row["before"]["end"] == datetime.fromisoformat(old[1]).isoformat()
    assert row["start_delta_hours"] == -0.75 and row["end_delta_hours"] == 0
    assert row["changed_fields"] == ["start"]
    batch = batches["data"]["items"][0]
    assert batch["finish_delta_hours"] == 0 and batch["delay_delta_hours"] == 0
    assert batch["operation_change_count"] == 1 and batch["change_kinds"] == ["operation_changed"]
    changed = detail(client, run_ref, "batches", **dict(query, change_kind="changed"))["data"]
    assert [item["batch_ref"] for item in changed["items"]] == [batch["batch_ref"]]
    assert before["data"]["baseline"]["baseline_ref"] == official["plan_ref"]
    assert before["data"]["baseline"]["source_candidate_ref"] == old_refs[0]
    reference_drill = row["drill"]["reference"]
    assert reference_drill["plan_ref"] == official["plan_ref"] and reference_drill["candidate_ref"] == selected
    assert reference_drill["reference_view"] is True
    assert workspace["data"]["tasks"][0]["end"] == row["before"]["end"]
    assert all(workspace["data"]["capabilities"][name] is False for name in ("edit_draft", "adopt", "report_actual"))
    case.conn.execute("UPDATE Schedule SET start_time='2026-09-12T08:00:00',end_time='2026-09-12T10:00:00' WHERE version=?",
                      (official["version"],))
    case.conn.execute("UPDATE Batches SET due_date='2026-09-01',quantity=9,part_name='Current batch part' WHERE batch_id='B1'")
    case.conn.execute("UPDATE Parts SET part_name='Current part' WHERE part_no='P1'")
    case.conn.execute("UPDATE Machines SET name='Current lathe' WHERE machine_id='M1'")
    case.batch("B-LATE")
    late_id = case.operation("B-LATE")
    case.conn.commit()
    for section, original in (("operations", before), ("batches", batches), ("reference-workspace", workspace)):
        reread = detail(client, run_ref, section, **dict(query, snapshot_ref=original["meta"]["snapshot_ref"]))
        assert reread["data"] == original["data"]
        assert reread["meta"]["snapshot_ref"] == original["meta"]["snapshot_ref"]
    assert operation_ref(case, late_id) not in before["data"]["scope"]["operation_refs"]
    assert case.ref("batch", "B-LATE") not in before["data"]["scope"]["batch_refs"]
    invalid_anchor = client.get(BASE + "/runs/" + run_ref + "/decision/batches", query_string=dict(query, selected_batch_ref=case.ref("batch", "B-LATE")))
    assert invalid_anchor.status_code == 404 and invalid_anchor.get_json()["error"]["code"] == "entity_not_found"


def test_absent_official_reference_preserves_selected_facts_and_null_differences(candidate_case):
    run_ref, refs = compute(candidate_case)
    client, _ = api(candidate_case)
    query = {"comparison_mode": "admission_official", "selected_candidate_ref": refs[0]}
    overview = decision(client, run_ref, **query)["data"]
    _assert_summary_matches_full(decision_summary(client, run_ref, **query)["data"], overview)
    candidate = next(row for row in overview["candidates"] if row["candidate_ref"] == refs[0])
    changes = {key: candidate["metrics"][key] for key in ("changed_operation_count", "machine_change_count")}
    for section in ("batches", "operations"):
        data = detail(client, run_ref, section, **query)["data"]
        assert data["scope"]["reference_state"] == data["baseline"]["state"] == "absent"
        assert data["baseline"]["baseline_ref"] is None
        assert data["summary"]["operation_changes"] == changes
        assert all(value["state"] == "unavailable" and value["value"] is value["known_subtotal"] is None
                   and value["reason"]["code"] == "reference_required" for value in changes.values())
        assert len(data["items"]) == 1
        row = data["items"][0]
        assert row["drill"]["reference"] is None and row["drill"]["selected"]["candidate_ref"] == refs[0]
        assert row["change_kinds"] == ["unknown"]
        if section == "batches":
            assert row["after"]["schedule_complete"] is True
            assert row["before"]["completion_state"] == row["before"]["risk"] == "unknown"
            assert "no_admission_baseline" in row["before"]["issues"]
            assert row["finish_delta_hours"] is row["delay_delta_hours"] is None
        else:
            assert row["after"]["end"] is not None and row["before"] is None
            assert row["start_delta_hours"] is row["end_delta_hours"] is None


def test_official_reference_uses_only_the_current_admission_scope(candidate_case):
    case = candidate_case
    case.conn.execute("INSERT INTO Machines(machine_id,name,op_type_id) VALUES ('M2','Other lathe','T1')")
    case.conn.execute("INSERT INTO Operators(operator_id,name) VALUES ('O2','Other operator')")
    case.conn.execute("INSERT INTO OperatorMachine(operator_id,machine_id) VALUES ('O2','M2')")
    case.batch("B2", quantity=1)
    second_id = case.operation("B2", unit_hours=6, machine_id="M2", operator_id="O2")
    case.conn.commit()
    _, first_refs = compute(case, case.settings("B1", "B2"))
    adopter = service(case.conn)
    preview = adopter.preview(first_refs[0])
    assert preview["validation"]["can_adopt"] is True, preview
    adopted = adopter.adopt(first_refs[0], preview["write_context"]["write_token"], "impact-scope-first-adopt", INTENT)
    official = adopted["data"]["official_plan"]
    saved = list(case.conn.execute(
        "SELECT op_id,machine_id,operator_id,start_time,end_time FROM Schedule WHERE version=? ORDER BY op_id",
        (official["version"],),
    ))
    assert len(saved) == 2
    other = next(row for row in saved if row["op_id"] == second_id)
    assert (other["machine_id"], other["operator_id"]) == ("M2", "O2")
    assert (datetime.fromisoformat(other["end_time"]) - datetime.fromisoformat(other["start_time"])).total_seconds() == 6 * 3600
    accepted = case.accept(key="impact-scope-b1-only-run", settings=case.settings("B1"))
    result = WorkbenchRunWorker(case.conn).execute(accepted["run_ref"])
    run_ref, selected = result["run_ref"], result["candidates"][0]["candidate_ref"]
    client, _ = api(case)
    query = {"comparison_mode": "admission_official", "selected_candidate_ref": selected}
    overview = decision(client, run_ref, **query)["data"]
    _assert_summary_matches_full(decision_summary(client, run_ref, **query)["data"], overview)
    expected_refs = {operation_ref(case, case.op_id)}
    assert set(overview["scope"]["operation_refs"]) == expected_refs
    assert overview["scope"]["batch_refs"] == [case.ref("batch", "B1")]
    assert overview["baseline"]["metrics"]["makespan_hours"]["value"] == 0.75
    candidate = next(row for row in overview["candidates"] if row["candidate_ref"] == selected)
    assert candidate["deltas"]["makespan_hours"]["value"] == 0
    assert candidate["deltas"]["makespan_hours"]["direction"] == "equal"
    top = {item["key"]: item for item in overview["difference_summary"]["items"]}
    assert list(top) == ["overdue_count", "total_tardiness_hours", "makespan_hours"]
    for key, item in top.items():
        assert {field: item[field] for field in ("value", "state", "direction", "reference_value", "selected_value")} == candidate["deltas"][key]
        assert item["value"] == 0 and item["direction"] == "equal"
        assert item["selected_coverage"]["unknown_count"] == item["reference_coverage"]["unknown_count"] == 0
    workspace = detail(client, run_ref, "reference-workspace", **query)["data"]
    assert workspace["task_count"] == workspace["candidate_task_count"] == 1
    assert {row["operation_ref"] for row in workspace["tasks"]} == expected_refs
    assert {row["batch_ref"] for row in workspace["tasks"]} == {case.ref("batch", "B1")}
    assert all(row["machine"]["ref"] == case.ref("machine", "M1") for row in workspace["tasks"])
    operations = detail(client, run_ref, "operations", **query)["data"]
    assert {row["operation_ref"] for row in operations["items"] if row["before"] is not None} == expected_refs
    assert {row["ref"] for row in operations["filter_options"]["resources"]} == {
        case.ref("machine", "M1"), case.ref("operator", "O1"),
    }
    # A changed source receipt forces the real fallback to read the full original
    # official arrangement, while the metric contract keeps the admitted scope.
    source_run = case.conn.execute(
        "SELECT run_ref FROM WorkbenchRunCandidates WHERE candidate_ref=?", (first_refs[0],),
    ).fetchone()[0]
    corrupt_update(case.conn, "WorkbenchRunReceipts",
                   "UPDATE WorkbenchRunReceipts SET receipt_ref=? WHERE run_ref=?", ("f" * 48, source_run))
    source = DecisionSource(CandidateStore(case.conn), run_ref, comparison_mode="admission_official",
                            selected_candidate_ref=selected)
    assert source.official_changed is True
    assert len(source.analysis_tasks("reference")) == 2
    fallback = decision(client, run_ref, **query)["data"]
    assert fallback["baseline"]["metrics"] == overview["baseline"]["metrics"]
    assert fallback["baseline"]["coverage"] == overview["baseline"]["coverage"]
    # An out-of-scope interval with unknown comparability must not invalidate the
    # known in-scope interval, nor enlarge its span to the other batch's 50 hours.
    tasks = [dict(row) for row in source.analysis_tasks("reference")]
    outside = next(row for row in tasks if row["operation_ref"] not in expected_refs)
    outside.update(end=(datetime.fromisoformat(outside["start"]) + timedelta(hours=50)).isoformat(),
                   interval_comparable=False)
    values, coverage = project_official_core_metrics(
        tasks, expected_refs, source.reference_deferred_refs, True, source.delivery("reference"))
    assert values["makespan_hours"] == overview["baseline"]["metrics"]["makespan_hours"]
    assert values["makespan_hours"]["value"] == 0.75
    assert coverage == {key: overview["baseline"]["coverage"][key] for key in ("schedulable", "delivery")}


def test_detail_snapshot_downloads_all_filtered_pages_and_preserves_cell_types(candidate_case):
    run_ref, refs, _ = many_batches(candidate_case)
    client, _ = api(candidate_case)
    query = {"selected_candidate_ref": refs[0], "reference_candidate_ref": refs[0],
             "change_kind": "unchanged", "sort": "batch_id", "order": "asc",
             "delivery_filter": "all", "selected_batch_ref": candidate_case.ref("batch", "B27")}
    first = detail(client, run_ref, "batches", **query)
    token = first["meta"]["snapshot_ref"]
    second = detail(client, run_ref, "batches", **dict(query, page=2, snapshot_ref=token))
    rows = first["data"]["items"] + second["data"]["items"]
    assert len(rows) == 25 and first["data"]["page"]["has_more"] is True
    assert first["data"]["selected_batch"]["batch_id"] == "B27"
    assert all(row["batch_id"] != "B27" for row in rows)
    assert second["meta"]["snapshot_ref"] == token
    path = BASE + "/runs/" + run_ref + "/decision/export"
    downloads = {}
    for fmt in ("csv", "xlsx"):
        response = client.get(path, query_string=dict(query, snapshot_ref=token, section="batches", format=fmt))
        assert response.status_code == 200, response.get_json()
        assert response.headers["X-Workbench-Snapshot-Ref"] == token
        assert int(response.headers["X-Workbench-Row-Count"]) == len(rows)
        if fmt == "csv":
            downloads[fmt] = list(csv.DictReader(io.StringIO(response.data.decode("utf-8-sig"))))
        else:
            book = load_workbook(io.BytesIO(response.data), read_only=True, data_only=False)
            try:
                iterator = book["批次影响"].iter_rows()
                headers = [cell.value for cell in next(iterator)]
                downloads[fmt] = [dict(zip(headers, cells)) for cells in iterator]
            finally:
                book.close()
    for fmt, exported in downloads.items():
        values = (lambda item: {key: cell.value for key, cell in item.items()}) if fmt == "xlsx" else (lambda item: item)
        actual = [values(item) for item in exported]
        assert [item["批次永久编号"] for item in actual] == [item["batch_ref"] for item in rows]
        assert actual[0]["批次编号"] == ("'000001" if fmt == "csv" else "000001")
        assert all(item["零件名称"] == r"零件'甲\N" for item in actual)
        assert all(item["已确认变化工序数"] == ("0" if fmt == "csv" else 0)
                   and item["变化待确认工序数"] == ("0" if fmt == "csv" else 0)
                   and item["工序变化覆盖状态"] == "完整" for item in actual)
        missing = next(item for item in actual if item["批次编号"] == "B25")
        assert missing["所选拖期（小时）"] == ("" if fmt == "csv" else None)
        if fmt == "xlsx":
            assert exported[0]["批次编号"].data_type == "s"
            assert exported[0]["数量"].data_type == "n" and exported[0]["数量"].value == 3
    operation_query = {"selected_candidate_ref": refs[0], "reference_candidate_ref": refs[0],
                       "resource_ref": candidate_case.ref("machine", "M1"), "sort": "sequence", "order": "asc"}
    first_operations = detail(client, run_ref, "operations", **operation_query)
    operation_token = first_operations["meta"]["snapshot_ref"]
    next_operations = detail(client, run_ref, "operations", **dict(operation_query, page=2, snapshot_ref=operation_token))
    operation_rows = first_operations["data"]["items"] + next_operations["data"]["items"]
    assert len(operation_rows) == 27 and first_operations["data"]["page"]["has_more"] is True
    for fmt in ("csv", "xlsx"):
        response = client.get(path, query_string=dict(operation_query, snapshot_ref=operation_token, section="operations", format=fmt))
        assert response.status_code == 200, response.get_json()
        assert int(response.headers["X-Workbench-Row-Count"]) == len(operation_rows)
        if fmt == "csv":
            exported = list(csv.DictReader(io.StringIO(response.data.decode("utf-8-sig"))))
        else:
            book = load_workbook(io.BytesIO(response.data), read_only=True, data_only=True)
            try:
                iterator = book["工序变化"].iter_rows(values_only=True)
                headers = next(iterator)
                exported = [dict(zip(headers, row)) for row in iterator]
            finally:
                book.close()
        assert [row["工序永久编号"] for row in exported] == [row["operation_ref"] for row in operation_rows]
        assert all(row["基准排产状态"] == row["所选排产状态"] == "已安排"
                   and row["排产状态变化"] == "否"
                   and row["已确认变化字段"] == ("" if fmt == "csv" else None) for row in exported)
        assert [row["所选结束"] for row in exported] == [row["after"]["end"] for row in operation_rows]
        assert all(row["零件名称"] == r"零件'甲\N" for row in exported)
        assert all(row["单件编号"] == ("" if fmt == "csv" else None) for row in exported)
    all_book = client.get(path, query_string=dict(query, snapshot_ref=token, section="all", filter_section="batches", format="xlsx"))
    assert all_book.status_code == 200, all_book.get_json()
    book = load_workbook(io.BytesIO(all_book.data), read_only=True, data_only=True)
    try:
        assert book.sheetnames == ["方案总览", "批次影响", "工序变化", "资源影响", "覆盖工序", "准备需求", "工种切换",
                                  "推荐说明", "工序安排", "剩余需求", "共同计算", "目标预算", "进步事件", "搜索动作", "搜索阶段",
                                  "评价范围", "范围成员", "真实桥接", "指标计算方式", "比较范围"]
        assert len(list(book["批次影响"].iter_rows(values_only=True))) - 1 == 25
        assert len(list(book["工序变化"].iter_rows(values_only=True))) - 1 == 27
    finally:
        book.close()
    stale = client.get(path, query_string=dict(query, snapshot_ref=token, section="batches", format="csv", change_kind="unknown"))
    assert stale.status_code == 409 and stale.get_json()["error"]["code"] == "snapshot_stale"
