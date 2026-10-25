"""Real report protection keeps complete public conflict identities through HTTP and files."""

from io import BytesIO

from openpyxl import load_workbook

from core.services.workbench.execution.field_report_files import FieldReportFileService
from core.services.workbench.execution.field_report_files_codec import encode_issues
from core.services.workbench.execution.production_report_prepare import ReportBatchRejected
from tests.workbench.execution_ledger_support import all_rows
from tests.workbench.field_workspace_support import BASE, _ledger_fixture, success  # noqa: F401
from tests.workbench.field_workspace_support import field_api as _field_api  # noqa: F401


def test_correction_preserves_both_downstream_official_arrangements_and_quantity_protection(field_api):
    api, case = field_api, field_api.case
    downstream = [case.op('OP2', seq=2), case.op('OP3', seq=3)]
    case.plan(2, [case.op_id] + downstream, start='2026-09-09T11:00:00', end='2026-09-09T12:00:00')
    api.create(case.values(10))
    report = api.task()['execution']['reports'][0]
    before = all_rows(case.conn)
    response = api.client.post(BASE + '/reports/' + report['report_ref'] + '/correct', json=api.body(
        report['write_context'], {'original_revision_ref': report['revision_ref'], 'reason': '核对数量',
                                 'completed_quantity': 8, 'effective_processing_hours': 1.2}))
    assert response.status_code == 409
    body = response.get_json()
    assert body['ok'] is False and body['committed'] is False
    assert body['error']['code'] == 'constraint_conflict' and body['error']['row_number'] == 1
    conflicts = body['error']['conflicts']
    assert len(conflicts) == 2
    assert {item['code'] for item in conflicts} == {'downstream_requires_completion'}
    downstream_tasks = [case.task(2, op) for op in downstream]
    headers = case.ledger.repo.task_headers(downstream_tasks)
    assert {item['operation_ref'] for item in conflicts} == {headers[task]['operation_ref'] for task in downstream_tasks}
    assert {item['operation_label'] for item in conflicts} == {'2 Turning', '3 Turning'}
    for item in conflicts:
        op = next(op for op in downstream if case.task(2, op) == item['plan']['task_ref'])
        assert item['batch_id'] == 'B1' and item['operation_code'] in ('OP2', 'OP3')
        assert item['message'] == '后道工序的执行或正式安排依赖本工序完工'
        assert item['piece_id'] is None
        assert item['plan'] == {'plan_ref': case.plan_ref(2), 'task_ref': case.task(2, op), 'version': 2,
                               'start': '2026-09-09T11:00:00', 'end': '2026-09-09T12:00:00'}
    assert all_rows(case.conn) == before
    assert api.task()['execution']['known_completed_quantity'] == 10
    preview = success(api.client.post(BASE + '/reports/' + report['report_ref'] + '/void-preview', json={
        'input': {'original_revision_ref': report['revision_ref'], 'reason': '核对撤销', 'declared_operator': ''}}))['data']
    assert not preview['can_confirm']
    assert sorted(preview['downstream_impacts'], key=lambda item: item['operation_ref']) == sorted(conflicts, key=lambda item: item['operation_ref'])
    assert all_rows(case.conn) == before

    # The file-preview boundary keeps the same facts and the original Excel row mapping.
    rows = [{'row': 4, 'result': 'pending', 'errors': []}]
    FieldReportFileService._preview_error(ReportBatchRejected('constraint_conflict', body['error']['message'],
        row_number=1, conflicts=conflicts), rows, [4])
    assert rows[0]['errors'][0]['conflicts'] == conflicts
    book = load_workbook(BytesIO(encode_issues(rows)), read_only=True)
    values = list(book.worksheets[0].iter_rows(values_only=True))
    assert values[0][:2] == ('Excel 行号', '问题')
    assert len(values) == 4 and {row[0] for row in values[1:]} == {4}
    assert {row[8] for row in values[2:]} == {item['operation_ref'] for item in conflicts}
    assert {row[9] for row in values[2:]} == {case.task(2, op) for op in downstream}
    assert all(row[5:8] == (2, '2026-09-09T11:00:00', '2026-09-09T12:00:00') for row in values[2:])
    book.close()
