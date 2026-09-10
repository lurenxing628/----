"""AN-only readback oracle. Never opens any database except an identified fixture."""

import csv
import hashlib
import io
import json
import sqlite3
import sys
from pathlib import Path

from live_environment import read_identity


def snapshot(root):
    root = Path(root).resolve()
    read_identity(root)
    db = root / "db/aps-live.db"
    with sqlite3.connect(db.as_uri() + "?mode=ro", uri=True) as conn:
        conn.row_factory = sqlite3.Row
        conn.execute("PRAGMA query_only=ON")
        tables = [row[0] for row in conn.execute("SELECT name FROM sqlite_master WHERE type='table' ORDER BY name")]
        schema = [list(row) for row in conn.execute("SELECT type,name,sql FROM sqlite_master ORDER BY type,name")]
        return {"schema": schema, "version": conn.execute("SELECT version FROM SchemaVersion WHERE id=1").fetchone()[0],
                "tables": {name: [dict(row) for row in conn.execute('SELECT rowid AS __oracle_rowid__,* FROM "' + name + '" ORDER BY rowid')]
                           for name in tables}, "integrity": conn.execute("PRAGMA integrity_check").fetchone()[0],
                "foreign_keys": [list(row) for row in conn.execute("PRAGMA foreign_key_check")]}


def download(file):
    file = Path(file)
    data = file.read_bytes()
    if file.suffix == ".csv":
        rows = list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))
        sheets = {"csv": rows}
    else:
        from openpyxl import load_workbook
        with file.open("rb") as stream:
            workbook = load_workbook(stream, read_only=True, data_only=False)
            sheets = {sheet.title: list(sheet.iter_rows(values_only=True)) for sheet in workbook}
            workbook.close()
    return {"file": str(file), "bytes": len(data), "sha256": hashlib.sha256(data).hexdigest(), "sheets": sheets}


def preservation(before, after):
    assert before["schema"] == after["schema"], "Schema changed during the UI run"
    assert after["integrity"] == "ok" and after["foreign_keys"] == []
    business = {"Parts": "part_no", "PartOperations": "part_no", "Batches": "batch_id",
                "BatchOperations": "batch_id", "OpTypes": "op_type_id"}
    metadata = {"WorkbenchEntityRefs", "WorkbenchCommandReceipts", "WorkbenchProcessWorkflow",
                "WorkbenchProcessOperationConfirmations", "WorkbenchPlanIdentityClock", "WorkbenchPlanSourceRefs", "sqlite_sequence"}
    changed = {}
    for table, rows in before["tables"].items():
        current = after["tables"][table]
        if rows != current:
            changed[table] = {"before_rows": len(rows), "after_rows": len(current)}
        if table in metadata:
            continue
        if table not in business:
            assert rows == current, "Unrelated table changed: " + table
            continue
        old = {row["__oracle_rowid__"]: row for row in rows}
        new = {row["__oracle_rowid__"]: row for row in current}
        for key, row in old.items():
            assert key in new, "Original row removed: " + table
            allowed = {"unit_hours"} if table == "PartOperations" and row["part_no"] == "PROC-001" and row["seq"] == 10 else set()
            assert {k: v for k, v in row.items() if k not in allowed} == {k: v for k, v in new[key].items() if k not in allowed}, (table, key)
        for key, row in new.items():
            if key not in old:
                assert row[business[table]].startswith("AN-"), (table, key)
    return {"passed": True, "tables_checked": len(before["tables"]), "changed_tables": changed,
            "all_original_rows_retained": True, "original_allowed_columns": {"PartOperations PROC-001 seq=10": ["unit_hours"]}}


if __name__ == "__main__":
    if sys.argv[1] == "preservation":
        root = Path(sys.argv[2]).resolve()
        read_identity(root)
        result = preservation(json.loads((root / "an-full-before.json").read_text()), json.loads((root / "an-full-after.json").read_text()))
    else:
        result = snapshot(sys.argv[2]) if sys.argv[1] == "snapshot" else download(sys.argv[2])
    print(json.dumps(result, ensure_ascii=False, default=str))
