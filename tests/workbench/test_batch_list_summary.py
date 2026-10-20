"""Batch summaries keep whole-scope decisions and defer operation DTO expansion."""

from io import BytesIO

from openpyxl import load_workbook

from core.models.workbench_batch_query import batch_scope
from core.models.workbench_execution import ExecutionProjection
from core.services.workbench.batch.projection import BatchProjection
from core.services.workbench.batch.queries import WorkbenchBatchQueryService
from tests.workbench.batch_support import BASE, batch_database
from tests.workbench.identity_metadata_support import insert_row

_batch_fixture = batch_database


def page_batches(conn):
    for code, quantity, due, ready in (
        ("PAGE-A", 3, None, "yes"),
        ("PAGE-B", 2, "2026-10-15", "yes"),
        ("PAGE-C", 2, "2026-10-16", "no"),
        ("PAGE-D", 7, "2026-10-14", "partial"),
    ):
        insert_row(conn, "Batches", {"batch_id": code, "part_no": "P1", "part_name": "retained name",
            "quantity": quantity, "due_date": due, "priority": "normal", "ready_status": ready, "status": "pending"})
        if code != "PAGE-A":
            insert_row(conn, "BatchOperations", {"op_code": code + "_01", "batch_id": code, "seq": 1,
                "op_type_id": "OT1", "op_type_name": "turning", "source": "internal", "setup_hours": 0,
                "unit_hours": 0, "machine_id": None if code == "PAGE-C" else "M1",
                "operator_id": None if code == "PAGE-C" else "O1", "status": "completed" if code == "PAGE-D" else "pending"})
    conn.commit()


def test_summary_keeps_complete_entity_values_and_whole_scope_pagination(batch_client, monkeypatch):
    conn = batch_client.batch_conn
    page_batches(conn)
    reader = WorkbenchBatchQueryService(conn)
    scope = batch_scope({"query": "PAGE-", "sort": "quantity", "size": 1})
    with reader.detached_read_snapshot():
        facts = reader.load()
        projection = BatchProjection(facts)
        full = {row["batch_id"]: projection.entity(row) for row in facts["Batches"] if row["batch_id"].startswith("PAGE-")}
        expanded, serialize = [], ExecutionProjection.to_dict

        def capture(row):
            expanded.append(row.operation_ref)
            return serialize(row)

        monkeypatch.setattr(ExecutionProjection, "to_dict", capture)
        order = ["PAGE-B", "PAGE-C", "PAGE-A", "PAGE-D"]
        for number, code in enumerate(order, 1):
            data = reader.page(dict(scope, page=number))
            assert data["page"]["total"] == 4 and data["page"]["pages"] == 4
            assert data["metrics"] == {"scope": "filtered", "total": 4, "completed": 1}
            assert data["entities"] == [{key: value for key, value in full[code].items() if key != "operations"}]
        assert [row["business_code"] for row in reader.matched(dict(scope, focus="gaps"))] == ["PAGE-C", "PAGE-A", "PAGE-D"]
        assert [row["business_code"] for row in reader.matched(dict(scope, focus="unready"))] == ["PAGE-C", "PAGE-D"]
        assert [row["business_code"] for row in reader.matched(dict(scope, status="completed"))] == ["PAGE-D"]
        assert [row["business_code"] for row in reader.matched(dict(scope, sort="due_date", direction="desc"))] == ["PAGE-A", "PAGE-C", "PAGE-B", "PAGE-D"]
        selection = reader.selection(scope)
        assert selection == {"refs": [full[code]["ref"] for code in order], "count": 4}
        assert expanded == []
        actual = reader.detail(full["PAGE-B"]["ref"])
        assert {key: value for key, value in actual.items() if key not in ("materials", "template")} == full["PAGE-B"]
        assert expanded == [full["PAGE-B"]["operations"][0]["ref"]]


def test_filtered_selection_facets_and_export_use_summaries(batch_client, monkeypatch):
    page_batches(batch_client.batch_conn)

    def no_details(_row):
        raise AssertionError("A batch list, selection, facet or basic export must not expand operation DTOs.")

    monkeypatch.setattr(ExecutionProjection, "to_dict", no_details)
    scope = batch_scope({"query": "PAGE-", "size": 1, "column_filters": {"quantity": [2]}})
    response = batch_client.post(BASE + "/query", json=scope)
    assert response.status_code == 200, response.get_json()
    listed = response.get_json()
    assert listed["data"]["page"]["total"] == 2
    assert listed["data"]["entities"][0]["business_code"] == "PAGE-B"
    assert "operations" not in listed["data"]["entities"][0]
    scope["snapshot_ref"] = listed["meta"]["snapshot_ref"]
    selected = batch_client.post(BASE + "/selection", json=scope).get_json()
    assert selected["data"]["count"] == 2
    facet = batch_client.post(BASE + "/facets", json={"scope": scope, "field": "quantity"}).get_json()
    assert facet["data"]["values"] == [2, 3, 7]
    preview = batch_client.post(BASE + "/export-preview", json={"scope": scope, "selection": "filtered"}).get_json()
    assert preview["data"]["count"] == 2
    exported = batch_client.get(BASE + "/export", query_string={"export_ref": preview["data"]["export_ref"]})
    assert exported.status_code == 200 and exported.headers["X-Workbench-Row-Count"] == "2"
    workbook = load_workbook(BytesIO(exported.data), read_only=True, data_only=True)
    try:
        rows = list(workbook["批次信息"].iter_rows(values_only=True))
        assert [row[0] for row in rows[1:]] == ["PAGE-B", "PAGE-C"]
        assert [row[2] for row in rows[1:]] == [2, 2]
    finally:
        workbook.close()
