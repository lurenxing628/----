"""Process transport contracts: sparse values, exact IDs, diagnostics and scale."""

import ast
import copy
import sqlite3
from datetime import datetime
from io import BytesIO
from pathlib import Path
from zipfile import ZIP_STORED, ZipFile

import openpyxl
import pytest

from core.errors import ValidationError
from core.models.workbench_process_file import COLUMNS, INT64_MAX, LABELS, ProcessFileDownload
from core.services.workbench import process_file_reader, process_file_writer
from core.services.workbench.process_file_codec import (
    INSTRUCTIONS,
    TEMPLATE_VERSION,
    decode_process_file,
    encode_process_file,
    file_columns,
    public_columns,
)
from tests.workbench.process_file_codec_support import (
    NS,
    SHEET,
    SinglePassRows,
    file_bytes,
    measure,
    mutate_xml,
    number_cell,
    source_cells,
)


@pytest.mark.parametrize("fmt", ("csv", "xlsx"))
@pytest.mark.parametrize("kind", ("route", "hours"))
def test_download_shape_chinese_headers_and_header_only_template(fmt, kind):
    result = encode_process_file(kind=kind, rows=[], fmt=fmt)
    assert isinstance(result, ProcessFileDownload)
    assert type(result.content) is bytes and result.row_count == 0
    assert result.filename == ("零件工艺路线" if kind == "route" else "零件工序工时") + "." + fmt
    assert result.mime_type == ("text/csv; charset=utf-8" if fmt == "csv"
                                else "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert tuple(next(source_cells(result.content, fmt))) == tuple(LABELS[field] for field in COLUMNS[kind])
    assert decode_process_file(kind=kind, content=result.content, fmt=fmt) == []
    assert TEMPLATE_VERSION == 1 and "\\N" in INSTRUCTIONS
    assert tuple(item["key"] for item in public_columns(kind)) == file_columns(kind)


@pytest.mark.parametrize("fmt", ("csv", "xlsx"))
@pytest.mark.parametrize("text", ("名称α🙂", " 前后空格 ", "  ", "甲\n乙\r\n丙\r丁", "=1+1", "+SUM(1,2)",
                                 "-2+1", "@SUM(A1)", "\t=1", "'原文", "''原文", "'", r"\N", r"\\N", r"\path"))
def test_text_roundtrip_never_strips_newlines_spaces_or_formula_prefixes(fmt, text):
    row = {"business_code": "000001", "label": text, "route_raw": text, "remark": text}
    result = encode_process_file("route", [row], fmt)
    assert decode_process_file("route", result.content, fmt) == [{"row": 2, "values": row, "errors": []}]
    if fmt == "xlsx":
        wb = openpyxl.load_workbook(BytesIO(result.content), read_only=True, data_only=False)
        try:
            assert all(cell.data_type == "s" for cell in next(wb.active.iter_rows(min_row=2)))
        finally:
            wb.close()
    else:
        assert result.content.startswith(b"\xef\xbb\xbf")
        assert all(cell.startswith("'") for cell in list(source_cells(result.content, fmt))[1])


@pytest.mark.parametrize("fmt", ("csv", "xlsx"))
def test_sparse_null_zero_and_empty_string_are_distinct(fmt):
    rows = [{"business_code": "甲", "sequence": 10, "setup_hours": 0, "unit_hours": None},
            {"business_code": "乙", "sequence": 20, "setup_hours": None, "unit_hours": 0}]
    result = decode_process_file("hours", encode_process_file("hours", rows, fmt).content, fmt)
    assert [row["values"] for row in result] == rows
    assert not any(row["errors"] for row in result)
    route = decode_process_file("route", file_bytes(["图号", "名称", "工艺路线字符串", "备注"],
                                                    [["甲", "", r"\N", "  "]], fmt), fmt)[0]
    assert route == {"row": 2, "values": {"business_code": "甲", "route_raw": None, "remark": "  "}, "errors": []}
    blank = encode_process_file("route", [{"business_code": "甲", "label": ""}], fmt)
    assert decode_process_file("route", blank.content, fmt)[0]["values"] == {"business_code": "甲"}


@pytest.mark.parametrize("fmt", ("csv", "xlsx"))
def test_old_chinese_templates_and_canonical_aliases(fmt):
    route = file_bytes(["图号", "名称", "工艺路线字符串"], [["0001", "零件", "10铣"]], fmt)
    assert decode_process_file("route", route, fmt)[0]["values"]["route_raw"] == "10铣"
    for headers in (("图号", "工序", "换型时间(h)", "单件工时(h)"),
                    ("business_code", "sequence", "setup_hours", "unit_hours")):
        row = decode_process_file("hours", file_bytes(headers, [["0001", 10, "", 0]], fmt), fmt)[0]
        assert row == {"row": 2, "values": {"business_code": "0001", "sequence": 10, "unit_hours": 0}, "errors": []}


@pytest.mark.parametrize("fmt", ("csv", "xlsx"))
@pytest.mark.parametrize("source,expected", (("自制", "internal"), ("外协", "external"),
                                           ("internal", "internal"), ("external", "external")))
def test_hours_source_assertions_and_group_columns_only_parse(fmt, source, expected):
    fields = COLUMNS["hours"]
    values = ["00001", 10, " 原工种\n甲 ", source, 0, 0.125, None, 10, 20, 2.5]
    row = decode_process_file("hours", file_bytes(fields, [values], fmt), fmt)[0]
    assert not row["errors"] and row["values"]["source"] == expected
    assert row["values"]["op_type_name"] == " 原工种\n甲 "
    assert row["values"]["group_start"] == 10 and row["values"]["group_end"] == 20
    assert "external_days" not in row["values"]
    assert row["values"]["group_total_days"] == 2.5
    back = decode_process_file("hours", encode_process_file("hours", [row["values"]], fmt).content, fmt)[0]
    assert back["values"] == row["values"] and not back["errors"]


@pytest.mark.parametrize("fmt", ("csv", "xlsx"))
@pytest.mark.parametrize("value", (1, 2 ** 53 - 1, 2 ** 53, 2 ** 53 + 1, INT64_MAX))
@pytest.mark.parametrize("field", ("sequence", "group_start", "group_end"))
def test_exact_int64_text_and_export_roundtrip(fmt, value, field):
    original = {"business_code": "001", "sequence": 1, field: value}
    encoded = encode_process_file("hours", [original], fmt)
    actual = decode_process_file("hours", encoded.content, fmt)[0]
    assert not actual["errors"] and actual["values"] == original
    assert type(actual["values"][field]) is int
    text = file_bytes(["图号", "工序", "外协组起序"], [["001", str(value), "00000001"]], fmt)
    parsed = decode_process_file("hours", text, fmt)[0]
    assert parsed["values"]["sequence"] == value and parsed["values"]["group_start"] == 1


@pytest.mark.parametrize("value,valid", ((str(2 ** 53), True), (str(2 ** 53 + 1), False),
                                        (str(INT64_MAX), False), ("1.0", True), ("1.25", False), ("1e309", False)))
def test_xlsx_numbers_above_safe_integer_range_are_not_guessed(value, valid):
    content = file_bytes(["图号", "工序"], [["001", "1"]], "xlsx")
    parsed = decode_process_file("hours", number_cell(content, "B2", value), "xlsx")[0]
    assert bool(parsed["errors"]) is not valid
    if not valid:
        assert "sequence" not in parsed["values"]


@pytest.mark.parametrize("fmt", ("csv", "xlsx"))
@pytest.mark.parametrize("value", ("0", "-1", "1.2", "1.0", "1e2", "1,000", "1序", " 1", "١", "True", str(INT64_MAX + 1)))
def test_integer_text_rejects_noncanonical_and_out_of_range(fmt, value):
    row = decode_process_file("hours", file_bytes(["图号", "工序"], [["甲", value]], fmt), fmt)[0]
    assert row["errors"][0]["field"] == "sequence" and "sequence" not in row["values"]


@pytest.mark.parametrize("fmt", ("csv", "xlsx"))
@pytest.mark.parametrize("value", ("NaN", "Inf", "-Infinity", "1e309", "True", "1,000", "1h", "=1+1", " 1"))
def test_hours_reject_nonfinite_boolean_unit_and_formula_text(fmt, value):
    row = decode_process_file("hours", file_bytes(["图号", "工序", "换型时间(h)"], [["甲", 1, value]], fmt), fmt)[0]
    assert row["errors"][0]["field"] == "setup_hours" and "setup_hours" not in row["values"]


@pytest.mark.parametrize("fmt", ("csv", "xlsx"))
@pytest.mark.parametrize("field,value", (("setup_hours", -1), ("unit_hours", -0.5),
                                        ("external_days", 0), ("group_total_days", -1)))
def test_sign_diagnostics_preserve_explicit_number_without_defaults(fmt, field, value):
    row = decode_process_file("hours", file_bytes(["图号", "工序", field], [["甲", 1, value]], fmt), fmt)[0]
    assert row["values"][field] == value and row["errors"][0]["field"] == field


@pytest.mark.parametrize("value", (True, False, 123, datetime(2026, 9, 9), "=1+1", "#DIV/0!"))
def test_xlsx_text_never_guesses_type_and_formula_or_error_cells_are_diagnosed(value):
    row = decode_process_file("route", file_bytes(["图号", "名称"], [["甲", value]], "xlsx"), "xlsx")[0]
    assert row["errors"][0]["field"] == "label"
    assert "label" not in row["values"]


@pytest.mark.parametrize("field", ("sequence", "setup_hours", "source"))
def test_xlsx_boolean_cells_are_not_numbers_or_sources(field):
    values = {"business_code": "甲", "sequence": 1, field: True}
    row = decode_process_file("hours", file_bytes(list(values), [list(values.values())], "xlsx"), "xlsx")[0]
    assert row["errors"][0]["field"] == field


@pytest.mark.parametrize("fmt", ("csv", "xlsx"))
@pytest.mark.parametrize("kind,headers", (("route", ["名称"]), ("hours", ["图号"]),
                                         ("route", ["图号", "business_code"]), ("hours", ["图号", "工序", "sequence"]),
                                         ("route", ["图号", "???"]), ("hours", ["图号", "工序", ""])))
def test_headers_fail_closed(fmt, kind, headers):
    with pytest.raises(ValidationError) as caught:
        decode_process_file(kind, file_bytes(headers, [], fmt), fmt)
    assert caught.value.field == "headers" and caught.value.details["row"] == 1


@pytest.mark.parametrize("fmt", ("csv", "xlsx"))
def test_empty_records_duplicate_keys_and_extra_columns_are_never_discarded(fmt):
    content = file_bytes(["图号", "名称"], [["甲", "首行\n次行"], [None, None], ["甲", "末行"]], fmt)
    rows = decode_process_file("route", content, fmt)
    assert len(rows) == 3 and all(row["errors"] for row in rows)
    assert rows[0]["errors"][0]["code"] == rows[2]["errors"][0]["code"] == "duplicate_entry"
    assert [row["row"] for row in rows] == ([2, 4, 5] if fmt == "csv" else [2, 3, 4])
    content = file_bytes(["图号", "工序"], [["甲|乙", "01"], ["甲|乙", 1], ["甲", 1], ["甲", 2]], fmt)
    rows = decode_process_file("hours", content, fmt)
    assert [bool(row["errors"]) for row in rows] == [True, True, False, False]
    extra = decode_process_file("route", file_bytes(["图号"], [["甲", "=1"]], fmt), fmt)[0]
    assert extra["errors"][0]["field"] == "columns"


def test_csv_physical_blank_lines_and_missing_bom():
    rows = decode_process_file("route", "图号\n\n甲\n\n".encode(), "csv")
    assert [row["row"] for row in rows] == [2, 3, 4]
    assert [bool(row["errors"]) for row in rows] == [True, False, True]


@pytest.mark.parametrize("content,fmt", ((b"", "csv"), (b"\xff", "csv"), (b'"unclosed', "csv"),
                                        (b"bad zip", "xlsx"), (b"PK\x03\x04", "csv"), (b"x", "xls")))
def test_invalid_encoding_format_and_bytes_raise_file_error(content, fmt):
    with pytest.raises(ValidationError):
        decode_process_file("route", content, fmt)


@pytest.mark.parametrize("kind", (None, [], {}, True, "process", "ROUTE"))
def test_invalid_kind_is_validation_error(kind):
    with pytest.raises(ValidationError):
        decode_process_file(kind, b"", "csv")
    with pytest.raises(ValidationError):
        encode_process_file(kind, [], "csv")


@pytest.mark.parametrize("fmt", (None, [], {}, True, "CSV"))
def test_invalid_format_is_validation_error(fmt):
    with pytest.raises(ValidationError):
        decode_process_file("route", b"", fmt)
    with pytest.raises(ValidationError):
        encode_process_file("route", [], fmt)


def test_multisheet_hidden_sheet_and_fake_dimensions():
    wb = openpyxl.Workbook()
    wb.active.append(["图号"])
    extra = wb.create_sheet("另一张")
    extra.sheet_state = "hidden"
    buffer = BytesIO()
    wb.save(buffer)
    wb.close()
    with pytest.raises(ValidationError, match="一张"):
        decode_process_file("route", buffer.getvalue(), "xlsx")
    content = file_bytes(["图号"], [["甲"], ["乙"]], "xlsx")
    content = mutate_xml(content, lambda root: root.find("s:dimension", NS).set("ref", "A1:A1"))
    assert [row["values"]["business_code"] for row in decode_process_file("route", content, "xlsx")] == ["甲", "乙"]


def test_mislabeled_macro_package_and_duplicate_zip_members_are_rejected():
    content = file_bytes(["图号"], [["甲"]], "xlsx")
    def macro(root):
        for node in root:
            if node.attrib.get("PartName") == "/xl/workbook.xml":
                node.set("ContentType", "application/vnd.ms-excel.sheet.macroEnabled.main+xml")
    with pytest.raises(ValidationError, match="标准 XLSX"):
        decode_process_file("route", mutate_xml(content, macro, "[Content_Types].xml"), "xlsx")
    buffer = BytesIO(content)
    with ZipFile(buffer, "a") as archive, pytest.warns(UserWarning):
        archive.writestr(SHEET, archive.read(SHEET))
    with pytest.raises(ValidationError, match="重复"):
        decode_process_file("route", buffer.getvalue(), "xlsx")


def test_byte_and_expanded_limits_before_openpyxl(monkeypatch):
    assert process_file_reader.IMPORT_BYTE_LIMIT == 16 * 1024 ** 2
    assert process_file_reader.XLSX_EXPANDED_BYTE_LIMIT == 64 * 1024 ** 2
    content = file_bytes(["图号"], [["甲"]], "xlsx")
    monkeypatch.setattr(process_file_reader, "IMPORT_BYTE_LIMIT", len(content) - 1)
    with pytest.raises(ValidationError, match="16 MiB"):
        decode_process_file("route", content, "xlsx")
    monkeypatch.setattr(process_file_reader, "IMPORT_BYTE_LIMIT", len(content))
    assert not decode_process_file("route", content, "xlsx")[0]["errors"]
    monkeypatch.setattr(process_file_reader, "XLSX_EXPANDED_BYTE_LIMIT", 1024)
    with pytest.raises(ValidationError, match="64 MiB"):
        decode_process_file("route", content, "xlsx")


@pytest.mark.parametrize("fmt", ("csv", "xlsx"))
@pytest.mark.parametrize("kind", ("route", "hours"))
def test_2000_import_allowed_2010_rejected_without_partial_return(fmt, kind):
    download = encode_process_file(kind, SinglePassRows(kind, 2000), fmt)
    parsed = measure(kind + "_import_2000_" + fmt, lambda: decode_process_file(kind, download.content, fmt))
    assert len(parsed) == 2000 and not any(row["errors"] for row in parsed)
    download = encode_process_file(kind, SinglePassRows(kind, 2010), fmt)
    with pytest.raises(ValidationError, match="2000") as caught:
        decode_process_file(kind, download.content, fmt)
    assert caught.value.details["row"] == (4002 if kind == "route" and fmt == "csv" else 2002)


@pytest.mark.parametrize("fmt", ("csv", "xlsx"))
@pytest.mark.parametrize("kind", ("route", "hours"))
def test_complete_10000_export_is_one_pass_and_not_import_capped(fmt, kind):
    source = SinglePassRows(kind, 10000)
    download = measure(kind + "_export_10000_" + fmt, lambda: encode_process_file(kind, source, fmt))
    assert download.row_count == 10000 and source.iterations == 1
    records = source_cells(download.content, fmt)
    next(records)
    count, first, last = 0, None, None
    for record in records:
        if first is None:
            first = record
        expected = [f"图{count:05d}", "零件", "10铣20检", "首行\n末行"] if kind == "route" else [
            f"图{count:05d}", str(count + 1), "铣", "自制", "0.0", "0.12345678901234566", r"\N", r"\N", r"\N", r"\N"]
        assert list(record) == ["'" + value if fmt == "csv" else value for value in expected]
        count += 1
        last = record
    prefix = "'" if fmt == "csv" else ""
    assert count == 10000 and first[0] == prefix + "图00000" and last[0] == prefix + "图09999"


def test_export_abort_cleans_tempfiles_and_never_truncates(monkeypatch):
    from openpyxl.worksheet._writer import ALL_TEMP_FILES
    before = list(ALL_TEMP_FILES)
    for value in ("x" * 32768, "\x00", "\uffff"):
        with pytest.raises(ValidationError):
            encode_process_file("route", [{"business_code": "甲", "remark": value}], "xlsx")
        assert list(ALL_TEMP_FILES) == before
    monkeypatch.setattr(process_file_writer, "XLSX_MAX_ROWS", 2)
    with pytest.raises(ValidationError, match="容量"):
        encode_process_file("route", SinglePassRows("route", 2), "xlsx")
    assert list(ALL_TEMP_FILES) == before


@pytest.mark.parametrize("fmt", ("csv", "xlsx"))
def test_no_database_access_and_python38_syntax(monkeypatch, fmt):
    def forbidden(*args, **kwargs):
        raise AssertionError("database access")
    monkeypatch.setattr(sqlite3, "connect", forbidden)
    result = encode_process_file("route", [{"business_code": "甲"}], fmt)
    assert not decode_process_file("route", result.content, fmt)[0]["errors"]
    root = Path(__file__).resolve().parents[2]
    # Domain/file API coordinators share the prefix but are not pure codecs.
    codec_modules = {
        "core.models.workbench_process_file",
        "core.services.workbench.process_file_codec",
        "core.services.workbench.process_file_reader",
        "core.services.workbench.process_file_values",
        "core.services.workbench.process_file_writer",
        "core.services.workbench.process_file_xml",
    }
    pure_services = codec_modules | {
        "core.services.workbench.resource_file_codec",
        "core.services.workbench.resource_file_writer",
    }
    paths = [root / (module.replace(".", "/") + ".py") for module in sorted(codec_modules)]
    for path in paths:
        tree = ast.parse(path.read_text(encoding="utf-8"), feature_version=8)
        imports = {node.module or "" for node in ast.walk(tree) if isinstance(node, ast.ImportFrom)}
        imports.update(alias.name for node in ast.walk(tree) if isinstance(node, ast.Import) for alias in node.names)
        assert not any(name.split(".")[0] in ("data", "web", "flask", "sqlite3")
                       or name.startswith("core.infrastructure") for name in imports), path
        assert {name for name in imports if name.startswith("core.services")} <= pure_services, path


@pytest.mark.parametrize("mutation", ("duplicate_row", "reverse_rows", "duplicate_cell", "reverse_cells", "wrong_cell_row", "invalid_row", "far_column"))
def test_physical_worksheet_records_cannot_be_hidden_or_overwritten(mutation):
    content = file_bytes(["图号", "名称"], [["甲", "甲名"], ["乙", "乙名"]], "xlsx")
    def change(root):
        data = root.find("s:sheetData", NS)
        row = data[1]
        if mutation == "duplicate_row":
            data.insert(2, copy.deepcopy(row))
        elif mutation == "reverse_rows":
            data[:] = [data[0], data[2], data[1]]
        elif mutation == "duplicate_cell":
            row.append(copy.deepcopy(row[0]))
        elif mutation == "reverse_cells":
            row[:] = list(reversed(row))
        elif mutation == "wrong_cell_row":
            row[0].set("r", "A3")
        elif mutation == "invalid_row":
            row.set("r", "0")
        else:
            row[1].set("r", "XFE2")
    with pytest.raises(ValidationError):
        decode_process_file("route", mutate_xml(content, change), "xlsx")


def test_crc_error_in_package_is_not_silently_ignored():
    buffer = BytesIO(file_bytes(["图号"], [["甲"]], "xlsx"))
    with ZipFile(buffer, "a", ZIP_STORED) as archive:
        archive.writestr("integrity.bin", b"UNIQUE-INTEGRITY-MARKER")
    content = bytearray(buffer.getvalue())
    content[content.index(b"UNIQUE-INTEGRITY-MARKER")] = 0
    with pytest.raises(ValidationError, match="完整性"):
        decode_process_file("route", bytes(content), "xlsx")


@pytest.mark.parametrize("fmt", ("csv", "xlsx"))
@pytest.mark.parametrize("value", ("1e-999", "-1e-999", "1" + "0" * 400))
def test_underflow_and_overflow_cannot_turn_hours_into_zero_or_infinity(fmt, value):
    row = decode_process_file("hours", file_bytes(["图号", "工序", "单件工时(h)"], [["甲", 1, value]], fmt), fmt)[0]
    assert row["errors"][0]["field"] == "unit_hours" and "unit_hours" not in row["values"]


@pytest.mark.parametrize("fmt", ("csv", "xlsx"))
@pytest.mark.parametrize("value", ("INTERNAL", " internal", "外包", "", r"\N"))
def test_source_is_fixed_mapping_or_explicit_omission_null(fmt, value):
    row = decode_process_file("hours", file_bytes(["图号", "工序", "归属"], [["甲", 1, value]], fmt), fmt)[0]
    if value == "":
        assert "source" not in row["values"] and not row["errors"]
    elif value == r"\N":
        assert row["values"]["source"] is None and not row["errors"]
    else:
        assert "source" not in row["values"] and row["errors"][0]["field"] == "source"


@pytest.mark.parametrize("fmt", ("csv", "xlsx"))
@pytest.mark.parametrize("value", (0.1, 0.30000000000000004, 0.12345678901234567, 1.7976931348623157e308, 5e-324))
def test_finite_float_precision_roundtrip(fmt, value):
    row = {"business_code": "甲", "sequence": 1, "unit_hours": value}
    assert decode_process_file("hours", encode_process_file("hours", [row], fmt).content, fmt)[0] == {
        "row": 2, "values": row, "errors": []}


@pytest.mark.parametrize("fmt", ("csv", "xlsx"))
@pytest.mark.parametrize("row", ({"business_code": "甲", "unknown": 1}, {"business_code": 123},
                                 {"business_code": "甲", "sequence": 10 ** 400},
                                 {"business_code": "甲", "unit_hours": True},
                                 {"business_code": "甲", "unit_hours": float("nan")}))
def test_invalid_export_rows_raise_without_payload_or_leftover_tempfiles(fmt, row):
    from openpyxl.worksheet._writer import ALL_TEMP_FILES
    before = list(ALL_TEMP_FILES)
    with pytest.raises(ValidationError):
        encode_process_file("hours", [row], fmt)
    assert list(ALL_TEMP_FILES) == before


def test_csv_nul_and_nonbytes_inputs_rejected_with_source_diagnostic():
    content = file_bytes(["图号", "备注"], [["甲", "x\x00y"]], "csv")
    with pytest.raises(ValidationError) as caught:
        decode_process_file("route", content, "csv")
    assert caught.value.details["row"] == 2
    for value in (None, "图号", bytearray(b"x")):
        with pytest.raises(ValidationError):
            decode_process_file("route", value, "csv")
    with pytest.raises(ValidationError):
        encode_process_file("route", [{"business_code": "甲", "remark": "x\x00y"}], "csv")


def test_all_blank_records_still_count_towards_import_limit():
    with pytest.raises(ValidationError, match="2000") as caught:
        decode_process_file("route", "图号\n".encode() + b"\n" * 2010, "csv")
    assert caught.value.details["row"] == 2002
