"""Execution reads keep their real cohort without unrelated plan projections."""

import json
from datetime import datetime, timedelta
from io import BytesIO

import pytest
from openpyxl import load_workbook

from core.models.workbench_command import WorkbenchCommandRejected
from core.models.workbench_plan_scope import PlanReadScope
from core.services.workbench.plan import queries
from tests.workbench.execution_ledger_support import all_rows
from tests.workbench.field_workspace_support import FieldAPI, _ledger_fixture  # noqa: F401
from tests.workbench.point_downstream_support import ACTUAL, FIELD, app_for, read


def _three_tasks(case):
    second, third = case.op('OP2', seq=2), case.op('OP3', seq=3)
    case.plan(2, [case.op_id, second, third])
    for op_id, start, end in ((case.op_id, '08:00:00', '10:00:00'),
                              (second, '10:00:00', '12:00:00'), (third, '12:00:00', '14:00:00')):
        case.conn.execute('UPDATE Schedule SET start_time=?,end_time=? WHERE version=2 AND op_id=?',
                          ('2026-09-09T' + start, '2026-09-09T' + end, op_id))
    case.conn.commit()
    case.install()
    case.path = case.conn.execute('PRAGMA database_list').fetchone()[2]
    return case.plan_ref(2), [case.task(2, op_id) for op_id in (case.op_id, second, third)]


def _unrelated_too_large(*args, **kwargs):
    raise WorkbenchCommandRejected('query_too_large', 'Oversized unrelated baseline comparison.', 413)


def test_unrelated_plan_projection_cannot_block_execution_cohort(ledger_case, monkeypatch):
    case = ledger_case
    plan_ref, refs = _three_tasks(case)
    api = FieldAPI(case)
    task = api.task(plan_ref=plan_ref)
    api.create(case.values(2), task)
    client = app_for(case).test_client()
    before = all_rows(case.conn)
    monkeypatch.setattr(queries, 'workspace_projections', _unrelated_too_large)
    full = client.get('/api/workbench/v1/plans/' + plan_ref + '/workspace')
    assert full.status_code == 413 and full.get_json()['error']['code'] == 'query_too_large'
    scope = dict(plan_ref=plan_ref, range_start='2026-09-09T08:00:00', range_end='2026-09-09T10:00:00')
    field = read(client, FIELD + '/tasks', **scope)
    actual = read(client, ACTUAL, **scope)
    row, = field['data']['tasks']
    item, = actual['data']['items']
    assert row['task_ref'] == item['task']['task_ref'] == refs[0]
    # This legacy plan did not capture adopted quantity evidence. Keep its
    # explicit unknown task quantity while the live ledger has a known target.
    assert row['quantity'] is item['task']['quantity'] is None
    assert row['quantity_basis'] == item['task']['quantity_basis'] == 'unknown'
    assert row['execution']['target_quantity'] == item['execution']['target_quantity'] == 10
    assert row['planned_start'] == item['task']['start'] == scope['range_start']
    assert row['planned_end'] == item['task']['end'] == scope['range_end']
    assert row['execution']['execution_state'] == item['execution']['execution_state'] == 'partial'
    assert row['execution']['reports'][0]['revision_ref'] == item['execution']['reports'][0]['revision_ref']
    detail = read(client, FIELD + '/tasks/' + refs[0], **scope, snapshot_ref=field['meta']['snapshot_ref'])
    assert detail['data']['task'] == row
    assert actual['data']['calendar']['time_scope']['range_start'] == scope['range_start']
    assert actual['data']['calendar']['time_scope']['range_end'] == scope['range_end']
    assert all_rows(case.conn) == before


def test_scoped_actual_chain_retains_outside_predecessors_and_snapshot(ledger_case):
    case = ledger_case
    plan_ref, refs = _three_tasks(case)
    client = app_for(case).test_client()
    scope = dict(plan_ref=plan_ref, range_start='2026-09-09T12:00:00', range_end='2026-09-09T14:00:00')
    actual = read(client, ACTUAL, **scope)
    chain = actual['data']['critical_chain']
    assert actual['data']['task_count'] == 1 and actual['data']['items_complete'] is True
    assert chain['state'] == 'available' and chain['scope'] == 'full_plan'
    assert chain['plan_task_count'] == 3 and chain['task_refs'] == refs
    assert [node['in_scope'] for node in chain['nodes']] == [False, False, True]
    token = actual['meta']['snapshot_ref']
    related = read(client, ACTUAL + '/chain', **scope, snapshot_ref=token, target_task_ref=refs[-1])
    assert related['data']['critical_chain']['task_refs'] == refs
    case.conn.execute("UPDATE Schedule SET end_time='2026-09-09T09:30:00' WHERE version=2 AND op_id=?", (case.op_id,))
    case.conn.commit()
    for path, extra in ((ACTUAL + '/chain', dict(target_task_ref=refs[-1])),
                        (ACTUAL + '/export', dict(format='csv'))):
        response = client.get(path, query_string=dict(scope, snapshot_ref=token, **extra))
        assert response.status_code == 409 and response.get_json()['error']['code'] == 'snapshot_stale'


def test_field_scoped_page_write_and_files_keep_same_snapshot(ledger_case, monkeypatch, record_property):
    from core.services.workbench.execution.field_workspace import FieldWorkspaceService

    case = ledger_case
    plan_ref, refs = _three_tasks(case)
    api = FieldAPI(case)
    scope = dict(plan_ref=plan_ref, range_start='2026-09-09T08:00:00', range_end='2026-09-09T14:00:00')
    original, calls = FieldWorkspaceService.cohort, []

    def cohort(reader, values):
        calls.append(values)
        return original(reader, values)

    with monkeypatch.context() as patch:
        patch.setattr(FieldWorkspaceService, 'cohort', cohort)
        direct = api.read(**scope, size=1, page=2)
    assert len(calls) == 1 and direct['data']['page']['number'] == 2
    assert direct['data']['tasks'][0]['task_ref'] == refs[1]
    record_property('direct_page', 2)
    record_property('cohort_query_calls', len(calls))
    record_property('data_fields', sorted(direct['data']))
    token = direct['meta']['snapshot_ref']
    second = api.read(**scope, size=1, page=2, snapshot_ref=token)
    assert second['data'] == direct['data'] and second['meta'].keys() == direct['meta'].keys()
    assert second['meta']['request_ref'] != direct['meta']['request_ref']
    assert all(second['meta'][key] == value for key, value in direct['meta'].items() if key != 'request_ref')
    first = api.read(**scope, size=1, snapshot_ref=token)
    assert second['meta']['snapshot_ref'] == token and second['meta']['as_of'] == first['meta']['as_of']
    assert second['data']['tasks'][0]['task_ref'] == refs[1]
    located = api.read(**scope, size=1, task_ref=refs[1], operation_ref=second['data']['tasks'][0]['operation_ref'])
    assert located['data']['page']['number'] == 2 and located['data']['tasks'][0]['task_ref'] == refs[1]
    mismatch = api.client.get(FIELD + '/tasks', query_string=dict(scope, size=1, task_ref=refs[1],
        operation_ref=first['data']['tasks'][0]['operation_ref']))
    assert mismatch.status_code == 409 and mismatch.get_json()['error']['code'] == 'constraint_conflict'
    response = api.client.get(FIELD + '/files/template', query_string=dict(scope, snapshot_ref=token))
    assert response.status_code == 200
    book = load_workbook(BytesIO(response.data)); assert book.worksheets[0].max_row == 4; book.close()
    changed = api.client.get(FIELD + '/tasks', query_string=dict(scope, range_end='2026-09-09T12:00:00', snapshot_ref=token))
    assert changed.status_code == 409 and changed.get_json()['error']['code'] == 'snapshot_stale'
    api.create(case.values(2), first['data']['tasks'][0])
    for path, extra in ((FIELD + '/tasks', dict(page=2, size=1)),
                        (FIELD + '/files/template', {}), (FIELD + '/files/export', {})):
        response = api.client.get(path, query_string=dict(scope, snapshot_ref=token, **extra))
        assert response.status_code == 409 and response.get_json()['error']['code'] == 'snapshot_stale'


def test_internal_execution_cohort_uses_actual_endpoint_budget(ledger_case):
    case = ledger_case
    plan_ref, _ = _three_tasks(case)
    # A full internal cohort may exceed the plan endpoint's 8 MiB while the
    # paged field response remains small. Only the actual HTTP output is bounded.
    case.conn.execute('UPDATE BatchOperations SET op_type_name=? WHERE id=?', ('工序' * (2 * 1024 * 1024), case.op_id))
    case.conn.commit()
    before = all_rows(case.conn)
    reader = queries.WorkbenchPlanQueryService(case.conn)
    with reader.read_snapshot():
        data, _ = reader.execution_workspace(PlanReadScope(plan_ref))
    assert data['task_count'] == 3 and data['tasks_complete'] is True
    client = app_for(case).test_client()
    field = client.get(FIELD + '/tasks', query_string=dict(plan_ref=plan_ref, size=1, page=2))
    assert field.status_code == 200 and field.get_json()['data']['page']['total'] == 3
    oversized = client.get(FIELD + '/tasks', query_string=dict(plan_ref=plan_ref, size=1, page=1))
    assert oversized.status_code == 413 and oversized.get_json()['error']['code'] == 'query_too_large'
    assert all_rows(case.conn) == before


def test_actual_large_plan_applies_finish_batch_and_overlap_before_ledger(ledger_case, monkeypatch):
    from flask import g

    from core.services.workbench.execution import actual_gantt

    case = ledger_case
    batches = ['SCALE-' + str(index).zfill(3) for index in range(200)]
    case.conn.executemany("INSERT INTO Batches(batch_id,part_no,quantity) VALUES (?,'P1',10)",
                          [(batch,) for batch in batches])
    case.conn.executemany("INSERT INTO BatchOperations(op_code,batch_id,seq,op_type_id,op_type_name,source) "
                          "VALUES (?,?,?,'T1','Turning','internal')",
                          [('SCALE-' + str(index), batch, sequence)
                           for index, (batch, sequence) in enumerate((batch, sequence)
                               for batch in batches for sequence in range(1, 101))])
    rows = list(case.conn.execute("SELECT id,batch_id,seq FROM BatchOperations WHERE batch_id LIKE 'SCALE-%' ORDER BY id"))
    assert len(rows) == 20000
    start = datetime(2026, 9, 9, 8)
    schedule = []
    for index, row in enumerate(rows):
        low = start + timedelta(days=index // 200, minutes=index % 200)
        schedule.append((row['id'], low.isoformat(), (low + timedelta(minutes=1)).isoformat()))
    case.conn.execute("INSERT INTO ScheduleHistory(version,strategy,result_status,result_summary,schedule_time) "
                      "VALUES (2,'actual-range','success','{}','2026-09-09T08:00:00')")
    case.conn.executemany("INSERT INTO Schedule(version,op_id,machine_id,operator_id,start_time,end_time) "
                          "VALUES (2,?,'M1','O1',?,?)", schedule)
    case.conn.commit()
    case.install()
    case.path = case.conn.execute('PRAGMA database_list').fetchone()[2]
    plan_ref, projected, statements = case.plan_ref(2), [], []
    original = actual_gantt.ledger_projection

    def project(conn, ref, tasks):
        projected.append(len(tasks))
        return original(conn, ref, tasks)

    monkeypatch.setattr(actual_gantt, 'ledger_projection', project)
    monkeypatch.setattr(queries.WorkbenchPlanQueryService, 'execution_workspace',
                        lambda *args, **kwargs: pytest.fail('Actual reconstructed an execution workspace'))
    app = app_for(case)

    @app.before_request
    def trace():
        g.db.set_trace_callback(statements.append)

    client = app.test_client()
    scopes = [({'plan_finish_date_from': '2026-09-09', 'plan_finish_date_to': '2026-09-09'}, 200),
              ({'batch_ids': json.dumps([batches[0]])}, 100),
              ({'range_start': '2026-09-09T08:00:00', 'range_end': '2026-09-09T08:10:00'}, 10)]
    for scope, expected in scopes:
        statements.clear()
        response = client.get(ACTUAL, query_string=dict(plan_ref=plan_ref, **scope))
        assert response.status_code == 200, response.get_json()
        data = response.get_json()['data']
        assert data['task_count'] == len(data['items']) == projected[-1] == expected
        assert data['availability']['state'] == 'available' and data['items_complete'] is True
        assert data['critical_chain']['state'] == 'available'
        assert data['critical_chain']['plan_task_count'] == len(data['critical_chain']['nodes']) == 20000
        assert data['critical_chain']['visible_node_count'] == expected
        source_reads = [' '.join(sql.upper().replace('"', '').split()) for sql in statements]
        assert sum(' FROM SCHEDULE R WHERE R.VERSION = 2 ' in sql for sql in source_reads) == 1
    response = client.get(ACTUAL, query_string={'plan_ref': plan_ref})
    assert response.status_code == 413 and projected[-1] == 20000
    error = response.get_json()['error']
    assert error['code'] == 'query_too_large' and '10000' in error['message']
    assert '计划完工日期' in error['message'] and '批次' in error['message']


def test_actual_resource_filter_keeps_planned_and_reported_owners(ledger_case):
    case = ledger_case
    plan_ref, refs = _three_tasks(case)
    case.conn.execute("INSERT INTO Machines(machine_id,name,op_type_id) VALUES ('M2','Reported machine','T1')")
    case.conn.execute("INSERT INTO OperatorMachine(operator_id,machine_id) VALUES ('O1','M2')")
    case.conn.commit()
    api = FieldAPI(case)
    task = next(row for row in api.read(plan_ref=plan_ref)['data']['tasks'] if row['task_ref'] == refs[0])
    api.create(case.values(2, actual_machine_ref=case.ref('machine', 'M2')), task)
    client = app_for(case).test_client()
    scope = dict(plan_ref=plan_ref, plan_finish_date_from='2026-09-09',
                 plan_finish_date_to='2026-09-09', batch_ids=json.dumps(['B1']), resource_type='machine')
    for key, expected in (('M1', refs), ('M2', refs[:1])):
        data = read(client, ACTUAL, resource_ref=case.ref('machine', key), **scope)['data']
        assert {item['task']['task_ref'] for item in data['items']} == set(expected)
