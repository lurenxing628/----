"""BytesIO fixtures and explicit CommandService integration for material files."""

import csv
import json
import tracemalloc
from io import BytesIO, StringIO
from itertools import zip_longest
from time import perf_counter
from xml.etree import ElementTree
from zipfile import ZipFile

import openpyxl

from core.infrastructure.transaction import TransactionManager
from core.models.workbench_material_file import COLUMNS, HEADERS, MATERIAL_COLUMNS
from core.services.workbench.commands import WorkbenchCommandService
from core.services.workbench.material_bulk import WorkbenchMaterialBulkService
from core.services.workbench.material_files import WorkbenchMaterialFileService
from tests.workbench.identity_metadata_support import business_snapshot

KEY = "material-file-test-000001"


def file_bytes(rows, file_format, headers=HEADERS, *, formulas=False):
    if file_format == "csv":
        output = StringIO(newline="")
        writer = csv.writer(output, lineterminator="\n")
        writer.writerow(headers)
        writer.writerows(rows)
        return output.getvalue().encode("utf-8-sig")
    wb = openpyxl.Workbook()
    try:
        ws = wb.active
        ws.append(headers)
        for number, row in enumerate(rows, 2):
            for column, value in enumerate(row, 1):
                cell = ws.cell(number, column, value)
                if type(value) is str and not formulas:
                    cell.data_type = "s"
        output = BytesIO()
        wb.save(output)
        return output.getvalue()
    finally:
        wb.close()


def confirm_import(conn, preview, content, file_format, *, key=KEY, guard=None, command=None):
    return (command or WorkbenchCommandService(conn)).execute(
        request_key=key, action="material.import", context_ref="material.collection",
        normalized_input=preview.intent(), guard=guard or (lambda: None),
        mutate=lambda _: WorkbenchMaterialFileService(conn).confirm_import(preview, content, file_format=file_format, scope={}),
    )


def confirm_delete(conn, preview, refs, *, key=KEY, scope=None, guard=None, command=None):
    return (command or WorkbenchCommandService(conn)).execute(
        request_key=key, action="material.bulk_delete", context_ref="material.collection",
        normalized_input=preview.intent(), guard=guard or (lambda: None),
        mutate=lambda _: WorkbenchMaterialBulkService(conn).confirm_delete(preview, refs, scope=scope or {}),
    )


def export_file(conn, file_format, **selection):
    with TransactionManager(conn).transaction():
        return WorkbenchMaterialFileService(conn).export(file_format, **selection)


def seed_many(conn, count):
    conn.executemany("INSERT INTO Materials (material_id, name, stock_qty, status) VALUES (?, ?, ?, ?)",
                     [(f"BULK{n:05d}", f"name{n:05d}", n / 4, "active") for n in range(count)])
    conn.commit()


def small_dimensions(content):
    output = BytesIO()
    with ZipFile(BytesIO(content)) as source, ZipFile(output, "w") as target:
        for entry in source.infolist():
            payload = source.read(entry.filename)
            if entry.filename == "xl/worksheets/sheet1.xml":
                tree = ElementTree.fromstring(payload)
                tree.find("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}dimension").set("ref", "A1")
                payload = ElementTree.tostring(tree)
            target.writestr(entry, payload)
    return output.getvalue()


def measure(operation, call, **details):
    """Incremental Python allocation peak, not total RSS; timing includes tracing overhead."""
    tracemalloc.start()
    started = perf_counter()
    try:
        result = call()
        elapsed = perf_counter() - started
        _, peak = tracemalloc.get_traced_memory()
    finally:
        tracemalloc.stop()
    metric = {"operation": operation, "elapsed_s": round(elapsed, 6), "python_peak_bytes": peak, **details}
    if hasattr(result, "content"):
        metric.update(file_bytes=len(result.content), rows=result.row_count)
    print("MATERIAL_METRIC " + json.dumps(metric, sort_keys=True))
    return result


def seed_export_scale(conn):
    """9999 diverse rows plus the shared MAT1 fixture: exactly 10000 materials."""
    codes = ("000123", "2026-01-02", "=1+1", "+cmd", "-2", "@SUM(A1)", "'literal")
    rows = []
    for n in range(9999):
        rows.append((codes[n] if n < len(codes) else f"BULK{n:05d}", f"同名{n % 17:02d}",
                     None if n % 3 == 0 else r"\N", None if n % 5 == 0 else "kg",
                     None if n % 7 == 0 else n / 4, ("active", "inactive", "legacy HOLD ", None)[n % 4],
                     "=中文,\"文字\"\n第二行" if n % 2 else None, f"2000-01-{n % 28 + 1:02d} 01:02:03"))
    conn.executemany("INSERT INTO Materials (material_id, name, spec, unit, stock_qty, status, remark, created_at) VALUES (?, ?, ?, ?, ?, ?, ?, ?)", rows)
    conn.commit()
    return [dict(row) for row in conn.execute("SELECT * FROM Materials ORDER BY material_id")]


def _expected_file_values(row, file_format):
    result = []
    for field, column in zip(COLUMNS, MATERIAL_COLUMNS):
        value = row[column]
        if field == "stock_qty":
            if file_format == "csv":
                value = "" if value is None else str(value)
        else:
            if value is None:
                value = r"\N" if field in ("spec", "unit", "remark") else None
            elif value.startswith("\\"):
                value = "\\" + value
            if file_format == "csv":
                value = "'" + value if value is not None else ""
        result.append(value)
    return result


def verify_download(download, expected, file_format):
    """Independent reader/oracle: do not invoke the import parser (which has a 2000-row cap)."""
    if file_format == "csv":
        rows = csv.reader(StringIO(download.content.decode("utf-8-sig"), newline=""))
        assert next(rows) == list(HEADERS)
        count = 0
        for count, (actual, source) in enumerate(zip_longest(rows, expected), 1):
            assert actual is not None and source is not None, count
            assert actual == _expected_file_values(source, file_format), count
    else:
        wb = openpyxl.load_workbook(BytesIO(download.content), read_only=True, data_only=False)
        try:
            rows = wb.worksheets[0].iter_rows(max_col=len(HEADERS))
            assert [cell.value for cell in next(rows)] == list(HEADERS)
            count = 0
            for count, (actual, source) in enumerate(zip_longest(rows, expected), 1):
                assert actual is not None and source is not None, count
                assert all(cell.data_type != "f" for cell in actual), count
                assert [cell.value for cell in actual] == _expected_file_values(source, file_format), count
        finally:
            wb.close()
    assert count == len(expected) == download.row_count


def exercise_scale_export(conn, file_format, selection):
    rows = seed_export_scale(conn)
    before, changes = business_snapshot(conn), conn.total_changes
    if selection == "all":
        expected, arguments = rows, {"scope": {}}
    elif selection == "filtered":
        expected = [row for row in rows if "bulk" in row["material_id"].lower() and row["status"] == "active"]
        expected.sort(key=lambda row: row["name"], reverse=True)
        arguments = {"scope": {"query": "BULK", "status": "active", "sort": "label", "direction": "desc"}}
    else:
        expected = list(reversed(rows[::3]))
        refs = dict(conn.execute("SELECT entity_key, ref FROM WorkbenchEntityRefs WHERE kind = 'material' AND active = 1"))
        arguments = {"selected_refs": [refs[row["material_id"]] for row in expected]}
    assert len(rows) == 10000 and len(expected) > 2000
    call = lambda: export_file(conn, file_format, **arguments)
    download = measure("export_all", call, format=file_format) if selection == "all" else call()
    verify_download(download, expected, file_format)
    assert business_snapshot(conn) == before and conn.total_changes == changes
    assert not conn.in_transaction


class OnePassRows:
    def __init__(self, rows):
        self.rows = iter(rows)
        self.iterations = 0

    def __iter__(self):
        self.iterations += 1
        assert self.iterations == 1, "export rows must not be consumed twice"
        return self.rows

    def __len__(self):
        raise AssertionError("an export stream has no len()")


def codec_rows(rows):
    return (dict(zip(COLUMNS, (row[key] for key in MATERIAL_COLUMNS))) for row in rows)
