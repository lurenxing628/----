"""回归测试：/reports/execution-review “计划和现场实际”页与 xlsx 导出——按计划身份精确匹配现场反馈，无反馈时显示“暂无现场反馈”，身份不匹配时忽略反馈，匹配时展示实际开工/完工偏差、暂停、异常、计划/实际资源等，且页面与工作簿都不泄漏 event_type/plan_role/source_table/schedule_id 等内部字段；含日期参数校验、报表索引/资源派工入口、场景导航禁用、流式导出与离线模板契约。"""

from __future__ import annotations

from io import BytesIO
from urllib.parse import unquote

import openpyxl
import pytest

from core.infrastructure.database import get_connection
from core.infrastructure.errors import ValidationError
from core.services.report import ReportEngine
from core.services.report.execution_review import _execution_scope
from core.services.scheduler.operation_execution_labels import action_to_event_type
from data.repositories.operation_execution_event_repo import OperationExecutionEventRepo
from tests.operation_execution_feedback_test_support import _build_app

EXPECTED_HEADERS = [
    "批次",
    "工序",
    "计划开始",
    "实际开始",
    "开工偏差",
    "计划结束",
    "实际结束",
    "完工偏差",
    "暂停时长",
    "异常原因",
    "严重程度",
    "预计影响时间",
    "影响设备",
    "影响人员",
    "处理状态",
    "是否建议重新排程",
    "计划资源",
    "实际资源",
    "现场反馈状态",
]


def _load_xlsx(resp):
    return openpyxl.load_workbook(BytesIO(resp.data), data_only=True)


def _insert_event(repo: OperationExecutionEventRepo, **overrides) -> None:
    event_type = action_to_event_type(overrides.get("event_type") or "start")
    reported_status = {
        "start": "processing",
        "pause": "paused",
        "resume": "processing",
        "exception": "exception",
        "finish": "completed",
    }[event_type]
    payload = {
        "schedule_version": 1,
        "schedule_id": 101,
        "op_id": 10,
        "batch_id": "B1",
        "source_table": "schedule",
        "effective_plan_role": "adopted",
        "scenario_id": None,
        "event_type": event_type,
        "reported_status": reported_status,
        "event_time": "2026-05-01 08:10:00",
        "actual_machine_id": None,
        "actual_operator_id": None,
        "quantity_done": None,
        "quantity_scrapped": None,
        "reason_code": None,
        "reason_detail": None,
        "severity": None,
        "impact_minutes": None,
        "affected_machine_id": None,
        "affected_operator_id": None,
        "handling_status": None,
        "suggest_reschedule": 0,
        "remark": None,
        "created_by": "pytest",
        "idempotency_key": f"execution-review-{event_type}-{overrides.get('previous_state_revision')}",
        "request_fingerprint": f"execution-review-{event_type}-{overrides.get('previous_state_revision')}",
        "previous_state_revision": "10:0:0",
    }
    payload.update(overrides)
    payload["event_type"] = action_to_event_type(payload["event_type"])
    payload["reported_status"] = {
        "start": "processing",
        "pause": "paused",
        "resume": "processing",
        "exception": "exception",
        "finish": "completed",
    }[payload["event_type"]]
    repo.insert_event(payload)


def _seed_execution_events(db_path: str, **identity) -> None:
    conn = get_connection(db_path)
    try:
        repo = OperationExecutionEventRepo(conn)
        _insert_event(
            repo,
            **identity,
            event_type="start",
            event_time="2026-05-01 08:10:00",
            actual_machine_id="M2",
            actual_operator_id="O2",
            previous_state_revision="10:0:0",
        )
        _insert_event(
            repo,
            **identity,
            event_type="pause",
            event_time="2026-05-01 08:20:00",
            reason_code="equipment",
            remark="设备检查",
            previous_state_revision="10:1:1",
        )
        _insert_event(
            repo,
            **identity,
            event_type="resume",
            event_time="2026-05-01 08:35:00",
            remark="继续生产",
            previous_state_revision="10:2:2",
        )
        _insert_event(
            repo,
            **identity,
            event_type="exception",
            event_time="2026-05-01 08:45:00",
            reason_code="equipment",
            reason_detail="主轴异常",
            severity="high",
            impact_minutes=30,
            affected_machine_id="M2",
            affected_operator_id="O2",
            handling_status="checking",
            suggest_reschedule=1,
            remark="主轴异常",
            previous_state_revision="10:3:3",
        )
        _insert_event(
            repo,
            **identity,
            event_type="finish",
            event_time="2026-05-01 09:05:00",
            quantity_done=10,
            quantity_scrapped=0,
            remark="完成",
            previous_state_revision="10:4:4",
        )
        conn.commit()
    finally:
        conn.close()


def _insert_out_of_range_plan_row(db_path: str) -> None:
    conn = get_connection(db_path)
    try:
        conn.execute(
            """
            INSERT INTO BatchOperations(id, op_code, batch_id, piece_id, seq, op_type_name, source, status)
            VALUES (11, 'OP20', 'B1', 'piece-a', 20, '磨削', 'internal', 'scheduled')
            """
        )
        conn.execute(
            """
            INSERT INTO Schedule(id, op_id, machine_id, operator_id, start_time, end_time, lock_status, version)
            VALUES (102, 11, 'M1', 'O1', '2026-05-02 08:00:00', '2026-05-02 09:00:00', 'unlocked', 2)
            """
        )
        conn.commit()
    finally:
        conn.close()


def _schedule_row(db_path: str, schedule_id: int = 100):
    conn = get_connection(db_path)
    try:
        row = conn.execute(
            """
            SELECT id, op_id, machine_id, operator_id, start_time, end_time, version
            FROM Schedule
            WHERE id = ?
            """,
            (int(schedule_id),),
        ).fetchone()
        return dict(row)
    finally:
        conn.close()


def _all_workbook_values(wb) -> str:
    values = []
    for ws in wb.worksheets:
        values.append(ws.title)
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None:
                    values.append(str(cell.value))
    return "\n".join(values)


def _summary_dict(wb):
    ws = wb["查询摘要"]
    out = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        key, value = row[0], row[1]
        if key:
            out[str(key)] = value
    return out


def _assert_no_feedback_page(body: str) -> None:
    assert "2026-05-01 08:10:00" not in body
    assert "2026-05-01 09:05:00" not in body
    assert "晚了 10 分钟" not in body
    assert "二号设备 / 李四" not in body
    assert "已完工" not in body
    assert "暂无现场反馈" in body


def _assert_feedback_page(body: str) -> None:
    for text in (
        "2026-05-01 08:10:00",
        "2026-05-01 09:05:00",
        "晚了 10 分钟",
        "晚了 5 分钟",
        "15 分钟",
        "设备问题",
        "严重",
        "预计影响 30 分钟",
        "二号设备 / 李四",
        "已完工",
    ):
        assert text in body
    assert "report_exception" not in body
    assert "event_type" not in body


def _assert_no_feedback_workbook(wb) -> None:
    ws = wb["计划和现场实际"]
    for cell in ("D2", "E2", "G2", "H2", "I2", "J2", "K2", "L2", "M2", "N2", "O2", "P2", "R2", "S2"):
        assert ws[cell].value == "暂无现场反馈"


def _assert_feedback_workbook(wb) -> None:
    ws = wb["计划和现场实际"]
    expected = {
        "D2": "2026-05-01 08:10:00",
        "E2": "晚了 10 分钟",
        "G2": "2026-05-01 09:05:00",
        "H2": "晚了 5 分钟",
        "I2": "15 分钟",
        "J2": "设备问题",
        "K2": "严重",
        "L2": "预计影响 30 分钟",
        "M2": "二号设备",
        "N2": "李四",
        "O2": "处理中",
        "P2": "建议重新排程",
        "R2": "二号设备 / 李四",
        "S2": "已完工",
    }
    for cell, value in expected.items():
        assert ws[cell].value == value


def _assert_matching_feedback_page(client, db_path: str) -> None:
    before_schedule = _schedule_row(db_path)
    page = client.get("/reports/execution-review?version=2&date_from=2026-05-01&date_to=2026-05-01&batch_id=B1")
    body = page.get_data(as_text=True)
    assert page.status_code == 200
    assert body.count("<tr>") == 2
    assert "2026-05-01 09:00:00" in body
    assert "2026-05-02 08:00:00" not in body
    _assert_feedback_page(body)
    assert _schedule_row(db_path) == before_schedule


def _assert_empty_feedback_page(client) -> None:
    empty_page = client.get("/reports/execution-review?version=2&batch_id=NO_SUCH")
    assert empty_page.status_code == 200
    assert "当前筛选条件下暂无计划任务" in empty_page.get_data(as_text=True)


def _assert_matching_export_response(export_resp) -> None:
    assert export_resp.status_code == 200
    disposition = unquote(export_resp.headers.get("Content-Disposition", ""))
    assert "计划和现场实际-正式采用方案-v2-2026-05-01 至 2026-05-01.xlsx" in disposition


def _assert_matching_export_summary(wb) -> None:
    assert "查询摘要" in wb.sheetnames
    summary = _summary_dict(wb)
    assert summary["报表"] == "计划和现场实际"
    assert summary["计划版本"] == "v2"
    assert summary["方案"] == "正式采用方案"
    assert summary["查询日期"] == "2026-05-01 至 2026-05-01"
    assert summary["批次"] == "B1"


def _assert_matching_planned_workbook_row(wb) -> None:
    ws = wb["计划和现场实际"]
    assert [cell.value for cell in ws[1]] == EXPECTED_HEADERS
    assert ws["A2"].value == "B1"
    assert ws["B2"].value == "OP10 / 车削"
    assert ws["C2"].value == "2026-05-01 08:00:00"
    assert ws["F2"].value == "2026-05-01 09:00:00"
    assert ws["Q2"].value == "一号设备 / 张三"


def _assert_no_internal_workbook_values(wb) -> None:
    all_values = _all_workbook_values(wb)
    assert "完整身份" not in all_values
    for token in ("plan_role", "source_table", "event_type", "report_exception", "schedule_id"):
        assert token not in all_values


def _assert_report_index_entry(client) -> None:
    index_page = client.get("/reports/")
    assert index_page.status_code == 200
    index_body = index_page.get_data(as_text=True)
    assert "计划和现场实际" in index_body
    assert "/reports/execution-review" in index_body


def _assert_dispatch_entry(client) -> None:
    dispatch_page = client.get("/scheduler/resource-dispatch?version=2&period_preset=custom&start_date=2026-05-01&end_date=2026-05-01")
    dispatch_body = dispatch_page.get_data(as_text=True)
    assert dispatch_page.status_code == 200
    assert "/reports/execution-review" in dispatch_body
    assert "/scheduler/execution-review" not in dispatch_body


def _assert_scenario_nav_disables_execution_review(client) -> None:
    nav_page = client.get(
        "/reports/utilization?version=2&plan_role=adopted&scenario_id=scenario-plain"
        "&date_from=2026-05-01&date_to=2026-05-01&batch_id=B1"
    )
    nav_body = nav_page.get_data(as_text=True)
    assert nav_page.status_code == 200
    assert "计划和现场实际只复盘正式采用方案" in nav_body
    assert 'aria-disabled="true"' in nav_body
    assert "/reports/execution-review?version=2&amp;plan_role" not in nav_body
    assert "/reports/execution-review?version=2&amp;scenario_id" not in nav_body


def _assert_stream_export(db_path: str) -> None:
    conn = get_connection(db_path)
    try:
        engine = ReportEngine(conn)
        engine.EXPORT_DIRECT_MAX_ROWS = 0
        export = engine.export_execution_review_xlsx(2, date_from="2026-05-01", date_to="2026-05-01", batch_id="B1")
        assert export.mode == "stream"
        export.data.seek(0)
        wb = openpyxl.load_workbook(BytesIO(export.data.read()), data_only=True)
        try:
            assert "计划和现场实际" in wb.sheetnames
            ws = wb["计划和现场实际"]
            assert [cell.value for cell in ws[1]] == EXPECTED_HEADERS
            assert ws["R2"].value == "二号设备 / 李四"
        finally:
            wb.close()
    finally:
        conn.close()


def test_execution_review_page_and_export_show_no_feedback_state(tmp_path, monkeypatch) -> None:
    app, _db_path = _build_app(tmp_path, monkeypatch)
    client = app.test_client()

    page = client.get("/reports/execution-review?version=2")
    body = page.get_data(as_text=True)
    assert page.status_code == 200
    assert "计划和现场实际" in body
    assert "暂无现场反馈" in body
    assert "一号设备 / 张三" in body
    assert "actual_machine_id" not in body
    assert "event_type" not in body

    export_resp = client.get("/reports/execution-review/export?version=2")
    assert export_resp.status_code == 200
    disposition = unquote(export_resp.headers.get("Content-Disposition", ""))
    assert "计划和现场实际-正式采用方案-v2.xlsx" in disposition
    wb = _load_xlsx(export_resp)
    try:
        ws = wb["计划和现场实际"]
        assert [cell.value for cell in ws[1]] == EXPECTED_HEADERS
        assert ws["A2"].value == "B1"
        assert ws["D2"].value == "暂无现场反馈"
        assert ws["E2"].value == "暂无现场反馈"
        assert ws["R2"].value == "暂无现场反馈"
        assert ws["S2"].value == "暂无现场反馈"
    finally:
        wb.close()


def test_execution_review_ignores_feedback_when_plan_identity_mismatches(tmp_path, monkeypatch) -> None:
    app, db_path = _build_app(tmp_path, monkeypatch)
    _seed_execution_events(db_path)
    client = app.test_client()

    page = client.get("/reports/execution-review?version=2&date_from=2026-05-01&date_to=2026-05-01&batch_id=B1")
    body = page.get_data(as_text=True)
    assert page.status_code == 200
    assert body.count("<tr>") == 2
    _assert_no_feedback_page(body)

    export_resp = client.get(
        "/reports/execution-review/export?version=2&date_from=2026-05-01&date_to=2026-05-01&batch_id=B1"
    )
    wb = _load_xlsx(export_resp)
    try:
        assert export_resp.status_code == 200
        _assert_no_feedback_workbook(wb)
    finally:
        wb.close()


def test_execution_review_uses_matching_execution_state_for_actual_times_and_resources(tmp_path, monkeypatch) -> None:
    app, db_path = _build_app(tmp_path, monkeypatch)
    _seed_execution_events(db_path, schedule_version=2, schedule_id=100)
    _insert_out_of_range_plan_row(db_path)
    client = app.test_client()
    before_schedule = _schedule_row(db_path)

    _assert_matching_feedback_page(client, db_path)
    _assert_empty_feedback_page(client)

    export_resp = client.get(
        "/reports/execution-review/export?version=2&date_from=2026-05-01&date_to=2026-05-01&batch_id=B1"
    )
    _assert_matching_export_response(export_resp)

    wb = _load_xlsx(export_resp)
    try:
        _assert_matching_export_summary(wb)
        _assert_matching_planned_workbook_row(wb)
        _assert_feedback_workbook(wb)
        _assert_no_internal_workbook_values(wb)
    finally:
        wb.close()
    assert _schedule_row(db_path) == before_schedule


def test_execution_review_validation_and_offline_template_contract(tmp_path, monkeypatch) -> None:
    app, _db_path = _build_app(tmp_path, monkeypatch)
    client = app.test_client()

    missing_date = client.get("/reports/execution-review?version=2&date_from=2026-05-01")
    assert missing_date.status_code == 400
    assert "开始日期和结束日期要一起填写" in missing_date.get_data(as_text=True)

    reversed_date = client.get("/reports/execution-review/export?version=2&date_from=2026-05-02&date_to=2026-05-01")
    assert reversed_date.status_code == 400
    assert "结束日期不能早于开始日期" in reversed_date.get_data(as_text=True)

    with open("templates/reports/execution_review.html", encoding="utf-8") as fh:
        template = fh.read()
    forbidden = ("http://", "https://", "cdn", "reason_code", "event_type", "report_exception")
    for token in forbidden:
        assert token not in template


def test_execution_review_scope_fails_loudly_when_plan_identity_is_incomplete() -> None:
    with pytest.raises(ValidationError, match="计划数据缺少现场执行身份字段"):
        _execution_scope({"op_id": 10, "batch_id": "B1"}, 10)


def test_execution_review_entry_boundaries_and_stream_export(tmp_path, monkeypatch) -> None:
    app, db_path = _build_app(tmp_path, monkeypatch)
    _seed_execution_events(db_path, schedule_version=2, schedule_id=100)
    client = app.test_client()

    _assert_report_index_entry(client)
    _assert_dispatch_entry(client)

    missing_scheduler_route = client.get("/scheduler/execution-review")
    assert missing_scheduler_route.status_code == 404

    _assert_scenario_nav_disables_execution_review(client)
    _assert_stream_export(db_path)
