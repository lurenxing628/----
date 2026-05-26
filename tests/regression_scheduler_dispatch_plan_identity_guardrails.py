from __future__ import annotations

import io
import json
import sqlite3
from pathlib import Path
from types import SimpleNamespace
from typing import Any, Dict, List, Optional, Set
from urllib.parse import unquote

import openpyxl
from flask import Flask, g

from core.infrastructure.database import ensure_schema, get_connection
from core.models.schedule_candidate import ScheduleCandidate, ScheduleCandidateRows, ScheduleCandidateSelection
from core.services.scheduler.resource_dispatch_excel import build_resource_dispatch_workbook
from core.services.scheduler.resource_dispatch_service import ResourceDispatchService
from core.services.scheduler.schedule_plan_query_service import ROLE_ADOPTED, ROLE_BASELINE_BEST, ROLE_CRITICAL_BEST
from core.services.scheduler.schedule_result_view_context import plan_role_filter_fields
from data.repositories.schedule_candidate_repo import ScheduleCandidateRepository
from data.repositories.schedule_plan_query_repo import SOURCE_CANDIDATE_ROWS, SOURCE_SCHEDULE
from web.viewmodels.scheduler_resource_dispatch import (
    build_resource_dispatch_filename,
    decorate_resource_dispatch_context,
    decorate_resource_dispatch_payload,
)

REPO_ROOT = Path(__file__).resolve().parents[1]
VERSION = 7
HISTORICAL_VERSION = 6

def _connect_fresh_schema(tmp_path: Path) -> sqlite3.Connection:
    db_path = tmp_path / "aps.db"
    ensure_schema(str(db_path), schema_path=str(REPO_ROOT / "schema.sql"), backup_dir=str(tmp_path / "backups"))
    return get_connection(str(db_path))

def _require_id(value: Any) -> int:
    assert value is not None
    return int(value)

def _seed_schedule_context(conn: sqlite3.Connection) -> None:
    conn.executescript(
        """
        INSERT INTO ResourceTeams(team_id, name, status)
        VALUES ('T-ADOPTED', '正式班组', 'active'), ('T-CANDIDATE', '对比班组', 'active');

        INSERT INTO Machines(machine_id, name, status, team_id)
        VALUES ('M-ADOPTED', '正式设备', 'active', 'T-ADOPTED'), ('M-CANDIDATE', '对比设备', 'active', 'T-CANDIDATE');

        INSERT INTO Operators(operator_id, name, status, team_id)
        VALUES ('O-ADOPTED', '正式人员', 'active', 'T-ADOPTED'), ('O-CANDIDATE', '对比人员', 'active', 'T-CANDIDATE');

        INSERT INTO Parts(part_no, part_name)
        VALUES ('P001', '零件一');

        INSERT INTO Batches(batch_id, part_no, part_name, quantity, due_date, priority, ready_status, status)
        VALUES ('B1', 'P001', '零件一', 1, '2026-05-20', 'normal', 'yes', 'scheduled');

        INSERT INTO BatchOperations(id, op_code, batch_id, piece_id, seq, op_type_name, source, status)
        VALUES (10, 'OP10', 'B1', 'piece-a', 10, '车削', 'internal', 'scheduled');

        INSERT INTO Schedule(id, op_id, machine_id, operator_id, start_time, end_time, lock_status, version)
        VALUES
          (60, 10, 'M-ADOPTED', 'O-ADOPTED', '2026-05-01 08:00:00', '2026-05-01 09:00:00', 'locked', 6),
          (70, 10, 'M-ADOPTED', 'O-ADOPTED', '2026-05-05 08:00:00', '2026-05-05 09:00:00', 'unlocked', 7);

        INSERT INTO ScheduleHistory(version, strategy, batch_count, op_count, result_status, result_summary, created_by)
        VALUES
          (6, 'greedy', 1, 1, 'success', '{}', 'pytest'),
          (7, 'greedy', 1, 1, 'success', '{}', 'pytest');
        """
    )

def _seed_candidates(conn: sqlite3.Connection) -> None:
    repo = ScheduleCandidateRepository(conn)
    adopted = repo.create_candidate(
        ScheduleCandidate(
            id=None,
            version=VERSION,
            candidate_key="adopted",
            candidate_label="最终采用",
            candidate_kind="baseline",
            status="completed",
            graph_enabled="no",
            detail_saved="no",
        )
    )
    baseline = repo.create_candidate(
        ScheduleCandidate(
            id=None,
            version=VERSION,
            candidate_key="baseline_best",
            candidate_label="原算法候选",
            candidate_kind="baseline",
            status="completed",
            graph_enabled="no",
            detail_saved="yes",
        )
    )
    baseline_id = _require_id(baseline.id)
    repo.bulk_create_candidate_rows(
        [
            ScheduleCandidateRows(
                id=None,
                version=VERSION,
                candidate_id=baseline_id,
                op_id=10,
                machine_id="M-CANDIDATE",
                operator_id="O-CANDIDATE",
                start_time="2026-05-05 13:00:00",
                end_time="2026-05-05 14:00:00",
                lock_status="locked",
            )
        ]
    )
    repo.create_selection(
        ScheduleCandidateSelection(
            id=None,
            version=VERSION,
            role=ROLE_ADOPTED,
            candidate_id=_require_id(adopted.id),
            source_table=SOURCE_SCHEDULE,
        )
    )
    repo.create_selection(
        ScheduleCandidateSelection(
            id=None,
            version=VERSION,
            role=ROLE_BASELINE_BEST,
            candidate_id=baseline_id,
            source_table=SOURCE_CANDIDATE_ROWS,
        )
    )

def _seed_scenario(conn: sqlite3.Connection) -> None:
    scenario_id = "scenario-guardrail"
    conn.execute(
        """
        INSERT INTO ScheduleAdjustmentScenario(
            scenario_id, source_draft_id, base_version, base_plan_role,
            base_source_table, base_candidate_id, base_candidate_key,
            scenario_name, status, validation_status, issue_count, row_count, created_by
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            scenario_id,
            "draft-guardrail",
            VERSION,
            ROLE_ADOPTED,
            SOURCE_SCHEDULE,
            None,
            "adopted",
            "",
            "active",
            "valid",
            0,
            1,
            "pytest",
        ),
    )
    conn.execute(
        """
        INSERT INTO ScheduleAdjustmentScenarioRow(
            scenario_id, source_table, source_row_id, op_id, machine_id, operator_id,
            start_time, end_time, lock_status, is_changed
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            scenario_id,
            SOURCE_SCHEDULE,
            70,
            10,
            "M-CANDIDATE",
            "O-CANDIDATE",
            "2026-05-06 08:00:00",
            "2026-05-06 09:00:00",
            "unlocked",
            "yes",
        ),
    )

def _seed_db(tmp_path: Path) -> sqlite3.Connection:
    conn = _connect_fresh_schema(tmp_path)
    _seed_schedule_context(conn)
    _seed_candidates(conn)
    _seed_scenario(conn)
    conn.commit()
    return conn

def _dispatch_payload(
    conn: sqlite3.Connection,
    *,
    version: int,
    plan_role: str = ROLE_ADOPTED,
    operator_id: str = "O-ADOPTED",
    query_date: str = "2026-05-05",
    scenario_id: Optional[str] = None,
) -> Dict[str, Any]:
    raw = ResourceDispatchService(conn, logger=None, op_logger=None).get_dispatch_payload(
        scope_type="operator",
        operator_id=operator_id,
        period_preset="week",
        query_date=query_date,
        version=version,
        plan_role=plan_role,
        scenario_id=scenario_id,
    )
    return decorate_resource_dispatch_payload(raw)

def _page_context(conn: sqlite3.Connection, **kwargs: Any) -> Dict[str, Any]:
    raw = ResourceDispatchService(conn, logger=None, op_logger=None).build_page_context(
        scope_type=kwargs.get("scope_type", "operator"),
        operator_id=kwargs.get("operator_id", "O-ADOPTED"),
        period_preset=kwargs.get("period_preset", "week"),
        query_date=kwargs.get("query_date", "2026-05-05"),
        version=kwargs.get("version", VERSION),
        plan_role=kwargs.get("plan_role", ROLE_ADOPTED),
        scenario_id=kwargs.get("scenario_id"),
    )
    return decorate_resource_dispatch_context(raw)

def _summary_values(payload: Dict[str, Any]) -> Dict[str, Any]:
    buffer = build_resource_dispatch_workbook(payload)
    wb = openpyxl.load_workbook(io.BytesIO(buffer.getvalue()), read_only=True, data_only=True)
    try:
        ws = wb["查询摘要"]
        return {str(row[0] or ""): row[1] for row in ws.iter_rows(values_only=True) if row and row[0]}
    finally:
        wb.close()

def _json_keys(value: Any) -> Set[str]:
    if isinstance(value, dict):
        keys = set(value)
        for child in value.values():
            keys.update(_json_keys(child))
        return keys
    if isinstance(value, list):
        keys: Set[str] = set()
        for child in value:
            keys.update(_json_keys(child))
        return keys
    return set()

def _workbook_text(payload: Dict[str, Any]) -> str:
    buffer = build_resource_dispatch_workbook(payload)
    wb = openpyxl.load_workbook(io.BytesIO(buffer.getvalue()), read_only=True, data_only=True)
    try:
        values: List[str] = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                values.extend(str(value or "") for value in row)
        return "\n".join(values)
    finally:
        wb.close()

def _count_rows(conn: sqlite3.Connection, table_name: str) -> Optional[int]:
    exists = conn.execute(
        "SELECT 1 FROM sqlite_master WHERE type = 'table' AND name = ?",
        (table_name,),
    ).fetchone()
    if exists is None:
        return None
    return int(conn.execute(f"SELECT COUNT(1) FROM {table_name}").fetchone()[0])

def _table_names(conn: sqlite3.Connection) -> Set[str]:
    rows = conn.execute("SELECT name FROM sqlite_master WHERE type = 'table'").fetchall()
    return {str(row[0]) for row in rows}

def _build_route_app() -> Flask:
    app = Flask(__name__)
    app.secret_key = "test"
    for path, endpoint in (
        ("/scheduler/resource-dispatch", "scheduler.resource_dispatch_page"),
        ("/scheduler/resource-dispatch/data", "scheduler.resource_dispatch_data"),
        ("/scheduler/resource-dispatch/export", "scheduler.resource_dispatch_export"),
    ):
        app.add_url_rule(path, endpoint=endpoint, view_func=lambda: "")
    return app

def _json_data(resp: Any) -> Dict[str, Any]:
    return json.loads(resp.data.decode("utf-8") or "{}").get("data") or {}


def test_current_official_dispatch_surfaces_write_guardrail_in_page_data_and_excel(tmp_path: Path) -> None:
    conn = _seed_db(tmp_path)
    try:
        page = _page_context(conn)
        assert page["plan_identity"]["kind_label"] == "正式采用方案"
        assert page["plan_identity"]["dispatch_feedback_label"] == "可用于派工和现场反馈"
        assert "可以用于派工和现场反馈" in page["plan_identity"]["guardrail_text"]

        data = _dispatch_payload(conn, version=VERSION)
        identity = data["plan_identity"]
        assert identity["label"] == "正式采用方案"
        assert identity["can_dispatch"] is True
        assert identity["can_write_feedback"] is True
        assert "可以用于派工和现场反馈" in identity["guardrail_text"]

        summary = _summary_values(data)
        assert summary["计划身份"] == "正式采用方案；可用于派工和现场反馈"
        assert summary["派工反馈说明"] == "这套是当前可执行的正式采用方案，可以用于派工和现场反馈。"

        template_source = (REPO_ROOT / "templates/scheduler/resource_dispatch.html").read_text(encoding="utf-8")
        assert "plan_identity.guardrail_text" in template_source
        assert "计划身份" in template_source
        assert "派工反馈" in template_source
    finally:
        conn.close()


def test_history_comparison_and_scenario_plans_are_read_only_with_plain_reasons(tmp_path: Path) -> None:
    conn = _seed_db(tmp_path)
    try:
        history_data = _dispatch_payload(
            conn,
            version=HISTORICAL_VERSION,
            operator_id="O-ADOPTED",
            query_date="2026-05-01",
        )
        history_identity = history_data["plan_identity"]
        assert history_identity["kind_label"] == "历史正式方案"
        assert history_identity["can_dispatch"] is False
        assert history_identity["can_write_feedback"] is False
        assert "历史正式方案" in history_identity["guardrail_text"]
        assert "只能查看" in history_identity["guardrail_text"]

        baseline_data = _dispatch_payload(
            conn,
            version=VERSION,
            plan_role=ROLE_BASELINE_BEST,
            operator_id="O-CANDIDATE",
            query_date="2026-05-05",
        )
        baseline_identity = baseline_data["plan_identity"]
        assert baseline_identity["kind_label"] == "对比参考方案"
        assert baseline_identity["can_dispatch"] is False
        assert baseline_identity["can_write_feedback"] is False
        assert "不能确认派工或写现场反馈" in baseline_identity["guardrail_text"]

        critical_data = _dispatch_payload(
            conn,
            version=VERSION,
            plan_role=ROLE_CRITICAL_BEST,
            operator_id="O-ADOPTED",
            query_date="2026-05-05",
        )
        assert critical_data["plan_identity"]["kind_label"] == "对比参考方案"
        assert critical_data["plan_identity"]["can_write_feedback"] is False

        scenario_id = "scenario-guardrail"
        scenario_data = _dispatch_payload(
            conn,
            version=VERSION,
            scenario_id=scenario_id,
            operator_id="O-CANDIDATE",
            query_date="2026-05-06",
        )
        scenario_identity = scenario_data["plan_identity"]
        assert scenario_identity["label"] == "模拟预览（未命名）"
        assert scenario_identity["kind_label"] == "模拟预览"
        assert scenario_identity["can_dispatch"] is False
        assert scenario_identity["can_write_feedback"] is False
        assert "正式计划还没有改变" in scenario_identity["guardrail_text"]
        assert "不能确认派工或写现场反馈" in scenario_identity["guardrail_text"]
    finally:
        conn.close()


def test_public_payload_and_export_never_show_internal_plan_identity_fields(tmp_path: Path) -> None:
    conn = _seed_db(tmp_path)
    try:
        scenario_id = "scenario-guardrail"
        data = _dispatch_payload(
            conn,
            version=VERSION,
            scenario_id=scenario_id,
            operator_id="O-CANDIDATE",
            query_date="2026-05-06",
        )
        public_json = json.dumps(data, ensure_ascii=False)
        public_keys = _json_keys(json.loads(public_json))
        forbidden_keys = {
            "source_table",
            "candidate_id",
            "candidate_key",
            "scenario_id",
            "requested_plan_role",
            "effective_plan_role",
            "plan_role_status",
            "trace_meta",
        }
        for key in forbidden_keys:
            assert key not in public_keys
        for text in ("source_table", "candidate_rows", "candidate_id", "scenario_id", scenario_id):
            assert text not in public_json

        filename = build_resource_dispatch_filename(data)
        assert scenario_id not in unquote(filename)
        assert "模拟预览（未命名）" in filename

        workbook_text = _workbook_text(data)
        assert "模拟预览（未命名）" in workbook_text
        assert "正式计划还没有改变" in workbook_text
        for text in ("source_table", "candidate_rows", "candidate_id", "scenario_id", scenario_id):
            assert text not in workbook_text
    finally:
        conn.close()


def test_plan_identity_write_flags_cannot_be_overridden_by_outer_fields() -> None:
    plan_identity = {
        "user_label": "正式采用方案",
        "can_dispatch": False,
        "can_write_feedback": False,
        "is_official": True,
        "is_preview": False,
        "is_current_executable_version": True,
        "is_current_executable_official_version": False,
        "is_superseded_by_newer_version": True,
    }
    fields = plan_role_filter_fields({
        "requested_role": ROLE_ADOPTED,
        "selected_role": ROLE_ADOPTED,
        "source_table": SOURCE_SCHEDULE,
        "status": "resolved_adopted",
        "plan_identity": plan_identity,
        "can_dispatch": True,
        "can_write_feedback": True,
        "is_current_executable_official_version": True,
        "is_superseded_by_newer_version": False,
    })

    assert fields["can_dispatch"] is False
    assert fields["can_write_feedback"] is False
    assert fields["is_current_executable_official_version"] is False
    assert fields["is_superseded_by_newer_version"] is True


def test_resource_dispatch_get_data_and_export_do_not_write_schedule_or_execution_events(
    tmp_path: Path,
    monkeypatch,
) -> None:
    conn = _seed_db(tmp_path)
    app = _build_route_app()
    query = (
        "scope_type=operator&operator_id=O-ADOPTED&period_preset=week"
        f"&query_date=2026-05-05&version={VERSION}&plan_role={ROLE_ADOPTED}"
    )
    try:
        from web.routes.domains.scheduler import scheduler_resource_dispatch as rd_routes

        before_schedule = _count_rows(conn, "Schedule")
        before_execution_events = _count_rows(conn, "OperationExecutionEvents")

        def _capture_template(template_name: str, **kwargs: Any) -> str:
            assert template_name == "scheduler/resource_dispatch.html"
            assert "可以用于派工和现场反馈" in (kwargs.get("plan_identity") or {}).get("guardrail_text", "")
            return "OK"

        monkeypatch.setattr(rd_routes, "render_template", _capture_template)
        monkeypatch.setattr(rd_routes, "log_excel_export", lambda **_: None)

        with app.test_request_context(f"/scheduler/resource-dispatch?{query}"):
            g.services = SimpleNamespace(resource_dispatch_service=ResourceDispatchService(conn, logger=None, op_logger=None))
            assert rd_routes.resource_dispatch_page() == "OK"

        with app.test_request_context(f"/scheduler/resource-dispatch/data?{query}"):
            g.services = SimpleNamespace(resource_dispatch_service=ResourceDispatchService(conn, logger=None, op_logger=None))
            data_resp = rd_routes.resource_dispatch_data()
            assert data_resp.status_code == 200
            assert _json_data(data_resp)["plan_identity"]["can_write_feedback"] is True

        with app.test_request_context(f"/scheduler/resource-dispatch/export?{query}"):
            g.services = SimpleNamespace(resource_dispatch_service=ResourceDispatchService(conn, logger=None, op_logger=None))
            g.op_logger = None
            export_resp = rd_routes.resource_dispatch_export()
            assert export_resp.status_code == 200

        assert _count_rows(conn, "Schedule") == before_schedule
        if before_execution_events is not None:
            assert _count_rows(conn, "OperationExecutionEvents") == before_execution_events

        table_names = _table_names(conn)
        assert not (table_names & {"ResourceDispatchConfirmations", "DispatchConfirmations", "ConfirmedDispatches"})
        assert not any("DispatchConfirm" in name or "ConfirmDispatch" in name for name in table_names)

        routes_source = (REPO_ROOT / "web/routes/domains/scheduler/scheduler_resource_dispatch.py").read_text(
            encoding="utf-8"
        )
        assert "/resource-dispatch/confirm" not in routes_source
        assert "resource_dispatch_confirm" not in routes_source
    finally:
        conn.close()
