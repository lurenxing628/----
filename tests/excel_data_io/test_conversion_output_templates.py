"""回归测试：随仓库交付的 templates_excel/转换输出/ 示例文件，必须和 get_template_definition 的注册定义一致。

这是 2026-09 旧模板整体退役后，excel_templates 这一侧仅剩的文件级契约：
12 张业务表的模板改由工作台按表描述现生成（见 tests/workbench/test_table_descriptors.py），
templates_excel/ 里只留下回转壳体单元产品数据和转换输出的示例，后者仍是交付物，要锁住表头与下拉。
"""

from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, List, Mapping

from openpyxl import load_workbook

from core.services.common.excel_templates import get_template_definition
from tests._support.paths import REPO_ROOT_STR

CONVERSION_OUTPUT_DIR = Path(REPO_ROOT_STR) / "templates_excel" / "转换输出"


def _definition_enum_values_by_header(definition: Mapping[str, Any]) -> Dict[str, List[str]]:
    headers = [str(item) for item in (definition.get("headers") or [])]
    enum_cols = ((definition.get("format_spec") or {}).get("enum_cols") or {})
    return {headers[int(col)]: [str(item).strip() for item in (values or []) if str(item).strip()]
            for col, values in enum_cols.items()}


def _inline_validation_values(formula1: Any) -> List[str]:
    raw = str(formula1 or "").strip()
    if raw.startswith('"') and raw.endswith('"'):
        raw = raw[1:-1]
    return [item.strip() for item in raw.split(",") if item.strip()]


def _validation_value_lists_for_column(ws: Any, one_based_col_idx: int) -> List[List[str]]:
    out: List[List[str]] = []
    for data_validation in ws.data_validations.dataValidation:
        if data_validation.type != "list":
            continue
        for cell_range in data_validation.sqref.ranges:
            if (cell_range.min_col <= one_based_col_idx <= cell_range.max_col
                    and cell_range.min_row <= 2 <= cell_range.max_row):
                out.append(_inline_validation_values(data_validation.formula1))
                break
    return out


def test_supplier_conversion_output_matches_the_registered_definition() -> None:
    filename = "供应商配置.xlsx"
    path = CONVERSION_OUTPUT_DIR / filename
    assert path.exists(), f"缺少供应商配置转换输出示例：{path}"

    definition = get_template_definition(filename)
    expected_headers = [str(item) for item in definition.get("headers") or []]
    expected_enums = _definition_enum_values_by_header(definition)

    workbook = load_workbook(path, data_only=True)
    try:
        ws = workbook.active
        actual_headers = [ws.cell(1, col).value for col in range(1, ws.max_column + 1)]
        assert actual_headers == expected_headers, "供应商配置转换输出仍是旧列，必须补齐状态和备注"

        for col, header in enumerate(expected_headers, start=1):
            expected_values = expected_enums.get(header)
            if not expected_values:
                continue
            # 一列只能挂一条下拉规则：多挂一条旧规则，用户看到的选项就不是当前定义了。
            assert _validation_value_lists_for_column(ws, col) == [expected_values], (
                filename + " 的 " + header + " 下拉规则必须只有一条且只包含当前选项")

        rows = 0
        for row_idx in range(2, ws.max_row + 1):
            values = [ws.cell(row_idx, col).value for col in range(1, len(expected_headers) + 1)]
            if not any(value not in (None, "") for value in values):
                continue
            rows += 1
            assert values[4] in ("启用", "停用"), f"第 {row_idx} 行状态必须是启用或停用"
        assert rows > 0, "供应商配置转换输出至少要保留一行示例或转换结果"
    finally:
        workbook.close()
