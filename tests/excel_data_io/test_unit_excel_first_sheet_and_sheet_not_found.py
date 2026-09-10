"""守护单元 Excel 工作表选择契约（审计 D07/D10）：
1) 不指定 sheet 时必须解析第一个工作表，而不是 Excel 上次保存时停留的活动页（wb.active）；
2) 工作簿含多个工作表时必须出 multiple_sheets_notice 诊断，样本带所用 sheet 名与全部 sheet 名，
   让"解析的是哪张表"对用户可见；单工作表不出该诊断；
3) --sheet 指定不存在的 sheet 名时，转换链抛 SheetNotFoundError（带可用清单），
   CLI 脚本打印友好中文报错并列出可用 Sheet、退出码非 0、不产出任何文件。"""

from __future__ import annotations

import contextlib
import io
import os
import sys
import tempfile

import openpyxl
import pytest

_HEADERS = [
    "图号",
    "名称",
    "是否关重件",
    "关键特性",
    "工艺路线",
    "材料牌号",
    "材料规格",
    "进单元可装夹直径",
    "3140124 胡凡 罗辉",
    "换型时间(min)",
    "单件加工时间(min)",
    "批次加工时间(min)",
]


def _build_two_sheet_xlsx_active_on_notes(path: str) -> None:
    """第一个 sheet 是数据表，第二个是说明页，且保存时活动页停在说明页。"""
    wb = openpyxl.Workbook()
    try:
        ws = wb.active
        assert ws is not None
        ws.title = "单元产品信息统计"
        ws.append(_HEADERS)
        ws.append(["P001", "壳体A", "是", "", "1车", "", "", "", "1-1车削", 20, 15, 35])
        notes = wb.create_sheet("说明")
        notes.append(["这是说明页，不是产品数据。"])
        wb.active = 1  # 模拟用户最后停留在说明页保存
        wb.save(path)
    finally:
        try:
            wb.close()
        except Exception:
            pass


def _build_single_sheet_xlsx(path: str) -> None:
    wb = openpyxl.Workbook()
    try:
        ws = wb.active
        assert ws is not None
        ws.title = "单元产品信息统计"
        ws.append(_HEADERS)
        ws.append(["P001", "壳体A", "是", "", "1车", "", "", "", "1-1车削", 20, 15, 35])
        wb.save(path)
    finally:
        try:
            wb.close()
        except Exception:
            pass


def test_default_parses_first_sheet_even_when_active_sheet_is_notes_page() -> None:
    from core.services.process.unit_excel_converter import UnitExcelConverter

    tmpdir = tempfile.mkdtemp(prefix="aps_reg_unit_first_sheet_")
    src_xlsx = os.path.join(tmpdir, "source.xlsx")
    _build_two_sheet_xlsx_active_on_notes(src_xlsx)

    # 前置条件自证：文件的活动页确实停在说明页（否则本测试测不到 wb.active 陷阱）。
    wb = openpyxl.load_workbook(src_xlsx, read_only=True)
    try:
        active = wb.active
        assert active is not None and active.title == "说明"
    finally:
        wb.close()

    converted = UnitExcelConverter().convert(src_xlsx)

    # 解析的是第一个工作表：数据表里的零件被转换出来（老行为会读到说明页、产物为空）。
    assert [r.get("图号") for r in converted.routes_rows] == ["P001"], converted.routes_rows

    diagnostics = dict(getattr(converted, "diagnostics", {}) or {})
    counters = dict(diagnostics.get("counters") or {})
    samples = dict(diagnostics.get("samples") or {})
    assert int(counters.get("multiple_sheets_notice") or 0) == 1, diagnostics
    sample = dict((samples.get("multiple_sheets_notice") or [{}])[0])
    assert sample.get("used_sheet") == "单元产品信息统计", sample
    assert sample.get("sheet_names") == ["单元产品信息统计", "说明"], sample


def test_explicit_sheet_on_multi_sheet_workbook_still_records_notice() -> None:
    from core.services.process.unit_excel_converter import UnitExcelConverter

    tmpdir = tempfile.mkdtemp(prefix="aps_reg_unit_explicit_sheet_")
    src_xlsx = os.path.join(tmpdir, "source.xlsx")
    _build_two_sheet_xlsx_active_on_notes(src_xlsx)

    converted = UnitExcelConverter().convert(src_xlsx, sheet_name="单元产品信息统计")

    counters = dict((getattr(converted, "diagnostics", {}) or {}).get("counters") or {})
    assert int(counters.get("multiple_sheets_notice") or 0) == 1, counters
    assert [r.get("图号") for r in converted.routes_rows] == ["P001"], converted.routes_rows


def test_single_sheet_workbook_has_no_multiple_sheets_notice() -> None:
    from core.services.process.unit_excel_converter import UnitExcelConverter

    tmpdir = tempfile.mkdtemp(prefix="aps_reg_unit_single_sheet_")
    src_xlsx = os.path.join(tmpdir, "source.xlsx")
    _build_single_sheet_xlsx(src_xlsx)

    converted = UnitExcelConverter().convert(src_xlsx)
    counters = dict((getattr(converted, "diagnostics", {}) or {}).get("counters") or {})
    assert int(counters.get("multiple_sheets_notice") or 0) == 0, counters


def test_missing_sheet_raises_sheet_not_found_error_with_available_sheets() -> None:
    from core.services.process.unit_excel_converter import SheetNotFoundError, UnitExcelConverter

    tmpdir = tempfile.mkdtemp(prefix="aps_reg_unit_sheet_missing_")
    src_xlsx = os.path.join(tmpdir, "source.xlsx")
    _build_single_sheet_xlsx(src_xlsx)

    with pytest.raises(SheetNotFoundError) as exc_info:
        UnitExcelConverter().convert(src_xlsx, sheet_name="单元产品信息统计表")

    exc = exc_info.value
    assert exc.sheet_name == "单元产品信息统计表"
    assert exc.available_sheets == ["单元产品信息统计"]
    assert "找不到指定的工作表" in str(exc)
    assert "单元产品信息统计" in str(exc)


def test_convert_script_prints_friendly_error_for_missing_sheet() -> None:
    from scripts import convert_rotary_shell_unit_excel as convert_script

    tmpdir = tempfile.mkdtemp(prefix="aps_reg_unit_sheet_cli_")
    src_xlsx = os.path.join(tmpdir, "source.xlsx")
    out_dir = os.path.join(tmpdir, "out")
    _build_single_sheet_xlsx(src_xlsx)

    stdout = io.StringIO()
    old_argv = list(sys.argv)
    try:
        sys.argv = [
            "convert_rotary_shell_unit_excel.py",
            "--input",
            src_xlsx,
            "--output-dir",
            out_dir,
            "--sheet",
            "单元产品信息",
        ]
        with contextlib.redirect_stdout(stdout):
            rc = convert_script.main()
    finally:
        sys.argv = old_argv

    output = stdout.getvalue()
    assert rc == 2, output
    assert "错误：找不到指定的 Sheet：单元产品信息" in output, output
    assert "可用 Sheet：单元产品信息统计" in output, output
    assert "Traceback" not in output, output
    # 失败发生在产出之前：不留半套输出文件误导用户。
    assert not os.path.exists(out_dir), os.listdir(out_dir) if os.path.exists(out_dir) else None
