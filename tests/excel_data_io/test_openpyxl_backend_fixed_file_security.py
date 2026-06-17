"""回归测试：OpenpyxlBackend.write 写固定 Excel 文件时拒绝软链接/硬链接。"""

from __future__ import annotations

import os
from pathlib import Path

import pytest
from openpyxl import load_workbook

from core.infrastructure.errors import AppError
from core.services.common.openpyxl_backend import OpenpyxlBackend


def _skip_without_symlink(tmp_path: Path) -> None:
    if not hasattr(os, "symlink"):
        pytest.skip("平台不支持软链接")
    target = tmp_path / "_symlink_probe_target"
    target.write_text("x", encoding="utf-8")
    link = tmp_path / "_symlink_probe"
    try:
        os.symlink(str(target), str(link))
    except (OSError, NotImplementedError) as exc:
        pytest.skip(f"平台不允许创建软链接：{exc}")
    finally:
        try:
            link.unlink()
        except OSError:
            pass


def _skip_without_hardlink(tmp_path: Path) -> None:
    target = tmp_path / "_hardlink_probe_target"
    target.write_text("x", encoding="utf-8")
    link = tmp_path / "_hardlink_probe"
    try:
        os.link(str(target), str(link))
    except (OSError, NotImplementedError) as exc:
        pytest.skip(f"平台不允许创建硬链接：{exc}")
    finally:
        try:
            link.unlink()
        except OSError:
            pass


def test_openpyxl_backend_write_refuses_symlink_without_touching_target(tmp_path: Path) -> None:
    _skip_without_symlink(tmp_path)
    victim = tmp_path / "victim.xlsx"
    victim.write_text("VICTIM-UNCHANGED", encoding="utf-8")
    output_path = tmp_path / "report.xlsx"
    os.symlink(str(victim), str(output_path))

    with pytest.raises(AppError, match="写入 Excel 文件失败"):
        OpenpyxlBackend().write([{"列": "值"}], str(output_path))

    assert os.path.islink(str(output_path))
    assert victim.read_text(encoding="utf-8") == "VICTIM-UNCHANGED"


def test_openpyxl_backend_write_refuses_hardlink_without_touching_target(tmp_path: Path) -> None:
    _skip_without_hardlink(tmp_path)
    victim = tmp_path / "victim.xlsx"
    victim.write_text("VICTIM-UNCHANGED", encoding="utf-8")
    output_path = tmp_path / "report.xlsx"
    os.link(str(victim), str(output_path))

    with pytest.raises(AppError, match="写入 Excel 文件失败"):
        OpenpyxlBackend().write([{"列": "值"}], str(output_path))

    assert output_path.exists()
    assert victim.read_text(encoding="utf-8") == "VICTIM-UNCHANGED"


def test_openpyxl_backend_write_sanitizes_formula_like_values(tmp_path: Path) -> None:
    output_path = tmp_path / "report.xlsx"

    OpenpyxlBackend().write([{"列": "=1+1"}], str(output_path))

    workbook = load_workbook(output_path, data_only=False)
    try:
        ws = workbook.active
        assert ws["A2"].value == "'=1+1"
        assert ws["A2"].data_type != "f"
    finally:
        workbook.close()
