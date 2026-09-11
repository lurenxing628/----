"""班组旧页面拒绝丢条件跳转；原领域筛选及 Excel 班组列、名称和下载合同保留。"""

from __future__ import annotations

import importlib
import io
import os
import sys
from pathlib import Path
from typing import List

import pytest

from core.errors import BusinessError
from core.services.equipment.machine_service import MachineService
from core.services.personnel.operator_service import OperatorService
from tests._support.legacy_http import assert_retired_response, confirmation_inputs, xlsx_download_rows
from tests._support.paths import REPO_ROOT
from tests._support.sqlite_snapshot import table_rows


def _assert_status(resp, name: str, expect: int = 200) -> None:
    if resp.status_code != expect:
        body = resp.data.decode("utf-8", errors="ignore") if getattr(resp, "data", None) else ""
        raise AssertionError(f"{name} 返回 {resp.status_code}，期望 {expect}，body={body[:500]}")


def _xlsx_headers(content: bytes) -> List[str]:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(content))
    try:
        ws = wb.active
        if ws is None:
            return []
        return [str(ws.cell(1, idx).value or "").strip() for idx in range(1, ws.max_column + 1)]
    finally:
        wb.close()


def test_team_pages_and_excel_routes_show_team_columns_and_headers(tmp_path, monkeypatch) -> None:
    repo_root = str(REPO_ROOT)
    if repo_root not in sys.path:
        sys.path.insert(0, repo_root)

    test_db = tmp_path / "aps_test.db"
    test_logs = tmp_path / "logs"
    test_backups = tmp_path / "backups"
    test_templates = tmp_path / "templates_excel"
    test_logs.mkdir(parents=True, exist_ok=True)
    test_backups.mkdir(parents=True, exist_ok=True)
    test_templates.mkdir(parents=True, exist_ok=True)

    monkeypatch.setenv("APS_ENV", "development")
    monkeypatch.setenv("APS_DB_PATH", str(test_db))
    monkeypatch.setenv("APS_LOG_DIR", str(test_logs))
    monkeypatch.setenv("APS_BACKUP_DIR", str(test_backups))
    monkeypatch.setenv("APS_EXCEL_TEMPLATE_DIR", str(test_templates))

    from core.infrastructure.database import ensure_schema, get_connection

    ensure_schema(str(test_db), logger=None, schema_path=os.path.join(repo_root, "schema.sql"))

    conn = get_connection(str(test_db))
    try:
        conn.execute(
            "INSERT INTO ResourceTeams (team_id, name, status) VALUES (?, ?, ?)",
            ("TEAM-01", "车工一组", "active"),
        )
        conn.execute(
            "INSERT INTO Operators (operator_id, name, status, team_id, remark) VALUES (?, ?, ?, ?, ?)",
            ("OP001", "张三", "active", "TEAM-01", "人员备注"),
        )
        conn.execute(
            "INSERT INTO Machines (machine_id, name, status, team_id, remark) VALUES (?, ?, ?, ?, ?)",
            ("MC001", "数控车床1", "active", "TEAM-01", "设备备注"),
        )
        conn.commit()
        before = {table: table_rows(conn, table) for table in ("ResourceTeams", "Operators", "Machines")}
    finally:
        conn.close()

    sys.modules.pop("app", None)
    app_mod = importlib.import_module("app")
    app = app_mod.create_app()
    client = app.test_client()

    assert _xlsx_headers((test_templates / "人员基本信息.xlsx").read_bytes()) == ["工号", "姓名", "状态", "班组", "备注"]
    assert _xlsx_headers((test_templates / "设备信息.xlsx").read_bytes()) == ["设备编号", "设备名称", "工种", "班组", "状态"]

    resp_team_page = client.get("/personnel/teams")
    assert_retired_response(resp_team_page)

    resp_personnel = client.get("/personnel/?team_id=TEAM-01")
    assert_retired_response(resp_personnel)

    resp_equipment = client.get("/equipment/?team_id=TEAM-01")
    assert_retired_response(resp_equipment)

    resp_personnel_invalid = client.get("/personnel/?team_id=TEAM-404")
    assert_retired_response(resp_personnel_invalid)

    resp_equipment_invalid = client.get("/equipment/?team_id=TEAM-404")
    assert_retired_response(resp_equipment_invalid)
    conn = get_connection(str(test_db))
    try:
        assert [row.operator_id for row in OperatorService(conn).list(team_id="TEAM-01")] == ["OP001"]
        assert [row.machine_id for row in MachineService(conn).list(team_id="TEAM-01")] == ["MC001"]
        for service in (OperatorService(conn), MachineService(conn)):
            with pytest.raises(BusinessError, match="TEAM-404"):
                service.list(team_id="TEAM-404")
    finally:
        conn.close()

    resp_operator_excel = client.get("/personnel/excel/operators")
    assert_retired_response(resp_operator_excel)

    resp_machine_excel = client.get("/equipment/excel/machines")
    assert_retired_response(resp_machine_excel)

    resp_operator_template = client.get("/personnel/excel/operators/template")
    _assert_status(resp_operator_template, "GET /personnel/excel/operators/template")
    assert _xlsx_headers(resp_operator_template.data) == ["工号", "姓名", "状态", "班组", "备注"]

    resp_machine_template = client.get("/equipment/excel/machines/template")
    _assert_status(resp_machine_template, "GET /equipment/excel/machines/template")
    assert _xlsx_headers(resp_machine_template.data) == ["设备编号", "设备名称", "工种", "班组", "状态"]

    resp_operator_export = client.get("/personnel/excel/operators/export")
    _assert_status(resp_operator_export, "GET /personnel/excel/operators/export")
    assert _xlsx_headers(resp_operator_export.data) == ["工号", "姓名", "状态", "班组", "备注"]
    operator_rows = xlsx_download_rows(resp_operator_export)
    assert len(operator_rows) == 1
    assert operator_rows[0]["工号"] == "OP001" and operator_rows[0]["班组"] == "车工一组"

    resp_machine_export = client.get("/equipment/excel/machines/export")
    _assert_status(resp_machine_export, "GET /equipment/excel/machines/export")
    assert _xlsx_headers(resp_machine_export.data) == ["设备编号", "设备名称", "工种", "班组", "状态"]
    machine_rows = xlsx_download_rows(resp_machine_export)
    assert len(machine_rows) == 1
    assert machine_rows[0]["设备编号"] == "MC001" and machine_rows[0]["班组"] == "车工一组"
    for path, response, filename in (
        ("/personnel/excel/operators", resp_operator_export, "operators.xlsx"),
        ("/equipment/excel/machines", resp_machine_export, "machines.xlsx"),
    ):
        preview = client.post(path + "/preview", data={
            "mode": "overwrite", "file": (io.BytesIO(response.data), filename),
        }, content_type="multipart/form-data")
        _assert_status(preview, "POST " + path + "/preview")
        body = preview.get_data(as_text=True)
        fields = confirmation_inputs(body, path + "/confirm")
        assert fields["mode"] == "overwrite" and "班组" in body and "车工一组" in body

    (test_templates / "人员基本信息.xlsx").unlink()
    resp_excel_demo_template = client.get("/excel-demo/template")
    _assert_status(resp_excel_demo_template, "GET /excel-demo/template")
    assert _xlsx_headers(resp_excel_demo_template.data) == ["工号", "姓名", "状态", "班组", "备注"]
    conn = get_connection(str(test_db))
    try:
        assert before == {table: table_rows(conn, table) for table in before}
    finally:
        conn.close()
