"""回归测试：资源排班对外输出契约——decorate_resource_dispatch_payload 与 build_resource_dispatch_workbook 必须把 source/lock_status 翻成中文公开标签（自制/外协/已锁定…）、不外泄 internal/external/locked 等内部枚举和 op_id/source_table 等内部字段，build_resource_dispatch_filename 须清洗 scope_id 中的非法文件名字符，且坏时间过滤用大白话标签呈现。"""

from __future__ import annotations

import io
from pathlib import Path
from typing import List

import openpyxl

from core.services.common.excel_templates import _sanitize_export_cell
from core.services.scheduler.resource_dispatch_excel import build_resource_dispatch_workbook
from web.viewmodels.scheduler_resource_dispatch import (
    build_resource_dispatch_filename,
    decorate_resource_dispatch_payload,
)

REPO_ROOT = Path(__file__).resolve().parents[1]


def _base_payload() -> dict:
    return {
        "filters": {
            "scope_type": "operator",
            "scope_id": 'OP/\\:*?"<>|001\r\n\t\x00',
            "scope_name": "张三",
            "period_preset": "week",
            "start_date": "2026-05-04",
            "end_date": "2026-05-10",
            "version": 7,
        },
        "summary": {
            "total_tasks": 3,
            "total_hours": 6,
            "cross_day_count": 0,
            "overdue_count": 0,
            "external_count": 1,
            "cross_team_count": 0,
            "degraded": False,
            "degradation_counters": {},
            "degradation_events": [],
        },
        "calendar_headers": ["2026-05-04"],
        "calendar_rows": [],
    }


def _detail_row(source: str, lock_status: str, suffix: str) -> dict:
    return {
        "schedule_id": suffix,
        "op_id": f"OPID-{suffix}",
        "op_code": f"OP{suffix}",
        "batch_id": f"B{suffix}",
        "part_no": f"P{suffix}",
        "seq": 10,
        "op_type_name": "车削",
        "start_time": "2026-05-04 08:00:00",
        "end_time": "2026-05-04 10:00:00",
        "duration_minutes": 120,
        "scope_type": "operator",
        "scope_id": f"OP{suffix}",
        "scope_name": "张三",
        "operator_id": f"OP{suffix}",
        "operator_name": "张三",
        "machine_id": f"MC{suffix}",
        "machine_name": "设备1",
        "source": source,
        "lock_status": lock_status,
        "is_cross_day": False,
        "is_overdue": False,
    }


def _workbook_cell_values(wb) -> List[str]:
    values: List[str] = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                values.append(str(cell.value or ""))
    return values


def test_resource_dispatch_payload_decorates_public_source_and_lock_labels_without_mutation() -> None:
    payload = _base_payload()
    payload["detail_rows"] = [
        _detail_row("internal", "locked", "001"),
        _detail_row("external", "unlocked", "002"),
        _detail_row("future_source", "future_lock", "003"),
    ]

    out = decorate_resource_dispatch_payload(payload)

    assert "source_label" not in payload["detail_rows"][0]
    assert out["detail_rows"][0]["source_label"] == "自制"
    assert out["detail_rows"][0]["lock_status_label"] == "已锁定"
    assert out["detail_rows"][1]["source_label"] == "外协"
    assert out["detail_rows"][1]["lock_status_label"] == "未锁定"
    assert out["detail_rows"][2]["source_label"] == "来源未识别"
    assert out["detail_rows"][2]["lock_status_label"] == "锁定状态未识别"


def test_resource_dispatch_excel_uses_public_labels_and_never_leaks_internal_enums() -> None:
    payload = _base_payload()
    payload["detail_rows"] = [
        _detail_row("internal", "locked", "001"),
        _detail_row("external", "unlocked", "002"),
        _detail_row("future_source", "future_lock", "003"),
    ]
    decorated = decorate_resource_dispatch_payload(payload)
    buffer = build_resource_dispatch_workbook(decorated)
    wb = openpyxl.load_workbook(io.BytesIO(buffer.getvalue()))
    ws = wb["任务明细"]
    headers = [ws.cell(1, idx).value for idx in range(1, ws.max_column + 1)]
    source_col = headers.index("来源") + 1
    lock_col = headers.index("锁定状态") + 1

    assert ws.cell(2, source_col).value == "自制"
    assert ws.cell(2, lock_col).value == "已锁定"
    assert ws.cell(3, source_col).value == "外协"
    assert ws.cell(3, lock_col).value == "未锁定"
    assert ws.cell(4, source_col).value == "来源未识别"
    assert ws.cell(4, lock_col).value == "锁定状态未识别"

    all_values = "\n".join(_workbook_cell_values(wb))
    for forbidden in ("internal", "external", "locked", "unlocked", "future_source", "future_lock"):
        assert forbidden not in all_values


def test_resource_dispatch_excel_never_uses_internal_op_id_when_op_code_missing() -> None:
    payload = _base_payload()
    row = _detail_row("internal", "locked", "001")
    row["op_id"] = "INTERNAL-OP-ID-987654"
    row["op_code"] = ""
    payload["detail_rows"] = [row]

    decorated = decorate_resource_dispatch_payload(payload)
    buffer = build_resource_dispatch_workbook(decorated)
    wb = openpyxl.load_workbook(io.BytesIO(buffer.getvalue()))
    all_values = "\n".join(_workbook_cell_values(wb))

    assert "INTERNAL-OP-ID-987654" not in all_values
    assert "op_INTERNAL" not in all_values
    assert "OPID-987654" not in all_values


def test_resource_dispatch_team_excel_uses_public_labels_in_all_detail_sheets() -> None:
    payload = _base_payload()
    payload["filters"]["scope_type"] = "team"
    payload["operator_rows"] = [_detail_row("internal", "locked", "team-op")]
    payload["machine_rows"] = [_detail_row("external", "unlocked", "team-machine")]
    payload["cross_team_rows"] = [_detail_row("future_source", "future_lock", "team-cross")]
    decorated = decorate_resource_dispatch_payload(payload)
    buffer = build_resource_dispatch_workbook(decorated)
    wb = openpyxl.load_workbook(io.BytesIO(buffer.getvalue()))

    assert {"班组人员任务明细", "班组设备任务明细", "跨班组"} <= set(wb.sheetnames)
    all_values = "\n".join(_workbook_cell_values(wb))
    for expected in ("自制", "外协", "来源未识别", "已锁定", "未锁定", "锁定状态未识别"):
        assert expected in all_values
    for forbidden in ("internal", "external", "locked", "unlocked", "future_source", "future_lock"):
        assert forbidden not in all_values


def test_resource_dispatch_export_filename_sanitizes_scope_id_for_public_downloads() -> None:
    filename = build_resource_dispatch_filename(decorate_resource_dispatch_payload(_base_payload()))

    assert filename.startswith("资源排班_人员_OP----001_2026-05-04_2026-05-10_v7")
    for forbidden in ('/', "\\", ":", "*", "?", '"', "<", ">", "|", "\r", "\n", "\t", "\x00"):
        assert forbidden not in filename


def test_resource_dispatch_excel_cell_sanitizer_removes_illegal_control_characters() -> None:
    assert _sanitize_export_cell('=公式\x00\x08') == "'=公式"
    assert _sanitize_export_cell("普通\n换行") == "普通\n换行"


def test_resource_dispatch_excel_summary_uses_plain_bad_time_label() -> None:
    payload = _base_payload()
    payload["summary"]["degradation_counters"] = {"bad_time_row_skipped": 2}
    buffer = build_resource_dispatch_workbook(decorate_resource_dispatch_payload(payload))
    wb = openpyxl.load_workbook(io.BytesIO(buffer.getvalue()))
    ws = wb["查询摘要"]
    labels = [str(ws.cell(row, 1).value or "") for row in range(1, ws.max_row + 1)]

    assert "开始或结束时间写法不对，已过滤的记录数" in labels
    assert "坏时间过滤数量" not in labels
    assert "bad_time_row_skipped" not in labels


def test_resource_dispatch_excel_adds_calendar_detail_sheet_from_structured_items() -> None:
    payload = _base_payload()
    payload["calendar_rows"] = [
        {
            "scope_type": "operator",
            "scope_id": "OP001",
            "scope_name": "张三",
            "current_resource_id": "OP001",
            "current_resource_name": "张三",
            "operator_id": "OP001",
            "operator_name": "张三",
            "cells": [
                {
                    "date": "2026-05-04",
                    "items": [
                        {
                            "schedule_id": "S-INTERNAL-1",
                            "op_id": "OP-INTERNAL-1",
                            "_row_identity": "ROW-INTERNAL-1",
                            "source_table": "ScheduleSecret",
                            "state_revision": "STATE-INTERNAL-1",
                            "execution_snapshot_revision": "SNAPSHOT-INTERNAL-1",
                            "start": "2026-05-04 08:00:00",
                            "end": "2026-05-04 10:00:00",
                            "time_label": "08:00-10:00",
                            "scope_type": "operator",
                            "batch_id": "B001",
                            "op_code": "OP10",
                            "seq": 10,
                            "part_no": "P001",
                            "part_name": "回转壳体",
                            "machine_id": "MC001",
                            "machine_name": "数控车床1",
                            "operator_id": "OP001",
                            "operator_name": "张三",
                            "is_overdue": True,
                        }
                    ],
                }
            ],
        }
    ]

    buffer = build_resource_dispatch_workbook(decorate_resource_dispatch_payload(payload))
    wb = openpyxl.load_workbook(io.BytesIO(buffer.getvalue()))

    assert "日历排班" in wb.sheetnames
    assert "日历明细" in wb.sheetnames
    ws = wb["日历明细"]
    headers = [ws.cell(1, idx).value for idx in range(1, ws.max_column + 1)]
    assert headers == [
        "日历来源",
        "查询对象",
        "日期",
        "时间",
        "批次",
        "工序",
        "图号 / 物料",
        "计划设备",
        "计划人员",
        "对应资源",
        "是否超期",
        "任务说明",
    ]
    assert ws.max_row == 2
    row = [ws.cell(2, idx).value for idx in range(1, ws.max_column + 1)]
    assert row[:9] == [
        "日历排班",
        "张三\n完整身份：OP001 张三",
        "2026-05-04",
        "08:00-10:00",
        "B001",
        "OP10",
        "P001 回转壳体",
        "MC001 数控车床1",
        "OP001 张三",
    ]
    assert row[10] == "是"
    assert "08:00-10:00 OP10 数控车床1 P001" in str(row[11] or "")

    all_values = "\n".join(_workbook_cell_values(wb))
    for forbidden in (
        "S-INTERNAL-1",
        "OP-INTERNAL-1",
        "ROW-INTERNAL-1",
        "ScheduleSecret",
        "STATE-INTERNAL-1",
        "SNAPSHOT-INTERNAL-1",
        "schedule_id",
        "op_id",
        "_row_identity",
        "source_table",
        "state_revision",
        "execution_snapshot_revision",
    ):
        assert forbidden not in all_values


def test_resource_dispatch_team_calendar_detail_sheet_marks_calendar_source() -> None:
    payload = _base_payload()
    payload["filters"]["scope_type"] = "team"
    calendar_item = {
        "time_label": "08:00-10:00",
        "scope_type": "operator",
        "batch_id": "B001",
        "op_code": "OP10",
        "part_no": "P001",
        "machine_id": "MC001",
        "machine_name": "数控车床1",
        "operator_id": "OP001",
        "operator_name": "张三",
        "is_overdue": False,
    }
    calendar_row = {
        "scope_type": "operator",
        "scope_id": "OP001",
        "scope_name": "张三",
        "current_resource_id": "OP001",
        "current_resource_name": "张三",
        "operator_id": "OP001",
        "operator_name": "张三",
        "cells": [{"date": "2026-05-04", "items": [calendar_item]}],
    }
    payload["operator_calendar_rows"] = [calendar_row]
    payload["machine_calendar_rows"] = [dict(calendar_row, current_resource_id="MC001", current_resource_name="数控车床1")]

    buffer = build_resource_dispatch_workbook(decorate_resource_dispatch_payload(payload))
    wb = openpyxl.load_workbook(io.BytesIO(buffer.getvalue()))
    ws = wb["日历明细"]
    source_values = [str(ws.cell(row, 1).value or "") for row in range(2, ws.max_row + 1)]

    assert source_values == ["班组人员日历", "班组设备日历"]


def test_resource_dispatch_calendar_detail_sheet_handles_empty_calendar_rows() -> None:
    payload = _base_payload()
    buffer = build_resource_dispatch_workbook(decorate_resource_dispatch_payload(payload))
    wb = openpyxl.load_workbook(io.BytesIO(buffer.getvalue()))
    ws = wb["日历明细"]

    assert ws.max_row == 1
    assert [ws.cell(1, idx).value for idx in range(1, ws.max_column + 1)][0] == "日历来源"


def test_resource_dispatch_template_version_summary_prefers_public_schedule_time() -> None:
    template = (REPO_ROOT / "templates" / "scheduler" / "resource_dispatch.html").read_text(encoding="utf-8")

    assert "schedule_time_display or selected_version_row.item.schedule_time" in template
    assert "ui.summary_item('时间', selected_version_row.item.schedule_time)" not in template
