"""Report numeric contracts and XLSX execution survive leaf extraction."""

from __future__ import annotations

from io import BytesIO

import openpyxl
import pytest

from core.infrastructure.errors import ValidationError
from core.services.report.exporters import xlsx
from core.services.report.values import number_parsing


def test_report_numeric_values_and_error_contracts():
    args = {"field": "amount", "label": "amount", "source_label": "rows"}
    assert number_parsing.parse_optional_report_float(None, **args) is None
    assert number_parsing.parse_optional_report_float("  ", **args) is None
    assert number_parsing.parse_optional_report_float("2.5", **args) == 2.5
    assert number_parsing.parse_report_float("", blank_default=3.0, **args) == 3.0
    assert number_parsing.parse_report_int("2", **args) == 2
    assert number_parsing.parse_report_nonnegative_int(None, **args) == 0
    for raw in ("bad", True, "1.0", -1):
        with pytest.raises(ValidationError) as error:
            number_parsing.parse_report_nonnegative_int(raw, **args)
        assert error.value.field == "amount"
        assert error.value.message == "amount不是有效整数，请检查rows。"


@pytest.mark.parametrize("write_only", (False, True))
def test_export_preserves_values_and_real_parser_patch(monkeypatch, write_only):
    calls = []
    parser = xlsx.parse_optional_report_float

    def observed(value, **kwargs):
        calls.append((value, kwargs))
        return parser(value, **kwargs)

    monkeypatch.setattr(xlsx, "parse_optional_report_float", observed)
    row = {"machine_id": "M1", "machine_name": "Machine", "hours": 1.25, "task_count": 2, "capacity_hours": 4, "utilization": "0.3125"}
    buf = xlsx.export_utilization_xlsx([row], [], write_only=write_only)
    try:
        wb = openpyxl.load_workbook(BytesIO(buf.read()), read_only=True, data_only=True)
        try:
            assert wb.sheetnames == ["设备负荷", "人员负荷"]
            rows = list(wb.worksheets[0].iter_rows(values_only=True))
            assert rows[1] == ("M1", "Machine", 1.25, 2, 4, 31.25, None, None, None, None, None)
        finally:
            wb.close()
    finally:
        buf.close()
    assert calls == [("0.3125", {"field": "utilization", "label": "利用率", "source_label": "资源负荷导出数据"})]


