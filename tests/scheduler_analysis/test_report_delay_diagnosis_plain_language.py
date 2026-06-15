"""回归测试：超期报表页 /reports/overdue 及其导出（含模拟方案预览 scenario_id 入参）把延期诊断以大白话中文呈现（查看为什么晚了/建议先复核/证据等级/证据缺口/下一步），不下结论式归因（不出现"物料不够导致延期"），且不向用户泄露 trace_meta、rule_version、source_table、scenario_id、根因 等内部术语；导出"诊断依据"sheet 表头固定为七列。"""

from __future__ import annotations

import os
import sys
from io import BytesIO
from pathlib import Path
from typing import Iterable
from urllib.parse import parse_qs, unquote, urlsplit

from flask import Flask
from openpyxl import load_workbook

from core.infrastructure.database import ensure_schema, get_connection
from tests._support.paths import REPO_ROOT

TESTS_ROOT = REPO_ROOT / "tests"
if str(TESTS_ROOT) not in sys.path:
    sys.path.insert(0, str(TESTS_ROOT))

from core.models.schedule_plan_identity import PlanIdentity  # noqa: E402
from core.services.scheduler.schedule_delay_diagnosis_utils import plan_link  # noqa: E402
from tests._support.excel_templates import point_env_at_shared  # noqa: E402
from tests.scheduler_analysis.test_scheduler_delay_diagnosis_contract import (  # noqa: E402
    VERSION,
    _seed_base,
    _seed_candidates,
    _seed_scenario,
)
from web.routes.domains.scheduler.scheduler_plan_context_token import (  # noqa: E402
    plan_context_token,
    scenario_id_from_plan_context_token,
)
from web.viewmodels.scheduler_reports_workbench import (  # noqa: E402
    build_report_context,
    decorate_delay_diagnosis_context,
)

INTERNAL_TERMS = (
    "trace_meta",
    "rule_version",
    "input_fingerprint",
    "source_table",
    "scenario_id",
    "primary_reason",
    "critical_chain",
    "根因",
    "主要原因",
)


def _build_app(tmp_path: Path, monkeypatch):
    db_path = tmp_path / "aps.db"
    monkeypatch.setenv("APS_ENV", "development")
    monkeypatch.setenv("APS_DB_PATH", str(db_path))
    monkeypatch.setenv("APS_LOG_DIR", str(tmp_path / "logs"))
    monkeypatch.setenv("APS_BACKUP_DIR", str(tmp_path / "backups"))
    point_env_at_shared(monkeypatch)
    ensure_schema(str(db_path), logger=None, schema_path=str(REPO_ROOT / "schema.sql"), backup_dir=None)
    conn = get_connection(str(db_path))
    try:
        _seed_base(conn)
        _seed_candidates(conn)
        scenario_id = _seed_scenario(conn)
        conn.commit()
    finally:
        conn.close()

    for name in ("app", "web.routes.reports", "web.routes.report_plan_preview"):
        sys.modules.pop(name, None)
    import app as app_mod  # noqa: E402

    return app_mod.create_app(), scenario_id


def _workbook_texts(xlsx_bytes: bytes) -> Iterable[str]:
    wb = load_workbook(BytesIO(xlsx_bytes), data_only=True)
    try:
        for ws in wb.worksheets:
            yield ws.title
            for row in ws.iter_rows(values_only=True):
                for value in row:
                    if value is not None:
                        yield str(value)
    finally:
        wb.close()


def _assert_no_internal_terms(text: str) -> None:
    for term in INTERNAL_TERMS:
        assert term not in text


def _insert_bad_finish_time_batch() -> None:
    db_path = Path(os.environ["APS_DB_PATH"])
    conn = get_connection(str(db_path))
    try:
        conn.executescript(
            """
            INSERT INTO Batches(batch_id, part_no, part_name, quantity, due_date, priority, ready_status, status)
            VALUES ('B_BAD_TIME', 'P001', '坏时间批次', 1, '2026-05-03', 'normal', 'yes', 'scheduled');

            INSERT INTO BatchOperations(id, op_code, batch_id, piece_id, seq, op_type_name, source, status)
            VALUES (50, 'OP-BAD-10', 'B_BAD_TIME', 'piece-bad', 10, '车削', 'internal', 'scheduled');

            INSERT INTO Schedule(id, op_id, machine_id, operator_id, start_time, end_time, lock_status, version)
            VALUES (150, 50, 'M1', 'O1', '2026-05-04 10:00:00', '坏结束时间', 'unlocked', 11);
            """
        )
        conn.commit()
    finally:
        conn.close()


def test_overdue_page_shows_delay_diagnosis_in_plain_chinese(tmp_path: Path, monkeypatch) -> None:
    app, _scenario_id = _build_app(tmp_path, monkeypatch)
    client = app.test_client()

    response = client.get(f"/reports/overdue?version={VERSION}&plan_role=adopted")
    assert response.status_code == 200
    html = response.get_data(as_text=True)

    assert "查看为什么晚了" in html
    assert "建议先复核" in html
    assert "证据等级" in html
    assert "证据缺口" in html
    assert "下一步" in html
    assert "没有现场执行反馈时，不判断现场做慢了" in html
    assert "缺少批次物料明细" in html
    assert "当前物料数据不足，只提示核对物料，不判断一定是物料造成延期" in html
    assert "物料不够导致延期" not in html
    _assert_no_internal_terms(html)


def test_overdue_page_and_export_show_bad_finish_time_as_data_issue(tmp_path: Path, monkeypatch) -> None:
    app, _scenario_id = _build_app(tmp_path, monkeypatch)
    _insert_bad_finish_time_batch()
    client = app.test_client()

    page_response = client.get(f"/reports/overdue?version={VERSION}&plan_role=adopted")
    assert page_response.status_code == 200
    html = page_response.get_data(as_text=True)

    assert "B_BAD_TIME" in html
    assert "排程时间异常" in html
    assert "计划完成时间写法不对" in html
    assert "有排程记录，但计划完成时间写法不对" in html

    export_response = client.get(f"/reports/overdue/export?version={VERSION}&plan_role=adopted")
    assert export_response.status_code == 200
    values = "\n".join(_workbook_texts(export_response.data))

    assert "B_BAD_TIME" in values
    assert "排程时间异常" in values
    assert "计划完成时间写法不对" in values
    assert "坏结束时间" not in values
    _assert_no_internal_terms(values)


def test_overdue_export_contains_trace_sheet_without_internal_terms(tmp_path: Path, monkeypatch) -> None:
    app, _scenario_id = _build_app(tmp_path, monkeypatch)
    client = app.test_client()

    response = client.get(f"/reports/overdue/export?version={VERSION}&plan_role=adopted")
    assert response.status_code == 200
    wb = load_workbook(BytesIO(response.data), data_only=True)
    try:
        assert "超期清单" in wb.sheetnames
        assert "诊断依据" in wb.sheetnames
        ws = wb["诊断依据"]
        headers = [ws.cell(row=1, column=i).value for i in range(1, 8)]
        assert headers == ["本次诊断编号", "生成依据摘要", "核对信息", "证据来源", "证据缺口", "生成时间", "筛选条件"]
        values = "\n".join(_workbook_texts(response.data))
    finally:
        wb.close()

    assert "延期诊断-v" in values
    assert "建议先复核" in values
    assert "证据等级" in values
    assert "缺少批次物料明细" in values
    _assert_no_internal_terms(values)


def test_overdue_export_supports_scenario_preview_without_internal_terms(tmp_path: Path, monkeypatch) -> None:
    app, scenario_id = _build_app(tmp_path, monkeypatch)
    client = app.test_client()

    response = client.get(f"/reports/overdue/export?version={VERSION}&plan_role=adopted&scenario_id={scenario_id}")
    assert response.status_code == 200
    disposition = unquote(str(response.headers.get("Content-Disposition") or ""))
    values = "\n".join(_workbook_texts(response.data))
    assert "延期诊断预览" in disposition
    assert scenario_id not in disposition
    assert scenario_id not in values
    _assert_no_internal_terms(values)


def test_delay_diagnosis_plan_link_does_not_expose_raw_scenario_id() -> None:
    identity = PlanIdentity(
        version=12,
        requested_plan_role="adopted",
        effective_plan_role="adopted",
        plan_resolution_status="resolved_scenario",
        source_table="adjustment_scenario_rows",
        source_row_id=None,
        candidate_id=None,
        candidate_key=None,
        scenario_id="SC-SECRET",
        schedule_result_status="success",
        result_summary_parse_failed=False,
        result_summary_parse_reason="",
        is_simulation=True,
        label="模拟预览",
        user_label="模拟预览",
        is_official=False,
        is_preview=True,
        is_current_executable_version=False,
        is_current_executable_official_version=False,
        is_superseded_by_newer_version=False,
        schedule_lock_status=None,
        can_dispatch=False,
        can_write_feedback=False,
        detail_saved=True,
    )

    link = plan_link(identity, "/scheduler/gantt", scenario_id="SC-SECRET")

    assert "version=12" in link
    assert "plan_role=adopted" in link
    assert "scenario_id" not in link
    assert "SC-SECRET" not in link


def test_delay_diagnosis_public_action_link_preserves_preview_context_as_token() -> None:
    app = Flask(__name__)
    app.secret_key = "delay-diagnosis-public-token"
    scenario_id = "SC-SECRET"
    with app.app_context():
        context = build_report_context(
            version=12,
            plan_resolution={
                "requested_role": "adopted",
                "selected_role": "adopted",
                "scenario_id": scenario_id,
                "scenario_display_name": "预览方案",
                "is_scenario_preview": True,
                "is_preview": True,
                "can_write_feedback": False,
                "can_dispatch": False,
            },
            plan_context_token=plan_context_token(scenario_id),
            date_from="2026-05-01",
            date_to="2026-05-07",
        )
        decorated = decorate_delay_diagnosis_context(
            {
                "items_by_batch": {
                    "B1": {
                        "batch_id": "B1",
                        "suggested_actions": [
                            {
                                "label": "查看甘特图",
                                "reason": "先核对计划里这批次排到了哪里。",
                                "link": "/scheduler/gantt?version=12&plan_role=adopted",
                            }
                        ],
                    }
                }
            },
            context,
        )

        link = decorated["items_by_batch"]["B1"]["suggested_actions"][0]["link"]
        query = parse_qs(urlsplit(link).query)
        token = query.get("plan_context_token", [""])[0]

        assert link.startswith("/scheduler/gantt?")
        assert "scenario_id" not in link
        assert scenario_id not in link
        assert token
        assert scenario_id_from_plan_context_token(token) == scenario_id
