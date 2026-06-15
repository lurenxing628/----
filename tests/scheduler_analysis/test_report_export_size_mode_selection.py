"""回归测试：报表导出（/reports/overdue/export、/reports/utilization/export）按行数在 EXPORT_DIRECT_MAX_ROWS/EXPORT_STREAM_MAX_ROWS 间选 direct 与 stream 模式，并在响应头回写 X-APS-Report-Export-Mode/Estimated-Rows；同时校验超期清单表头与列序、利用率以百分比导出、候选方案 stream 导出文件名用中文方案名且不泄露 plan_role/candidate_id 等内部字段。"""

from __future__ import annotations

from io import BytesIO
from typing import Any, Dict, List
from urllib.parse import unquote


def _assert_xlsx(resp, name: str, expect_version: int) -> None:
    if resp.status_code != 200:
        body = resp.data.decode("utf-8") if getattr(resp, "data", None) else ""
        raise RuntimeError(f"{name} 返回 {resp.status_code}，期望 200，body={body[:500]}")
    ct = resp.headers.get("Content-Type", "")
    if "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" not in ct:
        raise RuntimeError(f"{name} content-type 异常：{ct}")
    cd = resp.headers.get("Content-Disposition", "")
    if f"v{int(expect_version)}" not in cd:
        raise RuntimeError(f"{name} filename 未包含 v{expect_version}（Content-Disposition={cd!r}）")


def _assert_overdue_xlsx_columns(resp) -> None:
    import openpyxl

    wb = openpyxl.load_workbook(BytesIO(resp.data), data_only=True)
    try:
        ws = wb["超期清单"]
        headers = [ws.cell(row=1, column=i).value for i in range(1, 10)]
        if headers != ["类别", "批次号", "图号", "名称", "数量", "交期", "完工/截至时间", "超期(天)", "超期(小时)"]:
            raise RuntimeError(f"超期清单表头顺序异常：{headers!r}")
        row_tail = [ws.cell(row=2, column=8).value, ws.cell(row=2, column=9).value]
        if row_tail != [0.33, 8.0]:
            raise RuntimeError(f"超期清单最后两列数据顺序异常：{row_tail!r}")
    finally:
        wb.close()


def _assert_utilization_xlsx_percent(resp) -> None:
    import openpyxl

    wb = openpyxl.load_workbook(BytesIO(resp.data), data_only=True)
    try:
        ws = wb["设备负荷"]
        if ws["F1"].value != "利用率(%)":
            raise RuntimeError(f"利用率表头异常：{ws['F1'].value!r}")
        if ws["F2"].value != 50.0:
            raise RuntimeError(f"利用率应以百分比数值导出：{ws['F2'].value!r}")
    finally:
        wb.close()


def _assert_no_internal_terms_in_xlsx(resp) -> None:
    import openpyxl

    forbidden_terms = (
        "plan_role",
        "candidate_key",
        "candidate_id",
        "source_table",
        "baseline_best",
        "critical_best",
        "scenario_id",
        "candidate_rows",
    )
    wb = openpyxl.load_workbook(BytesIO(resp.data), data_only=True)
    values: List[str] = []
    try:
        for ws in wb.worksheets:
            values.append(str(ws.title or ""))
            for row in ws.iter_rows(values_only=True):
                values.extend(str(cell or "") for cell in row)
    finally:
        wb.close()
    text = resp.headers.get("Content-Disposition", "") + "\n" + "\n".join(values)
    leaked = [term for term in forbidden_terms if term in text]
    if leaked:
        raise RuntimeError(f"候选方案 stream 导出泄露内部字段：{leaked!r}")


def _make_overdue_items(count: int) -> List[Dict[str, Any]]:
    return [
        {
            "bucket_label": "已排程超期",
            "batch_id": f"B{i:03d}",
            "part_no": f"P{i:03d}",
            "part_name": f"零件{i}",
            "quantity": 1,
            "due_date": "2026-01-01",
            "finish_time": "2026-01-02 08:00:00",
            "as_of_time": None,
            "delay_hours": 8.0,
            "delay_days": 0.33,
        }
        for i in range(1, count + 1)
    ]


def _make_utilization_rows(count: int) -> List[Dict[str, Any]]:
    return [
        {
            "machine_id": f"MC{i:03d}",
            "machine_name": f"设备{i}",
            "hours": 8.0,
            "task_count": 1,
            "capacity_hours": 16.0,
            "utilization": 0.5,
        }
        for i in range(1, count + 1)
    ]


class _FakePlanResolution:
    def __init__(self, version: int, selected_role: str) -> None:
        self.version = int(version)
        self.selected_role = selected_role
        self.scenario_display_name = ""
        self.is_scenario_preview = False

    def to_dict(self) -> Dict[str, Any]:
        return {
            "requested_role": self.selected_role,
            "selected_role": self.selected_role,
            "status": "resolved_comparison" if self.selected_role == "baseline_best" else "resolved_adopted",
            "candidate_id": 2 if self.selected_role == "baseline_best" else None,
            "candidate_key": "baseline" if self.selected_role == "baseline_best" else None,
        }


def test_report_export_size_mode_selection(app_client, db_path, monkeypatch) -> None:
    from core.infrastructure.database import get_connection
    from core.services.report.report_engine import ReportEngine

    conn = get_connection(db_path)
    try:
        conn.execute(
            """
            INSERT INTO ScheduleHistory (version, strategy, batch_count, op_count, result_status, result_summary, created_by)
            VALUES (?, ?, ?, ?, ?, ?, ?)
            """,
            (7, "test", 0, 0, "ok", None, "regression"),
        )
        conn.commit()
    finally:
        conn.close()

    client = app_client

    monkeypatch.setattr(ReportEngine, "EXPORT_DIRECT_MAX_ROWS", 2)
    monkeypatch.setattr(ReportEngine, "EXPORT_STREAM_MAX_ROWS", 4)

    def fake_overdue_batches(self, version: int, plan_role: Any = None, scenario_id: Any = None) -> Dict[str, Any]:
        items = _make_overdue_items(2)
        return {
            "version": int(version),
            "count": len(items),
            "scheduled_count": len(items),
            "unscheduled_count": 0,
            "as_of_time": "2026-01-02 08:00:00",
            "items": items,
            "scheduled_items": list(items),
            "unscheduled_items": [],
        }

    def fake_utilization(self, version: int, start_date: Any, end_date: Any, plan_role: Any = None, scenario_id: Any = None) -> Dict[str, Any]:
        return {
            "version": int(version),
            "start_date": str(start_date),
            "end_date": str(end_date),
            "capacity_hours_per_resource": 16.0,
            "machines": _make_utilization_rows(3),
            "operators": [],
        }

    def fake_resolve_plan(self, version: int, plan_role: Any, scenario_id: Any = None):
        selected_role = "baseline_best" if str(plan_role or "").strip() == "baseline_best" else "adopted"
        return _FakePlanResolution(int(version), selected_role)

    def fake_overdue_diagnosis_export_rows(
        self,
        *,
        version: int,
        resolution,
        batch_ids: Any = None,
        resource_type: Any = None,
        resource_id: Any = None,
        batch_id: Any = None,
    ) -> List[Dict[str, Any]]:
        return [
            {
                "diagnosis_id": f"延期诊断-v{int(version)}-B001",
                "summary": "计划完成已经晚了 8.00 小时。建议先复核：测试诊断。",
                "check_info": "测试用诊断依据。",
                "evidence_sources": "超期清单：测试证据",
                "data_gaps": "当前没有明显证据缺口。",
                "generated_at": "2026-01-02 08:00:00",
                "filters": f"排产版本：v{int(version)}；排产方案：正式采用方案",
            }
        ]

    monkeypatch.setattr(ReportEngine, "overdue_batches", fake_overdue_batches)
    monkeypatch.setattr(ReportEngine, "_overdue_diagnosis_export_rows", fake_overdue_diagnosis_export_rows)
    monkeypatch.setattr(ReportEngine, "utilization", fake_utilization)
    monkeypatch.setattr(ReportEngine, "_resolve_plan", fake_resolve_plan)

    overdue_resp = client.get("/reports/overdue/export?version=7")
    _assert_xlsx(overdue_resp, "GET /reports/overdue/export", 7)
    _assert_overdue_xlsx_columns(overdue_resp)
    overdue_mode = overdue_resp.headers.get("X-APS-Report-Export-Mode", "")
    if overdue_mode != "direct":
        raise RuntimeError(f"超期清单导出模式错误：{overdue_mode!r}")
    overdue_rows = overdue_resp.headers.get("X-APS-Report-Estimated-Rows", "")
    if overdue_rows != "2":
        raise RuntimeError(f"超期清单导出行数估计错误：{overdue_rows!r}")

    util_resp = client.get("/reports/utilization/export?version=7&start_date=2026-01-01&end_date=2026-01-07")
    _assert_xlsx(util_resp, "GET /reports/utilization/export", 7)
    _assert_utilization_xlsx_percent(util_resp)
    util_mode = util_resp.headers.get("X-APS-Report-Export-Mode", "")
    if util_mode != "stream":
        raise RuntimeError(f"资源负荷导出模式错误：{util_mode!r}")
    util_rows = util_resp.headers.get("X-APS-Report-Estimated-Rows", "")
    if util_rows != "3":
        raise RuntimeError(f"资源负荷导出行数估计错误：{util_rows!r}")

    candidate_stream_resp = client.get(
        "/reports/utilization/export?version=7&plan_role=baseline_best&start_date=2026-01-01&end_date=2026-01-07"
    )
    _assert_xlsx(candidate_stream_resp, "GET /reports/utilization/export?plan_role=baseline_best", 7)
    if candidate_stream_resp.headers.get("X-APS-Report-Export-Mode", "") != "stream":
        raise RuntimeError("候选方案资源负荷导出没有走 stream 模式")
    cd = unquote(candidate_stream_resp.headers.get("Content-Disposition", ""))
    if "原算法代表方案" not in cd:
        raise RuntimeError(f"候选方案 stream 导出文件名缺少中文方案名：{cd!r}")
    _assert_no_internal_terms_in_xlsx(candidate_stream_resp)
