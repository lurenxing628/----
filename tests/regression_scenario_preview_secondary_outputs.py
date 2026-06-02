from __future__ import annotations

import json
import sys
from io import BytesIO
from pathlib import Path
from urllib.parse import unquote

import pytest
from openpyxl import load_workbook

REPO_ROOT = Path(__file__).resolve().parents[1]
TESTS_ROOT = REPO_ROOT / "tests"
if str(TESTS_ROOT) not in sys.path:
    sys.path.insert(0, str(TESTS_ROOT))

from regression_gantt_draft_save_and_preview import (  # noqa: E402
    VERSION,
    _build_app,
    _connect,
    _draft_with_change,
    _seed_base,
)

from core.infrastructure.errors import ValidationError  # noqa: E402
from core.services.report import ReportEngine  # noqa: E402
from core.services.scheduler.gantt_adjustment_scenario_service import GanttAdjustmentScenarioService  # noqa: E402
from core.services.scheduler.gantt_service import GanttService  # noqa: E402
from core.services.scheduler.resource_dispatch_service import ResourceDispatchService  # noqa: E402

EXPORT_INTERNAL_TERMS = (
    "plan_role",
    "candidate_key",
    "candidate_id",
    "source_table",
    "baseline_best",
    "critical_best",
    "scenario_id",
    "candidate_rows",
    "adjustment_scenario_rows",
)


def _build_secondary_output_app(tmp_path: Path, monkeypatch):
    for name in list(sys.modules):
        if name in {"web.routes.reports", "web.routes.report_plan_preview"}:
            sys.modules.pop(name, None)
    return _build_app(tmp_path, monkeypatch)


def _save_secondary_output_scenario(conn, scenario_name="二级页模拟", clear_name=False) -> str:
    conn.execute("UPDATE Batches SET due_date = '2026-05-05' WHERE batch_id = 'B2'")
    conn.commit()
    draft_id = _draft_with_change(
        conn,
        to_start="2026-05-06 11:00:00",
        to_end="2026-05-06 12:00:00",
    )
    scenario = GanttAdjustmentScenarioService(conn).save_scenario(
        draft_id=draft_id,
        scenario_name=scenario_name,
        created_by="pytest",
    )
    if clear_name:
        conn.execute(
            "UPDATE ScheduleAdjustmentScenario SET scenario_name = NULL WHERE scenario_id = ?",
            (scenario.scenario_id,),
        )
    conn.execute(
        """
        INSERT INTO MachineDowntimes(machine_id, start_time, end_time, reason_code, reason_detail, status)
        VALUES ('M2', '2026-05-06 11:15:00', '2026-05-06 11:45:00', 'maint', 'pytest', 'active')
        """
    )
    conn.commit()
    return scenario.scenario_id


def _json_data(response):
    payload = json.loads(response.get_data(as_text=True) or "{}")
    return payload.get("data") or {}


def _workbook_summary_values(xlsx_bytes):
    wb = load_workbook(BytesIO(xlsx_bytes), read_only=True, data_only=True)
    try:
        if "查询摘要" not in wb.sheetnames:
            return None
        ws = wb["查询摘要"]
        return {str(row[0] or ""): row[1] for row in ws.iter_rows(values_only=True) if row and row[0]}
    finally:
        wb.close()


def _workbook_text(xlsx_bytes) -> str:
    wb = load_workbook(BytesIO(xlsx_bytes), read_only=True, data_only=True)
    values = []
    try:
        for ws in wb.worksheets:
            values.append(ws.title)
            for row in ws.iter_rows(values_only=True):
                values.extend(str(cell or "") for cell in row)
    finally:
        wb.close()
    return "\n".join(values)


def _workbook_sheetnames(xlsx_bytes):
    wb = load_workbook(BytesIO(xlsx_bytes), read_only=True, data_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def _assert_report_export_uses_public_scenario_name(response, *, scenario_id: str, expected_name: str) -> str:
    disposition = unquote(str(response.headers.get("Content-Disposition") or ""))
    sheetnames = _workbook_sheetnames(response.data)
    workbook_text = _workbook_text(response.data)
    summary = _workbook_summary_values(response.data)
    assert summary is not None
    assert summary["导出类型"] == "模拟方案预览"
    assert summary["提示"] == "这是模拟方案预览，正式计划还没有改变。"
    assert summary["模拟方案"] == expected_name
    assert "Sheet" not in sheetnames
    assert all(str(name or "").strip() for name in sheetnames)
    assert scenario_id not in disposition
    assert scenario_id not in workbook_text
    assert expected_name in disposition
    for forbidden in EXPORT_INTERNAL_TERMS:
        assert forbidden not in disposition
        assert forbidden not in workbook_text
    return workbook_text


def _assert_scenario_filter_resets_when_plan_identity_changes(html: str) -> None:
    assert 'name="scenario_id"' in html
    assert "data-report-plan-version-select" in html
    assert "data-report-plan-role-select" in html
    assert "data-initial-version" in html
    assert "data-initial-plan-role" in html
    assert "report_plan_filter.js" in html


def test_secondary_output_services_read_scenario_rows_and_do_not_fallback(tmp_path: Path) -> None:
    conn = _connect(tmp_path)
    try:
        _seed_base(conn)
        scenario_id = _save_secondary_output_scenario(conn)

        week_plan = GanttService(conn).get_week_plan_rows(
            week_start="2026-05-04",
            version=VERSION,
            plan_role="adopted",
            scenario_id=scenario_id,
        )
        assert week_plan["is_scenario_preview"] is True
        assert any(row["批次号"] == "B2" and row["日期"] == "2026-05-06" and "11:00" in row["时段"] for row in week_plan["rows"])

        dispatch = ResourceDispatchService(conn).get_dispatch_payload(
            scope_type="operator",
            operator_id="O2",
            period_preset="week",
            query_date="2026-05-06",
            version=VERSION,
            plan_role="adopted",
            scenario_id=scenario_id,
        )
        detail = dispatch["detail_rows"][0]
        assert dispatch["filters"]["scenario_id"] == scenario_id
        assert dispatch["filters"]["is_scenario_preview"] is True
        assert detail["start_time"] == "2026-05-06 11:00:00"
        assert detail["is_overdue"] is True

        engine = ReportEngine(conn)
        overdue = engine.overdue_batches(VERSION, plan_role="adopted", scenario_id=scenario_id)
        assert overdue["is_scenario_preview"] is True
        assert any(row["batch_id"] == "B2" for row in overdue["scheduled_items"])

        utilization = engine.utilization(
            VERSION,
            "2026-05-06",
            "2026-05-06",
            plan_role="adopted",
            scenario_id=scenario_id,
        )
        assert any(row["machine_id"] == "M2" and row["hours"] == 1.0 for row in utilization["machines"])

        downtime = engine.downtime_impact(
            VERSION,
            "2026-05-06",
            "2026-05-06",
            plan_role="adopted",
            scenario_id=scenario_id,
        )
        assert any(row["machine_id"] == "M2" and row["schedule_overlap_hours"] == 0.5 for row in downtime["machines"])

        with pytest.raises(ValidationError, match="模拟方案不存在"):
            engine.utilization(VERSION, "2026-05-06", "2026-05-06", plan_role="adopted", scenario_id="missing")
        with pytest.raises(ValidationError, match="模拟方案不存在"):
            ResourceDispatchService(conn).get_dispatch_payload(
                scope_type="operator",
                operator_id="O2",
                period_preset="week",
                query_date="2026-05-06",
                version=VERSION,
                plan_role="adopted",
                scenario_id="missing",
            )
    finally:
        conn.close()


def test_secondary_output_pages_keep_scenario_context(tmp_path: Path, monkeypatch) -> None:
    conn = _connect(tmp_path)
    try:
        _seed_base(conn)
        scenario_id = _save_secondary_output_scenario(conn)
    finally:
        conn.close()

    app = _build_secondary_output_app(tmp_path, monkeypatch)
    client = app.test_client()
    scenario_query = f"version={VERSION}&plan_role=adopted&scenario_id={scenario_id}"

    week_html = client.get(f"/scheduler/week-plan?week_start=2026-05-04&{scenario_query}").get_data(as_text=True)
    assert "当前周计划正在预览“二级页模拟”" in week_html
    assert week_html.count("当前周计划正在预览") == 1
    assert f'name="scenario_id" value="{scenario_id}"' in week_html
    assert f"scenario_id={scenario_id}" in week_html
    assert f"/scheduler/resource-dispatch?version={VERSION}&amp;plan_role=adopted&amp;scenario_id={scenario_id}" in week_html
    assert f"/scheduler/gantt?view=machine&amp;version={VERSION}&amp;plan_role=adopted&amp;scenario_id={scenario_id}" in week_html
    _assert_scenario_filter_resets_when_plan_identity_changes(week_html)
    assert "这套方案只用于对比，不代表最终采用的排产结果" not in week_html
    assert "2026-05-06" in week_html
    assert "11:00" in week_html

    gantt_html = client.get(
        f"/scheduler/gantt?view=machine&start_date=2026-05-04&end_date=2026-05-10&{scenario_query}"
    ).get_data(as_text=True)
    _assert_scenario_filter_resets_when_plan_identity_changes(gantt_html)

    week_export = client.get(f"/scheduler/week-plan/export?week_start=2026-05-04&{scenario_query}")
    assert week_export.status_code == 200
    week_disposition = unquote(str(week_export.headers.get("Content-Disposition") or ""))
    assert scenario_id not in week_disposition
    assert "二级页模拟" in week_disposition
    week_summary = _workbook_summary_values(week_export.data)
    assert week_summary is not None
    assert week_summary["导出类型"] == "模拟方案预览"
    assert week_summary["提示"] == "这是模拟方案预览，正式计划还没有改变。"
    assert week_summary["模拟方案"] == "二级页模拟"
    assert "模拟方案编号" not in week_summary
    assert "模拟方案名称" not in week_summary
    assert scenario_id not in "\n".join(str(value or "") for value in week_summary.values())

    formal_week_export = client.get(f"/scheduler/week-plan/export?week_start=2026-05-04&version={VERSION}&plan_role=adopted")
    assert formal_week_export.status_code == 200
    assert _workbook_summary_values(formal_week_export.data) is None

    resource_html = client.get(
        f"/scheduler/resource-dispatch?scope_type=operator&operator_id=O2&period_preset=week&query_date=2026-05-06&{scenario_query}"
    ).get_data(as_text=True)
    assert "当前资源排班正在预览" in resource_html
    assert f'name="scenario_id" value="{scenario_id}"' in resource_html
    assert f"scenario_id={scenario_id}" in resource_html
    assert f"/scheduler/week-plan?version={VERSION}&amp;plan_role=adopted&amp;scenario_id={scenario_id}" in resource_html
    _assert_scenario_filter_resets_when_plan_identity_changes(resource_html)

    resource_data = client.get(
        f"/scheduler/resource-dispatch/data?scope_type=operator&operator_id=O2&period_preset=week&query_date=2026-05-06&{scenario_query}"
    )
    assert resource_data.status_code == 200
    data = _json_data(resource_data)
    assert "scenario_id" not in data["filters"]
    assert data["filters"]["is_scenario_preview"] is True
    assert data["filters"]["scenario_display_name"] == "二级页模拟"
    assert data["detail_rows"][0]["start_time"] == "2026-05-06 11:00:00"

    resource_export = client.get(
        f"/scheduler/resource-dispatch/export?scope_type=operator&operator_id=O2&period_preset=week&query_date=2026-05-06&{scenario_query}"
    )
    assert resource_export.status_code == 200
    resource_disposition = unquote(str(resource_export.headers.get("Content-Disposition") or ""))
    assert scenario_id not in resource_disposition
    assert "二级页模拟" in resource_disposition
    assert "正式采用方案" not in resource_disposition

    overdue_html = client.get(f"/reports/overdue?{scenario_query}").get_data(as_text=True)
    assert "当前超期清单正在预览" in overdue_html
    assert f'name="scenario_id" value="{scenario_id}"' in overdue_html
    _assert_scenario_filter_resets_when_plan_identity_changes(overdue_html)
    assert f"/reports/utilization?version={VERSION}&amp;plan_role=adopted&amp;scenario_id={scenario_id}" in overdue_html
    assert "模拟预览暂不支持导出" not in overdue_html
    assert f"/reports/overdue/export?version={VERSION}&amp;plan_role=adopted&amp;scenario_id={scenario_id}" in overdue_html
    assert "B2" in overdue_html

    utilization_html = client.get(
        f"/reports/utilization?start_date=2026-05-06&end_date=2026-05-06&{scenario_query}"
    ).get_data(as_text=True)
    assert "当前资源负荷与利用率正在预览" in utilization_html
    assert f'name="scenario_id" value="{scenario_id}"' in utilization_html
    _assert_scenario_filter_resets_when_plan_identity_changes(utilization_html)
    assert (
        f"/reports/downtime?version={VERSION}&amp;plan_role=adopted&amp;scenario_id={scenario_id}"
        "&amp;start_date=2026-05-06&amp;end_date=2026-05-06"
    ) in utilization_html
    assert (
        f"/reports/utilization/export?version={VERSION}&amp;plan_role=adopted&amp;scenario_id={scenario_id}"
        "&amp;start_date=2026-05-06&amp;end_date=2026-05-06"
    ) in utilization_html
    assert "M2" in utilization_html
    assert '<div class="aps-summary-label">排产方案</div>' in utilization_html
    assert '<div class="aps-summary-value">二级页模拟</div>' in utilization_html
    assert "当前方案：二级页模拟" in utilization_html
    assert "模拟预览暂不支持导出" not in utilization_html

    downtime_html = client.get(
        f"/reports/downtime?start_date=2026-05-06&end_date=2026-05-06&{scenario_query}"
    ).get_data(as_text=True)
    assert "当前停机影响统计正在预览" in downtime_html
    assert f'name="scenario_id" value="{scenario_id}"' in downtime_html
    _assert_scenario_filter_resets_when_plan_identity_changes(downtime_html)
    assert "0.5" in downtime_html
    assert '<div class="aps-summary-label">排产方案</div>' in downtime_html
    assert '<div class="aps-summary-value">二级页模拟</div>' in downtime_html
    assert "当前方案：二级页模拟" in downtime_html
    assert (
        f"/reports/downtime/export?version={VERSION}&amp;plan_role=adopted&amp;scenario_id={scenario_id}"
        "&amp;start_date=2026-05-06&amp;end_date=2026-05-06"
    ) in downtime_html
    assert "模拟预览暂不支持导出" not in downtime_html

    export_resp = client.get(f"/reports/overdue/export?{scenario_query}")
    assert export_resp.status_code == 200
    overdue_export_text = _assert_report_export_uses_public_scenario_name(
        export_resp,
        scenario_id=scenario_id,
        expected_name="二级页模拟",
    )
    assert "B2" in overdue_export_text

    utilization_export = client.get(
        f"/reports/utilization/export?start_date=2026-05-06&end_date=2026-05-06&{scenario_query}"
    )
    assert utilization_export.status_code == 200
    utilization_export_text = _assert_report_export_uses_public_scenario_name(
        utilization_export,
        scenario_id=scenario_id,
        expected_name="二级页模拟",
    )
    assert "M2" in utilization_export_text

    downtime_export = client.get(
        f"/reports/downtime/export?start_date=2026-05-06&end_date=2026-05-06&{scenario_query}"
    )
    assert downtime_export.status_code == 200
    downtime_export_text = _assert_report_export_uses_public_scenario_name(
        downtime_export,
        scenario_id=scenario_id,
        expected_name="二级页模拟",
    )
    assert "M2" in downtime_export_text
    assert "0.5" in downtime_export_text


def test_secondary_output_pages_use_plain_fallback_for_unnamed_scenario(tmp_path: Path, monkeypatch) -> None:
    conn = _connect(tmp_path)
    try:
        _seed_base(conn)
        scenario_id = _save_secondary_output_scenario(conn, scenario_name="", clear_name=True)
    finally:
        conn.close()

    app = _build_secondary_output_app(tmp_path, monkeypatch)
    client = app.test_client()
    scenario_query = f"version={VERSION}&plan_role=adopted&scenario_id={scenario_id}"

    page_urls = [
        f"/scheduler/week-plan?week_start=2026-05-04&{scenario_query}",
        f"/scheduler/resource-dispatch?scope_type=operator&operator_id=O2&period_preset=week&query_date=2026-05-06&{scenario_query}",
        f"/scheduler/gantt?view=machine&start_date=2026-05-04&end_date=2026-05-10&{scenario_query}",
        f"/reports/overdue?{scenario_query}",
        f"/reports/utilization?start_date=2026-05-06&end_date=2026-05-06&{scenario_query}",
        f"/reports/downtime?start_date=2026-05-06&end_date=2026-05-06&{scenario_query}",
    ]
    for url in page_urls:
        html = client.get(url).get_data(as_text=True)
        assert "模拟预览（未命名）" in html, url
        if "/reports/utilization?" in url or "/reports/downtime?" in url:
            assert "当前方案：模拟预览（未命名）" in html, url
            assert '<div class="aps-summary-label">排产方案</div>' in html, url
            assert '<div class="aps-summary-value">模拟预览（未命名）</div>' in html, url
            assert "当前方案：None" not in html, url
            if "/reports/utilization?" in url:
                assert "当前资源负荷与利用率正在预览“模拟预览（未命名）”，正式计划还没有改变。" in html, url
            if "/reports/downtime?" in url:
                assert "当前停机影响统计正在预览“模拟预览（未命名）”，正式计划还没有改变。" in html, url

    week_export = client.get(f"/scheduler/week-plan/export?week_start=2026-05-04&{scenario_query}")
    assert week_export.status_code == 200
    week_disposition = unquote(str(week_export.headers.get("Content-Disposition") or ""))
    assert scenario_id not in week_disposition
    assert "模拟预览（未命名）" in week_disposition
    week_summary = _workbook_summary_values(week_export.data)
    assert week_summary is not None
    assert week_summary["模拟方案"] == "模拟预览（未命名）"
    assert scenario_id not in "\n".join(str(value or "") for value in week_summary.values())

    resource_export = client.get(
        f"/scheduler/resource-dispatch/export?scope_type=operator&operator_id=O2&period_preset=week&query_date=2026-05-06&{scenario_query}"
    )
    assert resource_export.status_code == 200
    resource_disposition = unquote(str(resource_export.headers.get("Content-Disposition") or ""))
    assert scenario_id not in resource_disposition
    assert "模拟预览（未命名）" in resource_disposition
    assert "正式采用方案" not in resource_disposition

    report_exports = [
        client.get(f"/reports/overdue/export?{scenario_query}"),
        client.get(f"/reports/utilization/export?start_date=2026-05-06&end_date=2026-05-06&{scenario_query}"),
        client.get(f"/reports/downtime/export?start_date=2026-05-06&end_date=2026-05-06&{scenario_query}"),
    ]
    for response in report_exports:
        assert response.status_code == 200
        _assert_report_export_uses_public_scenario_name(
            response,
            scenario_id=scenario_id,
            expected_name="模拟预览（未命名）",
        )


def test_secondary_output_pages_reject_bad_scenario_without_showing_official_plan(tmp_path: Path, monkeypatch) -> None:
    conn = _connect(tmp_path)
    try:
        _seed_base(conn)
    finally:
        conn.close()

    app = _build_secondary_output_app(tmp_path, monkeypatch)
    client = app.test_client()
    bad_query = f"version={VERSION}&plan_role=adopted&scenario_id=scenario-missing"

    for url in (
        f"/scheduler/week-plan?week_start=2026-05-04&{bad_query}",
        f"/scheduler/resource-dispatch?scope_type=operator&operator_id=O2&period_preset=week&query_date=2026-05-04&{bad_query}",
        f"/reports/overdue?{bad_query}",
        f"/reports/utilization?start_date=2026-05-04&end_date=2026-05-04&{bad_query}",
        f"/reports/downtime?start_date=2026-05-04&end_date=2026-05-04&{bad_query}",
    ):
        response = client.get(url, follow_redirects=False)
        html = response.get_data(as_text=True)
        assert response.status_code == 400, url
        assert "模拟方案不存在" in html, url
        assert "正式设备" not in html, url

    data_resp = client.get(
        f"/scheduler/resource-dispatch/data?scope_type=operator&operator_id=O2&period_preset=week&query_date=2026-05-04&{bad_query}"
    )
    assert data_resp.status_code == 400
    data_payload = json.loads(data_resp.get_data(as_text=True))
    assert (data_payload.get("error") or {}).get("message") == "模拟方案不存在。"
    data_details = (data_payload.get("error") or {}).get("details") or {}
    assert "scenario_id" not in data_details.get("cleanup_query_keys", [])

    export_resp = client.get(
        f"/scheduler/resource-dispatch/export?scope_type=operator&operator_id=O2&period_preset=week&query_date=2026-05-04&{bad_query}",
        follow_redirects=False,
    )
    assert export_resp.status_code == 400
    assert "模拟方案不存在" in export_resp.get_data(as_text=True)

    for url in (
        f"/reports/overdue/export?{bad_query}",
        f"/reports/utilization/export?start_date=2026-05-04&end_date=2026-05-04&{bad_query}",
        f"/reports/downtime/export?start_date=2026-05-04&end_date=2026-05-04&{bad_query}",
    ):
        response = client.get(url, follow_redirects=False)
        html = response.get_data(as_text=True)
        assert response.status_code == 400, url
        assert "模拟方案不存在" in html, url
        assert "scenario_id" not in html, url
