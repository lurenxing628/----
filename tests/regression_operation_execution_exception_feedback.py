from __future__ import annotations

import io
from typing import Any, Dict

import openpyxl

from core.infrastructure.database import get_connection
from core.services.scheduler.resource_dispatch_excel import build_resource_dispatch_workbook
from core.services.scheduler.resource_dispatch_service import ResourceDispatchService
from tests.regression_operation_execution_feedback_routes import (
    RESOURCE_DISPATCH_JS,
    RESOURCE_DISPATCH_TEMPLATE,
    _base_payload,
    _build_app,
    _current_card,
    _event_count,
    _json,
    _post_controlled,
)
from web.viewmodels.scheduler_resource_dispatch import decorate_resource_dispatch_payload


def _action_by_name(card: Dict[str, Any]) -> Dict[str, Dict[str, Any]]:
    return {item["action"]: item for item in card.get("available_actions") or []}


def _post(client, card: Dict[str, Any], action_path: str, **overrides: Any):
    return _post_controlled(
        client,
        f"/scheduler/resource-dispatch/execution/{card['op_id']}/{action_path}",
        _base_payload(card, **overrides),
    )


def test_pause_resume_exception_flow_updates_task_card_and_actions(tmp_path, monkeypatch) -> None:
    app, db_path = _build_app(tmp_path, monkeypatch)
    client = app.test_client()
    card = _current_card(client)

    initial_actions = _action_by_name(card)
    assert initial_actions["start"]["enabled"] is True
    assert initial_actions["report_exception"]["enabled"] is False
    assert "不能报异常" in initial_actions["report_exception"]["disabled_reason"]

    start_resp = _post(client, card, "start", idempotency_key="exception-flow-start")
    start_data = _json(start_resp)["data"]
    processing_actions = _action_by_name(start_data["task_card"])
    assert start_resp.status_code == 200
    assert start_data["current_status_label"] == "生产中"
    assert processing_actions["pause"]["enabled"] is True
    assert processing_actions["finish"]["enabled"] is True
    assert processing_actions["report_exception"]["enabled"] is True

    pause_resp = _post(
        client,
        card,
        "pause",
        expected_state_revision=start_data["state_revision"],
        event_time="2026-05-01 08:20:00",
        idempotency_key="exception-flow-pause",
        reason_code="equipment",
        remark="设备需要检查",
    )
    pause_data = _json(pause_resp)["data"]
    paused_actions = _action_by_name(pause_data["task_card"])
    assert pause_resp.status_code == 200
    assert pause_data["event"]["action_label"] == "暂停"
    assert pause_data["event"]["reason_label"] == "设备问题"
    assert pause_data["current_status_label"] == "已暂停"
    assert paused_actions["resume"]["enabled"] is True
    assert paused_actions["finish"]["enabled"] is True
    assert paused_actions["report_exception"]["enabled"] is True

    resume_resp = _post(
        client,
        card,
        "resume",
        expected_state_revision=pause_data["state_revision"],
        event_time="2026-05-01 08:35:00",
        idempotency_key="exception-flow-resume",
        remark="继续生产",
    )
    resume_data = _json(resume_resp)["data"]
    assert resume_resp.status_code == 200
    assert resume_data["event"]["action_label"] == "继续生产"
    assert resume_data["current_status_label"] == "生产中"

    exception_resp = _post(
        client,
        card,
        "report-exception",
        expected_state_revision=resume_data["state_revision"],
        event_time="2026-05-01 08:45:00",
        idempotency_key="exception-flow-report",
        reason_code="equipment",
        severity="high",
        impact_minutes=30,
        affected_machine_id="M2",
        affected_operator_id="O2",
        handling_status="checking",
        suggest_reschedule=True,
        remark="主轴异常",
        reason_detail="主轴异常",
    )
    exception_data = _json(exception_resp)["data"]
    exception_actions = _action_by_name(exception_data["task_card"])

    assert exception_resp.status_code == 200
    assert exception_data["event"]["action"] == "report_exception"
    assert exception_data["event"]["action_label"] == "报异常"
    assert exception_data["event"]["reason_label"] == "设备问题"
    assert exception_data["event"]["severity_label"] == "严重"
    assert exception_data["event"]["impact_minutes_label"] == "预计影响 30 分钟"
    assert exception_data["event"]["affected_machine_label"] == "二号设备"
    assert exception_data["event"]["affected_operator_label"] == "李四"
    assert exception_data["event"]["handling_status_label"] == "处理中"
    assert exception_data["event"]["suggest_reschedule_label"] == "建议重新排程"
    assert exception_data["current_status_label"] == "异常中"
    assert exception_data["task_card"]["current_status_label"] == "异常中"
    assert exception_data["task_card"]["latest_exception_reason_label"] == "设备问题"
    assert exception_data["task_card"]["latest_exception_severity_label"] == "严重"
    assert exception_data["task_card"]["latest_exception_impact_minutes_label"] == "预计影响 30 分钟"
    assert exception_data["task_card"]["latest_exception_affected_machine_label"] == "二号设备"
    assert exception_data["task_card"]["latest_exception_affected_operator_label"] == "李四"
    assert exception_data["task_card"]["latest_exception_handling_status_label"] == "处理中"
    assert exception_data["task_card"]["latest_exception_suggest_reschedule_label"] == "建议重新排程"
    assert exception_data["task_card"]["latest_exception_remark"] == "主轴异常"
    assert exception_actions["resume"]["enabled"] is True
    assert exception_actions["finish"]["enabled"] is True
    assert exception_actions["report_exception"]["enabled"] is False
    assert exception_data["state_revision"] != resume_data["state_revision"]
    assert _event_count(db_path) == 4

    events_resp = client.get(
        f"/scheduler/resource-dispatch/execution/{card['op_id']}/events"
        "?scope_type=operator&operator_id=O1&period_preset=week&query_date=2026-05-01&version=2&plan_role=adopted"
    )
    events_payload = _json(events_resp)["data"]
    events = events_payload["events"]
    assert events_resp.status_code == 200
    assert [item["action_label"] for item in events] == ["开工", "暂停", "继续生产", "报异常"]
    assert events[-1]["action"] == "report_exception"
    assert events[-1]["reason_label"] == "设备问题"
    assert events[-1]["severity_label"] == "严重"
    assert events[-1]["affected_machine_label"] == "二号设备"
    assert events[-1]["affected_operator_label"] == "李四"
    assert events[-1]["remark"] == "主轴异常"
    assert "event_type" not in events[-1]
    assert events_payload["task_card"]["current_status_label"] == "异常中"


def test_exception_feedback_validation_and_conflicts_are_user_facing(tmp_path, monkeypatch) -> None:
    app, db_path = _build_app(tmp_path, monkeypatch)
    client = app.test_client()
    card = _current_card(client)

    not_started_exception = _post(
        client,
        card,
        "report-exception",
        idempotency_key="not-started-exception",
        reason_code="equipment",
        severity="high",
        remark="设备异常",
    )
    error = _json(not_started_exception)["error"]
    assert not_started_exception.status_code == 409
    assert error["details"]["reason"] == "invalid_state_transition"
    assert error["details"]["action_label"] == "报异常"

    start_data = _json(_post(client, card, "start", idempotency_key="validation-start"))["data"]
    cases = [
        ("pause", {"reason_code": "", "remark": "设备需要检查"}, "原因", "reason_code"),
        ("pause", {"reason_code": "equipment", "remark": ""}, "情况说明", "remark"),
        ("report-exception", {"reason_code": "equipment", "severity": "", "remark": "设备异常"}, "严重程度", "severity"),
        (
            "report-exception",
            {"reason_code": "equipment", "severity": "high", "impact_minutes": -1, "remark": "设备异常"},
            "预计影响时间",
            "impact_minutes",
        ),
        (
            "report-exception",
            {"reason_code": "equipment", "severity": "high", "affected_machine_id": "NO_SUCH", "remark": "设备异常"},
            "影响设备",
            "affected_machine_id",
        ),
        (
            "report-exception",
            {"reason_code": "equipment", "severity": "high", "affected_operator_id": "NO_SUCH", "remark": "设备异常"},
            "影响人员",
            "affected_operator_id",
        ),
    ]
    for index, (action_path, overrides, field_label, raw_field) in enumerate(cases):
        resp = _post(
            client,
            card,
            action_path,
            expected_state_revision=start_data["state_revision"],
            idempotency_key=f"bad-exception-feedback-{index}",
            **overrides,
        )
        payload = _json(resp)
        assert resp.status_code == 400
        assert payload["error"]["details"]["field_label"] == field_label
        assert raw_field not in payload["error"]["message"]

    stale_resp = _post(
        client,
        card,
        "report-exception",
        idempotency_key="stale-exception-feedback",
        reason_code="equipment",
        severity="high",
        remark="设备异常",
    )
    stale_error = _json(stale_resp)["error"]
    assert stale_resp.status_code == 409
    assert stale_error["details"]["reason"] == "stale_state_revision"
    assert stale_error["details"]["action_label"] == "报异常"
    assert _event_count(db_path) == 1


def test_reason_detail_only_is_visible_as_exception_remark(tmp_path, monkeypatch) -> None:
    app, db_path = _build_app(tmp_path, monkeypatch)
    client = app.test_client()
    card = _current_card(client)
    start_data = _json(_post(client, card, "start", idempotency_key="reason-detail-start"))["data"]

    exception_resp = _post(
        client,
        card,
        "report-exception",
        expected_state_revision=start_data["state_revision"],
        event_time="2026-05-01 08:30:00",
        idempotency_key="reason-detail-exception",
        reason_code="equipment",
        severity="high",
        remark="",
        reason_detail="只填原因详情",
    )
    exception_data = _json(exception_resp)["data"]

    assert exception_resp.status_code == 200
    assert exception_data["event"]["remark"] == "只填原因详情"
    assert exception_data["task_card"]["latest_exception_remark"] == "只填原因详情"

    data_resp = client.get(
        "/scheduler/resource-dispatch/data?scope_type=operator&operator_id=O1&period_preset=week&query_date=2026-05-01&version=2&plan_role=adopted"
    )
    detail_row = _json(data_resp)["data"]["detail_rows"][0]
    assert detail_row["latest_exception_remark"] == "只填原因详情"

    events_resp = client.get(
        f"/scheduler/resource-dispatch/execution/{card['op_id']}/events"
        "?scope_type=operator&operator_id=O1&period_preset=week&query_date=2026-05-01&version=2&plan_role=adopted"
    )
    assert _json(events_resp)["data"]["events"][-1]["remark"] == "只填原因详情"
    assert _event_count(db_path) == 2


def test_events_list_keeps_each_exception_own_impact_and_affected_resources(tmp_path, monkeypatch) -> None:
    app, _db_path = _build_app(tmp_path, monkeypatch)
    client = app.test_client()
    card = _current_card(client)
    start_data = _json(_post(client, card, "start", idempotency_key="two-exceptions-start"))["data"]
    first_exception = _json(
        _post(
            client,
            card,
            "report-exception",
            expected_state_revision=start_data["state_revision"],
            event_time="2026-05-01 08:30:00",
            idempotency_key="two-exceptions-first",
            reason_code="equipment",
            severity="high",
            impact_minutes=30,
            affected_machine_id="M2",
            affected_operator_id="O2",
            handling_status="new",
            suggest_reschedule=True,
            remark="第一次异常",
        )
    )["data"]
    resumed = _json(
        _post(
            client,
            card,
            "resume",
            expected_state_revision=first_exception["state_revision"],
            event_time="2026-05-01 08:45:00",
            idempotency_key="two-exceptions-resume",
            remark="继续生产",
        )
    )["data"]
    _post(
        client,
        card,
        "report-exception",
        expected_state_revision=resumed["state_revision"],
        event_time="2026-05-01 09:00:00",
        idempotency_key="two-exceptions-second",
        reason_code="person",
        severity="medium",
        impact_minutes=60,
        affected_machine_id="M1",
        affected_operator_id="O1",
        handling_status="checking",
        suggest_reschedule=False,
        remark="第二次异常",
    )

    events_resp = client.get(
        f"/scheduler/resource-dispatch/execution/{card['op_id']}/events"
        "?scope_type=operator&operator_id=O1&period_preset=week&query_date=2026-05-01&version=2&plan_role=adopted"
    )
    events = _json(events_resp)["data"]["events"]
    exception_events = [item for item in events if item["action"] == "report_exception"]

    assert events_resp.status_code == 200
    assert len(exception_events) == 2
    assert exception_events[0]["reason_label"] == "设备问题"
    assert exception_events[0]["impact_minutes_label"] == "预计影响 30 分钟"
    assert exception_events[0]["affected_machine_label"] == "二号设备"
    assert exception_events[0]["affected_operator_label"] == "李四"
    assert exception_events[0]["remark"] == "第一次异常"
    assert exception_events[1]["reason_label"] == "人员问题"
    assert exception_events[1]["impact_minutes_label"] == "预计影响 60 分钟"
    assert exception_events[1]["affected_machine_label"] == "一号设备"
    assert exception_events[1]["affected_operator_label"] == "张三"
    assert exception_events[1]["remark"] == "第二次异常"


def test_report_exception_rejects_candidate_scenario_and_history_plans(tmp_path, monkeypatch) -> None:
    app, db_path = _build_app(tmp_path, monkeypatch)
    client = app.test_client()
    card = _current_card(client)
    cases = (
        {"requested_plan_role": "baseline_best", "effective_plan_role": "adopted", "source_table": "schedule"},
        {"requested_plan_role": "adopted", "scenario_id": "scenario-plain"},
        {"version": 1},
    )

    for index, overrides in enumerate(cases):
        resp = _post(
            client,
            card,
            "report-exception",
            idempotency_key=f"reject-report-exception-{index}",
            reason_code="equipment",
            severity="high",
            remark="设备异常",
            **overrides,
        )
        payload = _json(resp)
        assert resp.status_code == 409
        assert payload["error"]["details"]["reason"] == "not_current_official_plan"
        assert payload["error"]["details"]["action_label"] == "报异常"
        assert payload["error"]["details"]["can_retry"] is False
    assert _event_count(db_path) == 0


def test_exception_details_are_visible_in_detail_rows_gantt_popup_source_and_export(tmp_path, monkeypatch) -> None:
    app, db_path = _build_app(tmp_path, monkeypatch)
    client = app.test_client()
    card = _current_card(client)
    start_data = _json(_post(client, card, "start", idempotency_key="surface-start"))["data"]
    _post(
        client,
        card,
        "report-exception",
        expected_state_revision=start_data["state_revision"],
        event_time="2026-05-01 08:30:00",
        idempotency_key="surface-report-exception",
        reason_code="equipment",
        severity="critical",
        impact_minutes="",
        affected_machine_id="M2",
        affected_operator_id="O2",
        handling_status="waiting",
        suggest_reschedule=False,
        remark="等待维修",
    )

    data_resp = client.get(
        "/scheduler/resource-dispatch/data?scope_type=operator&operator_id=O1&period_preset=week&query_date=2026-05-01&version=2&plan_role=adopted"
    )
    detail_row = _json(data_resp)["data"]["detail_rows"][0]
    assert detail_row["execution_status_label"] == "异常中"
    assert detail_row["latest_exception_reason_label"] == "设备问题"
    assert detail_row["latest_exception_severity_label"] == "紧急"
    assert detail_row["latest_exception_impact_minutes_label"] == "暂时不知道影响多久"
    assert detail_row["latest_exception_affected_machine_label"] == "二号设备"
    assert detail_row["latest_exception_affected_operator_label"] == "李四"
    assert detail_row["latest_exception_handling_status_label"] == "等待条件"
    assert detail_row["latest_exception_suggest_reschedule_label"] == "暂不建议重新排程"
    assert detail_row["latest_exception_remark"] == "等待维修"

    db_conn = get_connection(db_path)
    try:
        payload = ResourceDispatchService(db_conn, logger=None, op_logger=None).get_dispatch_payload(
            scope_type="operator",
            operator_id="O1",
            period_preset="week",
            query_date="2026-05-01",
            version=2,
            plan_role="adopted",
        )
        payload = decorate_resource_dispatch_payload(payload)
        buffer = build_resource_dispatch_workbook(payload)
    finally:
        db_conn.close()

    wb = openpyxl.load_workbook(io.BytesIO(buffer.getvalue()))
    ws = wb["任务明细"]
    headers = [cell.value for cell in ws[1]]
    values = [cell.value for cell in ws[2]]
    row_map = dict(zip(headers, values))
    assert row_map["现场状态"] == "异常中"
    assert row_map["最近异常原因"] == "设备问题"
    assert row_map["严重程度"] == "紧急"
    assert row_map["预计影响时间"] == "暂时不知道影响多久"
    assert row_map["影响设备"] == "二号设备"
    assert row_map["影响人员"] == "李四"
    assert row_map["处理状态"] == "等待条件"
    assert row_map["是否建议重排"] == "暂不建议重新排程"
    assert row_map["情况说明"] == "等待维修"


def test_frontend_static_contract_for_exception_feedback() -> None:
    source = RESOURCE_DISPATCH_JS.read_text(encoding="utf-8")
    template = RESOURCE_DISPATCH_TEMPLATE.read_text(encoding="utf-8")

    assert "report-exception" in source
    assert "请选择原因：设备问题、人员问题、物料问题、质量问题、工艺问题、外协问题、其他" in source
    assert "请选择严重程度：轻微、一般、严重、紧急" in source
    assert "请填写影响设备编号（可留空）" in source
    assert "请填写影响人员工号（可留空）" in source
    assert "payload.affected_machine_id = affectedMachine" in source
    assert "payload.affected_operator_id = affectedOperator" in source
    assert "是否建议重新排程？请输入 是 或 否" in source
    assert "最近异常" in source
    assert "影响设备" in source
    assert "影响人员" in source
    assert "现场状态" in source
    assert "紧急异常，请计划员尽快处理。" in source
    assert "现场状态" in template
    assert "最近异常" in template
    assert "影响资源" in template
