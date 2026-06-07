"""回归测试：超期报表页 /reports/overdue 及其导出（含模拟方案预览 scenario_id 入参）把延期诊断以大白话中文呈现（查看为什么晚了/建议先复核/证据等级/证据缺口/下一步），不下结论式归因（不出现"物料不够导致延期"），且不向用户泄露 trace_meta、rule_version、source_table、scenario_id、根因 等内部术语；导出"诊断依据"sheet 表头固定为七列。"""

from __future__ import annotations

import sys
from io import BytesIO
from pathlib import Path
from typing import Iterable
from urllib.parse import unquote

from openpyxl import load_workbook

from core.infrastructure.database import ensure_schema, get_connection
from tests._support.paths import REPO_ROOT

TESTS_ROOT = REPO_ROOT / "tests"
if str(TESTS_ROOT) not in sys.path:
    sys.path.insert(0, str(TESTS_ROOT))

from regression_scheduler_delay_diagnosis_contract import (  # noqa: E402
    VERSION,
    _seed_base,
    _seed_candidates,
    _seed_scenario,
)

INTERNAL_TERMS = (
    "trace_meta",
    "rule_version",
    "input_fingerprint",
    "source_table",
    "scenario_id",
    "primary_reason",
    "critical_chain",
    "根因",
    "主要原因",
)


def _build_app(tmp_path: Path, monkeypatch):
    db_path = tmp_path / "aps.db"
    monkeypatch.setenv("APS_ENV", "development")
    monkeypatch.setenv("APS_DB_PATH", str(db_path))
    monkeypatch.setenv("APS_LOG_DIR", str(tmp_path / "logs"))
    monkeypatch.setenv("APS_BACKUP_DIR", str(tmp_path / "backups"))
    monkeypatch.setenv("APS_EXCEL_TEMPLATE_DIR", str(tmp_path / "templates_excel"))
    ensure_schema(str(db_path), logger=None, schema_path=str(REPO_ROOT / "schema.sql"), backup_dir=None)
    conn = get_connection(str(db_path))
    try:
        _seed_base(conn)
        _seed_candidates(conn)
        scenario_id = _seed_scenario(conn)
        conn.commit()
    finally:
        conn.close()

    for name in ("app", "web.routes.reports", "web.routes.report_plan_preview"):
        sys.modules.pop(name, None)
    import app as app_mod  # noqa: E402

    return app_mod.create_app(), scenario_id


def _workbook_texts(xlsx_bytes: bytes) -> Iterable[str]:
    wb = load_workbook(BytesIO(xlsx_bytes), data_only=True)
    try:
        for ws in wb.worksheets:
            yield ws.title
            for row in ws.iter_rows(values_only=True):
                for value in row:
                    if value is not None:
                        yield str(value)
    finally:
        wb.close()


def _assert_no_internal_terms(text: str) -> None:
    for term in INTERNAL_TERMS:
        assert term not in text


def test_overdue_page_shows_delay_diagnosis_in_plain_chinese(tmp_path: Path, monkeypatch) -> None:
    app, _scenario_id = _build_app(tmp_path, monkeypatch)
    client = app.test_client()

    response = client.get(f"/reports/overdue?version={VERSION}&plan_role=adopted")
    assert response.status_code == 200
    html = response.get_data(as_text=True)

    assert "查看为什么晚了" in html
    assert "建议先复核" in html
    assert "证据等级" in html
    assert "证据缺口" in html
    assert "下一步" in html
    assert "没有现场执行反馈时，不判断现场做慢了" in html
    assert "缺少批次物料明细" in html
    assert "当前物料数据不足，只提示核对物料，不判断一定是物料造成延期" in html
    assert "物料不够导致延期" not in html
    _assert_no_internal_terms(html)


def test_overdue_export_contains_trace_sheet_without_internal_terms(tmp_path: Path, monkeypatch) -> None:
    app, _scenario_id = _build_app(tmp_path, monkeypatch)
    client = app.test_client()

    response = client.get(f"/reports/overdue/export?version={VERSION}&plan_role=adopted")
    assert response.status_code == 200
    wb = load_workbook(BytesIO(response.data), data_only=True)
    try:
        assert "超期清单" in wb.sheetnames
        assert "诊断依据" in wb.sheetnames
        ws = wb["诊断依据"]
        headers = [ws.cell(row=1, column=i).value for i in range(1, 8)]
        assert headers == ["本次诊断编号", "生成依据摘要", "核对信息", "证据来源", "证据缺口", "生成时间", "筛选条件"]
        values = "\n".join(_workbook_texts(response.data))
    finally:
        wb.close()

    assert "延期诊断-v" in values
    assert "建议先复核" in values
    assert "证据等级" in values
    assert "缺少批次物料明细" in values
    _assert_no_internal_terms(values)


def test_overdue_export_supports_scenario_preview_without_internal_terms(tmp_path: Path, monkeypatch) -> None:
    app, scenario_id = _build_app(tmp_path, monkeypatch)
    client = app.test_client()

    response = client.get(f"/reports/overdue/export?version={VERSION}&plan_role=adopted&scenario_id={scenario_id}")
    assert response.status_code == 200
    disposition = unquote(str(response.headers.get("Content-Disposition") or ""))
    values = "\n".join(_workbook_texts(response.data))
    assert "延期诊断预览" in disposition
    assert scenario_id not in disposition
    assert scenario_id not in values
    _assert_no_internal_terms(values)
