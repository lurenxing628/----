from __future__ import annotations

import io
import json
from datetime import datetime

import openpyxl
import pytest

from core.infrastructure.errors import ValidationError
from core.models.operation_execution_state import OperationExecutionState
from core.services.scheduler.resource_dispatch_actual_import import ResourceDispatchActualImportPreviewer
from core.services.scheduler.resource_dispatch_actual_records import TaskRef
from tests.operation_execution_feedback_test_support import (
    _base_payload,
    _build_app,
    _current_card,
    _current_query,
    _event_count,
    _events_for_op,
    _json,
)


def test_actual_template_uses_public_columns_without_internal_fields(tmp_path, monkeypatch) -> None:
    app, _db_path = _build_app(tmp_path, monkeypatch)
    client = app.test_client()

    resp = client.get(f"/scheduler/resource-dispatch/execution/actual-template?{_current_query()}")
    wb = openpyxl.load_workbook(io.BytesIO(resp.data))

    assert resp.status_code == 200
    assert set(wb.sheetnames) >= {"任务反馈", "暂停明细"}
    task_headers = [cell.value for cell in wb["任务反馈"][1]]
    pause_headers = [cell.value for cell in wb["暂停明细"][1]]
    assert "任务识别码" in task_headers
    assert "实际开工时间" in task_headers
    assert "实际完工时间" in task_headers
    assert "暂停开始时间" in pause_headers
    assert "暂停结束时间" in pause_headers
    assert "暂停时长分钟" in pause_headers
    all_values = "\n".join(str(cell.value or "") for ws in wb.worksheets for row in ws.iter_rows() for cell in row)
    for forbidden in ("op_id", "schedule_id", "state_revision", "execution_snapshot_revision"):
        assert forbidden not in all_values


def test_actual_import_rejects_wrong_workbook_shape(tmp_path, monkeypatch) -> None:
    app, db_path = _build_app(tmp_path, monkeypatch)
    client = app.test_client()
    wb = openpyxl.Workbook()
    wb.active.title = "不是模板"
    wb.active.append(["任务识别码", "实际开工时间"])
    output = io.BytesIO()
    wb.save(output)

    resp = _post_preview(client, output.getvalue())
    payload = _json(resp)

    assert resp.status_code == 400
    assert "请使用下载的现场实际情况填写模板" in payload["error"]["message"]
    assert _event_count(db_path) == 0


def test_actual_import_rejects_missing_task_code_header(tmp_path, monkeypatch) -> None:
    app, db_path = _build_app(tmp_path, monkeypatch)
    client = app.test_client()
    wb = openpyxl.Workbook()
    wb.active.title = "任务反馈"
    wb.active.append(["批次", "实际开工时间"])
    output = io.BytesIO()
    wb.save(output)

    resp = _post_preview(client, output.getvalue())
    payload = _json(resp)

    assert resp.status_code == 400
    assert "缺少任务识别码列" in payload["error"]["message"]
    assert _event_count(db_path) == 0


def test_actual_import_oversize_file_gives_friendly_chinese_limit(tmp_path, monkeypatch) -> None:
    """超过 Excel 体积上限时，导入入口给出友好的中文提示（而不是绕过这道校验）。"""
    app, db_path = _build_app(tmp_path, monkeypatch)
    app.config["EXCEL_MAX_UPLOAD_BYTES"] = 1 * 1024 * 1024
    app.config["MAX_CONTENT_LENGTH"] = 8 * 1024 * 1024
    client = app.test_client()

    resp = _post_import(client, b"x" * (app.config["EXCEL_MAX_UPLOAD_BYTES"] + 1))
    payload = _json(resp)

    assert resp.status_code == 413
    assert "上传文件超过 1MB" in payload["error"]["message"]
    assert _event_count(db_path) == 0


def test_actual_import_rejects_duplicate_task_codes_before_matching() -> None:
    task = TaskRef(
        task_code="重复任务",
        schedule_version=2,
        schedule_id=100,
        op_id=10,
        batch_id="B1",
        op_name="OP10",
        planned_machine_id="M1",
        planned_machine_label="一号设备",
        planned_operator_id="O1",
        planned_operator_label="张三",
        state=OperationExecutionState(op_id=10, batch_id="B1"),
    )
    previewer = ResourceDispatchActualImportPreviewer(feedback_service=object())

    with pytest.raises(ValidationError) as exc_info:
        previewer.preview([{"sheet": "任务反馈", "row_number": 2, "任务识别码": "重复任务"}], [task, task])

    assert "重复的任务识别码" in exc_info.value.message


def _workbook_bytes(rows, pause_rows=None) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "任务反馈"
    headers = [
        "任务识别码",
        "批次",
        "工序",
        "计划设备",
        "计划人员",
        "实际开工时间",
        "实际完工时间",
        "完成数量",
        "报废数量",
        "异常时间",
        "异常原因",
        "异常严重程度",
        "异常说明",
        "反馈人",
        "备注",
    ]
    ws.append(headers)
    for row in rows:
        ws.append([row.get(header) for header in headers])
    pause_ws = wb.create_sheet("暂停明细")
    pause_headers = ["任务识别码", "暂停开始时间", "暂停结束时间", "暂停时长分钟", "暂停原因", "暂停说明", "反馈人"]
    pause_ws.append(pause_headers)
    for row in pause_rows or []:
        pause_ws.append([row.get(header) for header in pause_headers])
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def _task_code_from_template(client) -> str:
    resp = client.get(f"/scheduler/resource-dispatch/execution/actual-template?{_current_query()}")
    wb = openpyxl.load_workbook(io.BytesIO(resp.data), data_only=True)
    return str(wb["任务反馈"].cell(2, 1).value or "")


def _post_preview(client, file_bytes: bytes):
    return client.post(
        f"/scheduler/resource-dispatch/execution/import/preview?{_current_query()}",
        data={"file": (io.BytesIO(file_bytes), "actual.xlsx")},
        content_type="multipart/form-data",
    )


def _post_import(client, file_bytes: bytes):
    return client.post(
        f"/scheduler/resource-dispatch/execution/import?{_current_query()}",
        data={"file": (io.BytesIO(file_bytes), "actual.xlsx")},
        content_type="multipart/form-data",
    )


def test_actual_template_and_import_reject_incomplete_query_context(tmp_path, monkeypatch) -> None:
    app, db_path = _build_app(tmp_path, monkeypatch)
    client = app.test_client()
    import_file = _workbook_bytes([{"任务识别码": "任意任务", "实际开工时间": "2026-05-01 08:00:00"}])
    incomplete_queries = (
        "version=2&plan_role=adopted",
        "operator_id=O1&period_preset=week&query_date=2026-05-01&date_from=2026-05-01&date_to=2026-05-07&version=2&plan_role=adopted",
        "scope_type=operator&operator_id=O1&query_date=2026-05-01&date_from=2026-05-01&date_to=2026-05-07&version=2&plan_role=adopted",
    )

    for index, incomplete_query in enumerate(incomplete_queries):
        template = client.get(f"/scheduler/resource-dispatch/execution/actual-template?{incomplete_query}")
        preview = client.post(
            f"/scheduler/resource-dispatch/execution/import/preview?{incomplete_query}",
            data={"file": (io.BytesIO(import_file), "actual.xlsx")},
            content_type="multipart/form-data",
        )
        direct_import = client.post(
            f"/scheduler/resource-dispatch/execution/import?{incomplete_query}",
            data={"file": (io.BytesIO(import_file), f"actual-{index}.xlsx")},
            content_type="multipart/form-data",
        )
        confirm = client.post(
            f"/scheduler/resource-dispatch/execution/import/confirm?{incomplete_query}",
            json={"preview_token": f"incomplete-{index}", "raw_rows": []},
        )

        for resp in (template, preview, direct_import, confirm):
            payload = _json(resp)
            assert resp.status_code == 400
            assert payload["error"]["details"]["field"] == "plan_identity"
    assert _event_count(db_path) == 0


def test_actual_import_preview_reports_errors_and_does_not_write_db(tmp_path, monkeypatch) -> None:
    app, db_path = _build_app(tmp_path, monkeypatch)
    client = app.test_client()
    task_code = _task_code_from_template(client)
    file_bytes = _workbook_bytes(
        [
            {
                "任务识别码": task_code,
                "实际开工时间": "2026-05-01 09:00:00",
                "实际完工时间": "2026-05-01 08:50:00",
                "完成数量": 10,
            },
            {"任务识别码": "找不到的任务", "实际开工时间": "2026-05-01 08:00:00"},
            {"任务识别码": task_code, "实际开工时间": "bad-time"},
        ],
        pause_rows=[
            {
                "任务识别码": task_code,
                "暂停开始时间": "2026-05-01 08:05:00",
                "暂停结束时间": "2026-05-01 08:25:00",
                "暂停时长分钟": 10,
                "暂停原因": "设备问题",
                "暂停说明": "设备点检",
            },
            {
                "任务识别码": task_code,
                "暂停开始时间": "2026-05-01 08:20:00",
                "暂停结束时间": "2026-05-01 08:40:00",
                "暂停原因": "设备问题",
                "暂停说明": "重叠暂停一",
            },
            {
                "任务识别码": task_code,
                "暂停开始时间": "2026-05-01 08:30:00",
                "暂停结束时间": "2026-05-01 08:45:00",
                "暂停原因": "设备问题",
                "暂停说明": "重叠暂停二",
            },
        ],
    )

    resp = _post_preview(client, file_bytes)
    payload = _json(resp)["data"]
    messages = "\n".join(row.get("message") or "" for row in payload["rows"])

    assert resp.status_code == 200
    assert payload["can_confirm"] is False
    assert payload["summary"]["matched_rows"] == 5
    assert payload["summary"]["total_rows"] == 6
    assert payload["summary"]["error_count"] == 6
    assert payload["summary"]["conflict_count"] == 3
    assert payload["summary"]["add_count"] == 0
    assert payload["summary"]["skip_count"] == 0
    assert payload["summary"]["can_confirm"] is False
    assert {row["status"] for row in payload["rows"]} == {"error"}
    missing_rows = [row for row in payload["rows"] if row["task_code"] == "找不到的任务"]
    assert len(missing_rows) == 1
    assert missing_rows[0]["messages"] == ["找不到任务，请重新下载填写模板后再导入。"]
    assert "找不到任务" in messages
    assert "反馈时间格式不正确" in messages
    assert "实际完工时间不能早于实际开工时间" in messages
    assert "暂停结束时间和暂停时长对不上" in messages
    assert "重叠" in messages
    assert _event_count(db_path) == 0
    raw_text = json.dumps(payload["raw_rows"], ensure_ascii=False)
    for forbidden in ("op_id", "schedule_id", "state_revision", "execution_snapshot_revision"):
        assert forbidden not in raw_text


def test_actual_import_submit_checks_then_writes_or_rejects_whole_file(tmp_path, monkeypatch) -> None:
    app, db_path = _build_app(tmp_path, monkeypatch)
    client = app.test_client()
    task_code = _task_code_from_template(client)

    bad_resp = _post_import(client, _workbook_bytes([{"任务识别码": task_code, "实际开工时间": "bad-time"}]))
    bad_payload = _json(bad_resp)

    assert bad_resp.status_code == 400
    assert bad_payload["error"]["details"]["reason"] == "actual_import_validation_failed"
    assert bad_payload["error"]["details"]["summary"]["error_count"] == 1
    assert "反馈时间格式不正确" in bad_payload["error"]["details"]["rows"][0]["message"]
    assert "将新增" not in json.dumps(bad_payload, ensure_ascii=False)
    assert "将跳过" not in json.dumps(bad_payload, ensure_ascii=False)
    assert _event_count(db_path) == 0

    good_resp = _post_import(
        client,
        _workbook_bytes(
            [
                {
                    "任务识别码": task_code,
                    "实际开工时间": "2026-05-01 08:00:00",
                    "实际完工时间": "2026-05-01 09:00:00",
                    "完成数量": 10,
                    "反馈人": "",
                }
            ]
        ),
    )
    good_payload = _json(good_resp)["data"]
    events = _events_for_op(db_path, 10)

    assert good_resp.status_code == 200
    assert good_payload["summary"]["added_events"] == 2
    assert good_payload["message"] == "导入完成：新增 2 条现场记录。"
    assert [row["event_type"] for row in events] == ["start", "finish"]
    assert events[0]["created_by"] == "未填写反馈人"


def test_actual_import_allows_total_quantity_over_planned_for_rework(tmp_path, monkeypatch) -> None:
    app, db_path = _build_app(tmp_path, monkeypatch)
    client = app.test_client()
    task_code = _task_code_from_template(client)
    over_file = _workbook_bytes(
        [
            {
                "任务识别码": task_code,
                "实际开工时间": "2026-05-01 08:00:00",
                "实际完工时间": "2026-05-01 09:00:00",
                "完成数量": 8,
                "报废数量": 5,
            }
        ]
    )

    preview = _json(_post_preview(client, over_file))["data"]
    assert preview["can_confirm"] is True

    import_resp = _post_import(client, over_file)
    assert import_resp.status_code == 200
    events = _events_for_op(db_path, 10)
    assert [row["event_type"] for row in events] == ["start", "finish"]
    assert events[1]["quantity_done"] == 8
    assert events[1]["quantity_scrapped"] == 5


def test_actual_import_confirm_writes_only_after_clean_preview_and_rejects_errors(tmp_path, monkeypatch) -> None:
    app, db_path = _build_app(tmp_path, monkeypatch)
    client = app.test_client()
    task_code = _task_code_from_template(client)
    good_file = _workbook_bytes(
        [
            {
                "任务识别码": task_code,
                "实际开工时间": "2026-05-01 08:00:00",
                "实际完工时间": "2026-05-01 09:00:00",
                "完成数量": 10,
                "反馈人": "",
            }
        ],
        pause_rows=[
            {
                "任务识别码": task_code,
                "暂停开始时间": "2026-05-01 08:20:00",
                "暂停时长分钟": 10,
                "暂停原因": "设备问题",
                "暂停说明": "设备点检",
            }
        ],
    )
    preview = _json(_post_preview(client, good_file))["data"]

    assert preview["can_confirm"] is True
    assert preview["summary"]["add_count"] == 4
    assert _event_count(db_path) == 0

    confirm = client.post(
        f"/scheduler/resource-dispatch/execution/import/confirm?{_current_query()}",
        json={"preview_token": preview["preview_token"], "raw_rows": preview["raw_rows"]},
    )
    confirm_payload = _json(confirm)["data"]
    events = _events_for_op(db_path, 10)

    assert confirm.status_code == 200
    assert confirm_payload["summary"]["added_events"] == 4
    assert [row["event_type"] for row in events] == ["start", "pause", "resume", "finish"]
    assert events[0]["created_by"] == "未填写反馈人"

    bad_file = _workbook_bytes([{"任务识别码": task_code, "实际开工时间": "bad-time"}])
    bad_preview = _json(_post_preview(client, bad_file))["data"]
    bad_confirm = client.post(
        f"/scheduler/resource-dispatch/execution/import/confirm?{_current_query()}",
        json={"preview_token": bad_preview["preview_token"], "raw_rows": bad_preview["raw_rows"]},
    )

    assert bad_preview["can_confirm"] is False
    assert bad_confirm.status_code == 400
    assert _event_count(db_path) == 4


def test_actual_import_confirm_handles_datetime_cells_round_trip(tmp_path, monkeypatch) -> None:
    """日期单元格是真正的 datetime 对象时，预览→确认两步的幂等 token 仍要一致。

    回归 P3：openpyxl 把日期读成 datetime，经接口下发再回传后序列化格式会变，曾导致 confirm
    误报“导入内容已经变化”。读取源头归一成字符串后，两次 token 应一致、确认能正常写入。
    """
    app, db_path = _build_app(tmp_path, monkeypatch)
    client = app.test_client()
    task_code = _task_code_from_template(client)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "任务反馈"
    headers = [
        "任务识别码", "批次", "工序", "计划设备", "计划人员",
        "实际开工时间", "实际完工时间", "完成数量", "报废数量",
        "异常时间", "异常原因", "异常严重程度", "异常说明", "反馈人", "备注",
    ]
    ws.append(headers)
    row = {
        "任务识别码": task_code,
        "实际开工时间": datetime(2026, 5, 1, 8, 0, 0),
        "实际完工时间": datetime(2026, 5, 1, 9, 0, 0),
        "完成数量": 10,
    }
    ws.append([row.get(header) for header in headers])
    wb.create_sheet("暂停明细").append(
        ["任务识别码", "暂停开始时间", "暂停结束时间", "暂停时长分钟", "暂停原因", "暂停说明", "反馈人"]
    )
    buffer = io.BytesIO()
    wb.save(buffer)
    file_bytes = buffer.getvalue()

    preview = _json(_post_preview(client, file_bytes))["data"]
    assert preview["can_confirm"] is True

    confirm = client.post(
        f"/scheduler/resource-dispatch/execution/import/confirm?{_current_query()}",
        json={"preview_token": preview["preview_token"], "raw_rows": preview["raw_rows"]},
    )

    assert confirm.status_code == 200, _json(confirm)
    assert _json(confirm)["data"]["summary"]["added_events"] == 2
    assert [row["event_type"] for row in _events_for_op(db_path, 10)] == ["start", "finish"]


def test_actual_record_and_import_reject_non_current_official_plan(tmp_path, monkeypatch) -> None:
    app, db_path = _build_app(tmp_path, monkeypatch)
    client = app.test_client()
    card = _current_card(client)
    import_file = _workbook_bytes([{"任务识别码": "任意任务", "实际开工时间": "2026-05-01 08:00:00"}])
    old_plan_query = (
        "scope_type=operator&operator_id=O1&period_preset=week&query_date=2026-05-01"
        "&date_from=2026-05-01&date_to=2026-05-07&version=1&plan_role=adopted"
    )

    direct = client.post(
        f"/scheduler/resource-dispatch/execution/{card['op_id']}/actual?{old_plan_query}",
        json=_base_payload(
            card,
            idempotency_key="actual-reject-history",
            version=1,
            actual_start_time="2026-05-01 08:00:00",
        ),
    )
    template = client.get(
        f"/scheduler/resource-dispatch/execution/actual-template?{old_plan_query}"
    )
    preview = client.post(
        f"/scheduler/resource-dispatch/execution/import/preview?{old_plan_query}",
        data={"file": (io.BytesIO(import_file), "actual.xlsx")},
        content_type="multipart/form-data",
    )
    direct_import = client.post(
        f"/scheduler/resource-dispatch/execution/import?{old_plan_query}",
        data={"file": (io.BytesIO(import_file), "actual.xlsx")},
        content_type="multipart/form-data",
    )
    confirm = client.post(
        f"/scheduler/resource-dispatch/execution/import/confirm?{old_plan_query}",
        json={
            "preview_token": "old-plan",
            "raw_rows": [{"sheet": "任务反馈", "row_number": 2, "任务识别码": "任意任务", "实际开工时间": "2026-05-01 08:00:00"}],
        },
    )

    assert direct.status_code == 409
    assert _json(direct)["error"]["details"]["reason"] == "not_current_official_plan"
    assert template.status_code == 409
    assert preview.status_code == 409
    assert _json(preview)["error"]["details"]["reason"] == "not_current_official_plan"
    assert direct_import.status_code == 409
    assert _json(direct_import)["error"]["details"]["reason"] == "not_current_official_plan"
    assert confirm.status_code == 409
    assert _json(confirm)["error"]["details"]["reason"] == "not_current_official_plan"
    assert _event_count(db_path) == 0


def test_actual_record_and_import_reject_candidate_and_scenario_plans(tmp_path, monkeypatch) -> None:
    app, db_path = _build_app(tmp_path, monkeypatch)
    client = app.test_client()
    card = _current_card(client)
    import_file = _workbook_bytes([{"任务识别码": "任意任务", "实际开工时间": "2026-05-01 08:00:00"}])
    cases = (
        (
            "candidate",
            {"requested_plan_role": "baseline_best", "effective_plan_role": "adopted", "source_table": "schedule"},
            "scope_type=operator&operator_id=O1&period_preset=week&query_date=2026-05-01&date_from=2026-05-01&date_to=2026-05-07&version=2&plan_role=baseline_best",
        ),
        (
            "scenario",
            {"requested_plan_role": "adopted", "scenario_id": "scenario-plain"},
            "scope_type=operator&operator_id=O1&period_preset=week&query_date=2026-05-01&date_from=2026-05-01&date_to=2026-05-07&version=2&plan_role=adopted&scenario_id=scenario-plain",
        ),
    )

    for label, overrides, query in cases:
        direct = client.post(
            f"/scheduler/resource-dispatch/execution/{card['op_id']}/actual?{query}",
            json=_base_payload(
                card,
                idempotency_key=f"actual-reject-{label}",
                actual_start_time="2026-05-01 08:00:00",
                **overrides,
            ),
        )
        template = client.get(f"/scheduler/resource-dispatch/execution/actual-template?{query}")
        preview = client.post(
            f"/scheduler/resource-dispatch/execution/import/preview?{query}",
            data={"file": (io.BytesIO(import_file), "actual.xlsx")},
            content_type="multipart/form-data",
        )
        direct_import = client.post(
            f"/scheduler/resource-dispatch/execution/import?{query}",
            data={"file": (io.BytesIO(import_file), "actual.xlsx")},
            content_type="multipart/form-data",
        )
        confirm = client.post(
            f"/scheduler/resource-dispatch/execution/import/confirm?{query}",
            json={
                "preview_token": f"reject-{label}",
                "raw_rows": [{"sheet": "任务反馈", "row_number": 2, "任务识别码": "任意任务", "实际开工时间": "2026-05-01 08:00:00"}],
            },
        )

        assert direct.status_code == 409
        assert _json(direct)["error"]["details"]["reason"] == "not_current_official_plan"
        assert template.status_code == 409
        assert preview.status_code == 409
        assert _json(preview)["error"]["details"]["reason"] == "not_current_official_plan"
        assert direct_import.status_code == 409
        assert _json(direct_import)["error"]["details"]["reason"] == "not_current_official_plan"
        assert confirm.status_code == 409
        assert _json(confirm)["error"]["details"]["reason"] == "not_current_official_plan"
    assert _event_count(db_path) == 0
