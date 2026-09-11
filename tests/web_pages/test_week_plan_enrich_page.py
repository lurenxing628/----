"""周计划增强页面契约（fusion-week-plan-enrich）：列渲染/汇总条/空周升级/Excel 列。

钉死点：现场状态列出现在预览表与 Excel；每日汇总条含容量来源明示文案；
空周+版本有计划行 → 升级文案+跳转链接（URL 带日期与版本身份）；
无计划行版本 → 现有文案零跳转；禁直调 calculations.capacity_hours（grep 钉死）。
"""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from urllib.parse import parse_qs, urlsplit

from tests._support.legacy_report_contract import assert_retired, business_rows, get_unchanged

REPO_ROOT = Path(__file__).resolve().parents[2]


def _week_model(client, db_path, week_start):
    from core.infrastructure.database import get_connection
    from core.services.report import ReportEngine
    from core.services.scheduler.gantt_service import GanttService
    from core.services.scheduler.schedule_plan_query_service import SchedulePlanQueryService
    from core.services.scheduler.week_plan_daily_summary import build_week_plan_daily_summary
    from web.routes.domains.scheduler.scheduler_week_plan_preview import (
        build_week_plan_preview_state,
        week_plan_span_jump,
    )

    response = get_unchanged(client, f"/scheduler/week-plan?version=7&week_start={week_start}", db_path)
    assert_retired(response, public=("7", "正式采用方案"))
    before = business_rows(db_path)
    conn = get_connection(db_path)
    try:
        data = GanttService(conn).get_week_plan_rows(version=7, week_start=week_start)
        services = SimpleNamespace(schedule_plan_query_service=SchedulePlanQueryService(conn))
        with client.application.app_context():
            preview = build_week_plan_preview_state(data, span_jump=week_plan_span_jump(services, data))
        daily = build_week_plan_daily_summary(data["daily_planned_minutes"], calendar=ReportEngine(conn).calendar,
                                               week_start=data["week_start"], week_end=data["week_end"])
    finally:
        conn.close()
    assert business_rows(db_path) == before
    return data, preview, daily


def _seed_week_plan(db_path: str, *, with_schedule: bool = True) -> None:
    from core.infrastructure.database import get_connection

    conn = get_connection(db_path)
    conn.execute(
        "INSERT INTO ScheduleHistory (version, schedule_time, strategy, batch_count, op_count, result_status, result_summary, created_by)"
        " VALUES (7, '2026-06-01 08:00:00', 'weighted', 1, 1, 'success', '{}', 'pytest')"
    )
    if with_schedule:
        conn.execute("INSERT INTO Parts (part_no, part_name) VALUES ('P1', '零件1')")
        conn.execute("INSERT INTO Batches (batch_id, part_no, quantity) VALUES ('B1', 'P1', 10)")
        conn.execute(
            "INSERT INTO BatchOperations (op_code, batch_id, seq, op_type_name) VALUES ('B1-10', 'B1', 10, '车')"
        )
        op_id = conn.execute("SELECT id FROM BatchOperations WHERE op_code='B1-10'").fetchone()[0]
        conn.execute(
            "INSERT INTO Schedule (op_id, start_time, end_time, version)"
            " VALUES (?, '2026-06-01 08:00:00', '2026-06-01 12:00:00', 7)",
            (op_id,),
        )
    conn.commit()
    conn.close()


def test_week_plan_renders_execution_status_column_and_daily_summary(app_client, db_env):
    _seed_week_plan(db_env)
    data, preview, daily = _week_model(app_client, db_env, "2026-06-01")
    assert len(preview["preview_rows"]) == 1
    assert preview["preview_rows"][0]["现场状态"] == "待开工"
    assert data["daily_planned_minutes"] == {"2026-06-01": 240}
    assert len(daily) == 1
    assert daily[0]["date"] == "2026-06-01"
    assert daily[0]["planned_hours_label"] == "4 小时"
    assert daily[0]["capacity_hours_label"] != "-"
    assert "machine_id" not in daily[0] and "operator_id" not in daily[0]
    assert not (REPO_ROOT / "templates/scheduler/week_plan.html").exists()


def test_week_plan_empty_week_upgraded_hint_with_jump_link(app_client, db_env):
    _seed_week_plan(db_env)
    data, preview, daily = _week_model(app_client, db_env, "2026-07-06")
    assert data["rows"] == [] and daily == []
    assert "这个版本的计划在 2026-06-01 ～ 2026-06-01" in preview["empty_message"]
    link = preview["empty_jump_link"]
    assert link["label"] == "跳到计划区间"
    query = parse_qs(urlsplit(link["url"]).query)
    assert query["version"] == ["7"] and query["plan_role"] == ["adopted"]
    assert query["week_start"] == ["2026-06-01"]


def test_week_plan_empty_week_no_plan_rows_keeps_plain_message(app_client, db_env):
    _seed_week_plan(db_env, with_schedule=False)  # 有历史但无计划行（失败/模拟）
    data, preview, daily = _week_model(app_client, db_env, "2026-07-06")
    assert data["rows"] == [] and daily == []
    assert "暂无数据" in preview["empty_message"]
    assert preview["empty_jump_link"] is None


def test_week_plan_export_contains_execution_status_column(app_client, db_env):
    from openpyxl import load_workbook

    _seed_week_plan(db_env)
    resp = app_client.get("/scheduler/week-plan/export?version=7&week_start=2026-06-01")
    assert resp.status_code == 200
    wb = load_workbook(BytesIO(resp.data))
    ws = wb.active
    headers = [c.value for c in ws[1]]
    assert headers == ["日期", "批次号", "图号", "工序", "设备", "人员", "时段", "现场状态"]
    assert ws.cell(row=2, column=8).value == "待开工"
    # 导出不含每日合计行（合计是页面阅读辅助，不破坏导出行语义）
    all_values = [str(c.value or "") for row in ws.iter_rows() for c in row]
    assert not any("容量" in v for v in all_values)


def test_no_direct_capacity_hours_call_in_week_plan_chain():
    # 4.6 红线：周计划链路禁直调 calculations.capacity_hours（midnight 采样错归属）
    for rel in (
        "core/services/scheduler/week_plan_daily_summary.py",
        "core/services/scheduler/gantt_week_plan.py",
        "web/routes/domains/scheduler/scheduler_week_plan.py",
        "web/routes/domains/scheduler/scheduler_week_plan_preview.py",
    ):
        text = (REPO_ROOT / rel).read_text(encoding="utf-8")
        assert "capacity_hours(" not in text.replace("_capacity_hours_at_noon(", ""), rel
