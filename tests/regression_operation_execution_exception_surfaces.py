from __future__ import annotations

import io

import openpyxl

from core.infrastructure.database import get_connection
from core.services.scheduler.resource_dispatch_excel import build_resource_dispatch_workbook
from core.services.scheduler.resource_dispatch_service import ResourceDispatchService
from tests.operation_execution_feedback_test_support import _build_app, _current_card, _json
from tests.regression_operation_execution_exception_feedback import _post
from web.viewmodels.scheduler_resource_dispatch import decorate_resource_dispatch_payload


def test_exception_details_are_visible_in_detail_rows_gantt_popup_source_and_export(tmp_path, monkeypatch) -> None:
    app, db_path = _build_app(tmp_path, monkeypatch)
    client = app.test_client()
    card = _current_card(client)
    start_data = _json(_post(client, card, "start", idempotency_key="surface-start"))["data"]
    _post(
        client,
        card,
        "report-exception",
        expected_state_revision=start_data["state_revision"],
        event_time="2026-05-01 08:30:00",
        idempotency_key="surface-report-exception",
        reason_code="equipment",
        severity="critical",
        impact_minutes="",
        affected_machine_id="M2",
        affected_operator_id="O2",
        handling_status="waiting",
        suggest_reschedule=False,
        remark="等待维修",
    )

    data_resp = client.get(
        "/scheduler/resource-dispatch/data?scope_type=operator&operator_id=O1&period_preset=week&query_date=2026-05-01&version=2&plan_role=adopted"
    )
    detail_row = _json(data_resp)["data"]["detail_rows"][0]
    assert detail_row["execution_status_label"] == "异常中"
    assert detail_row["latest_exception_reason_label"] == "设备问题"
    assert detail_row["latest_exception_severity_label"] == "紧急"
    assert detail_row["latest_exception_impact_minutes_label"] == "暂时不知道影响多久"
    assert detail_row["latest_exception_affected_machine_label"] == "M2 二号设备"
    assert detail_row["latest_exception_affected_operator_label"] == "O2 李四"
    assert detail_row["latest_exception_handling_status_label"] == "等待条件"
    assert detail_row["latest_exception_suggest_reschedule_label"] == "暂不建议重新排程"
    assert detail_row["latest_exception_remark"] == "等待维修"

    db_conn = get_connection(db_path)
    try:
        payload = ResourceDispatchService(db_conn, logger=None, op_logger=None).get_dispatch_payload(
            scope_type="operator",
            operator_id="O1",
            period_preset="week",
            query_date="2026-05-01",
            version=2,
            plan_role="adopted",
        )
        payload = decorate_resource_dispatch_payload(payload)
        buffer = build_resource_dispatch_workbook(payload)
    finally:
        db_conn.close()

    wb = openpyxl.load_workbook(io.BytesIO(buffer.getvalue()))
    ws = wb["任务明细"]
    headers = [cell.value for cell in ws[1]]
    values = [cell.value for cell in ws[2]]
    row_map = dict(zip(headers, values))
    assert row_map["现场状态"] == "异常中"
    assert row_map["最近异常原因"] == "设备问题"
    assert row_map["严重程度"] == "紧急"
    assert row_map["预计影响时间"] == "暂时不知道影响多久"
    assert row_map["影响设备"] == "二号设备\n完整身份：M2 二号设备"
    assert row_map["影响人员"] == "李四\n完整身份：O2 李四"
    assert row_map["处理状态"] == "等待条件"
    assert row_map["是否建议重排"] == "暂不建议重新排程"
    assert row_map["情况说明"] == "等待维修"

