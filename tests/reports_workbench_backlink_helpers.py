from __future__ import annotations

import html
import importlib
import json
import os
import sys
import tempfile
from html.parser import HTMLParser
from io import BytesIO
from pathlib import Path
from typing import Dict, List
from urllib.parse import parse_qs, urlparse

from core.infrastructure.database import ensure_schema, get_connection

REPO_ROOT = Path(__file__).resolve().parents[1]
SCHEMA_PATH = REPO_ROOT / "schema.sql"

INTERNAL_VISIBLE_TOKENS = (
    "plan_role",
    "scenario_id",
    "source_table",
    "candidate_id",
    "op_id",
    "schedule_id",
)
INTERNAL_PRIVATE_TOKENS = (
    "source_table",
    "candidate_id",
    "op_id",
    "schedule_id",
)
_URL_DATA_ATTRIBUTE_NAMES = {
    "data-url",
    "data-export-url",
    "data-execution-url",
    "data-actual-record-url-template",
    "data-actual-template-url",
    "data-actual-import-url",
}


def _is_url_data_attribute(name: str) -> bool:
    return str(name or "").lower() in _URL_DATA_ATTRIBUTE_NAMES


class _PageParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self._href = ""
        self._text: List[str] = []
        self._link_attrs: Dict[str, str] = {}
        self.links: List[Dict[str, str]] = []
        self.inputs: List[Dict[str, str]] = []
        self.visible_parts: List[str] = []
        self.public_payload_parts: List[str] = []
        self.public_attribute_parts: List[str] = []
        self._capture_public_json = False

    def handle_starttag(self, tag: str, attrs) -> None:
        attr_map = dict(attrs)
        self._capture_public_attributes(tag, attrs)
        self._start_script_capture(tag, attr_map)
        self._capture_input(tag, attr_map)
        self._start_anchor(tag, attr_map)

    def _capture_public_attributes(self, tag: str, attrs) -> None:
        for name, value in attrs:
            self.public_attribute_parts.append(f"{tag}.{name}={value or ''}")
            if str(name or "").lower().startswith("data-"):
                if not _is_url_data_attribute(str(name)):
                    self.public_payload_parts.append(f"{name}={value or ''}")

    def _start_script_capture(self, tag: str, attr_map: Dict[str, str]) -> None:
        if tag == "script":
            script_type = str(attr_map.get("type") or "").lower()
            self._capture_public_json = "json" in script_type

    def _capture_input(self, tag: str, attr_map: Dict[str, str]) -> None:
        if tag == "input":
            self.inputs.append({str(key): str(value or "") for key, value in attr_map.items()})

    def _start_anchor(self, tag: str, attr_map: Dict[str, str]) -> None:
        if tag == "a":
            self._href = str(attr_map.get("href") or "")
            self._link_attrs = {str(key): str(value or "") for key, value in attr_map.items()}
            self._text = []

    def handle_data(self, data: str) -> None:
        text = str(data or "").strip()
        if not text:
            return
        if self._capture_public_json:
            self.public_payload_parts.append(text)
        self.visible_parts.append(text)
        if self._href:
            self._text.append(text)

    def handle_endtag(self, tag: str) -> None:
        if tag == "a" and self._href:
            item = {"href": html.unescape(self._href), "text": " ".join(self._text)}
            item.update(self._link_attrs)
            self.links.append(item)
            self._href = ""
            self._link_attrs = {}
            self._text = []
        elif tag == "script":
            self._capture_public_json = False


def _prepare_env(tmpdir: str) -> str:
    db_path = str(Path(tmpdir) / "aps_reports_workbench.db")
    os.environ["APS_ENV"] = "development"
    os.environ["APS_DB_PATH"] = db_path
    os.environ["APS_LOG_DIR"] = str(Path(tmpdir) / "logs")
    os.environ["APS_BACKUP_DIR"] = str(Path(tmpdir) / "backups")
    os.environ["APS_EXCEL_TEMPLATE_DIR"] = str(Path(tmpdir) / "templates_excel")
    os.environ["SECRET_KEY"] = "reports-workbench-backlink"
    Path(os.environ["APS_LOG_DIR"]).mkdir(exist_ok=True)
    Path(os.environ["APS_BACKUP_DIR"]).mkdir(exist_ok=True)
    Path(os.environ["APS_EXCEL_TEMPLATE_DIR"]).mkdir(exist_ok=True)
    return db_path


def _load_app():
    if str(REPO_ROOT) not in sys.path:
        sys.path.insert(0, str(REPO_ROOT))
    for name in list(sys.modules):
        if (
            name == "app"
            or name.startswith("web.bootstrap.entrypoint")
            or name.startswith("web.bootstrap.factory")
            or name.startswith("web.routes.scheduler")
            or name.startswith("web.routes.domains.scheduler")
        ):
            sys.modules.pop(name, None)
    app_mod = importlib.import_module("app")
    return app_mod.create_app()


def _seed_reports_data(db_path: str) -> None:
    summary = {"algo": {"metrics": {"machine_util_avg": 0.8}}, "overdue_batches": {"count": 1}}
    conn = get_connection(db_path)
    try:
        conn.execute("INSERT INTO Parts (part_no, part_name) VALUES (?, ?)", ("P-RPT", "报表回跳零件"))
        conn.execute("INSERT INTO Machines (machine_id, name, status) VALUES (?, ?, ?)", ("M-RPT", "一号设备", "active"))
        conn.execute("INSERT INTO Machines (machine_id, name, status) VALUES (?, ?, ?)", ("M-OTHER", "二号设备", "active"))
        conn.execute("INSERT INTO Operators (operator_id, name, status) VALUES (?, ?, ?)", ("O-RPT", "张三", "active"))
        conn.execute("INSERT INTO Operators (operator_id, name, status) VALUES (?, ?, ?)", ("O-OTHER", "李四", "active"))
        conn.execute(
            """
            INSERT INTO Batches (batch_id, part_no, part_name, quantity, due_date, priority, status)
            VALUES (?, ?, ?, ?, ?, ?, ?)
            """,
            ("B-RPT", "P-RPT", "报表回跳零件", 1, "2026-05-01", "normal", "scheduled"),
        )
        cur = conn.execute(
            """
            INSERT INTO BatchOperations (op_code, batch_id, seq, op_type_name, source, status)
            VALUES (?, ?, ?, ?, ?, ?)
            """,
            ("OP-RPT", "B-RPT", 1, "加工", "internal", "pending"),
        )
        op_id = int(cur.lastrowid)
        conn.execute(
            """
            INSERT INTO Batches (batch_id, part_no, part_name, quantity, due_date, priority, status)
            VALUES (?, ?, ?, ?, ?, ?, ?)
            """,
            ("B-OTHER", "P-RPT", "报表回跳零件", 1, "2026-05-01", "normal", "scheduled"),
        )
        cur_other = conn.execute(
            """
            INSERT INTO BatchOperations (op_code, batch_id, seq, op_type_name, source, status)
            VALUES (?, ?, ?, ?, ?, ?)
            """,
            ("OP-OTHER", "B-OTHER", 1, "加工", "internal", "pending"),
        )
        other_op_id = int(cur_other.lastrowid)
        conn.execute(
            """
            INSERT INTO Batches (batch_id, part_no, part_name, quantity, due_date, priority, status)
            VALUES (?, ?, ?, ?, ?, ?, ?)
            """,
            ("B-SAME", "P-RPT", "报表回跳零件", 1, "2026-06-01", "normal", "scheduled"),
        )
        cur_same = conn.execute(
            """
            INSERT INTO BatchOperations (op_code, batch_id, seq, op_type_name, source, status)
            VALUES (?, ?, ?, ?, ?, ?)
            """,
            ("OP-SAME", "B-SAME", 1, "加工", "internal", "pending"),
        )
        same_op_id = int(cur_same.lastrowid)
        conn.execute(
            """
            INSERT INTO Schedule (op_id, machine_id, operator_id, start_time, end_time, lock_status, version)
            VALUES (?, ?, ?, ?, ?, ?, ?)
            """,
            (op_id, "M-RPT", "O-RPT", "2026-05-06 08:00:00", "2026-05-06 12:00:00", "unlocked", 12),
        )
        conn.execute(
            """
            INSERT INTO Schedule (op_id, machine_id, operator_id, start_time, end_time, lock_status, version)
            VALUES (?, ?, ?, ?, ?, ?, ?)
            """,
            (other_op_id, "M-OTHER", "O-OTHER", "2026-05-06 08:00:00", "2026-05-06 12:00:00", "unlocked", 12),
        )
        conn.execute(
            """
            INSERT INTO Schedule (op_id, machine_id, operator_id, start_time, end_time, lock_status, version)
            VALUES (?, ?, ?, ?, ?, ?, ?)
            """,
            (same_op_id, "M-RPT", "O-RPT", "2026-05-06 13:00:00", "2026-05-06 14:00:00", "unlocked", 12),
        )
        conn.execute(
            """
            INSERT INTO MachineDowntimes (machine_id, start_time, end_time, reason_code, reason_detail, status)
            VALUES (?, ?, ?, ?, ?, ?)
            """,
            ("M-RPT", "2026-05-06 09:00:00", "2026-05-06 10:00:00", "maintenance", "保养", "active"),
        )
        conn.execute(
            """
            INSERT INTO MachineDowntimes (machine_id, start_time, end_time, reason_code, reason_detail, status)
            VALUES (?, ?, ?, ?, ?, ?)
            """,
            ("M-OTHER", "2026-05-06 09:00:00", "2026-05-06 09:30:00", "maintenance", "二号设备保养", "active"),
        )
        conn.execute(
            """
            INSERT INTO MachineDowntimes (machine_id, start_time, end_time, reason_code, reason_detail, status)
            VALUES (?, ?, ?, ?, ?, ?)
            """,
            ("M-RPT", "2026-05-06 13:15:00", "2026-05-06 13:45:00", "maintenance", "同设备别的批次保养", "active"),
        )
        conn.execute(
            """
            INSERT INTO ScheduleAdjustmentScenario (
              scenario_id, source_draft_id, base_version, base_plan_role, base_source_table,
              scenario_name, status, validation_status, row_count, created_by
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            ("SCENARIO-RPT", "DRAFT-RPT", 12, "adopted", "schedule", "测试模拟方案", "active", "valid", 1, "pytest"),
        )
        conn.execute(
            """
            INSERT INTO ScheduleAdjustmentScenarioRow (
              scenario_id, source_table, op_id, machine_id, operator_id, start_time, end_time, lock_status, is_changed
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                "SCENARIO-RPT",
                "schedule",
                op_id,
                "M-RPT",
                "O-RPT",
                "2026-05-06 08:00:00",
                "2026-05-06 12:00:00",
                "unlocked",
                "no",
            ),
        )
        conn.execute(
            """
            INSERT INTO ScheduleHistory (version, strategy, batch_count, op_count, result_status, result_summary, created_by)
            VALUES (?, ?, ?, ?, ?, ?, ?)
            """,
            (12, "priority_first", 1, 1, "success", json.dumps(summary, ensure_ascii=False), "pytest"),
        )
        conn.commit()
    finally:
        conn.close()


def _client():
    tmpdir = tempfile.mkdtemp(prefix="aps_reports_workbench_")
    db_path = _prepare_env(tmpdir)
    ensure_schema(db_path, logger=None, schema_path=str(SCHEMA_PATH), backup_dir=None)
    _seed_reports_data(db_path)
    app = _load_app()
    client = app.test_client()
    try:
        client.set_cookie("aps_ui_mode", "v1", domain="localhost")
    except TypeError:
        client.set_cookie("localhost", "aps_ui_mode", "v1")
    return client


def _parser_for(client, path: str) -> _PageParser:
    parser = _PageParser()
    parser.feed(_html_for(client, path))
    return parser


def _html_for(client, path: str) -> str:
    resp = client.get(path)
    assert resp.status_code == 200, path
    return resp.get_data(as_text=True)


def _visible_text(parser: _PageParser) -> str:
    return "\n".join(parser.visible_parts)


def _href_with_text(parser: _PageParser, text: str, path: str) -> str:
    matches = []
    for link in parser.links:
        if text in link["text"] and urlparse(link["href"]).path == path:
            matches.append(link["href"])
    for href in matches:
        if "version=12" in href and "plan_role=adopted" in href:
            return href
    for href in matches:
        if "version=12" in href:
            return href
    if matches:
        return matches[0]
    raise AssertionError(f"没有找到链接：{text} -> {path}，links={parser.links!r}")


def _href_with_text_and_fragment(parser: _PageParser, text: str, path: str, fragment: str) -> str:
    for link in parser.links:
        if text in link["text"] and urlparse(link["href"]).path == path and fragment in link["href"]:
            return link["href"]
    raise AssertionError(f"没有找到带 {fragment!r} 的链接：{text} -> {path}，links={parser.links!r}")


def _href_with_text_and_class(parser: _PageParser, text: str, path: str, class_name: str, *, exact: bool = False) -> str:
    matches = []
    for link in parser.links:
        classes = set(str(link.get("class") or "").split())
        link_text = link["text"]
        text_matches = link_text == text if exact else text in link_text
        if class_name in classes and text_matches and urlparse(link["href"]).path == path:
            matches.append(link["href"])
    if len(matches) == 1:
        return matches[0]
    raise AssertionError(f"没有找到唯一链接：{text} -> {path} class={class_name}，matches={matches!r}，links={parser.links!r}")


_REPORT_CARD_TITLES = {
    "overdue": "超期清单",
    "utilization": "资源负荷与利用率",
    "execution_review": "计划和现场实际",
    "downtime": "停机影响统计",
}


def _href_for_report_card(parser: _PageParser, card_key: str, path: str) -> str:
    card_title = _REPORT_CARD_TITLES.get(card_key, card_key)
    for link in parser.links:
        if link.get("data-report-card") == card_title and urlparse(link["href"]).path == path:
            return link["href"]
    raise AssertionError(f"没有找到报表卡片链接：{card_key} -> {path}，links={parser.links!r}")


def _query(href: str) -> Dict[str, List[str]]:
    return parse_qs(urlparse(href).query)


def _input_values(parser: _PageParser) -> Dict[str, List[str]]:
    values: Dict[str, List[str]] = {}
    for item in parser.inputs:
        name = item.get("name") or ""
        if not name:
            continue
        values.setdefault(name, []).append(item.get("value") or "")
    return values


def _assert_public_visible_text(parser: _PageParser) -> None:
    visible = _visible_text(parser)
    for token in INTERNAL_VISIBLE_TOKENS:
        assert token not in visible


def _assert_public_payload(parser: _PageParser) -> None:
    payload = "\n".join(parser.public_payload_parts)
    for token in INTERNAL_VISIBLE_TOKENS:
        assert token not in payload


def _assert_public_attributes(parser: _PageParser) -> None:
    payload = "\n".join(parser.public_attribute_parts)
    for token in INTERNAL_PRIVATE_TOKENS:
        assert token not in payload


def _assert_public_links_and_forms(parser: _PageParser) -> None:
    public_url_and_form_payload = "\n".join(
        [link["href"] for link in parser.links]
        + [f"{item.get('name', '')}={item.get('value', '')}" for item in parser.inputs]
    )
    for token in INTERNAL_PRIVATE_TOKENS:
        assert token not in public_url_and_form_payload


def _assert_public_output_boundaries(parser: _PageParser) -> None:
    _assert_public_visible_text(parser)
    _assert_public_payload(parser)
    _assert_public_attributes(parser)
    _assert_public_links_and_forms(parser)


def _assert_date_from_to(query: Dict[str, List[str]]) -> None:
    assert query["date_from"] == ["2026-05-06"]
    assert query["date_to"] == ["2026-05-06"]


def _assert_start_end(query: Dict[str, List[str]]) -> None:
    assert query["start_date"] == ["2026-05-06"]
    assert query["end_date"] == ["2026-05-06"]


def _assert_export_headers_hide_internal_tokens(client) -> None:
    import openpyxl

    export_paths = (
        "/reports/overdue/export?version=12&plan_role=adopted",
        "/reports/utilization/export?version=12&plan_role=adopted&start_date=2026-05-06&end_date=2026-05-06",
        "/reports/execution-review/export?version=12&date_from=2026-05-06&date_to=2026-05-06",
        "/reports/downtime/export?version=12&plan_role=adopted&start_date=2026-05-06&end_date=2026-05-06",
    )
    for path in export_paths:
        resp = client.get(path)
        assert resp.status_code == 200, path
        wb = openpyxl.load_workbook(BytesIO(resp.data), read_only=True, data_only=True)
        try:
            workbook_text_parts = [ws.title for ws in wb.worksheets]
            for ws in wb.worksheets:
                first_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
                header_text = "\n".join(str(value or "") for value in first_row)
                for token in INTERNAL_VISIBLE_TOKENS:
                    assert token not in header_text, (path, ws.title, header_text)
                for row in ws.iter_rows(values_only=True):
                    workbook_text_parts.extend(str(value or "") for value in row)
            workbook_text = "\n".join(workbook_text_parts)
            for token in INTERNAL_VISIBLE_TOKENS:
                assert token not in workbook_text, path
        finally:
            wb.close()


def _xlsx_text(data: bytes) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(BytesIO(data), read_only=True, data_only=True)
    try:
        parts = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                parts.extend(str(value or "") for value in row)
        return "\n".join(parts)
    finally:
        wb.close()


def _xlsx_sheet_rows(data: bytes, sheet_name: str) -> List[List[object]]:
    import openpyxl

    wb = openpyxl.load_workbook(BytesIO(data), read_only=True, data_only=True)
    try:
        ws = wb[sheet_name]
        return [list(row) for row in ws.iter_rows(values_only=True)]
    finally:
        wb.close()
