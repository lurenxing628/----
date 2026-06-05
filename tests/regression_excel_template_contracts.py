"""为 Excel 模板契约校验提供断言库：核对工艺/页面/转换输出模板文件的表头、示例行、下拉枚举与 get_template_definition 注册定义一致，校验页面说明与总说明书的推荐值/兼容旧写法措辞，并验证 ensure_excel_templates 刷新旧工种模板下拉时不覆盖用户数据、坏模板或表头不匹配时不静默重建。供 main() 与其它回归测试调用。"""

from __future__ import annotations

import importlib
import re
import sys
import tempfile
from pathlib import Path
from typing import Any, Dict, List, Mapping, Match, Sequence, Tuple

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation

REPO_ROOT = Path(__file__).resolve().parents[1]

PROCESS_TEMPLATE_FILES = (
    "工种配置.xlsx",
    "供应商配置.xlsx",
    "零件工艺路线.xlsx",
    "零件工序工时.xlsx",
)

PAGE_MANUAL_TEMPLATE_COLUMN_CONTRACTS = {
    "excel_personnel": "人员基本信息.xlsx",
    "excel_personnel_link": "人员设备关联.xlsx",
    "excel_personnel_calendar": "人员专属工作日历.xlsx",
    "excel_equipment": "设备信息.xlsx",
    "excel_equipment_link": "设备人员关联.xlsx",
    "excel_op_types": "工种配置.xlsx",
    "excel_suppliers": "供应商配置.xlsx",
    "excel_routes": "零件工艺路线.xlsx",
    "excel_part_op_hours": "零件工序工时.xlsx",
    "excel_batches": "批次信息.xlsx",
    "excel_calendar": "工作日历.xlsx",
}

RELATION_REVERSE_HEADER_COPY_CONTRACTS: Mapping[str, Mapping[str, Any]] = {
    "excel_personnel_link": {
        "filename": "人员设备关联.xlsx",
        "required_phrases": [
            "列：工号(必填)、设备编号(必填)、技能等级、主操设备。",
            "设备侧模板列顺序是 设备编号 / 工号",
        ],
    },
    "excel_equipment_link": {
        "filename": "设备人员关联.xlsx",
        "required_phrases": [
            "列：设备编号(必填)、工号(必填)、技能等级、主操设备。",
            "和人员侧导入的区别主要是列顺序：这里是“设备编号在前、工号在后”",
        ],
    },
}

DROPDOWN_MANUAL_CONTRACTS: Mapping[str, Mapping[str, Any]] = {
    "工种配置.xlsx": {
        "manual_id": "excel_op_types",
        "manual_heading": "#### 1.5.1 工种配置",
        "dropdowns": {"归属": ["自制", "外协"]},
        "compatible_values": {"归属": ["内部", "外部"]},
        "page_phrases": ["归属可填：自制 / 外协", "新文件请只填这两个中文选项"],
        "manual_phrases": ["填 `自制` 或 `外协`", "新文件请只填这两个中文值"],
        "forbidden_page_phrases": ["内部/外部", "内部 / 外部", "内部、外部", "内部，外部"],
        "forbidden_manual_phrases": ["`内部`", "`外部`", "内部/外部", "内部 / 外部", "内部、外部", "内部，外部"],
    },
    "供应商配置.xlsx": {
        "manual_id": "excel_suppliers",
        "manual_heading": "#### 1.5.2 供应商配置",
        "dropdowns": {"状态": ["启用", "停用"]},
        "compatible_values": {"状态": ["在用", "正常", "禁用"]},
        "page_phrases": ["启用/停用", "对应工种", "状态和备注"],
        "manual_phrases": ["`启用`/`停用`", "`在用`/`正常`/`禁用`", "状态", "备注"],
        "forbidden_page_phrases": ["新填数据请使用自制/外协", "新填数据请使用 `自制`/`外协`"],
        "forbidden_manual_phrases": [],
    },
    "人员基本信息.xlsx": {
        "manual_id": "excel_personnel",
        "manual_heading": "#### 1.5.5 人员基本信息",
        "dropdowns": {"状态": ["在岗", "停用"]},
        "compatible_values": {"状态": ["启用", "可用", "正常", "休假", "离岗"]},
        "page_phrases": ["状态", "在岗", "停用", "状态模板下拉优先选：在岗 / 停用"],
        "manual_phrases": ["`在岗` 或 `停用`"],
        "forbidden_page_phrases": ["在岗 / 停用 / 休假", "`在岗/停用/休假`", "模板下拉优先选：在岗 / 停用 / 休假"],
        "forbidden_manual_phrases": [],
    },
    "人员设备关联.xlsx": {
        "manual_id": "excel_personnel_link",
        "manual_heading": "#### 1.6.7 人员设备关联",
        "dropdowns": {"技能等级": ["初级", "普通", "熟练"], "主操设备": ["是", "否"]},
        "compatible_values": {"技能等级": ["新手", "一般", "中级", "高级", "专家"], "主操设备": ["主操", "非主操", "主", "非主", "1", "0"]},
        "page_phrases": ["新手/一般/中级/高级/专家", "主/非主", "1/0"],
        "manual_phrases": ["`新手`、`一般`、`中级`、`高级`、`专家`", "`主操`、`非主操`、`主`、`非主`、`1`、`0`"],
        "forbidden_page_phrases": [],
        "forbidden_manual_phrases": [],
    },
    "人员专属工作日历.xlsx": {
        "manual_id": "excel_personnel_calendar",
        "manual_heading": "#### 1.6.9 个人工作日历",
        "dropdowns": {"类型": ["工作日", "假期"], "允许普通件": ["是", "否"], "允许急件": ["是", "否"]},
        "compatible_values": {"类型": ["周末", "节假日"], "允许普通件": ["1", "0"], "允许急件": ["1", "0"]},
        "page_phrases": ["类型建议填：工作日 / 假期", "以前写过周末 / 节假日", "允许普通件/急件模板下拉优先填 是/否", "以前文件里写过 1/0"],
        "manual_phrases": ["旧文件里的 `周末`、`节假日`", "旧文件里的 `1/0`"],
        "forbidden_page_phrases": [
            "类型可填：工作日 / 假期 / 周末 / 节假日",
            "可填写：工作日 / 假期 / 周末 / 节假日",
        ],
        "forbidden_manual_phrases": [
            "类型可填：工作日 / 假期 / 周末 / 节假日",
            "可填写：工作日 / 假期 / 周末 / 节假日",
        ],
    },
    "设备信息.xlsx": {
        "manual_id": "excel_equipment",
        "manual_heading": "#### 1.5.6 设备信息",
        "dropdowns": {"状态": ["可用", "停用", "维修"]},
        "compatible_values": {"状态": ["启用", "正常", "禁用", "不可用", "维护", "维护中", "维修中", "保养"]},
        "page_phrases": ["状态", "可用", "停用", "维修"],
        "manual_phrases": ["`可用`、`停用`、`维修`"],
        "forbidden_page_phrases": [],
        "forbidden_manual_phrases": [],
    },
    "批次信息.xlsx": {
        "manual_id": "excel_batches",
        "manual_heading": "#### 1.5.10 批次信息（排产核心）",
        "dropdowns": {"优先级": ["普通", "急件", "特急"], "齐套": ["齐套", "未齐套", "部分齐套"]},
        "compatible_values": {"优先级": ["急"], "齐套": ["是", "否"]},
        "page_phrases": ["优先级", "普通", "急件", "特急", "齐套", "未齐套", "部分齐套"],
        "manual_phrases": ["`普通`、`急件`、`特急`", "`齐套`、`未齐套`、`部分齐套`"],
        "forbidden_page_phrases": [
            "优先级可填：普通 / 急件 / 特急 / 急",
            "齐套可填：齐套 / 未齐套 / 部分齐套 / 是 / 否",
            "新文件请按这些中文选项填写",
        ],
        "forbidden_manual_phrases": [],
    },
    "工作日历.xlsx": {
        "manual_id": "excel_calendar",
        "manual_heading": "#### 1.5.8 工作日历",
        "dropdowns": {"类型": ["工作日", "假期"], "允许普通件": ["是", "否"], "允许急件": ["是", "否"]},
        "compatible_values": {"类型": ["周末", "节假日"]},
        "page_phrases": ["类型", "工作日", "假期", "允许普通件", "允许急件", "是", "否"],
        "manual_phrases": ["`工作日`", "`假期`", "`是` 或 `否`"],
        "forbidden_page_phrases": [
            "类型可填：工作日 / 假期 / 周末 / 节假日",
            "可填写：工作日 / 假期 / 周末 / 节假日",
        ],
        "forbidden_manual_phrases": [
            "类型可填：工作日 / 假期 / 周末 / 节假日",
            "可填写：工作日 / 假期 / 周末 / 节假日",
        ],
    },
}


def _ensure_repo_on_path() -> None:
    repo_root = str(REPO_ROOT)
    if repo_root not in sys.path:
        sys.path.insert(0, repo_root)


def _template_definition(filename: str) -> Mapping[str, Any]:
    _ensure_repo_on_path()
    excel_templates = importlib.import_module("core.services.common.excel_templates")
    return excel_templates.get_template_definition(filename)


def _read_static_manual() -> str:
    return (REPO_ROOT / "static" / "docs" / "scheduler_manual.md").read_text(encoding="utf-8")


def _manual_payload_text(payload: Mapping[str, Any]) -> str:
    parts = [str(payload.get("title") or ""), str(payload.get("summary") or "")]
    help_card = payload.get("help_card") or {}
    parts.append(str(help_card.get("title") or ""))
    parts.extend(str(item or "") for item in (help_card.get("items") or []))
    for section in payload.get("sections") or []:
        parts.append(str(section.get("title") or ""))
        parts.append(str(section.get("body_md") or ""))
    return "\n".join(part for part in parts if part)


def _extract_markdown_section(markdown_text: str, heading: str) -> str:
    start = markdown_text.find(heading)
    if start < 0 and heading.startswith("#### "):
        heading_title = re.sub(r"^####\s+\d+(?:\.\d+)*\s+", "", heading).strip()
        pattern = rf"^####\s+\d+(?:\.\d+)*\s+{re.escape(heading_title)}\s*$"
        m = re.search(pattern, markdown_text, flags=re.M)
        if m:
            start = m.start()
            heading = m.group(0)
    assert start >= 0, f"总说明书缺少章节标题：{heading}"
    rest = markdown_text[start + len(heading) :]
    next_heading = re.search(r"^####\s+\d+(?:\.\d+)*\s+", rest, flags=re.M)
    end = start + len(heading) + next_heading.start() if next_heading else len(markdown_text)
    return markdown_text[start:end]


def _is_allowed_legacy_op_type_note(section_text: str, phrase: str) -> bool:
    return (
        phrase == "内部/外部"
        and "旧模板中已有数据行不会被系统擅自改写" in section_text
        and "新填数据请使用 `自制`/`外协`" in section_text
    )


def _compact_text(text: str) -> str:
    return re.sub(r"\s+", "", str(text or ""))


def _value_has_compatibility_context(text: str, value: str) -> bool:
    markers = ("旧文件", "旧模板", "以前", "兼容", "也能识别", "尽量读懂", "也可以填", "按假期")
    for match in re.finditer(re.escape(value), text):
        start = max(0, match.start() - 50)
        end = min(len(text), match.end() + 50)
        context = text[start:end]
        if any(marker in context for marker in markers):
            return True
    return False


def _assert_dropdown_copy_distinguishes_recommended_and_compatible(
    text: str,
    *,
    label: str,
    dropdowns: Mapping[str, Sequence[str]],
    compatible_values: Mapping[str, Sequence[str]],
) -> None:
    compact = _compact_text(text)
    for column_name, recommended_values in dropdowns.items():
        assert column_name in text, f"{label} 缺少下拉字段名：{column_name}"
        for value in recommended_values:
            assert value in text, f"{label} 缺少模板推荐下拉值：{column_name}={value}"
        joined = "/".join(str(value) for value in recommended_values)
        assert joined in compact or any(marker in text for marker in ("新文件", "新表", "新填数据", "推荐填", "建议填")), (
            f"{label} 应把 {column_name} 的推荐值和兼容旧写法分开说明：{joined}"
        )

    for column_name, values in compatible_values.items():
        for value in values:
            if value not in text:
                continue
            assert _value_has_compatibility_context(text, value), (
                f"{label} 里的兼容写法没有标清是旧写法或兼容写法：{column_name}={value}"
            )


def _definition_enum_values_by_header(definition: Mapping[str, Any]) -> Dict[str, List[str]]:
    headers = [str(item) for item in (definition.get("headers") or [])]
    enum_cols = ((definition.get("format_spec") or {}).get("enum_cols") or {})
    out: Dict[str, List[str]] = {}
    for raw_col_idx, values in enum_cols.items():
        col_idx = int(raw_col_idx)
        header = headers[col_idx]
        out[header] = [str(item).strip() for item in (values or []) if str(item).strip()]
    return out


def _inline_validation_values(formula1: Any) -> List[str]:
    raw = str(formula1 or "").strip()
    if raw.startswith('"') and raw.endswith('"'):
        raw = raw[1:-1]
    return [item.strip() for item in raw.split(",") if item.strip()]


def _validation_values_for_column(ws: Any, one_based_col_idx: int) -> List[str]:
    value_lists = _validation_value_lists_for_column(ws, one_based_col_idx)
    return value_lists[0] if value_lists else []


def _validation_value_lists_for_column(ws: Any, one_based_col_idx: int) -> List[List[str]]:
    out: List[List[str]] = []
    for data_validation in ws.data_validations.dataValidation:
        if data_validation.type != "list":
            continue
        for cell_range in data_validation.sqref.ranges:
            if (
                cell_range.min_col <= one_based_col_idx <= cell_range.max_col
                and cell_range.min_row <= 2 <= cell_range.max_row
            ):
                out.append(_inline_validation_values(data_validation.formula1))
                break
    return out


def _actual_enum_values_by_header(ws: Any, headers: Sequence[str]) -> Dict[str, List[str]]:
    out: Dict[str, List[str]] = {}
    for one_based_col_idx, header in enumerate(headers, start=1):
        values = _validation_values_for_column(ws, one_based_col_idx)
        if values:
            out[str(header)] = values
    return out


def _assert_enum_validations_match_exactly(template_path: Path, definition: Mapping[str, Any]) -> None:
    expected_headers = [str(item) for item in definition.get("headers") or []]
    expected_enums = _definition_enum_values_by_header(definition)
    workbook = load_workbook(template_path, data_only=True)
    try:
        ws = workbook.active
        for one_based_col_idx, header in enumerate(expected_headers, start=1):
            expected_values = expected_enums.get(header)
            if not expected_values:
                continue
            actual_value_lists = _validation_value_lists_for_column(ws, one_based_col_idx)
            assert actual_value_lists == [expected_values], (
                f"{template_path.name} 的 {header} 下拉规则必须只有一条且只包含当前选项；"
                f"期望 {[expected_values]}，实际 {actual_value_lists}"
            )
    finally:
        workbook.close()


def _assert_enum_validation_refresh_keeps_other_columns() -> None:
    _ensure_repo_on_path()
    excel_templates = importlib.import_module("core.services.common.excel_templates")

    workbook = Workbook()
    try:
        ws = workbook.active
        shared_validation = DataValidation(type="list", formula1='"自制,外协"', allow_blank=True)
        shared_validation.add("C2:C500")
        shared_validation.add("F2:F500")
        ws.add_data_validation(shared_validation)
        other_validation = DataValidation(type="list", formula1='"是,否"', allow_blank=True)
        other_validation.add("G2:G500")
        ws.add_data_validation(other_validation)

        excel_templates._apply_sheet_layout(
            ws,
            format_spec={"enum_cols": {2: ["自制", "外协"]}},
            data_row_count=1,
        )

        assert _validation_value_lists_for_column(ws, 3) == [["自制", "外协"]]
        assert _validation_value_lists_for_column(ws, 6) == [["自制", "外协"]]
        assert _validation_value_lists_for_column(ws, 7) == [["是", "否"]]
    finally:
        workbook.close()


def _assert_enum_validation_refresh_splits_contiguous_ranges() -> None:
    _ensure_repo_on_path()
    excel_templates = importlib.import_module("core.services.common.excel_templates")

    workbook = Workbook()
    try:
        ws = workbook.active
        shared_validation = DataValidation(type="list", formula1='"旧一,旧二"', allow_blank=True)
        shared_validation.add("C2:F500")
        ws.add_data_validation(shared_validation)

        excel_templates._apply_sheet_layout(
            ws,
            format_spec={"enum_cols": {2: ["自制", "外协"]}},
            data_row_count=1,
        )

        assert _validation_value_lists_for_column(ws, 3) == [["自制", "外协"]]
        assert _validation_value_lists_for_column(ws, 4) == [["旧一", "旧二"]]
        assert _validation_value_lists_for_column(ws, 6) == [["旧一", "旧二"]]
    finally:
        workbook.close()


def _normalize_help_column_label(raw: str) -> str:
    text = re.sub(r"[`*]", "", str(raw or "")).strip().rstrip("。.")

    def _replace_note(match: Match[str]) -> str:
        inner = str(match.group(1) or match.group(2) or "").strip()
        return "(h)" if inner.lower() == "h" else ""

    text = re.sub(r"（([^）]*)）|\(([^)]*)\)", _replace_note, text)
    text = re.split(r"[；;]", text, maxsplit=1)[0]
    return text.strip()


def _split_help_columns(column_text: str) -> List[str]:
    columns: List[str] = []
    current: List[str] = []
    depth = 0
    for char in column_text:
        if char in "（(":
            depth += 1
        elif char in "）)" and depth > 0:
            depth -= 1
        if char in "、,，" and depth == 0:
            columns.append("".join(current).strip())
            current = []
            continue
        current.append(char)
    if current:
        columns.append("".join(current).strip())
    return columns


def _extract_help_card_columns(payload: Mapping[str, Any]) -> List[str]:
    help_card = payload.get("help_card") or {}
    for item in help_card.get("items") or []:
        text = str(item or "").strip()
        if text.startswith("列："):
            column_text = text[len("列：") :].strip()
        elif text.startswith("模板列："):
            column_text = text[len("模板列：") :].strip()
        else:
            continue
        column_text = re.split(r"[。]", column_text, maxsplit=1)[0].strip()
        raw_columns = _split_help_columns(column_text)
        return [_normalize_help_column_label(column) for column in raw_columns if _normalize_help_column_label(column)]
    raise AssertionError(f"{payload.get('title')} 页面帮助卡缺少“列：...”说明")


def _read_sample_rows(ws: Any, *, width: int, count: int) -> List[List[Any]]:
    rows: List[List[Any]] = []
    for row_idx in range(2, 2 + count):
        rows.append([ws.cell(row_idx, col_idx).value for col_idx in range(1, width + 1)])
    return rows


def _read_header_row(ws: Any) -> List[Any]:
    return [ws.cell(1, col_idx).value for col_idx in range(1, ws.max_column + 1)]


def _read_template_snapshot(
    template_path: Path,
    definition: Mapping[str, Any],
) -> Tuple[List[Any], List[List[Any]], Dict[str, List[str]]]:
    expected_headers = [str(item) for item in definition.get("headers") or []]
    expected_sample_rows = list(definition.get("sample_rows") or [])
    workbook = load_workbook(template_path, data_only=True)
    try:
        ws = workbook.active
        actual_headers = _read_header_row(ws)
        actual_sample_rows = _read_sample_rows(
            ws,
            width=len(expected_headers),
            count=len(expected_sample_rows),
        )
        actual_enums = _actual_enum_values_by_header(ws, expected_headers)
    finally:
        workbook.close()
    return actual_headers, actual_sample_rows, actual_enums


def _write_legacy_op_type_template(
    template_path: Path,
    *,
    sample_rows: Sequence[Sequence[Any]],
    dropdown_values: Sequence[str],
) -> None:
    definition = _template_definition("工种配置.xlsx")
    workbook = Workbook()
    try:
        ws = workbook.active
        ws.append(list(definition.get("headers") or []))
        for row in sample_rows:
            ws.append(list(row))
        if dropdown_values:
            validation = DataValidation(type="list", formula1=f"\"{','.join(dropdown_values)}\"", allow_blank=True)
            validation.add("C2:C500")
            ws.add_data_validation(validation)
        template_path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(template_path)
    finally:
        workbook.close()


def _assert_op_type_template_current(template_path: Path) -> None:
    definition = _template_definition("工种配置.xlsx")
    expected_headers = [str(item) for item in definition.get("headers") or []]
    expected_sample_rows = list(definition.get("sample_rows") or [])
    expected_enums = _definition_enum_values_by_header(definition)

    actual_headers, actual_sample_rows, actual_enums = _read_template_snapshot(template_path, definition)
    assert actual_headers == expected_headers, "工种配置旧模板刷新后表头不正确"
    assert actual_sample_rows == expected_sample_rows, "工种配置旧模板刷新后示例行不正确"
    assert actual_enums == expected_enums, "工种配置旧模板刷新后下拉值不正确"


def _assert_legacy_op_type_template_refreshes() -> None:
    _ensure_repo_on_path()
    excel_templates = importlib.import_module("core.services.common.excel_templates")

    scenarios = (
        (
            "旧示例行仍写内部/外部",
            [["OT001", "数车", "内部"], ["OT002", "标印", "外部"]],
            ["自制", "外协"],
        ),
        (
            "旧下拉仍写内部/外部",
            [["OT001", "数车", "自制"], ["OT002", "标印", "外协"]],
            ["内部", "外部"],
        ),
    )

    for label, sample_rows, dropdown_values in scenarios:
        with tempfile.TemporaryDirectory(prefix="aps_excel_template_contract_") as tmpdir:
            template_path = Path(tmpdir) / "工种配置.xlsx"
            _write_legacy_op_type_template(
                template_path,
                sample_rows=sample_rows,
                dropdown_values=dropdown_values,
            )

            result = excel_templates.ensure_excel_templates(tmpdir)
            assert "工种配置.xlsx" in result.get("created", []), f"{label} 时，旧模板没有触发自动刷新"
            definition = _template_definition("工种配置.xlsx")
            expected_headers = [str(item) for item in definition.get("headers") or []]
            expected_enums = _definition_enum_values_by_header(definition)
            actual_headers, actual_sample_rows, actual_enums = _read_template_snapshot(template_path, definition)
            assert actual_headers == expected_headers, f"{label} 时，表头应保持当前定义"
            assert actual_sample_rows == sample_rows, f"{label} 时，已有示例行不能被自动覆盖"
            assert actual_enums == expected_enums, f"{label} 时，归属下拉应刷新为自制/外协"


def _assert_legacy_op_type_template_with_extra_rows_repairs_dropdown_without_overwriting_data() -> None:
    _ensure_repo_on_path()
    excel_templates = importlib.import_module("core.services.common.excel_templates")

    original_rows = [
        ["OT001", "数车", "内部"],
        ["OT002", "标印", "外部"],
        ["OT003", "喷涂", "内部"],
    ]

    with tempfile.TemporaryDirectory(prefix="aps_excel_template_contract_") as tmpdir:
        template_path = Path(tmpdir) / "工种配置.xlsx"
        _write_legacy_op_type_template(
            template_path,
            sample_rows=original_rows,
            dropdown_values=["内部", "外部"],
        )

        result = excel_templates.ensure_excel_templates(tmpdir)
        assert "工种配置.xlsx" in result.get("created", []), "旧工种模板有额外数据行时，仍应触发下拉刷新"

        definition = _template_definition("工种配置.xlsx")
        expected_headers = [str(item) for item in definition.get("headers") or []]
        expected_enums = _definition_enum_values_by_header(definition)
        workbook = load_workbook(template_path, data_only=True)
        try:
            ws = workbook.active
            actual_headers = _read_header_row(ws)
            actual_rows = _read_sample_rows(ws, width=len(expected_headers), count=len(original_rows))
            actual_enums = _actual_enum_values_by_header(ws, expected_headers)
        finally:
            workbook.close()

        assert actual_headers == expected_headers, "旧工种模板有额外数据行时，表头应刷新为当前表头"
        assert actual_rows == original_rows, "旧工种模板有额外数据行时，用户数据行不能被覆盖"
        assert actual_enums == expected_enums, "旧工种模板有额外数据行时，归属下拉应刷新为自制/外协"


def _assert_legacy_op_type_template_custom_rows_repairs_dropdown_without_overwriting_data() -> None:
    _ensure_repo_on_path()
    excel_templates = importlib.import_module("core.services.common.excel_templates")

    original_rows = [
        ["OT777", "现场自定义工种", "内部"],
    ]

    with tempfile.TemporaryDirectory(prefix="aps_excel_template_contract_") as tmpdir:
        template_path = Path(tmpdir) / "工种配置.xlsx"
        _write_legacy_op_type_template(
            template_path,
            sample_rows=original_rows,
            dropdown_values=["内部", "外部"],
        )

        result = excel_templates.ensure_excel_templates(tmpdir)
        assert "工种配置.xlsx" in result.get("created", []), "旧工种模板有自定义数据行时，应只修下拉"

        definition = _template_definition("工种配置.xlsx")
        expected_headers = [str(item) for item in definition.get("headers") or []]
        expected_enums = _definition_enum_values_by_header(definition)
        workbook = load_workbook(template_path, data_only=True)
        try:
            ws = workbook.active
            actual_headers = _read_header_row(ws)
            actual_rows = _read_sample_rows(ws, width=len(expected_headers), count=len(original_rows))
            actual_enums = _actual_enum_values_by_header(ws, expected_headers)
        finally:
            workbook.close()

        assert actual_headers == expected_headers, "旧工种模板有自定义数据行时，表头应保持当前定义"
        assert actual_rows == original_rows, "旧工种模板有自定义数据行时，用户数据行不能被覆盖"
        assert actual_enums == expected_enums, "旧工种模板有自定义数据行时，归属下拉应刷新为自制/外协"


def _assert_broken_existing_template_is_not_silently_overwritten() -> None:
    _ensure_repo_on_path()
    excel_templates = importlib.import_module("core.services.common.excel_templates")

    with tempfile.TemporaryDirectory(prefix="aps_excel_template_contract_") as tmpdir:
        template_path = Path(tmpdir) / "人员基本信息.xlsx"
        original_bytes = b"not a real xlsx"
        template_path.write_bytes(original_bytes)

        try:
            excel_templates.ensure_excel_templates(tmpdir)
        except excel_templates.ExcelTemplateError as exc:
            assert "人员基本信息.xlsx" in str(exc), "坏模板报错应带文件名，方便现场定位"
        else:
            raise AssertionError("已有模板读取失败时不应被静默覆盖重建")
        assert template_path.read_bytes() == original_bytes, "坏模板读取失败时不能覆盖用户文件"


def _assert_mismatched_existing_template_is_not_silently_overwritten() -> None:
    _ensure_repo_on_path()
    excel_templates = importlib.import_module("core.services.common.excel_templates")

    with tempfile.TemporaryDirectory(prefix="aps_excel_template_contract_") as tmpdir:
        template_path = Path(tmpdir) / "人员基本信息.xlsx"
        workbook = Workbook()
        try:
            ws = workbook.active
            ws.append(["自定义列", "姓名", "备注"])
            ws.append(["现场保留值", "张三", "不要覆盖"])
            workbook.save(template_path)
        finally:
            workbook.close()
        original_bytes = template_path.read_bytes()

        try:
            excel_templates.ensure_excel_templates(tmpdir)
        except excel_templates.ExcelTemplateError as exc:
            assert "人员基本信息.xlsx" in str(exc), "表头不匹配报错应带文件名，方便现场定位"
        else:
            raise AssertionError("可打开但表头不匹配的已有模板不应被静默覆盖重建")
        assert template_path.read_bytes() == original_bytes, "表头不匹配时不能覆盖现场已有模板"


def _assert_mismatched_op_type_template_extra_header_is_not_repaired() -> None:
    _ensure_repo_on_path()
    excel_templates = importlib.import_module("core.services.common.excel_templates")

    with tempfile.TemporaryDirectory(prefix="aps_excel_template_contract_") as tmpdir:
        template_path = Path(tmpdir) / "工种配置.xlsx"
        definition = _template_definition("工种配置.xlsx")
        workbook = Workbook()
        try:
            ws = workbook.active
            ws.append(list(definition.get("headers") or []) + ["现场自定义列"])
            ws.append(["OT777", "现场工种", "内部", "不要覆盖"])
            validation = DataValidation(type="list", formula1="\"内部,外部\"", allow_blank=True)
            validation.add("C2:C500")
            ws.add_data_validation(validation)
            workbook.save(template_path)
        finally:
            workbook.close()
        original_bytes = template_path.read_bytes()

        try:
            excel_templates.ensure_excel_templates(tmpdir)
        except excel_templates.ExcelTemplateError as exc:
            assert "工种配置.xlsx" in str(exc), "表头多列报错应带文件名，方便现场定位"
        else:
            raise AssertionError("工种模板多出现场自定义列时，不应被自动修复保存")
        assert template_path.read_bytes() == original_bytes, "表头多列时不能改动现场已有模板"


def _assert_process_excel_template_files_match_registered_definitions() -> None:
    for filename in PROCESS_TEMPLATE_FILES:
        template_path = REPO_ROOT / "templates_excel" / filename
        assert template_path.exists(), f"缺少工艺 Excel 模板文件：{template_path}"

        definition = _template_definition(filename)
        expected_headers = [str(item) for item in definition.get("headers") or []]
        expected_sample_rows = list(definition.get("sample_rows") or [])
        expected_enums = _definition_enum_values_by_header(definition)
        actual_headers, actual_sample_rows, actual_enums = _read_template_snapshot(template_path, definition)

        assert actual_headers == expected_headers, f"{filename} 表头和模板定义不一致"
        assert actual_sample_rows == expected_sample_rows, f"{filename} 示例行和模板定义不一致"
        assert actual_enums == expected_enums, f"{filename} 下拉值和模板定义不一致"
        _assert_enum_validations_match_exactly(template_path, definition)


def _assert_page_manual_excel_template_files_match_registered_definitions() -> None:
    for filename in sorted(set(PAGE_MANUAL_TEMPLATE_COLUMN_CONTRACTS.values())):
        template_path = REPO_ROOT / "templates_excel" / filename
        assert template_path.exists(), f"页面说明提到的 Excel 模板文件不存在：{template_path}"

        definition = _template_definition(filename)
        expected_headers = [str(item) for item in definition.get("headers") or []]
        expected_enums = _definition_enum_values_by_header(definition)
        actual_headers, _actual_sample_rows, actual_enums = _read_template_snapshot(template_path, definition)

        assert actual_headers == expected_headers, f"{filename} 真实文件表头和页面说明引用的模板定义不一致"
        assert actual_enums == expected_enums, f"{filename} 真实文件下拉值和页面说明引用的模板定义不一致"


def _assert_supplier_conversion_output_matches_current_template_contract() -> None:
    filename = "供应商配置.xlsx"
    template_path = REPO_ROOT / "templates_excel" / "转换输出" / filename
    assert template_path.exists(), f"缺少供应商配置转换输出示例：{template_path}"

    definition = _template_definition(filename)
    expected_headers = [str(item) for item in definition.get("headers") or []]
    expected_enums = _definition_enum_values_by_header(definition)
    actual_headers, _actual_sample_rows, actual_enums = _read_template_snapshot(template_path, definition)

    assert actual_headers == expected_headers, "供应商配置转换输出仍是旧列，必须补齐状态和备注"
    assert actual_enums == expected_enums, "供应商配置转换输出的状态下拉必须和正式模板一致"
    _assert_enum_validations_match_exactly(template_path, definition)

    workbook = load_workbook(template_path, data_only=True)
    try:
        ws = workbook.active
        row_count = 0
        for row_idx in range(2, ws.max_row + 1):
            values = [ws.cell(row_idx, col_idx).value for col_idx in range(1, len(expected_headers) + 1)]
            if not any(value not in (None, "") for value in values):
                continue
            row_count += 1
            assert values[4] in ("启用", "停用"), f"供应商配置转换输出第 {row_idx} 行状态必须是启用或停用"
        assert row_count > 0, "供应商配置转换输出至少要保留一行示例或转换结果"
    finally:
        workbook.close()


def _assert_process_excel_dropdown_values_match_page_and_full_manuals() -> None:
    _ensure_repo_on_path()
    page_manuals = importlib.import_module("web.viewmodels.page_manuals")
    static_manual_text = _read_static_manual()

    for filename, contract in DROPDOWN_MANUAL_CONTRACTS.items():
        definition = _template_definition(filename)
        expected_enums = _definition_enum_values_by_header(definition)
        template_path = REPO_ROOT / "templates_excel" / filename
        workbook = load_workbook(template_path, data_only=True)
        try:
            ws = workbook.active
            headers = [str(item) for item in definition.get("headers") or []]
            actual_enums = _actual_enum_values_by_header(ws, headers)
        finally:
            workbook.close()

        for column_name, expected_values in (contract["dropdowns"] or {}).items():
            expected_values = list(expected_values)
            assert expected_enums.get(column_name) == expected_values, f"{filename} 模板定义中的 {column_name} 下拉值不正确"
            assert actual_enums.get(column_name) == expected_values, (
                f"{filename} 文件里的 {column_name} 下拉值不正确；"
                f"期望 {expected_values}，实际 {actual_enums.get(column_name)}"
            )

        payload = page_manuals.build_manual_payload(str(contract["manual_id"]), include_sections=True)
        assert payload is not None, f"{contract['manual_id']} 无法构建页面说明"
        page_text = _manual_payload_text(payload)
        full_manual_section = _extract_markdown_section(static_manual_text, str(contract["manual_heading"]))

        for phrase in contract["page_phrases"]:
            assert phrase in page_text, f"{contract['manual_id']} 页面说明缺少推荐值：{phrase}"
        for phrase in contract["manual_phrases"]:
            assert phrase in full_manual_section, f"{contract['manual_heading']} 说明书章节缺少推荐值：{phrase}"
        _assert_dropdown_copy_distinguishes_recommended_and_compatible(
            page_text,
            label=f"{contract['manual_id']} 页面说明",
            dropdowns=contract["dropdowns"],
            compatible_values=contract.get("compatible_values", {}),
        )
        _assert_dropdown_copy_distinguishes_recommended_and_compatible(
            full_manual_section,
            label=f"{contract['manual_heading']} 说明书章节",
            dropdowns=contract["dropdowns"],
            compatible_values=contract.get("compatible_values", {}),
        )
        for phrase in contract["forbidden_page_phrases"]:
            assert phrase not in page_text, f"{contract['manual_id']} 页面说明仍包含旧说法：{phrase}"
        for phrase in contract["forbidden_manual_phrases"]:
            if _is_allowed_legacy_op_type_note(full_manual_section, phrase):
                continue
            assert phrase not in full_manual_section, f"{contract['manual_heading']} 说明书章节仍包含旧说法：{phrase}"


def _assert_page_manual_help_columns_match_template_headers() -> None:
    _ensure_repo_on_path()
    page_manuals = importlib.import_module("web.viewmodels.page_manuals")

    for manual_id, filename in PAGE_MANUAL_TEMPLATE_COLUMN_CONTRACTS.items():
        definition = _template_definition(filename)
        expected_headers = [str(item) for item in definition.get("headers") or []]
        payload = page_manuals.build_manual_payload(manual_id, include_sections=True)
        assert payload is not None, f"{manual_id} 无法构建页面说明"
        actual_columns = _extract_help_card_columns(payload)
        assert actual_columns == expected_headers, (
            f"{manual_id} 帮助卡列顺序必须和 {filename} 模板表头一致；"
            f"期望 {expected_headers}，实际 {actual_columns}"
        )


def _assert_relation_reverse_header_copy_matches_templates() -> None:
    _ensure_repo_on_path()
    page_manuals = importlib.import_module("web.viewmodels.page_manuals")
    static_manual_text = _read_static_manual()

    assert "人员设备关联，也可以用设备人员关联，填的是同一类关系" in static_manual_text, (
        "总说明书必须写清人员设备关联和设备人员关联是同一类关系"
    )
    relation_section = _extract_markdown_section(static_manual_text, "#### 1.6.7 人员设备关联")
    assert "| 人员管理侧 | 工号、设备编号、技能等级、主操设备 |" in relation_section, (
        "总说明书必须写清人员管理侧人员设备关联模板列顺序"
    )
    assert "| 设备管理侧 | 设备编号、工号、技能等级、主操设备 |" in relation_section, (
        "总说明书必须写清设备管理侧设备人员关联模板列顺序"
    )

    for manual_id, contract in RELATION_REVERSE_HEADER_COPY_CONTRACTS.items():
        filename = str(contract["filename"])
        definition = _template_definition(filename)
        expected_headers = [str(item) for item in definition.get("headers") or []]
        payload = page_manuals.build_manual_payload(manual_id, include_sections=True)
        assert payload is not None, f"{manual_id} 无法构建设备/人员关联反向表头校验 payload"
        actual_columns = _extract_help_card_columns(payload)
        assert actual_columns == expected_headers, (
            f"{manual_id} 帮助卡列顺序必须和 {filename} 模板表头一致；"
            f"期望 {expected_headers}，实际 {actual_columns}"
        )

        payload_text = _manual_payload_text(payload)
        for phrase in contract["required_phrases"]:
            assert phrase in payload_text, f"{manual_id} 缺少设备/人员关联反向表头说明：{phrase}"


def main() -> None:
    _assert_enum_validation_refresh_keeps_other_columns()
    _assert_enum_validation_refresh_splits_contiguous_ranges()
    _assert_process_excel_template_files_match_registered_definitions()
    _assert_page_manual_excel_template_files_match_registered_definitions()
    _assert_supplier_conversion_output_matches_current_template_contract()
    _assert_process_excel_dropdown_values_match_page_and_full_manuals()
    _assert_page_manual_help_columns_match_template_headers()
    _assert_relation_reverse_header_copy_matches_templates()
    _assert_legacy_op_type_template_refreshes()
    _assert_legacy_op_type_template_with_extra_rows_repairs_dropdown_without_overwriting_data()
    _assert_legacy_op_type_template_custom_rows_repairs_dropdown_without_overwriting_data()
    _assert_broken_existing_template_is_not_silently_overwritten()
    _assert_mismatched_existing_template_is_not_silently_overwritten()
    _assert_mismatched_op_type_template_extra_header_is_not_repaired()
    print("OK")


if __name__ == "__main__":
    main()
