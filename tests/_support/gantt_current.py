"""Current Gantt HTTP fixtures and assertions; inputs are app clients and plan DTOs.

These helpers do not run the application at import time. They use only neutral
test support and production APIs, never another test module's fixtures.
"""

import csv
import json
from contextlib import closing
from io import BytesIO, StringIO

from openpyxl import load_workbook

from core.infrastructure.database import get_connection
from tests._support.gantt_retirement import _Boot, _business_state


def prepare_read_state(client):
    """Complete existing fixture configuration before measuring read-only calls."""
    from core.services.scheduler.config.config_service import ConfigService
    from core.services.system.system_config_service import SystemConfigService

    with closing(get_connection(client.application.config["DATABASE_PATH"])) as conn:
        ConfigService(conn).ensure_defaults()
        SystemConfigService(conn).ensure_defaults(client.application.config["BACKUP_KEEP_DAYS"])
        conn.commit()
    return _business_state(client)


def navigation(client, query):
    """Read an exact legacy redirect and its typed boot without changing scope."""
    response = client.get("/scheduler/gantt", query_string=query)
    assert response.status_code == 302, response.get_data(as_text=True)
    target = response.headers["Location"]
    page = client.get(target)
    assert page.status_code == 200, page.get_data(as_text=True)
    parser = _Boot()
    parser.feed(page.get_data(as_text=True))
    result = json.loads("".join(parser.parts))["navigation"]
    assert result["view"] == "gantt"
    parser = _Boot()
    parser.feed(client.get(target).get_data(as_text=True))
    assert json.loads("".join(parser.parts))["navigation"] == result
    return result["context"]


def assert_retired(client, query, message=None):
    """Require honest retirement, with no fallback redirect or stale Gantt boot."""
    response = client.get("/scheduler/gantt", query_string=query)
    html = response.get_data(as_text=True)
    assert response.status_code == 410, html
    assert "Location" not in response.headers
    assert "workbench-boot" not in html
    assert "gantt-container" not in html
    if message is not None:
        assert message in html
    return html


def read_workspace(client, context):
    """Read only the permanent plan ref and exact optional local-time scope."""
    scope = {key: value for key, value in context.items() if key != "plan_ref"}
    response = client.get("/api/workbench/v1/plans/" + context["plan_ref"] + "/workspace", query_string=scope)
    assert response.status_code == 200, response.get_data(as_text=True)
    payload = response.get_json()
    assert payload["ok"] is True
    assert payload["meta"]["source"] == "production"
    assert payload["meta"]["time_basis"] == "factory_local"
    assert payload["data"]["plan"]["plan_ref"] == context["plan_ref"]
    assert payload["data"]["tasks_complete"] is True
    assert payload["data"]["task_count"] == len(payload["data"]["tasks"])
    return payload


def plan_fixture(client, *, scenario=False):
    """Seed the existing three-task fixture; optionally save its real scenario."""
    from tests._support.gantt_scenario import VERSION, _saved_scenario, _seed_base

    with closing(get_connection(client.application.config["DATABASE_PATH"])) as conn:
        _seed_base(conn)
        query = {"version": VERSION}
        if scenario:
            query["scenario_id"] = _saved_scenario(conn).scenario_id
        conn.commit()
    before = prepare_read_state(client)
    context = navigation(client, query)
    return query, context, read_workspace(client, context), before


def invalid_calendar_fixture(client):
    """Keep a damaged explicit calendar row, then read its honest null capacity."""
    _, context, original, _ = plan_fixture(client)
    with closing(get_connection(client.application.config["DATABASE_PATH"])) as conn:
        conn.execute("INSERT INTO WorkCalendar(date,day_type,shift_hours,efficiency,allow_normal,allow_urgent) "
                     "VALUES ('2026-05-04','BROKEN_CALENDAR_CANARY',8,1,'yes','yes')")
        conn.commit()
    before = prepare_read_state(client)
    payload = read_workspace(client, context)
    assert payload["data"]["tasks"] == original["data"]["tasks"]
    calendar = payload["data"]["projections"]["calendar"]
    assert calendar["state"] == "unavailable"
    for row in [calendar["global"]] + calendar["resources"]:
        assert row["state"] == "unavailable"
        assert row["windows"] is None
        assert row["available_hours"] is None
        assert row["issues"]
    for row in payload["data"]["projections"]["occupancy"]["resources"]:
        assert row["available_hours"] is None
        assert row["utilization"] is None
    assert "BROKEN_CALENDAR_CANARY" not in json.dumps(payload, ensure_ascii=False)
    assert _business_state(client) == before
    return context, payload, before


def legacy_payload(**changes):
    """Build retained backend/public Gantt metadata without a legacy UI claim."""
    from core.services.scheduler.gantt_contract import build_gantt_contract

    inputs = dict(contract_version=2, view="machine", version=7, week_start="2026-01-26", week_end="2026-02-01",
                  tasks=[], calendar_days=[], critical_chain={"available": True, "ids": [], "edges": []})
    inputs.update(changes)
    return build_gantt_contract(**inputs)


def assert_plan_exports(client, context, payload):
    """Compare actual CSV and XLSX downloads against independent DTO facts."""
    scope = {key: value for key, value in context.items() if key != "plan_ref"}
    scope["snapshot_ref"] = payload["meta"]["snapshot_ref"]
    rows_by_format = {}
    for fmt in ("csv", "xlsx"):
        response = client.get("/api/workbench/v1/plans/" + context["plan_ref"] + "/export",
                              query_string=dict(scope, format=fmt))
        assert response.status_code == 200, response.get_data(as_text=True) if response.is_json else response.status
        assert response.headers["X-Workbench-Plan-Ref"] == context["plan_ref"]
        assert response.headers["X-Workbench-Snapshot-Ref"] == scope["snapshot_ref"]
        assert int(response.headers["X-Workbench-Row-Count"]) == payload["data"]["task_count"]
        if fmt == "csv":
            raw_rows = list(csv.reader(StringIO(response.data.decode("utf-8-sig"))))
            # The CSV contract guards every text cell, not only formula prefixes.
            assert all(cell.startswith("'") for row in raw_rows[1:] for index, cell in enumerate(row)
                       if index not in (29, 30))
            rows = [raw_rows[0]] + [[cell[1:] if cell.startswith("'") else float(cell)
                                     for cell in row] for row in raw_rows[1:]]
        else:
            book = load_workbook(BytesIO(response.data), read_only=True, data_only=False)
            try:
                sheet = book["计划任务"]
                cells = list(sheet.iter_rows())
                assert all(cell.data_type == "s" for row in cells for index, cell in enumerate(row)
                           if index not in (29, 30))
                rows = [[cell.value for cell in row] for row in cells]
            finally:
                book.close()
        rows_by_format[fmt] = rows
    assert rows_by_format["csv"] == rows_by_format["xlsx"]
    rows = rows_by_format["csv"]
    assert len(rows) == payload["data"]["task_count"] + 1
    assert len(rows[0]) == 32
    actual = {row[11]: row for row in rows[1:]}
    assert set(actual) == {task["task_ref"] for task in payload["data"]["tasks"]}
    for task in payload["data"]["tasks"]:
        row = actual[task["task_ref"]]
        assert row[0] == context["plan_ref"]
        assert row[6] == scope["snapshot_ref"]
        assert row[12] == task["operation_ref"]
        label = task["process_label"]
        assert row[15] == ("\\" + label if label.startswith("\\") else label)
        assert row[22:24] == [task["start"], task["end"]]
    return rows
