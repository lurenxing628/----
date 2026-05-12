from __future__ import annotations

import os
import sys
import tempfile
from datetime import date
from io import BytesIO


def find_repo_root() -> str:
    here = os.path.dirname(os.path.abspath(__file__))
    repo_root = os.path.abspath(os.path.join(here, ".."))
    if os.path.exists(os.path.join(repo_root, "app.py")) and os.path.exists(os.path.join(repo_root, "schema.sql")):
        return repo_root
    raise RuntimeError("未找到项目根目录：要求存在 app.py 与 schema.sql")


def _assert_xlsx(resp, name: str, expect_version: int) -> None:
    if resp.status_code != 200:
        body = resp.data.decode("utf-8", errors="ignore") if getattr(resp, "data", None) else ""
        raise RuntimeError(f"{name} 返回 {resp.status_code}，期望 200，body={body[:500]}")
    ct = resp.headers.get("Content-Type", "")
    if "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" not in ct:
        raise RuntimeError(f"{name} content-type 异常：{ct}")
    cd = resp.headers.get("Content-Disposition", "")
    if f"v{int(expect_version)}" not in cd:
        raise RuntimeError(f"{name} filename 未包含 v{expect_version}（Content-Disposition={cd!r}）")


def _load_xlsx(resp):
    import openpyxl

    return openpyxl.load_workbook(BytesIO(resp.data), data_only=True)


def _assert_number(actual, expect: float, name: str) -> None:
    try:
        actual_number = round(float(actual), 2)
    except Exception as exc:
        raise RuntimeError(f"{name} 不是数字：{actual!r}") from exc
    if actual_number != round(float(expect), 2):
        raise RuntimeError(f"{name} 异常：{actual!r}，期望 {expect!r}")


def _assert_single_data_row(ws, name: str) -> None:
    if ws.max_row != 2:
        raise RuntimeError(f"{name} 数据行数异常：max_row={ws.max_row}，期望 2")


def _assert_overdue_export(resp, name: str, expect_version: int, expect_finish_time: str, expect_delay_hours: float) -> None:
    _assert_xlsx(resp, name, expect_version)
    wb = _load_xlsx(resp)
    try:
        ws = wb["overdue"]
        _assert_single_data_row(ws, name)
        if ws["B2"].value != "B_REPORT":
            raise RuntimeError(f"{name} 超期批次异常：{ws['B2'].value!r}")
        if ws["G2"].value != expect_finish_time:
            raise RuntimeError(f"{name} 完工时间异常：{ws['G2'].value!r}，期望 {expect_finish_time!r}")
        _assert_number(ws["I2"].value, expect_delay_hours, f"{name} 超期小时")
    finally:
        wb.close()


def _assert_utilization_export(resp, name: str, expect_version: int, expect_machine_hours: float) -> None:
    _assert_xlsx(resp, name, expect_version)
    wb = _load_xlsx(resp)
    try:
        ws = wb["machines"]
        _assert_single_data_row(ws, name)
        if ws["A2"].value != "MC_REPORT":
            raise RuntimeError(f"{name} 设备编号异常：{ws['A2'].value!r}")
        _assert_number(ws["C2"].value, expect_machine_hours, f"{name} 设备负荷小时")
    finally:
        wb.close()


def _assert_downtime_export(resp, name: str, expect_version: int, expect_overlap_hours: float) -> None:
    _assert_xlsx(resp, name, expect_version)
    wb = _load_xlsx(resp)
    try:
        ws = wb["停机影响"]
        _assert_single_data_row(ws, name)
        if ws["A2"].value != "MC_REPORT":
            raise RuntimeError(f"{name} 停机设备编号异常：{ws['A2'].value!r}")
        _assert_number(ws["C2"].value, 2.0, f"{name} 停机小时")
        _assert_number(ws["E2"].value, expect_overlap_hours, f"{name} 排程重叠小时")
    finally:
        wb.close()


def _assert_invalid_version(resp, name: str) -> None:
    from core.services.scheduler.version_resolution import VERSION_ERROR_MESSAGE

    if resp.status_code != 400:
        raise RuntimeError(f"{name} 返回 {resp.status_code}，期望 400")
    body = resp.get_data(as_text=True)
    if VERSION_ERROR_MESSAGE not in body:
        raise RuntimeError(f"{name} 未返回统一版本错误文案：{body[:200]!r}")


def _assert_missing_version(resp, name: str) -> None:
    if resp.status_code != 404:
        raise RuntimeError(f"{name} 返回 {resp.status_code}，期望 404")
    body = resp.get_data(as_text=True)
    if "排产版本不存在，请先选择已有版本。" not in body:
        raise RuntimeError(f"{name} 未返回版本不存在文案：{body[:200]!r}")


def _assert_date_validation(resp, name: str, expected_text: str) -> None:
    if resp.status_code != 400:
        raise RuntimeError(f"{name} 返回 {resp.status_code}，期望 400")
    body = resp.get_data(as_text=True)
    if expected_text not in body:
        raise RuntimeError(f"{name} 未返回预期日期错误文案：{body[:200]!r}")


def _assert_no_data_export(resp, name: str) -> None:
    if resp.status_code != 400:
        raise RuntimeError(f"{name} 返回 {resp.status_code}，期望 400")
    body = resp.get_data(as_text=True)
    if "暂无数据，不能导出" not in body:
        raise RuntimeError(f"{name} 未返回无数据导出文案：{body[:200]!r}")


def _seed_distinguishable_report_data(conn) -> None:
    default_day = date.today().isoformat()
    conn.execute("INSERT INTO Parts(part_no, part_name) VALUES (?, ?)", ("P_REPORT", "报表测试零件"))
    conn.execute(
        "INSERT INTO Machines(machine_id, name, status) VALUES (?, ?, ?)",
        ("MC_REPORT", "报表测试设备", "active"),
    )
    conn.execute(
        "INSERT INTO Operators(operator_id, name, status) VALUES (?, ?, ?)",
        ("OP_REPORT", "报表测试人员", "active"),
    )
    conn.execute(
        """
        INSERT INTO Batches(batch_id, part_no, part_name, quantity, due_date, priority, ready_status, status, remark)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        ("B_REPORT", "P_REPORT", "报表测试零件", 1, "2026-01-01", "normal", "yes", "pending", "latest export"),
    )
    conn.execute(
        """
        INSERT INTO BatchOperations(op_code, batch_id, seq, op_type_name, source, machine_id, operator_id, setup_hours, unit_hours, status)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        ("B_REPORT_10", "B_REPORT", 10, "数铣", "internal", "MC_REPORT", "OP_REPORT", 0, 1, "scheduled"),
    )
    row = conn.execute("SELECT id FROM BatchOperations WHERE op_code=?", ("B_REPORT_10",)).fetchone()
    if not row:
        raise RuntimeError("未插入报表测试工序")
    op_id = int(row["id"])

    conn.executemany(
        """
        INSERT INTO Schedule(op_id, machine_id, operator_id, start_time, end_time, lock_status, version)
        VALUES (?, ?, ?, ?, ?, ?, ?)
        """,
        [
            (op_id, "MC_REPORT", "OP_REPORT", "2026-01-02 08:00:00", "2026-01-02 10:00:00", "unlocked", 6),
            (op_id, "MC_REPORT", "OP_REPORT", "2026-01-03 08:00:00", "2026-01-03 12:00:00", "unlocked", 7),
        ],
    )
    conn.executemany(
        """
        INSERT INTO ScheduleHistory (version, strategy, batch_count, op_count, result_status, result_summary, created_by)
        VALUES (?, ?, ?, ?, ?, ?, ?)
        """,
        [
            (5, "test", 0, 0, "ok", None, "regression"),
            (6, "test", 1, 1, "ok", None, "regression"),
            (7, "test", 1, 1, "ok", None, "regression"),
        ],
    )
    conn.execute(
        """
        INSERT INTO MachineDowntimes(machine_id, start_time, end_time, reason_code, reason_detail, status)
        VALUES (?, ?, ?, ?, ?, ?)
        """,
        ("MC_REPORT", "2026-01-03 09:00:00", "2026-01-03 11:00:00", "maintenance", "latest export", "active"),
    )
    conn.execute(
        """
        INSERT INTO MachineDowntimes(machine_id, start_time, end_time, reason_code, reason_detail, status)
        VALUES (?, ?, ?, ?, ?, ?)
        """,
        (
            "MC_REPORT",
            f"{default_day} 09:00:00",
            f"{default_day} 11:00:00",
            "maintenance",
            "default seven days must not leak into version without schedule",
            "active",
        ),
    )


def main() -> None:
    repo_root = find_repo_root()
    if repo_root not in sys.path:
        sys.path.insert(0, repo_root)

    # 隔离目录，避免污染真实 db/logs/backups/templates_excel
    root = tempfile.mkdtemp(prefix="aps_regression_reports_export_ver_")
    test_db = os.path.join(root, "aps_test.db")
    test_logs = os.path.join(root, "logs")
    test_backups = os.path.join(root, "backups")
    test_templates = os.path.join(root, "templates_excel")
    os.makedirs(test_logs, exist_ok=True)
    os.makedirs(test_backups, exist_ok=True)
    os.makedirs(test_templates, exist_ok=True)

    os.environ["APS_ENV"] = "development"
    os.environ["APS_DB_PATH"] = test_db
    os.environ["APS_LOG_DIR"] = test_logs
    os.environ["APS_BACKUP_DIR"] = test_backups
    os.environ["APS_EXCEL_TEMPLATE_DIR"] = test_templates

    from core.infrastructure.database import ensure_schema, get_connection

    ensure_schema(test_db, logger=None, schema_path=os.path.join(repo_root, "schema.sql"))

    # 准备 v6/v7 两套可区分报表数据，让 latest_version() 稳定落到 v7
    conn = get_connection(test_db)
    try:
        _seed_distinguishable_report_data(conn)
        conn.commit()
    finally:
        conn.close()

    # Flask test_client（不启动 server）
    import importlib

    app_mod = importlib.import_module("app")
    app = app_mod.create_app()
    client = app.test_client()

    sd = "2026-01-01"
    ed = "2026-01-07"

    # 1) overdue/export
    _assert_overdue_export(
        client.get("/reports/overdue/export?version=6"),
        "GET /reports/overdue/export（version=6）",
        6,
        "2026-01-02 10:00:00",
        10.0,
    )
    _assert_overdue_export(
        client.get("/reports/overdue/export"),
        "GET /reports/overdue/export（missing version）",
        7,
        "2026-01-03 12:00:00",
        36.0,
    )
    _assert_overdue_export(
        client.get("/reports/overdue/export?version="),
        "GET /reports/overdue/export（empty version）",
        7,
        "2026-01-03 12:00:00",
        36.0,
    )
    _assert_overdue_export(
        client.get("/reports/overdue/export?version=latest"),
        "GET /reports/overdue/export（version=latest）",
        7,
        "2026-01-03 12:00:00",
        36.0,
    )
    _assert_invalid_version(client.get("/reports/overdue/export?version=abc"), "GET /reports/overdue/export（invalid version）")
    _assert_invalid_version(client.get("/reports/overdue/export?version=0"), "GET /reports/overdue/export（version=0）")

    # 2) utilization/export（显式日期按参数，无日期按版本排程范围）
    _assert_utilization_export(
        client.get(f"/reports/utilization/export?version=6&start_date={sd}&end_date={ed}"),
        "GET /reports/utilization/export（version=6）",
        6,
        2.0,
    )
    _assert_utilization_export(
        client.get(f"/reports/utilization/export?start_date={sd}&end_date={ed}"),
        "GET /reports/utilization/export（missing version）",
        7,
        4.0,
    )
    _assert_utilization_export(
        client.get(f"/reports/utilization/export?version=&start_date={sd}&end_date={ed}"),
        "GET /reports/utilization/export（empty version）",
        7,
        4.0,
    )
    _assert_utilization_export(
        client.get(f"/reports/utilization/export?version=latest&start_date={sd}&end_date={ed}"),
        "GET /reports/utilization/export（version=latest）",
        7,
        4.0,
    )
    _assert_utilization_export(
        client.get("/reports/utilization/export?version=latest"),
        "GET /reports/utilization/export（version=latest，无日期）",
        7,
        4.0,
    )
    _assert_utilization_export(
        client.get("/reports/utilization/export"),
        "GET /reports/utilization/export（missing version，无日期）",
        7,
        4.0,
    )
    _assert_utilization_export(
        client.get("/reports/utilization/export?version="),
        "GET /reports/utilization/export（empty version，无日期）",
        7,
        4.0,
    )
    _assert_utilization_export(
        client.get("/reports/utilization/export?version=6"),
        "GET /reports/utilization/export（version=6，无日期）",
        6,
        2.0,
    )
    _assert_date_validation(
        client.get("/reports/utilization/export?version=latest&start_date=bad-date"),
        "GET /reports/utilization/export（只传坏开始日期）",
        "日期格式不正确",
    )
    _assert_date_validation(
        client.get(f"/reports/utilization/export?version=latest&start_date={sd}"),
        "GET /reports/utilization/export（只传开始日期）",
        "缺少开始日期或结束日期",
    )
    _assert_invalid_version(
        client.get(f"/reports/utilization/export?version=abc&start_date={sd}&end_date={ed}"),
        "GET /reports/utilization/export（invalid version）",
    )
    _assert_invalid_version(
        client.get(f"/reports/utilization/export?version=0&start_date={sd}&end_date={ed}"),
        "GET /reports/utilization/export（version=0）",
    )
    _assert_missing_version(
        client.get(f"/reports/utilization/export?version=999&start_date={sd}&end_date={ed}"),
        "GET /reports/utilization/export（version=999）",
    )

    # 3) downtime/export（显式日期按参数，无日期按版本排程范围）
    _assert_downtime_export(
        client.get(f"/reports/downtime/export?version=6&start_date={sd}&end_date={ed}"),
        "GET /reports/downtime/export（version=6）",
        6,
        0.0,
    )
    _assert_downtime_export(
        client.get(f"/reports/downtime/export?start_date={sd}&end_date={ed}"),
        "GET /reports/downtime/export（missing version）",
        7,
        2.0,
    )
    _assert_downtime_export(
        client.get(f"/reports/downtime/export?version=&start_date={sd}&end_date={ed}"),
        "GET /reports/downtime/export（empty version）",
        7,
        2.0,
    )
    _assert_downtime_export(
        client.get(f"/reports/downtime/export?version=latest&start_date={sd}&end_date={ed}"),
        "GET /reports/downtime/export（version=latest）",
        7,
        2.0,
    )
    _assert_downtime_export(
        client.get("/reports/downtime/export?version=latest"),
        "GET /reports/downtime/export（version=latest，无日期）",
        7,
        2.0,
    )
    _assert_downtime_export(
        client.get("/reports/downtime/export"),
        "GET /reports/downtime/export（missing version，无日期）",
        7,
        2.0,
    )
    _assert_downtime_export(
        client.get("/reports/downtime/export?version="),
        "GET /reports/downtime/export（empty version，无日期）",
        7,
        2.0,
    )
    _assert_no_data_export(
        client.get("/reports/downtime/export?version=6"),
        "GET /reports/downtime/export（version=6，无日期）",
    )
    _assert_no_data_export(
        client.get("/reports/downtime/export?version=5"),
        "GET /reports/downtime/export（version=5，无排程，无日期）",
    )
    _assert_date_validation(
        client.get("/reports/downtime/export?version=latest&end_date=bad-date"),
        "GET /reports/downtime/export（只传坏结束日期）",
        "缺少开始日期或结束日期",
    )
    _assert_date_validation(
        client.get(f"/reports/downtime/export?version=latest&end_date={ed}"),
        "GET /reports/downtime/export（只传结束日期）",
        "缺少开始日期或结束日期",
    )
    _assert_invalid_version(
        client.get(f"/reports/downtime/export?version=abc&start_date={sd}&end_date={ed}"),
        "GET /reports/downtime/export（invalid version）",
    )
    _assert_invalid_version(
        client.get(f"/reports/downtime/export?version=0&start_date={sd}&end_date={ed}"),
        "GET /reports/downtime/export（version=0）",
    )
    _assert_missing_version(
        client.get(f"/reports/downtime/export?version=999&start_date={sd}&end_date={ed}"),
        "GET /reports/downtime/export（version=999）",
    )

    print("OK")


if __name__ == "__main__":
    main()
