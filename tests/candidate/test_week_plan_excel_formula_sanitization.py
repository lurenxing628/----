from __future__ import annotations

from openpyxl import load_workbook

from core.services.scheduler.week_plan_excel import build_week_plan_export_workbook


def test_week_plan_summary_sheet_sanitizes_formula_like_values() -> None:
    workbook_bytes = build_week_plan_export_workbook(
        [],
        plan_resolution={"user_label": "=HYPERLINK(\"http://example.invalid\",\"点我\")"},
        export_context={"version": 7, "week_start": "2026-05-11", "week_end": "2026-05-17"},
    )

    workbook = load_workbook(workbook_bytes, data_only=False)
    try:
        ws = workbook["查询摘要"]
        assert ws["B2"].value == "'=HYPERLINK(\"http://example.invalid\",\"点我\")"
        assert ws["B2"].data_type != "f"
    finally:
        workbook.close()
