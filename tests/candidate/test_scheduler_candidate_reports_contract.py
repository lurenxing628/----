"""回归测试：/reports/overdue|utilization|downtime 及其 export 在带 plan_role 时按所选候选方案(候选行)取数并保持页内导出链接同方案——
default_plan_resolution 缺失角色回退 adopted 标 is_fallback/is_comparison；候选非 completed 或源不符回 400 提示「不是已完成状态」；导出文件名带方案标签、导出过滤器只记 public 摘要不记 candidate_key；页面与导出文本均不泄露 plan_role、candidate_id、candidate_key、scenario_id、source_table 等内部字段。"""

from __future__ import annotations

import importlib
import json
import os
import sys
from io import BytesIO
from pathlib import Path
from typing import Any
from urllib.parse import unquote

import openpyxl

from core.infrastructure.database import ensure_schema, get_connection
from core.services.common.excel_audit import public_export_filters
from core.services.report import ReportEngine
from data.repositories.schedule_plan_query_repo import SOURCE_SCHEDULE
from tests._support.excel_templates import point_env_at_shared
from tests._support.legacy_report_contract import (
    MISSING_ROLE,
    assert_rejected,
    assert_retired,
    get_unchanged,
)
from tests._support.paths import REPO_ROOT
from web.routes.report_plan_preview import default_plan_resolution

SCHEMA_PATH = REPO_ROOT / "schema.sql"
VERSION = 17
EXPORT_INTERNAL_TERMS = (
    "plan_role",
    "candidate_key",
    "candidate_id",
    "source_table",
    "baseline_best",
    "critical_best",
    "scenario_id",
    "candidate_rows",
)
EXPORT_LOG_FORBIDDEN_FILTER_KEYS = (
    "candidate_id",
    "source_table",
    "scenario_id",
    "schedule_id",
    "op_id",
    "node_id",
    "op:",
    "target_id",
    "candidate_key",
    "adopted_candidate_key",
    "baseline_best_candidate_key",
    "critical_best_candidate_key",
    "raw_score_best_candidate_key",
    "selected_candidate_key",
    "graph_w",
)


def _build_app(tmp_path, monkeypatch):
    repo_root = str(REPO_ROOT)
    if repo_root not in sys.path:
        sys.path.insert(0, repo_root)

    test_db = tmp_path / "aps_test.db"
    test_logs = tmp_path / "logs"
    test_backups = tmp_path / "backups"
    test_logs.mkdir(parents=True, exist_ok=True)
    test_backups.mkdir(parents=True, exist_ok=True)

    monkeypatch.setenv("APS_ENV", "development")
    monkeypatch.setenv("APS_DB_PATH", str(test_db))
    monkeypatch.setenv("APS_LOG_DIR", str(test_logs))
    monkeypatch.setenv("APS_BACKUP_DIR", str(test_backups))
    point_env_at_shared(monkeypatch)

    from core.infrastructure.database import ensure_schema, get_connection

    ensure_schema(str(test_db), logger=None, schema_path=str(SCHEMA_PATH), backup_dir=None)
    conn = get_connection(str(test_db))
    try:
        _seed_candidate_report_data(conn)
        conn.commit()
    finally:
        conn.close()

    sys.modules.pop("app", None)
    app_mod = importlib.import_module("app")
    return app_mod.create_app()


def _build_empty_app(tmp_path, monkeypatch):
    repo_root = str(REPO_ROOT)
    if repo_root not in sys.path:
        sys.path.insert(0, repo_root)

    test_db = tmp_path / "aps_empty.db"
    test_logs = tmp_path / "logs"
    test_backups = tmp_path / "backups"
    test_logs.mkdir(parents=True, exist_ok=True)
    test_backups.mkdir(parents=True, exist_ok=True)

    monkeypatch.setenv("APS_ENV", "development")
    monkeypatch.setenv("APS_DB_PATH", str(test_db))
    monkeypatch.setenv("APS_LOG_DIR", str(test_logs))
    monkeypatch.setenv("APS_BACKUP_DIR", str(test_backups))
    point_env_at_shared(monkeypatch)

    ensure_schema(str(test_db), logger=None, schema_path=str(SCHEMA_PATH), backup_dir=None)
    sys.modules.pop("app", None)
    app_mod = importlib.import_module("app")
    return app_mod.create_app()


def _seed_candidate_report_data(conn) -> None:
    conn.executescript(
        """
        INSERT INTO Parts(part_no, part_name)
        VALUES ('P_REPORT_PLAN', '方案报表零件');

        INSERT INTO Machines(machine_id, name, status)
        VALUES
          ('MC_ADOPTED', '最终采用设备', 'active'),
          ('MC_CANDIDATE', '候选方案设备', 'active');

        INSERT INTO Operators(operator_id, name, status)
        VALUES
          ('OP_ADOPTED', '最终采用人员', 'active'),
          ('OP_CANDIDATE', '候选方案人员', 'active');

        INSERT INTO Batches(batch_id, part_no, part_name, quantity, due_date, priority, ready_status, status, remark)
        VALUES ('B_PLAN_REPORT', 'P_REPORT_PLAN', '方案报表零件', 1, '2026-01-01', 'normal', 'yes', 'pending', 'candidate reports');

        INSERT INTO BatchOperations(op_code, batch_id, seq, op_type_name, source, machine_id, operator_id, setup_hours, unit_hours, status)
        VALUES ('B_PLAN_REPORT_10', 'B_PLAN_REPORT', 10, '数铣', 'internal', 'MC_ADOPTED', 'OP_ADOPTED', 0, 1, 'scheduled');

        INSERT INTO ScheduleHistory(version, strategy, batch_count, op_count, result_status, result_summary, created_by)
        VALUES (17, 'test', 1, 1, 'success', '{}', 'pytest');
        """
    )
    op_row = conn.execute("SELECT id FROM BatchOperations WHERE op_code = 'B_PLAN_REPORT_10'").fetchone()
    assert op_row is not None
    op_id = int(op_row["id"])

    conn.execute(
        """
        INSERT INTO Schedule(op_id, machine_id, operator_id, start_time, end_time, lock_status, version)
        VALUES (?, 'MC_ADOPTED', 'OP_ADOPTED', '2026-01-02 08:00:00', '2026-01-02 10:00:00', 'unlocked', 17)
        """,
        (op_id,),
    )
    adopted_candidate_id = _insert_candidate(conn, key="adopted", label="最终采用", kind="baseline", detail_saved="no")
    baseline_candidate_id = _insert_candidate(conn, key="baseline_best", label="原算法最好", kind="baseline", detail_saved="yes")
    conn.execute(
        """
        INSERT INTO ScheduleCandidateRows(version, candidate_id, op_id, machine_id, operator_id, start_time, end_time, lock_status)
        VALUES (17, ?, ?, 'MC_CANDIDATE', 'OP_CANDIDATE', '2026-01-03 08:00:00', '2026-01-03 12:00:00', 'unlocked')
        """,
        (baseline_candidate_id, op_id),
    )
    conn.execute(
        "INSERT INTO ScheduleCandidateSelection(version, role, candidate_id, source_table) VALUES (17, 'adopted', ?, 'schedule')",
        (adopted_candidate_id,),
    )
    conn.execute(
        "INSERT INTO ScheduleCandidateSelection(version, role, candidate_id, source_table) VALUES (17, 'baseline_best', ?, 'candidate_rows')",
        (baseline_candidate_id,),
    )
    conn.executescript(
        """
        INSERT INTO MachineDowntimes(machine_id, start_time, end_time, reason_code, reason_detail, status)
        VALUES ('MC_CANDIDATE', '2026-01-03 09:00:00', '2026-01-03 11:00:00', 'maintenance', 'candidate overlap', 'active');

        INSERT INTO MachineDowntimes(machine_id, start_time, end_time, reason_code, reason_detail, status)
        VALUES ('MC_ADOPTED', '2026-01-02 08:30:00', '2026-01-02 09:30:00', 'maintenance', 'adopted overlap', 'active');
        """
    )


def _insert_candidate(conn, *, key: str, label: str, kind: str, detail_saved: str) -> int:
    cur = conn.execute(
        """
        INSERT INTO ScheduleCandidate(version, candidate_key, candidate_label, candidate_kind, status, graph_enabled, detail_saved)
        VALUES (17, ?, ?, ?, 'completed', 'no', ?)
        """,
        (key, label, kind, detail_saved),
    )
    assert cur.lastrowid is not None
    return int(cur.lastrowid)


def _assert_status(resp, name: str, expect: int = 200) -> None:
    assert resp.status_code == expect, f"{name} 返回 {resp.status_code}，body={resp.get_data(as_text=True)[:500]}"


def _load_xlsx(resp):
    return openpyxl.load_workbook(BytesIO(resp.data), data_only=True)


def _content_disposition(resp) -> str:
    return unquote(resp.headers.get("Content-Disposition", ""))


def _workbook_text(resp) -> str:
    wb = _load_xlsx(resp)
    values = []
    try:
        for ws in wb.worksheets:
            values.append(ws.title)
            for row in ws.iter_rows(values_only=True):
                values.extend(str(cell or "") for cell in row)
    finally:
        wb.close()
    return "\n".join(values)


def _assert_export_public_text(resp) -> None:
    text = _content_disposition(resp) + "\n" + _workbook_text(resp)
    for forbidden in EXPORT_INTERNAL_TERMS:
        assert forbidden not in text


def _latest_report_export_filters(tmp_path, target_type: str):
    conn = get_connection(str(tmp_path / "aps_test.db"))
    try:
        row = conn.execute(
            """
            SELECT detail
            FROM OperationLogs
            WHERE module='reports' AND action='export' AND target_type=?
            ORDER BY id DESC
            LIMIT 1
            """,
            (target_type,),
        ).fetchone()
    finally:
        conn.close()
    assert row is not None
    return (json.loads(row["detail"]) or {}).get("filters") or {}


def _assert_public_export_log_filters(filters: Any) -> None:
    text = json.dumps(filters, ensure_ascii=False, sort_keys=True)
    for forbidden in EXPORT_LOG_FORBIDDEN_FILTER_KEYS:
        assert forbidden not in text


def _assert_number(actual: Any, expected: float) -> None:
    assert round(float(actual), 2) == round(float(expected), 2)


def test_public_export_filters_drop_internal_identity_keys() -> None:
    filters = public_export_filters(
        {
            "version": 17,
            "target_id": 987,
            "candidate_key": "graph_w1_of_5",
            "adopted_candidate_key": "graph_w1_of_5",
            "candidate_id": 99,
            "source_table": "candidate_rows",
            "op_ref": "前缀 op:SECRET 后缀；样本 op_code：OP010；裸编号 OP020",
            "json_like": '{"source_table": "candidate_rows", "candidate_id": 7, "node_id": "op:SECRET"}',
            "standalone": "candidate_rows 查询失败 attempts_debug graph_debug target_id=987 candidate_key=graph_w1_of_5",
            "candidate_rows": {"safe_label": "不应保留"},
            "attempts_debug": "不应保留",
            "nested": {
                "scenario_id": "SC-SECRET",
                "node_id": "op:SECRET-NODE",
                "safe_label": "可见说明",
            },
        }
    )

    assert filters == {
        "version": 17,
        "op_ref": "前缀 内部标识已省略 后缀；样本 内部标识已省略；裸编号 内部标识已省略",
        "json_like": "{内部标识已省略, 内部标识已省略, 内部标识已省略}",
        "standalone": "内部标识已省略 查询失败 内部标识已省略 内部标识已省略 内部标识已省略 内部标识已省略",
        "nested": {"safe_label": "可见说明"},
    }


def test_report_default_plan_resolution_uses_common_fallback_fields() -> None:
    resolution = default_plan_resolution(version=VERSION, raw_role="baseline_best")

    assert resolution["version"] == VERSION
    assert resolution["requested_role"] == "baseline_best"
    assert resolution["selected_role"] == "adopted"
    assert resolution["source_table"] == "schedule"
    assert resolution["candidate_id"] is None
    assert resolution["candidate_key"] is None
    assert resolution["status"] == "fallback_to_adopted"
    assert resolution["is_fallback"] is True
    assert resolution["is_comparison"] is True


def test_candidate_report_pages_receive_plan_role_and_keep_export_links_on_same_plan(tmp_path, monkeypatch) -> None:
    app = _build_app(tmp_path, monkeypatch)
    client = app.test_client()

    for topic in ("overdue", "utilization", "downtime"):
        query = "version=17&plan_role=baseline_best"
        if topic != "overdue":
            query += "&start_date=2026-01-03&end_date=2026-01-03"
        download = "/reports/" + topic + "/export?" + query
        response = get_unchanged(client, "/reports/" + topic + "?" + query, tmp_path / "aps_test.db")
        assert_retired(response, public=("17", "原算法代表方案"), downloads=(download,))
        exported = get_unchanged(client, download, tmp_path / "aps_test.db")
        _assert_status(exported, topic + " original download")
        _assert_export_public_text(exported)
        workbook = _load_xlsx(exported)
        try:
            if topic == "overdue":
                assert workbook["超期清单"]["G2"].value == "2026-01-03 12:00:00"
            elif topic == "utilization":
                sheet = workbook["设备负荷"]
                assert sheet["A2"].value == "MC_CANDIDATE"
                assert sheet["B2"].value == "候选方案设备"
                _assert_number(sheet["C2"].value, 4.0)
                assert "MC_ADOPTED" not in _workbook_text(exported)
            else:
                assert workbook["停机影响"]["A2"].value == "MC_CANDIDATE"
                _assert_number(workbook["停机影响"]["E2"].value, 2.0)
        finally:
            workbook.close()


def test_candidate_report_exports_use_selected_plan_rows_and_filename_label(tmp_path, monkeypatch) -> None:
    app = _build_app(tmp_path, monkeypatch)
    client = app.test_client()

    overdue_resp = client.get("/reports/overdue/export?version=17&plan_role=baseline_best")
    _assert_status(overdue_resp, "candidate overdue export")
    assert "原算法代表方案" in _content_disposition(overdue_resp)
    overdue_filters = _latest_report_export_filters(tmp_path, "overdue")
    assert overdue_filters.get("requested_plan_role") == "baseline_best"
    assert overdue_filters.get("effective_plan_role") == "baseline_best"
    assert overdue_filters.get("plan_role_status") == "resolved_comparison"
    assert "candidate_key" not in overdue_filters
    _assert_public_export_log_filters(overdue_filters)
    wb = _load_xlsx(overdue_resp)
    try:
        ws = wb["超期清单"]
        assert ws["G2"].value == "2026-01-03 12:00:00"
    finally:
        wb.close()
    _assert_export_public_text(overdue_resp)

    utilization_resp = client.get("/reports/utilization/export?version=17&plan_role=baseline_best&start_date=2026-01-03&end_date=2026-01-03")
    _assert_status(utilization_resp, "candidate utilization export")
    assert "原算法代表方案" in _content_disposition(utilization_resp)
    utilization_filters = _latest_report_export_filters(tmp_path, "utilization")
    assert utilization_filters.get("requested_plan_role") == "baseline_best"
    assert utilization_filters.get("effective_plan_role") == "baseline_best"
    assert utilization_filters.get("plan_role_status") == "resolved_comparison"
    assert "candidate_key" not in utilization_filters
    _assert_public_export_log_filters(utilization_filters)
    wb = _load_xlsx(utilization_resp)
    try:
        ws = wb["设备负荷"]
        assert ws["A2"].value == "MC_CANDIDATE"
        _assert_number(ws["C2"].value, 4.0)
    finally:
        wb.close()
    _assert_export_public_text(utilization_resp)

    downtime_resp = client.get("/reports/downtime/export?version=17&plan_role=baseline_best&start_date=2026-01-03&end_date=2026-01-03")
    _assert_status(downtime_resp, "candidate downtime export")
    assert "原算法代表方案" in _content_disposition(downtime_resp)
    downtime_filters = _latest_report_export_filters(tmp_path, "downtime")
    assert downtime_filters.get("requested_plan_role") == "baseline_best"
    assert downtime_filters.get("effective_plan_role") == "baseline_best"
    assert downtime_filters.get("plan_role_status") == "resolved_comparison"
    assert "candidate_key" not in downtime_filters
    _assert_public_export_log_filters(downtime_filters)
    wb = _load_xlsx(downtime_resp)
    try:
        ws = wb["停机影响"]
        assert ws["A2"].value == "MC_CANDIDATE"
        _assert_number(ws["E2"].value, 2.0)
    finally:
        wb.close()
    _assert_export_public_text(downtime_resp)


def test_candidate_report_pages_reject_non_completed_candidate_rows(tmp_path, monkeypatch) -> None:
    app = _build_app(tmp_path, monkeypatch)
    conn = get_connection(str(tmp_path / "aps_test.db"))
    try:
        conn.execute(
            "UPDATE ScheduleCandidate SET status = 'failed' WHERE version = ? AND candidate_key = ?",
            (VERSION, "baseline_best"),
        )
        conn.commit()
    finally:
        conn.close()

    client = app.test_client()
    urls = (
        "/reports/overdue?version=17&plan_role=baseline_best",
        "/reports/utilization?version=17&plan_role=baseline_best&start_date=2026-01-03&end_date=2026-01-03",
        "/reports/downtime?version=17&plan_role=baseline_best&start_date=2026-01-03&end_date=2026-01-03",
        "/reports/overdue/export?version=17&plan_role=baseline_best",
        "/reports/utilization/export?version=17&plan_role=baseline_best&start_date=2026-01-03&end_date=2026-01-03",
        "/reports/downtime/export?version=17&plan_role=baseline_best&start_date=2026-01-03&end_date=2026-01-03",
    )

    for url in urls:
        resp = get_unchanged(client, url, tmp_path / "aps_test.db")
        html = resp.get_data(as_text=True)
        assert resp.status_code == 400, url
        if "/export?" in url:
            assert "这套对比参考方案当前不是已完成状态，不能查看明细。" in html
        else:
            assert_rejected(resp)
        assert "candidate_rows" not in html
        assert "MC_CANDIDATE" not in html
        assert "候选方案设备" not in html


def test_candidate_report_pages_reject_non_completed_schedule_source_comparison(tmp_path, monkeypatch) -> None:
    app = _build_app(tmp_path, monkeypatch)
    conn = get_connection(str(tmp_path / "aps_test.db"))
    try:
        conn.execute(
            "UPDATE ScheduleCandidateSelection SET source_table = ? WHERE version = ? AND role = ?",
            (SOURCE_SCHEDULE, VERSION, "baseline_best"),
        )
        conn.execute(
            "UPDATE ScheduleCandidate SET status = 'skipped' WHERE version = ? AND candidate_key = ?",
            (VERSION, "baseline_best"),
        )
        conn.commit()
    finally:
        conn.close()

    client = app.test_client()
    for url in (
        "/reports/utilization?version=17&plan_role=baseline_best&start_date=2026-01-02&end_date=2026-01-02",
        "/reports/downtime/export?version=17&plan_role=baseline_best&start_date=2026-01-02&end_date=2026-01-02",
    ):
        resp = get_unchanged(client, url, tmp_path / "aps_test.db")
        html = resp.get_data(as_text=True)
        assert resp.status_code == 400, url
        if "/export?" in url:
            assert "这套对比参考方案当前不是已完成状态，不能查看明细。" in html
        else:
            assert_rejected(resp)
        assert "source_table" not in html
        assert "candidate_id" not in html
        assert "MC_ADOPTED" not in html


def test_candidate_report_missing_role_falls_back_to_adopted_with_visible_status(tmp_path, monkeypatch) -> None:
    app = _build_app(tmp_path, monkeypatch)
    client = app.test_client()

    for topic in ("overdue", "utilization", "downtime"):
        query = "version=17&plan_role=critical_best"
        if topic != "overdue":
            query += "&start_date=2026-01-02&end_date=2026-01-02"
        response = get_unchanged(client, "/reports/" + topic + "?" + query, tmp_path / "aps_test.db")
        page = assert_retired(response, public=("17",), message=MISSING_ROLE)
        assert "2026-01-02 10:00:00" not in page.text
        assert "2026-01-03 12:00:00" not in page.text
        assert "已显示正式采用方案" not in page.text

    export_resp = client.get("/reports/overdue/export?version=17&plan_role=critical_best")
    _assert_status(export_resp, "candidate fallback export")
    assert "正式采用方案" in _content_disposition(export_resp)
    wb = _load_xlsx(export_resp)
    try:
        ws = wb["超期清单"]
        assert ws["G2"].value == "2026-01-02 10:00:00"
    finally:
        wb.close()

    utilization_export = client.get("/reports/utilization/export?version=17&plan_role=critical_best&start_date=2026-01-02&end_date=2026-01-02")
    _assert_status(utilization_export, "candidate fallback utilization export")
    assert "正式采用方案" in _content_disposition(utilization_export)
    _assert_export_public_text(utilization_export)

    downtime_export = client.get("/reports/downtime/export?version=17&plan_role=critical_best&start_date=2026-01-02&end_date=2026-01-02")
    _assert_status(downtime_export, "candidate fallback downtime export")
    assert "正式采用方案" in _content_disposition(downtime_export)
    _assert_export_public_text(downtime_export)

    conn = get_connection(str(tmp_path / "aps_test.db"))
    try:
        diagnosis = ReportEngine(conn).overdue_delay_diagnosis_context(VERSION, plan_role="critical_best")
        links = [
            str(action.get("link") or "")
            for item in diagnosis["items_by_batch"].values()
            for action in item.get("suggested_actions") or []
        ]
        assert any("plan_role=critical_best" in link for link in links)
        assert all("plan_role=adopted" not in link for link in links if "/scheduler/gantt" in link)
    finally:
        conn.close()


def test_report_exports_reject_unknown_scenario_without_internal_field_name(tmp_path, monkeypatch) -> None:
    app = _build_app(tmp_path, monkeypatch)
    client = app.test_client()

    for url in (
        "/reports/overdue/export?version=17&scenario_id=scenario-1",
        "/reports/utilization/export?version=17&scenario_id=scenario-1&start_date=2026-01-02&end_date=2026-01-02",
        "/reports/downtime/export?version=17&scenario_id=scenario-1&start_date=2026-01-02&end_date=2026-01-02",
    ):
        resp = client.get(url)
        html = resp.get_data(as_text=True)
        assert resp.status_code == 400
        assert "模拟方案不存在" in html
        assert "scenario_id" not in html


def test_report_pages_reject_unknown_plan_role_without_history(tmp_path, monkeypatch) -> None:
    app = _build_empty_app(tmp_path, monkeypatch)
    client = app.test_client()

    resp = get_unchanged(client, "/reports/overdue?plan_role=bad_role", tmp_path / "aps_empty.db")
    html = resp.get_data(as_text=True)

    assert resp.status_code == 404
    assert "bad_role" not in html
    assert 'value="0"' not in html
    assert "Location" not in resp.headers
    exported = client.get("/reports/overdue/export?plan_role=bad_role")
    assert exported.status_code == 404
    assert "暂无排产历史" in exported.get_data(as_text=True)
